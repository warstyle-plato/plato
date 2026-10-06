"""Полная сверка с калькулятором ГлавАПУ: все листы новой выгрузки.

Выгрузка владельца от 06.10.2026 (КРТ Нагатино, 17,811 га, наш сценарий
в калькуляторе) принесла то, чего сверка ещё не читала:

* лист «Социальные объекты» — типовые здания, которые калькулятор поставил,
  и разделы «ДОО / Школы / Поликлиники» листа «ТЭП» (места, СПП, НП, участок);
* «Расчёт объектов обслуживания» (строки 30–41) — нормы на население;
* компенсацию по объектам со знаком (минус — места сверх потребности);
* лист «Машино-места» — ВСЕ виды мест по каждому ВРИ, не только приобъектные;
* лист «Параметры территории» — К1/К2, зона, нормативы, квартал, аренда,
  УПКС и базовые по типам;
* квартиры по размерам, баланс и элементы жилых территорий.

И две странности, которые сверка обязана ПОКАЗАТЬ: льгота за МПТ при
«Коэффициент МПТ: не включён» и школа на 2 500 мест при потребности 351.

Книга собирается здесь же из чисел выгрузки владельца (`_owner_book`): самой
книги в репозитории нет, и числа, которых владелец не назвал, в неё не
выдуманы — листа «МПТ» нет вовсе, и сверка обязана сказать «нет», а не ноль.

Запуск: python3 -m pytest tests/test_glavapu_full_check.py -q
"""

from __future__ import annotations

import copy
import json
import os
import sys
import time
from pathlib import Path

import pytest

ROOT = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(ROOT))
for _key in ("AUCTION_KRT_WEEKLY", "AUCTION_KRT_WATCH", "DEVELOPAID_WORKBOOK_CACHE",
             "NORMATIVES_WATCH", "NAGATINO_EGRN_READ"):
    os.environ.setdefault(_key, "0")

import glavapu_scenario as gs  # noqa: E402
import project_preset  # noqa: E402

NUMBERS = ["77:05:0004001:1"]

TEP_ROWS = [
    ("№", "Наименования", "Единицы измерения", "Показатель"),
    (1, "Площадь территории проектирования", "га", "17,811"),
    (4, "Население", "чел.", 3897),
    (5, "Количество квартир", "шт.", 1856),
    (None, "Квартиры по площади:", None, None),
    ("5.1", "квартиры до 70 кв.м.", "шт.", 1856),
    ("5.2", "квартиры от 70 до 100 кв.м.", "шт.", 0),
    ("5.3", "квартиры более 100 кв.м.", "шт.", 0),
    (None, "Баланс территории:", None, None),
    (12, "Территория жилых зданий, в т.ч.:", "га", "9,11 (51,1%)"),
    ("12.1", "участки многоквартирных жилых зданий", "га", "7,0 (39,3%)"),
    ("12.2", "незастраиваемая территория", "га", "2,11 (11,8%)"),
    (13, "Участки социальных объектов", "га", "4,3 (24,1%)"),
    (14, "Участки общественных, производственных объектов", "га", "4,401 (24,7%)"),
    (None, "ДОО:", None, None),
    (18, "количество мест", "мест", 250),
    (19, "СПП", "тыс.кв.м.", "5,000"),
    (20, "наземная площадь", "тыс.кв.м.", "4,500"),
    (21, "площадь земельного участка", "га", "0,8000"),
    (None, "Школы:", None, None),
    (22, "количество мест", "мест", 2500),
    (23, "СПП", "тыс.кв.м.", "36,111"),
    (24, "наземная площадь", "тыс.кв.м.", "32,500"),
    (25, "площадь земельного участка", "га", "3,5000"),
    (None, "Поликлиники:", None, None),
    (26, "мощность", "пос./см.", 0),
    (27, "СПП", "тыс.кв.м.", "0,000"),
    (28, "наземная площадь", "тыс.кв.м.", "0,000"),
    (29, "площадь земельного участка", "га", "0,0000"),
    (None, "Расчёт объектов обслуживания:", None, None),
    (30, "ДОО", "мест", 172),
    (31, "Школа", "мест", 351),
    (32, "Поликлиника смешанного типа", "пос./см.", 75),
    (33, "Поликлиника взрослая", "пос./см.", 52),
    (34, "Поликлиника детская", "пос./см.", 23),
    (35, "Плоскостные спортивные сооружения", "га", "0,3781"),
    (36, "Крытые объекты спорта (ННП), в т.ч.:", "тыс.кв.м.", "3,118"),
    ("36.1", "в радиусе пешеходной доступности до 500 м", "тыс.кв.м.", "1,248"),
    ("36.2", "в радиусе пешеходной доступности до 1500 м", "тыс.кв.м.", "1,871"),
    (37, "Объекты торговли (ННП)", "тыс.кв.м.", "1,053"),
    (38, "Объекты бытового обслуживания населения (ННП)", "тыс.кв.м.", "0,390"),
    (39, "Объекты общественного питания (ННП)", "тыс.кв.м.", "0,468"),
    (40, "Объекты культуры и досуга (ННП)", "тыс.кв.м.", "0,585"),
    (41, "Объекты для размещения городских служб (ННП)", "тыс.кв.м.", "0,351"),
    (None, "Расчёт машино-мест:", None, None),
    ("42", "Места хранения и паркирования, в т.ч.:", "м/м", 2574),
    ("42.1", "Постоянные парковки", "м/м", 1485),
    ("42.2", "Гостевые парковки", "м/м", 149),
    ("42.3", "Приобъектные парковки", "м/м", 940),
    ("43", "Места кратковременной остановки", "м/м", 28),
    (None, "Расчёт стоимости смены ВРИ:", "млн.руб.", "0,000"),
    (44, "Многоквартирная жилые здания", "млн.руб.", "9\xa0342,996"),
    (45, "Индивидуальные, блочные жилые здания", "млн.руб.", "0,000"),
    (52, "Льгота на стр-во жилья за создание МПТ", "млн.руб.", "23\xa0253,958"),
    (53, "Льгота на стр-во жилья за передачу жилых помещений в собственность города Москвы",
     "млн.руб.", "0,000"),
    (None, "Расчёт компенсации за социальные объекты:", "млн.руб.", "0,000"),
    (54, "ДОО", "млн.руб.", "-603,123"),
    (55, "Школа", "млн.руб.", "-12\xa0360,506"),
    (56, "Поликлиника", "млн.руб.", "780,121"),
    (None, "Элементы жилых территорий:", None, None),
    (57, "Озелененные территории ЖК, в т.ч.:", "га", "1,9485"),
    ("57.1", "зелёные насаждения", "га", "1,3640"),
    ("57.2", "площадки и ДТС", "га", "0,5846"),
    (58, "Озелененные территории общего пользования", "га", "0,2728"),
]

PARKING_HEADER = ("№", "Наименования", "Единицы измерения", "Всего", "Приобъектные",
                  "Постоянные", "Гостевые", "Кратковременные")
PARKING_ROWS = [
    (1, "Образование и просвещение (3.5)", "машино-места", 53, 35, 0, 0, 18),
    (2, "Деловое управление (4.1)", "машино-места", 857, 857, 0, 0, 0),
    (3, "Многоквартирный дом (2.1.1, 2.5, 2.6)", "машино-места", 1640, 0, 1485, 149, 6),
    (4, "Встроенно-пристроенные помещения многоквартирного дома (2.1.1, 2.5, 2.6)",
     "машино-места", 52, 48, 0, 0, 4),
]

SOCIAL_HEADER = ("№", "Наименования", "Площадь участка (га)", "Наземная площадь (тыс.кв.м.)",
                 "СПП в ГНС (тыс.кв.м.)", "Всего (мест | пос./см.)", "ДОО (мест)", "Школа (мест)")
SOCIAL_ROWS = [
    (1, "Дошкольное здание на 250 мест", 0.8, 4.5, 5.0, 250, 250, 0),
    (2, "Школьное здание на 2500 мест со спортивным ядром", 3.5, 32.5, 36.111, 2500, 0, 2500),
]

PARAMS_ROWS = [
    ("Параметр", "Значение", "Ед.изм."),
    ("Машино-места", None, None),
    ("К1 — доступность рельсового каркаса", "0,75", "—"),
    ("К2 — деловая активность (вне ТТК)", "0,5", "—"),
    ("Социальные объекты", None, None),
    ("Район", "Нагатино-Садовники", "—"),
    ("Расчётная зона", 1, "—"),
    ("Норматив ДОО", 44, "мест / 1000 жит."),
    ("Норматив школ", 90, "мест / 1000 жит."),
    ("Стоимость смены ВРИ", None, None),
    ("Кадастровый квартал", "77:05:0004001", "—"),
    ("Коэффициент аренды", "0,1184", "—"),
    ("Коэффициент МПТ", "не включён", "—"),
    ("УПКС и базовые стоимости по типам использования", None, None),
    ("Тип использования", "УПКС, руб/м²", "Базовая, тыс.руб/м²"),
    ("МКД (многоэтажный жилой дом)", "83\xa0789,85", "194\xa0324,67"),
    ("Гаражи", "21\xa0054,51", None),
    ("Торговля и многофункц.", "52\xa0008,65", None),
    ("Офисы", "44\xa0965,57", None),
    ("Социальные объекты", "17\xa0100,86", None),
]


def _owner_book(tmp_path, *, parking_extra=(), social_header=SOCIAL_HEADER,
                social_rows=SOCIAL_ROWS, params=PARAMS_ROWS, tep=TEP_ROWS) -> bytes:
    import openpyxl
    book = openpyxl.Workbook()
    sheet = book.active
    sheet.title = "ТЭП"
    for row in tep:
        sheet.append(list(row))
    parking = book.create_sheet("Машино-места")
    parking.append(list(PARKING_HEADER))
    for row in list(PARKING_ROWS) + list(parking_extra):
        parking.append(list(row))
    social = book.create_sheet("Социальные объекты")
    social.append(list(social_header))
    for row in social_rows:
        social.append(list(row))
    territory = book.create_sheet("Параметры территории")
    for row in params:
        territory.append(list(row))
    path = tmp_path / "owner.xlsx"
    book.save(path)
    return path.read_bytes()


def _side(data: bytes) -> dict:
    import main_legacy as legacy
    parsed = legacy.parse_glavapu_xlsx(data, "owner.xlsx")
    return gs.glavapu_side(parsed["normalized"], {}, data)


# ------------------------------------------------------- сторона ГлавАПУ --

def test_social_buildings_are_read_by_header(tmp_path) -> None:
    items = gs.social_buildings(_owner_book(tmp_path))
    school = next(i for i in items if i["kind"] == "school")
    assert school["places"] == 2500 and school["site_ha"] == 3.5 and school["spp_ths"] == 36.111
    assert school["np_ths"] == 32.5
    doo = next(i for i in items if i["kind"] == "kindergarten")
    assert doo["places"] == 250 and doo["site_ha"] == 0.8


def test_social_columns_are_found_by_header_not_position(tmp_path) -> None:
    """Контрпример: столбцы переставлены — участок не читается как площадь."""
    header = ("№", "Наименования", "Школа (мест)", "ДОО (мест)", "СПП в ГНС (тыс.кв.м.)",
              "Площадь участка (га)", "Всего (мест | пос./см.)", "Наземная площадь (тыс.кв.м.)")
    rows = [(1, "Школьное здание на 2500 мест", 2500, 0, 36.111, 3.5, 2500, 32.5)]
    items = gs.social_buildings(_owner_book(tmp_path, social_header=header, social_rows=rows))
    assert items == [{"name": "Школьное здание на 2500 мест", "school": 2500, "kindergarten": 0,
                      "spp_ths": 36.111, "site_ha": 3.5, "places": 2500, "np_ths": 32.5,
                      "kind": "school"}]


def test_every_parking_kind_is_read_per_vri(tmp_path) -> None:
    side = _side(_owner_book(tmp_path))
    pv = {k: v[0] for k, v in side["parking_vri"].items()}
    assert pv["3_5.total"] == 53 and pv["3_5.attached"] == 35 and pv["3_5.short_stop"] == 18
    assert pv["4_1.total"] == 857 and pv["4_1.attached"] == 857 and pv["4_1.permanent"] == 0
    assert (pv["2_1_1.total"], pv["2_1_1.permanent"], pv["2_1_1.guest"],
            pv["2_1_1.short_stop"]) == (1640, 1485, 149, 6)
    assert pv["built_in.attached"] == 48 and pv["built_in.short_stop"] == 4
    # Сумма по ВРИ сходится с итогом листа «ТЭП» (строки 42.3 и 43).
    assert sum(v for k, v in pv.items() if k.endswith(".attached")) == 940
    assert sum(v for k, v in pv.items() if k.endswith(".short_stop")) == 28


def test_service_social_objects_and_territory_are_read_by_name(tmp_path) -> None:
    side = _side(_owner_book(tmp_path))
    obj = {k: v[0] for k, v in side["social_obj"].items()}
    assert obj["school.places"] == 2500 and obj["school.site"] == 3.5
    assert obj["kindergarten.np"] == 4.5 and obj["clinic.places"] == 0
    service = {k: v[0] for k, v in side["service"].items()}
    assert service["clinic_adult"] == 52 and service["clinic_child"] == 23
    assert service["sport_flat"] == 0.3781 and service["sport_indoor_1500"] == 1.871
    assert service["commerce_need"] == pytest.approx(1.053 + 0.390 + 0.468 + 0.585 + 0.351)
    # «до 500 м» крытого спорта — не квартиры «до 70».
    assert {k: v[0] for k, v in side["flats"].items()} == {
        "total": 1856, "small": 1856, "medium": 0, "large": 0}
    assert side["territory"]["green_zhk"][0] == 1.9485
    assert side["territory"]["58"][0] == 0.2728
    params = {k: v[0] for k, v in side["params"].items()}
    assert params["k1"] == 0.75 and params["k2"] == 0.5 and params["rent"] == 0.1184
    assert params["kindergarten_norm"] == 44 and params["school_norm"] == 90
    assert params["quarter"] == "77:05:0004001" and params["mpt_coef"] == "не включён"
    assert params["upks_mkd"] == pytest.approx(83789.85) and params["base_mkd"] == pytest.approx(194324.67)
    assert params["upks_office"] == pytest.approx(44965.57)
    assert "вне ТТК" in side["params"]["k2"][1]


def test_a_missing_sheet_is_named_not_zero(tmp_path) -> None:
    side = _side(_owner_book(tmp_path))
    # Листа «МПТ» в книге нет: величины нет, и это не ноль.
    assert side["mpt"][0] is None


# ------------------------------------------------------------- странности --

def test_the_owner_book_strangeness_is_shown(tmp_path) -> None:
    found = {a["kind"]: a for a in _side(_owner_book(tmp_path))["anomalies"]}
    school = found["social_surplus.school"]["text"]
    assert "2 500" in school and "351" in school and "сверх потребности 2 149" in school
    assert "−" not in school and "-12 360,506" in school
    assert "Школьное здание на 2500 мест" in school
    assert "social_surplus.kindergarten" in found
    # Итог компенсации 0 при дефиците поликлиники: профицит школы его гасит.
    offset = found["social_comp_offset"]["text"]
    assert "Поликлиника +780,121" in offset and "max(0" in offset
    # Льгота за МПТ при «не включён» — с объяснением, почему она не ноль.
    flag = found["mpt_relief_flag"]["text"]
    assert "23 253,958" in flag and "не включён" in flag and "соцобъекты (0,3)" in flag
    over = found["relief_over_fee"]["text"]
    assert "23 253,958" in over and "9 342,996" in over
    assert all(a["where"] for a in found.values())


def test_no_strangeness_where_there_is_none(tmp_path) -> None:
    """Контрпример: школа по потребности, признак МПТ включён, льгота меньше
    платы — странностей нет. Проверка, которая молчит всегда, ничего не доказывает."""
    tep = [row for row in TEP_ROWS]
    swap = {22: (22, "количество мест", "мест", 300), 18: (18, "количество мест", "мест", 150),
            52: (52, "Льгота на стр-во жилья за создание МПТ", "млн.руб.", "1\xa0000,000"),
            54: (54, "ДОО", "млн.руб.", "100,000"), 55: (55, "Школа", "млн.руб.", "50,000")}
    tep = [swap.get(row[0], row) for row in tep]
    tep = [(None, row[1], row[2], "930,121") if row[1].startswith("Расчёт компенсации") else row
           for row in tep]
    params = [("Коэффициент МПТ", "включён", "—") if r[0] == "Коэффициент МПТ" else r
              for r in PARAMS_ROWS]
    side = _side(_owner_book(tmp_path, tep=tep, params=params))
    assert side["anomalies"] == []


# ------------------------------------------------------------------ маршрут --

@pytest.fixture()
def core(monkeypatch, tmp_path):
    import main as wrapper
    core = wrapper.core
    monkeypatch.setenv("DEVELOPAID_DATA_DIR", str(tmp_path))
    monkeypatch.setattr(core, "_core_api_url", lambda path: "")
    monkeypatch.setattr(core, "_glavapu_headless_available", lambda: True)
    return core


def _nagatino() -> tuple[dict, dict]:
    preview = project_preset.build_preview(
        json.loads((ROOT / "presets" / "КРТ_Нагатино.json").read_text(encoding="utf-8")))
    inputs = dict(preview["inputs"], site_area_ha=17.811, land_right="lease",
                  landscaping_area_per_person_sqm=11)
    inputs["_cadastral_analysis"] = {"recognized": NUMBERS}
    return inputs, preview["tep"]


def _run_with(core, monkeypatch, data: bytes, inputs: dict, tep: dict) -> dict:
    def run(numbers, area_ha, params=None):
        return {"rows": [], "scenario_report": {"applied": [], "refused": []},
                "scenario_rows": {"1": area_ha}, "scenario_xlsx": data, "timings": {}}
    monkeypatch.setattr(core, "_glavapu_headless_run", run)
    req = core.GlavapuScenarioRequest(inputs=inputs, tep=tep)
    core.glavapu_scenario_check(req)
    deadline = time.monotonic() + 10
    while time.monotonic() < deadline:
        answer = core.glavapu_scenario_check(req.model_copy(update={"poll": True}))
        if answer["state"] not in ("queued", "running"):
            return answer
        time.sleep(0.05)
    raise AssertionError("задание не кончилось")


def test_every_parking_kind_is_compared_by_vri(core, monkeypatch, tmp_path) -> None:
    inputs, tep = _nagatino()
    answer = _run_with(core, monkeypatch, _owner_book(tmp_path), inputs, tep)
    rows = {r["kind"]: r for r in answer["comparison"]["rows"]}
    # МКД: постоянные и гостевые — наша норма; места остановки — нет нашей.
    assert rows["parking_vri.2_1_1.permanent"]["glavapu"] == 1485
    assert "tep_derived_norms" in rows["parking_vri.2_1_1.permanent"]["ours_origin"]
    assert rows["parking_vri.2_1_1.guest"]["glavapu"] == 149
    short = rows["parking_vri.2_1_1.short_stop"]
    assert short["status"] == "ours_missing" and "остановки" in short["reason"]
    assert rows["parking_vri.2_1_1.total"]["status"] == "ours_missing"
    # Офисы: приобъектные — наша норма; нули калькулятора по другим видам — не строки.
    office = rows["parking_vri.4_1.attached"]
    assert office["glavapu"] == 857 and "parking_demand" in office["ours_origin"]
    assert "parking_vri.4_1.permanent" not in rows and "parking_vri.4_1.short_stop" not in rows
    # Образование: у соцобъектов своих мест в модели нет — так и сказано.
    assert rows["parking_vri.3_5.attached"]["status"] == "ours_missing"
    assert "соцобъект" in rows["parking_vri.3_5.short_stop"]["reason"]
    assert rows["parking_vri.built_in.attached"]["glavapu"] == 48
    # Порядок: по ВРИ, внутри — всего, приобъектные, постоянные, гостевые, остановка.
    mkd = [k for k in rows if k.startswith("parking_vri.2_1_1.")]
    assert mkd == ["parking_vri.2_1_1.total", "parking_vri.2_1_1.permanent",
                   "parking_vri.2_1_1.guest", "parking_vri.2_1_1.short_stop"]
    note = answer["comparison"]["notes"]["Машино-места по ВРИ"]
    assert "а не «уличные»" in note and "ПОТРЕБНОСТЬ, а не размещение" in note


def test_a_shopping_centre_row_is_compared_too(core, monkeypatch, tmp_path) -> None:
    """В выгрузке владельца ТЦ нет — строка подставлена: ТЦ нашего ТЭП уходит
    в калькулятор как 4.2 и сверяется по всем видам, а «Магазины (4.4)» —
    названы, а не приписаны ТЦ."""
    inputs, tep = _nagatino()
    extra = [(5, "Объекты торговли (4.2)", "машино-места", 120, 100, 0, 0, 20),
             (6, "Магазины (4.4)", "машино-места", 10, 10, 0, 0, 0)]
    answer = _run_with(core, monkeypatch, _owner_book(tmp_path, parking_extra=extra), inputs, tep)
    rows = {r["kind"]: r for r in answer["comparison"]["rows"]}
    retail = rows["parking_vri.4_2.attached"]
    assert retail["glavapu"] == 100 and retail["ours"] is not None
    assert "parking_demand" in retail["ours_origin"] and retail["status"] in ("match", "diff")
    assert rows["parking_vri.4_2.short_stop"]["glavapu"] == 20
    shops = rows["parking_vri.4_4.attached"]
    assert shops["status"] == "ours_missing" and "4.2" in shops["reason"]


def test_params_service_and_territory_are_compared(core, monkeypatch, tmp_path) -> None:
    inputs, tep = _nagatino()
    answer = _run_with(core, monkeypatch, _owner_book(tmp_path), inputs, tep)
    rows = {r["kind"]: r for r in answer["comparison"]["rows"]}
    # Нормативы соцобъектов — наш единый источник против листа калькулятора.
    assert rows["params.kindergarten_norm"]["status"] == "match"
    assert rows["params.school_norm"]["status"] == "match"
    assert "SOCIAL_NORMS_PER_1000" in rows["params.school_norm"]["ours_origin"]
    # Квартал — из кадастровых номеров сценария, текстом.
    assert rows["params.quarter"]["status"] == "match" and rows["params.quarter"]["delta"] is None
    assert rows["params.mpt_coef"]["status"] == "reference"
    # Встроенная коммерция покрывает ННП обслуживания; запас, а не равенство.
    commerce = rows["service.commerce_need"]
    assert commerce["status"] == "covered" and commerce["ours"] > commerce["glavapu"]
    assert rows["service.retail"]["status"] == "ours_missing"
    # Наш двор против озеленённых территорий ЖК.
    green = rows["territory.green_zhk"]
    assert green["glavapu"] == 1.9485 and green["status"] in ("covered", "short")
    assert "2152-ПП" in green["reason"] or green["status"] == "covered"
    assert rows["territory.58"]["status"] == "reference"
    # Школа калькулятора против нашей — расходится, с причиной.
    school = rows["social_obj.school.places"]
    assert school["glavapu"] == 2500 and school["ours"] == 1000 and school["status"] == "diff"
    assert "типовое здание" in school["reason"]
    assert rows["social_obj.school.site"]["status"] == "reference"
    assert rows["flats.small"]["glavapu"] == 1856 and rows["flats.small"]["status"] == "ours_missing"
    assert "средняя квартира" in rows["flats.small"]["reason"]
    kinds = {a["kind"] for a in answer["comparison"]["anomalies"]}
    assert {"social_surplus.school", "mpt_relief_flag", "relief_over_fee"} <= kinds


def test_a_short_yard_is_named_short(core, monkeypatch, tmp_path) -> None:
    """Контрпример к «покрывает»: двор 1 м² на жителя меньше 5,0 по городу."""
    inputs, tep = _nagatino()
    inputs = dict(inputs, landscaping_area_per_person_sqm=1)
    answer = _run_with(core, monkeypatch, _owner_book(tmp_path), inputs, tep)
    green = {r["kind"]: r for r in answer["comparison"]["rows"]}["territory.green_zhk"]
    assert green["status"] == "short" and green["reason"].startswith("не хватает")


def test_the_page_shows_anomalies_notes_and_new_statuses(core, monkeypatch, tmp_path) -> None:
    """Настоящий код PAGE рисует странности, пояснения групп и типовые здания."""
    import page_blocks

    inputs, tep = _nagatino()
    answer = _run_with(core, monkeypatch, _owner_book(tmp_path), inputs, copy.deepcopy(tep))
    prelude = "const answer=%s;" % json.dumps(answer, ensure_ascii=False, default=str)
    tail = ("console.log(JSON.stringify({html:glavapuScenarioHtml(answer),"
            "summary:glavapuScenarioSummaryHtml(answer)}));")
    out, _ = page_blocks.run(prelude, tail)
    got = json.loads(out)
    html = got["html"]
    anomalies = html.find("Странности выгрузки калькулятора")
    assert 0 <= anomalies < html.find("glavapu-scenario-table")
    assert "сверх потребности 2 149" in html and "23 253,958" in html
    assert "а не «уличные»" in html
    assert 'data-status="covered"' in html and "покрывает" in html
    assert "Школьное здание на 2500 мест со спортивным ядром" in html
    order = [html.find(t) for t in ("Параметры территории", "СПП и ГНС по видам",
                                    "Объекты обслуживания", "Машино-места по ВРИ",
                                    "Элементы жилых территорий")]
    assert all(i >= 0 for i in order) and order == sorted(order), order
    assert "Найдено странностей калькулятора: 5" in got["summary"]
