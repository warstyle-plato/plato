"""Монитор принимает финмодель (лист «cf_1_new») вместо финкниги с листом РСС.

Владелец (23.10.2026): «бюджет достройки берём из вкладки cf_1_new», «загрузить
финкнигу я не могу — он не распознаёт лист РСС, а его там и нет». Утверждённый
бюджет — «ИТОГО», оплачено — «ФАКТ», средства на завершение — «ПЛАН»,
помесячная потребность — колонки плана; главы статей — из карты «статьи_БДДС»
той же книги (решение владельца: разложить по главам 1/2/3 через карту).
"""

from __future__ import annotations

import datetime
import io
import os
import sys
from pathlib import Path

import openpyxl
import pytest

sys.path.insert(0, os.path.dirname(os.path.dirname(os.path.abspath(__file__))))

import developaid_monitor as monitor  # noqa: E402
import developaid_monitor_dashboard as dashboard  # noqa: E402

K = 1000.0
MONTHS = (datetime.datetime(2026, 10, 1), datetime.datetime(2026, 11, 1), datetime.datetime(2027, 4, 1))
# код: (название, ИТОГО, ФАКТ, ПЛАН, план по месяцам) — в тысячах, как лист
ROWS = {
    "1.2": ("Поступления", 900, 100, 800, (0, 0, 0)),
    "2.1": ("Стоимость приобретения проекта", 100, 60, 40, (10, 30, 0)),
    "2.1.1": ("Стоимость ЗУ", 100, 60, 40, (10, 30, 0)),
    "2.2": ("Стоимость строительства", 550, 200, 350, (100, 0, 200)),
    "2.2.1": ("Подготовка территории", 500, 200, 300, (100, 0, 200)),
    "2.2.11": ("Резервы", 50, 0, 50, (0, 0, 0)),
    "2.3": ("Коммерческие расходы", 90, 40, 50, (50, 0, 0)),
    "2.3.4": ("Расходы на продажу (КС)", 90, 40, 50, (50, 0, 0)),
    "2.4.9": ("Статья без карты", 10, 10, 0, (0, 0, 0)),
    "3.1.4": ("Проценты по кредитам", 70, 20, 50, (5, 5, 5)),
    "3.1.1.1.2": ("Поступления кредитов (ПФ)", 800, 500, 300, (100, 150, 50)),
}
# карта «статьи_БДДС»: строка под статьёй → код РСС
CROSSWALK = [("2.1.1.1", "Права застройщика", "1.1"), ("2.2.1.1", "Вырубка", "2.1"),
             ("2.2.11.1", "Резерв", "2.8"), ("2.3.4.1", "Аналитика", "3.2"),
             ("2.3.4.2", "Риэлторы", "3.2"), ("2.3.4.3", "Представительские", "2.6")]


def _finmodel(crosswalk: bool = True, credit: bool = True) -> bytes:
    book = openpyxl.Workbook()
    ws = book.active
    ws.title = "ОТЧЕТ"
    cf = book.create_sheet("cf_1_new")
    cf.cell(6, 9, "ИТОГО"), cf.cell(6, 10, "ФАКТ"), cf.cell(6, 11, "ПЛАН")
    # факт закрыт сентябрём, план — с октября
    cf.cell(5, 13, "факт"), cf.cell(6, 13, datetime.datetime(2026, 9, 1))
    for i, month in enumerate(MONTHS):
        cf.cell(5, 14 + i, "план"), cf.cell(6, 14 + i, month)
    for r, (code, (name, total, fact, plan, months)) in enumerate(ROWS.items(), 7):
        if not credit and code == "3.1.1.1.2":
            continue
        cf.cell(r, 3, code), cf.cell(r, 4, name)
        cf.cell(r, 9, total), cf.cell(r, 10, fact), cf.cell(r, 11, plan)
        cf.cell(r, 13, 7)
        for i, value in enumerate(months):
            cf.cell(r, 14 + i, value)
    if crosswalk:
        xw = book.create_sheet("статьи_БДДС")
        xw.cell(1, 7, "Код статьи"), xw.cell(2, 9, "Код РСС")
        for r, (bdds, name, est) in enumerate(CROSSWALK, 3):
            xw.cell(r, 7, bdds), xw.cell(r, 8, name), xw.cell(r, 9, est)
    out = io.BytesIO()
    book.save(out)
    return out.getvalue()


def _read(tmp_path: Path, data: bytes) -> dict:
    path = tmp_path / "finance.xlsx"
    path.write_bytes(data)
    return dashboard._read_finance_baseline(path)


def test_the_model_columns_become_the_baseline(tmp_path):
    got = _read(tmp_path, _finmodel())
    assert got["known"] and got["kind"] == "finmodel"
    # Только расходы «2.…»: поступления 1.2 и проценты 3.1.4 не бюджет стройки.
    assert got["approved"] == pytest.approx((100 + 500 + 50 + 90 + 10) * K)
    assert got["completion_need_at_baseline"] == pytest.approx((40 + 300 + 50 + 50) * K)
    assert got["paid_at_baseline"] == pytest.approx((60 + 200 + 40 + 10) * K)
    assert got["monthly_need"] == {"2026-10-01": pytest.approx(160 * K),
                                   "2026-11-01": pytest.approx(30 * K),
                                   "2027-04-01": pytest.approx(200 * K)}
    assert got["tail_after_apr"] == pytest.approx(200 * K)
    # Резерв — строка 2.2.11, не статья.
    assert got["reserve"] == pytest.approx(50 * K) and got["reserve_names"] == {"2.2.11": "Резервы"}
    assert set(got["articles"]) == {"2.1.1", "2.2.1", "2.3.4", "2.4.9"}
    article = got["articles"]["2.2.1"]
    assert article["rss_limit"] == pytest.approx(500 * K)
    assert article["paid_at_baseline"] == pytest.approx(200 * K)
    assert article["monthly_need"]["2027-04-01"] == pytest.approx(200 * K)
    assert article["has_programme"] is True
    # Сколько финмодель берёт у банка — строка «Поступления кредитов (ПФ)».
    credit = got["credit_plan"]
    assert credit["total"] == pytest.approx(800 * K) and credit["drawn"] == pytest.approx(500 * K)
    assert credit["plan"] == pytest.approx(300 * K)
    assert credit["monthly"] == {"2026-10-01": pytest.approx(100 * K),
                                 "2026-11-01": pytest.approx(150 * K),
                                 "2027-04-01": pytest.approx(50 * K)}


def test_the_chapters_come_from_the_crosswalk(tmp_path):
    got = _read(tmp_path, _finmodel())
    assert got["chapters"] == ["1", "2", "3"]
    assert got["approved_by_chapter"] == {"1": pytest.approx(100 * K),
                                          "2": pytest.approx(550 * K),
                                          "3": pytest.approx(90 * K)}
    assert got["articles"]["2.1.1"]["chapter"] == "1"
    # 2.3.4: две строки карты в главу 3, одна в главу 2 — глава большинства, и это сказано.
    assert got["articles"]["2.3.4"]["chapter"] == "3"
    assert "2.3.4" in got["chapters_reason"] and "разные главы" in got["chapters_reason"]
    # Статья без строк в карте — вне глав, и это тоже сказано.
    assert got["articles"]["2.4.9"]["chapter"] == ""
    assert "2.4.9" in got["chapters_reason"] and "не найдены" in got["chapters_reason"]


def test_without_the_crosswalk_the_model_is_a_reason_not_a_guess(tmp_path):
    got = _read(tmp_path, _finmodel(crosswalk=False))
    assert got["known"] is False and "статьи_БДДС" in got["reason"]


def test_a_book_with_neither_sheet_names_both(tmp_path):
    book = openpyxl.Workbook()
    book.active.title = "ОТЧЕТ"
    out = io.BytesIO()
    book.save(out)
    got = _read(tmp_path, out.getvalue())
    assert got["known"] is False
    assert "Расчет стоимости строительства" in got["reason"] and "cf_1_new" in got["reason"]


def test_the_replacement_takes_the_finmodel_and_refuses_a_stranger(tmp_path, monkeypatch):
    monkeypatch.setattr(monitor, "_SNAPSHOT_DIR", tmp_path)
    got = monitor.replace_finance_baseline("demo", _finmodel(), "2026-10-23")
    assert got["replaced"] is True
    assert dashboard._finance_baseline("demo")["kind"] == "finmodel"
    book = openpyxl.Workbook()
    book.active.title = "ОТЧЕТ"
    out = io.BytesIO()
    book.save(out)
    with pytest.raises(ValueError) as refused:
        monitor.replace_finance_baseline("demo", out.getvalue(), "2026-10-23")
    assert "cf_1_new" in str(refused.value) and "Расчет стоимости строительства" in str(refused.value)


def _finmodel_funding(monkeypatch, sections: dict | None, estimate: dict, paid: dict) -> dict:
    articles = {"2.2.1": {"code": "2.2.1", "chapter": "2", "name": "Подготовка", "rss_limit": 500.0,
                          "paid_at_baseline": 200.0, "has_programme": True,
                          "monthly_need": {"2026-11-01": 300.0}}}
    monkeypatch.setattr(dashboard, "_finance_baseline", lambda _: {
        "known": True, "kind": "finmodel", "source": "finance.xlsx",
        "approved": 1000.0, "completion_need_at_baseline": 300.0, "paid_at_baseline": 700.0,
        "reserve": 0.0, "reserve_parts": {}, "monthly_need": {}, "tail_after_apr": 0.0,
        "chapters": ["1", "2", "3"], "articles": articles, "sections": sections,
    })
    monkeypatch.setattr(dashboard.actuals, "read_estimate", lambda _: {"rows": [], "by_code": estimate})
    monkeypatch.setattr(dashboard, "_paid_in_chapters",
                        lambda rss, est, chapters=("2", "3"): {
                            "total": sum(paid[c][0] for c in chapters),
                            "borrowed": sum(paid[c][1] for c in chapters)})
    # Оплаты РСС по кодам в статьи финмодели не кладутся вовсе.
    monkeypatch.setattr(dashboard, "_payment_by_code",
                        lambda _: (_ for _ in ()).throw(AssertionError("оплаты РСС по коду не нужны")))
    monkeypatch.setattr(dashboard, "_unspent", lambda *a, **k: None)
    view = {"schedule": {"approved_end": "2027-09-30", "forecast_end": None}}
    return dashboard._funding_risk("demo", Path("rss.xlsx"), datetime.date(2026, 10, 20), view)


SECTIONS = {
    "2": {"code": "2.2", "chapter": "2", "label": "стройка", "label_acc": "стройку",
          "name": "Стоимость строительства",
          "total": 900.0, "paid": 500.0, "plan": 400.0,
          "monthly": {"2026-09-01": 50.0, "2026-11-01": 250.0, "2026-12-01": 100.0}},
    "3": {"code": "2.3", "chapter": "3", "label": "коммерция", "label_acc": "коммерцию",
          "name": "Коммерческие расходы",
          "total": 200.0, "paid": 120.0, "plan": 80.0,
          "monthly": {"2026-11-01": 30.0, "2027-01-01": 50.0}},
}
# главы сметы: лимит (с резервами); оплачено (всего, кредитом); глава 1 — вне вопроса
ESTIMATE = {"1": {"estimate": 300.0}, "2": {"estimate": 450.0}, "3": {"estimate": 80.0}}
PAID = {"1": (250.0, 200.0), "2": (400.0, 360.0), "3": (70.0, 40.0)}


def test_the_need_is_the_model_sections_against_chapters_two_and_three(monkeypatch):
    """«Сколько надо ещё денег в банке взять, чтобы достроить и продать (2–3
    главы)»; «тогда по модели» (владелец, 24.10.2026)."""
    got = _finmodel_funding(monkeypatch, SECTIONS, ESTIMATE, PAID)
    # Остаток лимита — главы 2–3, лимит минус всё оплаченное: (450 − 400) + (80 − 70).
    assert got["bank_remaining"] == pytest.approx(60.0)
    assert got["remaining_scope"] == {"chapters": "2–3", "limit": 530.0, "paid": 470.0, "borrowed": 400.0}
    # Нужно — план разделов с месяца среза (сентябрь уже прошёл): 350 + 80.
    assert got["remaining_need"] == pytest.approx(430.0)
    assert got["approved_remaining"] == pytest.approx(430.0)
    # Не хватает: стройка 350 − 50 = 300, коммерция 80 − 10 = 70; первый месяц — ноябрь.
    assert got["additional_financing"] == pytest.approx(370.0)
    assert got["additional_financing_from"] == "2026-11-01"
    assert got["deficit_ladder"]["total"] == pytest.approx(370.0)
    assert got["need_basis"] == "model"
    assert got["need_by_chapter"] == {
        "2": {"label": "стройка", "label_acc": "стройку", "need": pytest.approx(350.0),
              "remaining": pytest.approx(50.0), "gap": pytest.approx(300.0)},
        "3": {"label": "коммерция", "label_acc": "коммерцию", "need": pytest.approx(80.0),
              "remaining": pytest.approx(10.0), "gap": pytest.approx(70.0)}}
    assert [row["code"] for row in got["articles"]] == ["2", "3"]
    # Лимит глав 2–3 минус оплаченное против остатка водопада — одно число, без «разницы».
    assert got["remaining_check"]["book_gap"] == pytest.approx(0.0)
    assert "по финмодели" in got["method"]
    # Плашки «Лимит банка» и «Выбрано у банка» — по всем главам, как прежде.
    assert got["bank_limit"] == pytest.approx(830.0) and got["paid_borrowed"] == pytest.approx(600.0)


def test_without_the_sections_the_model_is_a_reason(monkeypatch, tmp_path):
    got = _finmodel_funding(monkeypatch, {"2": SECTIONS["2"]}, ESTIMATE, PAID)
    assert got["known"] is False and "«2.3»" in got["reason"]
    path = tmp_path / "finance.xlsx"
    path.write_bytes(_finmodel())
    sections = dashboard._read_finance_baseline(path)["sections"]
    assert set(sections) == {"2", "3"}
    assert sections["2"]["plan"] == pytest.approx(350 * K)
    assert sections["2"]["monthly"] == {"2026-10-01": pytest.approx(100 * K), "2026-11-01": 0.0,
                                        "2027-04-01": pytest.approx(200 * K)}
    assert sections["3"]["label_acc"] == "коммерцию" and sections["3"]["plan"] == pytest.approx(50 * K)


def test_the_report_names_the_bank_need_by_chapter_for_the_model(monkeypatch):
    import developaid_management_report as report
    from test_management_report_funding import _project, PROJECT, M

    _project()
    monkeypatch.setattr(report, "_monitor_view", lambda project, end: {"dashboard": {"funding": {
        "known": True, "need_basis": "model", "bank_limit": 6_810 * M, "paid_actual": 5_263 * M,
        "paid_borrowed": 4_832 * M, "bank_remaining": 1_278 * M, "reserve": 0.0,
        "remaining_scope": {"chapters": "2–3", "limit": 4_640 * M, "paid": 3_362 * M,
                            "borrowed": 3_270 * M},
        "need_by_chapter": {"2": {"label": "стройка", "label_acc": "стройку", "need": 3_052 * M,
                                  "remaining": 1_123 * M, "gap": 1_928 * M},
                            "3": {"label": "коммерция", "label_acc": "коммерцию", "need": 621 * M,
                                  "remaining": 155 * M, "gap": 465 * M}},
        "additional_financing": 2_393 * M, "additional_financing_from": "2026-11-01",
        "approved": 9_599 * M, "approved_remaining": 3_672 * M,
        "deficit_ladder": {"total": 2_393 * M}}}})
    built = report.build(PROJECT, "week", "2026-10-18")
    tiles = {item["label"]: item for item in report.tiles(built)}
    assert tiles["Остаток лимита банка"]["value"] == "1,28 млрд ₽"
    assert "главы 2–3: лимит 4,64 млрд ₽ без оплаченного 3,36 млрд ₽" in tiles["Остаток лимита банка"]["note"]
    assert "Дефицит на достройку" not in tiles
    tile = tiles["Потребность в допфинансировании банка"]
    assert tile["value"] == "2,39 млрд ₽"
    assert tile["note"] == ("на стройку 1,93 млрд ₽ (нужно 3,05 млрд ₽, лимита 1,12 млрд ₽), "
                            "на коммерцию 465,0 млн ₽ (нужно 621,0 млн ₽, лимита 155,0 млн ₽); "
                            "с ноября 2026")
    text = report.render_text(built)
    assert ("на стройку и коммерцию по финмодели нужно 3,67 млрд ₽ при остатке лимита глав 2–3 "
            "1,28 млрд ₽ — потребность в допфинансировании банка 2,39 млрд ₽: на стройку") in text
    assert "По финмодели не хватает: 2,39 млрд ₽, начиная" in text
    page = report.render_html(built)
    assert "Нужно на стройку по финмодели" in page and "Потребность в допфинансировании банка" in page


def test_the_need_tile_equals_the_shortfall_even_with_redistributable_limits(monkeypatch):
    """«Потребность» 2,27 против «не хватает» 2,39 на соседней плашке (владелец,
    24.10.2026): переброска лимита между статьями сметы вычиталась из водопада,
    который уже идёт по главам."""
    got_plain = _finmodel_funding(monkeypatch, SECTIONS, ESTIMATE, PAID)
    unspent = {"by_chapter": [{"chapter": "2", "shortage": 300.0, "sources": 120.0}],
               "sources": [1], "excluded": [], "needy": [], "retention": {"reason": ""}}
    monkeypatch.setattr(dashboard, "_unspent", lambda *a, **k: unspent)
    view = {"schedule": {"approved_end": "2027-09-30", "forecast_end": None}}
    got = dashboard._funding_risk("demo", Path("rss.xlsx"), datetime.date(2026, 10, 20), view)
    assert got["deficit_ladder"]["total"] == pytest.approx(got["additional_financing"])
    assert got["deficit_ladder"]["total"] == pytest.approx(got_plain["deficit_ladder"]["total"])
