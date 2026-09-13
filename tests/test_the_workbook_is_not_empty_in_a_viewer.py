"""Книга несёт сохранённые значения: она не пуста там, где формулы не считают.

Книга v4 — это сто двадцать тысяч формул и почти ни одного числа. Excel
пересчитывает её при открытии и показывает всё, а предпросмотр в телеграме,
Quick Look, телефон, Google Drive и Яндекс.Диск читают только сохранённые
значения. Их не было НИ ОДНОГО с 02.08.2026 — сборщик стирал их регуляркой на
каждом листе, — и во всех этих окнах книга выглядела пустой. Владелец открыл
выгрузку 13.09.2026 и сказал «модель пустая вообще»; сорок два дня до этого
файл в любом просмотрщике был чистым листом.

Стереть их было верным решением по доводу: смесь старых трёхочередных значений
с пустыми клонами четвёртой очереди выдавала бы непосчитанную книгу за
посчитанную. Не была названа цена. Поэтому здесь третий ответ: значение
СЧИТАЕТСЯ нашим вычислителем и кладётся рядом с живой формулой, а не вместо
неё, — `fullCalcOnLoad="1"` остаётся, и в Excel число живёт до первого открытия.

Проверяется то, что видно: книга читается как книга СО ЗНАЧЕНИЯМИ, и они те же,
что считает вычислитель. Проверка исходника здесь не значила бы ничего — строки
регулярок стоят в файле и у сломанного сборщика.

Запуск: python3 -m pytest tests/test_the_workbook_is_not_empty_in_a_viewer.py -q
"""

from __future__ import annotations

import io
import re
import sys
import zipfile
from pathlib import Path

import openpyxl
import pytest

ROOT = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(ROOT))

import main as wrapper  # noqa: E402
from xlsx_eval import Evaluator  # noqa: E402

core = wrapper.core

# Показатели, ради которых книгу и открывают. Пустая клетка здесь и есть
# «модель пустая вообще».
_HEADLINE = {"B5": "выручка", "B6": "CAPEX", "B8": "EBITDA",
             "B12": "чистая прибыль", "B19": "LLCR"}


@pytest.fixture(scope="module")
def workbook() -> bytes:
    sys.setrecursionlimit(400000)
    inputs = dict(core.DEFAULT_INPUTS)
    tep = {key: dict(value) for key, value in core.TEP_DEFAULT.items()}
    # Проход выключен в наборе (`conftest`) — здесь он и проверяется, значит
    # включается явно, а не через окружение: тест, зависящий от порядка
    # переменных, зелен по случайности.
    content, _, meta = core.build_project_workbook(
        inputs, tep, [], {}, cache_values=True)
    assert not [m for m in (meta.get("missing") or []) if "сохранённые значения" in str(m)], \
        meta.get("missing")
    return content


def test_a_viewer_that_does_not_compute_still_sees_the_numbers(workbook):
    """Ключевые показатели ОТЧЁТа читаются без пересчёта формул."""
    book = openpyxl.load_workbook(io.BytesIO(workbook), data_only=True)
    sheet = book["ОТЧЕТ"]
    for coord, label in _HEADLINE.items():
        value = sheet[coord].value
        assert isinstance(value, (int, float)), f"{label} ({coord}) пуста без пересчёта"


def test_the_saved_number_is_what_the_formula_computes(workbook):
    """Сохранённое значение равно посчитанному — кэш не врёт.

    Соврать здесь дороже, чем промолчать: число видит ровно тот, кто проверить
    его не может.
    """
    values = openpyxl.load_workbook(io.BytesIO(workbook), data_only=True)
    formulas = openpyxl.load_workbook(io.BytesIO(workbook), data_only=False)
    evaluator = Evaluator(formulas)
    checked = 0
    for sheet in ("ОТЧЕТ", "ПРОВЕРКИ", "CF_1"):
        for row in values[sheet].iter_rows():
            for cell in row:
                if not isinstance(cell.value, (int, float)):
                    continue
                source = formulas[sheet][cell.coordinate].value
                if not (isinstance(source, str) and source.startswith("=")):
                    continue
                expected = evaluator.cell(sheet, cell.coordinate)
                if not isinstance(expected, (int, float)):
                    continue
                assert cell.value == pytest.approx(float(expected), rel=1e-9, abs=1e-9), \
                    f"{sheet}!{cell.coordinate}"
                checked += 1
    assert checked > 500, f"сверено всего {checked} клеток — проверка ничего не значит"


def test_almost_every_formula_carries_its_value(workbook):
    """Значение стоит почти у каждой формулы, а не у горстки.

    Частичный кэш — это те же пустые клетки, только вперемешку с числами:
    читатель решит, что пустая клетка и есть ответ. На прежнем сборщике доля
    была РОВНО НОЛЬ, поэтому проверка валится на нём первой же строкой.
    """
    with zipfile.ZipFile(io.BytesIO(workbook)) as archive:
        formulas = cached = 0
        for name in archive.namelist():
            if not re.fullmatch(r"xl/worksheets/\w+\.xml", name):
                continue
            text = archive.read(name).decode("utf-8")
            for cell in re.findall(r"<(?:x:)?c\b[^>]*>(.*?)</(?:x:)?c>", text, re.S):
                if not re.search(r"<(?:x:)?f[\s/>]", cell):
                    continue
                formulas += 1
                cached += 1 if re.search(r"<(?:x:)?v>", cell) else 0
    assert formulas > 10000, f"формул всего {formulas} — книга собралась не та"
    assert cached / formulas > 0.9, (
        f"значение стоит у {cached} формул из {formulas}")


def test_excel_still_recomputes_on_open(workbook):
    """`fullCalcOnLoad` на месте: в Excel считает книга, а не наш кэш.

    Без него сохранённое число становится вторым ответом об одной величине — и
    устареет молча при первой же правке вводной.
    """
    with zipfile.ZipFile(io.BytesIO(workbook)) as archive:
        book_xml = archive.read("xl/workbook.xml").decode("utf-8")
    assert 'fullCalcOnLoad="1"' in book_xml
