"""Выгрузка торгов в Excel: текст площадки — данные, даты — даты, ссылки — http(s).

Строка источника, начинавшаяся с «=», писалась формулой и исполнялась в Excel;
даты заявок лежали текстом, и сортировка шла по строке («01.10» раньше
«21.09»); ссылка `javascript:` попадала в гиперссылку; стиль Hyperlink
сбрасывал перенос строк в «Адресе».
"""

from __future__ import annotations

import datetime
import io
import zipfile

import openpyxl

from auction_search.api import _xlsx

ROWS = [
    {"section": "Торги", "name": '=HYPERLINK("http://evil","клик")',
     "address": "Москва, ул. Тверская, 1", "application_deadline": "21.09.2026 18:00",
     "auction_date": "в течение месяца", "url": "javascript:alert(1)"},
    {"section": "Торги", "name": "Лот 2", "url": "https://torgi.gov.ru/lot/2",
     "application_deadline": "01.10.2026"},
]


def _sheet():
    data = _xlsx(ROWS, "auctions")
    ws = openpyxl.load_workbook(io.BytesIO(data)).active
    header = [cell.value for cell in ws[1]]
    return data, ws, (lambda row, name: ws[row][header.index(name)])


def test_a_leading_equals_is_text_not_a_formula():
    data, _ws, cell = _sheet()
    assert cell(2, "Название").data_type == "s"
    sheet_xml = zipfile.ZipFile(io.BytesIO(data)).read("xl/worksheets/sheet1.xml").decode()
    assert "<f>" not in sheet_xml, "текст площадки записан формулой"


def test_dates_are_dates_and_prose_stays_prose():
    _data, _ws, cell = _sheet()
    deadline = cell(2, "Окончание приёма заявок")
    assert deadline.value == datetime.datetime(2026, 9, 21, 18, 0)
    assert deadline.number_format == "DD.MM.YYYY HH:MM"
    assert cell(3, "Окончание приёма заявок").number_format == "DD.MM.YYYY"
    assert cell(2, "Дата торгов").value == "в течение месяца"


def test_only_http_links_become_hyperlinks_and_address_still_wraps():
    _data, _ws, cell = _sheet()
    assert cell(2, "Источник").hyperlink is None
    assert cell(3, "Источник").hyperlink.target == "https://torgi.gov.ru/lot/2"
    address = cell(2, "Адрес")
    assert address.hyperlink is not None
    assert address.alignment.wrap_text is True
