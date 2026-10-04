"""Стратегия реализации ТЦ и офисов: ДДУ, прямая продажа, доходный метод.

Владелец (03.10.2026): у объекта нежилья выбор — продажа по ДДУ 214-ФЗ с
эскроу (как было), прямая продажа без эскроу или доходный метод (аренда, NOI,
выход по ставке капитализации). Финансирование (б) и (в) — своё, без эскроу.
Выбор есть только у ТЦ и офисов.

Проверки держат три обещания:
* ДДУ — умолчание, и проект без поля считается ровно как раньше;
* объект вне ДДУ идёт мимо эскроу и ПФ проекта, а его деньги входят в прибыль,
  налог, НДС и поток капитала ОДИН раз;
* арифметика аренды и выхода — та, что названа (ручной пересчёт).

Запуск: python3 -m pytest tests/test_nonres_strategy_switch.py -q
"""

from __future__ import annotations

import copy
import sys
from datetime import date
from pathlib import Path

import pytest

ROOT = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(ROOT))

import developaid_nonres_strategy as ns  # noqa: E402
import main_legacy as core  # noqa: E402

OFFICE = dict(offices_enabled=True, offices_gba_sqm=40000.0, offices_saleable_sqm=24000.0,
              offices_parking_under_spaces=200, offices_parking_over_spaces=40,
              offices_parking_guest_pct=10)


def _spec(**over):
    x = copy.deepcopy(core.DEFAULT_INPUTS)
    t = copy.deepcopy(core.TEP_DEFAULT)
    x.update(OFFICE)
    x.update(over)
    for obj in core.STANDALONE_OBJECTS:
        gba = float(x.get(f"{obj.prefix}_gba_sqm") or 0)
        if not x.get(f"{obj.prefix}_enabled") or gba <= 0:
            continue
        saleable = float(x.get(f"{obj.prefix}_saleable_sqm") or 0)
        t.setdefault(obj.key, {}).update(gns=gba, total_area=round(gba * 0.94, 2),
                                         useful=saleable, saleable=saleable)
    return x, t


_CACHE: dict[tuple, dict] = {}


def _run(**over) -> dict:
    key = tuple(sorted(over.items()))
    if key not in _CACHE:
        x, t = _spec(**over)
        _CACHE[key] = core._run_authoritative_model(x, t, [], {})["consolidated"]
    return _CACHE[key]


def _rows_sum(result: dict, name: str) -> float:
    return sum(float(row.get(name, 0.0) or 0.0) for row in result["finance"]["rows"])


def _object(result: dict, key: str = "offices") -> dict:
    return next(o for o in result["finance"]["nonres"]["objects"] if o["key"] == key)


# --- реестр и умолчания ------------------------------------------------------

def test_only_retail_and_offices_have_a_strategy() -> None:
    with_choice = {obj.key for obj in core.STANDALONE_OBJECTS if obj.strategies}
    assert {"offices", "standalone_retail", "offices2", "standalone_retail5"} <= with_choice
    assert not any(key.startswith(("above_parking", "sports")) for key in with_choice)
    for obj in core.STANDALONE_OBJECTS:
        assert (f"{obj.prefix}_strategy" in core.DEFAULT_INPUTS) == obj.strategies


def test_the_default_strategy_is_ddu_and_a_missing_field_reads_as_ddu() -> None:
    assert core.DEFAULT_INPUTS["offices_strategy"] == ns.STRATEGY_DDU
    assert ns.object_strategy({}, "offices") == ns.STRATEGY_DDU
    assert ns.object_strategy({"offices_strategy": "что-то"}, "offices") == ns.STRATEGY_DDU
    explicit = _run(offices_strategy="ddu")
    x, t = _spec()
    x.pop("offices_strategy")
    missing = core._run_authoritative_model(x, t, [], {})["consolidated"]
    for field in ("net_profit", "npv", "llcr", "irr_equity"):
        assert missing["summary"][field] == explicit["summary"][field]
    assert explicit["finance"]["nonres"]["objects"] == []


def test_the_page_offers_the_switch_in_the_object_group() -> None:
    group = core.standalone_object_group(core.standalone_objects(("offices",))[0])
    fields = {field[0]: field for field in group[1]}
    switch = fields["offices_strategy"]
    assert switch[3] == "select"
    assert [value for value, _ in switch[4]] == ["ddu", "direct", "income"]
    assert core.standalone_object_section(
        core.standalone_objects(("offices",))[0], "offices_rent_th_per_sqm_month") == "Доходный метод"
    # У ТЦ места не сдаются и не продаются — поля аренды места у него нет.
    retail = core.standalone_object_group(core.standalone_objects(("standalone_retail",))[0])
    assert "retail_parking_rent_th_month" not in {field[0] for field in retail[1]}
    assert "retail_strategy" in {field[0] for field in retail[1]}


# --- мимо эскроу и ПФ ----------------------------------------------------------

@pytest.mark.parametrize("strategy", ["direct", "income"])
def test_an_object_outside_ddu_bypasses_the_escrow_and_the_project_loan(strategy) -> None:
    ddu = _run(offices_strategy="ddu")
    other = _run(offices_strategy=strategy)
    capex = _object(other)["totals"]["capex"]
    assert capex > 0
    # Эскроу не получает ни рубля объекта: вся выручка ДДУ — без офисов.
    assert _rows_sum(other, "sales") < _rows_sum(ddu, "sales")
    office_ddu = ddu["finance"]["tax_margin_by_product"]
    assert "offices" in office_ddu
    # Кредит объекта: выборка = доля × затраты стройки; к концу погашен.
    loan = _object(other)["totals"]
    share = core.DEFAULT_INPUTS["offices_loan_share_pct"] / 100
    assert loan["loan_draw"] == pytest.approx(capex * share, rel=1e-9)
    # Гасится тело и капитализированные до ввода проценты — не больше и не меньше.
    capitalized = (_rows_sum(other, "nonres_loan_interest")
                   - _rows_sum(other, "nonres_loan_interest_paid"))
    assert capitalized > 0
    assert loan["loan_repayment"] == pytest.approx(loan["loan_draw"] + capitalized, rel=1e-9)
    assert other["finance"]["rows"][-1]["nonres_loan_balance"] == pytest.approx(0.0, abs=1e-3)
    # ПФ проекта не финансирует объект: выборка ПФ меньше хотя бы на долю,
    # которую раньше давал объект.
    assert _rows_sum(other, "pf_draw") < _rows_sum(ddu, "pf_draw")


@pytest.mark.parametrize("strategy,extra", [
    ("ddu", {}), ("direct", {}), ("income", {}),
    ("income", {"offices_exit_mode": "hold"}),
    ("income", {"offices_debt_repayment": "bullet"}),
])
def test_the_object_money_enters_the_project_once(strategy, extra) -> None:
    """Поток проекта минус чистая прибыль — одна и та же величина при любой
    стратегии: НДС проекта (поток его не несёт). Объект, посчитанный дважды
    или забытый в потоке, сдвинул бы разницу на свою выручку."""
    result = _run(offices_strategy=strategy, **extra)
    finance = result["finance"]
    object_vat = finance["nonres"]["totals"].get("nonres_vat_paid", 0.0)
    gap = (sum(result["cashflow"]["project"]) - result["summary"]["net_profit"]
           - (finance["vat"] - object_vat))
    assert gap == pytest.approx(0.0, abs=1.0)
    if strategy != "ddu":
        assert result["revenue"]["offices"] == pytest.approx(
            _object(result)["totals"]["revenue"])


def test_the_object_revenue_moves_the_profit_and_the_tax() -> None:
    low = _run(offices_strategy="income", offices_rent_th_per_sqm_month=3.0)
    high = _run(offices_strategy="income", offices_rent_th_per_sqm_month=6.0)
    assert high["summary"]["net_profit"] > low["summary"]["net_profit"]
    assert high["finance"]["profit_tax"] > low["finance"]["profit_tax"]
    assert (high["finance"]["tax_margin_by_product"]["offices"]
            > low["finance"]["tax_margin_by_product"]["offices"])


def test_a_bullet_loan_costs_more_interest_than_a_sweep() -> None:
    sweep = _object(_run(offices_strategy="income"))["totals"]
    bullet = _object(_run(offices_strategy="income", offices_debt_repayment="bullet"))["totals"]
    assert bullet["loan_interest"] > sweep["loan_interest"]
    assert bullet["loan_draw"] == pytest.approx(sweep["loan_draw"])


def test_hold_shows_the_value_without_a_sale() -> None:
    sale = _object(_run(offices_strategy="income"))
    hold = _object(_run(offices_strategy="income", offices_exit_mode="hold"))
    assert sale["totals"]["exit_revenue"] > 0 and sale["totals"]["residual_value"] == 0
    assert hold["totals"]["exit_revenue"] == 0
    assert hold["totals"]["residual_value"] == pytest.approx(sale["kpi"]["exit_value"])
    # Сделки нет — нет ни затрат выхода, ни НДС с продажи.
    assert hold["totals"]["exit_cost"] == 0
    assert hold["totals"]["vat_charged"] < sale["totals"]["vat_charged"]


def test_the_peak_debt_counts_the_object_loan() -> None:
    result = _run(offices_strategy="income")
    peak = core._peak_total_debt(result)
    without = max(float(r.get("bridge_balance") or 0) + float(r.get("pf_balance") or 0)
                  for r in result["finance"]["rows"])
    assert _object(result)["totals"]["loan_peak"] > 0
    assert peak > without


# --- арифметика модуля (ручной пересчёт) -------------------------------------

def _plan(strategy: str, **params) -> dict:
    capex = {date(2027, m, 1): 100_000_000.0 for m in range(1, 13)}
    base = {suffix: pair[0] for suffix, pair in ns.STRATEGY_FIELD_DEFAULTS.items()}
    base.update(params)
    return {"strategy": strategy, "capex": capex, "commissioning": date(2028, 1, 1),
            "area_sqm": 10_000.0, "price_rub_sqm": 300_000.0,
            "price_start": date(2027, 6, 1), "growth_pre": 0.01, "growth_post": 0.0,
            "parking_spaces": 0.0, "parking_price_rub": 0.0, "selling_share": 0.03,
            "params": base}


def test_rent_noi_and_exit_follow_the_named_formulas() -> None:
    plan = _plan("income", rent_th_per_sqm_month=4.0, occupancy_start_pct=50,
                 occupancy_stable_pct=90, leaseup_months=7, rent_index_pct=6,
                 opex_pct=20, hold_years=3, exit_cap_pct=10, property_tax_pct=2.0)
    flows = ns.object_flows(plan, lambda m: 0.10, vat_rate=0.22)
    monthly = flows["monthly"]
    # Месяц 4 эксплуатации: загрузка 50 + (90−50)·3/6 = 70%, индексация 1,06^(3/12).
    month4 = date(2028, 4, 1)
    rent = 10_000 * 4000 * 1.06 ** 0.25 * 0.70
    assert monthly["rent_revenue"][month4] == pytest.approx(rent)
    assert monthly["opex"][month4] == pytest.approx(rent * 0.20)
    # Выход — NOI следующих 12 месяцев минус налог на имущество, делённые на ставку.
    basis = 1_200_000_000 / 1.22
    forward = sum(10_000 * 4000 * 1.06 ** ((36 + k - 1) / 12) * 0.90 * 0.80
                  for k in range(1, 13)) - basis * 0.02
    assert flows["kpi"]["exit_value"] == pytest.approx(forward / 0.10)
    assert monthly["exit_revenue"][date(2030, 12, 1)] == pytest.approx(forward / 0.10)


def test_direct_sale_starts_at_commissioning_and_sells_everything() -> None:
    plan = _plan("direct", direct_sale_months=4, direct_sale_offset_months=2)
    flows = ns.object_flows(plan, lambda m: 0.10, vat_rate=0.22)
    sales = flows["monthly"]["sale_revenue"]
    assert min(sales) == date(2028, 3, 1) and max(sales) == date(2028, 6, 1)
    # Цена растёт 1%/мес. от старта цены до ввода (7 мес.), дальше — 0%.
    assert sum(sales.values()) == pytest.approx(10_000 * 300_000 * 1.01 ** 7)
    assert flows["totals"]["loan_draw"] == pytest.approx(1_200_000_000 * 0.6)
    assert flows["monthly"]["loan_balance"].get(date(2028, 6, 1), 0.0) == pytest.approx(0.0, abs=1e-6)
    assert flows["monthly"]["loan_balance"][date(2027, 12, 1)] > 0
    # Проценты до ввода капитализируются, после — платятся.
    assert flows["monthly"]["loan_interest_cap"][date(2027, 12, 1)] > 0
    assert date(2028, 2, 1) in flows["monthly"]["loan_interest_paid"]


def test_vat_on_the_object_nets_input_against_output() -> None:
    flows = ns.object_flows(_plan("direct"), lambda m: 0.10, vat_rate=0.22)
    totals = flows["totals"]
    input_vat = 1_200_000_000 * 0.22 / 1.22
    assert totals["vat_paid"] == pytest.approx(totals["vat_charged"] - input_vat)


def test_direct_sale_can_start_before_commissioning_with_advances() -> None:
    """Предварительный ДКП: авансы до ввода, без эскроу (владелец, 03.10.2026).

    Деньги приходят сразу и гасят кредит объекта; НДС с аванса — в месяц
    оплаты; прибыль признаётся не раньше ввода."""
    after = ns.object_flows(_plan("direct", direct_sale_months=6), lambda m: 0.10, vat_rate=0.22)
    early = ns.object_flows(_plan("direct", direct_sale_months=6, direct_sale_offset_months=-6),
                            lambda m: 0.10, vat_rate=0.22)
    sales = early["monthly"]["sale_revenue"]
    assert min(sales) == date(2027, 7, 1) and max(sales) == date(2027, 12, 1)
    # Аванс гасит кредит до ввода — процентов меньше, чем при продаже после.
    assert early["monthly"]["loan_repayment"][date(2027, 7, 1)] > 0
    assert early["totals"]["loan_interest"] < after["totals"]["loan_interest"]
    # НДС с аванса начислен в месяц оплаты.
    assert early["monthly"]["vat_charged"][date(2027, 7, 1)] == pytest.approx(
        sales[date(2027, 7, 1)] * 0.22 / 1.22)
    # Налоговая маржа до ввода — только расходы на продажу, выручка — в месяц ввода.
    margin = early["monthly"]["tax_margin"]
    assert all(value <= 0 for when, value in margin.items() if when < date(2028, 1, 1))
    assert margin[date(2028, 1, 1)] > 0
    assert early["totals"]["tax_margin"] == pytest.approx(
        early["totals"]["sale_revenue"] - early["totals"]["vat_charged"]
        - 1_200_000_000 / 1.22 - early["totals"]["selling_cost"]
        - early["totals"]["property_tax"])
    # Горизонт не короче ввода: кредит закрывается не позже него.
    assert early["horizon_end"] == date(2028, 1, 1)


def test_the_project_takes_advances_before_commissioning() -> None:
    result = _run(offices_strategy="direct", offices_direct_sale_offset_months=-12)
    office = _object(result)
    assert office["totals"]["sale_revenue"] > 0
    first_sale = min(d for d, v in (
        (core.d(r["month"]), float(r.get("nonres_revenue") or 0)) for r in result["finance"]["rows"]) if v)
    assert first_sale < core.d(office["commissioning"])


# --- отчёт и PDF -------------------------------------------------------------

def test_the_report_table_reads_the_same_totals() -> None:
    result = _run(offices_strategy="income")
    [table] = result["report"]["nonres_strategy"]
    rows = {row["label"]: row["value"] for row in table["rows"]}
    office = _object(result)
    assert rows["Затраты стройки объекта"] == office["totals"]["capex"]
    assert rows["NOI за срок удержания"] == office["totals"]["noi"]
    assert rows["Выход: продажа по ставке капитализации"] == office["kpi"]["exit_value"]
    assert _run(offices_strategy="ddu")["report"]["nonres_strategy"] == []


def test_the_pdf_prints_the_strategy_table(tmp_path) -> None:
    pypdf = pytest.importorskip("pypdf")
    x, t = _spec(offices_strategy="income")
    result = core.calculate(core.CalcRequest(inputs=x, tep=t, rates=[]))
    content = core._build_developaid_pdf({"project_name": "Офис в аренду", "result": result,
                                          "inputs": x, "tep": t, "rates": []})
    path = tmp_path / "nonres.pdf"
    path.write_bytes(content)
    text = "\n".join(page.extract_text() or "" for page in pypdf.PdfReader(str(path)).pages)
    flat = " ".join(text.split())
    assert "Нежильё — стратегия реализации: Офисы" in flat
    assert "Доходный метод: аренда и выход" in flat
    exit_value = next(row["value"] for row in result["report"]["nonres_strategy"][0]["rows"]
                      if row["label"].startswith("Выход"))
    assert " ".join(core._pdf_money(exit_value).split()) in flat


# --- книга -------------------------------------------------------------------

def _book(**over):
    import io
    import openpyxl
    x, t = _spec(**over)
    content, _name, meta = core.build_project_workbook(x, t, [], {})
    return openpyxl.load_workbook(io.BytesIO(content)), meta


def test_the_book_carries_the_engine_result_sheet_and_names_the_gap() -> None:
    book, meta = _book(offices_strategy="income")
    assert "Нежильё — стратегия" in book.sheetnames
    sheet = book["Нежильё — стратегия"]
    assert str(sheet["A2"].value).startswith("СЧИТАЕТ ДВИЖОК")
    values = {sheet.cell(row, 1).value: sheet.cell(row, 2).value for row in range(1, sheet.max_row + 1)}
    report = _run(offices_strategy="income")["report"]["nonres_strategy"][0]
    for row in report["rows"]:
        if row["unit"] == "rub":
            assert values[row["label"]] == pytest.approx(row["value"], rel=1e-9)
    assert any("Нежильё — стратегия: Офисы" in item for item in meta["missing"])
    # Месячные ряды сходятся с итогом движка.
    header = next(r for r in range(1, sheet.max_row + 1) if sheet.cell(r, 1).value == "Месяц")
    revenue = sum(float(sheet.cell(r, 2).value or 0) for r in range(header + 1, sheet.max_row + 1))
    office = _object(_run(offices_strategy="income"))
    assert revenue == pytest.approx(office["totals"]["revenue"], rel=1e-9)


def test_a_ddu_book_has_no_strategy_sheet() -> None:
    book, meta = _book(offices_strategy="ddu")
    assert "Нежильё — стратегия" not in book.sheetnames
    assert not any("Нежильё — стратегия" in item for item in meta["missing"])


# --- кредит удержания и общие затраты нежилого проекта -----------------------

def test_the_hold_loan_amortizes_to_its_balloon() -> None:
    """После ввода — аннуитет на срок кредита, в конце баллон (владелец, 04.10.2026)."""
    plan = _plan("income", hold_years=12, loan_term_years=10, loan_balloon_pct=20,
                 debt_repayment="annuity", rent_th_per_sqm_month=6.0)
    flows = ns.object_flows(plan, lambda m: 0.10, vat_rate=0.22)
    monthly, kpi = flows["monthly"], flows["kpi"]
    at_commissioning = monthly["loan_balance"][date(2028, 1, 1)]
    assert kpi["loan_balloon"] == pytest.approx(at_commissioning * 0.20)
    assert kpi["loan_maturity"] == date(2038, 1, 1)
    # Платёж (проценты + тело) постоянен при постоянной ставке.
    payments = [monthly["loan_interest_paid"][mm] + monthly["loan_amortization"][mm]
                for mm in (date(2028, 2, 1), date(2031, 6, 1), date(2037, 12, 1))]
    assert payments[0] == pytest.approx(payments[1]) == pytest.approx(payments[2])
    # Перед баллоном долг ровно баллон, в срок он погашен.
    assert monthly["loan_balance"][date(2037, 12, 1)] == pytest.approx(kpi["loan_balloon"])
    assert monthly["loan_repayment"][date(2038, 1, 1)] == pytest.approx(kpi["loan_balloon"])
    assert monthly["loan_balance"].get(date(2038, 1, 1), 0.0) == pytest.approx(0.0, abs=1e-6)
    assert kpi["dscr_min"] is not None and kpi["dscr_min"] < kpi["icr_min"]


def test_a_sale_before_maturity_repays_the_rest_from_the_exit() -> None:
    flows = ns.object_flows(_plan("income", hold_years=5, loan_term_years=10),
                            lambda m: 0.10, vat_rate=0.22)
    exit_month = flows["kpi"]["exit_month"]
    assert flows["monthly"]["loan_repayment"][exit_month] > flows["kpi"]["loan_balloon"]
    assert flows["monthly"]["loan_balance"].get(exit_month, 0.0) == pytest.approx(0.0, abs=1e-6)


def _nonresidential(**over):
    x, t = _spec(**over)
    for key in core.MKD_PRODUCTS:
        for col in ("gns", "total_area", "useful", "saleable", "transfer", "units"):
            if key in t:
                t[key][col] = 0
    x["project_kind"] = core.PROJECT_KIND_NONRESIDENTIAL
    return core._run_authoritative_model(x, t, [], {})["consolidated"]


def test_a_nonresidential_project_lends_on_the_whole_cost() -> None:
    """Офисник без жилья: общие затраты (участок, надбавки) — в кредит объекта,
    БРИДЖа и ПФ без эскроу нет."""
    result = _nonresidential(offices_strategy="income")
    office = _object(result)["totals"]
    total_capex = result["finance"]["total_capex"]
    share = core.DEFAULT_INPUTS["offices_loan_share_pct"] / 100
    assert office["loan_draw"] > office["capex"] * share
    assert office["loan_draw"] <= total_capex * share + 1.0
    assert result["finance"]["peak_pf"] == pytest.approx(0.0, abs=1.0)
    assert _rows_sum(result, "pf_draw") == pytest.approx(0.0, abs=1.0)
    assert _rows_sum(result, "bridge_draw") == pytest.approx(0.0, abs=1.0)


def test_a_mixed_project_keeps_the_common_costs_on_the_housing_loan() -> None:
    result = _run(offices_strategy="income")
    office = _object(result)["totals"]
    share = core.DEFAULT_INPUTS["offices_loan_share_pct"] / 100
    assert office["loan_draw"] == pytest.approx(office["capex"] * share, rel=1e-9)
    assert _rows_sum(result, "pf_draw") > 0


def _pdf_text(result, x, t, tmp_path, name):
    pypdf = pytest.importorskip("pypdf")
    content = core._build_developaid_pdf({"project_name": name, "result": result,
                                          "inputs": x, "tep": t, "rates": []})
    path = tmp_path / f"{name}.pdf"
    path.write_bytes(content)
    return " ".join("\n".join(p.extract_text() or "" for p in pypdf.PdfReader(str(path)).pages).split())


def test_the_nonresidential_pdf_has_no_bank_blocks(tmp_path) -> None:
    x, t = _spec(offices_strategy="income")
    for key in core.MKD_PRODUCTS:
        for col in ("gns", "total_area", "useful", "saleable", "transfer", "units"):
            if key in t:
                t[key][col] = 0
    x["project_kind"] = core.PROJECT_KIND_NONRESIDENTIAL
    result = core._run_authoritative_model(x, t, [], {})["consolidated"]
    assert result["report"]["layout"]["project_finance"] is False
    text = _pdf_text(result, x, t, tmp_path, "nonres")
    assert "Проект без продаж по ДДУ" in text
    assert "DSCR — минимум по годам" in text
    assert "Расчётный БРИДЖ" not in text and "Эскроу против обязательств" not in text
    mixed = _run(offices_strategy="income")
    xm, tm = _spec(offices_strategy="income")
    mixed_text = _pdf_text(mixed, xm, tm, tmp_path, "mixed")
    assert "Расчётный БРИДЖ" in mixed_text and "Проект без продаж по ДДУ" not in mixed_text


def test_the_nonresidential_teaser_speaks_of_the_object_loan(tmp_path) -> None:
    pypdf = pytest.importorskip("pypdf")
    x, t = _spec(offices_strategy="income")
    for key in core.MKD_PRODUCTS:
        for col in ("gns", "total_area", "useful", "saleable", "transfer", "units"):
            if key in t:
                t[key][col] = 0
    x["project_kind"] = core.PROJECT_KIND_NONRESIDENTIAL

    def text_of(inputs, tep, name):
        bundle = core._run_authoritative_model(inputs, tep, [], {})
        pdf = core.build_teaser_pdf(bundle, inputs, tep, {})
        path = tmp_path / f"{name}.pdf"
        path.write_bytes(pdf)
        return " ".join("\n".join(p.extract_text() or "" for p in pypdf.PdfReader(str(path)).pages).split()), bundle

    text, bundle = text_of(x, t, "teaser_nonres")
    assert "DSCR — минимум по годам" in text
    assert "LLCR" not in text and "Пик эскроу" not in text and "Пик БРИДЖа" not in text
    presentation = core.project_presentation(bundle, x, t, {})
    assert not {"llcr", "peak_bridge_mln", "peak_pf_mln"} & {k["key"] for k in presentation["kpi"]}
    assert not {"llcr_below_target", "weakest_phase"} & {r["key"] for r in presentation["risks"]}
    xm, tm = _spec(offices_strategy="income")
    mixed, _ = text_of(xm, tm, "teaser_mixed")
    assert "LLCR" in mixed and "Пик эскроу" in mixed


def test_the_bot_card_of_a_nonresidential_project_names_the_object_loan(monkeypatch) -> None:
    x, t = _spec(offices_strategy="income")
    for key in core.MKD_PRODUCTS:
        for col in ("gns", "total_area", "useful", "saleable", "transfer", "units"):
            if key in t:
                t[key][col] = 0
    x["project_kind"] = core.PROJECT_KIND_NONRESIDENTIAL
    layout = core._run_authoritative_model(x, t, [], {})["consolidated"]["report"]["layout"]

    def card(summary):
        sent: list[str] = []
        monkeypatch.setattr(core, "_telegram_verify_session", lambda s: {"chat_id": 42, "cad": []})
        monkeypatch.setattr(core, "_telegram_user_allowed", lambda c: True)
        monkeypatch.setattr(core, "_telegram_send_message", lambda chat_id, text, **kw: sent.append(text))
        monkeypatch.setattr(core, "_telegram_web_app_url", lambda *a, **k: "https://example.org/")
        core.telegram_result(core.TelegramResultRequest(session="s", summary={
            "purchase_price_mln": 500, "net_profit_mln": 900, "llcr": 0.0,
            "revenue_mln": 12000, "total_expenses_mln": 11000, **summary}))
        return sent[0]

    text = card({"layout": layout})
    assert "LLCR" not in text and "БРИДЖ" not in text and "квартиры" not in text
    assert "dscr — минимум по годам" in text
    dscr = next(tile for tile in layout["nonres_tiles"] if tile["unit"] == "mult")
    assert core._telegram_tile(dscr) in text
    assert "LLCR" in card({})  # старый результат без решения — как прежде


@pytest.mark.parametrize("dscr,status", [(0.9, "negative"), (1.1, "review"), (1.5, "positive")])
def test_the_verdict_of_a_project_without_ddu_reads_the_dscr(dscr, status) -> None:
    verdict = core._purchase_feasibility(500, 900, 0.0, 0.0, 0.0, None, 0.0, None, dscr)
    assert verdict["status"] == status
    assert "DSCR" in verdict["text"] and "LLCR" not in verdict["text"]
    assert core._layout_dscr({"project_finance": True, "nonres_tiles": [{"unit": "mult", "value": 0.5}]}) is None
    assert core._layout_dscr({"project_finance": False, "nonres_tiles": [
        {"unit": "mult", "value": 1.4}, {"unit": "rub", "value": 9.0}]}) == 1.4
