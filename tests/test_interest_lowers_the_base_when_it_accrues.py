"""Проценты уменьшают базу налога в месяце начисления, а не выплаты.

П. 8 ст. 272 НК: проценты по договорам, действующим больше одного отчётного
периода, признаются расходом на конец КАЖДОГО МЕСЯЦА, независимо от даты
выплат. В этой модели платится всё разом в РВЭ — до раскрытия эскроу проценты
отсрочены, — то есть при признании «по уплате» вычет уезжал на год-полтора
вперёд от месяца, где расход возник.

Пока очереди прибыльны, разницы нет: база копится нарастающим итогом, и всё
признанное рано или поздно доходит до налога. Она появляется там, где есть
УБЫТОЧНАЯ очередь: половинное ограничение ст. 283 считает по годам, и год, в
котором расход признан, решает, сколько убытка можно зачесть.

Поэтому у теста два предохранителя. Первый: выплата обязана быть отложена —
на проекте, где проценты платятся в месяце начисления, обе меры совпадают, и
проверять нечего. Второй: две меры обязаны разойтись в НАЛОГЕ — иначе тест
зелен на любой из них.

Запуск: python3 -m pytest tests/test_interest_lowers_the_base_when_it_accrues.py -q
"""

from __future__ import annotations

import io
import sys
from collections import defaultdict
from pathlib import Path

import openpyxl
import pytest

ROOT = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(ROOT))
sys.path.insert(0, str(Path(__file__).resolve().parent))

import main_legacy as core  # noqa: E402
from xlsx_eval import Evaluator  # noqa: E402

ACCRUAL_KEYS = ("bridge_interest", "bridge_capitalization",
                "pf_interest", "pf_interest_capitalization", "limit_fee")


def _losing_scenario() -> tuple[dict, dict, dict]:
    """Тот же дорогой вход, что у `test_a_losing_queue_shares_its_loss`.

    Второго проекта под ту же ветку не заводим: у обоих тестов одна почва, и
    разойдясь, они однажды сказали бы про неё разное.
    """
    inputs = dict(core.DEFAULT_INPUTS)
    inputs.update(apartment_price_th=650, commercial_price_th=650,
                  parking_price_th=5000, purchase_price_mln=26000)
    tep = {key: dict(value) for key, value in core.TEP_DEFAULT.items()}
    phasing = {
        "enabled": True, "phase_count": 2, "phase_gap_months": 12,
        "phases": [
            {"name": "О1", "start_offset_months": 0, "construction_months": 24},
            {"name": "О2", "start_offset_months": 12, "construction_months": 24},
        ],
        "social_objects": [],
        "discrete": {"offices": 2, "standalone_retail": 2, "above_parking": 2},
    }
    return inputs, tep, phasing


@pytest.fixture(scope="module")
def single():
    inputs = dict(core.DEFAULT_INPUTS)
    inputs.update(apartment_price_th=650, commercial_price_th=650,
                  parking_price_th=5000)
    tep = {key: dict(value) for key, value in core.TEP_DEFAULT.items()}
    return core.calculate(core.CalcRequest(inputs=inputs, tep=tep))


@pytest.fixture(scope="module")
def phased():
    inputs, tep, phasing = _losing_scenario()
    return core.calculate_phased(core.PhasedCalcRequest(
        inputs=inputs, tep=tep, phasing=phasing))


def test_the_payment_really_is_deferred(single) -> None:
    """Предохранитель: платится позже, чем начисляется.

    Заплати проект проценты в месяце начисления — обе меры сойдутся, и всё
    ниже проверяло бы одно и то же дважды.
    """
    rows = single["finance"]["rows"]
    deferred = [row for row in rows
                if sum(float(row.get(key) or 0.0) for key in ACCRUAL_KEYS) > 1.0
                and float(row.get("interest_payment") or 0.0) == 0.0]
    assert len(deferred) > 12, (
        "проценты платятся в месяце начисления — проверять нечего")


def test_the_deduction_stands_in_the_month_it_accrued(single) -> None:
    """Вычет месяца равен начисленному этого месяца, а не уплаченному.

    Комиссии выдачи стоят своими датами и в месячные строки не входят —
    начало проекта и РнС их и несут.
    """
    finance = single["finance"]
    fees = {
        core.d(single["dates"]["project_start"]): float(finance["bridge_fee"]),
        core.d(single["dates"]["permit"]): float(finance["pf_reservation_fee"]),
    }
    for row in finance["rows"]:
        month = core.d(row["month"])
        accrued = sum(float(row.get(key) or 0.0) for key in ACCRUAL_KEYS)
        accrued += fees.get(month, 0.0)
        assert float(row.get("financing_tax_deduction") or 0.0) == pytest.approx(
            accrued, abs=1.0), f"месяц {row['month']}"


def test_the_deductions_add_up_to_the_cost_of_financing(single) -> None:
    """Сумма вычетов — ровно стоимость финансирования, без строки сверки.

    Прежде остаток (проценты БРИДЖа, оплаченные капиталом на РнС) доезжал до
    конца горизонта строкой сверки: она есть и сейчас, но при начислении ей
    нечего подбирать.
    """
    finance = single["finance"]
    assert finance["financing_tax_deductions"] == pytest.approx(
        single["summary"]["financing_cost"], rel=1e-9)
    assert abs(finance["financing_tax_reconciliation"]) < 1.0


def test_the_two_measures_disagree_on_a_losing_queue(phased) -> None:
    """Предохранитель и сама находка разом: меры обязаны разойтись.

    Считаются обе одной и той же функцией движка на одних и тех же строках —
    вторая реализация налога разошлась бы с первой молча.
    """
    first = phased["comparison"][0]
    assert first["net_profit"] < 0, "первая очередь должна быть убыточной"

    finance = phased["consolidated"]["finance"]
    months = [core.d(row["month"]) for row in finance["rows"]]
    margins = {core.d(row["month"]): float(row.get("taxable_margin") or 0.0)
               for row in finance["rows"]}
    rate = core._phase_tax_rate([phase["result"] for phase in phased["phases"]])
    gate = finance.get("first_taxable_month")
    assert rate > 0 and gate, "без ставки и гейта мерить нечего"

    def tax(by_month: dict) -> float:
        schedule, _ = core._profit_tax_schedule(months, margins, by_month, gate, rate)
        return sum(schedule.values())

    accrued: dict = defaultdict(float)
    paid: dict = defaultdict(float)
    for phase in phased["phases"]:
        for row in phase["result"]["finance"]["rows"]:
            month = core.d(row["month"])
            accrued[month] += sum(float(row.get(key) or 0.0) for key in ACCRUAL_KEYS)
            paid[month] += float(row.get("interest_payment") or 0.0)

    by_accrual, by_payment = tax(accrued), tax(paid)
    assert abs(by_accrual - by_payment) > 1e7, (
        f"меры совпали ({by_accrual/1e6:,.1f} и {by_payment/1e6:,.1f} млн) — "
        "проект их не различает, тест бесполезен")
    assert phased["consolidated"]["summary"]["profit_tax"] == pytest.approx(
        by_accrual, rel=1e-6), "свод обязан считать по начислению"


def test_accrual_in_the_base_is_not_capitalization_in_the_debt(single) -> None:
    """Начисление в базе и капитализация в долге — разные вещи.

    Отсроченные проценты копятся отдельным обязательством и лимит ПФ НЕ
    выбирают (решение владельца 04.08.2026). Признание расхода в базе этого
    не трогает: тело долга остаётся телом, обязательство — телом плюс
    начисленное.
    """
    limit = float(single["finance"].get("pf_limit_calculated")
                  or single["finance"].get("calculated_pf_limit") or 0.0)
    seen_payable = False
    for row in single["finance"]["rows"]:
        body = float(row.get("pf_balance") or 0.0)
        payable = float(row.get("pf_payable") or 0.0)
        if payable > 1.0:
            seen_payable = True
            assert float(row.get("pf_obligation") or 0.0) == pytest.approx(
                body + payable, abs=1.0), f"месяц {row['month']}"
        if limit:
            assert body <= limit + 1.0, (
                f"тело выбрало больше лимита в {row['month']} — "
                "начисленное в лимит попадать не должно")
    assert seen_payable, "отсроченных процентов нет — проверять нечего"


def test_the_workbook_recognizes_interest_in_the_same_month() -> None:
    """Книга и движок — одной правкой.

    Строка 53 остаётся кассовой (её читают вклад капитала, распределение и
    проверка фондирования), базу налога вычитает новая строка 58. За горизонт
    начисленное обязано сойтись с уплаченным: расходись они, книга признала бы
    расход, которого проект не несёт.
    """
    sys.setrecursionlimit(400000)
    inputs, tep, phasing = _losing_scenario()
    engine = core.calculate_phased(core.PhasedCalcRequest(
        inputs=inputs, tep=tep, phasing=phasing))
    content, _, meta = core.build_project_workbook(
        inputs, tep, [], phasing, project_name="Начисление процентов")
    assert meta["missing"] == []
    evaluator = Evaluator(openpyxl.load_workbook(io.BytesIO(content)))

    book_tax = evaluator.cell("КОНСОЛИДАТОР", "K8")
    engine_tax = engine["consolidated"]["summary"]["profit_tax"] / 1e6
    assert book_tax == pytest.approx(engine_tax, rel=0.005), (
        f"книга {book_tax:,.1f} против движка {engine_tax:,.1f}")

    for phase in (1, 2):
        accrued = evaluator.cell(f"CF_{phase}", "B58")
        paid = evaluator.cell(f"CF_{phase}", "B53")
        assert accrued > 0, f"CF_{phase}: строка начисления пуста"
        assert accrued == pytest.approx(paid, rel=0.005), (
            f"CF_{phase}: начислено {accrued:,.1f}, уплачено {paid:,.1f}")


def test_the_cash_rows_still_read_the_payment() -> None:
    """Вклад капитала, распределение и проверка фондирования — про кассу.

    Переведи их на начисление — и книга показала бы отток там, где денег ещё
    не платили.
    """
    import re
    source = ROOT.joinpath("main_legacy.py").read_text(encoding="utf-8")
    at = source.index("def _v4_apply_interest_accrual(")
    end = source.index("\ndef ", at + 1)
    body = source[at:end]
    assert re.search(r'-\{column\}53-', body), (
        "правка обязана искать хвост выплаты в строке базы налога")
    assert "49" not in re.findall(r'\{column\}(\d+)', body), \
        "строки кассы правке не подлежат"
