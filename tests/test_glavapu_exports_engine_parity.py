"""Все выгрузки ГлавАПУ из репозитория против движка без ГлавАПУ.

Когда выгрузки нет (калькулятор недоступен или правится ТЭП), движок считает
сам, и его ответ обязан совпадать с тем, что отдал бы калькулятор на тех же
вводных. Нагатино 06.10 (умолчания) сверяется в
`tests/test_nagatino_engine_vs_glavapu.py`; здесь — остальные выгрузки из
`tests/fixtures/glavapu_exports_2026.json`: Савёловский, Можайский,
Головинский (16.08.2026), Хамовники 77:01 (20.08.2026), книга владельца с нашим
сценарием (Нагатино, 3 897 жителей), Донской (прод 07.10.2026) и стендовая копия
калькулятора (14,62 га).

Две наши стороны:

* **норма** — `tep_derived_norms` (страница: `/tep/derived-by-site`, пересчёт
  правок; Telegram: свободный ТЭП через те же зону и 33 м²);
* **дубль** — `vri_tep_quick("msk")`, серверный ТЭП в формате калькулятора
  (кнопка бота, фолбэк сайта без браузера). Дубль воспроизводит только
  умолчания калькулятора (100/0, рекомендуемая плотность), поэтому книги со
  сценарием через него не сверяются.

Расхождение не прячется допуском: строка совпадает либо записана в `*_KNOWN`
с причиной; ушедшее расхождение без записи — тоже падение.

Запуск: python3 -m pytest tests/test_glavapu_exports_engine_parity.py -q
"""

from __future__ import annotations

import io
import json
import math
import re
import sys
from pathlib import Path

import pytest

ROOT = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(ROOT))

import glavapu_scenario as gs  # noqa: E402
import main_legacy as core  # noqa: E402

EXPORTS = json.loads((ROOT / "tests" / "fixtures" / "glavapu_exports_2026.json")
                     .read_text(encoding="utf-8"))["exports"]

# Допуск только на копейку компенсации (калькулятор округляет слагаемые до
# тысячных): 188,414 против 188,415 — одно и то же число.
MONEY_TOL = 0.0015

# ------------------------------------------------------------------ норма --

NORM_KEYS = ("population", "apartment_units", "kindergarten_places", "school_places",
             "clinic_capacity", "parking_permanent", "parking_guest", "parking_onsite", "jobs",
             "compensation_kindergarten_mln", "compensation_school_mln", "compensation_clinic_mln")

NORM_KNOWN = {
    ("vendored_stand_14_62", "parking_permanent"): (
        "стендовая копия калькулятора отстала: 820 = ⌈140 218 × 0,257 / 33 × 0,75⌉ — прежняя "
        "строка; живой калькулятор и наша норма считают по 2118-ПП (1 619–1 620)"),
    ("vendored_stand_14_62", "parking_guest"): "десятая часть постоянных — следствие строки выше",
    ("nagatino_owner_scenario_2026_10", "compensation_kindergarten_mln"): (
        "объект поставлен сверх потребности (250 мест при 172): калькулятор кредитует профицит "
        "со знаком минус без коэффициента 1,2; норма считает компенсацию только за дефицит"),
    ("nagatino_owner_scenario_2026_10", "compensation_school_mln"): (
        "то же: школа 2 500 мест при 351 — профицит 2 149 мест, земля 14 м²/место"),
    ("vendored_stand_14_62", "compensation_kindergarten_mln"): (
        "профицит 163 мест (350 при 187): минус без 1,2, земля 32 м²/место"),
    ("vendored_stand_14_62", "compensation_school_mln"): (
        "профицит 617 мест (1 000 при 383): минус без 1,2, земля 19 м²/место"),
}


def _norm(key: str) -> dict:
    e = EXPORTS[key]
    inputs, analysis = e["inputs"], e["analysis"]
    got = core.tep_derived_norms(
        apartment_area_sqm=inputs["apartment_area_sqm"] or 0.0,
        residential_living_spp_sqm=inputs["residential_living_spp_sqm"] or 0.0,
        nonresidential_np_sqm=inputs["built_in_np_sqm"] or 0.0,
        k1=analysis["k1"], k2=analysis["k2"],
        zone_two=core.district_zone_two(analysis["district"]),
        upks_rub=analysis["upks"] or 0.0,
        apartment_count=inputs["apartment_count"] or 0.0)
    out = {k: got.get(k) for k in NORM_KEYS if k in got}
    for kind, value in got["compensation_breakdown_mln"].items():
        out[f"compensation_{kind}_mln"] = value
    return out


def _norm_table(key: str) -> list[dict]:
    ours = _norm(key)
    rows = []
    for kind, theirs in EXPORTS[key]["calculator"].items():
        if kind not in NORM_KEYS or theirs is None:
            continue
        mine = ours.get(kind)
        tol = MONEY_TOL if kind.startswith("compensation") else 0.0
        same = mine is not None and abs(float(mine) - float(theirs)) <= tol
        rows.append({"kind": kind, "ours": mine, "glavapu": theirs, "same": same})
    return rows


NORM_EXPORTS = [k for k, e in EXPORTS.items()
                if e["inputs"] and e["inputs"].get("apartment_area_sqm")]


@pytest.mark.parametrize("key", NORM_EXPORTS)
def test_every_norm_difference_is_named_with_its_reason(key: str) -> None:
    bad = [r for r in _norm_table(key) if not r["same"] and (key, r["kind"]) not in NORM_KNOWN]
    assert not bad, f"{EXPORTS[key]['title']}: норма расходится без причины: " + "; ".join(
        f"{r['kind']}: наше {r['ours']} против {r['glavapu']}" for r in bad)


@pytest.mark.parametrize("key", NORM_EXPORTS)
def test_a_closed_norm_difference_leaves_the_list(key: str) -> None:
    same = {r["kind"] for r in _norm_table(key) if r["same"]}
    stale = {kind for (k, kind) in NORM_KNOWN if k == key} & same
    assert not stale, sorted(stale)


def test_the_norm_agrees_on_population_places_and_parking() -> None:
    """Сверка, которая молчит всегда, ничего не доказывает: где методики
    совпадают, движок обязан дать число города — на каждой выгрузке."""
    for key in NORM_EXPORTS:
        rows = {r["kind"]: r for r in _norm_table(key)}
        for kind in ("population", "kindergarten_places", "school_places", "clinic_capacity",
                     "parking_onsite", "compensation_clinic_mln"):
            if kind in rows:
                assert rows[kind]["same"], (key, rows[kind])
        if (key, "parking_permanent") not in NORM_KNOWN:
            assert rows["parking_permanent"]["same"] and rows["parking_guest"]["same"], key


# ------------------------------------------------------------------ дубль --

MIRROR_ROW_OF = {
    "population": "4", "apartment_units": "5", "kindergarten_places": "30", "school_places": "31",
    "clinic_capacity": "32", "clinic_adult": "33", "clinic_child": "34",
    "parking_permanent": "42.1", "parking_guest": "42.2", "parking_onsite": "42.3",
    "parking_short_stop": "43", "change_vri_mln": "44",
    "compensation_kindergarten_mln": "54", "compensation_school_mln": "55",
    "compensation_clinic_mln": "56",
}

MIRROR_KNOWN = {
    ("mozhaysky_2026_08_16", "change_vri_mln"): (
        "851,602 против 851,315 — 0,03 %: площадь участка в анализе округлена до 0,581 га, "
        "калькулятор считал от точной"),
}

# Дубль воспроизводит умолчания калькулятора: нужны площадь, плотность и ставки
# квартала. Книги со сценарием владельца и стендовая копия сюда не попадают.
MIRROR_EXPORTS = [k for k, e in EXPORTS.items()
                  if e["analysis"].get("area_ha") and e["analysis"].get("rent")
                  and e["analysis"].get("density_th_sqm_ha") and e["inputs"] is None
                  or k == "savyolovsky_2026_08_16"]


def _mirror(key: str) -> dict[str, float | None]:
    a = EXPORTS[key]["analysis"]
    k2_key = "business_inside_ttc" if a["inside_ttc"] else "business_outside_ttc"
    analysis = {
        "territory": {"area_ha": a["area_ha"], "inside_ttc": a["inside_ttc"],
                      "district": a["district"]},
        "coefficients": {"rail": a["k1"], k2_key: a["k2"], "rent": a["rent"],
                         "base_cost_zh_high": a["base_cost_mkd"], "upks_zh_high": a["upks"] or 0.0},
    }
    result = core.vri_tep_quick("msk", "77:00:0000000:1", analysis=analysis)
    import openpyxl
    book = openpyxl.load_workbook(io.BytesIO(result["file"]))
    rows: dict[str, float | None] = {}
    for row in book["ТЭП"].iter_rows(min_row=2, values_only=True):
        if row[0] not in (None, ""):
            rows[str(row[0]).strip()] = gs.ru_number(row[3])
    rows["jobs"] = gs.ru_number(book["МПТ"].cell(row=2, column=4).value)
    return rows


def _mirror_table(key: str) -> list[dict]:
    rows = _mirror(key)
    out = []
    for kind, theirs in EXPORTS[key]["calculator"].items():
        code = "jobs" if kind == "jobs" else MIRROR_ROW_OF.get(kind)
        if code is None or theirs is None:
            continue
        mine = rows.get(code)
        tol = MONEY_TOL if kind.startswith("compensation") or kind == "change_vri_mln" else 0.0
        same = mine is not None and abs(float(mine) - float(theirs)) <= tol
        out.append({"kind": kind, "row": code, "ours": mine, "glavapu": theirs, "same": same})
    return out


@pytest.mark.parametrize("key", MIRROR_EXPORTS)
def test_every_mirror_difference_is_named_with_its_reason(key: str) -> None:
    bad = [r for r in _mirror_table(key) if not r["same"] and (key, r["kind"]) not in MIRROR_KNOWN]
    assert not bad, f"{EXPORTS[key]['title']}: дубль расходится без причины: " + "; ".join(
        f"строка {r['row']} {r['kind']}: наше {r['ours']} против {r['glavapu']}" for r in bad)


@pytest.mark.parametrize("key", MIRROR_EXPORTS)
def test_a_closed_mirror_difference_leaves_the_list(key: str) -> None:
    same = {r["kind"] for r in _mirror_table(key) if r["same"]}
    stale = {kind for (k, kind) in MIRROR_KNOWN if k == key} & same
    assert not stale, sorted(stale)


def test_the_mirror_reproduces_the_city_rows_on_every_default_export() -> None:
    assert set(MIRROR_EXPORTS) == {"savyolovsky_2026_08_16", "mozhaysky_2026_08_16",
                                   "golovinsky_2026_08_16"}
    for key in MIRROR_EXPORTS:
        rows = {r["kind"]: r for r in _mirror_table(key)}
        for kind in ("population", "clinic_capacity", "clinic_adult", "clinic_child"):
            assert rows[kind]["same"], (key, rows[kind])


# -------------------------------------------- постоянные места по квартирам --

def test_donskoy_permanent_places_are_point_two_by_the_flat_count() -> None:
    """Донской: 55 квартир без состава — калькулятор берёт пункт 2 по числу
    квартир (все до 70 м²): ⌈55 × 0,8⌉ = 44, гостевые ⌈4,4⌉ = 5. Пункт 1 от
    площади дал бы другое число, и тест на нём падает."""
    c = EXPORTS["donskoy_2026_10_07"]["calculator"]
    mix = EXPORTS["donskoy_2026_10_07"]["inputs"]["flats_by_size"]
    assert core.moscow_permanent_parking_by_mix(mix["small"], mix["medium"], mix["large"]) \
        == c["parking_permanent"] == 44
    permanent, basis = core.moscow_permanent_parking_by_average(3400.0, c["apartment_units"])
    assert permanent == 44 and "пункт 2" in basis
    assert math.ceil(permanent / 10.0) == c["parking_guest"] == 5
    assert core.moscow_permanent_parking_2118(3400.0) != 44


def test_the_owner_book_flats_by_size_give_the_city_permanent_places() -> None:
    """Книга владельца несёт состав квартир (5.1–5.3): 1 856 до 70 м² → 1 485.
    Пункт 1 от площади квартир (3 897 × 33) даёт то же — методики сходятся."""
    e = EXPORTS["nagatino_owner_scenario_2026_10"]
    mix = e["inputs"]["flats_by_size"]
    assert core.moscow_permanent_parking_by_mix(mix["small"], mix["medium"], mix["large"]) == 1485
    assert core.moscow_permanent_parking_2118(e["inputs"]["apartment_area_sqm"]) == 1485
    assert e["calculator"]["parking_permanent"] == 1485


# ----------------------------------------------------- зона района и нормы --

ASSETS_JS = sorted((ROOT / "genplan_assets").glob("index-*.js"))


@pytest.mark.skipif(not ASSETS_JS, reason="локальной копии калькулятора нет")
def test_our_zone_two_districts_are_the_calculators() -> None:
    """Единственный коэффициент района в соцпотребности калькулятора — зона
    нормирования: 28 районов зоны 2 (63/124), остальные и неизвестный район —
    44/90. Наш список обязан быть тем же до района."""
    text = ASSETS_JS[0].read_text(encoding="utf-8")
    two = re.search(r"\{zone:2,needPlacesDOO:63,needPlacesSchool:124,districts:`([^`]*)`", text)
    one = re.search(r"\{zone:1,needPlacesDOO:44,needPlacesSchool:90,districts:`([^`]*)`", text)
    assert two and one, "в копии калькулятора нет таблицы зон"
    theirs = {d.strip().lower().replace("ё", "е") for d in two.group(1).split(".")}
    assert theirs == set(core.MOSCOW_ZONE_TWO_DISTRICTS)
    assert len(theirs) == 28
    assert all(core.district_zone_two(d) for d in two.group(1).split("."))
    assert not any(core.district_zone_two(d) for d in one.group(1).split("."))
    # Неизвестный район — зона 1 у обоих (функция cp калькулятора: t ? … : 44/90).
    assert core.district_zone_two("") is False and core.district_zone_two("Нет такого") is False
    assert core.SOCIAL_NORMS_PER_1000 == {"kindergarten": {"1": 44, "2": 63},
                                          "school": {"1": 90, "2": 124},
                                          "clinic": {"1": 19, "2": 19}}
    # Нормативов ОВЗ, повышающих или понижающих коэффициентов района в коде
    # калькулятора нет: только зона и три нормы поликлиники.
    assert "normPoliclinicAdult:13.2,normPoliclinicChild:5.8,normPoliclinicMixed:19" in text
    assert "needPlacesOVZ" not in text and "ОВЗ" not in text


def test_a_zone_two_district_raises_the_places_the_way_the_calculator_would() -> None:
    """Контрпример: тот же Савёловский ТЭП в Солнцеве (зона 2) — 63/124 на тысячу."""
    e = EXPORTS["savyolovsky_2026_08_16"]
    base = _norm("savyolovsky_2026_08_16")
    two = core.tep_derived_norms(apartment_area_sqm=e["inputs"]["apartment_area_sqm"],
                                 residential_living_spp_sqm=e["inputs"]["residential_living_spp_sqm"],
                                 zone_two=core.district_zone_two("Солнцево"))
    assert (base["kindergarten_places"], base["school_places"]) == (19, 38)
    assert (two["kindergarten_places"], two["school_places"]) == (
        math.ceil(63 * 422 / 1000), math.ceil(124 * 422 / 1000)) == (27, 53)
    assert two["clinic_capacity"] == base["clinic_capacity"] == 9


def test_the_telegram_free_text_tep_counts_population_and_places_like_the_city() -> None:
    """Свободный ТЭП бота: население по 33 м² и места по зоне района — те же
    числа, что у калькулятора. Квартиры там рыночные (60 м², решение владельца),
    паркинг — прежняя строка (PR #710), поэтому они здесь не сверяются."""
    for key in ("savyolovsky_2026_08_16", "khamovniki_2026_08_20"):
        e = EXPORTS[key]
        summary = core.build_freeform_tep("", {
            "apartments_saleable_sqm": e["inputs"]["apartment_area_sqm"],
            "site_area_ha": e["analysis"]["area_ha"] or 1.0,
            "district": e["analysis"]["district"],
            "commercial_gns_sqm": e["inputs"]["built_in_spp_sqm"]})["summary"]
        c = e["calculator"]
        assert summary["population"] == c["population"], key
        assert (summary["required_kindergarten_places"], summary["required_school_places"],
                summary["required_clinic_capacity"]) == (
            c["kindergarten_places"], c["school_places"], c["clinic_capacity"]), key


# ------------------------------------------------- типовые объекты города --

def test_the_typical_social_objects_use_our_area_per_place() -> None:
    """Калькулятор ставит типовой объект по ступени ёмкости: ДОО 350 мест —
    16 м² НП на место (5 600 м²), школа 1 000 — 15 (15 000). Наша ступень — та же."""
    objects = EXPORTS["vendored_stand_14_62"]["calculator"]["social_objects"]
    for kind, obj in objects.items():
        per_place = obj["np_th"] * 1000 / obj["places"]
        assert core.moscow_social_area_per_place(kind, obj["places"]) == per_place
