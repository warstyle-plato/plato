"""ГПР-факт по пакетам: та же кнопка «ГПР-факт» Монитора, укрупнённый файл.

Владелец, 29.10.2026, о «ГПР до ввода» (около 40 пакетов с кодом РСС и колонкой
«Факт 20.10.2026»): «можно ли использовать этот ГПР относительно факта
выполнения?», «не хотелось бы плодить доп формы для загрузки», «давай
попробуем научить». Раньше такой файл ложился снимком с нулём работ — молча, и
проценты в Мониторе и сводке не менялись. Теперь процент пакета разносится на
работы утверждённого ГПР по коду сметы (самый длинный совпавший код, при равных
— по названию), пакеты без кода — по названию; что не сопоставилось, названо.

Запуск: python3 -m pytest tests/test_gpr_fact_by_packages.py -q
"""
from __future__ import annotations

import datetime
import io
import sys
from pathlib import Path

import pytest
from openpyxl import Workbook

ROOT = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(ROOT))

import developaid_gpr_packages as packages  # noqa: E402
import developaid_management_report as report  # noqa: E402
import developaid_monitor as monitor  # noqa: E402
import developaid_monitor_manager as manager  # noqa: E402

PROJECT = "Пакетный"


@pytest.fixture(autouse=True)
def _isolated(tmp_path, monkeypatch):
    monkeypatch.setattr(monitor, "_SNAPSHOT_DIR", tmp_path / "monitor")
    monkeypatch.setenv("DATA_DIR", str(tmp_path / "data"))


# id, объект, название, код РСС, процент, статус в утверждённом ГПР
WORKS = [
    ("1", "Корпус 1", "Монолит надземной части", "2.2.2.1.1.", 0.30, ""),
    ("2", "Корпус 2", "Монолит надземной части", "2.2.2.1.2.", 1.0, "Завершено"),
    ("3", "Корпус 1", "Кладка стен", "2.2.2.2.1.", 0.10, ""),
    ("4", "Площадка", "Ограждения из металлоконструкций", "", 0.0, ""),
    ("5", "Корпус 1", "Электромонтажные работы ЭОМ", "2.2.3.5.1.", 0.0, ""),
    ("6", "Площадка", "Наружные сети", "2.4.3.1.", 0.05, ""),
    ("7", "Площадка", "Содержание площадки", "2.2.2.8.1.", 0.20, ""),
]


def _baseline() -> bytes:
    book = Workbook()
    gpr = book.active
    gpr.title = "ГПР"
    header = ["ID", "WBS", "Раздел", "Объект", "Наименование работ",
              "Тип строки", "% выполнения", "Начало", "Окончание", "Статус",
              "Длительность, р.д.", "Предшественники", "Связанный тендер",
              "Окончание тендера", "Резерв", "Увязка", "Код РСС",
              "Статья РСС", "Основание привязки"]
    for col, value in enumerate(header, 1):
        gpr.cell(row=4, column=col, value=value)
    for line, (tid, obj, name, code, progress, status) in enumerate(WORKS, 5):
        row = [tid, tid, "СМР", obj, name, "Работа", progress,
               datetime.date(2026, 7, 1), datetime.date(2027, 6, 30), status, 10,
               "", "", "", "", "", code, "", ""]
        for col, value in enumerate(row, 1):
            gpr.cell(row=line, column=col, value=value)
    blob = io.BytesIO()
    book.save(blob)
    return blob.getvalue()


def _packages_book() -> bytes:
    """Как у владельца: заголовок в 3-й строке, строки-группы, № у пакетов."""
    book = Workbook()
    sheet = book.active
    sheet.title = "ГПР"
    sheet["A1"] = "ГПР до ввода в эксплуатацию"
    for col, value in enumerate(["№", "Раздел / пакет", "Код РСС", "Стадия", "Статус",
                                 "Приоритет", "Факт 20.10.2026", "Прогноз 30.07.2027",
                                 "Старт по ГПР", "Финиш по ГПР"], 1):
        sheet.cell(row=3, column=col, value=value)
    rows = [
        ("Этап 1. Проектирование",),
        (1, "Корректировка проекта", "—", "П", "Не начато", "B", 0.0, 1.0),
        ("СМР",),
        (2, "Монолит надземной части", "2.2.2.1.", "СМР", "В работе", "A", 0.9, 1.0),
        (3, "Кладка стен и перегородок", "2.2.2.2.", "СМР", "В работе", "A", 0.55, 1.0),
        (4, "Содержание стройплощадки", "2.2.1.10 + 2.2.2.8", "СМР", "В работе", "B",
         None, None),
        (5, "Лифты", "2.2.3.1.", "СМР", "В работе", "A", 0.0, 1.0),
        (6, "ЭОМ + слаботочные системы", "2.2.3.5.", "СМР", "Не начато", "A", 0.1, 1.0),
        (7, "ПНР индивидуальные", "2.2.3.5.", "СМР", "Не начато", "A", 0.0, 1.0),
        (8, "Наружное теплоснабжение", "2.4.3.", "СМР", "В работе", "A", 0.3, 1.0),
        (9, "Наружная дождевая канализация", "2.4.3.", "СМР", "В работе", "A", 0.6, 1.0),
        (10, "Ограждения (металлоконструкции)", "", "СМР", "Не начато", "B", 0.2, 1.0),
    ]
    for line, row in enumerate(rows, 4):
        for col, value in enumerate(row, 1):
            sheet.cell(row=line, column=col, value=value)
    blob = io.BytesIO()
    book.save(blob)
    return blob.getvalue()


def test_the_package_sheet_is_read_with_its_codes_and_fact_day():
    got = packages.read(_packages_book())
    assert got["fact_day"] == "2026-10-20"
    by_name = {item["name"]: item for item in got["packages"]}
    assert len(by_name) == 10
    assert by_name["Корректировка проекта"]["group"] == "Этап 1. Проектирование"
    assert by_name["Корректировка проекта"]["codes"] == []
    assert by_name["Монолит надземной части"]["codes"] == ["2.2.2.1"]
    assert by_name["Содержание стройплощадки"]["codes"] == ["2.2.1.10", "2.2.2.8"]
    # Пустой «Факт» — нет числа, а не ноль.
    assert by_name["Содержание стройплощадки"]["fact"] is None
    assert by_name["Лифты"]["fact"] == 0.0
    # Подробный ГПР по работам — не пакеты.
    assert packages.read(_baseline()) is None


def test_the_monitor_puts_package_percents_on_the_approved_works():
    monitor.store_schedule(PROJECT, _baseline(), None, "2026-07-01")
    stored = monitor.store_schedule_fact(PROJECT, _packages_book(), "2026-10-20")
    info = stored["packages"]
    assert info["packages"] == 10 and info["fact_day"] == "2026-10-20"
    # Без своих работ в утверждённом ГПР — названы, а не пропали.
    assert info["unmatched"] == ["Корректировка проекта", "Лифты", "ПНР индивидуальные"]
    assert info["no_fact"] == ["Содержание стройплощадки"]
    assert info["ambiguous"] == [{"work": "Наружные сети", "packages": [
        "Наружное теплоснабжение", "Наружная дождевая канализация"]}]
    assert stored["works"] == 5

    status = manager._baseline_status(PROJECT, "2026-10-25")
    progress = {tid: item["progress"] for tid, item in status.items()}
    assert progress == {
        "1": pytest.approx(0.9),   # по коду пакета 2.2.2.1
        "2": pytest.approx(1.0),   # закрытая работа не «раскрывается» до 90%
        "3": pytest.approx(0.55),
        "4": pytest.approx(0.2),   # пакет без кода — по названию
        "5": pytest.approx(0.1),   # общий код 2.2.3.5 — ближе по названию ЭОМ
        "6": pytest.approx(0.05),  # два равных пакета — оставлено из ГПР
        "7": pytest.approx(0.20),  # у пакета нет числа — оставлено из ГПР
    }
    assert status["2"]["closed"] is True
    # До даты снимка — утверждённый ГПР, как и для обычного ГПР-факта.
    assert manager._baseline_status(PROJECT, "2026-10-19")["1"]["progress"] == \
        pytest.approx(0.30)


def test_a_file_without_works_or_packages_is_refused_not_stored_empty():
    monitor.store_schedule(PROJECT, _baseline(), None, "2026-07-01")
    book = Workbook()
    book.active.title = "ГПР"
    book.active["A1"] = "что-то другое"
    blob = io.BytesIO()
    book.save(blob)
    with pytest.raises(ValueError, match="Это не график работ"):
        monitor.store_schedule_fact(PROJECT, blob.getvalue(), "2026-10-20")
    assert monitor.latest_schedule_fact(PROJECT) is None


def test_packages_need_the_approved_schedule_first():
    with pytest.raises(ValueError, match="в проекте «Пакетный» он в Мониторе не загружен"):
        monitor.store_schedule_fact(PROJECT, _packages_book(), "2026-10-20")
    assert monitor.latest_schedule_fact(PROJECT) is None


def test_the_report_says_the_percents_come_from_packages():
    monitor.store_schedule(PROJECT, _baseline(), None, "2026-07-01")
    monitor.store_schedule_fact(PROJECT, _packages_book(), "2026-10-20")
    block = report.physical(PROJECT, "week", "2026-10-25")
    assert block["source"]["label"] == "ГПР-факт по пакетам от 20.10.2026"
    note = report._physical_notes(block)[-1]
    assert note.startswith("График загружен по пакетам работ (факт на 20.10.2026): процент "
                           "пакета поставлен всем работам утверждённого ГПР")
    assert "Сопоставлено пакетов: 4 из 10." in note
    assert ("Не нашлось работ в утверждённом ГПР для пакетов: «Корректировка проекта», "
            "«Лифты», «ПНР индивидуальные»; их процент не учтён.") in note
    assert "1 работа подходит сразу к нескольким пакетам" in note
    assert "Без процента в графике: «Содержание стройплощадки»." in note


def test_the_monitor_button_says_what_the_packages_gave(tmp_path, monkeypatch):
    """Отрисованная страница Монитора: та же кнопка «ГПР-факт», сообщение о пакетах."""
    sys.path.insert(0, str(ROOT / "tests"))
    from browser import chromium_or_skip, serve

    chrome = chromium_or_skip()
    from playwright.sync_api import sync_playwright

    import main_legacy as core

    monkeypatch.setattr(core, "_require_web_access", lambda *_a, **_k: None)
    monkeypatch.setattr(core, "_require_team_access", lambda *_a, **_k: None)
    monitor.store_schedule(PROJECT, _baseline(), None, "2026-07-01")
    upload = tmp_path / "ГПР по пакетам.xlsx"
    upload.write_bytes(_packages_book())
    with serve(core.app, 18806) as base, sync_playwright() as pw:
        browser = pw.chromium.launch(executable_path=str(chrome))
        try:
            tab = browser.new_page()
            tab.add_init_script("localStorage.setItem('developaid_web_session','проба')")
            tab.goto(f"{base}/monitor", wait_until="domcontentloaded")
            tab.wait_for_function("document.getElementById('gprFactBtn')!==null")
            if not tab.is_visible("#gprFactBtn"):
                tab.click("#setupHead")
            tab.fill("#project", "")
            tab.evaluate("document.getElementById('cut').value='2026-10-20'")
            tab.set_input_files("#gprFact", str(upload))
            tab.click("#gprFactBtn")
            empty = tab.text_content("#gprFactMsg")
            tab.fill("#project", PROJECT)
            tab.click("#gprFactBtn")
            tab.wait_for_function(
                "(document.getElementById('gprFactMsg').textContent||'').includes('ГПР-факт')")
            text = tab.text_content("#gprFactMsg")
            # Тот же файл на ту же дату ещё раз: отказ говорит, что уже лежит.
            tab.click("#gprFactBtn")
            tab.wait_for_function(
                "(document.getElementById('gprFactMsg').textContent||'').includes('уже загружен')")
            again = tab.text_content("#gprFactMsg")
        finally:
            browser.close()
    assert again == (
        "ГПР-факт на 20.10.2026 уже загружен: это ГПР по пакетам (пакетов 10, факт на "
        "20.10.2026, с работами утверждённого ГПР сопоставлено 4). Загруженный снимок не "
        "заменяется: уберите его кнопкой «Убрать ГПР-факт на эту дату» или поставьте вверху "
        "страницы другую дату.")
    assert empty == ("Не указан проект: впишите название проекта в поле «Проект» вверху "
                     "страницы.")
    assert text == (
        "ГПР-факт по пакетам загружен: пакетов 10, из них нашли свои работы в утверждённом "
        "ГПР 4; работ с процентом 5, завершено 1. Не нашлось работ для пакетов: "
        "Корректировка проекта, Лифты, ПНР индивидуальные — их процент не учтён.")


def test_an_empty_project_name_is_named_as_the_reason():
    """Владелец, 29.10.2026: поле «Проект» на телефоне было пустым, а отказ
    говорил «baseline не загружен» — файл ушёл в безымянный проект."""
    monitor.store_schedule(PROJECT, _baseline(), None, "2026-07-01")
    with pytest.raises(ValueError, match="Не указан проект"):
        monitor.store_schedule_fact("  ", _packages_book(), "2026-10-20")
    assert not (monitor._SNAPSHOT_DIR / "project").exists()


def _book(sheets: dict[str, list[tuple]]) -> bytes:
    book = Workbook()
    book.remove(book.active)
    for title, rows in sheets.items():
        sheet = book.create_sheet(title)
        for line, row in enumerate(rows, 1):
            for col, value in enumerate(row, 1):
                sheet.cell(row=line, column=col, value=value)
    blob = io.BytesIO()
    book.save(blob)
    return blob.getvalue()


def test_a_wrong_file_is_refused_with_what_was_seen_in_it():
    """Владелец, 29.10.2026, загрузил кнопкой «ГПР-факт» справку по договорам
    подрядчиков (три листа, по договору на лист). Отказ говорил только, чего
    нет; теперь он называет листы файла и куда такие справки загружают."""
    monitor.store_schedule(PROJECT, _baseline(), None, "2026-07-01")
    data = _book({"Подрядчик А": [("Справка по договору № 1",), ("Стоимость", 100)],
                  "Подрядчик Б": [("Справка по договору № 2",)]})
    with pytest.raises(ValueError) as caught:
        monitor.store_schedule_fact(PROJECT, data, "2026-10-20")
    assert str(caught.value) == (
        "Это не график работ: в файле нет ни листа «ГПР» со строками типа «Работа», ни "
        "таблицы пакетов с колонками «Раздел / пакет» и «Факт <дата>». Листы в файле: "
        "«Подрядчик А», «Подрядчик Б». Справки по договорам и спорам загружают в блок "
        "«Споры и претензии» на странице сводки.")


def test_a_package_table_without_the_fact_column_says_so():
    monitor.store_schedule(PROJECT, _baseline(), None, "2026-07-01")
    data = _book({"План": [("№", "Раздел / пакет", "Код РСС", "Прогноз 30.07.2027"),
                           (1, "Монолит надземной части", "2.2.2.1.", 1.0)]})
    with pytest.raises(ValueError, match="На листе «План» есть колонка «Раздел / пакет», но "
                                         "нет колонки «Факт …» с процентами"):
        monitor.store_schedule_fact(PROJECT, data, "2026-10-20")


def test_package_headers_written_a_little_differently_are_read():
    """«Раздел/пакет» без пробелов и «Факт на 20.10.2026» — те же колонки."""
    data = _book({"Пакеты": [("Укрупнённый график",), (),
                             ("№", "Раздел/пакет", "Код РСС", "Статус", "Факт на 20.10.2026"),
                             (1, "Монолит надземной части", "2.2.2.1.", "В работе", 0.9)]})
    got = packages.read(data)
    assert got["fact_day"] == "2026-10-20"
    assert [(p["name"], p["codes"], p["fact"]) for p in got["packages"]] == [
        ("Монолит надземной части", ["2.2.2.1"], 0.9)]


def test_a_second_upload_on_the_same_day_says_what_is_already_there():
    """Владелец, 30.10.2026: «ГПР-факт на эту дату уже есть — открываю» не
    говорило, какой файл уже лежит, — тот же по пакетам или другой."""
    monitor.store_schedule(PROJECT, _baseline(), None, "2026-07-01")
    monitor.store_schedule_fact(PROJECT, _baseline(), "2026-10-20")
    with pytest.raises(FileExistsError) as caught:
        monitor.store_schedule_fact(PROJECT, _packages_book(), "2026-10-20")
    assert str(caught.value).startswith(
        "ГПР-факт на 20.10.2026 уже загружен: в нём работ ГПР — 7. Загруженный снимок не "
        "заменяется")


def test_a_loaded_fact_can_be_withdrawn_to_the_archive():
    """Владелец, 30.10.2026: ГПР по пакетам «не подходит» — убрать из расчёта,
    но не стереть: файл уходит в архив с записью, кто убрал."""
    monitor.store_schedule(PROJECT, _baseline(), None, "2026-07-01")
    monitor.store_schedule_fact(PROJECT, _packages_book(), "2026-10-20")
    assert manager._baseline_status(PROJECT, "2026-10-25")["1"]["progress"] == pytest.approx(0.9)
    got = monitor.withdraw_schedule_fact(PROJECT, "2026-10-20", "Иванов")
    assert got == {"taken_at": "2026-10-20", "withdrawn": True}
    assert monitor.latest_schedule_fact(PROJECT) is None
    assert manager._baseline_status(PROJECT, "2026-10-25")["1"]["progress"] == pytest.approx(0.3)
    archive = monitor._project_dir(PROJECT) / "schedule_fact" / "withdrawn"
    assert len(list(archive.glob("2026-10-20-*.xlsx"))) == 1
    assert "Иванов" in next(archive.glob("2026-10-20-*.json")).read_text(encoding="utf-8")
    # Теперь на ту же дату можно загрузить другой файл.
    monitor.store_schedule_fact(PROJECT, _baseline(), "2026-10-20")
    with pytest.raises(FileNotFoundError, match="ГПР-факта на 21.10.2026 в проекте «Пакетный» "
                                                "нет. Загружены: 20.10.2026."):
        monitor.withdraw_schedule_fact(PROJECT, "2026-10-21")


def test_the_weekly_blank_is_the_approved_schedule_grouped_by_object():
    """Бланк из утверждённого ГПР: корпуса по номеру, затем общие объекты; ID те
    же, процент предзаполнен текущим — загруженный обратно, он меняет только
    поправленные работы."""
    import developaid_gpr_template as template
    from openpyxl import load_workbook

    monitor.store_schedule(PROJECT, _baseline(), None, "2026-07-01")
    data = template.build(PROJECT, "2026-10-25")
    sheet = load_workbook(io.BytesIO(data))["ГПР"]
    rows = [[cell.value for cell in row] for row in sheet.iter_rows(min_row=5)]
    groups = [row[4] for row in rows if row[0] is None]
    assert groups == ["Корпус 1", "Корпус 2", "Площадка"]
    by_id = {row[0]: row for row in rows if row[0] is not None}
    assert sorted(by_id) == ["1", "2", "3", "4", "5", "6", "7"]
    assert by_id["1"][6] == pytest.approx(0.30) and by_id["2"][9] == "Завершено"
    assert 0 < by_id["1"][19] < 1  # план к дате
    # Читается тем же разборщиком, что и утверждённый ГПР: все семь работ.
    assert len(monitor._read_baseline_gpr_bytes(data)["works"]) == 7

    # Поправили одну работу — меняется только она.
    book = load_workbook(io.BytesIO(data))
    line = next(row[0].row for row in book["ГПР"].iter_rows(min_row=5) if row[0].value == "3")
    book["ГПР"].cell(row=line, column=7, value=0.8)
    blob = io.BytesIO()
    book.save(blob)
    stored = monitor.store_schedule_fact(PROJECT, blob.getvalue(), "2026-10-25")
    assert stored["works"] == 7 and "packages" not in stored
    status = manager._baseline_status(PROJECT, "2026-10-26")
    assert status["3"]["progress"] == pytest.approx(0.8)
    assert {tid: item["progress"] for tid, item in status.items() if tid != "3"} == {
        tid: pytest.approx(progress) for tid, _, _, _, progress, _ in WORKS if tid != "3"}


def test_the_blank_and_withdraw_buttons_on_the_rendered_page(tmp_path, monkeypatch):
    sys.path.insert(0, str(ROOT / "tests"))
    from browser import chromium_or_skip, serve

    chrome = chromium_or_skip()
    from openpyxl import load_workbook
    from playwright.sync_api import sync_playwright

    import main_legacy as core

    monkeypatch.setattr(core, "_require_web_access", lambda *_a, **_k: None)
    monkeypatch.setattr(core, "_require_team_access", lambda *_a, **_k: None)
    monitor.store_schedule(PROJECT, _baseline(), None, "2026-07-01")
    monitor.store_schedule_fact(PROJECT, _packages_book(), "2026-10-20")
    with serve(core.app, 18807) as base, sync_playwright() as pw:
        browser = pw.chromium.launch(executable_path=str(chrome))
        try:
            tab = browser.new_page(accept_downloads=True)
            tab.add_init_script("localStorage.setItem('developaid_web_session','проба')")
            tab.goto(f"{base}/monitor", wait_until="domcontentloaded")
            tab.wait_for_function("document.getElementById('gprTemplateBtn')!==null")
            if not tab.is_visible("#gprTemplateBtn"):
                tab.click("#setupHead")
            tab.fill("#project", PROJECT)
            tab.evaluate("document.getElementById('cut').value='2026-10-20'")
            with tab.expect_download() as caught:
                tab.click("#gprTemplateBtn")
            blank = tmp_path / "blank.xlsx"
            caught.value.save_as(blank)
            tab.on("dialog", lambda dialog: dialog.accept())
            tab.click("#gprWithdrawBtn")
            tab.wait_for_function(
                "(document.getElementById('gprFactMsg').textContent||'').includes('убран')")
            text = tab.text_content("#gprFactMsg")
        finally:
            browser.close()
    assert load_workbook(blank)["ГПР"]["A1"].value == f"ГПР с процентами — {PROJECT}"
    assert text == ("ГПР-факт на 20.10.2026 убран из расчёта, файл сохранён в архиве. "
                    "Пересчитываю.")
    assert monitor.latest_schedule_fact(PROJECT) is None
