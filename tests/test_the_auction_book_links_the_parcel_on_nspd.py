"""Ссылка из выгрузки торгов ведёт на участок в НСПД, а не на общую карту.

НСПД номер из адреса не читает: `?query=<КН>` открывает карту там, где она
была, с чужим выбранным участком (проверено владельцем 27.09.2026). Участок
открывает только адрес с координатами — его собирает движок (`_nspd_map_url`)
после поиска в НСПД. Поиск идёт фоном, выгрузка читает сохранённую точку, а
без точки клетка называет причину.
"""

from __future__ import annotations

import io
import sys
import types

import openpyxl
from fastapi import FastAPI

from auction_search import api
from auction_search import lot_notes as rules

SOON = "2099-01-10T12:00:00+03:00"
ARBAT = "77:01:0001058:2944"
HREF = ("https://nspd.gov.ru/map?thematic=PKK&zoom=17"
        "&coordinate_x=4184314.62&coordinate_y=7508348.33")
NSPD = "Участок на карте НСПД"


def _app(tmp_path, monkeypatch, lookup, *, core_url=""):
    monkeypatch.setenv("DATA_DIR", str(tmp_path))
    monkeypatch.delenv("AUCTIONS_VIEW_KEY", raising=False)
    core = types.ModuleType("developaid_core")
    core._land_lookup_by_numbers = lookup
    core._core_api_url = lambda path: core_url + path if core_url else ""
    core.forwarded = []

    def core_post(url, payload, timeout):
        core.forwarded.append((url, payload))
        return {"results": lookup([payload["query"]])}

    core._core_post = core_post
    monkeypatch.setitem(sys.modules, "developaid_core", core)
    app = FastAPI()
    app.state.market_discovery_service = types.SimpleNamespace(search=None, krt=None)
    api.install(app)
    return app, core


def _book(rows):
    ws = openpyxl.load_workbook(io.BytesIO(api._xlsx(rows, "auctions"))).active
    header = [cell.value for cell in ws[1]]
    return lambda row, name: ws.cell(row, header.index(name) + 1)


def _rows(*cadastres):
    return [{"section": "Торги", "name": f"Лот {index}", "cadastre": cadastre,
             "application_deadline_iso": SOON, "url": f"https://etp/lot/{index}"}
            for index, cadastre in enumerate(cadastres, start=1)]


def test_the_book_links_the_parcel_by_its_coordinates(tmp_path, monkeypatch):
    found = lambda numbers: [{"found": True, "cadastral_number": numbers[0], "map_url": HREF}]
    app, _core = _app(tmp_path, monkeypatch, found)
    store = rules.LotNotes(tmp_path / "market")
    rows = _rows(f"{ARBAT}, 77:01:0001058:1", "")
    store.request(rows)

    get = _book(rows)
    assert get(2, "Кадастровые номера").hyperlink is None, "общая карта — не ссылка на участок"
    assert get(2, NSPD).value.startswith("Точка у НСПД ещё не получена")
    assert get(3, NSPD).value.startswith("Площадка не указала кадастровый номер")

    assert app.state.auction_nspd_points_once() == 2
    get = _book(rows)
    assert get(2, "Кадастровые номера").hyperlink.target == HREF
    assert get(2, NSPD).hyperlink.target == HREF
    assert get(2, NSPD).value == f"{ARBAT} на карте НСПД (первый из 2 номеров)"
    assert get(3, NSPD).hyperlink is None
    assert rules.LotNotes(tmp_path / "market").status()["nspd_points"]["found"] == 2


def test_a_map_without_a_point_is_not_a_link_and_says_why(tmp_path, monkeypatch):
    no_point = lambda numbers: [{"found": True, "cadastral_number": numbers[0],
                                 "map_url": f"https://nspd.gov.ru/map?thematic=PKK&query={numbers[0]}"}]
    app, _core = _app(tmp_path, monkeypatch, no_point)
    rows = _rows(ARBAT)
    rules.LotNotes(tmp_path / "market").request(rows)
    app.state.auction_nspd_points_once()
    get = _book(rows)
    assert get(2, "Кадастровые номера").hyperlink is None
    assert get(2, NSPD).hyperlink is None
    assert get(2, NSPD).value.startswith("Точка не получена: НСПД нашёл номер, но не отдал")


def test_an_nspd_refusal_is_named_and_retried_later(tmp_path, monkeypatch):
    refused = lambda numbers: [{"found": False, "lookup_failed": True,
                                "cadastral_number": numbers[0], "note": "НСПД ответил 503"}]
    app, _core = _app(tmp_path, monkeypatch, refused)
    store = rules.LotNotes(tmp_path / "market")
    rows = _rows(ARBAT)
    store.request(rows)
    app.state.auction_nspd_points_once()
    assert _book(rows)(2, NSPD).value == "Точка не получена: НСПД не ответил: НСПД ответил 503"
    assert store.points_queue() == [], "отказ не переспрашивается каждый проход"
    import time
    assert store.points_queue(now=time.time() + rules.NSPD_FAILED_RETRY_SECONDS + 1) == [ARBAT]


def test_where_nspd_is_closed_the_point_is_asked_from_the_core(tmp_path, monkeypatch):
    found = lambda numbers: [{"found": True, "map_url": HREF}]
    app, core = _app(tmp_path, monkeypatch, found, core_url="https://core.test")
    rows = _rows(ARBAT)
    rules.LotNotes(tmp_path / "market").request(rows)
    app.state.auction_nspd_points_once()
    assert core.forwarded == [("https://core.test/land/lookup", {"query": ARBAT, "limit": 1})]
    assert _book(rows)(2, NSPD).hyperlink.target == HREF


def test_a_number_recovered_from_the_title_is_queued_too(tmp_path, monkeypatch):
    found = lambda numbers: [{"found": True, "map_url": HREF}]
    app, _core = _app(tmp_path, monkeypatch, found)
    rows = [{"section": "Торги", "name": f"Здание, к/н {ARBAT}", "cadastre": "",
             "application_deadline_iso": SOON, "url": "https://etp/lot/1"}]
    rules.LotNotes(tmp_path / "market").request(rows)
    assert app.state.auction_nspd_points_once() == 1
    assert _book(rows)(2, NSPD).hyperlink.target == HREF


def test_the_status_shows_when_the_background_last_ran_and_why_it_stood(tmp_path, monkeypatch):
    app, _core = _app(tmp_path, monkeypatch, lambda numbers: [])
    store = rules.LotNotes(tmp_path / "market")
    assert store.status()["last_run_ago_seconds"] is None
    store.request(_rows(ARBAT))
    assert app.state.auction_lot_note_once() is False
    status = store.status()
    assert status["last_run_ago_seconds"] is not None
    assert status["runs"]["notes"]["result"].startswith("Платон недоступен"), (
        "стоящая очередь называет причину, а не молчит")
    app.state.auction_nspd_points_once()
    status = store.status()
    assert status["nspd_points"]["last_failure"] == "НСПД не нашёл участок по этому номеру"
    assert "спрошено 1, получено 0" in status["runs"]["nspd_points"]["result"]


def test_no_lot_row_of_the_real_export_has_an_empty_comment_or_map_cell(tmp_path, monkeypatch):
    """Маршрут выгрузки целиком: у каждой строки лота — комментарий или причина,
    ссылка на участок или причина. Пустая клетка читалась бы как «пропало»."""
    from fastapi.testclient import TestClient

    app, _core = _app(tmp_path, monkeypatch, lambda numbers: [{"found": True, "map_url": HREF}])
    rows = _rows(ARBAT, "") + [{"section": "Торги", "name": "Без срока", "url": "https://etp/lot/9"}]
    got = TestClient(app).post("/auctions/export.xlsx", json={"rows": rows, "kind": "auctions"})
    assert got.status_code == 200
    ws = openpyxl.load_workbook(io.BytesIO(got.content)).active
    header = [cell.value for cell in ws[1]]
    comment = header.index("Комментарий Платона: чем интересен, чем опасен, что пишут") + 1
    nspd = header.index(NSPD) + 1
    for row in range(2, ws.max_row + 1):
        assert str(ws.cell(row, comment).value or "").strip(), f"строка {row}: пустой комментарий"
        assert str(ws.cell(row, nspd).value or "").strip(), f"строка {row}: пустая клетка НСПД"
