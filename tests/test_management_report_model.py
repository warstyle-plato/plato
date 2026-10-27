"""Управленческая сводка собирается из тех же источников, что мониторинг и продажи.

Проверки:
- физическое выполнение берётся из процентов ГПР-факта, прирост — разница
  снимков на начало и конец периода; без снимка на начало прирост не
  выдаётся нулём;
- бюджет месяца — план из план-факта БДДС, факт — платежи РСС по явной
  карте кодов (стройка — глава 2 без 2.6, коммерция — 3.1, 3.2, 3.4, 3.7);
- продажи — договоры, подписанные в периоде, против плана финмодели и банка;
  договоры вне периода в факт не попадают;
- отсутствие источника названо причиной, а не нулём.

Запуск: python3 -m pytest tests/test_management_report_model.py -q
"""
from __future__ import annotations

import datetime
import io

import pytest
from openpyxl import Workbook

import developaid_management_report as report
import developaid_monitor as monitor
import developaid_monthly_budget as monthly_budget
from market_search import sales_store

M = 1e6
PROJECT = "Гродненская"


@pytest.fixture(autouse=True)
def _isolated(tmp_path, monkeypatch):
    monkeypatch.setattr(monitor, "_SNAPSHOT_DIR", tmp_path / "monitor")
    monkeypatch.setenv("DATA_DIR", str(tmp_path / "data"))


def _gpr_book(progress: dict[str, float], cash: bool = True) -> bytes:
    book = Workbook()
    gpr = book.active
    gpr.title = "ГПР"
    header = ["ID", "WBS", "Раздел", "Объект", "Наименование работ",
              "Тип строки", "% выполнения", "Начало", "Окончание", "Статус",
              "Длительность, р.д.", "Предшественники", "Связанный тендер",
              "Окончание тендера", "Резерв", "Увязка", "Код РСС",
              "Статья РСС", "Основание привязки"]
    for col, value in enumerate(header, 1):
        gpr.cell(row=4, column=col, value=value)
    works = [
        # id, объект, начало, конец: «Монолит» вдвое длиннее «Фасада».
        ("1", "Корпус 1", "Монолит", datetime.date(2026, 7, 1), datetime.date(2026, 12, 28)),
        ("2", "Корпус 2", "Фасад", datetime.date(2026, 9, 1), datetime.date(2026, 11, 29)),
    ]
    for line, (tid, obj, name, start, finish) in enumerate(works, 5):
        row = [tid, tid, "СМР", obj, name, "Работа", progress.get(tid, 0.0),
               start, finish, "", 10, "", "", "", "", "", "2.2.1.4", "", ""]
        for col, value in enumerate(row, 1):
            gpr.cell(row=line, column=col, value=value)
    if cash:
        sheet = book.create_sheet("CF ПЛАН-ФАКТ")
        for i, (month, plan) in enumerate(((datetime.date(2026, 9, 1), 300.0),
                                           (datetime.date(2026, 10, 1), 400.0))):
            col = 2 + i * 4
            sheet.cell(row=2, column=col, value=month)
            sheet.cell(row=4, column=col, value="План")
            sheet.cell(row=4, column=col + 1, value="Факт")
            sheet.cell(row=6, column=col, value=plan)
        sheet.cell(row=6, column=1, value="ИТОГО ПРОЕКТ")
    blob = io.BytesIO()
    book.save(blob)
    return blob.getvalue()


def _rss_book(payments) -> bytes:
    """Платежи: (дата, сумма) — стройка 2.2.1.4, или (дата, сумма, код РСС)."""
    book = Workbook()
    sheet = book.active
    sheet.title = "Расчет стоимости строительства"
    sheet.cell(row=10, column=1, value="2.2.1.4")
    sheet.cell(row=10, column=4, value="Монолит")
    sheet.cell(row=10, column=5, value=1_000 * M)
    sheet.cell(row=12, column=4, value="Всего инвестиционные расходы")
    sheet.cell(row=12, column=5, value=1_000 * M)
    book.create_sheet("Реестр договоров")
    paid = book.create_sheet("Реестр платежей")
    for offset, (day, amount, *code) in enumerate(payments):
        r = 10 + offset
        paid.cell(row=r, column=2, value="СтройКо")
        paid.cell(row=r, column=3, value="Д-1")
        paid.cell(row=r, column=7, value=code[0] if code else "2.2.1.4")
        paid.cell(row=r, column=9, value=day)
        paid.cell(row=r, column=10, value=amount)
    book.create_sheet("Реестр выполненных работ")
    blob = io.BytesIO()
    book.save(blob)
    return blob.getvalue()


def _contract(signed: str, product: str, area: float, amount: float) -> dict:
    return {"month": signed[:7], "signed": signed, "product": product,
            "unit": "", "kind": "", "rooms": "", "area": area, "contract": "",
            "contract_type": "", "state": "", "amount": amount, "units": 1.0,
            "payment_variant": "", "broker": "", "broker_fee": 0.0,
            "broker_rate": None, "sales_bonus_paid": 0.0, "escrow_paid": 0.0,
            "company_buyer": False, "escrow_schedule": []}


def _sales(contracts, bank: bool = True) -> None:
    parts = {
        "contracting": {"project": "ЖК Гродненская", "missing": [], "rows": contracts},
        "fm_plan": {"sheet": "ФМ", "plan": {
            "Итого": {"2026-10": {"amount": 100 * M}, "2026-11": {"amount": 150 * M},
                      "2026-12": {"amount": 50 * M}},
            "Квартира": {"2026-10": {"area": 400.0}, "2026-11": {"area": 600.0}}}},
    }
    if bank:
        parts["bank_plan"] = {"sheet": "Банк",
                              "gross_by_quarter": {"2026 Q3": 900 * M, "2026 Q4": 280 * M},
                              "area_by_quarter": {"2026 Q4": 1_100.0},
                              "price_by_quarter": {"2026 Q4": 255_000.0},
                              "revenue_by_quarter": {}}
    sales_store.save(sales_store.default_dir(), "ЖК Гродненская", parts, "test.xlsx")


def _project(with_before: bool = True) -> None:
    monitor.store_schedule(PROJECT, _gpr_book({}), None, "2026-07-01")
    if with_before:
        monitor.store_schedule_fact(PROJECT, _gpr_book({"1": 0.40, "2": 0.10}, cash=False),
                                    "2026-10-09")
    monitor.store_schedule_fact(PROJECT, _gpr_book({"1": 0.55, "2": 0.40}, cash=False),
                                "2026-10-16")
    monitor.store_estimate(PROJECT, _rss_book([
        (datetime.date(2026, 9, 20), 280 * M),
        (datetime.date(2026, 10, 5), 120 * M),
        (datetime.date(2026, 10, 14), 30 * M),
        (datetime.date(2026, 10, 15), 20 * M, "2.6"),
        (datetime.date(2026, 10, 15), 9 * M, "3.1"),
        (datetime.date(2026, 10, 15), 1 * M, "3.7.2"),
        # Проценты банка и налоги — не коммерческие расходы.
        (datetime.date(2026, 10, 15), 16 * M, "3.5"),
        (datetime.date(2026, 10, 15), 3 * M, "3.3"),
    ]), "2026-10-16")


def test_physical_progress_is_the_gpr_percent_and_its_weekly_growth():
    _project()
    got = report.physical(PROJECT, "week", "2026-10-18")
    assert got["state"] == "ok"
    assert got["snapshot"] == "2026-10-16"
    assert got["snapshot_before"] == "2026-10-09"
    assert got["fresh"] is True
    # Монолит 180 дней, фасад 89: вес по длительности работ.
    w1, w2 = 180.0, 89.0
    fact = (w1 * 0.55 + w2 * 0.40) / (w1 + w2)
    before = (w1 * 0.40 + w2 * 0.10) / (w1 + w2)
    assert got["fact"] == pytest.approx(fact)
    assert got["delta"] == pytest.approx(fact - before)
    plan1 = (datetime.date(2026, 10, 18) - datetime.date(2026, 7, 1)).days / w1
    plan2 = (datetime.date(2026, 10, 18) - datetime.date(2026, 9, 1)).days / w2
    assert got["plan"] == pytest.approx((w1 * plan1 + w2 * plan2) / (w1 + w2))
    assert [o["object"] for o in got["objects"]] == ["Корпус 1", "Корпус 2"]
    assert got["objects"][1]["fact"] == pytest.approx(0.40)


def test_without_a_snapshot_before_the_period_growth_is_unknown_not_zero():
    _project(with_before=False)
    got = report.physical(PROJECT, "week", "2026-10-18")
    assert got["delta"] is None
    text = report.render_text(report.build(PROJECT, "week", "2026-10-18"))
    assert "посчитать не из чего" in text


def test_an_old_snapshot_is_flagged_as_not_refreshed():
    _project()
    got = report.physical(PROJECT, "week", "2026-10-25")
    assert got["fresh"] is False
    assert got["delta"] == pytest.approx(0.0)
    text = report.render_text(report.build(PROJECT, "week", "2026-10-25"))
    assert "новый график с процентами не загружали" in text


def test_without_a_schedule_fact_the_report_shows_what_the_monitor_shows():
    """Нет недельного ГПР-факта — проценты из утверждённого ГПР, как в Мониторе,
    и сводка прямо говорит, откуда они и от какой даты версия графика."""
    monitor.store_schedule(PROJECT, _gpr_book({"1": 0.30}), None, "2026-07-01")
    got = report.physical(PROJECT, "week", "2026-10-18")
    assert got["state"] == "ok" and got["fact_snapshot"] is None
    assert got["source"]["label"] == "утверждённый ГПР от 01.07.2026"
    text = report.render_text(report.build(PROJECT, "week", "2026-10-18"))
    assert "Проценты — утверждённый ГПР от 01.07.2026" in text
    assert "Недельный ГПР-факт с процентами в Монитор не загружали" in text


def budget_book(months=(("Сентябрь 2026", (-300, -10, -20), (-280, -9, -21)),
                        ("Октябрь 2026", (-410, -10, -25), (None, None, None)))) -> bytes:
    """План-факт БДДС в виде выгрузки: шапка месяцев, «ПЛАН»/«ФАКТ», статьи."""
    book = Workbook()
    sheet = book.active
    sheet.cell(row=2, column=1, value="План-фактный анализ_Гродненская_30.09.2026")
    sheet.cell(row=4, column=1, value="Статья")
    sheet.cell(row=5, column=1, value="Контрагент")
    for i, (name, _plan, _fact) in enumerate(months):
        col = 5 + i * 3
        sheet.cell(row=4, column=col, value=name)
        sheet.cell(row=5, column=col, value="ПЛАН (СЗ Гродненская 18)")
        sheet.cell(row=5, column=col + 1, value="ФАКТ")
        sheet.cell(row=5, column=col + 2, value="Отклонение абс. (ФАКТ)")
    total = 5 + len(months) * 3
    sheet.cell(row=4, column=total, value="Итого")
    sheet.cell(row=5, column=total, value="ПЛАН (СЗ Гродненская 18)")
    sheet.cell(row=5, column=total + 1, value="ФАКТ")
    names = ["Стоимость строительства и проектирования",
             "Содержание службы заказчика", "Коммерческие расходы"]
    sheet.cell(row=6, column=1, value="РАСХОДЫ")
    for r, name in enumerate(names, 7):
        sheet.cell(row=r, column=1, value=name)
        for i, (_name, plan, fact) in enumerate(months):
            col = 5 + i * 3
            if plan[r - 7] is not None:
                sheet.cell(row=r, column=col, value=plan[r - 7] * M)
            if fact[r - 7] is not None:
                sheet.cell(row=r, column=col + 1, value=fact[r - 7] * M)
        # Итого — ловушка: план за все месяцы сразу.
        sheet.cell(row=r, column=total, value=-9_999 * M)
    blob = io.BytesIO()
    book.save(blob)
    return blob.getvalue()


def test_month_budget_plan_comes_from_the_budget_and_fact_from_rss():
    _project()
    monthly_budget.store_budget(PROJECT, budget_book(), "2026-10-02", "бюджет.xlsx")
    got = report.budget(PROJECT, "month", "2026-10-31")
    assert got["state"] == "ok"
    assert got["month"] == "2026-10-01"
    assert got["budget_file"] == "2026-10-02"
    lines = {line["key"]: line for line in got["lines"]}
    # Стройка: план 410 без службы заказчика 10; факт — глава 2 без 2.6.
    assert lines["construction"]["plan"] == pytest.approx(400 * M)
    assert lines["construction"]["fact"] == pytest.approx(150 * M)
    assert lines["construction"]["remaining"] == pytest.approx(250 * M)
    assert lines["commercial"]["plan"] == pytest.approx(25 * M)
    assert lines["commercial"]["fact"] == pytest.approx(10 * M)
    # План службы заказчика — её строка бюджета (владелец, 22.10.2026); факт
    # РСС 2.6 шире плана, это объяснено примечанием строки.
    assert lines["service"]["plan"] == pytest.approx(10 * M)
    assert lines["service"]["fact"] == pytest.approx(20 * M)


def test_month_budget_without_budget_file_shows_fact_and_names_the_gap():
    _project()
    got = report.budget(PROJECT, "month", "2026-10-31")
    assert got["state"] == "partial"
    assert all(line["plan"] is None for line in got["lines"])
    assert "Не загружен бюджет" in got["reason"]
    text = report.render_text(report.build(PROJECT, "month", "2026-10-31"))
    assert "Стройка и проектирование за октябрь 2026: оплачено 150,0 млн ₽." in text


def test_a_month_missing_from_the_newest_budget_is_taken_from_an_older_one():
    _project()
    monthly_budget.store_budget(PROJECT, budget_book(), "2026-10-02")
    monthly_budget.store_budget(PROJECT, budget_book(months=(
        ("Ноябрь 2026", (-500, -10, -30), (None, None, None)),)), "2026-10-10")
    got = report.budget(PROJECT, "week", "2026-10-18")
    assert got["budget_file"] == "2026-10-02"
    # Бюджет, загруженный позже конца периода, в сводку периода не идёт.
    monthly_budget.store_budget(PROJECT, budget_book(months=(
        ("Октябрь 2026", (-999, -10, -30), (None, None, None)),)), "2026-10-20")
    assert report.budget(PROJECT, "week", "2026-10-18")["lines"][0]["plan"] == \
        pytest.approx(400 * M)


def test_the_total_column_is_not_the_last_month_plan():
    """У последнего месяца нет «ПЛАН» — план «Итого» за него не сходит."""
    book = Workbook()
    sheet = book.active
    sheet.cell(row=4, column=5, value="Октябрь 2026")
    sheet.cell(row=5, column=5, value="ФАКТ")
    sheet.cell(row=4, column=7, value="Итого")
    sheet.cell(row=5, column=7, value="ПЛАН")
    sheet.cell(row=5, column=8, value="ФАКТ")
    blob = io.BytesIO()
    book.save(blob)
    with pytest.raises(ValueError, match="нет колонок"):
        monthly_budget.read_budget(blob.getvalue())


def test_budget_without_a_needed_row_is_refused_by_name():
    book = Workbook()
    sheet = book.active
    sheet.cell(row=4, column=5, value="Октябрь 2026")
    sheet.cell(row=5, column=5, value="ПЛАН")
    sheet.cell(row=5, column=6, value="ФАКТ")
    sheet.cell(row=7, column=1, value="Коммерческие расходы")
    blob = io.BytesIO()
    book.save(blob)
    with pytest.raises(ValueError, match="Стоимость строительства и проектирования"):
        monthly_budget.read_budget(blob.getvalue())


@pytest.mark.parametrize("code, line", [
    ("2.2.1.4", "construction"), ("2.7", "construction"), ("2.6", "service"),
    ("2.6.1", "service"), ("3.1", "commercial"), ("3.7.2", "commercial"),
    ("3.5", None), ("3.3", None), ("1.6", None), ("3.10", None), ("2", None)])
def test_rss_code_map(code, line):
    assert monthly_budget.line_of(code) == line


def test_sales_take_only_contracts_signed_in_the_period():
    _project()
    _sales([
        _contract("2026-10-13", "Квартира", 50.0, 12.5 * M),
        _contract("2026-10-17", "Квартира", 70.0, 17.5 * M),
        _contract("2026-10-05", "Квартира", 100.0, 25 * M),
        _contract("2026-09-20", "Машиноместо", 0.0, 3 * M),
        # После конца периода — в сводку недели не входит.
        _contract("2026-10-19", "Квартира", 40.0, 10 * M),
    ])
    got = report.sales("ЖК Гродненская", "week", "2026-10-18")
    assert got["state"] == "ok"
    assert got["period"]["contracts"] == 2
    assert got["period"]["amount"] == pytest.approx(30 * M)
    assert got["period"]["area"] == pytest.approx(120.0)
    assert got["period"]["price_flats"] == pytest.approx(250_000.0)
    assert got["month"]["fact"]["contracts"] == 3
    assert got["month"]["fm"] == {"amount": 100 * M, "area": 400.0}
    assert got["quarter"]["key"] == "2026 Q4"
    assert got["quarter"]["fm"]["amount"] == pytest.approx(300 * M)
    assert got["quarter"]["bank"]["amount"] == pytest.approx(280 * M)
    assert got["quarter"]["fact"]["amount"] == pytest.approx(55 * M)


def test_sales_without_bank_plan_say_so():
    _project()
    _sales([_contract("2026-10-13", "Квартира", 50.0, 12.5 * M)], bank=False)
    got = report.sales("ЖК Гродненская", "week", "2026-10-18")
    assert got["quarter"]["bank"] is None
    assert got["quarter"]["bank_reason"] == "План банка не загружен."


def test_unlinked_or_empty_sales_are_named():
    assert report.sales("", "week", "2026-10-18")["state"] == "missing"
    got = report.sales("Нет такого", "week", "2026-10-18")
    assert got["state"] == "missing"
    assert "контрактация" in got["reason"]


def test_one_broken_block_does_not_hide_the_others(monkeypatch):
    _project()

    def broken(*_args, **_kwargs):
        raise ValueError("лист испорчен")

    monkeypatch.setattr(report, "budget", broken)
    got = report.build(PROJECT, "week", "2026-10-18")
    assert got["physical"]["state"] == "ok"
    assert got["budget"]["state"] == "missing"
    assert "лист испорчен" in got["budget"]["reason"]
    assert "лист испорчен" in report.render_text(got)


def test_message_carries_all_three_blocks():
    _project()
    _sales([_contract("2026-10-13", "Квартира", 50.0, 12.5 * M)])
    text = report.render_text(report.build(PROJECT, "month", "2026-10-31",
                                           sales_project="ЖК Гродненская"))
    assert "за октябрь 2026" in text
    assert "Стройка и проектирование за октябрь 2026: оплачено 150,0 млн ₽." in text
    assert "План банка на квартал — 280,0 млн ₽" in text
    assert "Выполнено по графику работ" in text


def test_period_bounds():
    week = report.period_bounds("week", "2026-10-18")
    assert (week["start"], week["before"]) == (datetime.date(2026, 10, 12),
                                               datetime.date(2026, 10, 11))
    month = report.period_bounds("month", "2026-09-30")
    assert month["start"] == datetime.date(2026, 9, 1)
    with pytest.raises(ValueError):
        report.period_bounds("year", "2026-09-30")


def test_budget_route_stores_and_refuses_by_name(monkeypatch):
    import base64

    from fastapi import HTTPException

    import main_legacy as core

    monkeypatch.setattr(core, "_require_web_access", lambda *_a, **_k: None)
    good = core.MonitorEstimateRequest(
        project=PROJECT, taken_at="2026-10-02", filename="бюджет.xlsx",
        content_base64=base64.b64encode(budget_book()).decode())
    stored = core.monitor_store_budget(good)
    assert stored["months"] == ["2026-09", "2026-10"]
    # Тот же период тем же днём — замена, прежний файл в архиве.
    assert core.monitor_store_budget(good)["replaced"] is True
    broken = core.MonitorEstimateRequest(
        project=PROJECT, taken_at="2026-10-03",
        content_base64=base64.b64encode(_rss_book([])).decode())
    with pytest.raises(HTTPException) as bad:
        core.monitor_store_budget(broken)
    assert bad.value.status_code == 400
    assert "шапки с месяцами" in bad.value.detail


def test_several_periods_load_the_same_day_and_the_latest_period_wins():
    # «Эти я загрузить через форму не смог: он не принял новые файлы вообще,
    # сказал что октябрь уже загружен» (владелец, 21.10.2026). Файлы БДДС за
    # прошлые периоды догружаются пачкой в один день.
    _project()
    monthly_budget.store_budget(PROJECT, budget_book(), "2026-10-02")  # сентябрь—октябрь
    for months in ((("Июль 2026", (-100, -10, -30), (-90, -10, -30)),
                    ("Август 2026", (-111, -10, -30), (None, None, None))),
                   (("Август 2026", (-222, -10, -30), (-200, -10, -30)),
                    ("Сентябрь 2026", (-777, -10, -30), (None, None, None)))):
        stored = monthly_budget.store_budget(PROJECT, budget_book(months=months), "2026-10-21")
        assert stored["replaced"] is False
    # Август — из файла за август—сентябрь (период позже), а не за июль—август.
    plan = monthly_budget._plan(PROJECT, "2026-08-01", "2026-10-31")
    assert plan["plan"]["construction"] == pytest.approx(212 * M)
    # Сентябрь — из файла за сентябрь—октябрь, хотя он загружен раньше.
    plan = monthly_budget._plan(PROJECT, "2026-09-01", "2026-10-31")
    assert plan["file"] == "2026-10-02"


def test_past_months_loaded_later_reach_an_earlier_report():
    # «Почему бюджеты, которые я положил с апреля по август, не учитываются»
    # (владелец, 21.10.2026): файлы прошлых месяцев, загруженные 21.10, не
    # попадали в сводку на 20.10 — сводка брала только загруженное к её дате.
    _project()
    for months in ((("Апрель 2026", (-101, -10, -30), (-90, -10, -30)),
                    ("Май 2026", (-102, -10, -30), (-95, -10, -30))),
                   (("Июль 2026", (-104, -10, -30), (-99, -10, -30)),
                    ("Август 2026", (-105, -10, -30), (-98, -10, -30)))):
        monthly_budget.store_budget(PROJECT, budget_book(months=months), "2026-10-21")
    # А будущий период, загруженный позже даты сводки, в неё не попадает.
    monthly_budget.store_budget(PROJECT, budget_book(months=(
        ("Ноябрь 2026", (-500, -10, -30), (None, None, None)),)), "2026-10-21")
    plan = monthly_budget._plan(PROJECT, "2026-04-01", "2026-10-20")
    assert plan["plan"]["construction"] == pytest.approx(91 * M)
    assert monthly_budget._plan(PROJECT, "2026-08-01", "2026-10-20") is not None
    assert monthly_budget._plan(PROJECT, "2026-11-01", "2026-10-20") is None
    # Отклонения за прошлый закрытый месяц сводки на 15.09 — из файла с августом.
    articles = monthly_budget._articles(PROJECT, datetime.date(2026, 9, 15))
    assert articles["known"] or "нет месяца" not in articles["reason"]
