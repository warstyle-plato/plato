"""Пустой месяц основного листа графика — план с листа ближних месяцев.

Владелец, 28.10.2026, о графике от 23.09: «он же с августа реально». На листе
«ГФ_до конца стр» август — пустая колонка и 0 в итоге, а план августа стоит на
листе «ГФ_до конца 2026» того же файла. Выбор владельца: «Со второго листа».

- пустой месяц берётся со второго листа по коду статьи, итог — его ИТОГО;
- месяц, где на основном листе план есть, второй лист не трогает;
- сверка и сводка называют, откуда план месяца.

Запуск: python3 -m pytest tests/test_the_funding_schedule_takes_an_empty_month_from_the_near_sheet.py -q
"""
from __future__ import annotations

import io
import sys
from pathlib import Path

import pytest

ROOT = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(ROOT))
sys.path.insert(0, str(ROOT / "tests"))

import developaid_funding_schedule as fs  # noqa: E402
import developaid_management_report as report  # noqa: E402
import developaid_monitor as monitor  # noqa: E402
import funding_schedule_fixture as fx  # noqa: E402
from test_the_funding_schedule_is_checked_against_payments import (  # noqa: E402
    END, PAYMENTS, PROJECT)

M = fx.M
# Август на втором листе: код, план (млн ₽); сентябрь там иной, чем на основном.
NEAR = [("2.2.2.6.", 146.88, 999.0), ("2.6.", 19.08, 999.0), ("2.2.1.7.", 20.97, 999.0)]


@pytest.fixture(autouse=True)
def _isolated(tmp_path, monkeypatch):
    monkeypatch.setattr(monitor, "_SNAPSHOT_DIR", tmp_path / "monitor")
    monkeypatch.setenv("DATA_DIR", str(tmp_path / "data"))


def _book(august_total=0.0) -> bytes:
    from openpyxl import load_workbook

    wb = load_workbook(io.BytesIO(fx.book()))
    main = wb["ГФ_до конца стр"]
    total_row = next(r for r in range(5, main.max_row + 1) if main.cell(r, 2).value == "ИТОГО")
    main.cell(total_row, 14, august_total)          # август основного листа: 0 в итоге
    near = wb["ГФ_до конца 2026"]
    for c, text in {1: "Код РСС", 2: "Работа", 10: "Сумма договора", 11: "Выплачено"}.items():
        near.cell(2, c, text)
    near["N2"] = "2026 год"
    near.merge_cells("N2:R2")
    for offset, name in enumerate(fx.MONTHS_2026):
        near.cell(3, 14 + offset, name)
    for r, (code, aug, sep) in enumerate(NEAR, start=5):
        near.cell(r, 1, code)
        near.cell(r, 2, "статья")
        near.cell(r, 14, aug * M)
        near.cell(r, 15, sep * M)
    last = 5 + len(NEAR)
    near.cell(last, 2, "ИТОГО")
    near.cell(last, 14, sum(aug for _, aug, _ in NEAR) * M)
    near.cell(last, 15, 999 * M)
    blob = io.BytesIO()
    wb.save(blob)
    return blob.getvalue()


def test_an_empty_month_is_taken_from_the_near_sheet_by_code():
    got = fs.read_schedule(_book())
    assert got["month_sources"] == {"2026-08": "ГФ_до конца 2026"}
    rows = {row["code"]: row for row in got["rows"]}
    assert rows["2.2.2.6"]["plan"]["2026-08"] == pytest.approx(146.88 * M)
    assert rows["2.2.1.7"]["plan"]["2026-08"] == pytest.approx(20.97 * M)
    assert rows["2.2.3.4"]["plan"]["2026-08"] is None
    assert got["total"]["plan"]["2026-08"] == pytest.approx(186.93 * M)
    # Сентябрь на основном листе есть — второй лист его не трогает.
    assert rows["2.2.2.6"]["plan"]["2026-09"] == pytest.approx(123.11 * M)
    assert got["total"]["plan"]["2026-09"] == pytest.approx(231.85 * M)


def test_a_month_with_a_plan_on_the_main_sheet_is_kept():
    got = fs.read_schedule(_book(august_total=5.0 * M))
    assert got["month_sources"] == {}
    assert got["total"]["plan"]["2026-08"] == pytest.approx(5.0 * M)


def test_the_check_and_the_report_name_the_sheet():
    monitor.store_estimate(PROJECT, fx.rss(PAYMENTS), "2026-10-25")
    fs.store(PROJECT, _book(), "2026-08-01", "График финансирования Кутузов Сити 23.09.xlsx")
    got = fs.execution(PROJECT, END)
    august = next(m for m in got["months"] if m["month"] == "2026-08")
    assert not august["no_plan"] and august["plan"] == pytest.approx(186.93 * M)
    assert august["fact"] == pytest.approx(10 * M)
    assert august["plan_sheet"] == "ГФ_до конца 2026"
    page = report._schedule_html(got)
    assert "план месяца — с листа «ГФ_до конца 2026»" in page
