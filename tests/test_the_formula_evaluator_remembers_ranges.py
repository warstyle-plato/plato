"""Вычислитель помнит диапазон целиком, и от этого не меняется ни одно число.

Проверки паритета считают собранную книгу нашим вычислителем, и он же —
единственное место, где книгу вообще можно пересчитать: openpyxl формул не
считает, LibreOffice в образе сломан. Считает он лениво и ячейки кэширует, а
СПИСОК диапазона собирал заново на каждое обращение. Замер на умолчаниях: 29 437
обращений к диапазонам, различных 3 293; `Продажи!D9:GA9` спрашивался 1 620 раз,
и каждый раз это 180 обращений к ячейке.

Стоило это времени всего набора. После протяжки сетки книги со 120 месяцев до
180 (0.22.91) счёт подорожал вместе с ней, и прогон на CI вырос с 74 минут до
108 при том же числе проверок.

Утверждений здесь два, и второе важнее первого: кэш РАБОТАЕТ и кэш НИЧЕГО НЕ
МЕНЯЕТ. Проверка «числа совпали» без первого зелена и на выключенном кэше.

Запуск: python3 -m pytest tests/test_the_formula_evaluator_remembers_ranges.py -q
"""

from __future__ import annotations

import sys
from pathlib import Path

import openpyxl
import pytest

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT))
sys.path.insert(0, str(ROOT / "tests"))

from xlsx_eval import Evaluator, FormulaError  # noqa: E402


def _book():
    """Лист, где один диапазон спрашивают многократно и разными функциями."""
    book = openpyxl.Workbook()
    sheet = book.active
    sheet.title = "Лист"
    for index in range(1, 13):
        sheet.cell(row=1, column=index, value=index * 1.5)
    sheet["A3"] = "=SUM(A1:L1)"
    sheet["A4"] = "=SUM(A1:L1)+SUM(A1:L1)+MAX(A1:L1)+MIN(A1:L1)"
    sheet["A5"] = "=INDEX(A1:L1,1,4)"
    sheet["A6"] = "=A3+A4+A5"
    return book


def _values(evaluator: Evaluator) -> dict[str, object]:
    return {coord: evaluator.cell("Лист", coord) for coord in ("A3", "A4", "A5", "A6")}


def test_the_cache_changes_no_number() -> None:
    """Ответ с памятью о диапазонах — тот же, что без неё."""
    remembered = _values(Evaluator(_book()))

    plain = Evaluator(_book())
    original = Evaluator._range

    def forgetful(self, sheet, start, end):
        self._ranges.clear()
        return original(self, sheet, start, end)

    plain._range = forgetful.__get__(plain)
    assert remembered == _values(plain)


def test_the_cache_is_actually_used() -> None:
    """Диапазон собирается один раз на адрес, а не на каждое обращение.

    Без этого утверждения соседнее ничего не значит: совпадение чисел
    подтверждается и выключенным кэшем.
    """
    evaluator = Evaluator(_book())
    built: list[tuple[str, str, str]] = []
    original = Evaluator._range

    def counted(self, sheet, start, end):
        before = len(self._ranges)
        value = original(self, sheet, start, end)
        if len(self._ranges) > before:
            built.append((sheet, start, end))
        return value

    evaluator._range = counted.__get__(evaluator)
    _values(evaluator)
    assert len(built) == len(set(built)), f"диапазон собран дважды: {built}"
    assert len(built) < 5, f"диапазонов собрано {len(built)} — память не работает"


def test_a_circular_range_still_breaks() -> None:
    """Круг не запоминается как посчитанный диапазон.

    Кэш кладётся только досчитанным: иначе полусобранный список стал бы
    ответом, и книга с круговой ссылкой прошла бы проверку.
    """
    book = openpyxl.Workbook()
    sheet = book.active
    sheet.title = "Лист"
    sheet["A1"] = 1
    sheet["A2"] = "=SUM(A1:A3)"
    sheet["A3"] = "=A2"
    evaluator = Evaluator(book)
    with pytest.raises(FormulaError):
        evaluator.cell("Лист", "A2")
    assert not evaluator._ranges, "недосчитанный диапазон попал в память"
