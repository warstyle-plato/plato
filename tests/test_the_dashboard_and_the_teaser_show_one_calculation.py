"""Дашборд книги и тизер PDF показывают тот же расчёт, что движок.

ТЗ владельца (19.09.2026): одна модель представления на две поверхности —
инвестор читает Excel, банк и партнёр одну страницу. Здесь проверяется вся
цепочка: движок → `presentation_numbers` → модель представления → скрытый
`Dashboard_Data` (через вычислитель формул) → «Дашборд» → тизер (текст PDF).

Сверяются те величины, за которые книга уже отвечает строками паритета
(ПРОВЕРКИ 76–87): Dashboard_Data читает их же колонку B, и второй источник
«где в книге выручка» здесь не заводится. NPV в книге считается СВОЕЙ
формулой (КОНСОЛИДАТОР!O8 — NPV помесячного CF капитала) и с движком на
умолчаниях расходится на 200 млн ₽ — строки паритета у него нет, поэтому здесь
он не сверяется, а назван в задачах.

Предохранитель: проект с нехваткой одобренного лимита и с дефолтом в РВЭ —
иначе риски не проверяются вовсе, и проверка зелена на любом коде.

Запуск: python3 -m pytest tests/test_the_dashboard_and_the_teaser_show_one_calculation.py -q
"""

from __future__ import annotations

import ast
import copy
import io
import re
import sys
import zipfile
from pathlib import Path

import openpyxl
import pytest

sys.path.insert(0, str(Path(__file__).resolve().parent.parent))
sys.path.insert(0, str(Path(__file__).resolve().parent))

from xlsx_eval import Evaluator  # noqa: E402

import main as _wrapper  # noqa: E402
import presentation  # noqa: E402
import v4_dashboard  # noqa: E402

core = _wrapper.core
ROOT = Path(__file__).resolve().parent.parent

# Величины, за которые книга отвечает строками паритета, — и их допуски.
SWORN = ("revenue_mln", "capex_mln", "ebitda_mln", "financing_mln", "tax_mln",
         "net_profit_mln", "llcr", "peak_bridge_mln", "peak_pf_mln", "vat_mln",
         "pf_shortfall_mln", "commercial_mln", "rve_unpaid_mln", "ending_pf_mln",
         "term_months", "margin")


def _tolerance(key: str, target: float, evaluator=None) -> float:
    """Допуск — тот, которым книга сама судит свой паритет.

    У строк ПРОВЕРОК он стоит в колонке E: у производных (налог, чистая
    прибыль) это сумма допусков слагаемых, а не полпроцента от себя —
    второй список допусков здесь разошёлся бы с книгой и краснел бы на её
    же шуме.
    """
    if evaluator is not None and key in v4_dashboard._PARITY:
        return float(evaluator.cell("ПРОВЕРКИ", f"E{v4_dashboard._PARITY[key]}") or 0)
    if key == "llcr":
        return core._V4_PARITY_LLCR_TOLERANCE
    if key == "margin":
        return 0.001
    if key == "term_months":
        return 0.0
    share = core._V4_PARITY_SHARE.get(key, core._V4_PARITY_SHARE_DEFAULT)
    return max(1.0, abs(target) * share)


def _starved() -> tuple[dict, dict, None]:
    """Одобренный лимит меньше потребности: нехватка ПФ и LLCR ниже цели."""
    inputs = copy.deepcopy(core.DEFAULT_INPUTS)
    inputs.update(pf_limit_approved_mln=15000, project_name="Проект дашборда")
    return inputs, copy.deepcopy(core.TEP_DEFAULT), None


def _defaulted() -> tuple[dict, dict, None]:
    """Дефолт в РВЭ: раскрытого эскроу не хватает."""
    inputs = copy.deepcopy(core.DEFAULT_INPUTS)
    inputs.update(project_name="Проект с дефолтом")
    return inputs, copy.deepcopy(core.TEP_DEFAULT), None


def _build(shape):
    inputs, tep, phasing = shape()
    bundle = core._run_authoritative_model(inputs, tep, [], phasing)
    numbers = core.presentation_numbers(bundle["consolidated"])
    model = core.project_presentation(bundle, inputs, tep, phasing)
    sys.setrecursionlimit(400000)
    content, _, meta = core.build_project_workbook(inputs, tep, [], phasing)
    assert meta["missing"] == [], meta["missing"]
    evaluator = Evaluator(openpyxl.load_workbook(io.BytesIO(content), data_only=False))
    teaser = core.build_teaser_pdf(bundle, inputs, tep, phasing)
    return bundle, numbers, model, content, evaluator, teaser


@pytest.fixture(scope="module")
def starved():
    return _build(_starved)


@pytest.fixture(scope="module")
def defaulted():
    return _build(_defaulted)


def _pdf_text(pdf: bytes) -> str:
    from pypdf import PdfReader
    reader = PdfReader(io.BytesIO(pdf))
    assert len(reader.pages) == 1, f"тизер — одна страница, а вышло {len(reader.pages)}"
    return " ".join(page.extract_text() for page in reader.pages)


# --- модель представления не считает ------------------------------------------

def test_the_presentation_layer_contains_no_arithmetic():
    tree = ast.parse((ROOT / "presentation.py").read_text(encoding="utf-8"))
    arithmetic = (ast.Add, ast.Sub, ast.Mult, ast.Div, ast.FloorDiv, ast.Mod, ast.Pow)
    found = [f"строка {node.lineno}: {type(node.op).__name__}"
             for node in ast.walk(tree)
             if isinstance(node, ast.BinOp) and isinstance(node.op, arithmetic)]
    assert not found, "в модели представления появилась арифметика: " + "; ".join(found)


def test_the_fixtures_carry_the_risks(starved, defaulted):
    """Предохранитель: без нехватки лимита и дефолта риски не проверяются."""
    _, numbers, model, *_ = starved
    assert numbers["pf_shortfall_mln"] > 1.0
    assert {r["key"] for r in model["risks"] if r["active"]} >= {"pf_shortfall", "llcr_below_target"}
    _, numbers, model, *_ = defaulted
    assert numbers["rve_unpaid_mln"] > 1.0
    assert "default_rve" in {r["key"] for r in model["risks"] if r["active"]}


# --- движок → Dashboard_Data ---------------------------------------------------

@pytest.mark.parametrize("shape", ["starved", "defaulted"])
def test_the_data_sheet_reads_the_same_numbers_as_the_engine(shape, request):
    _, numbers, _, _, evaluator, _ = request.getfixturevalue(shape)
    for key in SWORN:
        row = v4_dashboard.DATA_ROWS[key]
        book = float(evaluator.cell(v4_dashboard.DATA_SHEET, f"C{row}") or 0)
        engine = float(numbers[key] or 0)
        assert abs(book - engine) <= _tolerance(key, engine, evaluator), (key, book, engine)


def test_the_product_rows_match_the_engine(starved):
    bundle, numbers, model, _, evaluator, _ = starved
    by_key = {p["key"]: p for p in model["products"]}
    checked = 0
    for key, row in v4_dashboard.PRODUCT_ROWS.items():
        product = by_key.get(key)
        if not product:
            continue
        book_revenue = float(evaluator.cell(v4_dashboard.DATA_SHEET, f"F{row}") or 0)
        assert abs(book_revenue - float(product["revenue_mln"] or 0)) <= max(1.0, abs(book_revenue) * 0.0005), key
        if product.get("saleable"):
            book_saleable = float(evaluator.cell(v4_dashboard.DATA_SHEET, f"D{row}") or 0)
            assert abs(book_saleable - float(product["saleable"])) <= 1.0, key
        if product.get("gns"):
            book_gns = float(evaluator.cell(v4_dashboard.DATA_SHEET, f"C{row}") or 0)
            assert abs(book_gns - float(product["gns"])) <= 1.0, key
        checked += 1
    assert checked >= 3, "продуктов с выручкой должно быть хотя бы три"


@pytest.mark.parametrize("shape", ["starved", "defaulted"])
def test_the_risks_agree(shape, request):
    _, _, model, _, evaluator, _ = request.getfixturevalue(shape)
    for risk in model["risks"]:
        row = v4_dashboard.RISK_ROWS[risk["key"]]
        book_flag = float(evaluator.cell(v4_dashboard.DATA_SHEET, f"C{row}") or 0)
        assert book_flag == (1.0 if risk["active"] else 0.0), (risk["key"], book_flag, risk)


# --- Dashboard_Data → Дашборд ---------------------------------------------------

def test_the_dashboard_is_formulas_only_and_the_source_is_hidden(starved):
    _, _, _, content, evaluator, _ = starved
    archive = zipfile.ZipFile(io.BytesIO(content))
    workbook = archive.read("xl/workbook.xml").decode("utf-8")
    assert re.search(r'<x:sheet name="Dashboard_Data"[^>]*state="hidden"', workbook)
    sheet = archive.read(core._v4_sheet_path(archive, "Дашборд")).decode("utf-8")
    numbers = re.findall(r'<x:c r="([A-Z]+\d+)"[^>]*>(?:(?!</x:c>).)*?<x:v>', sheet)
    assert numbers == [], f"на «Дашборде» числа вместо формул: {numbers[:5]}"
    formulas = re.findall(r"<x:f>([^<]*)</x:f>", sheet)
    foreign = [f for f in formulas if "Dashboard_Data" not in f and "'" in f]
    assert not foreign, f"«Дашборд» читает мимо источника: {foreign[:5]}"
    data = archive.read(v4_dashboard.DATA_SHEET_PATH).decode("utf-8")
    literals = re.findall(r'<x:c r="([A-Z]+\d+)"[^>]*>\s*<x:v>', data)
    assert literals == [f"C{v4_dashboard.DATA_ROWS['llcr_target']}"], literals
    # Карточки читают те же строки источника, что и паритет.
    first_card = evaluator.cell("Дашборд", "A5")
    assert abs(float(first_card) - float(evaluator.cell("ПРОВЕРКИ", "B76"))) < 1e-6


def test_the_charts_read_the_source_ranges(starved):
    _, _, _, content, _, _ = starved
    archive = zipfile.ZipFile(io.BytesIO(content))
    types = archive.read("[Content_Types].xml").decode("utf-8")
    charts = re.findall(r'PartName="(/xl/drawings/charts/[^"]+)"', types)
    assert len(charts) == 3, charts
    for part in charts:
        xml = archive.read(part.lstrip("/")).decode("utf-8")
        refs = re.findall(r"<c:f>([^<]*)</c:f>", xml)
        assert refs and all("Dashboard_Data" in ref for ref in refs), (part, refs)
    drawing_rels = archive.read("xl/drawings/_rels/drawing1.xml.rels").decode("utf-8")
    assert len(re.findall(r"relationships/chart", drawing_rels)) == 3
    assert "chart4" not in types and "xl/drawings/charts/chart4.xml" not in archive.namelist()


# --- модель представления → тизер ----------------------------------------------

@pytest.mark.parametrize("shape", ["starved", "defaulted"])
def test_the_teaser_prints_the_cards_of_the_model(shape, request):
    _, _, model, _, _, teaser = request.getfixturevalue(shape)
    text = _pdf_text(teaser)
    for card in model["cards"]:
        unit = card["unit"]
        value = card["value"]
        shown = (core._pdf_num(value * 100, 1) + "%" if unit == "%"
                 else core._pdf_num(value, 2) + "x" if unit == "x"
                 else core._pdf_num(value, 1))
        assert shown in text, (card["key"], shown)
    for risk in model["risks"]:
        if risk["active"]:
            assert risk["label"] in text, risk["label"]
    assert model["origin"]["calculation_id"] in text
    assert core.VERSION in text


def test_the_page_offers_the_teaser_next_to_the_pdf():
    page = core.PAGE
    assert 'onclick="exportTeaserPdf()"' in page
    pdf_at = page.index('onclick="exportReportPdf()"')
    assert abs(page.index('onclick="exportTeaserPdf()"') - pdf_at) < 400
    assert "fetch('/report/teaser'" in page


def test_the_teaser_route_answers_with_one_page(monkeypatch):
    from fastapi.testclient import TestClient
    monkeypatch.delenv("TELEGRAM_BOT_TOKEN", raising=False)
    inputs, tep, _ = _starved()
    client = TestClient(core.app)
    response = client.post("/report/teaser", json={"inputs": inputs, "tep": tep, "rates": [],
                                                   "phasing": {}, "project_name": "Маршрут"})
    assert response.status_code == 200, response.text[:300]
    assert response.headers["content-type"].startswith("application/pdf")
    text = _pdf_text(response.content)
    assert "Маршрут" in text
    bad = client.post("/report/teaser", json={"tep": tep})
    assert bad.status_code == 400
