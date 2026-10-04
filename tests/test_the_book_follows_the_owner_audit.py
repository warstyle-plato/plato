"""Книга модели после ревизии 29.09.2026 — решения владельца, как их видит читатель.

Решения (docs/excel_book_audit_2026-09-29.md): книга открывается на Дашборде,
итоги впереди; объектов и очередей, которых нет в проекте, не видно; вместо
«СБОЙ» при непогашенном долге — «ДЕФОЛТ ПО ДОЛГУ»; стиль итога — на строке
итога; подпись «на м²» называет делитель; на помесячных листах закреплены
подписи и шапка.

Проверяется собранная книга, а не исходник сборщика: строка кода стоит в файле
и у сломанной сборки. Каждое утверждение падает на прежней книге — это
проверено на ней же при написании.

Запуск: python3 -m pytest tests/test_the_book_follows_the_owner_audit.py -q
"""

from __future__ import annotations

import io
import sys
import zipfile
from pathlib import Path

import openpyxl
import pytest

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT))

import main as wrapper  # noqa: E402
import v4_book_polish  # noqa: E402

core = wrapper.core


def _build(inputs=None, cache=False):
    sys.setrecursionlimit(400000)
    tep = {key: dict(value) for key, value in core.TEP_DEFAULT.items()}
    content, _, meta = core.build_project_workbook(
        {**core.DEFAULT_INPUTS, **(inputs or {})}, tep, [], {}, cache_values=cache)
    return content, meta


@pytest.fixture(scope="module")
def default_book():
    """Одна очередь, без отдельно стоящих объектов, со значениями для просмотра."""
    content, meta = _build(cache=True)
    return (openpyxl.load_workbook(io.BytesIO(content)),
            openpyxl.load_workbook(io.BytesIO(content), data_only=True), meta, content)


def test_the_book_opens_on_the_dashboard_with_totals_first(default_book):
    book, _, meta, _ = default_book
    assert book.sheetnames[:4] == ["Дашборд", "ОТЧЕТ", "Вводные", "ИНСТРУКЦИЯ"]
    assert book.sheetnames[-3:] == ["ПРОВЕРКИ", "Источники", "Dashboard_Data"]
    assert book.active.title == "Дашборд"
    assert not [m for m in meta["missing"] if "оформлен" in str(m)], meta["missing"]


def test_absent_queues_and_objects_are_hidden_not_deleted(default_book):
    book, _, _, _ = default_book
    hidden = {ws.title for ws in book.worksheets if ws.sheet_state != "visible"}
    assert {"CF_2", "CF_3", "CF_4"} <= hidden
    assert "CF_1" not in hidden
    report = book["ОТЧЕТ"]
    # Очереди 2–4 и объекты, которых нет, скрыты; квартиры — видны.
    for row in (26, 27, 28, 50, 51, 52, 53, 91, 92, 93):
        assert report.row_dimensions[row].hidden, row
    for row in (25, 46, 47, 48, 54, 87):
        assert not report.row_dimensions[row].hidden, row
    # Скрыто, но не удалено: формулы итогов по-прежнему ссылаются на строки.
    assert report["A50"].value and report["B54"].value.startswith("=SUM(")


def test_the_dashboard_shows_only_products_of_the_project(default_book):
    _, values, _, _ = default_book
    # Подписи Дашборда — формулы на источник: читаются сохранённые значения.
    labels = [str(c.value) for row in values["Дашборд"].iter_rows() for c in row
              if isinstance(c.value, str)]
    text = " ".join(labels)
    assert "Квартиры" in text
    for absent in ("МФОЦ", "Торговый центр", "Наземный паркинг"):
        assert absent not in text, absent


def test_an_unpaid_debt_is_a_default_not_a_failure(default_book):
    """Дефолтный проект к концу не гасит долг — модель при этом цела."""
    _, values, _, _ = default_book
    assert values["ПРОВЕРКИ"]["F23"].value == "FAIL"
    assert values["ПРОВЕРКИ"]["B3"].value == "ДЕФОЛТ ПО ДОЛГУ"
    assert values["ОТЧЕТ"]["B20"].value == "ДЕФОЛТ ПО ДОЛГУ"


def test_the_total_style_sits_on_the_total_row(default_book):
    book, _, _, _ = default_book
    report, tep = book["ОТЧЕТ"], book["ТЭП"]
    assert report["A54"].value == "ИТОГО ПРОДУКТЫ" and report["A54"].font.b
    assert not report["A53"].font.b
    assert tep["A35"].value == "ИТОГО ОБЪЕКТЫ" and tep["A35"].font.b
    assert not tep["A34"].font.b


def test_every_per_sqm_label_names_its_divisor(default_book):
    book, _, _, _ = default_book
    report = book["ОТЧЕТ"]
    assert report["A62"].value == "CAPEX на м² GBA продуктов"
    assert "подземной" in report["C75"].value and "соцобъект" in report["C75"].value
    assert "строительного объёма" not in report["A62"].value


def test_monthly_sheets_keep_labels_and_months_in_view(default_book):
    book, _, _, _ = default_book
    for name in ("Продажи", "CAPEX", "CF_1", "CF"):
        assert book[name].freeze_panes == "D5", name
    cf = book["CF_1"]
    assert cf["C10"].value == "млн ₽"
    assert cf["B10"].number_format.startswith("#,##0.0")
    assert book["Продажи"]["D12"].number_format == "0"   # флаг 0/1, не «100%»
    assert book["ОТЧЕТ"].page_setup.orientation == "landscape"


def test_the_risks_read_as_numbers_a_person_can_read(default_book):
    _, values, _, _ = default_book
    risks = [values["Дашборд"].cell(r, 11).value for r in range(1, 160)]
    shown = [str(v) for v in risks if isinstance(v, str) and v.startswith("ДА · ")]
    assert shown, risks
    for text in shown:
        assert "." not in text, text          # «0,89x», «980 млн ₽», а не 979.5634112254247


def test_a_row_is_not_hidden_when_its_label_is_not_the_expected_one(default_book):
    """Скрыть не ту строку хуже, чем показать лишнюю: подпись — часть адреса."""
    _, _, _, content = default_book
    missing: list[str] = []
    out = v4_book_polish.polish(content, hidden_rows={"ОТЧЕТ": [(46, "A", "Кладовые")]},
                                missing=missing)
    report = openpyxl.load_workbook(io.BytesIO(out))["ОТЧЕТ"]
    assert not report.row_dimensions[46].hidden
    assert any("ОТЧЕТ!A46" in m for m in missing), missing


def test_the_extra_object_header_is_its_own_and_merged():
    """ФОК дописан копией блока ТЦ — его заголовок не несёт подписи ТЦ."""
    content, meta = _build({"sports_enabled": True, "sports_gba_sqm": 5000,
                            "sports_saleable_sqm": 3500})
    objects = openpyxl.load_workbook(io.BytesIO(content))["ОБЪЕКТЫ"]
    texts = {objects.cell(124, c).value for c in range(1, 184)} - {None}
    assert texts == {"ФОК / МЕДЦЕНТР"}, texts
    assert "A124:GA124" in {str(m) for m in objects.merged_cells.ranges}
    assert not objects.row_dimensions[124].hidden


def test_api_keys_are_quiet_and_the_rest_of_the_column_is_not(default_book):
    """Колонку ключей скрыть нельзя — в ней сценарии и даты; ключ гасится шрифтом."""
    book, _, _, _ = default_book
    entry = book["Вводные"]
    assert entry["D14"].value == "purchase_price_mln"
    assert entry["D14"].font.sz == 8 and entry["D14"].font.color.rgb == "FFA6A6A6"
    assert entry["D4"].font.sz != 8          # «Сценарный драйвер» — не ключ
