"""Подземная часть, которой под участком нет места, — причина словами, а не тысяча уровней.

Жалоба владельца (26.10.2026): КРТ с школой и садом, подземная часть 32 137,5 м²; в одном
браузере схема писала «Под участком — 8381 подземных уровней» и тянула колонку подписей
«−1: 4 м² … −1064: 4 м²» на всю страницу. Участки школы и сада почти равны площади
участка — на одном уровне остаётся около 4 м², и оценка уровней делила площадь на
них без потолка. Теперь оценка не выдаёт больше `UNDERGROUND_LEVELS_MAX` уровней:
не помещается — уровней в оценке нет, и движок говорит, сколько места на одном
уровне, сколько уровней понадобилось бы и что поправить. Книга считает так же.

На старом коде: уровней 8 717, причины нет, на странице — сотни подписей уровней.

Запуск: python3 -m pytest tests/test_site_massing_levels_ceiling.py -q
"""

from __future__ import annotations

import copy
import io
import sys
from pathlib import Path

import pytest

ROOT = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(ROOT))
sys.path.insert(0, str(ROOT / "tests"))

import main as _wrapper  # noqa: E402

core = _wrapper.core


def _krt(site_ha: float = 4.30046, **extra):
    """Жильё со школой на 2 500 мест и садом на 250: их участки — 43 000 м²."""
    x = copy.deepcopy(core.DEFAULT_INPUTS)
    x.update(project_region="msk", vri_region="msk", site_area_ha=site_ha, **extra)
    t = {"apartments": {"gns": 100000.0, "saleable": 70000.0, "units": 1500},
         "school": {"gns": 6000.0, "units": 2500},
         "kindergarten": {"gns": 4640.0, "units": 250},
         "underground_parking": {"units": 900, "gns": 32137.5}}
    return x, t


def test_no_room_under_the_site_gives_a_reason_not_thousands_of_levels():
    x, t = _krt()
    m = core.site_massing(x, t)
    assert m["social_plot_sqm"] == pytest.approx(43000.0)
    assert m["underground_contour_sqm"] < 10
    assert m["underground_levels"] == 0, m["underground_levels"]
    assert m["underground_levels_origin"] == "estimate"
    assert m["underground_levels_unfit"] is True
    assert m["underground_level_rows"] == [] and m["underground_depth_factor"] == 1.0
    reason = m["underground_levels_reason"]
    assert reason in m["warnings"]
    for needle in ("32 138 м²", "не помещается", "на одном подземном уровне здесь около 4 м²",
                   "участки школы, детского сада", "43 000 м²", "отступа",
                   "больше 7", "«Подземных уровней»"):
        assert needle in reason, (needle, reason)
    assert "под участком не помещается" in m["summary"], m["summary"]


def test_the_ceiling_is_the_limit_and_a_fitting_part_is_counted():
    assert core.UNDERGROUND_LEVELS_MAX == core._massing_book.UNDERGROUND_LEVELS_MAX == 7
    # Тот же проект на участке, где школа и сад оставляют место, — уровни считаются.
    x, t = _krt(site_ha=6.5)
    m = core.site_massing(x, t)
    assert 1 <= m["underground_levels"] <= core.UNDERGROUND_LEVELS_MAX
    assert not m["underground_levels_unfit"] and not m["underground_levels_reason"]
    capacity = m["underground_contour_sqm"] * m["underground_use_factor"]
    assert m["underground_levels"] == -(-m["underground_sqm"] // capacity)
    # Ровно на потолке — ещё уровни, на один больше — уже причина.
    for levels, fits in ((7, True), (8, False)):
        x, t = _krt(site_ha=6.5)
        t["underground_parking"]["gns"] = capacity * (levels - 0.5)
        m = core.site_massing(x, t)
        assert (m["underground_levels"] == levels) is fits, (levels, m["underground_levels"])
        assert m["underground_levels_unfit"] is (not fits)


def test_hand_levels_stay_and_the_scheme_draws_one_slab():
    x, t = _krt(underground_levels=40)
    m = core.site_massing(x, t)
    assert (m["underground_levels"], m["underground_levels_origin"]) == (40, "manual")
    assert not m["underground_levels_unfit"]
    layout = core.site_massing_layout(x, m, [{"phase": 1, "above_gns_sqm": 110640.0}])
    assert len(layout["underground_level_rows"]) == 1
    assert any("задано вручную 40" in n for n in layout["notes"]), layout["notes"]
    ax = core.site_massing_axonometry(layout)
    assert len(ax["texts"]) <= 1, len(ax["texts"])


def test_the_scheme_of_an_unfit_part_has_no_levels_and_says_why():
    x, t = _krt()
    m = core.site_massing(x, t)
    ax = core.site_massing_axonometry(core.site_massing_layout(x, m, [{"phase": 1, "above_gns_sqm": 110640.0}]))
    assert ax["texts"] == []
    assert not [s for s in ax["shapes"] if s["kind"] == "underground_level"]
    assert "под участком не помещается" in ax["caption"], ax["caption"]
    assert "подземных уровней с площадью" not in ax["caption"]


def test_the_report_says_not_counted_instead_of_zero():
    x, t = _krt()
    m = core.site_massing(x, t)
    rows = core._site_massing_block({"tep": {"site_massing": m}})[0]["rows"]
    row = next(r for r in rows if r["label"] == "Подземных уровней")
    assert "не помещается" in str(row["value"]), row


def test_the_book_leaves_the_levels_out_like_the_engine():
    pytest.importorskip("openpyxl")
    from openpyxl import Workbook
    from xlsx_eval import Evaluator

    for site_ha in (4.30046, 6.5):
        x, t = _krt(site_ha=site_ha)
        m = core.site_massing(x, t)
        values = core._massing_book.input_values(m, x)
        book = Workbook()
        sheet = book.active
        sheet.title = "Блок"
        rows = {key: i for i, (key, _, _) in enumerate(core._massing_book.ROWS, start=1)}
        formulas = core._massing_book.formulas(lambda key: f"$B${rows[key]}")
        for key, row in rows.items():
            sheet.cell(row, 2).value = ("=" + formulas[key]) if key in formulas else values[key]
        buf = io.BytesIO()
        book.save(buf)
        from openpyxl import load_workbook
        ev = Evaluator(load_workbook(io.BytesIO(buf.getvalue())))
        got = lambda key: float(ev.cell("Блок", f"B{rows[key]}") or 0)  # noqa: E731
        assert got("underground_levels") == m["underground_levels"], site_ha
        assert (got("underground_levels") == 0) is (site_ha < 5), site_ha
        assert got("depth_factor") == pytest.approx(m["underground_depth_factor"], rel=1e-9)


# Страница — настоящий PAGE: офис на 0,243 га с 68 местами под землёй и отступом
# 24 м — внутри отступов почти ничего не остаётся, на одном уровне около 1,5 м².
SETUP = """async ()=>{
 applyProjectKind('nonresidential'); closeProjectKindDialog(); await settledCalculate();
 setSiteArea(0.243); setSiteDensity(22842/0.243);
 applyDensityToObject('offices'); await settledCalculate();
 openTab('inputs'); renderInputs();
 document.querySelectorAll('#inputs details').forEach(d=>{d.open=true});
}"""

READ = """()=>{
 openTab('tep'); renderTep();
 const body=document.getElementById('siteMassingBody');
 const scheme=document.getElementById('siteMassingScheme');
 return {text:body?body.innerText:'', scheme:scheme?scheme.innerText:'',
         levels:document.querySelectorAll('#siteMassingScheme .site-massing-level').length,
         slabs:document.querySelectorAll('#siteMassingScheme [data-kind="underground_level"]').length,
         warnings:Array.from(document.querySelectorAll('.site-massing-warning')).map(n=>n.innerText),
         m:(lastResult.tep||{}).site_massing||null};
}"""


@pytest.fixture(scope="module")
def seen():
    pytest.importorskip("playwright.sync_api")
    from playwright.sync_api import sync_playwright

    import browser

    path = browser.chromium_or_skip()
    out: dict = {}
    with browser.serve(core.app, 18187) as base, sync_playwright() as pw:
        with pw.chromium.launch(executable_path=str(path)) as engine:
            page = engine.new_page(viewport={"width": 1440, "height": 900})
            page.goto(base, wait_until="domcontentloaded")
            browser.wait_for_result(page)
            page.evaluate(SETUP)
            page.fill("#f_offices_parking_under_spaces", "68")
            page.dispatch_event("#f_offices_parking_under_spaces", "change")
            browser.wait_for_result(page)
            out["fits"] = page.evaluate(READ)
            page.evaluate("()=>{openTab('tep');renderTep();"
                          "document.getElementById('siteMassingBox').open=true}")
            page.fill("#massing_site_setback_m", "24")
            page.dispatch_event("#massing_site_setback_m", "change")
            browser.wait_for_result(page)
            out["tiny"] = page.evaluate(READ)
    return out


def test_the_page_draws_no_column_of_levels_and_names_the_reason(seen):
    fits, tiny = seen["fits"], seen["tiny"]
    # Предохранитель: с обычным отступом два уровня видны на схеме подписями.
    assert fits["m"]["underground_levels"] == 2 and fits["levels"] == 2, fits["levels"]
    m = tiny["m"]
    assert m["setback_m"] == 24 and m["underground_contour_sqm"] < 5, m["underground_contour_sqm"]
    assert tiny["levels"] == 0 and tiny["slabs"] == 0, (tiny["levels"], tiny["slabs"])
    assert m["underground_levels"] == 0 and m["underground_levels_unfit"]
    warning = " ".join(tiny["warnings"]).replace("\xa0", " ")
    assert "2 380 м² под этим участком не помещается" in warning, warning
    assert "отступ" in warning and "«Подземных уровней»" in warning, warning
    text = tiny["text"].replace("\xa0", " ")
    assert "не посчитано" in text, text
    assert "под участком не помещается" in tiny["scheme"], tiny["scheme"]
