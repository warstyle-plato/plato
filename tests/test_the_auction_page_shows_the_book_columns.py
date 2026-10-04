"""Таблица торгов на странице — те же колонки, что в книге Excel.

«Почему комментарии Платона в Excel мы вставляем, а в окне веб их нет?» и
«возможно ли сделать интерфейс таблицы в веб аналогичным Excel? Там дней до
окончания и т.д., адреса» (владелец, 04.10.2026).

Список колонок один (`auction_search/export_columns.py`): книга берёт из него
заголовки, страница получает его маршрутом `/auctions/table` вместе с клетками,
подготовленными тем же кодом, что строки книги. Проверка сверяет ОТРИСОВАННУЮ
страницу с книгой, собранной из тех же строк, и падает на подделке: убранной
колонке комментария и переставленном порядке.

Запуск: python3 -m pytest tests/test_the_auction_page_shows_the_book_columns.py -q
"""

from __future__ import annotations

import io
import json
import sys
import time
from pathlib import Path

import pytest

from browser import chromium_or_skip, serve

sys.path.insert(0, str(Path(__file__).resolve().parent.parent))

PORT = 18811
LIVE_URL = "https://torgi.gov.ru/new/public/lots/lot/LIVE-1"
PAST_URL = "https://torgi.gov.ru/new/public/lots/lot/PAST-1"
NOTE_TEXT = "Интересен: участок у метро. Опасен: обременение сетями."


def _iso(days: int) -> str:
    return time.strftime("%Y-%m-%dT18:00:00+03:00", time.gmtime(time.time() + days * 86400))


def _ru(days: int) -> str:
    return time.strftime("%d.%m.%Y 18:00", time.gmtime(time.time() + days * 86400))


def _lot(url: str, title: str, days: int, status: str) -> dict:
    return {
        "title": title, "address": "", "cadastral_numbers": ["77:01:0001001:1001"],
        "lot_kind": "land_lease", "origin": "city", "land_area_sqm": 12500,
        "current_price_rub": 45_000_000, "application_start": _ru(-20),
        "application_deadline": _ru(days), "application_deadline_iso": _iso(days),
        "auction_date": _ru(days + 5), "status": status, "documents": [],
        "source": {"lot_url": url, "source_name": "ГИС Торги", "external_lot_id": url[-6:]},
        "screening": {"why_here": "земля под застройку"}, "quality": {"accepted": True},
        "main_selection": True,
    }


LOTS = [
    _lot(LIVE_URL, "Аренда земельного участка по адресу: г. Москва, ул. Тверская, д. 1", 9,
         "APPLICATIONS_SUBMISSION"),
    # Площадка не сменила статус, а срок прошёл: это не «идут торги».
    _lot(PAST_URL, "Аренда земельного участка по адресу: г. Москва, ул. Арбат, д. 2", -3,
         "APPLICATIONS_SUBMISSION"),
]


def _app(tmp_path, monkeypatch):
    from fastapi import FastAPI

    monkeypatch.setenv("DATA_DIR", str(tmp_path))
    from auction_search import lot_notes
    from auction_search.api import install

    store = lot_notes.LotNotes(tmp_path / "market")
    store.remember_lots([{"url": LIVE_URL, "title": LOTS[0]["title"]}])
    store.save_note(LIVE_URL, {"text": NOTE_TEXT, "asked_at": time.time(), "sources": []})
    app = FastAPI()
    install(app)
    return app


def _book_headers_and_rows(payload: list[dict]) -> tuple[list[str], list[list]]:
    from openpyxl import load_workbook

    from auction_search.api import _xlsx

    ws = load_workbook(io.BytesIO(_xlsx(payload))).active
    rows = [[cell.value for cell in row] for row in ws.iter_rows()]
    return [str(value) for value in rows[0]], rows[1:]


READ = """() => ({
  heads: [...document.querySelectorAll('#sheetHead th')].map(th => th.dataset.key + '|' + th.childNodes[0].textContent),
  rows: [...document.querySelectorAll('#sheetRows tr.sheetrow')].map(tr =>
    [...tr.children].map(td => td.textContent)),
  payload: exportPayload(state.filtered, 'auctions'),
  scores: state.filtered.map(l => String(lotScore(l).score)),
})"""


def check_page_against_book(seen: dict) -> None:
    """Колонки страницы = колонки книги, в том же порядке; комментарий на месте."""
    book_heads, book_rows = _book_headers_and_rows(seen["payload"])
    page_heads = [head.split("|", 1)[1] for head in seen["heads"]]
    assert page_heads == book_heads, f"страница {page_heads} ≠ книга {book_heads}"
    keys = [head.split("|", 1)[0] for head in seen["heads"]]
    assert "plato_comment" in keys, "колонки комментария Платона на странице нет"
    comment = keys.index("plato_comment")
    for page_row, book_row in zip(seen["rows"], book_rows):
        assert page_row[comment].strip() == str(book_row[comment]).strip(), (
            "комментарий на странице не тот, что в книге")


def _open(page, base: str) -> None:
    page.route("**/auctions/discover*", lambda route: route.fulfill(
        status=200, content_type="application/json",
        body=json.dumps({"lots": LOTS, "coverage": [], "quality": {}, "count": len(LOTS)})))
    page.goto(f"{base}/auctions", wait_until="domcontentloaded")
    page.click("#refresh")
    page.wait_for_function("document.querySelectorAll('#sheetRows tr.sheetrow').length === 2")


def test_the_book_and_the_table_route_read_one_column_list(tmp_path, monkeypatch):
    """Маршрут страницы и книга отдают один список колонок — из одного места."""
    from fastapi.testclient import TestClient

    from auction_search import export_columns

    client = TestClient(_app(tmp_path, monkeypatch))
    payload = [{"section": "Торги", "name": LOTS[0]["title"], "url": LIVE_URL,
                "application_deadline_iso": _iso(9), "status": "APPLICATIONS_SUBMISSION"},
               {"section": "Торги", "name": LOTS[1]["title"], "url": PAST_URL,
                "application_deadline_iso": _iso(-3), "status": "APPLICATIONS_SUBMISSION"}]
    answer = client.post("/auctions/table", json={"rows": payload, "kind": "auctions"}).json()
    titles = [column["title"] for column in answer["columns"]]
    assert titles == [column.title for column in export_columns.AUCTION_COLUMNS]
    assert titles == _book_headers_and_rows(payload)[0]
    live, past = answer["rows"]
    assert NOTE_TEXT in live["cells"]["plato_comment"]["text"]
    assert live["cells"]["plato_comment"]["ready"] is True
    # Неразобранный лот — причина, а не пустота.
    assert past["cells"]["plato_comment"]["text"].startswith("Не разбирался: приём заявок закончился")
    assert live["cells"]["days_to_deadline"]["value"] == 9
    assert past["cells"]["days_to_deadline"]["text"].startswith("Приём заявок закончился")
    assert past["cells"]["status"]["text"] == "Приём заявок закончился"
    assert live["cells"]["status"]["text"] == "Приём заявок"
    # В книге — тот же статус: просроченный лот и там не «приём заявок».
    _, book_rows = _book_headers_and_rows(payload)
    status = titles.index("Статус")
    assert book_rows[1][status] == "Приём заявок закончился"


@pytest.mark.timeout(240)
def test_in_a_real_browser_the_page_has_the_book_columns(tmp_path, monkeypatch):
    chrome = chromium_or_skip()
    from playwright.sync_api import sync_playwright

    shots = Path(__import__("os").environ.get("AUCTION_SHOTS_DIR") or tmp_path)
    with serve(_app(tmp_path, monkeypatch), PORT) as base, sync_playwright() as pw:
        browser = pw.chromium.launch(executable_path=str(chrome))
        try:
            page = browser.new_page(viewport={"width": 1440, "height": 900})
            errors: list[str] = []
            page.on("pageerror", lambda exc: errors.append(str(exc)))
            _open(page, base)
            seen = page.evaluate(READ)
            check_page_against_book(seen)
            keys = [head.split("|", 1)[0] for head in seen["heads"]]
            days, status = keys.index("days_to_deadline"), keys.index("status")
            live, past = seen["rows"]
            assert live[days] == "9", live[days]
            assert past[days].startswith("Приём заявок закончился"), past[days]
            assert past[status] == "Приём заявок закончился", past[status]
            # Балл в таблице — тот же lotScore, что в строке подборки и карточке лота.
            score = keys.index("score")
            assert [row[score] for row in seen["rows"]] == seen["scores"]
            assert page.is_visible("#sheetRows td.k-plato_comment"), "комментарий на десктопе скрыт"
            page.screenshot(path=str(shots / "auctions-desktop.png"), full_page=False)

            # Подделка обязана ронять сверку: колонка убрана, порядок переставлен.
            page.evaluate("state.sheet.columns = state.sheet.columns.filter(c => c.key !== 'plato_comment'); renderSheet()")
            with pytest.raises(AssertionError):
                check_page_against_book(page.evaluate(READ))
            page.evaluate("loadSheet()")
            page.wait_for_function("document.querySelectorAll('#sheetHead th.k-plato_comment').length === 1")
            page.evaluate("const c = state.sheet.columns; [c[1], c[2]] = [c[2], c[1]]; renderSheet()")
            with pytest.raises(AssertionError):
                check_page_against_book(page.evaluate(READ))
            assert not errors, errors

            # Десктоп: окно карточки открывается строкой, закрывается Esc,
            # выделенная строка — та, чья карточка открыта.
            page.click("#sheetRows tr.sheetrow >> nth=1")
            assert page.is_visible("#lotModal")
            assert page.evaluate("state.selected.source.lot_url") == PAST_URL
            assert page.evaluate("document.querySelector('#sheetRows tr.sel td.k-url').textContent") == PAST_URL
            assert "приём заявок закончился" in page.inner_text("#lotModalNote")
            # Просроченный лот очередь не берёт — и кнопки заказа у него нет.
            assert page.locator("#lotModalNote .askNote").count() == 0
            page.screenshot(path=str(shots / "auctions-desktop-card.png"))
            page.keyboard.press("Escape")
            page.wait_for_function("document.getElementById('lotModal').classList.contains('hidden')")
            assert page.evaluate("!!document.querySelector('#auctionLayout #side')")

            # Телефон: таблица не ломает ширину, комментарий — раскрытием строки.
            phone = browser.new_page(viewport={"width": 390, "height": 844}, is_mobile=True)
            phone.on("pageerror", lambda exc: errors.append(str(exc)))
            _open(phone, base)
            width = phone.evaluate("[document.documentElement.scrollWidth, innerWidth]")
            assert width[0] <= width[1], f"страница шире экрана: {width}"
            assert not phone.is_visible("#sheetRows td.k-plato_comment")
            phone.screenshot(path=str(shots / "auctions-phone.png"))
            # Строка открывает карточку лота окном, как у КРТ.
            phone.click("#sheetRows tr.sheetrow >> nth=0")
            modal = phone.locator("#lotModal")
            assert modal.is_visible()
            text = modal.inner_text()
            assert NOTE_TEXT in text, "в окне нет комментария Платона"
            # Название в строке обрезано до трёх строк; целиком — в окне.
            assert LOTS[0]["title"] in text
            # В окне — та же карточка подборки (#side), а не вторая копия.
            assert phone.evaluate("!!document.querySelector('#lotModal #side .actions')")
            box = phone.locator("#lotModalShell").bounding_box()
            assert box and box["x"] >= 0 and box["x"] + box["width"] <= 390, box
            width = phone.evaluate("[document.documentElement.scrollWidth, innerWidth]")
            assert width[0] <= width[1], f"окно расширило страницу: {width}"
            phone.screenshot(path=str(shots / "auctions-phone-open.png"))
            # «Назад» закрывает окно и возвращает карточку на её место.
            phone.go_back()
            phone.wait_for_function("document.getElementById('lotModal').classList.contains('hidden')")
            assert phone.evaluate("!!document.querySelector('#auctionLayout #side')")
            assert phone.url.endswith("/auctions"), phone.url
            assert not errors, errors
        finally:
            browser.close()
