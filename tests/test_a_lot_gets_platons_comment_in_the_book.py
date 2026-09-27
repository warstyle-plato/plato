"""У каждого лота торгов в выгрузке — комментарий Платона.

Владелец, 26.09.2026: «чем опасен, чем интересен, поиск информации из открытых
источников в интернете», искать для каждого лота — «из тех, что попали в
выгрузку». Очередь наполняет выгрузка; разбор идёт фоном по одному
лоту (платный поиск + ответ модели), ответ ложится на диск, выгрузка читает
сохранённое. Здесь идут настоящие маршрут сбора, фоновый шаг и выгрузка;
подменены только внешние поиск и модель.
"""

from __future__ import annotations

import io
import sys
import time
import types

import openpyxl
from fastapi import FastAPI
from fastapi.testclient import TestClient

from auction_search import api
from auction_search import lot_notes as rules

SOON = "2099-01-10T12:00:00+03:00"
LATER = "2099-03-10T12:00:00+03:00"


def _lot(url, *, deadline_iso=SOON, cadastre=("77:01:0001001:1",), address="Москва, ул. Тверская, 1"):
    return {"title": f"Лот {url[-1]}", "address": address, "cadastral_numbers": list(cadastre),
            "current_price_rub": 10_000_000, "application_deadline": "10.01.2099 12:00",
            "application_deadline_iso": deadline_iso,
            "source": {"lot_url": url, "source_name": "РАД", "platform": "lot_online"}}


def test_only_exported_lots_are_queued(tmp_path):
    store = rules.LotNotes(tmp_path)
    store.remember_lots([_lot("https://etp/lot/1"), _lot("https://etp/lot/2")])
    assert store.queue() == [], "собранный, но не выгруженный лот в платный разбор не идёт"
    store.request([{"url": "https://etp/lot/2", "name": "Лот 2"}])
    queued = store.queue()
    assert [lot["url"] for lot in queued] == ["https://etp/lot/2"]
    assert queued[0]["title"] == "Лот 2" and queued[0]["cadastral_numbers"], (
        "факты сбора не затёрты строкой выгрузки")


def test_a_row_the_collector_never_saw_is_queued_from_its_own_facts(tmp_path):
    store = rules.LotNotes(tmp_path)
    store.request([{"url": "https://etp/lot/7", "name": "Склад", "address": "Москва, ул. А, 1",
                    "cadastre": "77:01:1:1, 77:01:1:2", "price": 5_000_000,
                    "application_deadline_iso": SOON}])
    lot = store.queue()[0]
    assert lot["cadastral_numbers"] == ["77:01:1:1", "77:01:1:2"]
    assert lot["current_price_rub"] == 5_000_000


def test_the_queue_goes_unanswered_first_and_soonest_first(tmp_path):
    store = rules.LotNotes(tmp_path)
    lots = [
        _lot("https://etp/lot/1", deadline_iso=LATER),
        _lot("https://etp/lot/2", deadline_iso=SOON),
        _lot("https://etp/lot/3", deadline_iso="2000-01-01T00:00:00+03:00"),
    ]
    store.remember_lots(lots)
    store.request([{"url": lot["source"]["lot_url"],
                    "application_deadline_iso": lot["application_deadline_iso"]}
                   for lot in lots])
    order = [lot["url"] for lot in store.queue()]
    assert order == ["https://etp/lot/2", "https://etp/lot/1"], "прошедший срок в очередь не идёт"
    store.save_note("https://etp/lot/2", {"asked_at": int(time.time()), "text": "ok"})
    assert [lot["url"] for lot in store.queue()] == ["https://etp/lot/1"]


def test_search_asks_by_cadastre_and_by_address():
    asked = rules.search_queries(_lot("https://etp/lot/1"))
    assert asked == ['"77:01:0001001:1"', "Москва, ул. Тверская, 1 торги продажа"]


class _Doc:
    def __init__(self, url):
        self.url = url

    def to_dict(self):
        return {"title": "Новость", "url": self.url, "snippet": "Участок в залоге у банка"}


def _app(tmp_path, monkeypatch, answer):
    monkeypatch.setenv("DATA_DIR", str(tmp_path))
    monkeypatch.delenv("AUCTIONS_VIEW_KEY", raising=False)
    asked = []

    class Request:
        def __init__(self, message, inputs, tep):
            self.message = message

    def plato_answer(req, request):
        asked.append((req.message, request.client.host))
        return answer(req)

    core = types.ModuleType("developaid_core")
    core.plato_answer = plato_answer
    core.AgentChatRequest = Request
    core.DEFAULT_INPUTS = {}
    core.TEP_DEFAULT = {}
    monkeypatch.setitem(sys.modules, "developaid_core", core)
    search = types.SimpleNamespace(
        configured=True, search=lambda query: [_Doc("https://news/1"), _Doc("https://etp/lot/1")])
    app = FastAPI()
    app.state.market_discovery_service = types.SimpleNamespace(search=search, krt=None)
    api.install(app)
    return app, asked


def test_one_background_step_writes_a_comment_the_book_prints(tmp_path, monkeypatch):
    app, asked = _app(tmp_path, monkeypatch, lambda req: {
        "answer": "Чем интересен: центр.\nЧем опасен: залог [1].\nЧто пишут: залог у банка [1]."})
    store = rules.LotNotes(tmp_path / "market")
    store.remember_lots([_lot("https://etp/lot/1"), _lot("https://etp/lot/2", deadline_iso=LATER)])
    client = TestClient(app)
    rows = [{"section": "Торги", "name": "Лот 1", "url": "https://etp/lot/1"},
            {"section": "Торги", "name": "Лот 2", "url": "https://etp/lot/2"}]
    first_book = client.post("/auctions/export.xlsx", json={"rows": rows, "kind": "auctions"})
    assert first_book.status_code == 200
    ws = openpyxl.load_workbook(io.BytesIO(first_book.content)).active
    column = [cell.value for cell in ws[1]].index(
        "Комментарий Платона: чем интересен, чем опасен, что пишут")
    assert ws[2][column].value.startswith("Ещё не разобран: в очереди 1-й")
    assert ws[3][column].value.startswith("Ещё не разобран: в очереди 2-й")

    assert app.state.auction_lot_note_once() is True
    message, client = asked[0]
    assert "Участок в залоге у банка" in message and "77:01:0001001:1" in message
    assert client == "auctions-lot-notes", "фон тратит свой лимит Платона, а не чужой"
    note = store.note("https://etp/lot/1")
    assert [item["url"] for item in note["sources"]] == ["https://news/1"], (
        "карточка самого лота — не «что о нём пишут»")

    rows = rows + [{"section": "Торги", "name": "Чужой", "url": "https://etp/lot/9"}]
    ws = openpyxl.load_workbook(io.BytesIO(api._xlsx(rows, "auctions"))).active
    first, second, foreign = (ws[r][column].value for r in (2, 3, 4))
    assert first.startswith("Чем интересен: центр.") and "[1] https://news/1" in first
    assert second.startswith("Ещё не разобран: в очереди 1-й")
    assert foreign.startswith("Не разбирался")


def test_a_failed_answer_is_named_in_the_book(tmp_path, monkeypatch):
    def boom(req):
        raise RuntimeError("лимит AI-запросов исчерпан")

    app, _asked = _app(tmp_path, monkeypatch, boom)
    store = rules.LotNotes(tmp_path / "market")
    store.remember_lots([_lot("https://etp/lot/1")])
    store.request([{"url": "https://etp/lot/1"}])
    assert app.state.auction_lot_note_once() is True
    rows = [{"section": "Торги", "name": "Лот 1", "url": "https://etp/lot/1"}]
    ws = openpyxl.load_workbook(io.BytesIO(api._xlsx(rows, "auctions"))).active
    header = [cell.value for cell in ws[1]]
    cell = ws[2][header.index("Комментарий Платона: чем интересен, чем опасен, что пишут")].value
    assert cell.startswith("Разбор не удался") and "лимит" in cell
    assert store.queue() == [], "неудачный разбор не повторяется каждую минуту"


def test_the_krt_book_has_no_lot_comment_column():
    ws = openpyxl.load_workbook(io.BytesIO(api._xlsx([{"name": "КРТ"}], "krt"))).active
    assert not any("Комментарий Платона" in str(cell.value) for cell in ws[1])
