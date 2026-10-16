"""Книга v4 считает как движок там, где повторная ревизия 15.10.2026 нашла разрыв.

Ревизия пересчитала книги восьми эталонов и «Донского» (LibreOffice) и
сверила ОТЧЕТ и ПРОВЕРКИ с движком. Четыре разрыва были ошибками реализации;
каждый тест ниже падает на прежней книге:

1. «Реализованный объём» объекта с местами на первых этажах вычитал их из
   продаваемой второй раз и не той площадью — FAIL 23 400 против 22 000 и
   статус «СБОЙ» на исправном расчёте.
2. Денежная компенсация без введённой суммы: движок платит стоимость стройки
   программы, книга — ноль (Донской: 49,75 млн ₽ соцнагрузки, 57 млн CAPEX).
   Строка «Социальная нагрузка» структуры расходов показывала ноль.
3. Перенесённые в РнС проценты БРИДЖа капитализировались по ставке ПФ со
   следующего месяца, а не с месяца переноса — 2,5 млн ₽ стоимости
   финансирования на умолчаниях.
4. «Оценка потребности в финансировании» прибавляла покупку, ВРИ и
   соцнагрузку к «Расходам проекта всего», где они уже есть.

Полный разбор — docs/excel_book_audit_2026-10-15.md.

Запуск: python3 -m pytest tests/test_the_book_audit_2026_10_15.py -q
"""

from __future__ import annotations

import copy
import io
import json
import sys
from pathlib import Path

import pytest

ROOT = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(ROOT))
sys.path.insert(0, str(ROOT / "tests"))

import main_legacy as core  # noqa: E402

openpyxl = pytest.importorskip("openpyxl")

DONSKOY = json.loads((ROOT / "tests" / "fixtures" / "project_settings_donskoy_0_25_78.json")
                     .read_text(encoding="utf-8"))


def _evaluate(content: bytes):
    from xlsx_eval import Evaluator

    sys.setrecursionlimit(400000)
    return Evaluator(openpyxl.load_workbook(io.BytesIO(content), data_only=False))


def _book(x: dict, t: dict, name: str):
    content, _name, meta = core.build_project_workbook(
        dict(x), copy.deepcopy(t), [], {}, project_name=name)
    assert list(meta.get("missing") or []) == [], meta.get("missing")
    return _evaluate(content)


def _engine(x: dict, t: dict) -> dict:
    return core._run_authoritative_model(dict(x), copy.deepcopy(t), [], {})["consolidated"]


# --- 1. места первых этажей вычитаются из продаваемой один раз ----------------

def _office_with_floor_places() -> tuple[dict, dict]:
    x = dict(core.DEFAULT_INPUTS)
    x.update(offices_enabled=True, offices_gba_sqm=40000.0, offices_saleable_sqm=24000.0,
             offices_parking_under_spaces=200, offices_parking_over_spaces=40,
             offices_parking_guest_pct=10)
    t = copy.deepcopy(core.TEP_DEFAULT)
    t["offices"].update(gns=40000, total_area=37600, useful=24000, saleable=24000)
    return x, t


@pytest.fixture(scope="module")
def office_book():
    x, t = _office_with_floor_places()
    return _engine(x, t), _book(x, t, "Офис с местами на первых этажах")


def test_the_office_really_has_floor_places(office_book) -> None:
    """Предохранитель: без мест первых этажей проверка ниже ничего не значит."""
    engine, _book_ = office_book
    row = next(r for r in engine["tep"]["rows"] if r["key"] == "offices")
    assert row["parking_over_units"] == 40
    assert row["parking_saleable_after_sqm"] < row["saleable"]


def test_the_volume_check_expects_what_the_engine_sells(office_book) -> None:
    engine, book = office_book
    sold = next(r for r in engine["tep"]["rows"] if r["key"] == "offices")[
        "parking_saleable_after_sqm"]
    assert book.cell("ПРОВЕРКИ", "C39") == pytest.approx(sold, rel=1e-9)
    assert book.cell("ПРОВЕРКИ", "B39") == pytest.approx(sold, rel=1e-9)
    assert book.cell("ПРОВЕРКИ", "F39") == "OK"


def test_no_check_of_the_office_book_fails(office_book) -> None:
    _engine_, book = office_book
    checks = book.workbook["ПРОВЕРКИ"]
    failed = [f"{checks[f'A{row}'].value}: {book.cell('ПРОВЕРКИ', f'B{row}')} "
              f"против {book.cell('ПРОВЕРКИ', f'C{row}')}"
              for row in range(4, 130)
              if checks[f"A{row}"].value and book.cell("ПРОВЕРКИ", f"F{row}") == "FAIL"]
    assert failed == []


# --- 2. денежная компенсация без суммы — стоимость стройки, как у движка ------

@pytest.fixture(scope="module")
def donskoy():
    x = {**core.DEFAULT_INPUTS, **DONSKOY["inputs"]}
    t = copy.deepcopy(DONSKOY["tep"])
    return _engine(x, t), _book(x, t, "Донской")


def test_donskoy_pays_money_without_a_typed_amount(donskoy) -> None:
    """Предохранитель: режим денежный, суммы нет, а движок платит."""
    engine, _book_ = donskoy
    assert DONSKOY["inputs"]["social_mode"] == "Денежная компенсация"
    assert not float(DONSKOY["inputs"].get("social_compensation_mln") or 0)
    assert engine["capex"]["social"] > 1e6


def test_the_book_capex_carries_the_engine_social_payment(donskoy) -> None:
    engine, book = donskoy
    social = sum(book.cell("CAPEX", f"B{row}") or 0 for row in (31, 65, 99, 133))
    assert social == pytest.approx(engine["capex"]["social"] / 1e6, rel=1e-9)
    assert book.cell("ОТЧЕТ", "B6") == pytest.approx(engine["summary"]["capex"] / 1e6, rel=1e-6)


def test_the_expense_structure_shows_the_social_payment(donskoy) -> None:
    """Строка «Социальная нагрузка» показывает то, что заплачено, а не пустое поле."""
    engine, book = donskoy
    assert book.cell("ОТЧЕТ", "B36") == pytest.approx(engine["capex"]["social"] / 1e6, rel=1e-9)
    assert book.cell("ОТЧЕТ", "H37") == pytest.approx(book.cell("ОТЧЕТ", "B36"), rel=1e-9)


def test_the_typed_compensation_still_wins() -> None:
    """Введённая сумма платится как есть — стоимость стройки её не подменяет."""
    x = {**core.DEFAULT_INPUTS, **DONSKOY["inputs"], "social_compensation_mln": 120.0}
    assert core.social_cash_compensation(x) == pytest.approx(120e6)
    x["social_compensation_mln"] = 0
    assert core.social_cash_compensation(x) == pytest.approx(
        sum(core.social_construction_costs(x).values()))


# --- 3 и 4. стоимость финансирования и оценка потребности на умолчаниях --------

@pytest.fixture(scope="module")
def default_book():
    x, t = dict(core.DEFAULT_INPUTS), copy.deepcopy(core.TEP_DEFAULT)
    return _engine(x, t), _book(x, t, "Умолчания")


def test_the_bridge_interest_compounds_from_the_permit_month(default_book) -> None:
    engine, book = default_book
    finance = engine["finance"]
    assert finance["transferred_bridge_interest"] > 1e6, "переносить нечего — тест пуст"
    assert book.cell("ОТЧЕТ", "B9") == pytest.approx(
        finance["financing_cost"] / 1e6, rel=2e-5)


def test_the_financing_need_counts_each_cost_once(default_book) -> None:
    engine, book = default_book
    assert book.cell("Параметры модели", "B24") == pytest.approx(
        engine["finance"]["pf_limit_required"] / 1e6)


# --- 5. книга нежилого: IRR и NPV при непогашенном долге — «N/A», как у движка --

@pytest.fixture(scope="module")
def nonres_unrepaid():
    sys.path.insert(0, str(ROOT / "scripts"))
    import golden_snapshot

    spec = golden_snapshot.s_nonresidential(core, ROOT)
    engine = core._run_authoritative_model(
        dict(spec["inputs"]), copy.deepcopy(spec["tep"]), [], {})["consolidated"]
    content, _name, _meta = core.build_project_workbook(
        dict(spec["inputs"]), copy.deepcopy(spec["tep"]), [], {}, project_name="нежилой")
    return engine, _evaluate(content)


def test_the_nonres_engine_takes_the_returns_away(nonres_unrepaid) -> None:
    """Предохранитель: у эталона долг ПФ не погашен, и движок снял IRR/NPV."""
    engine, _book_ = nonres_unrepaid
    assert engine["summary"]["equity_returns"]["status"] == "debt_unrepaid"
    assert engine["summary"]["npv"] is None


def test_the_nonres_book_says_na_like_the_engine(nonres_unrepaid) -> None:
    engine, book = nonres_unrepaid
    label = engine["summary"]["equity_returns"]["label"]
    report = book.workbook["Отчёт"]
    shown = {report[f"A{row}"].value: book.cell("Отчёт", f"B{row}")
             for row in range(1, report.max_row + 1)
             if str(report[f"A{row}"].value or "").startswith(("NPV", "IRR"))}
    assert shown and all(value == label for value in shown.values()), shown
    assert book.cell("Сверка", "I3") == "ПРОЙДЕНО"


# --- 6. хвост продаж объекта — от РВЭ, даже когда продажи стартуют позже стройки --

@pytest.fixture(scope="module")
def late_office():
    x, t = _office_with_floor_places()
    x.update(offices_start="2029-01-01", offices_sales_start="2029-07-01",
             offices_residual_months=12, offices_months=30)
    return _engine(x, t), _book(x, t, "Офис: продажи позже стройки")


def test_the_late_office_sells_in_the_engine_window(late_office) -> None:
    import datetime as _dt
    from openpyxl.utils import get_column_letter

    engine, book = late_office
    product = next(p for p in engine["report"]["products"] if p["key"] == "offices")
    assert product["sales_start"] == "2029-07-01" and product["sales_end"] == "2032-06-01"
    sold = []
    for column in map(get_column_letter, range(4, 184)):
        month = book.cell("ОБЪЕКТЫ", f"{column}3")
        if isinstance(month, (int, float)):            # серийная дата Excel
            month = _dt.datetime(1899, 12, 30) + _dt.timedelta(days=float(month))
        if isinstance(month, _dt.datetime):
            month = month.date()
        if isinstance(month, _dt.date) and (book.cell("ОБЪЕКТЫ", f"{column}22") or 0) > 1e-9:
            sold.append(month.isoformat())
    assert sold and sold[0] == "2029-07-01" and sold[-1] == "2032-06-01", sold[-3:]
    assert book.cell("ОТЧЕТ", "B5") == pytest.approx(engine["summary"]["revenue"] / 1e6, rel=1e-6)


# --- 7. налог в потоках единственной очереди — годами с переносом убытка ----------

@pytest.fixture(scope="module")
def long_tail():
    """Убыточный проект с годом прибыли в хвосте продаж: половинное правило."""
    x, t = dict(core.DEFAULT_INPUTS), copy.deepcopy(core.TEP_DEFAULT)
    x["residual_sales_months"] = 12
    return _engine(x, t), _book(x, t, "Длинный хвост продаж")


def test_the_long_tail_pays_tax_under_the_half_rule(long_tail) -> None:
    """Предохранитель: движок платит налог, хотя накопленная база — убыток."""
    engine, _book_ = long_tail
    assert engine["summary"]["profit_tax"] > 1e6
    assert engine["summary"]["profit_before_tax"] < 0


def test_the_queue_cash_flow_pays_the_engine_tax(long_tail) -> None:
    engine, book = long_tail
    tax = engine["summary"]["profit_tax"] / 1e6
    assert book.cell("CF", "B14") == pytest.approx(tax, rel=1e-4)
    contributed = sum(-v for v in engine["cashflow"]["equity"] if v < 0) / 1e6
    assert book.cell("CF", "B23") == pytest.approx(contributed, rel=1e-5)


# --- 8. оформление: числа с форматом, текст не раздувает строки, ось по горизонту ---

@pytest.fixture(scope="module")
def styled_book():
    x = {**core.DEFAULT_INPUTS, **DONSKOY["inputs"]}
    content, _name, meta = core.build_project_workbook(
        x, copy.deepcopy(DONSKOY["tep"]), [], {}, project_name="Донской")
    assert list(meta.get("missing") or []) == [], meta.get("missing")
    return content, openpyxl.load_workbook(io.BytesIO(content))


@pytest.mark.parametrize("sheet, coords", [
    ("CF", ["E32", "E33", "E36", "E37", "B32", "B37"]),
    ("ПРОВЕРКИ", ["B76", "C76", "D76", "B82", "C82"]),
    ("Продажи", ["E10", "F10"]),
    ("ОБЪЕКТЫ", ["E15", "F15"]),
    ("КОНСОЛИДАТОР", ["B34", "C30", "B38"]),
])
def test_numbers_on_visible_sheets_have_a_format(styled_book, sheet, coords) -> None:
    """Строки, где «-28,42323685» и «0,8783333333» печатались как есть."""
    _content, book = styled_book
    bare = [coord for coord in coords if book[sheet][coord].number_format == "General"]
    assert bare == [], f"{sheet}: без формата {bare}"


def test_report_comments_do_not_blow_up_the_rows(styled_book) -> None:
    _content, book = styled_book
    report = book["ОТЧЕТ"]
    wrapped = [f"H{row}" for row in range(57, 74) if report[f"H{row}"].alignment.wrap_text]
    assert wrapped == [], wrapped
    assert report["A21"].alignment.wrap_text and report["A22"].alignment.wrap_text


def test_the_entry_sheet_has_no_empty_looking_fields(styled_book) -> None:
    """Подпись и единица без значения выглядели полем, которое забыли заполнить."""
    _content, book = styled_book
    entry = book["Вводные"]
    empty = [row for row in range(5, entry.max_row + 1)
             if isinstance(entry[f"A{row}"].value, str) and entry[f"C{row}"].value
             and not str(entry[f"A{row}"].value).isupper() and entry[f"B{row}"].value is None]
    assert empty == [], [entry[f"A{row}"].value for row in empty]


def test_the_debt_chart_ends_near_the_horizon(styled_book) -> None:
    import re
    import zipfile

    content, _book = styled_book
    with zipfile.ZipFile(io.BytesIO(content)) as archive:
        charts = [archive.read(name).decode("utf-8") for name in archive.namelist()
                  if "/charts/chart" in name and name.endswith(".xml")]
    debt = next(text for text in charts if "Долг и эскроу" in text)
    ends = set(re.findall(r"Dashboard_Data'!\$D\$\d+:\$([A-Z]+)\$\d+", debt))
    assert ends and "GA" not in ends, ends


def test_the_social_header_spans_its_table(styled_book) -> None:
    _content, book = styled_book
    assert any(str(rng).startswith("A38:") for rng in book["ТЭП"].merged_cells.ranges)


# --- 9. решения владельца 16.10.2026: НДС в полных расходах, CAPEX на объём, IRR --

def test_full_costs_carry_the_vat_like_the_engine(default_book) -> None:
    engine, book = default_book
    summary = engine["summary"]
    assert summary["vat"] > 1e6, "НДС нет — тест пуст"
    full = summary["total_expenses"] / 1e6
    assert book.cell("ОТЧЕТ", "B82") == pytest.approx(full, rel=1e-5)
    assert book.cell("ОТЧЕТ", "B40") == pytest.approx(full, rel=1e-5)
    assert abs(book.cell("ОТЧЕТ", "B42")) < 1e-6
    # Выручка − полные расходы = чистая прибыль: внутри ОТЧЕТа одно равенство.
    assert (book.cell("ОТЧЕТ", "B5") - book.cell("ОТЧЕТ", "B82")) == pytest.approx(
        book.cell("ОТЧЕТ", "B12"), abs=1e-6)
    assert book.cell("ОТЧЕТ", "F63") == pytest.approx(
        summary["full_cost_per_saleable_th"], rel=1e-5)


def test_capex_per_metre_uses_the_construction_volume(donskoy) -> None:
    """Донской: подземные кладовые входят в объём, как у движка."""
    engine, book = donskoy
    summary = engine["summary"]
    assert book.cell("ОТЧЕТ", "F58") == pytest.approx(summary["construction_volume_sqm"], rel=1e-6)
    assert book.cell("ОТЧЕТ", "F62") == pytest.approx(summary["capex_per_total_area_th"], rel=1e-5)


def test_the_irr_is_monthly_like_the_engine() -> None:
    """Две очереди с объектами: XIRR по дням давал 61,4% против 61,8% движка."""
    sys.path.insert(0, str(ROOT / "scripts"))
    import golden_snapshot

    spec = golden_snapshot.s_instances(core, ROOT)
    engine = core._run_authoritative_model(
        dict(spec["inputs"]), copy.deepcopy(spec["tep"]), [], copy.deepcopy(spec["phasing"]))
    content, _name, _meta = core.build_project_workbook(
        dict(spec["inputs"]), copy.deepcopy(spec["tep"]), [], copy.deepcopy(spec["phasing"]),
        project_name="instances")
    irr = engine["consolidated"]["summary"]["irr_equity"]
    assert isinstance(irr, float), "IRR нет — тест пуст"
    # Потоки книги и движка расходятся на десятые доли млн (ГНС с ROUND шаблона),
    # XIRR же уводил на 0,7% — допуск между ними.
    assert _evaluate(content).cell("ОТЧЕТ", "B14") == pytest.approx(irr, rel=1e-3)


def test_the_sold_sports_centre_counts_in_the_saleable_area() -> None:
    x = dict(core.DEFAULT_INPUTS)
    x.update(sports_enabled=True, sports_gba_sqm=20000.0, sports_saleable_sqm=12000.0,
             sports_disposition="sale", sports_purpose="sport")
    t = copy.deepcopy(core.TEP_DEFAULT)
    t.setdefault("sports", {}).update(gns=20000, total_area=18800, useful=12000, saleable=12000)
    engine, book = _engine(x, t), _book(x, t, "ФОК на продажу")
    assert book.cell("ОТЧЕТ", "F57") == pytest.approx(
        engine["summary"]["monetizable_saleable_sqm"], rel=1e-6)
