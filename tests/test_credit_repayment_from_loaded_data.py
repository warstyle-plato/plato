"""Погашение кредитов считается из уже загруженного — без ручных файлов.

Владелец, 27.10.2026: «Откуда я возьму эти файлы и данные? Они все есть или в
НКЛ, или в РСС, или в модели». Здесь закреплено:

- проект «Кутузов Сити» в Мониторе сам находит условия своих НКЛ;
- выдачи и долг — из реестра платежей РСС (платежи из заёмных средств),
  выдача ложится на первую НКЛ в пределах её лимита на дату, остаток — дальше;
- проценты и плата за лимит до даты РСС досчитаны от первой выдачи;
- эскроу — из графиков договоров отчёта о продажах, новые продажи — только
  превышение плана финмодели над договорами (договор не считается дважды);
- прогнозные выдачи — план остатка ПФ финмодели сверх выбранного;
- ручной набор — поправка поверх найденного, с подписью;
- чего нет — названо: что и где не нашлось; «загрузите данные» не пишется.

Данные вымышленные. Запуск:
python3 -m pytest tests/test_credit_repayment_from_loaded_data.py -q
"""
from __future__ import annotations

import datetime
import io
import sys
from pathlib import Path

import pytest

ROOT = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(ROOT))
sys.path.insert(0, str(ROOT / "tests"))

import developaid_credit_repayment as credit  # noqa: E402
import developaid_credit_sources as sources  # noqa: E402
import developaid_management_report as report  # noqa: E402
import developaid_monitor as monitor  # noqa: E402
from market_search import sales_store  # noqa: E402
from test_management_report_model import _contract, _rss_book  # noqa: E402

M = credit.M
PROJECT = "Кутузов Сити"
SALES = "ЖК Кутузов"
RSS_DAY = "2026-10-18"
# Выдачи: в 2025 лимит НКЛ 2025 — 2 600, в 2026 — 5 000; НКЛ 2026 открыта с 26.06.2026.
DRAWS = [(datetime.date(2025, 5, 10), 1000 * M), (datetime.date(2025, 12, 10), 1500 * M),
         (datetime.date(2026, 3, 10), 1000 * M), (datetime.date(2026, 8, 10), 1800 * M)]


@pytest.fixture(autouse=True)
def _isolated(tmp_path, monkeypatch):
    monkeypatch.setattr(monitor, "_SNAPSHOT_DIR", tmp_path / "monitor")
    monkeypatch.setenv("DATA_DIR", str(tmp_path / "data"))


def _months(first: str, last: str) -> list[str]:
    out, year, month = [], int(first[:4]), int(first[5:7])
    while f"{year}-{month:02d}" <= last:
        out.append(f"{year}-{month:02d}")
        month += 1
        if month > 12:
            year, month = year + 1, 1
    return out


def _rss() -> bytes:
    from openpyxl import load_workbook

    book = load_workbook(io.BytesIO(_rss_book(DRAWS + [(datetime.date(2026, 9, 1), 7 * M,
                                                         "3.1")])))
    paid = book["Реестр платежей"]
    for r in range(10, paid.max_row + 1):
        if paid.cell(row=r, column=10).value is not None:
            own = paid.cell(row=r, column=7).value == "3.1"
            paid.cell(row=r, column=11, value="Собственные средства" if own
                      else "Заемные средства")
    blob = io.BytesIO()
    book.save(blob)
    return blob.getvalue()


SCHEDULE = {m: 40 * M for m in _months("2025-06", "2027-06")}


def _sales(credit_plan: bool = True) -> None:
    row = _contract("2025-06-10", "Квартира", 1000.0, 1000 * M)
    row["contract"] = "ДДУ-1"
    row["escrow_schedule"] = [{"month": m, "amount": v} for m, v in SCHEDULE.items()]
    parts = {"contracting": {"project": SALES, "missing": [], "rows": [row]}}
    if credit_plan:
        months = _months("2025-05", "2029-12")
        pf = {m: min(6400 * M, 1000 * M + 200 * M * i) for i, m in enumerate(months)}
        escrow = {m: 60 * M * (i + 1) for i, m in enumerate(months)}
        parts["credit"] = {"sheet": "КРЕДИТЫ", "queues": [{
            "queue": "Кутузов", "escrow": escrow, "pf": pf, "coverage": {},
            "rate": {m: 0.2 for m in months}, "drawn_from": "2025-05", "repay_from": "2027-09"}]}
    sales_store.save(sales_store.default_dir(), SALES, parts, "test.xlsx")


def _project(credit_plan: bool = True) -> None:
    monitor.store_estimate(PROJECT, _rss(), RSS_DAY)
    _sales(credit_plan)
    report.link_sales(PROJECT, SALES, "проверка")


def _run(**extra):
    return credit.for_project(PROJECT, SALES, end="2026-10-20", commissioning={
        "date": "2027-09-15", "source": "прогноз Монитора"}, **extra)


def test_the_project_finds_its_loan_terms_by_explicit_binding():
    assert sources.preset_for(PROJECT) == "kutuzov_city"
    assert sources.preset_for("Другой проект") is None
    got = credit.for_project("Другой проект")
    assert got["state"] == "missing"
    assert got["reason"].startswith("Условия кредитов проекта «Другой проект» в системе не "
                                    "заведены")


def test_debt_draws_and_interest_come_from_the_rss_and_the_terms():
    _project()
    got = _run()
    assert got["state"] == "ok"
    base = got["base"]
    loans = {loan["id"]: loan for loan in base["loans"]}
    # 1000 + 1500 + 1000 на НКЛ 2025, август: 1500 до лимита 5 000, остаток 300 — на НКЛ 2026.
    assert loans["nkl2025"]["debt_now"] == pytest.approx(5000 * M)
    assert loans["nkl2026"]["debt_now"] == pytest.approx(300 * M)
    assert base["top"]["debt_now"]["value"] == pytest.approx(5300 * M)
    assert "реестр платежей РСС от 18.10.2026" in base["top"]["debt_now"]["inputs"][0]
    # Отложенные проценты за полтора года — досчитаны, а не ноль и не «нет данных».
    data = {d["name"]: d for d in base["datasets"]}
    assert data["deferred.nkl2025"]["source"] == "досчитано от первой выдачи по условиям договора"
    assert any("РСС не называет, какой НКЛ платил" in note for note in base["notes"])
    # Прогнозные выдачи — план ПФ финмодели сверх выбранного (до 6 400 в сумме).
    drawn = sum(d["amount"] for name in ("plan.nkl2025", "plan.nkl2026")
                for d in credit.inputs_for(PROJECT, SALES, end="2026-10-20")["datasets"]
                [name]["data"]["draws"])
    assert drawn == pytest.approx(1100 * M)


def test_escrow_comes_from_contract_schedules_and_new_sales_only_above_them():
    _project()
    base = _run()["base"]
    on_rss = sum(v for m, v in SCHEDULE.items() if m <= RSS_DAY[:7])
    assert base["top"]["escrow_now"]["value"] == pytest.approx(on_rss)
    # План эскроу модели — весь проект; в прогноз идёт только его превышение над договорами.
    plan = {"queues": [{"escrow": {"2026-11": 100 * M, "2026-12": 200 * M, "2027-01": 300 * M}}]}
    paid = [{"date": "2026-11-30", "amount": 50 * M}, {"date": "2026-12-31", "amount": 50 * M},
            {"date": "2027-01-31", "amount": 50 * M}]
    new = sources._new_sales(plan, paid, "2026-10-18")
    assert [round(x["amount"] / M) for x in new] == [50, 50, 50]
    assert sum(x["amount"] for x in new) + 150 * M == pytest.approx(300 * M)


def test_a_manual_set_is_a_correction_over_what_was_found():
    _project()
    credit.save(PROJECT, [{"name": "key_rate", "as_of": "2026-10-20", "source": "письмо банка",
                           "data": {"path": [{"from": "2025-01-01", "rate": 21.0}]}}],
                "проверка")
    base = _run()["base"]
    shown = {d["name"]: d for d in base["datasets"]}
    assert shown["key_rate"]["source"] == "поправка вручную: письмо банка"


def test_what_is_not_in_the_system_is_named_with_its_place():
    _project(credit_plan=False)
    got = _run()
    text = " ".join(got["base"]["gaps"])
    assert "нет финмодели (лист «КРЕДИТЫ»)" in text
    assert "загрузите" not in text.lower() and "не загружены" not in text
    # Без связи с отчётом о продажах — названо, где выбрать отчёт.
    bare = credit.for_project(PROJECT, "", end="2026-10-20")
    assert any("проект не связан с отчётом о продажах" in g for g in bare["base"]["gaps"])


def test_the_report_block_reads_the_loaded_data():
    _project()
    built = report.build(PROJECT, "week", "2026-10-18")
    block = built["credit"]
    assert block["state"] == "ok"
    page = report.render_html(built)
    assert "Данные по кредитам не загружены" not in page
    assert "Текущий совокупный долг" in page
