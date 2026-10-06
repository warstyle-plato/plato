"""Итог объекта вне ДДУ: свой капитал, поток, прибыль, IRR, NPV, окупаемость.

Владелец (05.10.2026): у объекта вне ДДУ не было СВОЕГО итога, а помесячные
ряды книги шли одним блоком на все объекты. Проверки держат обещания:

* итог считает движок один раз (`object_result` в `finance.nonres`), таблица
  `nonres_report` его только читает — страница, PDF и книга печатают её;
* сумма итогов объектов сходится с тем, что нежильё внесло в проект
  (`nonres_reconciliation`), и сверка ловит подделку;
* при удержании IRR показан с оценкой и без неё, денежный итог — без оценки;
* в книге у каждого объекта свои помесячные ряды, и их сумма — итог;
* страница рисует блок «Итог объекта» числами движка (стенд на node).

Запуск: python3 -m pytest tests/test_nonres_object_result.py -q
"""

from __future__ import annotations

import copy
import io
import json
import re
import sys
from pathlib import Path

import pytest

ROOT = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(ROOT))
sys.path.insert(0, str(ROOT / "tests"))

import main_legacy as core  # noqa: E402
import page_blocks  # noqa: E402

TWO_OBJECTS = dict(
    offices_enabled=True, offices_gba_sqm=40000.0, offices_saleable_sqm=24000.0,
    offices_parking_under_spaces=200, offices_parking_over_spaces=40,
    offices_parking_guest_pct=10, offices_strategy="income",
    retail_enabled=True, retail_gba_sqm=20000.0, retail_saleable_sqm=14000.0,
    retail_strategy="direct",
)


def _spec(**over):
    x = copy.deepcopy(core.DEFAULT_INPUTS)
    t = copy.deepcopy(core.TEP_DEFAULT)
    x.update(TWO_OBJECTS)
    x.update(over)
    for obj in core.STANDALONE_OBJECTS:
        gba = float(x.get(f"{obj.prefix}_gba_sqm") or 0)
        if not x.get(f"{obj.prefix}_enabled") or gba <= 0:
            continue
        saleable = float(x.get(f"{obj.prefix}_saleable_sqm") or 0)
        t.setdefault(obj.key, {}).update(gns=gba, total_area=round(gba * 0.94, 2),
                                         useful=saleable, saleable=saleable)
    return x, t


_CACHE: dict[tuple, dict] = {}


def _run(**over) -> dict:
    key = tuple(sorted(over.items()))
    if key not in _CACHE:
        x, t = _spec(**over)
        _CACHE[key] = core._run_authoritative_model(x, t, [], {})["consolidated"]
    return _CACHE[key]


def _objects(result: dict) -> dict[str, dict]:
    return {o["key"]: o for o in result["finance"]["nonres"]["objects"]}


def _table(result: dict, key: str) -> dict[str, dict]:
    item = next(i for i in result["report"]["nonres_strategy"] if i["key"] == key)
    return {row["label"]: row for row in item["rows"]}


# --- движок -------------------------------------------------------------------

def test_each_object_has_its_own_result() -> None:
    objects = _objects(_run())
    assert set(objects) == {"offices", "standalone_retail"}
    for item in objects.values():
        result = item["result"]
        assert result["equity_invested"] == pytest.approx(
            result["cost_total"] - item["totals"]["loan_draw"], rel=1e-9)
        assert result["profit_after_tax"] == pytest.approx(
            result["profit_before_tax"] - result["profit_tax"], rel=1e-9)
        assert result["irr"] is not None
        assert result["discount_rate"] == pytest.approx(
            core.DEFAULT_INPUTS["discount_rate_pct"] / 100)


def test_the_sum_of_objects_matches_what_the_project_took() -> None:
    result = _run()
    rec = result["finance"]["nonres"]["reconciliation"]
    assert rec["ok"] is True
    objects = _objects(result)
    rows = result["finance"]["rows"]
    project_equity = (sum(float(r.get("nonres_cash_to_equity") or 0) for r in rows)
                      - sum(o["result"]["capex"] + o["result"]["common_capex"]
                            for o in objects.values()))
    assert sum(o["result"]["equity_pre_tax"] for o in objects.values()) == pytest.approx(
        project_equity, abs=1.0)
    base = (sum(result["finance"]["tax_margin_by_product"][k] for k in objects)
            - sum(float(r.get("nonres_loan_interest") or 0) + float(r.get("nonres_loan_fee") or 0)
                  for r in rows))
    assert sum(o["result"]["profit_before_tax"] + o["result"]["common_recognized"]
               for o in objects.values()) == pytest.approx(base, abs=1.0)


def test_the_reconciliation_catches_a_forged_object_result() -> None:
    """Подделка: итог объекта посчитан не из того, что вошло в проект."""
    x, t = _spec()
    op = core.build_operating_model(x, t, [])
    nonres = core.nonres_overlay(x, [], op)
    rows = [{"month": m.isoformat(), **{k: v.get(m, 0.0) for k, v in nonres["monthly"].items()}}
            for m in sorted({m for v in nonres["monthly"].values() for m in v})]
    margins = {k: sum(f["monthly"].get("tax_margin", {}).values())
               for k, f in nonres["objects"].items()}
    assert core.nonres_reconciliation(nonres, rows, margins)["ok"] is True
    forged = copy.deepcopy(nonres)
    forged["objects"]["offices"]["result"]["equity_pre_tax"] += 1_000_000.0
    assert core.nonres_reconciliation(forged, rows, margins)["ok"] is False
    report = core.nonres_report(core.nonres_summary(
        forged, core.nonres_reconciliation(forged, rows, margins)))
    assert any("не сходится с вкладом нежилья" in w for w in report[0]["warnings"])


def test_ddu_objects_have_no_result_and_no_table() -> None:
    result = _run(offices_strategy="ddu", retail_strategy="ddu")
    assert result["report"]["nonres_strategy"] == []
    assert result["finance"]["nonres"]["objects"] == []


def test_the_table_prints_the_engine_result() -> None:
    result = _run()
    for key, item in _objects(result).items():
        rows = _table(result, key)
        assert rows["Итог объекта"]["unit"] == "section"
        res = item["result"]
        assert rows["Собственный капитал объекта (затраты − кредит объекта)"]["value"] == res["equity_invested"]
        assert rows["Прибыль объекта до налога"]["value"] == res["profit_before_tax"]
        assert rows["Прибыль объекта после налога"]["value"] == res["profit_after_tax"]
        assert rows["Денежный поток капитала объекта — итог"]["value"] == res["equity_cash"]
        assert rows["IRR капитала объекта"]["value"] == res["irr"]
        assert rows["NPV капитала объекта @20%"]["value"] == res["npv"]
        payback = rows["Срок окупаемости капитала"]["value"]
        assert (payback.startswith(f"{res['payback_months']} мес.") if res["payback_months"] is not None
                else payback == res["payback_reason"])


def test_hold_separates_the_valuation_from_the_cash() -> None:
    result = _run(offices_exit_mode="hold")
    office = _objects(result)["offices"]["result"]
    rows = _table(result, "offices")
    assert office["residual_value"] > 0
    assert office["equity_with_value"] == pytest.approx(
        office["equity_cash"] + office["residual_value"], rel=1e-9)
    assert rows["Оценка объекта на конец горизонта — без сделки"]["value"] == office["residual_value"]
    assert "IRR капитала объекта — с оценкой (часть результата без сделки)" in rows
    cash_irr = rows["IRR капитала объекта — только деньги, без оценки"]
    if office["irr_cash"] is None:
        assert cash_irr["unit"] == "text" and cash_irr["value"].startswith("не считается")
    assert rows["Денежный поток капитала объекта — итог без оценки"]["value"] == office["equity_cash"]
    assert "Срок окупаемости капитала — без оценки" in rows
    # Продажа — ни строки оценки, ни «без оценки».
    sale = _table(_run(), "offices")
    assert not any("оценк" in label for label in sale if label.startswith(("IRR", "Денежный")))


def test_a_nonresidential_project_charges_the_object_its_share_of_common_costs() -> None:
    x, t = _spec()
    for key in core.MKD_PRODUCTS:
        for col in ("gns", "total_area", "useful", "saleable", "transfer", "units"):
            if key in t:
                t[key][col] = 0
    x["project_kind"] = core.PROJECT_KIND_NONRESIDENTIAL
    result = core._run_authoritative_model(x, t, [], {})["consolidated"]
    assert result["finance"]["nonres"]["reconciliation"]["ok"] is True
    for key, item in _objects(result).items():
        assert item["result"]["common_capex"] > 0
        rows = _table(result, key)
        assert rows["в т.ч. доля общих затрат проекта (участок, проект, сети)"]["value"] == \
            item["result"]["common_capex"]


# --- книга -------------------------------------------------------------------

def test_the_book_has_monthly_rows_per_object_that_add_up_to_the_total() -> None:
    import openpyxl
    x, t = _spec()
    content, _name, _meta = core.build_project_workbook(x, t, [], {})
    sheet = openpyxl.load_workbook(io.BytesIO(content))["Нежильё — стратегия"]
    blocks: dict[str, dict[str, float]] = {}
    title = None
    for r in range(1, sheet.max_row + 1):
        a = sheet.cell(r, 1).value
        if isinstance(a, str) and a.startswith("Помесячно — "):
            title = a.removeprefix("Помесячно — ")
            blocks[title] = {}
        elif title and isinstance(a, str) and re.fullmatch(r"\d{4}-\d{2}", a):
            for c in range(2, sheet.max_column + 1):
                header = sheet.cell(next(h for h in range(r, 0, -1)
                                         if sheet.cell(h, 1).value == "Месяц"), c).value
                if header:
                    blocks[title][header] = blocks[title].get(header, 0.0) + float(sheet.cell(r, c).value or 0)
    assert set(blocks) == {"все объекты вне ДДУ вместе", "Офисы", "Коммерция ОСЗ"}
    total = blocks.pop("все объекты вне ДДУ вместе")
    for header, value in total.items():
        assert sum(b.get(header, 0.0) for b in blocks.values()) == pytest.approx(value, rel=1e-9, abs=1.0)
    objects = _objects(_run())
    flows = sum(o["result"]["equity_cash"] for o in objects.values())
    assert sum(b["Поток капитала объекта после налога, ₽"] for b in blocks.values()) == pytest.approx(
        flows, rel=1e-9)


# --- страница ----------------------------------------------------------------

def _render(report: list[dict]) -> list[dict]:
    prelude = """
const BOX={innerHTML:''}, CARD={hidden:true};
const document={getElementById:id=>id==='nonresStrategyCard'?CARD:BOX};
const R=%s;
""" % json.dumps({"report": {"nonres_strategy": report}}, ensure_ascii=False)
    tail = """
renderNonresStrategy(R);
const tables=[...BOX.innerHTML.matchAll(/<table[^>]*data-object="([^"]*)"[^>]*>([\\s\\S]*?)<\\/table>/g)]
 .map(m=>({object:m[1],rows:[...m[2].matchAll(/<tr( class="section")?><td>([\\s\\S]*?)<\\/td><td>([\\s\\S]*?)<\\/td><\\/tr>/g)]
   .map(r=>({section:!!r[1],label:r[2],cell:r[3]}))}));
const expected=R.report.nonres_strategy.map(o=>o.rows.map(r=>nonresCell(r)));
console.log(JSON.stringify({hidden:CARD.hidden,tables,expected,
 fmt:{pct:nonresCell({unit:'pct',value:0.1234})}}));
"""
    out, _ = page_blocks.run(prelude, tail)
    return json.loads(out)


def _mismatches(drawn: dict, report: list[dict]) -> list[str]:
    """Чем нарисованное расходится с таблицей движка; пусто — сошлось."""
    out: list[str] = []
    if len(drawn["tables"]) != len(report):
        return [f"таблиц {len(drawn['tables'])} вместо {len(report)}"]
    for table, item, cells in zip(drawn["tables"], report, drawn["expected"]):
        if table["object"] != item["key"]:
            out.append(f"объект {table['object']} вместо {item['key']}")
        labels = [r["label"] for r in table["rows"]]
        if labels != [r["label"] for r in item["rows"]]:
            out.append(f"{item['key']}: строки не те")
        if [r["cell"] for r in table["rows"]] != cells:
            out.append(f"{item['key']}: значения не те")
        sections = [r["label"] for r in table["rows"] if r["section"]]
        if sections != [r["label"] for r in item["rows"] if r["unit"] == "section"]:
            out.append(f"{item['key']}: блок «Итог объекта» не выделен")
    return out


def test_the_page_draws_the_object_result_block() -> None:
    report = _run()["report"]["nonres_strategy"]
    drawn = _render(report)
    assert drawn["hidden"] is False
    assert _mismatches(drawn, report) == []
    assert drawn["fmt"]["pct"] == "12,3%"
    for table, item in zip(drawn["tables"], report):
        rows = {r["label"]: r for r in table["rows"]}
        assert rows["Итог объекта"]["section"] and rows["Итог объекта"]["cell"] == ""
        irr = next(r for r in item["rows"] if r["label"] == "IRR капитала объекта")
        shown = f"{irr['value'] * 100:.1f}".rstrip("0").rstrip(".").replace(".", ",") + "%"
        assert rows["IRR капитала объекта"]["cell"] == shown


def test_the_page_check_fails_on_a_forgery() -> None:
    """Стенд падает на подделке: нарисованное с другим числом не сходится."""
    report = _run()["report"]["nonres_strategy"]
    drawn = _render(report)
    forged = copy.deepcopy(report)
    row = next(r for r in forged[0]["rows"] if r["label"] == "Прибыль объекта до налога")
    row["value"] = float(row["value"]) + 5e8
    forged_drawn = _render(forged)
    # Подделка движка видна сравнением с правдой.
    assert [r["cell"] for r in forged_drawn["tables"][0]["rows"]] != \
        [r["cell"] for r in drawn["tables"][0]["rows"]]
    # Подделка рисунка: блок без выделения и с чужим числом — сверка ловит оба.
    tampered = copy.deepcopy(drawn)
    tampered["tables"][0]["rows"][[r["label"] for r in tampered["tables"][0]["rows"]]
                                  .index("Итог объекта")]["section"] = False
    tampered["tables"][1]["rows"][-1]["cell"] = "0 ₽"
    assert len(_mismatches(tampered, report)) == 2
