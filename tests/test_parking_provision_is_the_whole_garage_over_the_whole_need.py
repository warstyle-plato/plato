"""Обеспеченность парковкой — весь гараж проекта на всю нормативную потребность.

Решение владельца (#656, 11.10.2026): числитель — все машино-места проекта,
знаменатель — вся нормативная потребность: жители (постоянные и гостевые),
встроенная коммерция, отдельные объекты и соцобъекты. Прежде 1 199 мест гаража
делились на 97 мест встроенных магазинов, и выходило 12,36× — числитель и
знаменатель относились к разным потребителям.

Одно число на всех поверхностях: делит `parking_provision`, а страница, отчёт,
тизер, Платон, сводка в чат и книга читают его ответ с подписью
«N мест / M по норме = K×».
"""

from __future__ import annotations

import copy
import io
import sys
from pathlib import Path

import pytest

sys.path.insert(0, str(Path(__file__).resolve().parent.parent))
sys.path.insert(0, str(Path(__file__).resolve().parent))

import main as _wrapper  # noqa: E402
import page_blocks  # noqa: E402
import v4_dashboard  # noqa: E402

core = _wrapper.core


def _shape():
    return copy.deepcopy(core.DEFAULT_INPUTS), {k: dict(v) for k, v in core.TEP_DEFAULT.items()}


@pytest.fixture(scope="module")
def bundle():
    inputs, tep = _shape()
    return core._run_authoritative_model(inputs, tep, [], None)


def test_the_denominator_is_the_whole_need_and_not_the_shops(bundle):
    """Проект из #550: 1 199 мест — на всю потребность, а не на 97 мест магазинов."""
    parking = bundle["consolidated"]["parking"]
    need = parking["project_required"]
    provision = parking["provision"]
    built = core.tep_parking_spaces(bundle["consolidated"]["tep"]["rows"])
    attached = float(parking["required_total"])
    # Предохранитель: стенд тот же, что у #550, иначе проверка ничего не ловит.
    assert built == 1199.0 and attached == 97.0, (built, attached)
    old_ratio = built / attached
    assert round(old_ratio, 2) == 12.36
    # Прежний знаменатель — красный: обеспеченность не 12,36×.
    assert provision["required"] != attached
    assert abs(provision["ratio"] - old_ratio) > 1.0, provision
    # Знаменатель — сумма частей, и в ней есть все потребители.
    parts = {part["key"]: part["spaces"] for part in need["parts"]}
    assert set(parts) == {"residents", "attached", "social"}
    assert parts["attached"] == attached
    assert parts["residents"] == need["residents_permanent"] + need["residents_guest"] > 0
    assert parts["social"] > 0, "ДОО стенда — соцобъект с нормативом 945-ПП"
    assert provision["required"] == need["total"] == sum(parts.values())
    assert provision["ratio"] == pytest.approx(built / need["total"])


def test_the_label_names_both_quantities(bundle):
    provision = bundle["consolidated"]["parking"]["provision"]
    built = f"{int(provision['built']):,}".replace(",", " ")
    need = f"{int(provision['required']):,}".replace(",", " ")
    ratio = f"{provision['ratio']:.2f}".replace(".", ",")
    assert provision["label"] == f"{built} мест / {need} по норме = {ratio}×"


def test_a_missing_need_is_not_a_ratio():
    """Нет потребности — нет и обеспеченности: ноль в знаменателе не ответ."""
    got = core.parking_provision(120.0, {"total": 0})
    assert got["ratio"] is None
    assert "не посчитана" in got["label"]


def test_moscow_oblast_names_the_social_gap():
    """У области норматива соцобъектов в модели нет — пробел назван, а не ноль."""
    inputs, tep = _shape()
    inputs["vri_region"] = "mo"
    demand = core.apply_object_parking(inputs, copy.deepcopy(tep))
    need = demand["project_required"]
    assert need["social"] == 0
    assert any("ДОО" in line and "не заведён" in line for line in need["missing"]), need


def test_presentation_and_teaser_read_the_engine(bundle):
    inputs, tep = _shape()
    provision = bundle["consolidated"]["parking"]["provision"]
    numbers = core.presentation_numbers(bundle["consolidated"])
    assert numbers["parking"]["provision"] == provision["ratio"]
    assert numbers["parking"]["required_total"] == provision["required"]
    assert numbers["parking"]["provision_label"] == provision["label"]
    model = core.project_presentation(bundle, inputs, tep, None)
    assert model["tep"]["parking_provision_label"] == provision["label"]
    from pypdf import PdfReader
    pdf = core.build_teaser_pdf(bundle, inputs, tep, None)
    text = " ".join(page.extract_text() for page in PdfReader(io.BytesIO(pdf)).pages)
    assert "Обеспеченность парковкой" in text
    assert provision["label"].split(" = ")[-1] in text


def test_the_report_pdf_prints_the_same_label(bundle):
    pytest.importorskip("reportlab", reason="reportlab нужен только для PDF")
    from market_search.krt_requirements import pdf_text

    inputs, tep = _shape()
    pdf = core._build_developaid_pdf({"result": bundle["consolidated"], "inputs": inputs,
                                      "tep": tep, "project_name": "Обеспеченность"})
    text = " ".join(pdf_text(pdf).split())
    label = " ".join(bundle["consolidated"]["parking"]["provision"]["label"]
                     .replace(" ", " ").split())
    assert "Обеспеченность парковкой: " + label in text.replace(" ", " ")


def test_plato_sees_the_same_label(bundle):
    snapshot = core._result_snapshot(bundle["consolidated"])
    assert snapshot["parking_provision"] == bundle["consolidated"]["parking"]["provision"]["label"]


def test_the_page_reads_the_engine_label():
    """Страница не делит сама: плашка и сводка в чат берут подпись расчёта."""
    prelude = ("let phaseBundle=null;"
               "let lastResult={parking:{provision:{label:'1 199 мест / 1 308 по норме = 0,92×'}}};")
    out, _ = page_blocks.run(prelude, "console.log(parkingProvisionText())")
    assert out.strip() == "1 199 мест / 1 308 по норме = 0,92×"
    prelude = ("let phaseBundle={mode:'phased',consolidated:{parking:{provision:{label:'свод'}}}};"
               "let lastResult={parking:{provision:{label:'очередь'}}};")
    out, _ = page_blocks.run(prelude, "console.log(parkingProvisionText())")
    assert out.strip() == "свод", "обеспеченность — величина проекта, а не вкладки"


def test_the_telegram_card_prints_the_label(monkeypatch):
    sent: list[str] = []
    monkeypatch.setattr(core, "_telegram_verify_session", lambda s: {"chat_id": 42, "cad": []})
    monkeypatch.setattr(core, "_telegram_user_allowed", lambda c: True)
    monkeypatch.setattr(core, "_telegram_send_message", lambda chat_id, text, **kw: sent.append(text))
    monkeypatch.setattr(core, "_telegram_web_app_url", lambda *a, **k: "https://example.org/")
    core.telegram_result(core.TelegramResultRequest(session="s", summary={
        "revenue_mln": 100, "total_expenses_mln": 80, "parking_spaces": 1199,
        "parking_provision_label": "1 199 мест / 1 308 по норме = 0,92×",
    }))
    assert sent and "обеспеченность парковкой — 1 199 мест / 1 308 по норме = 0,92×" in sent[0]


def test_the_book_divides_its_own_garage_by_the_engine_need(bundle):
    """Книга: места и потребность — числа движка, отношение — формула книги."""
    import openpyxl
    from xlsx_eval import Evaluator

    inputs, tep = _shape()
    sys.setrecursionlimit(400000)
    content, _, meta = core.build_project_workbook(inputs, tep, [], None)
    assert meta["missing"] == [], meta["missing"]
    evaluator = Evaluator(openpyxl.load_workbook(io.BytesIO(content), data_only=False))
    provision = bundle["consolidated"]["parking"]["provision"]
    sheet = v4_dashboard.DATA_SHEET
    need = evaluator.cell(sheet, f"C{v4_dashboard.DATA_ROWS['parking_required_total']}")
    ratio = evaluator.cell(sheet, f"C{v4_dashboard.DATA_ROWS['parking_provision']}")
    assert float(need) == provision["required"]
    assert float(ratio) == pytest.approx(provision["ratio"], abs=1e-6)
    # Прежний знаменатель книга тоже не несёт.
    assert abs(float(ratio) - provision["built"] / 97.0) > 1.0
