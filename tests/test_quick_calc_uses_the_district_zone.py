"""Быстрый расчёт (`vri_tep_quick`, дубль калькулятора ГлавАПУ) считает ДОО и
школы по расчётной зоне района, а не по зоне 1 для всех.

2151-ПП, табл. 1.4.1/1.4.2: зона 1 — 44 места ДОО и 90 мест СОШ на 1 000
жителей, зона 2 (28 районов, п. 1.3.2.2 в ред. 1762-ПП) — 63 и 124. Прежде в
быстром расчёте стояли литералы 44/90 для любого района — Некрасовка считалась
как Пресня. Решение владельца 16.10.2026
(docs/normative_audit_msk_mo_2026-10.md, § 6 п. 5): один владелец чисел —
`SOCIAL_NORMS_PER_1000` и предикат `district_zone_two`, как в
`tep_derived_norms`.

Числа здесь намеренно литералами: проверка должна падать на старом коде, а
не повторять словарь движка.
"""

from __future__ import annotations

import io
import math
import sys
from pathlib import Path

import openpyxl
import pytest

sys.path.insert(0, str(Path(__file__).resolve().parent.parent))

import main as wrapper  # noqa: E402

core = wrapper.core


def _quick(monkeypatch, district: str | None) -> dict:
    territory = {"area_ha": 1.4968, "cadastral_quarter": "77:03:0006003"}
    if district is not None:
        territory["district"] = district
    monkeypatch.setattr(core, "analyze_cadastral_territory", lambda req: {
        "territory": territory,
        "coefficients": {"rail": 0.75, "business_outside_ttc": 0.5,
                         "rent": 0.1281, "base_cost_zh_high": 229036.29},
    })
    return core.vri_tep_quick("msk", "77:03:0006003:98")


def _tep(result: dict) -> dict[str, int]:
    sheet = openpyxl.load_workbook(io.BytesIO(result["file"]))["ТЭП"]
    labels = {str(sheet.cell(row=r, column=1).value): r for r in range(2, sheet.max_row + 1)}
    cell = lambda code: int(str(sheet.cell(row=labels[code], column=4).value).replace("\xa0", ""))
    return {"population": cell("4"), "dou": cell("30"), "school": cell("31"), "clinic": cell("32")}


def _params(result: dict) -> dict[str, tuple[str, str]]:
    sheet = openpyxl.load_workbook(io.BytesIO(result["file"]))["Параметры территории"]
    return {str(row[0]): (str(row[1] or ""), str(row[2] or ""))
            for row in sheet.iter_rows(min_row=2, values_only=True) if row and row[0]}


def _norm_rows(params: dict[str, tuple[str, str]]) -> dict[str, tuple[str, str, str]]:
    """Строки нормативов по началу подписи — так же их ищет импорт выгрузки."""
    out = {}
    for label, (value, unit) in params.items():
        for key in ("Норматив ДОО", "Норматив школ"):
            if label.startswith(key):
                out[key] = (label, value, unit)
    return out


def test_a_zone_two_district_takes_63_and_124(monkeypatch):
    """Некрасовка — зона 2: 63/124 на тысячу, не 44/90. На старом коде
    (население 970) выходило 43/88 — здесь должно быть 62/121."""
    result = _quick(monkeypatch, "Некрасовка")
    tep = _tep(result)
    population = tep["population"]
    assert population == 970, "точка та же, что и у смешанной поликлиники"
    assert tep["dou"] == math.ceil(population * 63 / 1000) == 62
    assert tep["school"] == math.ceil(population * 124 / 1000) == 121
    # Смешанная поликлиника — 19 в обеих зонах.
    assert tep["clinic"] == math.ceil(population * 19 / 1000) == 19

    params = _params(result)
    norms = _norm_rows(params)
    label, value, unit = norms["Норматив ДОО"]
    assert value == "63" and "зона 2" in label and unit == "мест / 1000 жит."
    label, value, unit = norms["Норматив школ"]
    assert value == "124" and "зона 2" in label
    assert params["Расчётная зона"][0] == "2"


def test_a_zone_one_district_keeps_44_and_90(monkeypatch):
    """Головинский — зона 1: нормы прежние, подпись называет зону."""
    result = _quick(monkeypatch, "Головинский")
    tep = _tep(result)
    population = tep["population"]
    assert tep["dou"] == math.ceil(population * 44 / 1000) == 43
    assert tep["school"] == math.ceil(population * 90 / 1000) == 88

    params = _params(result)
    norms = _norm_rows(params)
    assert norms["Норматив ДОО"][1] == "44" and "зона 1" in norms["Норматив ДОО"][0]
    assert norms["Норматив школ"][1] == "90" and "зона 1" in norms["Норматив школ"][0]
    assert params["Расчётная зона"][0] == "1"


def test_the_zone_follows_the_engine_predicate(monkeypatch):
    """Решение о зоне — одно на движок: «ё» в названии не делает район зоной 1,
    как и в `district_zone_two`."""
    with_yo = _tep(_quick(monkeypatch, "Бирюлёво Восточное"))
    plain = _tep(_quick(monkeypatch, "Бирюлево Восточное"))
    assert with_yo == plain
    assert with_yo["dou"] == 62, "район зоны 2 через «ё» прочитан как зона 1"


def test_the_export_reimports_with_the_zone(monkeypatch):
    """Выгрузка быстрого расчёта читается импортом ГлавАПУ той же зоной и теми
    же нормами: зона в скобках подписи не ломает поиск по началу имени."""
    result = _quick(monkeypatch, "Солнцево")
    normalized = core.parse_glavapu_xlsx(result["file"], result["filename"])["normalized"]
    assert str(normalized["calculation_zone"]).strip() == "2"
    assert normalized["kindergarten_norm_per_1000"] == pytest.approx(63)
    assert normalized["school_norm_per_1000"] == pytest.approx(124)


def test_no_district_is_zone_one_and_says_so(monkeypatch):
    """Район не назван — зона 1 по умолчанию, и лист параметров говорит об
    этом вслух: неизвестный район и район зоны 1 на экране иначе неотличимы."""
    result = _quick(monkeypatch, None)
    params = _params(result)
    assert params["Расчётная зона"][0] == "1 (район не назван)"
    assert _norm_rows(params)["Норматив ДОО"][1] == "44"
