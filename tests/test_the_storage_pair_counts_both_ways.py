"""Штуки и метры кладовых — одна величина в двух видах.

Жили они независимо, и владелец увидел это первым же вводом (14.09.2026:
«ввёл в ТЭП 800 штук — ничего не изменилось в площади и наоборот»). Замер на
живых функциях страницы подтвердил дословно: 800 штук давали 0 м², 8 000 м²
давали 0 штук. Цена обеих половин на умолчаниях посчитана и стоит в CLAUDE.md:
800 кладовых без метров — выручка без СМР (LLCR 0,9656 → 0,9893), 8 000 м² без
штук — СМР без выручки (0,9333), верная пара 800 × 4,3 — 0,9750. Настоящий
ответ лежал МЕЖДУ двумя молчаливыми ошибками, и ни одна ошибкой не выглядела.

Ответ на пару обязан быть один на три поверхности: движок приводит строку,
страница считает её тем же правилом, книга считает метры формулой от штук —
иначе экран, отчёт и книга разойдутся молча, как уже расходились на паре
«места ↔ площадь» гаража.

Запуск: python3 -m pytest tests/test_the_storage_pair_counts_both_ways.py -q
"""

from __future__ import annotations

import copy
import io
import json
import subprocess
import sys
from pathlib import Path

import openpyxl
import pytest

sys.path.insert(0, str(Path(__file__).resolve().parent.parent))
sys.path.insert(0, str(Path(__file__).resolve().parent))

import main_legacy as core  # noqa: E402
import page_blocks  # noqa: E402
from xlsx_eval import Evaluator  # noqa: E402


def _tep(**storage):
    tep = copy.deepcopy(core.TEP_DEFAULT)
    tep["storage"] = dict(tep["storage"], **storage)
    return tep


def _applied(units=0.0, gns=0.0, inputs=None):
    """Строка кладовых, приведённая движком к вводным."""
    x = dict(core.DEFAULT_INPUTS)
    x.update(inputs or {})
    tep = _tep(units=units, gns=gns)
    core.apply_storage_tep_row(x, tep)
    return tep["storage"]


def test_units_bring_their_metres():
    """Вписали штуки — площадь считается нормативом, а не остаётся нулём."""
    per = core.storage_area_per_unit(core.DEFAULT_INPUTS)
    row = _applied(units=800)
    assert row["gns"] == pytest.approx(800 * per, abs=0.1)
    assert row["total_area"] == pytest.approx(row["gns"], abs=0.1)


def test_metres_bring_their_units():
    """И обратно: вписали площадь — считаются штуки. Кладовая неделима."""
    per = core.storage_area_per_unit(core.DEFAULT_INPUTS)
    row = _applied(gns=8000)
    assert row["units"] == pytest.approx(round(8000 / per))
    # Заданная руками площадь сильнее норматива — её не перетирают округлённым
    # произведением: это правило пары «места ↔ площадь» гаража.
    assert row["gns"] == pytest.approx(8000)


def test_a_hand_made_pair_is_left_alone():
    """Заданы оба числа — это решение человека, и трогать его нечем."""
    row = _applied(units=800, gns=6000)
    assert (row["units"], row["gns"]) == (800, 6000)


def test_the_metres_reach_the_underground_area():
    """Метры кладовых доезжают до подземной площади, а не считаются отдельно.

    Без этого пара была бы косметикой: владелец смотрит на «Итого · МКД», а
    не на строку.
    """
    def under(units):
        result = core._run_authoritative_model(
            dict(core.DEFAULT_INPUTS), _tep(units=units), [], {})
        return result["consolidated"]["summary"]["underground_gns_sqm"]

    per = core.storage_area_per_unit(core.DEFAULT_INPUTS)
    grew = under(800) - under(0)
    assert grew == pytest.approx(800 * per, abs=1.0)


def test_the_page_answers_exactly_as_the_engine():
    """Страница считает ту же пару — настоящим своим кодом, а не пересказом."""
    stand = page_blocks.tep_cell_stand() + """
tep={storage:{label:'Кладовые',gns:0,total_area:0,useful:0,saleable:0,transfer:0,units:0}};
tepCellChanged('storage','units',800);
const byUnits={units:tep.storage.units,gns:tep.storage.gns};
tep.storage={label:'Кладовые',gns:0,total_area:0,useful:0,saleable:0,transfer:0,units:0};
tepCellChanged('storage','gns',8000);
console.log(JSON.stringify({byUnits,byArea:{units:tep.storage.units,gns:tep.storage.gns}}));
"""
    done = subprocess.run([
        "node", "-e", stand], capture_output=True, text=True, timeout=120)
    assert done.returncode == 0, done.stderr[-2000:]
    got = json.loads(done.stdout.strip().splitlines()[-1])
    engine_units = _applied(units=800)
    engine_area = _applied(gns=8000)
    assert got["byUnits"]["gns"] == pytest.approx(engine_units["gns"], abs=0.1)
    assert got["byArea"]["units"] == pytest.approx(engine_area["units"])
    assert got["byUnits"]["units"] == pytest.approx(engine_units["units"])


def test_the_storage_price_follows_the_class():
    """Цена кладовой классовая: смена класса её двигает.

    Прежде она была одним числом вне профиля класса — «ничего не меняется в
    настройках классов» было верным описанием того, что видно.
    """
    prices = {name: preset["storage_price_th"]
              for name, preset in core.PROJECT_CLASS_PRESETS.items()}
    assert len(set(prices.values())) == len(prices), prices
    assert prices["comfort"] < prices["business"] < prices["elite"]
    # База «Комфорта» равна умолчаниям движка — иначе расчёт на умолчаниях
    # показывал бы отклонение от класса, которого никто не задавал.
    assert prices["comfort"] == core.DEFAULT_INPUTS["storage_price_th"]


def test_the_workbook_counts_metres_from_units():
    """Правка штук ВНУТРИ книги двигает метры: книга самостоятельна.

    Числом площадь стояла мёртвой — «если величина в книге не выражается
    формулой, значит не хватает не формулы, а вводной».
    """
    per = core.storage_area_per_unit(core.DEFAULT_INPUTS)
    content, _, meta = core.build_project_workbook(
        dict(core.DEFAULT_INPUTS), _tep(units=800, gns=round(800 * per, 1)),
        [], {})
    assert not [item for item in meta.get("missing", [])
                if "кладов" in item.lower()], meta.get("missing")
    book = openpyxl.load_workbook(io.BytesIO(content), data_only=False)
    before = Evaluator(book).cell("ТЭП", "C7")
    assert before == pytest.approx(800 * per, abs=0.1)
    # Человек вписал вдвое больше кладовых прямо в книге.
    book["Параметры модели"]["AC88"] = 1600
    after = Evaluator(book).cell("ТЭП", "C7")
    assert after == pytest.approx(1600 * per, abs=0.1)
    assert after > before
