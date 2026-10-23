"""Планы продаж — с листов «Продажи П-Ф_ФМ» и «Продажи П-Ф_БАНК» книги cf.

Владелец (23.10.2026): «план продаж по банку и ФМ отсюда с вкладок Продажи П-Ф
Банк и Продажи П-Ф ФМ». Оба листа помесячные и устроены как прежний лист
финмодели, но у листа банка подписи стоят в другой колонке, шапка на строку
ниже, а рядом с месяцами лежат годовые колонки «2025» без дат. Листы берутся по
явной карте имён: «Продажи П-Ф_old» и «…_сдвиг сроков» — не планы.
"""

from __future__ import annotations

import datetime
import io
import os
import sys

import openpyxl
import pytest

sys.path.insert(0, os.path.dirname(os.path.dirname(os.path.abspath(__file__))))

from market_search import contracting  # noqa: E402

M = 1e6
MONTHS = [datetime.date(2026, m, 1) for m in range(7, 13)]
# план по месяцам (тыс. ₽ договоров, м², тыс. ₽ эскроу) — квартиры
FLATS = {7: (150_000, 200, 90_000), 8: (160_000, 210, 95_000), 9: (170_000, 220, 100_000),
         10: (180_000, 230, 110_000), 11: (190_000, 240, 120_000), 12: (200_000, 250, 130_000)}
PARKING = {m: (6_000, 14, 3_000) for m in range(7, 13)}  # 1 место = 14 м² в месяц


def _pf_sheet(sheet, label_col: int, date_row: int, kind_row: int, years: bool) -> None:
    """Лист «план/факт по месяцам» в заданной раскладке."""
    col = label_col + 2  # первая колонка чисел (1-based) сразу после подписей
    if years:
        # Годовые колонки «2025» (план/факт) перед месяцами — число, не дата.
        sheet.cell(date_row, col, 2025)
        sheet.cell(date_row, col + 1, 2025)
        sheet.cell(kind_row, col, "план")
        sheet.cell(kind_row, col + 1, "факт")
        col += 2
    for month in MONTHS:
        last = (month.replace(day=28) + datetime.timedelta(days=4)).replace(day=1) - datetime.timedelta(days=1)
        sheet.cell(date_row, col, datetime.datetime.combine(last, datetime.time()))
        sheet.cell(date_row, col + 1, datetime.datetime.combine(last, datetime.time()))
        sheet.cell(kind_row, col, "план")
        sheet.cell(kind_row, col + 1, "факт")
        col += 2
    row = kind_row + 1
    first = label_col + 2 + (2 if years else 0)  # колонка первого «план» месяца (1-based)
    for product, plan in (("Квартира", FLATS), ("Машиноместа", PARKING), ("Итого", None)):
        sheet.cell(row, label_col + 1, product)
        row += 1
        for label, pick in (("Эскроу, тыс. руб ", 2), ("м2", 1), ("шт", None),
                            ("Заключенные договоры, тыс. руб", 0), ("Цена продажи кв.м, тыс. руб", None)):
            if product == "Итого" and label not in ("Эскроу, тыс. руб ", "Заключенные договоры, тыс. руб"):
                continue
            if label == "шт" and product != "Машиноместа":
                continue
            sheet.cell(row, label_col + 1, label)
            for i, month in enumerate(MONTHS):
                c = first + 2 * i
                if product == "Итого":
                    value = FLATS[month.month][pick] + PARKING[month.month][pick]
                elif label == "шт":
                    value = 1
                elif label.startswith("Цена"):
                    value = plan[month.month][0] / plan[month.month][1]
                else:
                    value = plan[month.month][pick]
                sheet.cell(row, c, value)
                sheet.cell(row, c + 1, value if month.month <= 9 else None)
            row += 1
    sheet.cell(row + 1, label_col + 1, "Продажи по годам")
    sheet.cell(row + 3, label_col + 1, "Квартира")
    sheet.cell(row + 4, label_col + 1, "Заключенные договоры, тыс. руб")
    sheet.cell(row + 4, first, 9_999_999)  # годовой блок, который нельзя смешать с месячным


def _book(fm: bool = True, bank: bool = True, decoys: bool = True) -> bytes:
    book = openpyxl.Workbook()
    book.active.title = "Раб.файл"
    if decoys:
        _pf_sheet(book.create_sheet("Продажи П-Ф_old"), 2, 3, 4, False)
    if fm:
        _pf_sheet(book.create_sheet("Продажи П-Ф_ФМ"), 2, 3, 4, False)
    if bank:
        _pf_sheet(book.create_sheet("Продажи П-Ф_БАНК"), 1, 4, 5, True)
    out = io.BytesIO()
    book.save(out)
    return out.getvalue()


def _q4(values: dict[str, float]) -> float:
    return sum(v for m, v in values.items() if contracting.quarter_of(m) == "2026 Q4")


def test_the_fm_plan_comes_from_the_cf_sheet():
    got = contracting.read_fm_plan(_book())
    assert got["sheet"] == "Продажи П-Ф_ФМ"
    by_month = contracting.fm_plan_by_month(got)
    assert set(by_month) == {m.strftime("%Y-%m") for m in MONTHS}
    expected = sum(FLATS[m][0] + PARKING[m][0] for m in (10, 11, 12)) * 1000.0
    assert _q4({m: v["amount"] for m, v in by_month.items()}) == pytest.approx(expected)
    # Факт заполнен по сентябрь; октябрь без факта — не ноль.
    assert "2026-10" not in got["fact"]["Квартира"]


def test_the_bank_plan_is_summed_from_the_monthly_sheet():
    bank = contracting.read_bank_plan(_book())
    assert bank["sheet"] == "Продажи П-Ф_БАНК" and bank["monthly"] is True
    assert bank["quarters"] == ["2026 Q3", "2026 Q4"], "годовая колонка «2025» — не квартал 1905 года"
    gross = sum(FLATS[m][0] + PARKING[m][0] for m in (10, 11, 12)) * 1000.0
    assert bank["gross_by_quarter"]["2026 Q4"] == pytest.approx(gross)
    assert bank["revenue_by_quarter"]["2026 Q4"] == pytest.approx(
        sum(FLATS[m][2] + PARKING[m][2] for m in (10, 11, 12)) * 1000.0)
    # Метры — без машино-мест; цена — договоры квартир квартала на их метры.
    assert bank["area_by_quarter"]["2026 Q4"] == pytest.approx(sum(FLATS[m][1] for m in (10, 11, 12)))
    flats_amount = sum(FLATS[m][0] for m in (10, 11, 12)) * 1000.0
    assert bank["price_by_quarter"]["2026 Q4"] == pytest.approx(flats_amount / sum(FLATS[m][1] for m in (10, 11, 12)))
    assert bank["products"]["Машиноместо"]["volume"]["2026 Q4"] == 3, "машино-места — штуками"
    assert bank["products"]["Машиноместо"]["price"]["2026 Q4"] == pytest.approx(6_000_000.0)
    # Годовой блок ниже не удвоил план.
    assert bank["gross_by_quarter"]["2026 Q4"] < 9_999_999 * 1000.0


def test_the_escrow_plan_accumulates_the_monthly_inflow():
    bank = contracting.read_bank_plan(_book())
    plan = contracting.bank_escrow_plan(bank)
    q3 = sum(FLATS[m][2] + PARKING[m][2] for m in (7, 8, 9)) * 1000.0
    q4 = sum(FLATS[m][2] + PARKING[m][2] for m in (10, 11, 12)) * 1000.0
    assert plan["2026 Q3"] == pytest.approx(q3) and plan["2026 Q4"] == pytest.approx(q3 + q4)


def test_a_monthly_bank_plan_has_no_opening_remainder():
    """Первый квартал помесячного плана — обычный квартал, а не остаток на дату."""
    bank = contracting.read_bank_plan(_book())
    fact = {"2025 Q4": 300 * M, "2026 Q1": 320 * M, "2026 Q2": 310 * M, "2026 Q3": 330 * M}
    summary = {
        "by_quarter": [{"quarter": q, "amount": v, "area": v / 250_000, "price_per_sqm": 250_000}
                       for q, v in fact.items()],
        "by_quarter_flats": {}, "dynamics": [],
        "fm_plan": contracting.read_fm_plan(_book()), "bank_plan": bank,
    }
    plans = contracting.plan_comparison(summary)
    assert plans["bank_opening"] is None
    first = next(r for r in plans["quarters"] if r["label"] == "2026 Q3")
    assert first["bank_amount"] == pytest.approx(bank["gross_by_quarter"]["2026 Q3"])


def test_the_old_sheets_stay_as_fallback_and_decoys_are_not_plans():
    with pytest.raises(KeyError) as missing:
        contracting.read_fm_plan(_book(fm=False, bank=False))
    assert "«Продажи П-Ф_ФМ»" in str(missing.value) and "«Продажи ФМ_new Банников»" in str(missing.value)
    with pytest.raises(KeyError) as missing:
        contracting.read_bank_plan(_book(fm=False, bank=False))
    assert "«Продажи П-Ф_БАНК»" in str(missing.value) and "«План банка_new»" in str(missing.value)
    # Прежний лист финмодели читается тем же читателем.
    book = openpyxl.Workbook()
    book.active.title = "Раб.файл"
    _pf_sheet(book.create_sheet(contracting.FM_SHEET), 2, 3, 4, False)
    out = io.BytesIO()
    book.save(out)
    assert contracting.read_fm_plan(out.getvalue())["sheet"] == contracting.FM_SHEET
