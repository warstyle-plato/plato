"""Число в книге показано форматом своего поля, а не General.

Стиль Excel несёт формат числа вместе с цветом, и три места теряли его молча:
вводные «Вводных» (цель ставки 9% читалась как 0,09, дата старта — как
46227), читалки `='Вводные'!X` на «Параметрах модели» (формульный стиль
шаблона — General) и дописанные ячейки ОТЧЁТа (2 349,1 млн — как 2349.07318).

Решение о формате одно: `_v4_input_format_code` по типу и единице поля из
FIELD_GROUPS, для вводных без ключа — явная карта сборщика, а у читалки —
формат её источника. Проверяется собранная книга: формат живёт в стиле
ячейки, и строка в коде о нём ничего не доказывает.

Запуск: python3 -m pytest tests/test_the_book_shows_numbers_in_their_format.py -q
"""

from __future__ import annotations

import io
import re
import sys
import zipfile
from pathlib import Path

import openpyxl
import pytest

ROOT = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(ROOT))

import main_legacy as core  # noqa: E402

import v4_entry_sheet as ves  # noqa: E402

TEMPLATE = ROOT / "templates" / "DevelopAid_model_v4.xlsx"
pytestmark = pytest.mark.skipif(not TEMPLATE.is_file(), reason="шаблон v4 не поставляется")

_MIRROR = re.compile(r"^='Вводные'!\$?([A-Z]{1,3})\$?(\d+)$")
_INPUT_FILL = "FFFFF2CC"


@pytest.fixture(scope="module")
def book():
    inputs = dict(core.DEFAULT_INPUTS)
    # Ступени ставки ПФ дописывают свой блок только когда заданы.
    content, _, meta = core.build_project_workbook(
        inputs, core.TEP_DEFAULT, [], {}, project_name="Формат")
    return openpyxl.load_workbook(io.BytesIO(content)), meta


def _keyed_inputs(entry):
    meta = core._v4_input_field_meta()
    for key_col, value_col in core._V4_INPUT_BLOCKS:
        for row in range(1, entry.max_row + 1):
            key = entry[f"{key_col}{row}"].value
            if key in meta or key in core._V4_KEYED_INPUT_FORMATS:
                yield key, entry[f"{value_col}{row}"], meta.get(key)


def test_every_keyed_input_has_the_format_of_its_field(book):
    workbook, _meta = book
    entry = workbook[ves.ENTRY_SHEET]
    checked, wrong = 0, []
    for key, cell, field in _keyed_inputs(entry):
        value = cell.value
        if value is None or isinstance(value, str) or cell.fill.fgColor.rgb != _INPUT_FILL:
            continue
        if field:
            wanted = core._v4_input_format_code(
                *field, 1 if hasattr(value, "year") else value)
        else:
            wanted = core._V4_KEYED_INPUT_FORMATS[key]
        if wanted is None:
            continue
        checked += 1
        if ves.format_family(cell.number_format) != ves.format_family(wanted):
            wrong.append((cell.coordinate, key, cell.number_format, wanted))
    assert checked > 60, f"проверено {checked} вводных — лист не разобран"
    assert not wrong, f"вводная показана не форматом своего поля: {wrong[:8]}"


def test_no_numeric_input_is_left_in_general(book):
    """Жёлтая числовая ячейка без формата — вводная, которую сборщик забыл.

    Ноль в поле даты — «не задана»: формат даты показал бы 00.01.1900, а
    openpyxl прочитал бы его временем суток.
    """
    workbook, _meta = book
    entry = workbook[ves.ENTRY_SHEET]
    bare = [cell.coordinate for row in entry.iter_rows() for cell in row
            if isinstance(cell.value, (int, float)) and not isinstance(cell.value, bool)
            and cell.value != 0 and cell.fill.fgColor.rgb == _INPUT_FILL
            and cell.number_format == "General"]
    assert not bare, f"вводные в General: {bare[:10]}"


def test_every_mirror_shows_its_source_format(book):
    workbook, _meta = book
    entry, params = workbook[ves.ENTRY_SHEET], workbook[ves.PARAMS_SHEET]
    mirrors, wrong = 0, []
    for row in params.iter_rows():
        for cell in row:
            link = _MIRROR.match(str(cell.value or ""))
            if not link:
                continue
            mirrors += 1
            source = entry[f"{link.group(1)}{link.group(2)}"]
            if cell.number_format != source.number_format:
                wrong.append((cell.coordinate, source.coordinate,
                              source.number_format, cell.number_format))
    assert mirrors > 150, f"читалок {mirrors} — лист не разделён"
    assert not wrong, f"читалка теряет формат вводной: {wrong[:8]}"


def test_the_report_numbers_we_write_are_formatted(book):
    workbook, _meta = book
    report = workbook["ОТЧЕТ"]
    total = core._V4_PRODUCT_STRUCTURE_TOTAL_ROW
    coords = ([f"B{core._V4_REPORT_DEFAULT_ROW}", "F33", "H33", "H36", "H37",
               f"B{total}", f"E{total}", f"{core._V4_UNDER_COLUMN}{total}"]
              + [f"{core._V4_UNDER_COLUMN}{row}"
                 for row in range(core._V4_PRODUCT_STRUCTURE_FIRST_ROW, total)]
              + [f"{column}73" for column in "BCDEF"])
    bare = [coord for coord in coords if report[coord].number_format == "General"]
    assert not bare, f"числа ОТЧЁТа в General: {bare}"
    consolidator = workbook["КОНСОЛИДАТОР"]
    followed = 0
    for row in report.iter_rows(min_row=24, max_row=29):
        for cell in row:
            link = re.match(r"^='КОНСОЛИДАТОР'!\$?([A-Z]+)\$?(\d+)$", str(cell.value or ""))
            if link:
                followed += 1
                source = consolidator[f"{link.group(1)}{link.group(2)}"]
                assert cell.number_format == source.number_format, (
                    cell.coordinate, source.number_format, cell.number_format)
    assert followed > 20, followed


# --- сам выбор формата: проверка обязана падать на подделке ----------------

def _sheet(cells: str) -> str:
    return f'<x:worksheet><x:sheetData><x:row r="5">{cells}</x:row></x:sheetData></x:worksheet>'


def _template_styles() -> str:
    with zipfile.ZipFile(TEMPLATE) as source:
        return source.read("xl/styles.xml").decode("utf-8")


def test_a_general_percent_input_becomes_a_percent_and_keeps_its_colour():
    styles = _template_styles()
    entry = ves.style_map(styles)["entry_default"]
    xml = _sheet(f'<x:c r="B5" s="{entry}"><x:v>0.09</x:v></x:c>'
                 '<x:c r="D5" t="inlineStr"><x:is><x:t>rate_target_base_pct</x:t></x:is></x:c>')
    typed, new_styles = core._v4_type_input_formats(xml, styles)
    style = int(re.search(r'<x:c r="B5" s="(\d+)"', typed).group(1))
    assert "%" in ves.format_code(new_styles, ves.num_fmt_id(new_styles, style))
    assert style in ves.style_map(new_styles)["entry"], "клон потерял цвет ввода"


def test_a_money_format_on_a_percent_field_is_replaced():
    """Спред БРИДЖа 0,06 форматом #,##0.0 показывался как 0,1."""
    styles = _template_styles()
    money = next(index for index in sorted(ves.style_map(styles)["entry"])
                 if ves.format_family(ves.format_code(styles, ves.num_fmt_id(styles, index)))
                 == "number")
    xml = _sheet(f'<x:c r="B5" s="{money}"><x:v>0.06</x:v></x:c>'
                 '<x:c r="D5" t="inlineStr"><x:is><x:t>bridge_spread_pp</x:t></x:is></x:c>')
    typed, new_styles = core._v4_type_input_formats(xml, styles)
    style = int(re.search(r'<x:c r="B5" s="(\d+)"', typed).group(1))
    assert ves.format_family(ves.format_code(new_styles, ves.num_fmt_id(new_styles, style))) == "pct"


def test_the_fraction_is_not_rounded_away():
    """2,75 тыс. ₽/м² форматом #,##0.0 показалось бы 2,8."""
    assert core._v4_input_format_code("number", "тыс. ₽/м²", 2.75) == core._V4_FMT_AMOUNT_FINE
    assert core._v4_input_format_code("number", "тыс. ₽/м²", 350) == core._V4_FMT_AMOUNT
    assert core._v4_input_format_code("number", "%", 0.0425) == core._V4_FMT_PCT_FINE
    assert core._v4_input_format_code("number", "мес.", 24) == core._V4_FMT_COUNT
    assert core._v4_input_format_code("text", "режим", "Да") is None
