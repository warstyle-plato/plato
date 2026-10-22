"""Бюджет месяца в сводке — по статьям БДДС, а не тремя агрегатами.

Задача контроля (20.10.2026, по образцу PDF): каждая статья план-факта БДДС —
план месяца, оплачено по РСС, остаток, процент; подытоги по главам и итог;
перерасход и самый большой остаток выделены; оплаты кодов РСС вне карты —
строкой «без плана». Статьи ищутся по точному имени, группировка строк
(outline) не нужна. Карта статей и кодов принята владельцем по сверке сентября.

Месяц без данных в тренде — «нет данных», а не «0 ₽»: реестр РСС, который
начинается в сентябре, ничего не говорит об августе.

Запуск: python3 -m pytest tests/test_management_report_budget_articles.py -q
"""
from __future__ import annotations

import datetime
import io
import sys
from pathlib import Path

import pytest
from openpyxl import Workbook, load_workbook

ROOT = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(ROOT))
sys.path.insert(0, str(ROOT / "tests"))

import developaid_management_report as report  # noqa: E402
import developaid_monitor as monitor  # noqa: E402
import developaid_monthly_budget as monthly_budget  # noqa: E402
from test_management_report_model import M, PROJECT, _gpr_book, _rss_book  # noqa: E402


@pytest.fixture(autouse=True)
def _isolated(tmp_path, monkeypatch):
    monkeypatch.setattr(monitor, "_SNAPSHOT_DIR", tmp_path / "monitor")
    monkeypatch.setenv("DATA_DIR", str(tmp_path / "data"))


# Строки бюджета: (имя, план августа, план октября, факт августа).
# «Прочие работы и затраты» сначала стоит в чужой главе — ловушка: статья
# стройки ищется только после строки своей главы.
BUDGET_ROWS = [
    ("Стоимость приобретения проекта", None, -50, None),
    ("Прочие работы и затраты", None, -50, None),
    ("Стоимость строительства и проектирования", -110, -233, -100),
    ("Проектные и изыскательские работы, авторский надзор", -5, -10, -5),
    ("Основные объекты строительства", -100, -200, -95),
    ("Прочие работы и затраты", -5, -5, None),
    ("Наружные сети и сооружения", None, -5, None),
    ("Временные здания и сооружения", None, None, None),
    ("Содержание службы заказчика", None, -13, None),
    ("Коммерческие расходы", -6, -12, -4),
    ("Расходы на персонал отдела продаж (КС) - Группа", -1, -2, -1),
    ("Расходы отдела продаж на содержание офиса (КС) - Группа", -1, -1, -1),
    ("Реклама и PR (КС) - Группа", -4, -6, -2),
    ("Расходы на продажу (КС) - Группа", None, -3, None),
    ("Административно-хозяйственные расходы", -9, -9, -9),
]


def articles_budget() -> bytes:
    book = Workbook()
    sheet = book.active
    sheet.cell(row=2, column=1, value="План-фактный анализ_Гродненская_01.10.2026")
    for col, name in ((5, "Август 2026"), (8, "Октябрь 2026")):
        sheet.cell(row=4, column=col, value=name)
        sheet.cell(row=5, column=col, value="ПЛАН (СЗ Гродненская 18)")
        sheet.cell(row=5, column=col + 1, value="ФАКТ")
    sheet.cell(row=4, column=11, value="Итого")
    for r, (name, august, october, fact) in enumerate(BUDGET_ROWS, 6):
        sheet.cell(row=r, column=1, value=name)
        for col, value in ((5, august), (6, fact), (8, october)):
            if value is not None:
                sheet.cell(row=r, column=col, value=value * M)
    blob = io.BytesIO()
    book.save(blob)
    return blob.getvalue()


# Оплаты октября: (день, сумма в млн, код РСС, имя статьи в листе расчёта).
PAYMENTS = [
    (datetime.date(2026, 9, 10), 50, "2.2.1.4", None),
    (datetime.date(2026, 10, 3), 60, "2.2.1.4", None),
    (datetime.date(2026, 10, 6), 8, "2.2.1.10", "Содержание стройплощадки"),
    (datetime.date(2026, 10, 7), 2, "2.7", "Проектные работы"),
    (datetime.date(2026, 10, 8), 4, "2.8", "Технологическое присоединение"),
    (datetime.date(2026, 10, 9), 5, "3.1", "Реклама"),
    (datetime.date(2026, 10, 9), 1, "3.7", "Прочие"),
    (datetime.date(2026, 10, 10), 20, "2.6", "Содержание службы заказчика"),
    (datetime.date(2026, 10, 10), 7, "3.5", "Услуги банка"),
]


def articles_rss() -> bytes:
    data = _rss_book([(day, amount * M, code) for day, amount, code, _ in PAYMENTS])
    book = load_workbook(io.BytesIO(data))
    sheet = book["Расчет стоимости строительства"]
    names = {code: name for _, _, code, name in PAYMENTS if name}
    for r, (code, name) in enumerate(sorted(names.items()), 20):
        sheet.cell(row=r, column=1, value=code)
        sheet.cell(row=r, column=4, value=name)
    blob = io.BytesIO()
    book.save(blob)
    return blob.getvalue()


def _setup() -> dict:
    monitor.store_schedule(PROJECT, _gpr_book({}), None, "2026-07-01")
    monitor.store_estimate(PROJECT, articles_rss(), "2026-10-18")
    monthly_budget.store_budget(PROJECT, articles_budget(), "2026-10-01")
    return report.build(PROJECT, "week", "2026-10-18")


def test_codes_go_to_the_longest_matching_article():
    assert monthly_budget.article_of("2.2.1.4") == "objects"
    # Содержание стройплощадки лежит внутри 2.2, но в бюджете — «Прочие работы».
    assert monthly_budget.article_of("2.2.1.10") == "upkeep"
    assert monthly_budget.article_of("2.2.2.8.1") == "upkeep"
    assert monthly_budget.article_of("2.2.10") == "objects"
    assert monthly_budget.article_of("3.4") == "selling"
    # Вне карты — не в соседнюю статью, а «без плана»; вне сводки — None.
    assert monthly_budget.article_of("2.8") is None
    assert monthly_budget.article_of("3.5") is None


def test_the_month_by_budget_articles():
    built = _setup()
    table = built["budget"]["by_budget_article"]
    chapters = {chapter["key"]: chapter for chapter in table["chapters"]}
    build = chapters["construction"]
    rows = {row["key"]: row for row in build["rows"]}
    # План главы — строка главы бюджета (233 без службы заказчика 13 = 220).
    assert build["total"]["plan"] == pytest.approx(220 * M)
    assert build["total"]["fact"] == pytest.approx(74 * M)
    assert rows["objects"]["plan"] == pytest.approx(200 * M)
    assert rows["objects"]["fact"] == pytest.approx(60 * M)
    assert rows["objects"]["remaining"] == pytest.approx(140 * M)
    # Ловушка: «Прочие работы» из чужой главы (50) не взяты.
    assert rows["upkeep"]["plan"] == pytest.approx(5 * M)
    assert rows["upkeep"]["fact"] == pytest.approx(8 * M)
    assert rows["upkeep"]["over"] is True
    # Статья без плана в бюджете — пустая ячейка, а не ноль.
    assert rows["temporary"]["plan"] is None
    unplanned = [row for row in build["rows"] if row["key"] == "unplanned"]
    assert [row["label"] for row in unplanned] == [
        "Без плана: 2.8 Технологическое присоединение"]
    assert unplanned[0]["fact"] == pytest.approx(4 * M)
    selling = {row["key"]: row for row in chapters["commercial"]["rows"]}
    assert selling["ads"]["plan"] == pytest.approx(6 * M)
    assert selling["sales_staff"]["plan"] == pytest.approx(3 * M)
    assert selling["sales_staff"]["fact"] == pytest.approx(1 * M)
    # Итог — стройка и коммерция; служба заказчика и услуги банка вне итога.
    assert table["total"]["plan"] == pytest.approx(232 * M)
    assert table["total"]["fact"] == pytest.approx(80 * M)
    assert table["overrun"] == ["upkeep"]
    assert table["largest_left"] == "objects"


def test_the_finding_names_the_overrun_and_the_largest_remainder():
    built = _setup()
    found = {item["topic"]: item for item in built["findings"]}
    text = found["Бюджет месяца"]["text"]
    assert text.startswith("По стройке и коммерческим расходам оплачено 80,0 млн ₽ из "
                           "232,0 млн ₽ (34,5%)")
    assert "перерасход: «Прочие работы и затраты (содержание стройплощадки)» на 3,0 млн ₽" in text
    assert ("больше всего не оплачено по статье «Основные объекты строительства» — "
            "140,0 млн ₽ из 200,0 млн ₽") in text
    assert "без плана в бюджете оплачено 4,0 млн ₽" in text
    assert found["Бюджет месяца"]["tone"] == "bad"


def test_the_document_shows_every_article_and_colours_the_overrun():
    built = _setup()
    page = report.render_html(built)
    part = page[page.index("<h2>Бюджет месяца</h2>"):]
    table = part[part.index("<table class='articles'>"):part.index("</table>")]
    for label in ("Проектные и изыскательские работы, авторский надзор",
                  "Основные объекты строительства", "Реклама и PR",
                  "Персонал и офис отдела продаж",
                  "Без плана: 2.8 Технологическое присоединение",
                  "Итого стройка и коммерческие расходы"):
        assert label in table, label
    assert "<tr class='over'><td>&nbsp;&nbsp;Прочие работы и затраты" in table
    assert "<tr class='left'><td>&nbsp;&nbsp;Основные объекты строительства" in table
    assert "<span class='neg'>−3,0 млн ₽</span>" in table
    # Глава без статей в карте не пропадает: «Наружные сети» с планом и без оплат.
    assert "Наружные сети и сооружения" in table


def test_months_without_data_are_not_zero():
    built = _setup()
    trend = built["budget"]["trend"]
    # Июль: ни плана, ни оплат — не показан. Август: план есть, реестр РСС
    # начинается в сентябре — факт None, а не ноль.
    assert [point["month"] for point in trend] == ["2026-08-01", "2026-09-01", "2026-10-01"]
    assert trend[0]["fact"]["construction"] is None
    assert trend[0]["plan"]["construction"] == pytest.approx(110 * M)
    assert trend[1]["fact"]["construction"] == pytest.approx(50 * M)
    page = report.render_html(built)
    row = page[page.index("<td>август 2026</td>"):]
    row = row[:row.index("</tr>")]
    assert "нет данных" in row and "0 ₽" not in row


def test_the_tiles_lead_the_first_page():
    built = _setup()
    items = report.tiles(built)
    assert len(items) == 12
    by_label = {item["label"]: item for item in items}
    build = by_label["Бюджет стройки за месяц"]
    assert build["value"] == "33,6%"
    assert build["tone"] == "bad" and "перерасход по 1 статье" in build["note"]
    # Нет отчёта о продажах — плитка так и говорит, а не показывает ноль.
    assert by_label["Продажи квартала к плану финмодели"]["value"] == "нет данных"
    page = report.render_html(built)
    assert page.index("class='tiles'") < page.index("Главное за период")
    empty = report.tiles(report.build("Пусто", "week", "2026-10-18"))
    assert {item["value"] for item in empty} == {"нет данных"}


def test_the_whole_month_budget_matches_the_file():
    # «Сумма 505, 1 глава 78, 2 — 398, 3 — 27, АУР 1,6» (владелец, 22.10.2026):
    # сводка показывала только стройку и коммерцию — 419 млн. Всего по бюджету
    # месяца — все его главы: стройка 220, коммерция 12, служба заказчика 13,
    # земля 50, АХР 9.
    built = _setup()
    table = built["budget"]["by_budget_article"]
    assert table["total"]["plan"] == pytest.approx(232 * M)
    assert table["whole"]["plan"] == pytest.approx(304 * M)
    chapters = {chapter["key"]: chapter["total"] for chapter in table["chapters"]}
    assert chapters["land"]["plan"] == pytest.approx(50 * M)
    assert chapters["admin"]["plan"] == pytest.approx(9 * M)
    # Оплат АХР в РСС нет — это не ноль.
    assert chapters["admin"]["fact"] is None
    page = report.render_html(built)
    part = page[page.index("<h2>Бюджет месяца</h2>"):]
    table_html = part[part.index("<table class='articles'>"):part.index("</table>")]
    assert table_html.index("Итого стройка и коммерческие расходы") \
        < table_html.index("Приобретение проекта (земля)") \
        < table_html.index("Всего по бюджету месяца")
    assert "нет в РСС" in table_html
    assert "304,0 млн ₽" in table_html


def test_the_whole_comes_from_the_expenses_row_when_the_file_has_it():
    rows = [("РАСХОДЫ", -130, -505, -120)] + BUDGET_ROWS
    book = Workbook()
    sheet = book.active
    for col, name in ((5, "Август 2026"), (8, "Октябрь 2026")):
        sheet.cell(row=4, column=col, value=name)
        sheet.cell(row=5, column=col, value="ПЛАН")
        sheet.cell(row=5, column=col + 1, value="ФАКТ")
    for r, (name, august, october, fact) in enumerate(rows, 6):
        sheet.cell(row=r, column=1, value=name)
        for col, value in ((5, august), (6, fact), (8, october)):
            if value is not None:
                sheet.cell(row=r, column=col, value=value * M)
    blob = io.BytesIO()
    book.save(blob)
    monitor.store_schedule(PROJECT, _gpr_book({}), None, "2026-07-01")
    monitor.store_estimate(PROJECT, articles_rss(), "2026-10-18")
    monthly_budget.store_budget(PROJECT, blob.getvalue(), "2026-10-01")
    whole = report.build(PROJECT, "week", "2026-10-18")["budget"]["by_budget_article"]["whole"]
    assert whole["plan"] == pytest.approx(505 * M)
