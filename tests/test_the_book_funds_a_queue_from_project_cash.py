"""Единая касса проекта и база вычета НДС в книге v4 — как в движке.

Книга «Нагатино КРТ с выкупом» (прод 0.25.91, 16.10.2026) показывала пять FAIL
паритета с движком при неизменном листе ввода. Разбор 17.10.2026 нашёл два
источника, оба — методика книги, а не ошибка движка:

- стоимость финансирования +401,6 млн и пик ПФ +930,5 млн — целиком в
  очереди 4: при стратегии `unified_project_cash` движок покрывает её расходы
  до РнС свободной кассой предыдущих очередей (930,5 млн) и берёт БРИДЖем
  только остаток; книга стратегии не знала и вела БРИДЖ каждой очереди от нуля
  (backlog 30.08.2026);
- НДС к уплате +324,7 млн — только в очереди 1: формула исключала из базы
  вычета НАЧИСЛЕННУЮ плату за ВРИ (B16 × доля), а при льготе 100 % уплачено 0.

Закреплено на этом проекте как эталоне (`tests/fixtures/project_settings_nagatino_krt_0_25_91.json`):
- строки 88–89 листов CF: покрытие кассой и доступная касса; сумма покрытия
  четвёртой очереди, её пик ПФ и пик БРИДЖа равны движку;
- база НДС очереди 1 — по уплаченной плате за ВРИ: НДС равен движку; движок
  сам исключает из базы вычета уплаченное, а не начисленное — и у одиночного
  проекта тоже (решение владельца 18.10.2026: Донской 113,8 → 99,8 млн);
- признак единой кассы — ввод на «Вводных» рядом с переносом долга; при
  независимом финансировании он «Нет»;
- дашборд не форматирует LLCR через TEXT(…,"0.00") — в русском Excel это #VALUE!.

Запуск: python3 -m pytest tests/test_the_book_funds_a_queue_from_project_cash.py -q
"""

from __future__ import annotations

import copy
import io
import json
import sys
from pathlib import Path

import pytest

ROOT = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(ROOT))
sys.path.insert(0, str(ROOT / "tests"))

import v4_inputs

import main_legacy as core

openpyxl = pytest.importorskip("openpyxl")
from openpyxl.utils import get_column_letter

from xlsx_eval import Evaluator

FIXTURE = ROOT / "tests" / "fixtures" / "project_settings_nagatino_krt_0_25_91.json"
PROJECT_CASH_ROW, POOL_ROW = 88, 89
STANDALONE_TAX_ROW, FREE_CASH_ROW = 90, 91
MONTHS = [get_column_letter(c) for c in range(4, 184)]


def _spec(strategy: str | None = None) -> tuple[dict, dict, dict]:
    saved = json.loads(FIXTURE.read_text(encoding="utf-8"))
    inputs = {**copy.deepcopy(core.DEFAULT_INPUTS), **saved["inputs"]}
    phasing = copy.deepcopy(saved["phasing"])
    if strategy:
        phasing["financing_strategy"] = strategy
    return inputs, saved["tep"], phasing


def _build(strategy: str | None = None):
    inputs, tep, phasing = _spec(strategy)
    bundle = core._run_authoritative_model(inputs, tep, [], copy.deepcopy(phasing))
    content, _, meta = core.build_project_workbook(
        inputs, tep, [], copy.deepcopy(phasing), project_name="Нагатино КРТ",
        finance_hints=core._v4_finance_hints(bundle))
    assert meta["missing"] == [], meta["missing"]
    return bundle, openpyxl.load_workbook(io.BytesIO(content), data_only=False)


@pytest.fixture(scope="module")
def unified():
    # Вычислитель идёт по цепочке месяцев рекурсией — как в соседних сверках CF.
    sys.setrecursionlimit(400000)
    bundle, book = _build()
    assert bundle["phases"][3]["result"]["finance"]["project_cash_used"] > 5e8, \
        "эталон перестал брать кассу в четвёртой очереди — проверять нечего"
    return bundle, book, Evaluator(book)


def _row(evaluator: Evaluator, sheet: str, row: int) -> list[float]:
    return [float(evaluator.cell(sheet, f"{column}{row}") or 0.0) for column in MONTHS]


def test_the_fourth_queue_is_funded_from_project_cash_like_the_engine(unified) -> None:
    bundle, _, evaluator = unified
    finance = bundle["phases"][3]["result"]["finance"]
    covered = float(evaluator.cell("CF_4", f"B{PROJECT_CASH_ROW}"))
    assert covered == pytest.approx(finance["project_cash_used"] / 1e6, abs=0.05)
    assert max(_row(evaluator, "CF_4", 47)) == pytest.approx(finance["peak_pf"] / 1e6, abs=0.1)
    assert max(_row(evaluator, "CF_4", 36)) == pytest.approx(finance["peak_bridge"] / 1e6, abs=0.05)
    # Первая очередь кассы брать не у кого: её строки — нули, остальное как было.
    assert float(evaluator.cell("CF_1", f"B{PROJECT_CASH_ROW}")) == 0.0
    assert max(_row(evaluator, "CF_1", 36)) == pytest.approx(
        bundle["phases"][0]["result"]["finance"]["peak_bridge"] / 1e6, abs=0.05)


def test_the_project_cash_pool_is_the_earlier_queues_distributions(unified) -> None:
    """Доступная касса — положительный поток капитала предыдущих очередей к
    месяцу, за вычетом уже взятого. Контрпример: касса «из ниоткуда» или без
    вычета взятого дала бы другой ряд.

    Поток — на налоге очереди в одиночку (строка 91), как копит кассу движок:
    строка 50 платит долю налога свода, а он зависит от процентов следующих
    очередей — касса на ней замкнула бы книгу в круг."""
    bundle, _, evaluator = unified
    for queue in (1, 2, 3):
        alone = bundle["phases"][queue - 1]["result"]["finance"]["standalone_profit_tax"]
        assert float(evaluator.cell(f"CF_{queue}", f"B{STANDALONE_TAX_ROW}")) == pytest.approx(
            alone / 1e6, abs=0.05), queue
    pool = _row(evaluator, "CF_4", POOL_ROW)
    draws = _row(evaluator, "CF_4", PROJECT_CASH_ROW)
    expected, taken = [], 0.0
    for index, column in enumerate(MONTHS):
        earlier = sum(
            sum(float(evaluator.cell(f"CF_{queue}", f"{letter}{FREE_CASH_ROW}") or 0.0)
                for letter in MONTHS[:index + 1])
            for queue in (1, 2, 3))
        expected.append(max(0.0, earlier - taken))
        taken += draws[index]
    assert pool == pytest.approx(expected, abs=0.01)


def test_the_vat_base_excludes_the_vri_actually_paid(unified) -> None:
    bundle, _, evaluator = unified
    for index, phase in enumerate(bundle["phases"], start=1):
        book = float(evaluator.cell(f"CF_{index}", "B21"))
        assert book == pytest.approx(phase["result"]["finance"]["vat"] / 1e6, abs=0.05), index


def test_the_strategy_is_an_input_next_to_the_debt_carry_flag(unified) -> None:
    _, book, _ = unified
    sheet = v4_inputs.inputs(book)
    assert sheet["K92"].value == "Да"
    assert sheet["M92"].value == "financing_strategy"
    _, independent = _build("independent")
    assert v4_inputs.inputs(independent)["K92"].value == "Нет"


def test_the_dashboard_does_not_format_llcr_with_a_dotted_text_code(unified) -> None:
    _, book, _ = unified
    formulas = [str(cell.value) for row in book["Дашборд"].iter_rows() for cell in row
                if isinstance(cell.value, str) and cell.value.startswith("=")]
    assert not [f for f in formulas if 'TEXT(' in f and '"0.00"' in f], "TEXT с кодом «0.00»"
    assert any("FIXED(" in f and '&"x"' in f for f in formulas), "LLCR без формулы FIXED"


def test_the_engine_excludes_only_the_vri_actually_paid() -> None:
    """Одиночный расчёт вычитал из базы вычета НАЧИСЛЕННУЮ плату за ВРИ: при
    льготе 100 % база урезалась на сумму, которой в CAPEX нет, и НДС выходил
    на 14 млн выше, чем у тех же вводных без начисления (Донской)."""
    saved = json.loads((ROOT / "tests" / "fixtures" / "project_settings_donskoy_0_25_78.json")
                       .read_text(encoding="utf-8"))
    inputs = {**copy.deepcopy(core.DEFAULT_INPUTS), **saved["inputs"]}
    assert inputs["vri_relief_mode"] == "percent" and float(inputs["vri_relief_pct"]) == 100
    assert float(inputs["land_rights_cost_mln"]) > 100
    charged = core.calculate(core.CalcRequest(inputs=inputs, tep=saved["tep"], rates=[]))
    free = core.calculate(core.CalcRequest(
        inputs={**inputs, "land_rights_cost_mln": 0.0}, tep=saved["tep"], rates=[]))
    assert charged["finance"]["total_capex"] == pytest.approx(free["finance"]["total_capex"])
    assert charged["finance"]["vat"] == pytest.approx(free["finance"]["vat"], rel=1e-9)
