"""Очередь объекта у книги — та, в которой его построил движок.

Очередь отдельно стоящего объекта решали двое и по-разному. У движка лестница
«объявлено продуктами очереди → `discrete` → умолчание»; у сборщика книги —
только `discrete` с зашитым умолчанием («ТЦ и наземный паркинг во второй»).
Проект, который размещает объект ПРОДУКТАМИ очереди — так это делает и вкладка
«Очередность», и пресеты, — до `discrete` не доходит вовсе, и книга ставила
объект по своему умолчанию.

На пресете КРТ Нагатино (ТЦ объявлен в четвёртой очереди) книга строила его во
ВТОРОЙ, то есть на два года раньше: он продавался и строился дешевле, и на
одних вводных выручка расходилась на 7 910 млн ₽, CAPEX на 6 692, налог на 584.
Пик БРИДЖа при этом сходился до копейки, а LLCR проходил допуск — то есть блок
паритета выглядел «почти сошедшимся» ровно там, где очередь строит не тот
объект, и поймать это им было нельзя.

Запуск: python3 -m pytest tests/test_the_book_builds_the_object_in_the_engines_queue.py -q
"""

from __future__ import annotations

import io
import sys
from pathlib import Path

import openpyxl
import pytest

ROOT = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(ROOT))
sys.path.insert(0, str(ROOT / "tests"))

import main as wrapper  # noqa: E402
from xlsx_eval import Evaluator  # noqa: E402

core = wrapper.core

# Ячейка «очередь объекта» на листе ОБЪЕКТЫ — та же, что читает КОНСОЛИДАТОР.
_RETAIL_QUEUE_CELL = "B36"
_OFFICES_QUEUE_CELL = "B8"


def _project(place_by_products: bool):
    """Четыре очереди; ТЦ в четвёртой, офисы в третьей.

    `place_by_products` — размещение объявлено продуктами очереди (так делают
    страница и пресеты) либо через `discrete` (выпадающий список).
    """
    inputs = dict(core.DEFAULT_INPUTS)
    inputs["offices_enabled"] = True
    inputs["retail_enabled"] = True
    tep = {key: dict(value) for key, value in core.TEP_DEFAULT.items()}
    for key, gns in (("offices", 20000.0), ("standalone_retail", 18000.0)):
        tep.setdefault(key, {})
        tep[key].update({"gns": gns, "total_area": gns * 0.94,
                         "saleable": gns * 0.6, "units": 0})
    phasing = {"enabled": True, "phase_count": 4, "phase_gap_months": 12,
               "phases": [{"name": f"Очередь {i + 1}", "start_offset_months": 12 * i,
                           "construction_months": 24} for i in range(4)]}
    if place_by_products:
        phasing["phases"][2]["products"] = {"offices": dict(tep["offices"])}
        phasing["phases"][3]["products"] = {
            "standalone_retail": dict(tep["standalone_retail"])}
    else:
        phasing["discrete"] = {"offices": 3, "standalone_retail": 4}
    return inputs, tep, phasing


def _built(inputs, tep, phasing):
    sys.setrecursionlimit(400000)
    content, _, meta = core.build_project_workbook(inputs, tep, [], phasing)
    book = openpyxl.load_workbook(io.BytesIO(content), data_only=False)
    return Evaluator(book), meta


def test_an_object_placed_by_products_lands_in_that_queue():
    """Объявленный продуктами четвёртой очереди ТЦ книга строит в четвёртой.

    На прежнем коде здесь стояла ВТОРАЯ — зашитое умолчание сборщика.
    """
    evaluator, _ = _built(*_project(place_by_products=True))
    assert int(float(evaluator.cell("ОБЪЕКТЫ", _RETAIL_QUEUE_CELL))) == 4
    assert int(float(evaluator.cell("ОБЪЕКТЫ", _OFFICES_QUEUE_CELL))) == 3


def test_the_dropdown_form_is_not_broken_by_the_fix():
    """Та же расстановка через `discrete` даёт то же самое.

    Формы объявления две, и правка обязана работать на обеих: починив одну и
    сломав вторую, мы поменяли бы один молчаливый разъезд на другой.
    """
    evaluator, _ = _built(*_project(place_by_products=False))
    assert int(float(evaluator.cell("ОБЪЕКТЫ", _RETAIL_QUEUE_CELL))) == 4
    assert int(float(evaluator.cell("ОБЪЕКТЫ", _OFFICES_QUEUE_CELL))) == 3


def test_the_two_forms_agree_cell_for_cell():
    """Одно размещение, объявленное двумя способами, — одна книга.

    Утверждение здесь именно это, а не «обе не упали»: пока числа сверяются
    только с самими собой, расхождение форм читается как свойство проекта.
    """
    by_products, _ = _built(*_project(place_by_products=True))
    by_discrete, _ = _built(*_project(place_by_products=False))
    for row in range(76, 85):
        left = float(by_products.cell("ПРОВЕРКИ", f"B{row}") or 0)
        right = float(by_discrete.cell("ПРОВЕРКИ", f"B{row}") or 0)
        assert left == pytest.approx(right, abs=0.01), f"строка {row}"


def test_a_silent_engine_names_the_queue_it_guessed():
    """Движок не ответил — книга берёт умолчание и НАЗЫВАЕТ его.

    Молчание тут хуже промаха: разъехавшееся размещение на экране выглядит
    ровно так же, как посчитанное.
    """
    inputs, tep, phasing = _project(place_by_products=True)
    sys.setrecursionlimit(400000)
    _, _, meta = core.build_project_workbook(
        inputs, tep, [], phasing, finance_hints={})
    said = "; ".join(str(item) for item in (meta.get("missing") or []))
    assert "очередь объекта" in said, said
