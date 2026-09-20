"""Одобренный лимит ПФ — потолок и в книге, а не только в движке.

Поле `pf_limit_approved_mln` работало ровно в одной поверхности. Движок им
режет выборку, считает непокрытую потребность и запоминает месяц первой
нехватки — это проверено отдельно (`test_pf_limit_is_a_ceiling_not_a_note`).
Книга о нём не знала вовсе: она считала лимит СВОЙ (B26 — округление выборки
вверх до 10 млн) и потолком его не ограничивала. Причина стояла прямо в коде,
в `V4_INPUTS_SHOWN_ONLY`, — то есть про пробел знали и оставили.

Цена: «требуется 40 880, банк одобрил 24 528» — движок обрежет выборку и
покажет дыру в 16 351 млн ₽, а книга посчитает так, будто все 40 880 доступны.
Это не допуск в полпроцента: два файла одного расчёта считают разные проекты.

Лимит применяется К КАЖДОЙ ОЧЕРЕДИ целиком — так делает движок, и книга
обязана делать так же. Что «одобренный лимит» значит у проекта с очередями —
общий потолок или потолок каждой — вопрос методики, а не реализации.

Запуск: python3 -m pytest tests/test_the_book_obeys_the_approved_pf_limit.py -q
"""

from __future__ import annotations

import copy
import io
import sys
from pathlib import Path

import openpyxl
import pytest

sys.path.insert(0, str(Path(__file__).resolve().parent.parent))
sys.path.insert(0, str(Path(__file__).resolve().parent))

from xlsx_eval import Evaluator  # noqa: E402

import main as _wrapper  # noqa: E402

core = _wrapper.core

DRAW_TOTAL = "B45"          # ПФ — выборка за весь срок
UNCOVERED = f"B{core._V4_PF_UNCOVERED_ROW}"


def _tep() -> dict:
    return {key: dict(row) for key, row in core.TEP_DEFAULT.items()}


def _inputs(approved: float) -> dict:
    return {**core.DEFAULT_INPUTS, "purchase_price_mln": 3000,
            "project_start": "2027-01-01", "ird_months": 12,
            "construction_months": 24, "apartment_price_th": 500,
            "pf_limit_approved_mln": approved}


def _case(approved: float):
    sys.setrecursionlimit(400000)
    inputs = _inputs(approved)
    bundle = core._run_authoritative_model(inputs, _tep(), [], {})
    content, _, meta = core.build_project_workbook(inputs, _tep(), [], {})
    assert meta["missing"] == [], meta["missing"]
    book = openpyxl.load_workbook(io.BytesIO(content), data_only=False)
    return bundle["consolidated"]["finance"], Evaluator(book), book


@pytest.fixture(scope="module")
def free():
    return _case(0)


@pytest.fixture(scope="module")
def starved(free):
    """Одобрено 60% требуемого."""
    finance = free[0]
    return _case(round(float(finance["pf_limit_required"]) / 1e6 * 0.6))


def test_without_an_approved_limit_the_book_is_untouched(free):
    """Ноль в поле — прежнее поведение до последней формулы.

    Иначе правка задевала бы каждый проект, а не тот, где лимит задан.
    """
    finance, evaluator, _ = free
    assert float(finance["pf_shortfall"]) == 0.0
    assert float(evaluator.cell("CF_1", UNCOVERED) or 0) == pytest.approx(0.0, abs=0.01)
    assert float(evaluator.cell("CF_1", DRAW_TOTAL) or 0) == pytest.approx(
        float(finance["pf_draw_total"]) / 1e6, abs=0.5)


def test_the_book_stops_drawing_at_the_approved_limit(starved, free):
    """Выборка книги упирается в потолок — и совпадает с движковой."""
    finance, evaluator, _ = starved
    free_finance, _, _ = free
    approved_mln = float(finance["pf_limit_approved"]) / 1e6
    assert approved_mln > 0, "предохранитель: лимит не задан — потолка нет"
    assert float(free_finance["pf_limit_required"]) / 1e6 > approved_mln, (
        "предохранитель: одобрено больше, чем нужно, — потолок не работает")

    book_draw = float(evaluator.cell("CF_1", DRAW_TOTAL) or 0)
    assert book_draw == pytest.approx(approved_mln, abs=0.5)
    assert book_draw == pytest.approx(float(finance["pf_draw_total"]) / 1e6, abs=0.5)


def test_the_uncovered_need_matches_the_engine(starved):
    """Дыра названа величиной — и той же, что у движка."""
    finance, evaluator, _ = starved
    engine = float(finance["pf_shortfall"]) / 1e6
    assert engine > 0, "предохранитель: дыры нет — сверять нечего"
    assert float(evaluator.cell("CF_1", UNCOVERED) or 0) == pytest.approx(engine, abs=0.5)


def test_the_parity_row_watches_the_gap(starved):
    """Строка паритета непокрытой потребности зелёная — и она существует."""
    _, evaluator, _ = starved
    row = core._V4_PF_UNCOVERED_PARITY_ROW
    assert "непокрыт" in str(evaluator.cell("ПРОВЕРКИ", f"A{row}") or "").lower()
    assert str(evaluator.cell("ПРОВЕРКИ", f"F{row}") or "") == "OK"


def test_the_approved_limit_is_a_working_input_now(starved):
    """Поле перестало быть справочным — и у него есть читатель.

    Пока оно стояло в `V4_INPUTS_SHOWN_ONLY`, сторож вводных требовал
    ОБРАТНОГО: ноль читателей. Теперь читатель обязан быть, иначе поле
    обещает то, чего не делает.
    """
    assert "pf_limit_approved_mln" not in core.V4_INPUTS_SHOWN_ONLY
    _, _, book = starved
    # Читатель — цепочка из двух формул, а не сам лист ввода: потолок очереди
    # стоит в её клетке блока очередей (AT), у одиночного проекта она читает
    # одобренный лимит F26, а строка выборки ПФ листа CF читает клетку.
    # «Вводные» при записи в архив переименовываются в «Параметры модели»,
    # и ссылки едут вместе с именем.
    cap_cell = f"{core._V4_PF_QUEUE_CAP_COL}{core._V4_CF_QUEUE_ENABLED_ROW}"
    cap_formula = str(book["Параметры модели"][cap_cell].value or "")
    cell = core._V4_PF_APPROVED_CELL
    reference = f"${cell[0]}${cell[1:]}"
    assert reference in cap_formula, (
        f"потолок очереди не читает одобренный лимит: {cap_formula}")
    formula = str(book["CF_1"]["D45"].value or "")
    assert f"${core._V4_PF_QUEUE_CAP_COL}${core._V4_CF_QUEUE_ENABLED_ROW}" in formula, (
        f"выборка ПФ не читает потолок очереди: {formula}")
    assert "Параметры модели" in formula, (
        f"ссылка не переименована вместе с листом: {formula}")


def test_the_funding_check_warns_instead_of_failing(starved):
    """«Потребность профинансирована полностью» больше не кричит зря.

    Проверка (строка 32) писалась тогда, когда выборка равнялась потребности
    ВСЕГДА: потолка у книги не было, и разрыв означал ровно одно — сломалась
    арифметика. Отсюда FAIL и «СБОЙ» на весь лист. С одобренным лимитом
    разрыв законен: банк дал меньше, чем нужно. FAIL здесь загорался бы у
    каждого, кто впишет настоящий лимит из term sheet, — та самая кричащая
    зря проверка, которую перестают читать.
    """
    _, evaluator, _ = starved
    assert str(evaluator.cell("ПРОВЕРКИ", "F32") or "") == "WARN"
    assert str(evaluator.cell("ПРОВЕРКИ", "B3") or "") != "СБОЙ"
    assert "WARN" in str(evaluator.cell("ПРОВЕРКИ", "A32") or ""), (
        "подпись не объясняет предупреждения — «не профинансирована» без "
        "причины читается как поломка книги")


def test_an_unexplained_gap_is_still_a_failure(starved, free):
    """Разрыв БЕЗ заданного лимита объяснить нечем — значит FAIL.

    Иначе смягчение сняло бы проверку вовсе: она молчала бы и там, где
    арифметика книги правда разъехалась.
    """
    _, _, book = starved
    formula = str(book["ПРОВЕРКИ"]["F32"].value or "")
    assert '"FAIL"' in formula, f"ветка FAIL исчезла: {formula}"
    cell = core._V4_PF_APPROVED_CELL
    assert f"${cell[0]}${cell[1:]}" in formula, (
        f"смягчение не привязано к одобренному лимиту: {formula}")
    # Без лимита лист по-прежнему без единого FAIL.
    _, free_evaluator, _ = free
    assert str(free_evaluator.cell("ПРОВЕРКИ", "F32") or "") == "OK"


def test_the_gap_is_counted_once(starved):
    """Факт строки 32 берётся из строк CF, а не вторым таким же вычитанием.

    Два ответа на «сколько не покрыто» в одной книге однажды разойдутся, и
    оба будут выглядеть верными.
    """
    _, evaluator, book = starved
    formula = str(book["ПРОВЕРКИ"]["B32"].value or "")
    assert f"B${core._V4_PF_UNCOVERED_ROW}" in formula, formula
    assert "D44" not in formula, f"строка 32 считает дыру заново: {formula}"
    total = sum(float(evaluator.cell(f"CF_{phase}", UNCOVERED) or 0)
                for phase in range(1, 5))
    assert float(evaluator.cell("ПРОВЕРКИ", "B32") or 0) == pytest.approx(total, abs=0.01)
