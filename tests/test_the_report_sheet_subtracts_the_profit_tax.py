"""Чистая прибыль ОТЧЁТа собирается из строк того же листа.

В шаблоне `ОТЧЕТ!B12` читал `КОНСОЛИДАТОР!L8` — там чистая собирается по
очередям, и из каждой вычитается её доля налога. У владельца Excel посчитал эти
доли нулями, и налог не вычелся ни у одной очереди: на экране стояло 56 598,6
вместо 42 128,6 — ровно на налог 14 470,0 больше (14.09.2026). Наш вычислитель
на том же файле давал верное число, то есть формулы законные, а расхождение
живёт между Excel и цепочкой раздачи.

Поэтому величина, которую читает человек, собирается из величин рядом с ней:
прибыль до налога минус налог минус НДС. Проверка держит оба конца — и то, что
формула ссылается на соседние строки, и то, что число не разошлось с прежней
цепочкой.
"""
from __future__ import annotations

import copy
import importlib
import io
import sys
from pathlib import Path

import openpyxl
import pytest

ROOT = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(ROOT))

core = importlib.import_module("main_legacy")
from xlsx_eval import Evaluator  # noqa: E402


def _workbook():
    inputs = dict(core.DEFAULT_INPUTS)
    # Проект обязан быть прибыльным: на убыточном налог равен нулю, и проверка
    # «налог вычтен» зелена при любом коде.
    inputs["apartment_price_th"] = 650
    inputs["commercial_price_th"] = 650
    phasing = {
        "enabled": True,
        "phase_gap_months": 12,
        "phases": [
            {"name": f"О{index}", "start_offset_months": 12 * (index - 1),
             "construction_months": 24}
            for index in range(1, 5)
        ],
    }
    content, _, _ = core.build_project_workbook(
        inputs, copy.deepcopy(core.TEP_DEFAULT), [], phasing,
        project_name="П", cache_values=False)
    return openpyxl.load_workbook(io.BytesIO(content), data_only=False)


def test_net_profit_is_built_from_the_rows_beside_it():
    """Формула ссылается на прибыль до налога и на налог, а не на чужой лист."""
    formula = _workbook()["ОТЧЕТ"]["B12"].value
    assert isinstance(formula, str)
    assert "B10" in formula and "B11" in formula
    assert "КОНСОЛИДАТОР" not in formula


def test_the_profit_tax_is_actually_subtracted():
    """Чистая = прибыль до налога − налог − НДС, и налог при этом не нулевой."""
    workbook = _workbook()
    sys.setrecursionlimit(400000)
    evaluator = Evaluator(workbook)
    before_tax = evaluator.cell("ОТЧЕТ", "B10")
    tax = evaluator.cell("ОТЧЕТ", "B11")
    net = evaluator.cell("ОТЧЕТ", "B12")
    vat = sum(evaluator.cell(f"CF_{phase}", "B21") for phase in range(1, 5))
    # Предохранитель: без налога и без НДС сравнивать нечего.
    assert tax > 1.0, "проект оказался без налога — проверка ничего не значит"
    assert vat > 1.0, "проект оказался без НДС — проверка ничего не значит"
    assert net == pytest.approx(before_tax - tax - vat, abs=0.01)


def test_the_new_formula_agrees_with_the_queue_allocation():
    """Число не разошлось с прежней цепочкой — правка про путь, не про методику."""
    workbook = _workbook()
    sys.setrecursionlimit(400000)
    evaluator = Evaluator(workbook)
    assert evaluator.cell("ОТЧЕТ", "B12") == pytest.approx(
        evaluator.cell("КОНСОЛИДАТОР", "L8"), abs=0.01)
