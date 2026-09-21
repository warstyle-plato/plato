"""Книга видит метры кладовых: они прибавляют подземную площадь.

Кладовые прибавляют подземную площадь, а не вычитаются из гаража (решение
владельца, 13.09.2026: «конечно они добавляют подземную площадь, иначе на
машиноместа и хватит просто площадей»), и движок считает базу подземной части
как «паркинг плюс кладовые». Книга кладовых не знала ВООБЩЕ: колонка «ГНС
подземная» блока очередей несёт только площадь паркинга, а строки «Кладовые»
листов ТЭП и ОТЧЕТ стоят в шаблоне владельца пустыми — при том что итоги НАД
ними их складывают.

Мера потери на умолчаниях с 400 кладовыми по 4 м² (величина владельца): 1 600 м²
подземной части и 140,8 млн ₽ СМР подземной части — 3 692,92 против 3 833,72.
Молча: файл скачивается и выглядит целым.

Проверка держит четыре утверждения, и у каждого свой диверсант:
  · строка «Кладовые» ТЭП несёт их метры (а не пусто);
  · строка «Кладовые» ОТЧЕТа — тоже;
  · подземная часть книги равна подземной части движка;
  · СМР подземной части и строительный объём выросли РОВНО на метры кладовых,
    а площадь гаража не шелохнулась — норматив 35 м²/место остался нормативом.

Предохранитель обязателен: без кладовых обе стороны совпадают на любом коде,
и проверка была бы зелена на сломанной книге. Поэтому пример держит кладовые
ненулевыми, а сравнение идёт ПАРОЙ прогонов — с ними и без них.

Запуск: python3 -m pytest tests/test_the_workbook_counts_the_storage_metres.py -q
"""

from __future__ import annotations

import io
import sys
from pathlib import Path

import openpyxl
import pytest

sys.path.insert(0, str(Path(__file__).resolve().parent.parent))
sys.path.insert(0, str(Path(__file__).resolve().parent))

import main as wrapper  # noqa: E402
from xlsx_eval import Evaluator  # noqa: E402

core = wrapper.core

# Величина владельца (13.09.2026): «Ну кладовые каким обычно площадью? Метра
# 3-4?». Своего замера площади кладовой у нас нет ни одного — ни в одном
# пресете нет пары «штуки + метры», — и это записано открытым вопросом.
STORAGE_UNITS = 400.0
STORAGE_SQM = 1600.0


def _build(storage_units: float, storage_sqm: float):
    """Расчёт и книга на одних вводных; кладовые задаются строкой ТЭП."""
    sys.setrecursionlimit(400000)
    inputs = dict(core.DEFAULT_INPUTS)
    tep = {key: dict(value) for key, value in core.TEP_DEFAULT.items()}
    tep["storage"] = dict(tep["storage"], units=storage_units, gns=storage_sqm,
                          total_area=storage_sqm, useful=storage_sqm,
                          saleable=storage_sqm)
    bundle = core._run_authoritative_model(inputs, tep, [], None)
    report = bundle.get("consolidated") or bundle
    content, _filename, meta = core.build_project_workbook(
        inputs, tep, [], bundle.get("phasing"), finance_hints={})
    book = openpyxl.load_workbook(io.BytesIO(content), data_only=False)
    return inputs, tep, report, meta, Evaluator(book)


@pytest.fixture(scope="module")
def with_storage():
    return _build(STORAGE_UNITS, STORAGE_SQM)


@pytest.fixture(scope="module")
def without_storage():
    return _build(0.0, 0.0)


def test_the_storage_row_of_the_tep_sheet_carries_its_metres(with_storage):
    """Строка «Кладовые» листа ТЭП несёт площадь, а не пустоту."""
    _inputs, _tep, _report, _meta, book = with_storage
    assert book.cell("ТЭП", "C7") == pytest.approx(STORAGE_SQM, abs=1.0)
    # Итог очереди = SUM(C4:C7) и кладовые складывает — значит пустая ячейка
    # была потерей ровно в том месте, где книга показывает площади.
    assert book.cell("ТЭП", "C8") == pytest.approx(
        sum(book.cell("ТЭП", f"C{row}") or 0 for row in (4, 5, 6, 7)), abs=1.0)


def test_the_storage_row_of_the_report_carries_its_metres(with_storage):
    """Строка «Кладовые» структуры продукта в ОТЧЕТе — тоже."""
    _inputs, _tep, _report, _meta, book = with_storage
    assert book.cell("ОТЧЕТ", "B49") == pytest.approx(STORAGE_SQM, abs=1.0)


def test_the_underground_area_of_the_book_matches_the_engine(with_storage):
    """Подземная часть книги равна подземной части движка."""
    _inputs, _tep, report, _meta, book = with_storage
    engine = (report.get("summary") or {}).get("underground_gns_sqm")
    assert engine is not None, "движок не назвал подземную площадь"
    assert book.cell("ТЭП", "C6") + book.cell("ТЭП", "C7") == pytest.approx(
        engine, abs=1.0)


def test_the_storage_adds_metres_and_the_garage_keeps_the_norm(
        with_storage, without_storage):
    """Кладовые ПРИБАВЛЯЮТ метры: гараж не шелохнулся, подземное СМР выросло
    ровно на их площадь, и строительный объём — тоже."""
    inputs, _tep, _report, _meta, with_book = with_storage
    _i0, _t0, _r0, _m0, bare = without_storage

    # Предохранитель: без кладовых сравнивать нечего, и проверка не значит
    # ничего — так уже зеленели наборы, ни разу не зашедшие в ветку.
    assert STORAGE_SQM > 0 and bare.cell("ТЭП", "C7") in (0, 0.0, None)

    # Площадь гаража — ровно «места × норматив», вычитания нет. Мест в книге
    # ДВА числа: N — проданные, AN — гостевые; строится их сумма, и делить на
    # проданные значило бы мерить норматив не тем, что построено.
    assert with_book.cell("ТЭП", "C6") == pytest.approx(bare.cell("ТЭП", "C6"), abs=0.5)
    norm = float(inputs["underground_area_per_space_sqm"])
    built = (with_book.cell("Параметры модели", "N88")
             + with_book.cell("Параметры модели", "AN88"))
    assert with_book.cell("ТЭП", "C6") == pytest.approx(built * norm, rel=0.002), (
        f"площадь гаража {with_book.cell('ТЭП', 'C6')} против "
        f"{built} мест × {norm} м²")

    # СМР подземной части выросло ровно на метры кладовых.
    rate = float(inputs["main_under_th_per_sqm"]) / 1000.0
    grew = with_book.cell("CAPEX", "B22") - bare.cell("CAPEX", "B22")
    assert grew == pytest.approx(STORAGE_SQM * rate, rel=0.01), (
        f"подземное СМР выросло на {grew:.2f} млн ₽ вместо "
        f"{STORAGE_SQM * rate:.2f}")

    # Строительный объём — тоже: на нём стоят ИРД, проектирование, сети.
    assert with_book.cell("ТЭП", "C36") - bare.cell("ТЭП", "C36") == pytest.approx(
        STORAGE_SQM, abs=1.0)


def test_the_volume_articles_see_the_storage(with_storage, without_storage):
    """Статьи на строительный объём выросли вместе с ним, а не остались на
    прежней базе: ИРД, П, РД, подготовка, сети, сдача, содержание."""
    inputs, _tep, _report, _meta, with_book = with_storage
    _i0, _t0, _r0, _m0, bare = without_storage
    rates = {16: "ird_th_per_sqm", 17: "design_p_th_per_sqm",
             18: "design_rd_th_per_sqm", 20: "preparation_th_per_sqm",
             23: "utilities_th_per_sqm", 25: "commissioning_th_per_sqm",
             26: "site_maintenance_th_per_sqm"}
    for row, key in rates.items():
        rate = float(inputs[key]) / 1000.0
        if rate <= 0:
            continue
        grew = with_book.cell("CAPEX", f"B{row}") - bare.cell("CAPEX", f"B{row}")
        assert grew > 0, f"строка {row} ({key}) не заметила метров кладовых"
        assert grew == pytest.approx(STORAGE_SQM * rate, rel=0.02), (
            f"строка {row} ({key}) выросла на {grew:.2f} млн ₽ вместо "
            f"{STORAGE_SQM * rate:.2f}")


def test_the_workbook_says_so_when_it_does_not_recognise_the_formula(with_storage):
    """Ни одна правленая формула не ушла в молчание: непонятая уходит в
    `missing`, а не считается по прежней базе."""
    _inputs, _tep, _report, meta, _book = with_storage
    noise = [line for line in (meta.get("missing") or [])
             if "кладов" in str(line) or "объём" in str(line)
             or "подземной части" in str(line)]
    assert noise == [], noise
