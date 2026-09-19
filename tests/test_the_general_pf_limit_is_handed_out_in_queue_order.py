"""Одобренный лимит у очередей — два поля, и оба настоящие.

Решение владельца (19.09.2026): «как правило дают общий лимит на всё в рамках
генеральных условий, а разбивку на очереди — когда одобряют индивидуальные
условия НКЛ». Значит:

- лимит НКЛ очереди задан (`phases[i].pf_limit_approved_mln`) — он и потолок;
- не задан, а общий (`pf_limit_approved_mln` проекта) задан — очереди
  достаётся ОСТАТОК общего после своей выборки предыдущих очередей: лимит
  раздаётся по порядку открытия ПФ, как банк выделяет НКЛ из генерального
  соглашения;
- не задано ничего — потолка нет, как и было.

Прежде заданное число доставалось КАЖДОЙ очереди целиком — ни то ни другое:
при 31 818 на проект и двух очередях свод печатал «Лимит ПФ» 63 636.

Книга считает то же теми же строками: потолок очереди стоит в её клетке блока
очередей (колонка AT «Вводных»), и остаток берётся из строк выборки предыдущих
листов CF. Своя выборка предыдущих от этой очереди не зависит — круга нет.

Запуск: python3 -m pytest tests/test_the_general_pf_limit_is_handed_out_in_queue_order.py -q
"""

from __future__ import annotations

import copy
import io
import sys
from pathlib import Path

import openpyxl
import pytest

sys.path.insert(0, str(Path(__file__).resolve().parent.parent))
sys.path.insert(0, str(Path(__file__).resolve().parent))

from xlsx_eval import Evaluator  # noqa: E402

import main as _wrapper  # noqa: E402

core = _wrapper.core

GENERAL = 31818
INDIVIDUAL = 12000


def _phasing(individual: float | None = None) -> dict:
    phasing = {
        "enabled": True, "mode": "phased", "user_enabled": True,
        "phase_count": 2, "phase_gap_months": 12,
        "phases": [{"name": "О1", "start_offset_months": 0, "construction_months": 24},
                   {"name": "О2", "start_offset_months": 12, "construction_months": 24}],
        "products": {key: [50, 50] for key in
                     ("apartments", "ground_commercial", "underground_parking", "storage")},
        "shared_cash": {}, "shared_allocation": {}, "social_objects": [],
        "carry_debt_forward": False,
    }
    if individual:
        phasing["phases"][1]["pf_limit_approved_mln"] = individual
    return phasing


def _inputs(general: float) -> dict:
    return {**core.DEFAULT_INPUTS, "purchase_price_mln": 12000,
            "project_start": "2027-01-01", "ird_months": 12,
            "construction_months": 24, "apartment_price_th": 700,
            "pf_limit_approved_mln": general}


def _tep() -> dict:
    return {key: dict(row) for key, row in core.TEP_DEFAULT.items()}


def _engine(general: float, individual: float | None = None) -> dict:
    return core.calculate_phased(core.PhasedCalcRequest(
        inputs=_inputs(general), tep=_tep(), rates=[], phasing=_phasing(individual)))


def _book(general: float, individual: float | None = None):
    sys.setrecursionlimit(400000)
    content, _, meta = core.build_project_workbook(
        _inputs(general), _tep(), [], _phasing(individual))
    assert meta["missing"] == [], meta["missing"]
    return Evaluator(openpyxl.load_workbook(io.BytesIO(content), data_only=False))


def _finance(bundle: dict, index: int) -> dict:
    return bundle["phases"][index]["result"]["finance"]


@pytest.fixture(scope="module")
def free():
    return _engine(0)


@pytest.fixture(scope="module")
def general_only():
    return _engine(GENERAL)


@pytest.fixture(scope="module")
def mixed():
    return _engine(GENERAL, INDIVIDUAL)


def test_the_fixture_needs_more_than_the_general_limit(free):
    """Предохранитель: общий лимит обязан быть меньше потребности первой
    очереди — иначе остатка второй хватает, и порядок раздачи не проверяется."""
    need_first = float(_finance(free, 0)["pf_limit_required"]) / 1e6
    assert need_first > GENERAL, (need_first, GENERAL)


def test_the_general_limit_is_a_remainder_not_a_copy(general_only):
    """Первой очереди — весь общий лимит, второй — что осталось (здесь ноль)."""
    first, second = _finance(general_only, 0), _finance(general_only, 1)
    assert first["pf_limit_cap_source"] == "general_remainder"
    assert float(first["pf_limit_approved"]) == pytest.approx(GENERAL * 1e6)
    assert float(first["peak_pf"]) <= GENERAL * 1e6 + 1.0
    assert second["pf_limit_cap_source"] == "general_remainder"
    assert float(second["pf_limit_approved"]) == 0.0
    # Ноль потолка — ноль выборки, а не «потолка нет».
    assert float(second["pf_draw_total"]) == 0.0
    assert float(second["pf_shortfall"]) > 0


def test_the_summary_prints_the_general_limit_once(general_only):
    """Свод не складывает остатки: у первой очереди остаток равен всему лимиту,
    и сумма вышла бы больше одобренного."""
    finance = general_only["consolidated"]["finance"]
    assert float(finance["pf_limit_general"]) == pytest.approx(GENERAL * 1e6)
    assert float(finance["pf_limit_approved"]) == pytest.approx(GENERAL * 1e6)
    assert float(finance["pf_limit"]) == pytest.approx(GENERAL * 1e6)


def test_an_individual_limit_wins_over_the_remainder(mixed, general_only):
    """Лимит НКЛ очереди задан — он и потолок, остаток общего не при чём."""
    second = _finance(mixed, 1)
    assert second["pf_limit_cap_source"] == "individual"
    assert float(second["pf_limit_approved"]) == pytest.approx(INDIVIDUAL * 1e6)
    assert float(second["pf_draw_total"]) == pytest.approx(INDIVIDUAL * 1e6, abs=1.0)
    # Первая очередь свой лимит не назвала — ей по-прежнему остаток общего.
    assert float(_finance(mixed, 0)["pf_limit_approved"]) == pytest.approx(
        float(_finance(general_only, 0)["pf_limit_approved"]))


def test_without_a_general_limit_the_other_queue_is_free():
    """Индивидуальный лимит одной очереди не ограничивает соседнюю."""
    bundle = _engine(0, INDIVIDUAL)
    assert _finance(bundle, 0)["pf_limit_cap_source"] == ""
    assert float(_finance(bundle, 0)["pf_shortfall"]) == 0.0
    assert _finance(bundle, 1)["pf_limit_cap_source"] == "individual"


@pytest.mark.parametrize("general,individual", [(GENERAL, None), (GENERAL, INDIVIDUAL), (0, INDIVIDUAL)])
def test_the_book_hands_the_limit_out_the_same_way(general, individual):
    """Книга: потолок каждой очереди, её выборка и дыра — как у движка."""
    bundle = _engine(general, individual)
    evaluator = _book(general, individual)
    for index in range(2):
        finance = _finance(bundle, index)
        row = core._V4_CF_QUEUE_ENABLED_ROW + index
        cap = float(evaluator.cell("Параметры модели", f"{core._V4_PF_QUEUE_CAP_COL}{row}") or 0)
        assert cap == pytest.approx(float(finance["pf_limit_approved"]) / 1e6, abs=0.5), index
        draw = float(evaluator.cell(f"CF_{index + 1}", "B45") or 0)
        assert draw == pytest.approx(float(finance["pf_draw_total"]) / 1e6, abs=0.5), index
        gap = float(evaluator.cell(f"CF_{index + 1}", f"B{core._V4_PF_UNCOVERED_ROW}") or 0)
        assert gap == pytest.approx(float(finance["pf_shortfall"]) / 1e6, abs=0.5), index
    assert str(evaluator.cell("ПРОВЕРКИ", f"F{core._V4_PF_UNCOVERED_PARITY_ROW}") or "") == "OK"
    assert str(evaluator.cell("ПРОВЕРКИ", "B3") or "") != "СБОЙ"


def test_the_individual_limit_is_an_input_cell_of_the_queue_block():
    """Поле очереди лежит в блоке очередей со стилем ввода — значит лист ввода
    его заберёт, а книга, где его правят, пересчитает потолок сама."""
    evaluator = _book(GENERAL, INDIVIDUAL)
    row = core._V4_CF_QUEUE_ENABLED_ROW + 1
    assert float(evaluator.cell("Параметры модели", f"{core._V4_PF_QUEUE_INDIVIDUAL_COL}{row}") or 0) == INDIVIDUAL
    header = str(evaluator.cell("Параметры модели", f"{core._V4_PF_QUEUE_INDIVIDUAL_COL}87") or "")
    assert "НКЛ" in header, header
