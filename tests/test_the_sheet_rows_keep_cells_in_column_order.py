"""Ячейки строки листа идут по возрастанию колонки — иначе Excel их теряет.

Блок «РАЗДАЧА СВОДНОГО НАЛОГА ПО ГОДАМ» писал строку как
`A B C D E F G L M N O P H I J K`: сперва запас, потом налог. Excel требует
возрастающего порядка ссылок и молча выбрасывает ячейки, стоящие после старшей
колонки, — колонки «Налог О1…О4» у владельца были ПУСТЫМИ, «Итого» под ними
нулём, и чистая прибыль ОТЧЁТа шла без налога на прибыль: на его проекте
56 598,6 вместо 42 128,6 при налоге 14 470,0.

Наш вычислитель читает ячейки по координате и порядка не замечает — то есть
проверки были зелёными ровно на той поломке, которую видно в Excel. Поэтому
сторож смотрит на САМ ФАЙЛ, а не на посчитанные значения, и проверяет все
листы: следующий блок попадёт под него тем, что он появился.

Рядом стоит проверка самого приёма: на подделанной строке сторож обязан
падать, иначе «нарушений нет» значит, что не сработал он сам.
"""

from __future__ import annotations

import io
import re
import sys
import zipfile

import openpyxl
import pytest

import main_legacy as core
from xlsx_eval import Evaluator


def _column_number(letters: str) -> int:
    number = 0
    for letter in letters:
        number = number * 26 + ord(letter) - 64
    return number


def _rows_out_of_order(sheet_xml: str) -> list[tuple[int, list[str]]]:
    broken: list[tuple[int, list[str]]] = []
    for row in re.finditer(r'<x:row r="(\d+)"[^>]*>(.*?)</x:row>', sheet_xml, re.S):
        coords = re.findall(r'<x:c r="([A-Z]+)\d+"', row.group(2))
        numbers = [_column_number(coord) for coord in coords]
        if numbers != sorted(numbers):
            broken.append((int(row.group(1)), coords))
    return broken


def _workbook(**overrides: float) -> bytes:
    inputs = dict(core.DEFAULT_INPUTS)
    inputs.update(overrides)
    phasing = {
        "enabled": True,
        "phases": [
            {"name": f"О{index + 1}", "start_offset_months": index * 24, "share": 0.25}
            for index in range(4)
        ],
    }
    content, _name, _report = core.build_project_workbook(
        inputs, core.TEP_DEFAULT, phasing=phasing)
    return content


def _sheets(content: bytes) -> dict[str, str]:
    archive = zipfile.ZipFile(io.BytesIO(content))
    return {
        name: archive.read(name).decode("utf-8")
        for name in archive.namelist()
        if name.startswith("xl/worksheets/")
    }


def test_no_row_of_the_workbook_puts_cells_out_of_column_order() -> None:
    sheets = _sheets(_workbook())
    assert len(sheets) >= 15, "листы не прочитаны — сторож судил бы о пустоте"

    broken: list[str] = []
    for name, sheet_xml in sorted(sheets.items()):
        for row, coords in _rows_out_of_order(sheet_xml):
            broken.append(f"{name} строка {row}: {' '.join(coords)}")
    assert not broken, "Excel выбросит ячейки, стоящие не по возрастанию:\n" + "\n".join(broken)


def test_the_guard_falls_on_a_shuffled_row() -> None:
    """Проверка самого приёма: переставленную строку сторож обязан найти."""
    shuffled = '<x:row r="7">' + "".join(
        f'<x:c r="{coord}7"><x:v>1</x:v></x:c>' for coord in ("A", "B", "L", "H")
    ) + "</x:row>"
    assert _rows_out_of_order(shuffled) == [(7, ["A", "B", "L", "H"])]
    ordered = '<x:row r="7">' + "".join(
        f'<x:c r="{coord}7"><x:v>1</x:v></x:c>' for coord in ("A", "B", "H", "L")
    ) + "</x:row>"
    assert _rows_out_of_order(ordered) == []


def test_the_queue_tax_columns_carry_the_consolidated_tax() -> None:
    """Колонки «Налог О1…О4» несут налог — та величина, что терялась в Excel.

    Предохранитель обязателен: на умолчаниях проект убыточный, налога нет
    вовсе, и проверка была бы зелёной при любом порядке ячеек.
    """
    content = _workbook(apartment_price_th=650.0, commercial_price_th=650.0)

    total_row = 0
    for sheet_xml in _sheets(content).values():
        if "РАЗДАЧА СВОДНОГО НАЛОГА" not in sheet_xml:
            continue
        rows = [int(number) for number in re.findall(r'<x:row r="(\d+)"', sheet_xml)]
        total_row = max(rows) - 1        # «Итого», под ним строка расхождения
    assert total_row, "блок раздачи налога на листе не найден"

    sys.setrecursionlimit(400000)
    evaluator = Evaluator(openpyxl.load_workbook(io.BytesIO(content), data_only=False))
    consolidated = evaluator.cell("КОНСОЛИДАТОР", f"B{total_row}") or 0.0
    assert consolidated > 1.0, (
        "проект без налога — проверка не отличила бы верный порядок от сломанного")

    by_queue = sum(
        evaluator.cell("КОНСОЛИДАТОР", f"{column}{total_row}") or 0.0
        for column in core._V4_TAX_SHARE_COLUMNS
    )
    assert by_queue == pytest.approx(consolidated, rel=1e-9), (
        f"раздача по очередям {by_queue} не равна сводному налогу {consolidated}")
