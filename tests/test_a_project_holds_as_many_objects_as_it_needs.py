"""Проект держит столько объектов одного типа, сколько нужно.

Владелец (28.09.2026): «Речь шла о возможности добавлять дубли — хоть 2, хоть
5 объектов». До этого вторые объекты были тремя зашитыми строками реестра:
третий офис или второй ФОК завести было нельзя. Теперь тип объекта порождает
экземпляры (`object_instance`), а какие из них есть в проекте, говорит явный
список `object_instances` — со своими площадями, ценой, очередью, паркингом
и календарём у каждого.

Запуск: python3 -m pytest tests/test_a_project_holds_as_many_objects_as_it_needs.py -q
"""

from __future__ import annotations

import copy
import io
import sys
from pathlib import Path

import openpyxl
import pytest

ROOT = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(ROOT))
sys.path.insert(0, str(ROOT / "tests"))

import main as wrapper  # noqa: E402
from xlsx_eval import Evaluator  # noqa: E402

core = wrapper.core

PORT = 18973
OFFICES = ("offices", "offices2", "offices3", "offices4", "offices5")
# Пять офисов в трёх очередях, у каждого свои метры и гараж; второй ФОК
# продаётся и назначен медцентром — при том что первый ФОК передаётся городу.
QUEUES = {"offices": 1, "offices2": 1, "offices3": 2, "offices4": 3, "offices5": 3,
          "sports2": 2}


def _five_offices() -> dict:
    x = dict(core.DEFAULT_INPUTS)
    x.update(apartment_price_th=650, commercial_price_th=650, parking_price_th=5000,
             offices_enabled=True)
    for number, key in enumerate(OFFICES[1:], start=2):
        prefix = core._BY_KEY[key].prefix
        x.update({f"{prefix}_enabled": True, f"{prefix}_gba_sqm": 6000 + 1000 * number,
                  f"{prefix}_saleable_sqm": 4000 + 500 * number,
                  f"{prefix}_parking_under_spaces": 10 * number})
    x.update(sports2_enabled=True, sports2_disposition="sale",
             sports2_purpose="healthcare", sports_disposition="transfer")
    x["object_instances"] = ["offices2", "offices3", "offices4", "offices5", "sports2"]
    return x


def _phasing() -> dict:
    return {"enabled": True, "phase_count": 3, "phase_gap_months": 12,
            "cost_inflation_pct": 8,
            "phases": [{"name": f"О{i + 1}", "start_offset_months": 12 * i,
                        "construction_months": 24} for i in range(3)],
            "discrete": dict(QUEUES)}


def _revenue(result: dict, key: str) -> float:
    return float(next((p.get("revenue") or 0.0) for p in result["report"]["products"]
                      if p["key"] == key) or 0.0) if any(
        p["key"] == key for p in result["report"]["products"]) else 0.0


# --- реестр ------------------------------------------------------------------

def test_every_type_has_its_instances() -> None:
    for base in core.OBJECT_TYPES:
        own = [o for o in core.STANDALONE_OBJECTS if o.family == base.key]
        assert len(own) == core.OBJECT_INSTANCES_MAX - 1, base.key
        for number, obj in enumerate(own, start=2):
            assert obj.key == f"{base.key}{number}" and obj.prefix == f"{base.prefix}{number}"
            assert obj.product == base.key
            assert (obj.measure, obj.garage, obj.garage_sellable) == (
                base.measure, base.garage, base.garage_sellable)
            # Своё поле цены, себестоимости и признака продажи — не поле типа.
            assert obj.rate_price.startswith(obj.prefix + "_")
            assert obj.rate_cost.startswith(obj.prefix + "_")
            assert not base.sale_gate or obj.sale_gate == f"{obj.prefix}_disposition"
            assert core.DEFAULT_INPUTS[obj.enabled_key] is False
    # Прежние «вторые» объекты — те же ключи, сохранённые проекты читаются.
    for key in core.LEGACY_SECOND_OBJECTS:
        assert key in core.OBJECT_INSTANCES
    assert core.OBJECT_INSTANCES["standalone_retail2"].prefix == "retail2"


def test_the_instance_field_must_belong_to_its_type() -> None:
    bad = core.OBJECT_TYPES[0]._replace(rate_cost="foreign_cost_th_per_sqm")
    with pytest.raises(ValueError, match="не принадлежит"):
        core.object_instance(bad, 2)


# --- состав проекта и миграция -----------------------------------------------

def test_a_saved_project_keeps_its_second_objects() -> None:
    """Проект до экземпляров списка не несёт: состав — из того, что он сказал."""
    x = dict(core.DEFAULT_INPUTS)
    assert core.project_object_instances(x) == ()
    x.update(offices2_enabled=True, above_parking2_spaces=120)
    # Выключенный второй паркинг с набранными местами — данные человека.
    assert core.project_object_instances(x) == ("offices2", "above_parking2")
    # Выключенный второй ТЦ с умолчаниями — не объект проекта.
    assert "standalone_retail2" not in core.project_object_instances(x)
    # Отсутствующий ключ — не пустой список: явный [] значит «экземпляров нет».
    x["object_instances"] = []
    assert core.project_object_instances(x) == ()


def test_the_list_decides_and_the_rest_is_named() -> None:
    x = dict(core.DEFAULT_INPUTS)
    x.update(offices3_enabled=True, object_instances=["offices2", "offices9"])
    assert core.project_object_instances(x) == ("offices2",)
    notes = core.object_instance_notes(x)
    assert any("offices9" in note and "предел" in note for note in notes), notes
    assert any("МФОЦ / офисы 3" in note and "не считается" in note for note in notes), notes
    applied, _ = core.object_instances_applied(x, copy.deepcopy(core.TEP_DEFAULT))
    assert applied["offices3_enabled"] is False
    assert x["offices3_enabled"] is True, "присланное не правится"


def test_adding_stops_at_the_limit_with_a_reason() -> None:
    x = dict(core.DEFAULT_INPUTS)
    added = [core.object_instance_add(x, "sports") for _ in range(core.OBJECT_INSTANCES_MAX - 1)]
    assert added == [f"sports{n}" for n in range(2, core.OBJECT_INSTANCES_MAX + 1)]
    assert all(x[f"{key}_enabled"] for key in added)
    reason = core.object_instance_refusal(x, "sports")
    assert "ФОК / медцентр" in reason and str(core.OBJECT_INSTANCES_MAX) in reason
    with pytest.raises(ValueError):
        core.object_instance_add(x, "sports")
    core.object_instance_remove(x, "sports3")
    assert "sports3" not in core.project_object_instances(x)
    assert core.object_instance_refusal(x, "sports") == ""
    # Снова добавленный встаёт умолчаниями, а не значениями удалённого.
    x["sports3_gba_sqm"] = 777
    assert core.object_instance_add(x, "sports") == "sports3"
    assert x["sports3_gba_sqm"] == core.DEFAULT_INPUTS["sports3_gba_sqm"]
    with pytest.raises(ValueError):
        core.object_instance_remove(x, "sports")


# --- движок ------------------------------------------------------------------

@pytest.fixture(scope="module")
def phased() -> dict:
    return core._run_authoritative_model(
        _five_offices(), copy.deepcopy(core.TEP_DEFAULT), [], _phasing())


def test_each_office_is_built_in_its_own_queue(phased) -> None:
    for key, queue in QUEUES.items():
        by_phase = [_revenue(item["result"], key) for item in phased["phases"]]
        assert [index + 1 for index, value in enumerate(by_phase) if value > 0] == [queue], (
            key, by_phase)
    summary = phased["consolidated"]["summary"]
    assert summary["object_instances"] == ["offices2", "offices3", "offices4", "offices5",
                                           "sports2"]
    assert summary["object_instance_notes"] == []


def test_the_second_sports_object_is_sold_by_its_own_switch(phased) -> None:
    """Продажу второго ФОКа читал признак ПЕРВОГО: переданный первый гасил
    выручку всех ФОКов, и книга с движком расходились на весь второй."""
    assert sum(_revenue(item["result"], "sports2") for item in phased["phases"]) > 0
    assert sum(_revenue(item["result"], "sports") for item in phased["phases"]) == 0


def test_an_instance_outside_the_project_is_not_counted() -> None:
    x = _five_offices()
    single = core._run_authoritative_model(x, copy.deepcopy(core.TEP_DEFAULT), [], {})
    x["object_instances"] = ["offices2", "offices3", "offices4", "sports2"]
    without = core._run_authoritative_model(x, copy.deepcopy(core.TEP_DEFAULT), [], {})
    assert _revenue(single["consolidated"], "offices5") > 0
    assert _revenue(without["consolidated"], "offices5") == 0
    assert without["consolidated"]["summary"]["revenue"] < single["consolidated"]["summary"]["revenue"]
    assert any("офисы 5" in note for note in
               without["consolidated"]["summary"]["object_instance_notes"])


# --- книга: паритет -----------------------------------------------------------

@pytest.fixture(scope="module")
def phased_book():
    sys.setrecursionlimit(400000)
    content, _, meta = core.build_project_workbook(
        _five_offices(), copy.deepcopy(core.TEP_DEFAULT), [], _phasing(),
        project_name="Пять офисов")
    book = openpyxl.load_workbook(io.BytesIO(content), data_only=False)
    return book, Evaluator(book), meta


def test_the_book_agrees_with_the_engine_on_five_offices(phased_book) -> None:
    book, evaluator, meta = phased_book
    assert meta["missing"] == [], meta["missing"]
    for row in range(76, 85):
        assert book["ПРОВЕРКИ"][f"C{row}"].value is not None, f"цель строки {row}"
        assert evaluator.cell("ПРОВЕРКИ", f"F{row}") == "OK", (
            str(book["ПРОВЕРКИ"][f"A{row}"].value),
            evaluator.cell("ПРОВЕРКИ", f"B{row}"), evaluator.cell("ПРОВЕРКИ", f"C{row}"))
    assert evaluator.cell("ПРОВЕРКИ", "B3") in ("ПРОЙДЕНО", "ПРОЙДЕНО С ПРЕДУПРЕЖДЕНИЯМИ")


def test_the_book_writes_only_the_project_instances() -> None:
    """Блок экземпляра пишется в книгу, только если экземпляр в проекте:
    включённый, но не заведённый, книгу не раздувает и числа не несёт."""
    x = _five_offices()
    x["offices2_enabled"] = False
    x["retail3_enabled"] = True  # включён, но в составе проекта его нет
    applied, _ = core.object_instances_applied(x, copy.deepcopy(core.TEP_DEFAULT))
    got = core._v4_book_objects(applied)
    assert {"offices3", "offices4", "offices5", "sports2"} <= got
    assert not {"offices2", "standalone_retail3", "standalone_retail2"} & got


# --- страница ----------------------------------------------------------------

MIRROR = """(cases) => cases.map(c => {
  const saved = inputs; inputs = Object.assign(cloneValue(INPUT_DEFAULT), c);
  if (!('object_instances' in c)) delete inputs.object_instances;
  const got = projectInstances(); inputs = saved; return got;
})"""

MIRROR_CASES = [
    {},
    {"offices2_enabled": True, "above_parking2_spaces": 120},
    {"offices3_enabled": True, "object_instances": ["offices2", "offices9"]},
    {"object_instances": "sports2, retail3"},
    {"object_instances": []},
]

STATE = """() => ({
  instances: projectInstances(),
  groups: Array.from(document.querySelectorAll('#inputGroups details[data-group]')).map(d=>d.dataset.group),
  tep: Array.from(document.querySelectorAll('#tepBody td:first-child')).map(td=>td.textContent),
  queues: Array.from(document.querySelectorAll('#assignObjects select[data-object]')).map(s=>s.dataset.object),
  officeOption: !!document.querySelector('#objectAddType option[value="offices"]:not([disabled])'),
  bar: (document.getElementById('objectAddBar')||{}).textContent||'',
  gba4: Number(inputs.offices4_gba_sqm||0),
})"""


@pytest.fixture(scope="module")
def page_run():
    pytest.importorskip("playwright.sync_api")
    from playwright.sync_api import sync_playwright

    import browser

    path = browser.chromium_or_skip()
    errors: list[str] = []
    out: dict = {}
    with browser.serve(core.app, PORT) as base, sync_playwright() as pw:
        with pw.chromium.launch(executable_path=str(path)) as engine:
            page = engine.new_page(viewport={"width": 1440, "height": 900})
            page.on("pageerror", lambda exc: errors.append(str(exc)))
            page.on("dialog", lambda dialog: dialog.accept())
            page.goto(base, wait_until="domcontentloaded")
            page.evaluate("localStorage.removeItem('plato_v04')")
            page.reload(wait_until="domcontentloaded")
            page.wait_for_timeout(1500)
            out["mirror"] = page.evaluate(MIRROR, MIRROR_CASES)
            page.evaluate("renderPhasing&&renderPhasing()")
            out["start"] = page.evaluate(STATE)
            # Четыре офиса сверх первого — кнопкой, как человек.
            for _ in range(core.OBJECT_INSTANCES_MAX - 1):
                page.select_option("#objectAddType", "offices")
                with page.expect_response(lambda r: "/calculate" in r.url) as resp:
                    page.click("#objectAddButton")
                out.setdefault("summaries", []).append(
                    (resp.value.json().get("summary")
                     or (resp.value.json().get("consolidated") or {}).get("summary") or {}))
            # Своё число у четвёртого — живёт через перерисовку и пересчёт.
            page.fill("#f_offices4_gba_sqm", "12345")
            page.dispatch_event("#f_offices4_gba_sqm", "change")
            page.evaluate("renderInputs();renderTep();renderPhasing&&renderPhasing()")
            out["full"] = page.evaluate(STATE)
            with page.expect_response(lambda r: "/calculate" in r.url):
                page.click('button.object-remove[data-object="offices3"]')
            page.evaluate("renderPhasing&&renderPhasing();persistLocalSilently()")
            out["removed"] = page.evaluate(STATE)
            page.reload(wait_until="domcontentloaded")
            page.wait_for_timeout(1500)
            page.evaluate("renderPhasing&&renderPhasing()")
            out["reloaded"] = page.evaluate(STATE)
            page.close()
    out["errors"] = errors
    return out


def test_the_page_mirrors_the_engine_list(page_run) -> None:
    assert page_run["errors"] == [], page_run["errors"]
    expected = []
    for case in MIRROR_CASES:
        x = dict(core.DEFAULT_INPUTS)
        x.update(case)
        expected.append(list(core.project_object_instances(x)))
    assert page_run["mirror"] == expected


def test_the_page_adds_and_removes_objects(page_run) -> None:
    assert page_run["errors"] == [], page_run["errors"]
    start, full = page_run["start"], page_run["full"]
    assert start["instances"] == []
    assert "МФОЦ / офисы 2" not in start["groups"]
    assert not any("Офисы 2" in text for text in start["tep"])
    assert start["officeOption"]
    assert full["instances"] == ["offices2", "offices3", "offices4", "offices5"]
    for number in range(2, 6):
        assert f"МФОЦ / офисы {number}" in full["groups"]
        assert any(text.startswith(f"Офисы {number}") for text in full["tep"]), full["tep"]
    # Шестой офис не обрезается молча: тип заперт, причина названа.
    assert not full["officeOption"]
    assert "МФОЦ / офисы" in full["bar"] and "больше реестр не заводит" in full["bar"]
    assert full["gba4"] == 12345
    # Расчёт видит тот же состав, что страница.
    assert page_run["summaries"][-1].get("object_instances") == full["instances"]
    removed = page_run["removed"]
    assert removed["instances"] == ["offices2", "offices4", "offices5"]
    assert "МФОЦ / офисы 3" not in removed["groups"]
    assert not any(text.startswith("Офисы 3") for text in removed["tep"])
    assert removed["officeOption"]


def test_the_page_keeps_objects_across_reload(page_run) -> None:
    reloaded = page_run["reloaded"]
    assert reloaded["instances"] == ["offices2", "offices4", "offices5"]
    assert "МФОЦ / офисы 4" in reloaded["groups"] and "МФОЦ / офисы 3" not in reloaded["groups"]
    assert reloaded["gba4"] == 12345
