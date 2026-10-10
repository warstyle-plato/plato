"""Скачанная гостиничная модель пересчитывается от вводных, как движок."""
import copy
import io
import sys
from datetime import date

import pytest
from openpyxl import load_workbook

import hotel_presets
import hotel_workbook as hw
import hotel_workbook_capex as capex
import main_legacy as core
from xlsx_eval import Evaluator


@pytest.fixture(scope="module")
def original():
    x = hotel_presets.apply_preset(copy.deepcopy(core.DEFAULT_INPUTS), "hotel1")
    x.update(project_kind=core.PROJECT_KIND_HOTEL, hotel_hold_years=6)
    t = copy.deepcopy(core.TEP_DEFAULT)
    content, _, _ = core.build_project_workbook(x, t, [], {}, project_name="Отель 1")
    return x, t, content


def evaluate(content):
    sys.setrecursionlimit(max(sys.getrecursionlimit(), 100000))
    book = load_workbook(io.BytesIO(content))
    return book, Evaluator(book)


@pytest.mark.parametrize("field,cell,value", [
    ("hotel_adr_rub", f"C{hw._ROWS['adr']}", 22000),
    ("hotel_keys", f"C{hw._ROWS['keys']}", 240),
    ("hotel_exit_multiple", f"C{hw._ROWS['multiple']}", 6),
    ("hotel_financing", f"C{hw._ROWS['financing']}", "Без кредита"),
    ("hotel_repayment", f"C{hw._ROWS['repayment']}", "Аннуитет"),
    ("hotel_hold_years", f"C{hw._ROWS['hold']}", 8),
    ("hotel_cost_th_per_sqm", f"C{capex.INPUT_ROWS['hotel_cost_th_per_sqm']}", 115),
    ("hotel_gba_sqm", f"C{capex.INPUT_ROWS['hotel_gba_sqm']}", 33000),
    ("construction_months", f"C{capex.INPUT_ROWS['construction_months']}", 18),
    ("scenario_cost_multiplier", f"C{capex.INPUT_ROWS['scenario_cost_multiplier']}", 1.1),
    ("project_start", f"C{hw._ROWS['start']}", date(2028, 1, 1)),
])
def test_editing_the_downloaded_book_reaches_the_same_answer_as_the_engine(original, field, cell, value):
    x, t, content = original
    book, evaluator = evaluate(content)
    old_irr = evaluator.cell(hw.TOTALS_SHEET, "B7")
    book[hw.INPUTS_SHEET][cell] = value
    evaluator = Evaluator(book)
    inputs = copy.deepcopy(x)
    choices = {"Без кредита": "none", "Аннуитет": "annuity"}
    inputs[field] = choices.get(value, value)
    engine = core._run_authoritative_model(inputs, copy.deepcopy(t), [], {})["consolidated"]["finance"]["hotel"]
    for row, key in ((4, "npv_project"), (5, "irr_project"), (6, "npv_equity"),
                     (7, "irr_equity"), (14, "loan_peak"), (18, "equity_peak")):
        actual = evaluator.cell(hw.TOTALS_SHEET, f"B{row}")
        assert actual == pytest.approx(engine["kpi"][key], rel=1e-6, abs=1e-6), (field, key, actual)
    assert evaluator.cell(hw.TOTALS_SHEET, "B7") != pytest.approx(old_irr)
    assert evaluator.cell(hw.CHECK_SHEET, "H3") == "ВВОДНЫЕ ИЗМЕНЕНЫ — НУЖНА НОВАЯ СВЕРКА"


def test_the_forecast_is_an_input_to_the_loan_and_returns(original):
    x, t, content = original
    book, evaluator = evaluate(content)
    for row in range(120, 120 + book[hw.CALC_SHEET].max_row - hw.FIRST_ROW + 1):
        book[hw.INPUTS_SHEET][f"F{row}"] = 0.08
    evaluator = Evaluator(book)
    rates = [{"date": x["project_start"], "low": 8, "base": 8, "high": 8}]
    engine = core._run_authoritative_model(x, t, rates, {})["consolidated"]["finance"]["hotel"]
    assert evaluator.cell(hw.TOTALS_SHEET, "B7") == pytest.approx(engine["kpi"]["irr_equity"], rel=1e-6)


def test_a_formula_error_is_visible_even_after_editing_inputs(original):
    book, _ = evaluate(original[2])
    book[hw.INPUTS_SHEET][f"C{hw._ROWS['adr']}"] = 22000
    book[hw.CALC_SHEET][f"{hw._c('project_cf')}{hw.FIRST_ROW+42}"] = "=1"
    evaluator = Evaluator(book)
    assert evaluator.cell(hw.CHECK_SHEET, "H3") == "ЕСТЬ РАСХОЖДЕНИЯ"


def test_the_workbook_is_presentable_and_uses_recalculating_returns(original):
    book, _ = evaluate(original[2])
    assert {"Затраты", "ОПУ", "Денежный поток", "Кредит", "Налоги", "Чувствительность"} <= set(book.sheetnames)
    assert "IRR(" in book[hw.TOTALS_SHEET]["B5"].value
    assert "IRR(" in book[hw.TOTALS_SHEET]["B7"].value
    assert book[hw.CALC_SHEET][f"{hw._c('capex')}{hw.FIRST_ROW+42}"].data_type == "f"
    assert book[hw.TOTALS_SHEET]["B5"].number_format == hw.PERCENT
    assert book[hw.TOTALS_SHEET]["A1"].font.name == "Montserrat"
    assert book.calculation.fullCalcOnLoad
    cached = load_workbook(io.BytesIO(original[2]), data_only=True)
    assert isinstance(cached[hw.TOTALS_SHEET]["B7"].value, float)
    assert cached[hw.CHECK_SHEET]["H3"].value == "ПРОЙДЕНО"
    for name in (hw.TOTALS_SHEET, hw.INPUTS_SHEET, "Затраты", "Денежный поток", "Кредит", "Налоги"):
        assert book[name].freeze_panes
        assert book[name].print_title_rows
        assert book[name].sheet_view.showGridLines is False


def test_sensitivity_recalculates_the_model_and_follows_changed_base_inputs(original):
    x, t, content = original
    book, evaluator = evaluate(content)
    inputs = copy.deepcopy(x)
    inputs["hotel_adr_rub"] *= 0.9
    engine = core._run_authoritative_model(inputs, t, [], {})["consolidated"]["finance"]["hotel"]
    assert evaluator.cell("Чувствительность", "E5") == pytest.approx(engine["kpi"]["irr_equity"], rel=1e-6)
    book[hw.INPUTS_SHEET][f"C{hw._ROWS['adr']}"] = 22000
    evaluator = Evaluator(book)
    inputs["hotel_adr_rub"] = 22000 * 1.1
    engine = core._run_authoritative_model(inputs, t, [], {})["consolidated"]["finance"]["hotel"]
    assert evaluator.cell("Чувствительность", "E6") == pytest.approx(engine["kpi"]["irr_equity"], rel=1e-6)


def test_a_purchase_schedule_can_be_edited_without_regenerating_the_book(original):
    x, t, content = original
    book, evaluator = evaluate(content)
    book[hw.INPUTS_SHEET][f"C{capex.INPUT_ROWS['purchase_price_mln']}"] = 900
    book[hw.INPUTS_SHEET]["G77"] = 0.3
    book[hw.INPUTS_SHEET]["F78"] = 8
    book[hw.INPUTS_SHEET]["G78"] = 0.7
    evaluator = Evaluator(book)
    inputs = {**x, "purchase_price_mln": 900, "purchase_schedule": "30%@0; 70%@8"}
    engine = core._run_authoritative_model(inputs, t, [], {})["consolidated"]["finance"]["hotel"]
    assert evaluator.cell(hw.TOTALS_SHEET, "B7") == pytest.approx(engine["kpi"]["irr_equity"], rel=1e-6)
    assert evaluator.cell(hw.CHECK_SHEET, "H3") == "ВВОДНЫЕ ИЗМЕНЕНЫ — НУЖНА НОВАЯ СВЕРКА"
