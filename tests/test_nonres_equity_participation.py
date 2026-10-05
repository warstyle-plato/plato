"""«Собственное участие» нежилого проекта: вложения, возврат, результат.

Владелец (05.10.2026): в блоке «Финансирование» нежилого проекта не было
собственного участия — экономики проекта при продаже по ДКП или доходным
методом не видно. Проверки держат обещания:

* таблица читает поток капитала проекта — тот же, по которому считан IRR;
  части вложений и возврата складываются в этот поток;
* оценка удержанного объекта — не деньги: названа отдельно, денежный итог
  без неё;
* в проекте с ПФ (поток несёт жильё) таблицы нет;
* страница (стенд на node) и PDF (раздел «Финансирование») печатают ту же
  таблицу, а книга нежилого проекта считает её формулами и сверяет с движком.

Запуск: python3 -m pytest tests/test_nonres_equity_participation.py -q
"""

from __future__ import annotations

import copy
import io
import json
import sys
from pathlib import Path

import pytest

ROOT = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(ROOT))
sys.path.insert(0, str(ROOT / "tests"))

import main_legacy as core  # noqa: E402
import page_blocks  # noqa: E402
from test_nonres_debt_metric import _nonres, _nonres_inputs  # noqa: E402
from test_nonres_object_result import _spec  # noqa: E402


def _rows(result: dict) -> dict[str, dict]:
    return {r["label"]: r for r in result["report"]["equity_participation"]["rows"]}


@pytest.mark.parametrize("over", [{}, {"offices_strategy": "direct"},
                                  {"offices_exit_mode": "hold"}])
def test_the_parts_add_up_to_the_equity_flow_of_the_irr(over) -> None:
    result = _nonres(**over)
    eq = result["report"]["equity_participation"]
    flows = result["cashflow"]["equity"]
    parts = eq["parts"]
    contributed = parts["in_build"] + parts["in_operation"] + parts["in_exit"]
    returned = parts["out_vat"] + parts["out_running"] + parts["out_exit"]
    assert contributed == pytest.approx(eq["contributed"], rel=1e-12)
    assert returned - contributed + eq["residual"] == pytest.approx(sum(flows), rel=1e-9, abs=1.0)
    assert sum(y["net"] for y in eq["years"]) == pytest.approx(returned - contributed, rel=1e-9, abs=1.0)
    rows = _rows(result)
    irr_label = next(label for label in rows if label.startswith("IRR собственного капитала"))
    assert rows[irr_label]["value"] == pytest.approx(result["summary"]["irr_equity"])


def test_a_direct_sale_returns_money_along_the_way_and_pays_back() -> None:
    result = _nonres(offices_strategy="direct")
    parts = result["report"]["equity_participation"]["parts"]
    assert parts["in_exit"] == 0 and parts["out_exit"] == 0 and parts["out_running"] > 0
    rows = _rows(result)
    assert rows["Окупаемость собственных средств"]["value"] != "не окупается за горизонт проекта"
    assert not any("выход" in label for label in rows)


def test_a_rent_project_names_the_operating_deficit_the_owner_covers() -> None:
    rows = _rows(_nonres())
    deficit = next(r for label, r in rows.items() if label.startswith("в т.ч. после ввода — дефицит"))
    assert deficit["value"] > 0


def test_hold_keeps_the_valuation_out_of_the_cash() -> None:
    result = _nonres(offices_exit_mode="hold")
    rows = _rows(result)
    eq = result["report"]["equity_participation"]
    assert rows["Оценка удержанного объекта — без сделки, не деньги"]["value"] == pytest.approx(eq["residual"])
    assert "Получено всего — деньгами, без оценки" in rows
    assert "IRR собственного капитала — только деньги" in rows


def test_a_project_with_pf_has_no_such_table() -> None:
    x, t = _spec()
    mixed = core.calculate(core.CalcRequest(inputs=x, tep=t, rates=[]))
    assert mixed["report"]["equity_participation"] == {}


def test_the_pdf_prints_it_in_the_financing_section(tmp_path) -> None:
    pypdf = pytest.importorskip("pypdf")
    from test_nonres_debt_metric import _nonres_inputs
    x, t = _nonres_inputs()
    result = core.calculate(core.CalcRequest(inputs=x, tep=t, rates=[]))
    content = core._build_developaid_pdf({"project_name": "Офис", "result": result,
                                          "inputs": x, "tep": t, "rates": []})
    path = tmp_path / "e.pdf"
    path.write_bytes(content)
    flat = " ".join(" ".join(p.extract_text() or "" for p in pypdf.PdfReader(str(path)).pages).split())
    loan = flat.index("Финансирование объекта: Офисы")
    equity = flat.index("Собственное участие")
    assert loan < equity
    for label in ("ВЛОЖЕНИЯ СОБСТВЕННИКА", "ВОЗВРАТ СОБСТВЕННИКУ", "РЕЗУЛЬТАТ СОБСТВЕННИКА",
                  "Собственный капитал по годам"):
        assert label in flat[equity:]
    value = _rows(result)["Вложено всего"]
    assert core._pdf_nonres_value(value) in flat[equity:]


def test_the_nonresidential_book_counts_it_with_formulas() -> None:
    """«Собственное участие» книги — формулы раздела «Отчёта» от потока
    капитала листа «Денежный поток», и «Сверка» сравнивает их с таблицей
    движка (`equity_participation_report`)."""
    import openpyxl

    import nonres_workbook as nw
    from xlsx_eval import Evaluator
    x, t = _nonres_inputs()
    content, _, meta = core.build_project_workbook(x, t, [], {}, project_name="Офис")
    assert meta.get("nonres_book") is True
    book = openpyxl.load_workbook(io.BytesIO(content))
    evaluator = Evaluator(book)
    sys.setrecursionlimit(max(sys.getrecursionlimit(), 100000))
    for name in book.sheetnames:
        for row in book[name].iter_rows(min_row=nw.FIRST_ROW):
            for cell in row:
                if isinstance(cell.value, str) and cell.value.startswith("="):
                    evaluator.cell(name, cell.coordinate)
    checks = book[nw.CHECK_SHEET]
    equity = {checks[f"B{r}"].value: evaluator.cell(nw.CHECK_SHEET, f"G{r}")
              for r in range(3, checks.max_row + 1)
              if checks[f"A{r}"].value == "Собственное участие"}
    assert {"Вложено всего", "Получено всего", "Вложено до ввода", "Пик потребности",
            "NPV собственного капитала"} <= set(equity)
    assert set(equity.values()) <= {"сходится", "—"}
    report = book[nw.REPORT_SHEET]
    labels = [report[f"A{r}"].value for r in range(1, report.max_row + 1)]
    assert "Собственное участие" in labels and "Вложено всего" in labels


# --- страница ----------------------------------------------------------------

def _render(report: dict) -> dict:
    prelude = """
const BOXES=[{innerHTML:''},{innerHTML:''}], CARDS=[{hidden:true},{hidden:true}];
const document={querySelectorAll:sel=>sel==='.equity-box'?BOXES:CARDS};
const R=%s;
""" % json.dumps({"report": report}, ensure_ascii=False)
    tail = """
renderEquityParticipation(R);
const html=BOXES[0].innerHTML;
const main=html.split('</table>')[0];
const rows=[...main.matchAll(/<tr( class="section")?><td>([\\s\\S]*?)<\\/td><td>([\\s\\S]*?)<\\/td><\\/tr>/g)]
 .map(r=>({section:!!r[1],label:r[2],cell:r[3]}));
const eq=R.report.equity_participation||{};
const expected=(eq.rows||[]).map(r=>nonresCell(r));
const head=[...html.matchAll(/<th>([^<]*)<\\/th>/g)].map(x=>x[1]);
console.log(JSON.stringify({cards:CARDS.map(c=>c.hidden),same:BOXES[0].innerHTML===BOXES[1].innerHTML,
 rows,expected,head}));
"""
    out, _ = page_blocks.run(prelude, tail)
    return json.loads(out)


def _mismatches(drawn: dict, report: dict) -> list[str]:
    eq = report.get("equity_participation") or {}
    out = []
    if [r["label"] for r in drawn["rows"]] != [r["label"] for r in eq.get("rows") or []]:
        out.append("строки не те")
    if [r["cell"] for r in drawn["rows"]] != drawn["expected"]:
        out.append("значения не те")
    if [r["label"] for r in drawn["rows"] if r["section"]] != \
            [r["label"] for r in eq.get("rows") or [] if r["unit"] == "section"]:
        out.append("блоки не выделены")
    return out


def test_the_page_draws_the_equity_table() -> None:
    report = _nonres()["report"]
    drawn = _render(report)
    assert drawn["cards"] == [False, False] and drawn["same"] is True
    assert _mismatches(drawn, report) == []
    assert drawn["head"] == [c[1] for c in report["equity_participation"]["columns"]]


def test_the_page_check_fails_on_a_forgery() -> None:
    report = _nonres()["report"]
    tampered = copy.deepcopy(_render(report))
    tampered["rows"][1]["cell"] = "0 млрд ₽"
    tampered["rows"][0]["section"] = False
    assert len(_mismatches(tampered, report)) == 2


def test_a_project_with_pf_hides_the_card() -> None:
    x, t = _spec()
    mixed = core.calculate(core.CalcRequest(inputs=x, tep=t, rates=[]))
    drawn = _render(mixed["report"])
    assert drawn["cards"] == [True, True] and drawn["rows"] == []
