"""Сводка руководству — документ для решений: «Главное», графики, динамика.

«А где графики? Аналитика?» (владелец, 20.10.2026, о PDF из четырёх таблиц).
Первая страница — выводы простым языком, каждый с числом и источником,
посчитанные из тех же блоков; дальше блоки с графиками (встроенный SVG — одна
разметка для страницы и PDF) и динамикой к предыдущему периоду.

Запуск: python3 -m pytest tests/test_management_report_analytics.py -q
"""
from __future__ import annotations

import datetime
import re
import sys
from pathlib import Path

import pytest

ROOT = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(ROOT))
sys.path.insert(0, str(ROOT / "tests"))

import developaid_management_report as report  # noqa: E402
import developaid_monitor as monitor  # noqa: E402
import developaid_monthly_budget as monthly_budget  # noqa: E402
import developaid_project_notes as notes  # noqa: E402
import developaid_report_charts as charts  # noqa: E402
from test_management_report_funding import _rss_with_sources, _sales as _credit_sales  # noqa: E402
from test_management_report_model import (  # noqa: E402
    M, PROJECT, _contract, _gpr_book, _project, budget_book)


@pytest.fixture(autouse=True)
def _isolated(tmp_path, monkeypatch):
    monkeypatch.setattr(monitor, "_SNAPSHOT_DIR", tmp_path / "monitor")
    monkeypatch.setenv("DATA_DIR", str(tmp_path / "data"))


def _full() -> dict:
    _project()
    monitor.store_estimate(PROJECT, _rss_with_sources(), "2026-10-18")
    monitor.store_schedule_fact(PROJECT, _gpr_book({"1": 0.30, "2": 0.0}, cash=False),
                                "2026-10-02")
    monthly_budget.store_budget(PROJECT, budget_book(), "2026-10-02")
    _credit_sales()
    notes.add(PROJECT, "Арбитраж с подрядчиком", "Иск подан", "Иванов",
              at="2026-10-14T09:00:00+00:00")
    return report.build(PROJECT, "week", "2026-10-18", sales_project="ЖК")


def test_the_lag_and_its_change_over_the_period():
    built = _full()
    block = built["physical"]
    assert block["lag"] == pytest.approx(block["fact"] - block["plan"])
    assert block["lag_before"] == pytest.approx(block["fact_before"] - block["plan_before"])
    # Три ГПР-факта — три точки S-кривой, каждая со своим планом на дату.
    assert [p["day"] for p in block["series"]["fact"]] == [
        "2026-10-02", "2026-10-09", "2026-10-16"]
    first = block["series"]["fact"][0]
    assert first["plan"] < block["plan"]


def test_findings_carry_numbers_and_sources():
    built = _full()
    found = {item["topic"]: item for item in built["findings"]}
    assert list(found) == ["Стройка", "Люди на площадке", "Отстающие работы", "Срок",
                           "Корпуса", "Бюджет месяца", "Продажи", "Финансирование", "Споры",
                           "Открытые вопросы"]
    assert "прогноз Монитора по темпу работ" == found["Срок"]["source"]
    lag = built["physical"]["lag"]
    assert f"отставание {report._pp(-lag)[1:]}" in found["Стройка"]["text"]
    assert "ГПР-факт от 16.10.2026" == found["Стройка"]["source"]
    # Оплата 19.10 уже в реестре, но после конца недели — в сводку не входит.
    assert "оплачено 150,0 млн ₽ из 400,0 млн ₽ (37,5%)" in found["Бюджет месяца"]["text"]
    assert "Выборка кредита за неделю — 48,0 млн ₽ (за предыдущий такой же период — " \
           "120,0 млн ₽)" in found["Финансирование"]["text"]
    assert "«Арбитраж с подрядчиком»" in found["Открытые вопросы"]["text"]
    # Нет данных — вывод говорит об этом, а не молчит.
    empty = report.build("Пусто", "week", "2026-10-18")
    assert {item["tone"] for item in empty["findings"]} == {"missing"}


def test_the_document_opens_with_findings_and_has_charts():
    built = _full()
    page = report.render_html(built)
    assert page.index("Главное за период") < page.index(">Стройка</h2>")
    titles = re.findall(r"<figcaption>([^<]+)</figcaption>", page)
    assert titles == [
        "Выполнение по графику работ: план и факт на дату каждого ГПР-факта",
        "Продажи нарастающим итогом по кварталам, сумма договоров",
        "Темп продаж по месяцам: м² договоров и цена квартир",
        "Темп продаж по месяцам: лотов всех типов (квартиры, помещения, машиноместа, кладовые) "
        "и цена квартир",
        "Темп продаж по месяцам: сумма договоров, млн ₽ и цена квартир",
        "Выборка кредита и эскроу нарастающим итогом по месяцам",
        "Покрытие кредита эскроу: факт и план по месяцам",
        "Бюджет на октябрь 2026: план и оплачено по статьям",
        "Стройка и проектирование: план и оплачено по месяцам",
    ]
    # Числа графика — те же, что в таблице блока: столбец «Оплачено» стройки.
    budget = page[page.index("Бюджет на октябрь 2026"):]
    assert ">150,0<" in budget and ">400,0<" in budget
    text = report.render_text(built)
    assert text.index("<b>Главное</b>") < text.index("<b>Стройка</b>")
    assert "п. п.." not in text and "п. п.." not in page


def test_charts_survive_missing_data():
    built = report.build("Пусто", "week", "2026-10-18")
    assert "<svg" not in report.render_html(built)
    assert charts.line_chart("t", [], [], "%") == ""
    assert charts.bar_chart("t", ["a"], [{"name": "x", "values": [None]}], "%") == ""


def test_axis_ticks_cover_the_data():
    for high in (0.73, 412.5, 17, 1, 2_999):
        ticks = charts._ticks(0, high)
        assert ticks[0] <= 0 and ticks[-1] >= high, (high, ticks)


REPORT = """Добрый день.
1. СПМ : Итр- 3чел., Рабочие - 6 чел.
2. НУР:ИТР- 11чел.;
Рабочие - 54чел.;
3. Клодо ( кладка) - Итр- 2 чел.Рабочие - 12чел.
По работам:
1. НУР: Корпус 1 монолит
"""


def test_people_on_site_from_the_bots_reports():
    import developaid_monitor_daily as daily

    _full()
    for day in ("2026-10-06", "2026-10-07", "2026-10-13", "2026-10-15"):
        daily.store_daily_report(PROJECT, REPORT, day)
    daily.store_daily_report(PROJECT, REPORT.replace("54чел", "24чел"), "2026-10-16")
    got = report.crew(PROJECT, "week", "2026-10-18")
    # Три отчёта за неделю: 72, 72, 42 рабочих; дни без отчёта не нули.
    assert got["now"]["days"] == 3
    assert got["now"]["workers_avg"] == pytest.approx((72 + 72 + 42) / 3)
    assert got["now"]["workers_max"] == 72
    assert got["previous"]["days"] == 2
    assert got["previous"]["workers_avg"] == pytest.approx(72)
    nur = next(row for row in got["by_contractor"] if "НУР" in row["name"])
    assert nur["workers_avg"] == pytest.approx((54 + 54 + 24) / 3)
    points = {p["day"]: p["workers"] for p in got["series"]}
    assert points["2026-10-14"] is None and points["2026-10-16"] == 42
    built = report.build(PROJECT, "week", "2026-10-18", sales_project="ЖК")
    found = {item["topic"]: item for item in built["findings"]}
    assert "В среднем 62 рабочих" in found["Люди на площадке"]["text"]
    assert "в предыдущий такой же период — 72 (−14%)" in found["Люди на площадке"]["text"]
    page = report.render_html(built)
    assert "Рабочих на площадке по дням" in page
    assert "Плановой численности в графике работ нет" in page


def test_no_reports_is_a_reason_not_zero():
    _full()
    got = report.crew(PROJECT, "week", "2026-10-18")
    assert got["state"] == "missing"
    assert "отчёты со стройки не приходили" in got["reason"].lower()


def test_breakdown_by_works():
    built = _full()
    breakdown = built["physical"]["breakdown"]
    # Обе работы отстают: «Фасад» (40% при плане ~54%) сильнее «Монолита».
    names = [row["name"] for row in breakdown["laggards"]]
    assert names[0] == "Фасад"
    assert breakdown["laggards"][0]["lag"] < breakdown["laggards"][1]["lag"]
    found = {item["topic"]: item for item in built["findings"]}
    assert found["Отстающие работы"]["text"].startswith("Больше всего отстают: Фасад (Корпус 2)")


def test_a_work_without_growth_is_flagged():
    _project()
    # Неделя назад и сейчас у «Монолита» 55% — прироста нет при идущем плане.
    monitor.store_schedule_fact(PROJECT, _gpr_book({"1": 0.55, "2": 0.30}, cash=False),
                                "2026-10-23")
    built = report.build(PROJECT, "week", "2026-10-25")
    stalled = built["physical"]["breakdown"]["stalled"]
    assert [row["name"] for row in stalled] == ["Монолит"]


def test_budget_sales_and_funding_depth():
    built = _full()
    budget = built["budget"]
    # Тренд: месяцы раньше первой оплаты и без плана не показаны (реестр
    # начинается в сентябре, бюджет — тоже), «нет данных» не рисуется нулём.
    assert [p["month"] for p in budget["trend"]] == ["2026-09-01", "2026-10-01"]
    assert budget["trend"][-1]["plan"]["construction"] == pytest.approx(400e6)
    assert budget["trend"][-1]["fact"]["construction"] == pytest.approx(150e6)
    # Прогноз до конца месяца по темпу: 150 млн за 18 из 31 дня.
    forecast = report._budget_forecast(budget, datetime.date(2026, 10, 18))
    assert forecast[0]["forecast"] == pytest.approx(150e6 / (18 / 31))
    assert budget["recipients"][0]["name"] == "СтройКо"
    # В тестовой книге из статей — только служба заказчика; остальной план
    # главы стройки — строкой «не разнесены по карте», план главы не теряется.
    build = budget["by_budget_article"]["chapters"][0]
    assert build["total"]["plan"] == pytest.approx(400e6)
    assert build["rows"][-1]["key"] == "rest"
    sales = built["sales"]
    flats = next(item for item in sales["products"] if item["product"] == "Квартира")
    assert flats["total"]["area"] == pytest.approx(200.0)
    assert [p["month"] for p in sales["pace"]][-1] <= "2026-10"
    escrow = built["funding"]["escrow"]
    assert escrow["inflow_month"] == pytest.approx(5e6)
    assert escrow["inflow_total"] == pytest.approx(65e6)
    found = {item["topic"]: item for item in built["findings"]}
    assert "по нынешнему темпу к концу месяца стройке оплатят около" in found["Бюджет месяца"]["text"]
    assert "на эскроу за октябрь 2026 поступило 5,0 млн ₽" in found["Финансирование"]["text"]


def test_closed_month_deviations_by_article_name_without_outline():
    from openpyxl import Workbook
    import io

    book = Workbook()
    sheet = book.active
    sheet.cell(row=4, column=5, value="Сентябрь 2026")
    sheet.cell(row=5, column=5, value="ПЛАН")
    sheet.cell(row=5, column=6, value="ФАКТ")
    rows = [("Стоимость строительства и проектирования", -300e6, -200e6),
            ("Основные объекты строительства", -280e6, -190e6),
            ("Подрядчик", -280e6, -190e6),
            ("Наружные сети и сооружения", -20e6, None),
            ("Содержание службы заказчика", -6e6, -6e6),
            ("Коммерческие расходы", -10e6, -9e6),
            ("Реклама и PR (КС) - Группа", -10e6, -9e6),
            ("Административно-хозяйственные расходы", -5e6, -5e6)]
    # Группировки строк нет — статьи всё равно находятся по имени.
    for r, (name, plan, fact) in enumerate(rows, 6):
        sheet.cell(row=r, column=1, value=name)
        sheet.cell(row=r, column=5, value=plan)
        if fact is not None:
            sheet.cell(row=r, column=6, value=fact)
    blob = io.BytesIO()
    book.save(blob)
    monthly_budget.store_budget(PROJECT, blob.getvalue(), "2026-10-02")
    got = monthly_budget._articles(PROJECT, datetime.date(2026, 10, 18))
    assert got["known"] and got["month"] == "2026-09-01"
    by_name = {row["article"]: row for row in got["rows"]}
    # Подрядчик и АХР — не статьи карты.
    assert set(by_name) == {"Основные объекты строительства", "Наружные сети и сооружения",
                            "Служба заказчика, охрана и офис", "Реклама и PR"}
    assert by_name["Основные объекты строительства"]["delta"] == pytest.approx(-90e6)
    assert by_name["Наружные сети и сооружения"]["fact_blank"] is True
    assert got["rows"][0]["article"] == "Основные объекты строительства"

def test_works_of_the_week_from_the_bots_reports():
    """«А выполняемые за неделю работы?» (владелец, 20.10.2026) — из отчётов бота."""
    import developaid_monitor_daily as daily

    _full()
    text = REPORT + "2. СПМ: Корпус 2 кладка 3 эт.\n"
    daily.store_daily_report(PROJECT, text, "2026-10-13")
    daily.store_daily_report(PROJECT, text.replace("кладка 3 эт.", "кладка 4 эт."), "2026-10-15")
    # Отчёт до начала недели — не её работа.
    daily.store_daily_report(PROJECT, REPORT.replace("монолит", "котлован"), "2026-10-09")
    built = report.build(PROJECT, "week", "2026-10-18")
    works = {group["contractor"]: group for group in built["crew"]["works"]}
    assert set(works) == {"НУР", "СПМ"}
    # Одна строка за два дня — одна строка с двумя датами, без имени подрядчика.
    assert works["НУР"]["lines"] == [{"line": "Корпус 1 монолит",
                                      "days": ["2026-10-13", "2026-10-15"]}]
    assert [e["line"] for e in works["СПМ"]["lines"]] == ["Корпус 2 кладка 3 эт",
                                                           "Корпус 2 кладка 4 эт"]
    page = report.render_html(built)
    part = page[page.index("Выполненные работы за неделю"):]
    assert "Корпус 1 монолит" in part and "13.10, 15.10" in part and "котлован" not in part
    assert "Работы НУР: Корпус 1 монолит." in report.render_text(built)


def test_reports_without_works_say_so_and_do_not_go_blank():
    import developaid_monitor_daily as daily

    _full()
    people_only = REPORT[:REPORT.index("По работам")]
    daily.store_daily_report(PROJECT, people_only, "2026-10-14")
    built = report.build(PROJECT, "week", "2026-10-18")
    assert built["crew"]["works"] == []
    page = report.render_html(built)
    assert ("Работы не перечислены: в отчётах за период нет раздела «По работам»"
            in page)


def test_deviations_are_for_last_month_not_the_last_month_with_a_fact():
    # «В блоке Бюджет месяца неверный план апреля, не могу понять, откуда он
    # подтянут» (финансист, 21.10.2026): в октябрьской сводке стояла таблица за
    # апрель — последний месяц файла с заполненным фактом.
    from openpyxl import Workbook
    import io

    book = Workbook()
    sheet = book.active
    for column, month in ((5, "Апрель 2026"), (7, "Сентябрь 2026")):
        sheet.cell(row=4, column=column, value=month)
        sheet.cell(row=5, column=column, value="ПЛАН")
        sheet.cell(row=5, column=column + 1, value="ФАКТ")
    rows = [("Стоимость строительства и проектирования", -300e6, -200e6),
            ("Основные объекты строительства", -280e6, -190e6),
            ("Содержание службы заказчика", -6e6, -6e6),
            ("Коммерческие расходы", -10e6, -9e6),
            ("Расходы на продажу (КС) - Группа", -5e6, -4e6)]
    for r, (name, plan, fact) in enumerate(rows, 6):
        sheet.cell(row=r, column=1, value=name)
        sheet.cell(row=r, column=5, value=plan)
        sheet.cell(row=r, column=6, value=fact)
        sheet.cell(row=r, column=7, value=plan * 4)  # сентябрь: план есть, факта нет
    blob = io.BytesIO()
    book.save(blob)
    monthly_budget.store_budget(PROJECT, blob.getvalue(), "2026-10-02")
    got = monthly_budget._articles(PROJECT, datetime.date(2026, 10, 18))
    assert got == {"known": False,
                   "reason": "в бюджете от 2026-10-02 не заполнен факт за 09.2026"}
    got = monthly_budget._articles(PROJECT, datetime.date(2026, 12, 5))
    assert got == {"known": False, "reason": "в бюджете от 2026-10-02 нет месяца 11.2026"}


def test_sales_area_is_of_goods_sold_by_the_metre():
    # «Сумма учитывает машиноместа, а площадь машиномест не учитывается»
    # (финансист, 21.10.2026). Машиноместа и кладовые — штуками рядом.
    rows = [_contract("2026-10-10", "Квартира", 50.0, 20 * M),
            _contract("2026-10-11", "Машиноместо", 14.0, 3 * M),
            _contract("2026-10-12", "Машиноместо", 0.0, 3 * M)]
    fact = report._fact(rows)
    assert fact["area"] == pytest.approx(50.0) and fact["amount"] == pytest.approx(26 * M)
    assert fact["pieces"] == {"Машиноместо": 2.0} and fact["area_unknown"] == 0
    assert report._sales_line("За неделю", fact) == (
        "За неделю: 3 договора на 26,0 млн ₽: квартиры и помещения — 50 м², 2 машиноместа, "
        "квартиры в среднем по 400 000 ₽ за м².")


def test_the_sales_pace_is_one_subject_in_three_measures():
    """«Может совместить для экономии места? и отдельно лотами и млн рублей
    сделать такие же графики?» (владелец, 22.10.2026): цена — линией на каждом
    графике темпа, меры — панели одного переключателя; первая открыта."""
    built = _full()
    page = report.render_html(built)
    block = page[page.index('<div class="switch-charts" data-switch="sales-pace">'):]
    block = block[:block.index("<h3>") if "<h3>" in block else len(block)]
    panes = re.findall(r'<div class="pane( on)?" data-pane="(\w+)">', block)
    assert panes == [(" on", "area"), ("", "units"), ("", "amount")]
    assert re.findall(r'data-pane="(\w+)" class="(on)?">([^<]+)</button>', block) == [
        ("area", "on", "м²"), ("units", "", "лоты"), ("amount", "", "млн ₽")]
    # На каждом графике — линия цены на правой шкале и её значение у последней точки.
    assert block.count('class="line"') == 3 and block.count(">тыс. ₽/м²<") == 3
    pace = built["sales"]["pace"]
    last = next(p for p in reversed(pace) if p.get("price_flats"))
    assert f">{charts._num(last['price_flats'] / 1e3)}<" in block
    # Лоты и миллионы — те же числа, что в данных блока, не второй счёт.
    assert f">{charts._num(sum(p['units'] for p in pace))}<" not in block  # не итог, а по месяцам
    assert f">{charts._num(pace[-1]['units'])}<" in block
    assert f">{charts._num(pace[-1]['amount'] / 1e6, 1)}<" in block


def test_a_bar_chart_carries_a_price_line_on_its_own_scale():
    svg = charts.bar_chart("t", ["01.26", "02.26"], [{"name": "м²", "values": [100.0, 200.0]}],
                           "м²", line={"name": "Цена", "values": [650.0, 817.0],
                                       "unit": "тыс. ₽/м²", "digits": 0})
    assert 'class="line"' in svg and ">817<" in svg and ">тыс. ₽/м²<" in svg
    # Левая шкала — от нуля; правая — от своего минимума, а не от нуля.
    assert ">0<" in svg and ">650<" in svg and ">850<" in svg
    # Без цены — обычные столбики, без правой шкалы.
    plain = charts.bar_chart("t", ["01.26"], [{"name": "м²", "values": [100.0]}], "м²",
                             line={"name": "Цена", "values": [None], "unit": "x"})
    assert 'class="line"' not in plain
