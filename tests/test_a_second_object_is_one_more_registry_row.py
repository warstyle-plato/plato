"""Вторые ОСЗ — строка реестра, а не пятый поимённый блок.

Второй офисник, второй ТЦ и второй наземный паркинг нужны для случая «два
объекта в ОДНОЙ очереди» (решение владельца 13.09.2026: инфляция объекта — от
старта его очереди). Книга обобщена первым шагом (`_v4_object_layouts`), и
объект встаёт в неё копией блока двойника. Этот файл держит второй шаг: всё,
что перечисляло объекты поимённо, читает реестр.

Главная поломка, ради которой шаг сделан: размещение объектов по очередям
перечисляло объекты шаблона по именам, и объект вне списка оставался
включённым в КАЖДОЙ очереди — его выручка считалась столько раз, сколько
очередей. На пробе с тремя очередями книга расходилась с движком ровно на
утроенную выручку новых объектов.

Запуск: python3 -m pytest tests/test_a_second_object_is_one_more_registry_row.py -q
"""

from __future__ import annotations

import copy
import io
import sys
from pathlib import Path

import openpyxl
import pytest

ROOT = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(ROOT))
sys.path.insert(0, str(ROOT / "tests"))

import main as wrapper  # noqa: E402
import presentation  # noqa: E402
import v4_dashboard  # noqa: E402
from xlsx_eval import Evaluator  # noqa: E402

core = wrapper.core

SECONDS = {"offices2": "offices", "standalone_retail2": "standalone_retail",
           "above_parking2": "above_parking"}
PORT = 18971


def _inputs() -> dict:
    x = dict(core.DEFAULT_INPUTS)
    x.update(apartment_price_th=650, commercial_price_th=650, parking_price_th=5000,
             offices_enabled=True, retail_enabled=True, above_parking_enabled=True,
             offices2_enabled=True, offices2_gba_sqm=8000, offices2_saleable_sqm=5000,
             offices2_parking_under_spaces=40, offices2_parking_over_spaces=10,
             retail2_enabled=True, retail2_gba_sqm=7000, retail2_saleable_sqm=4000,
             retail2_parking_under_spaces=20,
             above_parking2_enabled=True, above_parking2_spaces=300)
    return x


def _phasing() -> dict:
    # Второй офисник — в ОДНОЙ очереди с первым (ради этого он и заведён),
    # второй ТЦ — в первой, второй паркинг — в третьей.
    return {"enabled": True, "phase_count": 3, "phase_gap_months": 12,
            "cost_inflation_pct": 8,
            "phases": [{"name": f"О{i + 1}", "start_offset_months": 12 * i,
                        "construction_months": 24} for i in range(3)],
            "discrete": {"offices": 3, "offices2": 3, "standalone_retail": 2,
                         "standalone_retail2": 1, "above_parking": 2,
                         "above_parking2": 3}}


# --- реестр ------------------------------------------------------------------

def test_each_second_object_is_a_copy_of_its_twin() -> None:
    by_key = {o.key: o for o in core.STANDALONE_OBJECTS}
    for key, twin in SECONDS.items():
        obj, first = by_key[key], by_key[twin]
        assert obj.book_twin == twin and obj.product == twin, key
        assert (obj.measure, obj.garage, obj.garage_sellable, obj.default_queue) == (
            first.measure, first.garage, first.garage_sellable, first.default_queue), key
        assert obj.prefix != first.prefix and not obj.prefix.startswith(first.prefix + "_")
        # Умолчания — те же числа, что у первого, и объект выключен: проект без
        # вторых объектов считается как прежде.
        own = core.standalone_object_defaults(obj)
        twin_defaults = core.standalone_object_defaults(first)
        for field, value in twin_defaults.items():
            assert own[obj.prefix + field[len(first.prefix):]] == value, (key, field)
        assert core.DEFAULT_INPUTS[obj.enabled_key] is False
    # ФОК остаётся первым дописанным объектом: его строки в книге постоянны.
    # Дальше — все экземпляры реестра (`object_instance`), вторые включительно.
    extras = [lay.obj.key for lay in core._V4_EXTRA_OBJECT_LAYOUTS]
    assert extras[0] == "sports" and set(extras[1:]) == set(core.OBJECT_INSTANCES)
    assert set(SECONDS) <= set(core.OBJECT_INSTANCES)


def test_the_literals_place_the_second_object_after_its_twin() -> None:
    rows = list(core.TEP_DEFAULT)
    groups = [group[0] for group in core.FIELD_GROUPS]
    by_key = {o.key: o for o in core.STANDALONE_OBJECTS}
    for key, twin in SECONDS.items():
        assert rows.index(key) == rows.index(twin) + 1, rows
        assert core.TEP_DEFAULT[key]["label"] == by_key[key].tep_label
        assert groups.index(by_key[key].group_label) == groups.index(by_key[twin].group_label) + 1
    assert not any(str(g).startswith(core._OBJECT_PLACEHOLDER) for g in groups)


def test_a_second_object_follows_its_product() -> None:
    """Доли площадей, цена класса и норматив паркинга — те же, что у продукта."""
    for key, twin in SECONDS.items():
        if twin in core.TEP_RATIOS:
            assert core.TEP_RATIOS[key]["saleable_of_gns"] == core.TEP_RATIOS[twin]["saleable_of_gns"]
    for preset in core.PROJECT_CLASS_PRESETS.values():
        assert preset["offices2_price_th_per_sqm"] == preset["offices_price_th_per_sqm"]
        assert preset["retail2_price_th_per_sqm"] == preset["retail_price_th_per_sqm"]
    # Мера счёта — продукта: второй паркинг считается машино-местами, второй
    # офисник штук не имеет вовсе, а не уходит в «шт.» последнего прибежища.
    for key, twin in SECONDS.items():
        assert key in core.TEP_COUNT_MEASURE, key
        assert core.tep_count_measure(key) == core.tep_count_measure(twin), key
    assert core.tep_count_measure("above_parking2") == core.COUNT_PARKING
    functions = {row[0]: row[1:] for row in core._PARKING_DEMAND_PRODUCTS}
    assert functions["offices2"] == functions["offices"]
    assert functions["standalone_retail2"] == functions["standalone_retail"]


# --- движок: очереди ----------------------------------------------------------

@pytest.fixture(scope="module")
def phased() -> dict:
    return core._run_authoritative_model(
        _inputs(), copy.deepcopy(core.TEP_DEFAULT), [], _phasing())


def test_each_second_object_is_built_in_one_queue_only(phased) -> None:
    """Объект вне поимённого списка был включён в каждой очереди."""
    by_key = {o.key: o for o in core.STANDALONE_OBJECTS}
    for key, queue in _phasing()["discrete"].items():
        switch = by_key[key].enabled_key
        on = [index + 1 for index, item in enumerate(phased["phases"])
              if item["inputs"].get(switch)]
        assert on == [queue], (key, on)


def test_the_second_object_revenue_is_counted_once(phased) -> None:
    single = core._run_authoritative_model(
        _inputs(), copy.deepcopy(core.TEP_DEFAULT), [], {})
    for key in SECONDS:
        by_phase = [
            float(next((p.get("revenue") or 0.0) for p in item["result"]["report"]["products"]
                       if p["key"] == key) or 0.0)
            for item in phased["phases"]]
        assert sum(1 for value in by_phase if value > 0) == 1, (key, by_phase)
        # Очередь дорожает на инфляцию своего сдвига — выручка той же величины,
        # а не втрое больше одиночного расчёта.
        alone = float(next(p.get("revenue") or 0.0 for p in single["consolidated"]["report"]["products"]
                           if p["key"] == key))
        assert alone > 0
        assert max(by_phase) < alone * 1.5, (key, by_phase, alone)


def test_a_garage_without_metres_in_the_tep_row_stays_in_its_queue() -> None:
    """Строка ТЭП пустая (проект пришёл не со страницы) — гараж всё равно строится.

    Доля объекта в очереди бралась у метров строки, и при пустой строке она
    выходила нулём во всех очередях: движок терял гараж, а книга его строила.
    Так было и у первого офисника.
    """
    x = dict(core.DEFAULT_INPUTS)
    x.update(offices_enabled=True, offices_parking_under_spaces=40)
    got = core._run_authoritative_model(
        x, copy.deepcopy(core.TEP_DEFAULT), [],
        {**_phasing(), "discrete": {"offices": 3}})
    spaces = [item["inputs"].get("offices_parking_under_spaces") for item in got["phases"]]
    assert spaces == [0.0, 0.0, 40.0], spaces


# --- книга: паритет ------------------------------------------------------------

def _book(phasing: dict):
    sys.setrecursionlimit(400000)
    content, _, meta = core.build_project_workbook(
        _inputs(), copy.deepcopy(core.TEP_DEFAULT), [], phasing, project_name="Вторые ОСЗ")
    book = openpyxl.load_workbook(io.BytesIO(content), data_only=False)
    return book, Evaluator(book), meta


@pytest.fixture(scope="module")
def flat_book():
    return _book({})


@pytest.fixture(scope="module")
def phased_book():
    return _book(_phasing())


@pytest.mark.parametrize("which", ["flat_book", "phased_book"])
def test_the_book_agrees_with_the_engine(which, request) -> None:
    book, evaluator, meta = request.getfixturevalue(which)
    assert meta["missing"] == [], meta["missing"]
    for row in range(76, 85):
        assert book["ПРОВЕРКИ"][f"C{row}"].value is not None, f"цель строки {row}"
        assert evaluator.cell("ПРОВЕРКИ", f"F{row}") == "OK", (
            which, str(book["ПРОВЕРКИ"][f"A{row}"].value),
            evaluator.cell("ПРОВЕРКИ", f"B{row}"), evaluator.cell("ПРОВЕРКИ", f"C{row}"))
    assert evaluator.cell("ПРОВЕРКИ", "B3") in ("ПРОЙДЕНО", "ПРОЙДЕНО С ПРЕДУПРЕЖДЕНИЯМИ")


def test_the_dashboard_names_each_second_object(flat_book) -> None:
    """Дашборд книги знает дописанные объекты своими строками — и ФОК тоже."""
    book, evaluator, _ = flat_book
    # Строки дашборда — объектов ЭТОЙ книги, а не всего пула экземпляров.
    token = core._V4_BOOK_OBJECTS.set(core._v4_book_objects(
        core.object_instances_applied(_inputs(), {})[0]))
    try:
        extras = core._v4_dashboard_extra_products()
    finally:
        core._V4_BOOK_OBJECTS.reset(token)
    rows = v4_dashboard.product_rows(extras)
    assert [item[0] for item in extras][0] == "sports"
    engine = core._run_authoritative_model(
        _inputs(), copy.deepcopy(core.TEP_DEFAULT), [], {})["consolidated"]
    numbers = core.presentation_numbers(engine)["products"]
    for key in SECONDS:
        row = rows[key]
        assert book[v4_dashboard.DATA_SHEET][f"A{row}"].value == key
        book_revenue = float(evaluator.cell(v4_dashboard.DATA_SHEET, f"F{row}") or 0)
        assert book_revenue > 0, key
        # Строка ТЭП объекта в книге несёт и выручку его гаража — так же, как
        # строка первого офисника (G31 = блок + гараж); у движка гараж — свой
        # продукт объекта (`OBJECT_PARKING_PRODUCT_KEYS`). Места продаются
        # только у офисника, но строку гаража берём у каждого, у кого она есть.
        engine_revenue = numbers[key]["revenue_mln"]
        parking = core.OBJECT_PARKING_PRODUCT_KEYS.get(key)
        if parking in numbers:
            engine_revenue += numbers[parking]["revenue_mln"]
        assert book_revenue == pytest.approx(engine_revenue, rel=5e-4, abs=1.0), key
    # Итог продуктов складывает и дописанные строки.
    total = float(evaluator.cell(v4_dashboard.DATA_SHEET, f"F{v4_dashboard.PRODUCT_TOTAL_ROW}") or 0)
    assert total == pytest.approx(engine["summary"]["revenue"] / 1e6, rel=5e-4)


def test_the_second_garage_is_its_own_product_in_engine_and_book(phased_book, phased) -> None:
    """Паркинг второго офисника — СВОЙ продукт (реестр паркинга объектов), а не
    слагаемое чужой строки: в отчёте движка своей строкой, в КОНСОЛИДАТОРЕ —
    своей колонкой, читающей строку гаража именно этого объекта."""
    key = core.OBJECT_PARKING_PRODUCT_KEYS["offices2"]
    assert key != core.OBJECT_PARKING_PRODUCT_KEYS["offices"]
    products = {p["key"]: p for p in phased["consolidated"]["report"]["products"]}
    assert products[key]["object"] == "offices2"
    assert float(products[key].get("revenue") or 0) > 0
    # Очередь гаража — очередь его объекта (третья), в других его нет.
    by_phase = [float(next((p.get("revenue") or 0.0) for p in item["result"]["report"]["products"]
                           if p["key"] == key) or 0.0) if any(
                    p["key"] == key for p in item["result"]["report"]["products"]) else 0.0
                for item in phased["phases"]]
    assert [value > 0 for value in by_phase] == [False, False, True], by_phase
    book, _evaluator, meta = phased_book
    assert meta["missing"] == [], meta["missing"]
    sheet = book["КОНСОЛИДАТОР"]
    label = core.NON_TEP_PRODUCT_LABELS[key]
    column = next(cell.column_letter for cell in sheet[3]
                  if isinstance(cell.value, str) and label in cell.value)
    formula = str(sheet[f"{column}4"].value)
    own = core._V4_OBJECT_PARKING_BY_KEY["offices2"]
    assert f"$B${own[3]}" in formula, formula
    for other in core._V4_OBJECT_PARKING:
        if other is not own:
            assert f"$B${other[3]}," not in formula, (other[0], formula)


def test_the_second_garage_counts_as_parking_spaces(phased) -> None:
    """Гараж второго офисника — машино-места в разборе штук свода, а спутать
    места с метрами офиса или сложить с квартирами нельзя."""
    total = phased["consolidated"]["tep"]["total"]
    rows = phased["consolidated"]["tep"]["rows"]
    by_measure = total["units_by_measure"]
    garage = sum(float(row.get("parking_units") or 0) for row in rows
                 if row.get("key") == "offices2")
    assert garage > 0
    assert by_measure[core.COUNT_PARKING] == pytest.approx(
        core.tep_parking_spaces(rows)), by_measure
    assert core.COUNT_PIECES not in by_measure or all(
        core.tep_count_measure(row.get("key")) != core.COUNT_PIECES
        for row in rows if str(row.get("key", "")).endswith("2")), by_measure


# --- поверхности -----------------------------------------------------------------

def test_the_teaser_keeps_a_product_that_is_not_in_its_order() -> None:
    """Порядок и состав тизера — у отчёта движка: второй офисник не выпадает."""
    report = {"products": [{"key": key, "label": key} for key in
                           ("apartments", "offices", "offices2", "sports", "above_parking2")]}
    model = presentation.build_project_presentation(
        {"products": {}}, {"report": report}, [], {}, {}, 1.2)
    keys = [p["key"] for p in model["products"]]
    assert sorted(keys) == sorted(p["key"] for p in report["products"]), keys
    # Встаёт за тем, за кем его поставил расчёт.
    assert keys.index("offices2") == keys.index("offices") + 1, keys
    assert keys.index("above_parking2") == keys.index("sports") + 1, keys


def test_the_capex_tables_name_the_second_object() -> None:
    labels = dict(core._MODEL_CAPEX_LABELS)
    for key in SECONDS:
        assert key in labels and key in core._MONTHLY_CAPEX_LABELS
    keys = [key for key, _ in core._MODEL_CAPEX_LABELS]
    assert keys.index("offices2") == keys.index("offices") + 1


# --- страница ----------------------------------------------------------------------

PROBE_PAGE = """() => {
  inputs.offices2_enabled=true; inputs.offices2_gba_sqm=8000; inputs.offices2_saleable_sqm=0;
  inputs.above_parking2_enabled=true; inputs.above_parking2_spaces=300;
  // Поля второго офисника — его вкладкой в блоке «МФОЦ / офисы».
  OBJECT_TAB.offices='offices2';
  syncTep(false); renderInputs(); renderTep();
  // Правка мест гаража отмечает их «заданными руками» у своего объекта.
  const garage=document.getElementById('f_offices2_parking_under_spaces');
  garage.value='30'; garage.onchange();
  return {
    offices2_gns: Number((tep.offices2||{}).gns||0),
    offices2_saleable: Number((tep.offices2||{}).saleable||0),
    offices2_filled: Number(inputs.offices2_saleable_sqm||0),
    parking2_units: Number((tep.above_parking2||{}).units||0),
    derived: TEP_DERIVED_INPUTS.filter(k=>k.startsWith('offices2_')||k.startsWith('above_parking2_')),
    switch2: (TEP_ROW_SWITCH.offices2||[])[0]||'',
    norm_cell: !!document.getElementById('parkNorm_offices2'),
    group_field: !!document.getElementById('f_offices2_gba_sqm'),
    peek: GROUP_PEEK['МФОЦ / офисы 2']||[],
    by_hand: (inputs._parking_by_hand||[]).includes('offices2'),
  };
}"""


@pytest.fixture(scope="module")
def page_probe():
    pytest.importorskip("playwright.sync_api")
    from playwright.sync_api import sync_playwright

    import browser

    path = browser.chromium_or_skip()
    errors: list[str] = []
    with browser.serve(core.app, PORT) as base, sync_playwright() as pw:
        with pw.chromium.launch(executable_path=str(path)) as engine:
            page = engine.new_page(viewport={"width": 1440, "height": 900})
            page.on("pageerror", lambda exc: errors.append(str(exc)))
            page.goto(base, wait_until="domcontentloaded")
            page.wait_for_timeout(1800)
            got = page.evaluate(PROBE_PAGE)
            page.close()
    got["errors"] = errors
    return got


def test_the_page_syncs_the_second_object_row(page_probe) -> None:
    """Строку ТЭП второго объекта ведут его вводные — как у первого."""
    assert page_probe["errors"] == [], page_probe["errors"]
    ratio = core.TEP_RATIOS["offices2"]
    assert page_probe["offices2_gns"] == pytest.approx(8000)
    assert page_probe["offices2_saleable"] == pytest.approx(8000 * ratio["saleable_of_gns"], rel=1e-6)
    # Известна только ГНС — продаваемая вернулась в поле, а не осталась нулём.
    assert page_probe["offices2_filled"] == pytest.approx(page_probe["offices2_saleable"])
    assert page_probe["parking2_units"] == 300


def test_the_page_lists_read_the_registry(page_probe) -> None:
    assert page_probe["errors"] == [], page_probe["errors"]
    assert set(page_probe["derived"]) >= {
        "offices2_enabled", "offices2_gba_sqm", "offices2_saleable_sqm",
        "above_parking2_enabled", "above_parking2_spaces"}
    assert page_probe["switch2"] == "offices2_enabled"
    assert page_probe["group_field"], "поля второго офисника нет на форме"
    assert page_probe["norm_cell"], "под местами второго офисника нет места для нормы"
    assert page_probe["peek"] == ["offices2_enabled", "offices2_gba_sqm"]
    assert page_probe["by_hand"], "места второго офисника не отмечены заданными руками"
