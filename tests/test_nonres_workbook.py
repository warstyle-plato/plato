"""Книга нежилого проекта — финансовая модель формулами, сверенная с движком.

Владелец (05.10.2026): «Excel в данном случае не модель, а убогий свод». Книга
считает всё сама от листа «Вводные»: смету, объекты, кредит, НДС и налог на
прибыль с переносом убытка, поток проекта и собственника, NPV, IRR, отчёт; лист
«Сверка» сравнивает каждый итог с авторитетным расчётом движка.

С 23.10.2026 эту книгу получают нежилой проект без офиса и проект с объектом
по ДДУ; офис без ДДУ получает книгу офиса (`tests/test_office_workbook.py`).
Поэтому сценарии без ДДУ здесь — у ТЦ, а не у офиса.

Формулы вычисляет `xlsx_eval.Evaluator` — без Excel. Проверка обязана падать
на подделке: испорченная формула, статья сметы, вводная и налог краснят сверку.

Запуск: python3 -m pytest tests/test_nonres_workbook.py -q
"""

from __future__ import annotations

import copy
import io
import sys
from pathlib import Path

import pytest

ROOT = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(ROOT))

import main_legacy as core  # noqa: E402
import nonres_workbook as nw  # noqa: E402
from xlsx_eval import Evaluator  # noqa: E402

openpyxl = pytest.importorskip("openpyxl")

# Аренда задана числами: умолчания класса B (22.10.2026) слишком малы, чтобы
# сценарии удержания платили налог и гасили кредит из NOI — а это и проверяется.
OFFICE = dict(offices_enabled=True, offices_gba_sqm=40000.0, offices_saleable_sqm=24000.0,
              offices_parking_under_spaces=200, offices_parking_over_spaces=40,
              offices_parking_guest_pct=10, offices_rent_th_per_sqm_month=4.5,
              offices_occupancy_stable_pct=92.0, offices_leaseup_months=12,
              offices_opex_pct=18.0, offices_exit_cap_pct=11.0)
RETAIL = dict(retail_enabled=True, retail_gba_sqm=15000.0, retail_saleable_sqm=11000.0)


def _project(nonresidential: bool = True, **over):
    x = copy.deepcopy(core.DEFAULT_INPUTS)
    t = copy.deepcopy(core.TEP_DEFAULT)
    x.update(OFFICE)
    x.update(over)
    for obj in core.STANDALONE_OBJECTS:
        gba = float(x.get(f"{obj.prefix}_gba_sqm") or 0)
        if not x.get(f"{obj.prefix}_enabled") or gba <= 0:
            continue
        saleable = float(x.get(f"{obj.prefix}_saleable_sqm") or 0)
        t.setdefault(obj.key, {}).update(gns=gba, total_area=round(gba * 0.94, 2),
                                         useful=saleable, saleable=saleable)
    if nonresidential:
        for key in core.MKD_PRODUCTS:
            for col in ("gns", "total_area", "useful", "saleable", "transfer", "units"):
                if key in t:
                    t[key][col] = 0
        x["project_kind"] = core.PROJECT_KIND_NONRESIDENTIAL
        # Ставка метрового объекта нежилого проекта — СМР; умолчание то же,
        # что ставит страница при выборе вида (`object_smr_rate`).
        for key, turnkey in core.OBJECT_SMR_RATE_DEFAULTS.items():
            if key not in over and float(x.get(key) or 0) == turnkey:
                x[key] = core.object_smr_rate(turnkey, x)
    return x, t


# Сценарии без ДДУ — у ТЦ без офиса: проект с офисом получает книгу офиса.
RETAIL_ONLY = dict(offices_enabled=False, retail_enabled=True, retail_gba_sqm=40000.0,
                   retail_saleable_sqm=24000.0, retail_parking_under_spaces=200,
                   retail_parking_over_spaces=40, retail_rent_th_per_sqm_month=4.5,
                   retail_occupancy_stable_pct=92.0, retail_leaseup_months=12,
                   retail_opex_pct=18.0, retail_exit_cap_pct=11.0, retail_rent_date="2025-12-01")
VARIANTS = {
    "ТЦ ДКП, продажи до ввода": dict(RETAIL_ONLY, retail_strategy="direct",
                                     retail_direct_sale_offset_months=-6,
                                     retail_direct_sale_curve="bell"),
    "ТЦ аренда и продажа, аннуитет": dict(RETAIL_ONLY, retail_strategy="income"),
    "аренда и удержание, погашение из NOI": dict(RETAIL_ONLY, retail_strategy="income",
                                                 retail_exit_mode="hold", retail_debt_repayment="sweep",
                                                 retail_rent_th_per_sqm_month=7.0),
    "баллон в срок, участок рассрочкой, сценарий": dict(
        RETAIL_ONLY, retail_strategy="income", retail_loan_term_years=4, retail_hold_years=8,
        retail_loan_balloon_pct=10, purchase_price_mln=900, purchase_schedule="30%@0; 70%@8",
        land_buyout_mln=120, scenario_cost_multiplier=1.1, scenario_revenue_multiplier=0.9,
        rate_scenario="high"),
}

_BOOKS: dict[str, tuple[bytes, dict]] = {}


def _book(name: str) -> tuple[bytes, dict]:
    if name not in _BOOKS:
        x, t = _project(**VARIANTS[name])
        content, filename, meta = core.build_project_workbook(x, t, [], {}, project_name="Офис")
        assert meta.get("nonres_book") is True and "нежилой" in filename
        assert not meta.get("office_book")
        _BOOKS[name] = content, meta
    return _BOOKS[name]


def _evaluate(content: bytes, tamper=None) -> tuple[str, list, Evaluator, object]:
    book = openpyxl.load_workbook(io.BytesIO(content))
    if tamper:
        tamper(book)
    sys.setrecursionlimit(max(sys.getrecursionlimit(), 100000))
    evaluator = Evaluator(book)
    # Долг строки читает предыдущую — сверху вниз, чтобы не упереться в
    # глубину рекурсии вычислителя.
    for name in book.sheetnames:
        for row in book[name].iter_rows(min_row=nw.FIRST_ROW):
            for cell in row:
                if isinstance(cell.value, str) and cell.value.startswith("="):
                    evaluator.cell(name, cell.coordinate)
    # Каждая формула книги — и вводные-ссылки, и отчёт с ОПУ — понятна
    # вычислителю: непонятая упала бы здесь, а не пролезла.
    for name in book.sheetnames:
        for row in book[name].iter_rows(max_row=nw.FIRST_ROW):
            for cell in row:
                if isinstance(cell.value, str) and cell.value.startswith("="):
                    evaluator.cell(name, cell.coordinate)
    checks = book[nw.CHECK_SHEET]
    bad = [(checks[f"A{r}"].value, checks[f"B{r}"].value, evaluator.cell(nw.CHECK_SHEET, f"C{r}"),
            checks[f"D{r}"].value)
           for r in range(3, checks.max_row + 1)
           if evaluator.cell(nw.CHECK_SHEET, f"G{r}") == "РАСХОЖДЕНИЕ"]
    return evaluator.cell(nw.CHECK_SHEET, "I3"), bad, evaluator, book


@pytest.mark.parametrize("name", list(VARIANTS))
def test_the_book_formulas_agree_with_the_engine(name) -> None:
    content, meta = _book(name)
    assert meta["missing"] == []
    verdict, bad, evaluator, book = _evaluate(content)
    assert bad == [] and verdict == nw.PASSED
    checks = book[nw.CHECK_SHEET]
    sections = {checks[f"A{r}"].value for r in range(3, checks.max_row + 1)}
    assert {"Проект", "Смета", "Собственное участие", "По годам"} <= sections
    # Сверка не пустая: сравниваются настоящие деньги, а не нули.
    compared = {checks[f"B{r}"].value: checks[f"D{r}"].value for r in range(3, checks.max_row + 1)
                if checks[f"A{r}"].value == "Проект"}
    for label in ("Выручка проекта", "CAPEX", "EBITDA", "Налог на прибыль", "НДС", "NPV проекта",
                  "Расходы всего", "Чистая прибыль", "Кредит объектов — пик"):
        assert compared[label] not in ("", None, 0)
    # Свод «на глаз»: выручка − расходы всего = чистая прибыль, EBITDA −
    # проценты = прибыль до налога; обе разности стоят в сверке нулём.
    assert compared["Выручка − расходы всего − чистая прибыль"] == 0.0
    assert compared["EBITDA − проценты − прибыль до налога"] == 0.0


def _operating_row(sheet) -> int:
    """Строка помесячной таблицы за полтора года до её конца — объект в
    эксплуатации."""
    last = max(r for r in range(nw.FIRST_ROW, sheet.max_row + 1)
               if isinstance(sheet[f"A{r}"].value, int)
               and sheet[f"B{r}"].value and str(sheet[f"B{r}"].value).startswith("=EDATE"))
    return last - nw.EXTRA_MONTHS - 6


def _tampered_labels(name: str, tamper) -> set[str]:
    content, _ = _book(name)
    verdict, bad, _, _ = _evaluate(content, tamper)
    assert verdict == nw.FAILED
    return {label for _, label, _, _ in bad}


def test_a_tampered_formula_is_caught() -> None:
    """Подделка формулы: аренда одного месяца завышена на процент."""
    letter = nw.COLUMNS["rent"][0]

    def tamper(book):
        sheet = book["Объект 1"]
        row = _operating_row(sheet)
        sheet[f"{letter}{row}"] = sheet[f"{letter}{row}"].value + "*1.01"

    labels = _tampered_labels("ТЦ аренда и продажа, аннуитет", tamper)
    assert {"Арендная выручка", "NOI за срок удержания", "Выручка проекта"} <= labels


def test_a_tampered_cost_article_is_caught() -> None:
    """Подделка статьи сметы: ставка резерва в «Затратах» не та.

    Генподряд у нежилого проекта без ядра — ноль (ставка объекта «под
    ключ»), и подделка его ставки ничего бы не сдвинула: подделывается резерв.
    """
    def tamper(book):
        sheet = book[nw.COSTS_SHEET]
        row = next(r for r in range(1, nw.FIRST_ROW) if sheet[f"A{r}"].value == "Резерв")
        sheet[f"D{row}"] = sheet[f"D{row}"].value + "*1.2"

    labels = _tampered_labels("ТЦ аренда и продажа, аннуитет", tamper)
    assert {"Резерв", "CAPEX", "Кредит — выборка"} <= labels


def test_a_tampered_input_is_caught() -> None:
    """Подделка вводной: ставка аренды на «Вводных» не та, что у движка."""
    def tamper(book):
        sheet = book[nw.INPUTS_SHEET]
        cell = sheet[f"C{nw.O_ROW['rent']}"]
        cell.value = cell.value * 1.1

    labels = _tampered_labels("ТЦ аренда и продажа, аннуитет", tamper)
    assert {"Арендная выручка", "NOI за срок удержания"} <= labels


def test_the_rent_date_moves_the_book_like_the_engine() -> None:
    """Ставка аренды дана на свою дату и до ввода индексируется: стёртая в
    книге дата снимает индексацию до ввода, и аренда расходится с движком."""
    def tamper(book):
        book[nw.INPUTS_SHEET][f"C{nw.O_ROW['rent_date']}"].value = None

    labels = _tampered_labels("ТЦ аренда и продажа, аннуитет", tamper)
    assert {"Арендная выручка", "NOI за срок удержания"} <= labels


def test_a_tampered_tax_rule_is_caught() -> None:
    """Подделка налога: убыток прошлых лет зачитывается целиком, без
    половинного ограничения, — налог проекта расходится с движком."""
    def tamper(book):
        book[nw.INPUTS_SHEET][f"B{nw.P_ROW['loss_limit']}"] = 1.0

    labels = _tampered_labels("аренда и удержание, погашение из NOI", tamper)
    assert "Налог на прибыль" in labels


def test_the_model_counts_with_formulas_from_one_input_sheet() -> None:
    """Числа на листах модели — формулы; значения стоят только на «Вводных»
    и в столбцах, подписанных «СЧИТАЕТ ДВИЖОК»."""
    content, _ = _book("ТЦ ДКП, продажи до ввода")
    book = openpyxl.load_workbook(io.BytesIO(content))
    assert book.sheetnames[:2] == [nw.INPUTS_SHEET, nw.COSTS_SHEET]
    assert book.sheetnames[-5:] == [nw.CREDIT_SHEET, nw.TAX_SHEET, nw.CASH_SHEET,
                                    nw.REPORT_SHEET, nw.CHECK_SHEET]
    assert "Объект 1" in book.sheetnames
    for name in book.sheetnames:
        if name in (nw.INPUTS_SHEET, nw.CHECK_SHEET):
            continue
        sheet = book[name]
        for row in sheet.iter_rows(min_row=nw.FIRST_ROW, min_col=3):
            for cell in row:
                if isinstance(cell.value, (int, float)) and cell.fill != nw.ENGINE_FILL:
                    assert cell.value == 0 or cell.fill.fgColor.rgb == nw.ENGINE_FILL.fgColor.rgb, (
                        f"{name}!{cell.coordinate} = {cell.value} — число вместо формулы")
    sheet = book["Объект 1"]
    for key, (row, _) in nw.INPUTS.items():
        assert str(sheet[f"B{row}"].value).startswith("="), key
    assert f"'{nw.COSTS_SHEET}'!" in sheet[f"{nw.COLUMNS['capex'][0]}{nw.FIRST_ROW}"].value
    assert f"'{nw.CREDIT_SHEET}'!" in sheet[f"{nw.COLUMNS['key_rate'][0]}{nw.FIRST_ROW}"].value
    # Блок другой стратегии свёрнут: у прямой продажи — аренда, у аренды — ДКП.
    direct = sheet
    retail = openpyxl.load_workbook(io.BytesIO(_book("ТЦ аренда и продажа, аннуитет")[0]))["Объект 1"]
    assert direct.column_dimensions[nw.COLUMNS["rent"][0]].hidden
    assert not direct.column_dimensions[nw.COLUMNS["sale"][0]].hidden
    assert retail.column_dimensions[nw.COLUMNS["sale"][0]].hidden
    assert not retail.column_dimensions[nw.COLUMNS["rent"][0]].hidden
    report = book[nw.REPORT_SHEET]
    for r in range(1, report.max_row + 1):
        value = report[f"B{r}"].value
        assert not isinstance(value, (int, float)), f"Отчёт!B{r} = {value}"


def test_choices_are_russian_lists_and_dates_are_dates() -> None:
    content, _ = _book("ТЦ аренда и продажа, аннуитет")
    book = openpyxl.load_workbook(io.BytesIO(content))
    sheet = book[nw.INPUTS_SHEET]
    assert sheet[f"C{nw.O_ROW['strategy']}"].value == nw.STRATEGY["income"]
    assert sheet[f"C{nw.O_ROW['repayment']}"].value in nw.REPAY.values()
    assert sheet[f"B{nw.P_ROW['rate_scenario']}"].value in nw.SCENARIO.values()
    lists = {cell for dv in sheet.data_validations.dataValidation for rng in dv.sqref.ranges
             for cell in (rng.coord,)}
    for key in ("strategy", "exit_mode", "repayment", "curve"):
        assert f"C{nw.O_ROW[key]}" in lists
    assert f"B{nw.P_ROW['rate_scenario']}" in lists
    for key in ("start", "sales_start"):
        assert hasattr(sheet[f"C{nw.O_ROW[key]}"].value, "year")
    assert hasattr(sheet[f"B{nw.P_ROW['start']}"].value, "year")
    # Чьё умолчание роста цены до ввода — подписано, а не угадывается.
    assert "умолчание объекта" in str(sheet[f"C{nw.O_ROW['growth_pre_source']}"].value)
    codes = {"income", "direct", "sweep", "annuity", "bullet", "hold", "sale", "flat", "bell"}
    for row in sheet.iter_rows():
        for cell in row:
            assert cell.value not in codes, f"{cell.coordinate}: код «{cell.value}» вместо подписи"


def test_switching_a_choice_in_the_book_moves_the_model() -> None:
    """Список — не подпись: выбор «Удержание с оценкой» в книге убирает
    выручку выхода и ставит оценку, как у движка с тем же выбором."""
    content, _ = _book("ТЦ аренда и продажа, аннуитет")

    def tamper(book):
        book[nw.INPUTS_SHEET][f"C{nw.O_ROW['exit_mode']}"] = nw.EXIT["hold"]

    labels = _tampered_labels("ТЦ аренда и продажа, аннуитет", tamper)
    assert {"Выход — продажа", "Удержание — оценка"} <= labels


def test_a_mixed_project_keeps_the_v4_book() -> None:
    x, t = _project(nonresidential=False, offices_strategy="income")
    _, filename, meta = core.build_project_workbook(x, t, [], {})
    assert not meta.get("nonres_book") and "нежилой" not in filename


# --- объект по ДДУ: продажи, эскроу, БРИДЖ и ПФ формулами --------------------
#
# Владелец (13.10.2026): «мне нужна рабочая книга с живыми формулами, а не
# выгрузка в Excel хардов». Нежилой проект с объектом по ДДУ прежде получал
# жилую книгу v4 целиком; теперь — книгу нежилого проекта, где продажи по ДДУ,
# эскроу, БРИДЖ, ПФ со ступенями ставки и LLCR считают формулы листов «ДДУ N»
# и «ПФ», а «Сверка» сравнивает их с движком.

DDU_VARIANTS = {
    "офис по ДДУ": dict(offices_strategy="ddu"),
    "офис по ДДУ и ТЦ в аренде": dict(offices_strategy="ddu", retail_strategy="income", **RETAIL),
    "ТЦ по ДДУ, офис ДКП": dict(offices_strategy="direct", retail_strategy="ddu", **RETAIL),
    "офис по ДДУ: проценты капиталом, свои деньги, лимит, сценарий": dict(
        offices_strategy="ddu", bridge_interest_mode="Оплата капиталом", pre_pf_own_funds_mln=300,
        pf_limit_approved_mln=9000, scenario_revenue_multiplier=0.95, scenario_cost_multiplier=1.05,
        rate_scenario="high"),
    "офис по ДДУ: профиль продаж и лестница цены": dict(
        offices_strategy="ddu", offices_sales_profile="30%@0; 50%@12; 20%@30",
        offices_growth_stage1_pct=5, offices_growth_stage3_pct=7, offices_share_before_rve_pct=70,
        offices_residual_months=9),
}


def _ddu_book(name: str) -> tuple[bytes, dict]:
    key = "ДДУ: " + name
    if key not in _BOOKS:
        x, t = _project(**DDU_VARIANTS[name])
        content, filename, meta = core.build_project_workbook(x, t, [], {}, project_name="Офис")
        assert meta.get("nonres_book") is True and "нежилой" in filename
        _BOOKS[key] = content, meta
    return _BOOKS[key]


@pytest.mark.parametrize("name", list(DDU_VARIANTS))
def test_a_ddu_object_is_counted_by_the_book_formulas(name) -> None:
    content, meta = _ddu_book(name)
    assert meta["missing"] == []
    verdict, bad, _, book = _evaluate(content)
    assert bad == [] and verdict == nw.PASSED
    assert nw.PF_SHEET in book.sheetnames and "ДДУ 1" in book.sheetnames
    checks = book[nw.CHECK_SHEET]
    compared = {checks[f"B{r}"].value: checks[f"D{r}"].value for r in range(3, checks.max_row + 1)
                if checks[f"A{r}"].value == "ПФ"}
    # Сверка ПФ сравнивает настоящие деньги, а не нули.
    for label in ("Выборка ПФ всего, ₽", "ПФ — пик, ₽", "Эскроу — пик, ₽"):
        assert compared[label] not in ("", None, 0), label
    # Ни одного числа движка в листах ДДУ и ПФ: всё — формулы.
    for sheet in ("ДДУ 1", nw.PF_SHEET):
        values = [cell for row in book[sheet].iter_rows(min_row=2) for cell in row
                  if isinstance(cell.value, (int, float)) and cell.column > 1]
        # Числом стоит только номер месяца k (столбец A помесячной таблицы).
        assert [cell.coordinate for cell in values if cell.column != 1] == [], sheet


def test_a_tampered_pf_formula_is_caught() -> None:
    """Подделка ставки ПФ на один месяц краснит сверку ПФ и итогов."""
    letter = nw.PF_COLUMNS["p_rate"][0]

    def tamper(book):
        sheet = book[nw.PF_SHEET]
        for row in range(nw.FIRST_ROW, sheet.max_row + 1):
            value = sheet[f"{letter}{row}"].value
            if isinstance(value, str) and value.startswith("=") and row > nw.FIRST_ROW + 20:
                sheet[f"{letter}{row}"] = value + "*1.5"
                break

    content, _ = _ddu_book("офис по ДДУ")
    verdict, bad, _, _ = _evaluate(content, tamper)
    assert verdict == nw.FAILED
    assert "Проценты и комиссии" in {label for _, label, _, _ in bad}


def test_a_tampered_ddu_price_is_caught() -> None:
    """Вводная цены объекта по ДДУ в книге двигает выручку, эскроу и ПФ."""
    def tamper(book):
        cell = book[nw.INPUTS_SHEET][f"C{nw.O_ROW['price']}"]
        cell.value = cell.value * 1.1

    content, _ = _ddu_book("офис по ДДУ")
    verdict, bad, _, _ = _evaluate(content, tamper)
    assert verdict == nw.FAILED
    assert {"Выручка проекта", "Эскроу — пик, ₽"} <= {label for _, label, _, _ in bad}


def test_the_evaluator_irr_matches_the_engine() -> None:
    import xlsx_eval

    flows = [v * 1e6 for v in (-100.0, -50.0, 10.0, 20.0, 40.0, 60.0, 80.0, 30.0)]
    monthly = xlsx_eval.FUNCTIONS["IRR"]([flows])
    assert (1 + monthly) ** 12 - 1 == pytest.approx(core._monthly_irr(flows), abs=1e-9)
    assert xlsx_eval.FUNCTIONS["INT"]([-6.5]) == -7.0


def test_a_longer_project_than_the_tables_is_caught() -> None:
    """Срок удержания длиннее, чем строк в помесячных таблицах: книга не
    обрезает поток молча — горизонт в сверке красный."""
    def tamper(book):
        cell = book[nw.INPUTS_SHEET][f"C{nw.O_ROW['hold_years']}"]
        cell.value = cell.value + 2

    labels = _tampered_labels("ТЦ аренда и продажа, аннуитет", tamper)
    assert "Горизонт модели, мес." in labels


def test_the_ddu_vat_leaves_the_engine_cash_flows(monkeypatch) -> None:
    """НДС продаж по ДДУ уходит из потоков проекта и капитала (решение владельца 13.10.2026).

    Прежде он был только в чистой прибыли, и NPV/IRR его не видели; книга
    вычитает его в своих потоках, поэтому сверка выше сходится лишь с движком,
    который вычитает его тоже.
    """
    x, t = _project(offices_strategy="ddu")
    result = core.calculate(core.CalcRequest(inputs=x, tep=t))
    original = core.simulate_financing
    seen: list[float] = []

    def without_vat(*args, **kwargs):
        fin = original(*args, **kwargs)
        seen.append(sum(fin.get("vat_ddu_schedule", {}).values()))
        fin["vat_ddu_schedule"] = {}
        return fin

    monkeypatch.setattr(core, "simulate_financing", without_vat)
    blind = core.calculate(core.CalcRequest(inputs=x, tep=t))
    # Первый вызов — основной расчёт; следом считаются пути офиса
    # (`office_paths`), у которых ДДУ нет.
    assert seen and seen[0] > 1e6
    for flow in ("project", "equity"):
        gap = sum(blind["cashflow"][flow]) - sum(result["cashflow"][flow])
        assert abs(gap - seen[0]) < 1.0, flow


def test_the_v4_book_takes_the_vat_out_of_its_flows() -> None:
    """Книга v4 вычитает НДС к уплате (строка 21) из потоков, как движок."""
    import openpyxl

    x, t = _project(nonresidential=False, offices_strategy="ddu")
    content, _, meta = core.build_project_workbook(x, t, [], {})
    assert not [m for m in meta["missing"] if "CF_" in m and ("49" in m or "54" in m or "63" in m)]
    sheet = openpyxl.load_workbook(io.BytesIO(content))["CF_1"]
    column = core._v4_cf_columns()[0]
    for row, _, replacement in core._V4_VAT_CASH_ROWS:
        assert replacement.format(c=column) in str(sheet[f"{column}{row}"].value), row
