"""Что включает ставка объекта — по виду проекта; резерв не берётся с ВРИ.

Решение владельца 06.10.2026 («середина»), один предикат движка —
`object_rate_is_turnkey`:

- объект внутри жилого или смешанного проекта — ставка «под ключ»: ИРД,
  проект, сети, ввод, генподряд и техзаказчик уже в ней; генподряд и
  техзаказчик на стройку объекта не начисляются, резерв — начисляется;
- чисто нежилой проект — полная постатейная смета, как у жилья: ставка
  метрового объекта — СМР здания, статьи проекта идут от суммарной площади
  объектов в ГНС, генподряд, техзаказчик, управление и резерв — и на СМР
  объекта. Наземный паркинг (ставка за место) и тут «под ключ»;
- резерв ни у одного проекта не начисляется на плату за смену ВРИ.

Умолчание ставки СМР в нежилом проекте — «под ключ», уменьшенное на долю
статей проекта (`object_smr_rate`), чтобы на умолчаниях итог сохранился.

Методика одна на движок и обе книги. Каждая проверка падает на прежней
методике — это доказывают подделки ниже.

Запуск: python3 -m pytest tests/test_object_turnkey_overheads.py -q
"""

from __future__ import annotations

import copy
import io
import sys
from pathlib import Path

import openpyxl
import pytest

sys.path.insert(0, str(Path(__file__).resolve().parent.parent))
sys.path.insert(0, str(Path(__file__).resolve().parent))

import main as wrapper  # noqa: E402
import nonres_workbook as nw  # noqa: E402
from xlsx_eval import Evaluator  # noqa: E402

core = wrapper.core


def _amounts(**over) -> dict[str, float]:
    x = copy.deepcopy(core.DEFAULT_INPUTS)
    x.update(offices_enabled=True)
    x.update(over)
    return core.build_operating_model(x, copy.deepcopy(core.TEP_DEFAULT))["capex_amounts"]


def test_the_object_rate_carries_no_gc_fee_and_no_technical_customer() -> None:
    """Дороже объект — генподряд и техзаказчик те же; резерв растёт на
    процент от прибавки. На прежней методике оба процента выросли бы."""
    cheap = _amounts(offices_cost_th_per_sqm=150)
    dear = _amounts(offices_cost_th_per_sqm=250)
    added = dear["offices"] - cheap["offices"]
    assert added > 0
    assert dear["gc_fee"] == pytest.approx(cheap["gc_fee"])
    assert dear["technical_supervision"] == pytest.approx(cheap["technical_supervision"])
    reserve_pct = core.DEFAULT_INPUTS["reserve_pct"] / 100
    assert dear["reserve"] - cheap["reserve"] == pytest.approx(added * reserve_pct)
    # Базы — ровно СМР ядра и соцобъекты, без объекта.
    works = dear["main_above"] + dear["main_under"] + dear["social"]
    assert dear["gc_fee"] == pytest.approx(works * core.DEFAULT_INPUTS["gc_fee_pct"] / 100)
    assert dear["technical_supervision"] == pytest.approx(
        works * core.DEFAULT_INPUTS["technical_supervision_pct"] / 100)


def test_the_reserve_ignores_the_vri_fee() -> None:
    """Плата за ВРИ меняется — резерв нет. На прежней методике резерв брал 5%
    и с платы: 2,86 млрд ₽ умолчаний давали 143 млн ₽ резерва из ниоткуда."""
    low = _amounts(land_rights_cost_mln=1000.0)
    high = _amounts(land_rights_cost_mln=3000.0)
    assert high["land_rights"] - low["land_rights"] == pytest.approx(2000e6)
    assert high["reserve"] == pytest.approx(low["reserve"])


def test_the_reserve_base_is_every_article_but_the_excluded() -> None:
    """Резерв — процент от всех статей до него, кроме списка исключений."""
    a = _amounts()
    shadow = {"reserve", "vri_interest", "vri_security", "land_rights_gross",
              "land_rights_relief"}
    base = sum(v for k, v in a.items() if k not in shadow | core.RESERVE_EXCLUDED_ARTICLES)
    assert a["land_rights"] > 0 and a["offices"] > 0
    assert a["reserve"] == pytest.approx(base * core.DEFAULT_INPUTS["reserve_pct"] / 100)
    assert core.RESERVE_EXCLUDED_ARTICLES <= core._M2_RESERVE_EXCLUDED


def test_the_object_rate_hint_follows_the_project_kind() -> None:
    """Подпись ставки объекта — из того же предиката: «под ключ» в смешанном
    проекте и у паркинга всегда, «СМР» у метрового объекта нежилого."""
    mixed = {"project_kind": core.PROJECT_KIND_MIXED}
    nonres = {"project_kind": core.PROJECT_KIND_NONRESIDENTIAL}
    for obj in core.STANDALONE_OBJECTS:
        _title, fields = core.standalone_object_group(obj)
        hints = {row[0]: row[2] for row in fields}
        assert core.OBJECT_TURNKEY_HINT in hints[obj.rate_cost], obj.key
        assert core.object_rate_hint(mixed, obj) == core.OBJECT_TURNKEY_HINT
        expected = core.OBJECT_TURNKEY_HINT if obj.measure == "spaces" else core.OBJECT_SMR_HINT
        assert core.object_rate_hint(nonres, obj) == expected, obj.key
    assert "__DEVELOPAID_OBJECT_RATE_HINTS__" not in core.PAGE
    assert core.OBJECT_SMR_HINT in core.PAGE


# --- Чисто нежилой проект: полная смета ---------------------------------------

def _nonres_inputs(**over) -> tuple[dict, dict]:
    x = copy.deepcopy(core.DEFAULT_INPUTS)
    t = copy.deepcopy(core.TEP_DEFAULT)
    x.update(offices_enabled=True, offices_gba_sqm=40000.0, offices_saleable_sqm=24000.0,
             offices_parking_under_spaces=0, offices_parking_over_spaces=0,
             _parking_by_hand=["offices"],  # без гаража: метры — только здание
             offices_strategy="income", project_kind=core.PROJECT_KIND_NONRESIDENTIAL)
    x.update(over)
    t.setdefault("offices", {}).update(gns=40000.0, total_area=37600.0,
                                       useful=24000.0, saleable=24000.0)
    for key in core.MKD_PRODUCTS:
        for col in ("gns", "total_area", "useful", "saleable", "transfer", "units"):
            if key in t:
                t[key][col] = 0
    return x, t


def _nonres_amounts(**over) -> dict[str, float]:
    x, t = _nonres_inputs(**over)
    prepared = core.prepared_calculation(x, t, [])
    return core.build_operating_model(prepared["x"], prepared["t"])["capex_amounts"]


def test_the_nonresidential_project_counts_every_article() -> None:
    """Статьи проекта — от площади объектов, генподряд, техзаказчик и
    управление — и на СМР объекта. На прежней методике статьи были нулём."""
    a = _nonres_amounts()
    area = 40000.0
    d = core.DEFAULT_INPUTS
    for key in ("ird", "design_p", "design_rd", "preparation", "utilities",
                "commissioning", "site_maintenance"):
        assert a[key] == pytest.approx(area * d[f"{key}_th_per_sqm"] * 1000), key
    assert a["gc_fee"] == pytest.approx(a["offices"] * d["gc_fee_pct"] / 100)
    assert a["technical_supervision"] == pytest.approx(
        a["offices"] * d["technical_supervision_pct"] / 100)
    management = (a["ird"] + a["design_p"] + a["design_rd"] + a["author_supervision"]
                  + a["preparation"] + a["offices"] + a["utilities"] + a["landscaping"]
                  + a["site_maintenance"])
    assert a["project_management"] == pytest.approx(
        management * d["project_management_pct"] / 100)


def test_a_parking_priced_per_space_stays_turnkey() -> None:
    """Наземный паркинг меряется местами — и в нежилом проекте «под ключ»."""
    a = _nonres_amounts(offices_enabled=False, above_parking_enabled=True,
                        above_parking_spaces=400)
    assert a["above_parking"] > 0
    for key in ("ird", "gc_fee", "technical_supervision"):
        assert a[key] == 0.0, key


def test_the_smr_default_keeps_the_turnkey_total() -> None:
    """Умолчание СМР нежилого — «под ключ», уменьшенное на статьи проекта:
    на умолчаниях стройка здания без резерва та же, что «под ключ»."""
    for cls in ("comfort", "business", "elite"):
        turnkey = core.PROJECT_CLASS_PRESETS[cls]["offices_cost_th_per_sqm"]
        smr = core.class_base_preset(cls, "msk", core.PROJECT_KIND_NONRESIDENTIAL)[
            "offices_cost_th_per_sqm"]
        assert smr < turnkey
        a = _nonres_amounts(offices_cost_th_per_sqm=smr)
        building = sum(v for k, v in a.items()
                       if k not in ("reserve", "land_rights", "land_rights_gross",
                                    "land_rights_relief", "vri_interest", "vri_security",
                                    "total"))
        # Округление ставки до 0,1 тыс ₽/м² — не дальше полутора десятых на метр.
        assert building / 40000 / 1000 == pytest.approx(turnkey, abs=0.15), cls
    # У смешанного проекта база класса — прежняя ставка «под ключ».
    assert core.class_base_preset("comfort", "msk", core.PROJECT_KIND_MIXED)[
        "offices_cost_th_per_sqm"] == core.PROJECT_CLASS_PRESETS["comfort"]["offices_cost_th_per_sqm"]


def _two_objects() -> tuple[dict, dict]:
    x, t = _nonres_inputs(retail_enabled=True, retail_gba_sqm=15000.0,
                          retail_saleable_sqm=11000.0, retail_strategy="income",
                          _parking_by_hand=["offices", "retail"])
    t.setdefault("standalone_retail", {}).update(gns=15000.0, total_area=14100.0,
                                                 useful=11000.0, saleable=11000.0)
    return x, t


def test_two_objects_share_one_construction_block() -> None:
    """Офис + ТЦ: один блок «Строительство» на проект. Статьи проекта и
    проценты считаются от суммарных баз, а объекты получают их долей своей
    стройки — тем же `common_capex`, что делит общие затраты нежилого проекта
    между кредитами объектов. Сумма долей равна статьям проекта; доля каждого
    — его стройка к стройке всех объектов (подделка 50/50 здесь краснеет:
    стройки 7 000 и 3 000 млн)."""
    x, t = _two_objects()
    result = core.calculate(core.CalcRequest(inputs=x, tep=t, rates=[]))
    capex = result["capex"]
    # Один блок: статьи проекта от суммы площадей, своих ставок у объекта нет.
    assert capex["ird"] == pytest.approx((40000 + 15000) * x["ird_th_per_sqm"] * 1000)
    assert capex["gc_fee"] == pytest.approx(
        (capex["offices"] + capex["standalone_retail"]) * x["gc_fee_pct"] / 100)
    objects = {o["key"]: o for o in result["finance"]["nonres"]["objects"]}
    common = {key: float(objects[key]["result"]["common_capex"])
              for key in ("offices", "standalone_retail")}
    project_articles = capex["total"] - capex["offices"] - capex["standalone_retail"]
    assert sum(common.values()) == pytest.approx(project_articles, rel=1e-9)
    built = capex["offices"] + capex["standalone_retail"]
    for key, value in common.items():
        assert value == pytest.approx(project_articles * capex[key] / built, rel=1e-9), key
    assert abs(common["offices"] - common["standalone_retail"]) > 0.3 * project_articles


def test_a_turnkey_object_takes_no_share_of_the_gc_fee() -> None:
    """Смета объекта (`object_cost_parts`): у объекта «под ключ» в смешанном
    проекте доли генподряда и техзаказчика нет — они в его ставке. Прежняя
    раскладка own / works_base дала бы офису долю генподряда дома."""
    x = copy.deepcopy(core.DEFAULT_INPUTS)
    x.update(offices_enabled=True)
    parts = core.build_operating_model(x, copy.deepcopy(core.TEP_DEFAULT))["object_cost_parts"]
    assert parts["offices"]["building"] > 0
    assert parts["offices"]["gc_fee"] == 0.0
    assert parts["offices"]["technical_supervision"] == 0.0
    assert parts["offices"]["reserve"] == pytest.approx(
        (parts["offices"]["building"] + parts["offices"]["garage"])
        * core.DEFAULT_INPUTS["reserve_pct"] / 100)


# --- Книга ПЛАТО v4 ----------------------------------------------------------

_V4_ROWS = {27: "technical_supervision", 29: "gc_fee", 30: "reserve"}


def _v4_book():
    inputs = dict(core.DEFAULT_INPUTS, offices_enabled=True)
    tep = {key: dict(value) for key, value in core.TEP_DEFAULT.items()}
    op = core.build_operating_model(dict(inputs), copy.deepcopy(tep))
    content, _, meta = core.build_project_workbook(inputs, tep, [], None, finance_hints={})
    assert not [m for m in meta.get("missing", []) if "CAPEX" in m]
    return openpyxl.load_workbook(io.BytesIO(content), data_only=False), op["capex_amounts"]


def test_the_v4_book_counts_the_same_bases() -> None:
    """Книга v4 считает генподряд, техзаказчика и резерв на базах движка;
    вернуть шаблонную формулу — получить расхождение."""
    sys.setrecursionlimit(400000)
    book, engine = _v4_book()
    sheet = book["CAPEX"]
    for row, key in _V4_ROWS.items():
        assert "ОБЪЕКТЫ" not in sheet[f"B{row}"].value or key == "reserve"
        value = float(Evaluator(book).cell("CAPEX", f"B{row}") or 0)
        assert value == pytest.approx(engine[key] / 1e6, abs=0.05), key
    # Подделка: прежние формулы шаблона — объект в базе процентов, ВРИ в резерве.
    sheet["B29"] = sheet["B29"].value.replace(")*", "+'ОБЪЕКТЫ'!$B$96)*", 1)
    sheet["B30"] = sheet["B30"].value.replace("SUM(B16:B29)", "SUM(B15:B29)", 1)
    evaluator = Evaluator(book)
    assert float(evaluator.cell("CAPEX", "B29")) > engine["gc_fee"] / 1e6 + 1
    assert float(evaluator.cell("CAPEX", "B30")) > engine["reserve"] / 1e6 + 1


# --- Книга нежилого проекта --------------------------------------------------

def _nonres_content() -> bytes:
    x = copy.deepcopy(core.DEFAULT_INPUTS)
    t = copy.deepcopy(core.TEP_DEFAULT)
    x.update(offices_enabled=True, offices_gba_sqm=40000.0, offices_saleable_sqm=24000.0,
             offices_strategy="income", project_kind=core.PROJECT_KIND_NONRESIDENTIAL)
    t.setdefault("offices", {}).update(gns=40000.0, total_area=37600.0,
                                       useful=24000.0, saleable=24000.0)
    for key in core.MKD_PRODUCTS:
        for col in ("gns", "total_area", "useful", "saleable", "transfer", "units"):
            if key in t:
                t[key][col] = 0
    content, _, meta = core.build_project_workbook(x, t, [], {}, project_name="Офис")
    assert meta.get("nonres_book") is True and meta["missing"] == []
    return content


def _nonres_check(book) -> tuple[str, set[str], Evaluator]:
    sys.setrecursionlimit(max(sys.getrecursionlimit(), 100000))
    evaluator = Evaluator(book)
    checks = book[nw.CHECK_SHEET]
    bad = {checks[f"B{r}"].value for r in range(3, checks.max_row + 1)
           if evaluator.cell(nw.CHECK_SHEET, f"G{r}") == "РАСХОЖДЕНИЕ"}
    return evaluator.cell(nw.CHECK_SHEET, "I3"), bad, evaluator


def test_the_nonres_book_counts_the_same_bases() -> None:
    """Книга нежилого проекта — полная смета формулами: статьи проекта от
    площади объектов, генподряд и техзаказчик на СМР объекта, резерв без
    строки платы за ВРИ; «Сверка» ПРОЙДЕНО. Подделка — объект «под ключ»
    (прежняя методика) — ловится сверкой."""
    content = _nonres_content()
    book = openpyxl.load_workbook(io.BytesIO(content))
    sheet = book[nw.COSTS_SHEET]
    rows = {sheet[f"A{r}"].value: r for r in range(1, nw.FIRST_ROW) if sheet[f"A{r}"].value}
    verdict, bad, evaluator = _nonres_check(book)
    assert verdict == nw.PASSED and not bad
    for label in ("ИРД", "Сети", "Генподряд", "Технический заказчик"):
        assert float(evaluator.cell(nw.COSTS_SHEET, f"E{rows[label]}") or 0) > 0, label
    building = next(r for label, r in rows.items() if str(label).endswith(": здание"))
    vri_row = next(r for label, r in rows.items() if str(label).startswith("Плата за смену ВРИ"))
    reserve_base = sheet[f"B{rows['Резерв']}"].value.lstrip("=").split("+")
    assert f"E{building}" in reserve_base
    assert f"E{vri_row}" not in reserve_base

    tampered = openpyxl.load_workbook(io.BytesIO(content))
    tampered[nw.INPUTS_SHEET][f"C{nw.O_ROW['turnkey']}"] = nw.YES
    verdict, bad, _ = _nonres_check(tampered)
    assert verdict == nw.FAILED and {"Генподряд", "ИРД"} <= bad


# --- Страница: умолчание СМР и подпись ставки ---------------------------------

PORT = 19773


def test_the_page_follows_the_project_kind() -> None:
    """На отрисованной странице: выбор нежилого вида ставит умолчание СМР
    (`object_smr_rate`) и подпись «основное строительство здания (СМР)»,
    возврат к смешанному — прежние «под ключ»; нежилой проект, сохранённый
    до решения со ставкой «под ключ», при загрузке получает СМР."""
    from browser import chromium_or_skip, serve
    from playwright.sync_api import sync_playwright

    smr = core.class_base_preset("comfort", "msk", core.PROJECT_KIND_NONRESIDENTIAL)[
        "offices_cost_th_per_sqm"]
    turnkey = core.PROJECT_CLASS_PRESETS["comfort"]["offices_cost_th_per_sqm"]
    retail_smr = core.object_smr_rate(core.OBJECT_SMR_RATE_DEFAULTS["retail_cost_th_per_sqm"],
                                      core.PROJECT_CLASS_PRESETS["comfort"])
    unit = ("()=>{const s=document.querySelector("
            "'.field[data-field=\"offices_cost_th_per_sqm\"] > label > .unit');"
            "return s?s.textContent:''}")
    with serve(wrapper.app, PORT):
        with sync_playwright() as pw:
            browser = pw.chromium.launch(executable_path=chromium_or_skip())
            page = browser.new_page()
            errors: list[str] = []
            page.on("pageerror", lambda exc: errors.append(str(exc)))
            page.goto(f"http://127.0.0.1:{PORT}/", wait_until="load")
            page.wait_for_function("()=>typeof applyProjectKind==='function'", timeout=60000)
            page.evaluate("()=>{inputs.project_class='comfort';inputs.offices_enabled=true}")
            page.evaluate("()=>{try{applyProjectKind('nonresidential')}catch(e){}"
                          "renderInputs()}")
            got = page.evaluate("()=>[inputs.offices_cost_th_per_sqm,inputs.retail_cost_th_per_sqm,"
                                "inputs._object_rate_kind]")
            hint_nonres = page.evaluate(unit)
            page.evaluate("()=>{try{applyProjectKind('mixed')}catch(e){}renderInputs()}")
            back = page.evaluate("()=>[inputs.offices_cost_th_per_sqm,inputs.retail_cost_th_per_sqm]")
            hint_mixed = page.evaluate(unit)
            # Сохранённый до решения нежилой проект: пометки вида нет.
            page.evaluate(
                "t=>{localStorage.setItem('plato_v04',JSON.stringify({inputs:{"
                "project_kind:'nonresidential',project_class:'comfort',offices_enabled:true,"
                "offices_cost_th_per_sqm:t}}));loadLocal();}", turnkey)
            loaded = page.evaluate("()=>[inputs.offices_cost_th_per_sqm,inputs._object_rate_kind]")
            browser.close()
    assert errors == []
    assert got == [smr, retail_smr, core.PROJECT_KIND_NONRESIDENTIAL]
    assert core.OBJECT_SMR_HINT in hint_nonres, hint_nonres
    assert back == [turnkey, core.OBJECT_SMR_RATE_DEFAULTS["retail_cost_th_per_sqm"]]
    assert core.OBJECT_TURNKEY_HINT in hint_mixed, hint_mixed
    assert loaded == [smr, core.PROJECT_KIND_NONRESIDENTIAL]


def test_the_book_shares_the_block_by_construction() -> None:
    """Книга двух объектов: доля общих затрат — стройка объекта к стройке
    всех, «Сверка» ПРОЙДЕНО; поддельная доля 50/50 краснеет."""
    x, t = _two_objects()
    content, _, meta = core.build_project_workbook(x, t, [], {}, project_name="Два")
    assert meta.get("nonres_book") is True and meta["missing"] == []
    book = openpyxl.load_workbook(io.BytesIO(content))
    verdict, bad, _ = _nonres_check(book)
    assert verdict == nw.PASSED and not bad
    sheet = book[nw.COSTS_SHEET]
    head = next(r for r in range(1, nw.FIRST_ROW)
                if str(sheet[f"A{r}"].value or "").startswith("Доля объекта в общих затратах"))
    tampered = openpyxl.load_workbook(io.BytesIO(content))
    for row in (head + 1, head + 2):
        tampered[nw.COSTS_SHEET][f"B{row}"] = 0.5
    verdict, bad, _ = _nonres_check(tampered)
    assert verdict == nw.FAILED and bad
