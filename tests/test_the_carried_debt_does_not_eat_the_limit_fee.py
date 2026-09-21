"""Принятый долг лимита не выбирает — значит и платы за невыбранный не снижает.

Правило владельца от 27.08.2026: долг, переоформленный с предыдущей очереди,
приходит на линию, НЕ выбирая лимита, — банк переносит обязательство, а не
выдаёт новые деньги. Проценты он несёт и покрытие разбавляет, но свободного
лимита не уменьшает.

Плата за невыбранный лимит считалась от общего остатка, в котором принятый долг
уже лежит, — то есть свободный лимит выходил меньше на всю принятую сумму, а
комиссия меньше положенной. Книга повторяла это ТОЙ ЖЕ формулой (в строку 43
листов CF принятый долг добавлялся намеренно, одним комментарием с процентами и
покрытием, где он нужен), поэтому паритет молчал: обе поверхности ошибались
одинаково, и «Excel сходится с движком» ничего не значило.

Замер на проверочном проекте: комиссия принявшей очереди 131,7 млн ₽ вместо
160,3 — занижение на 28,5 млн, и ровно на его три четверти завышена чистая
прибыль.

Проверяется ИНВАРИАНТОМ, а не числом: при неизменной своей выборке и неизменном
своём лимите добавление принятого долга меняет проценты и покрытие, но плату за
невыбранный лимит не трогает вовсе.

Запуск: python3 -m pytest tests/test_the_carried_debt_does_not_eat_the_limit_fee.py -q
"""

from __future__ import annotations

import copy
import io
import sys
from pathlib import Path

import openpyxl
import pytest

sys.path.insert(0, str(Path(__file__).resolve().parent.parent))
sys.path.insert(0, str(Path(__file__).resolve().parent))

from xlsx_eval import Evaluator  # noqa: E402

import main as _wrapper  # noqa: E402

core = _wrapper.core

# Строка «ПФ — плата за лимит» листов CF: её база и есть предмет спора.
LIMIT_FEE_ROW = "B43"


def _phasing(carry: bool) -> dict:
    return {
        "enabled": True, "mode": "phased", "user_enabled": True,
        "phase_count": 2, "phase_gap_months": 12,
        "phases": [{"name": "О1", "start_offset_months": 0, "construction_months": 24},
                   {"name": "О2", "start_offset_months": 12, "construction_months": 24}],
        # Первая очередь мала, а земля и социалка на ней целиком: ровно тот
        # перекос, при котором она не гасит свой ПФ раскрытым эскроу.
        "products": {key: [35, 65] for key in
                     ("apartments", "ground_commercial", "underground_parking", "storage")},
        "shared_cash": {}, "shared_allocation": {}, "social_objects": [],
        "carry_debt_forward": carry,
    }


def _inputs() -> dict:
    return {**core.DEFAULT_INPUTS, "purchase_price_mln": 12000,
            "project_start": "2027-01-01", "ird_months": 12,
            "construction_months": 24, "apartment_price_th": 700}


def _tep() -> dict:
    return {key: dict(row) for key, row in core.TEP_DEFAULT.items()}


@pytest.fixture(scope="module")
def pair() -> dict:
    """Один и тот же проект с переносом и без. Больше ничего не меняется."""
    out = {}
    for carry in (True, False):
        bundle = core.calculate_phased(core.PhasedCalcRequest(
            inputs=_inputs(), tep=_tep(), rates=[], phasing=_phasing(carry)))
        out[carry] = bundle
    return out


def test_the_fixture_actually_carries_the_debt(pair):
    """Предохранитель: не сработал перенос — проверять нечего.

    Инвариант «плата не изменилась» на проекте БЕЗ переноса верен даром, и
    зелёным он был бы при любой ошибке в базе комиссии.
    """
    assert (pair[True].get("debt_carry") or {}).get("applied") is True
    accepted = pair[True]["phases"][1]["result"]["finance"]["carried_debt_in"]
    assert accepted > 1e9, "принятый долг обязан быть заметной величиной"
    assert pair[False]["phases"][1]["result"]["finance"]["carried_debt_in"] == 0.0


def test_the_carried_debt_leaves_the_limit_fee_alone(pair):
    """Своя выборка та же, свой лимит тот же — значит и комиссия та же."""
    with_carry = pair[True]["phases"][1]["result"]["finance"]
    without = pair[False]["phases"][1]["result"]["finance"]

    # Предпосылка инварианта: принятый долг ничего не финансирует, поэтому
    # своя выборка от него не зависит. Разъедься она — сравнивать было бы
    # нечего, и «комиссия совпала» ничего бы не значило.
    assert with_carry["pf_draw_total"] == pytest.approx(without["pf_draw_total"])
    assert with_carry["pf_limit"] == pytest.approx(without["pf_limit"])

    assert with_carry["pf_limit_fee"] == pytest.approx(without["pf_limit_fee"], abs=1.0)


def test_the_carried_debt_still_costs_interest_and_dilutes_coverage(pair):
    """Обратная половина правила: лимита не выбирает — но тело несёт.

    Без неё «комиссия не изменилась» достигалось бы тем, что принятый долг
    перестал существовать вовсе.
    """
    with_carry = pair[True]["phases"][1]["result"]["finance"]
    without = pair[False]["phases"][1]["result"]["finance"]

    assert with_carry["peak_pf"] > without["peak_pf"]
    assert with_carry["pf_interest"] > without["pf_interest"]


@pytest.fixture(scope="module")
def book_with_carry():
    """Книга того же проекта — та, которую скачивают, а не пересказ формул."""
    sys.setrecursionlimit(400000)
    inputs, tep, phasing = _inputs(), _tep(), _phasing(True)
    bundle = core._run_authoritative_model(
        inputs, copy.deepcopy(tep), [], copy.deepcopy(phasing))
    content, _, meta = core.build_project_workbook(
        inputs, copy.deepcopy(tep), [], copy.deepcopy(phasing))
    assert meta["missing"] == [], meta["missing"]
    book = openpyxl.load_workbook(io.BytesIO(content), data_only=False)
    return bundle, Evaluator(book)


def test_the_book_charges_the_same_limit_fee_as_the_engine(book_with_carry):
    """Обе поверхности считают одинаково — и теперь одинаково ВЕРНО.

    Прежде они тоже совпадали: книга повторяла ошибку движка формулой, и
    паритет был зелёным на неверном числе. Поэтому рядом стоит проверка
    инварианта выше — совпадение двух поверхностей само по себе ничего не
    доказывает, пока не доказано, с чем они совпали.
    """
    bundle, evaluator = book_with_carry
    assert (bundle.get("debt_carry") or {}).get("applied") is True, (
        "предохранитель: книга собрана без переноса — сверять нечего")
    for index, sheet in enumerate(("CF_1", "CF_2")):
        engine = float(bundle["phases"][index]["result"]["finance"]["pf_limit_fee"]) / 1e6
        book = float(evaluator.cell(sheet, LIMIT_FEE_ROW) or 0)
        assert book == pytest.approx(engine, abs=0.01), (
            f"{sheet}: книга {book:,.3f} против движка {engine:,.3f} млн ₽")


def test_the_receiving_queue_really_is_the_one_with_the_carried_debt(book_with_carry):
    """Предохранитель книги: принятый долг обязан стоять в её строке 64.

    Без него сверка выше идёт по очереди, которая ничего не принимала, и
    сравнивает две одинаково простые величины.
    """
    _, evaluator = book_with_carry
    accepted = float(evaluator.cell("CF_2", "B64") or 0)
    assert accepted > 1000, "строка 64 CF_2 пуста — книга долг не приняла"
