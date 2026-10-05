"""«Финансирование объекта»: кредит объекта вне ДДУ своим отчётом.

Владелец (05.10.2026): в нежилом проекте раздел «Финансирование» показывал
БРИДЖ и ПФ нулями и LLCR −8,68x, а кредита объекта — аннуитета, баллона,
схемы погашения — не было видно вовсе. Проверки держат обещания:

* движок раскладывает кредит объекта один раз (`object_financing`):
  выборка, долг на ввод, платёж, тело, баллон, выход — и части погашения
  складываются во всё погашение;
* таблица `nonres_financing` печатает только строки своей схемы: у прямой
  продажи нет аннуитета и баллона, а поля «погашение / срок / баллон» при ней
  не рисуются в форме — движок их не читает;
* таблица кредита живёт в одном месте: в таблице стратегии её строк нет;
* страница рисует блок числами движка (стенд на node), PDF и книга печатают
  ту же таблицу.

Запуск: python3 -m pytest tests/test_nonres_object_financing.py -q
"""

from __future__ import annotations

import copy
import io
import json
import sys
from pathlib import Path

import pytest

ROOT = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(ROOT))
sys.path.insert(0, str(ROOT / "tests"))

import developaid_nonres_strategy as ns  # noqa: E402
import main_legacy as core  # noqa: E402
import page_blocks  # noqa: E402
from test_nonres_object_result import _objects, _run, _spec  # noqa: E402


def _financing(result: dict, key: str) -> dict:
    return next(i for i in result["report"]["nonres_financing"] if i["key"] == key)


def _rows(item: dict) -> dict[str, dict]:
    return {row["label"]: row for row in item["rows"]}


# --- движок -------------------------------------------------------------------

@pytest.mark.parametrize("over", [{}, {"offices_exit_mode": "hold"},
                                  {"offices_debt_repayment": "sweep"},
                                  {"offices_debt_repayment": "bullet"},
                                  {"offices_hold_years": 12}])
def test_the_repayment_parts_add_up_to_the_whole_debt(over) -> None:
    for item in _objects(_run(**over)).values():
        f = item["financing"]
        parts = (f["repaid_scheduled"] + f["balloon_paid"] + f["repaid_at_exit"]
                 + f["repaid_from_cash"])
        assert parts == pytest.approx(f["repaid_total"], rel=1e-9)
        assert f["repaid_total"] == pytest.approx(
            f["draw_total"] + f["interest_capitalized"], rel=1e-9)
        assert f["years"][-1]["balance_end"] == pytest.approx(0.0, abs=1.0)
        assert sum(y["draw"] for y in f["years"]) == pytest.approx(f["draw_total"], rel=1e-9)


def test_the_annuity_shows_its_payment_and_balloon() -> None:
    office = _objects(_run())["offices"]
    f = office["financing"]
    month = f["first_payment_month"]
    series = office["rows"]
    row = next(r for r in series if r["month"] == month)
    assert f["first_payment"] == pytest.approx(
        row["nonres_loan_interest_paid"] + row["nonres_loan_repayment"], rel=1e-9)
    assert f["balloon_planned"] == pytest.approx(
        f["debt_at_commissioning"] * core.DEFAULT_INPUTS["offices_loan_balloon_pct"] / 100, rel=1e-9)
    rows = _rows(_financing(_run(), "offices"))
    assert rows["Схема погашения"]["value"] == core.REPAYMENT_LABELS[ns.REPAY_ANNUITY]
    assert rows["Аннуитетный платёж — средний, в месяц"]["unit"] == "mln"
    # Удержание 5 лет короче кредита 10 лет: баллона в срок нет, остаток — при выходе.
    assert "позже выхода" in rows["Срок кредита"]["value"]
    assert any(label.startswith("Остаток долга погашен при выходе") for label in rows)
    assert not any(label.startswith("Баллон погашен в срок") for label in rows)


def test_a_loan_maturing_before_the_exit_pays_its_balloon() -> None:
    result = _run(offices_hold_years=12)
    f = _objects(result)["offices"]["financing"]
    assert f["balloon_paid"] > 0 and f["balloon_month"] == f["maturity"]
    assert f["balloon_paid"] == pytest.approx(f["balloon_planned"], rel=0.05)
    rows = _rows(_financing(result, "offices"))
    assert any(label.startswith("Баллон погашен в срок") for label in rows)


def test_a_direct_sale_repays_from_the_sales_and_has_no_annuity() -> None:
    rows = _rows(_financing(_run(), "standalone_retail"))
    assert rows["Схема погашения"]["value"] == core.DIRECT_REPAYMENT_LABEL
    assert "Погашено из выручки ДКП" in rows
    assert not any("ннуитет" in label or "аллон" in label or label == "Срок кредита"
                   for label in rows)
    item = _financing(_run(), "standalone_retail")
    assert [c[0] for c in item["columns"]] == ["year", "draw", "interest", "repayment", "balance_end"]


def test_the_form_hides_the_annuity_fields_for_a_direct_sale() -> None:
    readers = core.strategy_field_readers()
    for prefix in ("offices", "retail"):
        for suffix in ("debt_repayment", "loan_term_years", "loan_balloon_pct"):
            assert readers[f"{prefix}_{suffix}"][1] == [ns.STRATEGY_INCOME]
        assert set(readers[f"{prefix}_loan_share_pct"][1]) == {ns.STRATEGY_DIRECT, ns.STRATEGY_INCOME}


def test_the_strategy_table_no_longer_repeats_the_loan() -> None:
    for item in _run()["report"]["nonres_strategy"]:
        assert not any(row["label"].startswith("Кредит объекта") or "DSCR" in row["label"]
                       for row in item["rows"])


# --- PDF и книга ----------------------------------------------------------------

def test_the_pdf_prints_the_object_financing(tmp_path) -> None:
    pypdf = pytest.importorskip("pypdf")
    x, t = _spec()
    result = core.calculate(core.CalcRequest(inputs=x, tep=t, rates=[]))
    content = core._build_developaid_pdf({"project_name": "Офис", "result": result,
                                          "inputs": x, "tep": t, "rates": []})
    path = tmp_path / "f.pdf"
    path.write_bytes(content)
    flat = " ".join(" ".join(p.extract_text() or "" for p in pypdf.PdfReader(str(path)).pages).split())
    assert "Финансирование объекта: Офисы" in flat
    # Таблица стоит в разделе «Финансирование» — за графиком долга кредита
    # объектов, который есть только там, — а не среди продуктов.
    chart = flat.index("Кредит объектов — остаток долга")
    first = flat.index("Финансирование объекта: ")
    assert chart < first and first - chart < 600
    assert "Финансирование объекта: Коммерция ОСЗ" in flat
    payment = _rows(_financing(result, "offices"))["Аннуитетный платёж — средний, в месяц"]
    assert core._pdf_nonres_value(payment) in flat


def test_the_nonresidential_pdf_financing_section_is_not_empty(tmp_path) -> None:
    """Снимок владельца: раздел «Финансирование» нежилого PDF был пуст —
    одна фраза со ссылкой на другую таблицу."""
    pypdf = pytest.importorskip("pypdf")
    from test_nonres_debt_metric import _nonres_inputs
    x, t = _nonres_inputs()
    result = core.calculate(core.CalcRequest(inputs=x, tep=t, rates=[]))
    content = core._build_developaid_pdf({"project_name": "Офис", "result": result,
                                          "inputs": x, "tep": t, "rates": []})
    path = tmp_path / "n.pdf"
    path.write_bytes(content)
    flat = " ".join(" ".join(p.extract_text() or "" for p in pypdf.PdfReader(str(path)).pages).split())
    intro = flat.index("Проект без продаж по ДДУ")
    assert intro < flat.index("Кредит объектов — остаток долга") < flat.index("Финансирование объекта: Офисы")
    assert "в таблице «Нежильё — стратегия реализации»" not in flat
    for label in ("Аннуитетный платёж — средний, в месяц", "Баллон по графику (доля долга на ввод)",
                  "Обслуживание кредита по годам"):
        assert label in flat


def test_the_book_prints_the_object_financing() -> None:
    import openpyxl
    x, t = _spec()
    content, _name, _meta = core.build_project_workbook(x, t, [], {})
    sheet = openpyxl.load_workbook(io.BytesIO(content))["Нежильё — стратегия"]
    cells = [[sheet.cell(r, c).value for c in range(1, 9)] for r in range(1, sheet.max_row + 1)]
    start = next(i for i, row in enumerate(cells) if row[0] == "Финансирование объекта: Офисы")
    values = {row[0]: row[1] for row in cells[start:start + 40]}
    item = _financing(_run(), "offices")
    for row in item["rows"]:
        if row["unit"] in ("rub", "mln"):
            assert values[row["label"]] == pytest.approx(row["value"], rel=1e-9)
    header = next(i for i in range(start, len(cells)) if cells[i][0] == "Год")
    assert cells[header][:len(item["columns"])] == [c[1] for c in item["columns"]]
    assert cells[header + 1][0] == str(item["years"][0]["year"])


# --- страница ----------------------------------------------------------------

def _render(report: dict) -> dict:
    prelude = """
const BOXES=[{innerHTML:''},{innerHTML:''}], CARDS=[{hidden:true},{hidden:true}];
const document={querySelectorAll:sel=>sel==='.nonres-finance-box'?BOXES:CARDS};
const R=%s;
""" % json.dumps({"report": report}, ensure_ascii=False)
    tail = """
renderNonresFinancing(R);
const html=BOXES[0].innerHTML;
const tables=[...html.matchAll(/<table class="nonres-strategy nonres-finance" data-object="([^"]*)"[^>]*>([\\s\\S]*?)<\\/table>/g)]
 .map(m=>({object:m[1],rows:[...m[2].matchAll(/<tr( class="section")?><td>([\\s\\S]*?)<\\/td><td>([\\s\\S]*?)<\\/td><\\/tr>/g)]
   .map(r=>({section:!!r[1],label:r[2],cell:r[3]}))}));
const years=[...html.matchAll(/<table class="nonres-finance-years" data-object="([^"]*)">([\\s\\S]*?)<\\/table>/g)]
 .map(m=>({object:m[1],head:[...m[2].matchAll(/<th>([^<]*)<\\/th>/g)].map(x=>x[1]),
   body:[...m[2].matchAll(/<tr><td>([\\s\\S]*?)<\\/tr>/g)].map(x=>x[0].replace(/<\\/?tr>/g,'').split('</td><td>').map(c=>c.replace(/<\\/?td>/g,'')))}));
const expected=(R.report.nonres_financing||[]).map(o=>({rows:o.rows.map(r=>nonresCell(r)),
 years:o.years.map(y=>o.columns.map(c=>c[2]==='text'?String(y[c[0]]):nonresCell({unit:c[2],value:y[c[0]]})))}));
console.log(JSON.stringify({cards:CARDS.map(c=>c.hidden),same:BOXES[0].innerHTML===BOXES[1].innerHTML,tables,years,expected}));
"""
    out, _ = page_blocks.run(prelude, tail)
    return json.loads(out)


def _mismatches(drawn: dict, report: dict) -> list[str]:
    items = report.get("nonres_financing") or []
    out: list[str] = []
    if [t["object"] for t in drawn["tables"]] != [i["key"] for i in items]:
        return ["объекты не те"]
    for table, years, item, exp in zip(drawn["tables"], drawn["years"], items, drawn["expected"]):
        if [r["label"] for r in table["rows"]] != [r["label"] for r in item["rows"]]:
            out.append(f"{item['key']}: строки не те")
        if [r["cell"] for r in table["rows"]] != exp["rows"]:
            out.append(f"{item['key']}: значения не те")
        if [r["label"] for r in table["rows"] if r["section"]] != \
                [r["label"] for r in item["rows"] if r["unit"] == "section"]:
            out.append(f"{item['key']}: блоки не выделены")
        if years["head"] != [c[1] for c in item["columns"]] or years["body"] != exp["years"]:
            out.append(f"{item['key']}: таблица по годам не та")
    return out


def test_the_page_draws_the_object_financing() -> None:
    report = _run()["report"]
    drawn = _render(report)
    assert drawn["cards"] == [False, False] and drawn["same"] is True
    assert _mismatches(drawn, report) == []
    office = next(t for t in drawn["tables"] if t["object"] == "offices")
    payment = next(r for r in office["rows"] if r["label"] == "Аннуитетный платёж — средний, в месяц")
    assert payment["cell"].endswith(" млн ₽")
    sections = [r["label"] for r in office["rows"] if r["section"]]
    assert sections == ["Условия кредита", "Выборка и долг на ввод", "Погашение", "Обслуживание"]


def test_the_page_check_fails_on_a_forgery() -> None:
    report = _run()["report"]
    drawn = _render(report)
    tampered = copy.deepcopy(drawn)
    tampered["years"][0]["body"][0][1] = "0,0 млн ₽"
    tampered["tables"][1]["rows"][0]["section"] = False
    assert len(_mismatches(tampered, report)) == 2


def test_no_objects_outside_ddu_hide_the_card() -> None:
    drawn = _render(_run(offices_strategy="ddu", retail_strategy="ddu")["report"])
    assert drawn["cards"] == [True, True] and drawn["tables"] == []


def test_a_pure_nonresidential_project_reserves_the_object_line_not_a_bridge() -> None:
    """Владелец (05.10.2026): чисто нежилой проект БРИДЖа не открывает —
    участок и проект в доле кредита берёт НКЛ объекта, и плата за
    резервирование считается с её лимита (вся плановая выборка) в месяц
    первой выдачи. Прежде движок брал её ещё и с «лимита БРИДЖа» от оплат
    участка — вторая комиссия за кредит, которого нет. Жилой проект с ПФ
    комиссию БРИДЖа сохраняет."""
    from test_nonres_debt_metric import _nonres_inputs

    x, t = _nonres_inputs(offices_strategy="direct", purchase_price_mln=900)
    result = core.calculate(core.CalcRequest(inputs=x, tep=t, rates=[]))
    assert result["finance"]["bridge_fee"] == 0.0
    for item in result["finance"]["nonres"]["objects"]:
        f = item["financing"]
        assert f["limit"] == pytest.approx(f["draw_total"])
        assert f["reservation_fee"] == pytest.approx(
            f["limit"] * float(x["reservation_fee_pct"]) / 100, rel=1e-12)
        assert f["reservation_fee"] + f["fee"] == pytest.approx(item["totals"]["loan_fee"], rel=1e-12)
        rows = _rows({"rows": next(r["rows"] for r in result["report"]["nonres_financing"]
                                   if r["key"] == item["key"])})
        assert rows["Плата за резервирование лимита НКЛ"]["value"] == pytest.approx(f["reservation_fee"])
    housing = core.calculate(core.CalcRequest(inputs=copy.deepcopy(core.DEFAULT_INPUTS),
                                              tep=copy.deepcopy(core.TEP_DEFAULT), rates=[]))
    assert housing["finance"]["bridge_fee"] > 0
