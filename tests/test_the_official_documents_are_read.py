"""Выписки ЕГРН и извещение о торгах: разбор первоисточников по КРТ Нагатино.

Владелец прислал 59 выписок Роскадастра от 18.08.2026 и извещение о торгах
ДГП-Р-54/26 от 14.08.2026. Они отвечают на то, чего публичный ответ НСПД не
даёт: какой участок под зданием, кто собственник, чем обременена земля.

Здесь закреплены ловушки, каждая из которых меняла ответ.

**Перечень лежит в одном поле через запятую** — разбор по узлам терял три
участка из четырёх.

**Личность — это ИНН, а не написание имени**: одна компания приходит и капсом,
и обычным письмом, а «Автокомбинат № 19» — то ЗАО, то АО.

**Оперативное управление — не собственность.** Девять строений записаны за
городом Москвой, а держит их ГБУ «Жилищник» на праве оперативного управления.

**Слово в графе числа — это не число.** У `:2091` в проекте решения стоит
«часть»; разбор, подбирающий похожий токен, дал 4,0 м².

**Объект стоит на нескольких участках** и повторяется строкой у каждого: 47
строк извещения — это 39 объектов.

Запуск: python3 -m pytest tests/test_the_official_documents_are_read.py -q
"""

from __future__ import annotations

import sys
from pathlib import Path

import pytest

ROOT = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(ROOT))

from auction_search import egrn_extracts, krt_notice  # noqa: E402
from auction_search import nagatino_parcels as parcels  # noqa: E402

EXTRACTS = ROOT / "docs" / "krt" / "egrn"
NOTICE = ROOT / "docs" / "krt" / "nagatino-auction-notice-2026-08-14.pdf"


def _read(number: str):
    return egrn_extracts.read((EXTRACTS / f"{number.replace(':', '_')}.xml").read_bytes())


def test_the_primary_sources_are_in_the_repository():
    """Файл, а не ссылка: ссылки протухают, выписка нет."""
    assert len(list(EXTRACTS.glob("*.xml"))) == 59
    assert NOTICE.exists()


def test_a_list_inside_one_field_is_not_one_number():
    """У 77:05:0004001:1098 в одном узле четыре участка через запятую."""
    record = _read("77:05:0004001:1098")
    assert record["lands"] == ["77:05:0004001:15", "77:05:0004001:40",
                               "77:05:0004001:2472", "77:05:0004001:3844"]


def test_a_public_owner_is_not_an_unnamed_one():
    """У города нет ни ИНН, ни ОГРН: разбор, ищущий их, выдал бы наш пробел
    за молчание документа."""
    owner = egrn_extracts.owner_of(_read("77:05:0004001:7"))
    assert owner and owner["kind"] == "public" and owner["name"] == "Москва"


def test_operational_management_is_not_ownership():
    """Собственник — город, держатель — ГБУ. Назвать учреждение собственником
    значит показать не то лицо: договариваться будут с разными."""
    record = _read("77:05:0004001:2078")
    assert egrn_extracts.owner_of(record)["name"] == "город Москва"
    others = egrn_extracts.other_rights(record)
    assert others and others[0]["right_type"] == "Оперативное управление"
    assert "Жилищник" in others[0]["name"]


def test_identity_is_the_inn_and_not_the_spelling():
    caps = {"kind": "legal", "name": 'ОБЩЕСТВО С ОГРАНИЧЕННОЙ ОТВЕТСТВЕННОСТЬЮ "УНИКС"',
            "inn": "9724179743", "ogrn": "1247700153667"}
    plain = {"kind": "legal", "name": 'Общество с ограниченной ответственностью "УНИКС"',
             "inn": "9724179743", "ogrn": "1247700153667"}
    assert egrn_extracts.holder_key(caps) == egrn_extracts.holder_key(plain)


def test_the_city_is_one_owner_however_it_is_spelled():
    """«Москва» и «город Москва» — одно лицо (владелец, 07.09.2026).

    В выписке по земле стоит субъект РФ 77 с кодом и именем «Москва», в
    выписке по зданию — тот же субъект без кода и с именем «город Москва».
    Ключ «по коду, а иначе по имени» разводил бы их на двоих ровно там, где
    надо свести: у второго кода нет вовсе.
    """
    with_code = {"kind": "public", "public_kind": "subject_of_rf",
                 "code": "77", "name": "Москва"}
    without = {"kind": "public", "public_kind": "subject_of_rf",
               "code": "", "name": "город Москва"}
    assert egrn_extracts.holder_key(with_code) == egrn_extracts.holder_key(without)
    # Разные субъекты при этом не слипаются.
    other = {"kind": "public", "public_kind": "subject_of_rf",
             "code": "50", "name": "Московская область"}
    assert egrn_extracts.holder_key(other) != egrn_extracts.holder_key(with_code)


def test_the_city_appears_once_in_the_summary():
    rows = [row for row in parcels.owners_summary()
            if egrn_extracts.public_name_key(row["name"]) == "москва"]
    assert len(rows) == 1, "город разошёлся на двух владельцев по написанию имени"
    assert rows[0]["lands"] == 4 and rows[0]["objects"] == 9
    assert rows[0]["name"] == "Москва", "показано имя без кода"


def test_no_two_codes_hide_under_one_public_name():
    """Ключ считается по имени, а код остаётся для сверки: два разных кода под
    одним именем — это столкновение, и оно обязано быть видно."""
    from collections import defaultdict

    codes = defaultdict(set)
    for path in EXTRACTS.glob("*.xml"):
        record = egrn_extracts.read(path.read_bytes())
        for right in record["rights"] + record["restrictions"]:
            for holder in right["holders"]:
                if holder.get("kind") == "public" and holder.get("code"):
                    codes[egrn_extracts.holder_key(holder)].add(holder["code"])
    clashes = {key: sorted(value) for key, value in codes.items() if len(value) > 1}
    assert not clashes, clashes


def test_an_unregistered_ownership_is_an_answer_not_a_gap():
    """У 14 участков из 20 собственность не зарегистрирована — это ответ."""
    lands = [egrn_extracts.read(path.read_bytes()) for path in EXTRACTS.glob("*.xml")]
    lands = [item for item in lands if item["kind"] == "land"]
    assert len(lands) == 20
    assert len([item for item in lands if egrn_extracts.owner_of(item) is None]) == 14


def test_the_lease_carries_its_term_and_tenant():
    leases = egrn_extracts.leases(_read("77:05:0004001:2475"))
    assert leases and any("НОВЫЙ ПРОЕКТ" in item["name"].upper() for item in leases)
    assert any(item["until"] or item["term"] for item in leases), "срок аренды не прочитан"


def test_a_foreign_document_is_refused_not_read_as_empty():
    with pytest.raises(ValueError):
        egrn_extracts.read("<html><body>не выписка</body></html>".encode("utf-8"))


# --- извещение о торгах ------------------------------------------------------

@pytest.fixture(scope="module")
def notice():
    return krt_notice.read(NOTICE)


def test_the_notice_names_the_composition_of_the_territory(notice):
    assert len(notice["lands"]) == 20
    assert len(notice["objects"]) == 39
    # Объект на нескольких участках повторяется строкой у каждого.
    assert notice["rows"] == 47


def test_an_object_on_several_parcels_is_counted_once(notice):
    found = {item["cadastral_number"]: item for item in notice["objects"]}
    assert len(found["77:05:0004001:2100"]["lands"]) == 2
    area = sum(item["area_sqm"] or 0 for item in notice["objects"])
    assert round(area, 1) == 54871.9, "строки сложены вместо объектов"


def test_a_word_in_the_number_column_is_not_a_number():
    """«часть» — это не 4,0 м². Подобранный похожий токен завысил бы итог."""
    assert krt_notice._number("часть") is None
    assert krt_notice._number("отсутствуют") is None
    assert krt_notice._number("2 522,8") == 2522.8


def test_a_missing_table_is_a_refusal_not_an_empty_territory(tmp_path):
    blank = tmp_path / "blank.pdf"
    import pymupdf

    document = pymupdf.open()
    document.new_page()
    document.save(blank)
    with pytest.raises(krt_notice.NoticeProblem):
        krt_notice.read(blank)


# --- свод --------------------------------------------------------------------

def test_the_territory_joins_the_notice_with_the_extracts():
    view = parcels.territory()
    totals = view["totals"]
    assert totals["lands"] == 20 and totals["objects"] == 39
    assert totals["rows_in_notice"] == 47
    # Земля и строения не складываются: это разные величины.
    assert totals["land_area_sqm"] != totals["objects_area_sqm"]
    assert round(totals["land_area_sqm"], 1) == 186860.1
    assert round(totals["objects_area_sqm"], 1) == 52381.0


def test_the_documents_disagree_and_it_is_named():
    """Выписка есть, а в извещении объекта нет — это ответ о составе."""
    view = parcels.territory()
    outside = [item["cadastral_number"] for item in view["objects_outside_notice"]]
    assert outside == ["77:05:0004001:2077"]
    without = [item for item in view["objects"] if not item["extract"]]
    assert [item["cadastral_number"] for item in without] == ["77:05:0004001:1951"]


def test_the_owners_are_folded_by_inn():
    rows = parcels.owners_summary()
    uniks = [row for row in rows if row["inn"] == "9724179743"]
    assert len(uniks) == 1, "одна компания в двух написаниях стала двумя владельцами"
    assert uniks[0]["objects"] == 10 and round(uniks[0]["objects_area_sqm"]) == 13832
    assert uniks[0]["group"] == "bryntsalov", "группа владельца не доехала"
    # Автокомбинат приходит и ЗАО, и АО — по ИНН это одно лицо.
    auto = [row for row in rows if row["inn"] == "7724020727"]
    assert len(auto) == 1 and auto[0]["lands"] == 2


def test_the_land_owner_is_shown_with_its_lease_not_instead_of_it():
    view = parcels.territory()
    land = next(item for item in view["lands"]
                if item["cadastral_number"] == "77:05:0004001:2475")
    assert land["owner"]["name"] == "", "у этого участка собственность не зарегистрирована"
    assert land["owner"]["note"]
    assert land["leases"], "аренда — единственное, что о нём известно, и она обязана быть"


# --- выгрузка ----------------------------------------------------------------

def test_the_workbook_is_built_from_the_same_numbers_as_the_screen():
    """Второй сборки нет: разойдясь, книга и страница дали бы два достоверных
    на вид ответа об одной территории."""
    import openpyxl
    from io import BytesIO

    from auction_search import nagatino_export

    view = parcels.territory()
    book = openpyxl.load_workbook(BytesIO(nagatino_export.build(view, parcels.owners_summary())))
    assert [sheet.title for sheet in book.worksheets] == [
        "ЗУ и объекты", "Кто чем владеет", "Источники"]
    sheet = book["ЗУ и объекты"]
    rows = list(sheet.iter_rows(values_only=True))[1:]
    lands = [row for row in rows if row[0] == "участок"]
    assert len(lands) == view["totals"]["lands"]


def test_the_column_adds_up_to_the_total_line():
    """Строка итога сходится с колонкой. Объект на нескольких участках стоит у
    каждого, и его площадь напечатана ОДИН раз: иначе сумма колонки (56 323,3)
    разошлась бы с итогом, и обе выглядели бы верными."""
    import openpyxl
    from io import BytesIO

    from auction_search import nagatino_export

    book = openpyxl.load_workbook(BytesIO(
        nagatino_export.build(parcels.territory(), parcels.owners_summary())))
    sheet = book["ЗУ и объекты"]
    # Колонку ищем по ЗАГОЛОВКУ, а не по номеру: номер держится за соседнюю
    # колонку и ломается, стоит рядом появиться новой.
    head = [cell.value for cell in sheet[1]]
    area = head.index("Площадь строения, м²")
    note = head.index("Примечание")
    rows = list(sheet.iter_rows(values_only=True))[1:]
    column = round(sum(row[area] for row in rows
                       if row[0] == "строение" and isinstance(row[area], (int, float))), 1)
    total = rows[-1]
    assert total[0] == "итого"
    assert round(total[area], 1) == column, "итог не сходится с колонкой"
    repeats = [row for row in rows if row[note] and "повтор" in str(row[note])]
    assert repeats and all(row[area] in (None, "") for row in repeats), \
        "у повтора напечатана площадь — она посчитается дважды"


def test_every_encumbrance_reaches_the_workbook_not_only_the_lease():
    """Ипотека Совкомбанка висит и на участке 77:05:0004001:2045, и на здании
    77:05:0004001:1046. Молча выброшенное ограничение читается как его
    отсутствие, а у залога это худшее из молчаний."""
    import openpyxl
    from io import BytesIO

    from auction_search import nagatino_export

    book = openpyxl.load_workbook(BytesIO(
        nagatino_export.build(parcels.territory(), parcels.owners_summary())))
    sheet = book["ЗУ и объекты"]
    head = [cell.value for cell in sheet[1]]
    column = head.index("Иные обременения (ипотека, ограничения)")
    number = head.index("Кадастровый номер")
    found = {str(row[number]).split()[0]: str(row[column])
             for row in sheet.iter_rows(values_only=True) if row[column]}
    assert "Ипотека" in found.get("77:05:0004001:2045", "")
    assert "Ипотека" in found.get("77:05:0004001:1046", "")
    assert "Совкомбанк" in found["77:05:0004001:1046"]
    assert "2034-05-02" in found["77:05:0004001:1046"], "срок залога не показан"


def test_a_term_without_a_date_is_named_not_printed_as_a_dangling_word():
    """У 77:05:0004001:1093 сама запись ЕГРН обрывается на слове «до».
    Печатать это как срок значит выдать обрыв документа за ответ."""
    from auction_search import nagatino_export

    text = nagatino_export._burden_text(
        [{"kind": "Аренда", "name": "физическое лицо", "until": "", "term": "до"}],
        with_kind=False)
    assert "срок в записи ЕГРН не указан" in text
    assert not text.rstrip().endswith("до")


def test_the_export_route_asks_the_owner(monkeypatch):
    from fastapi import FastAPI
    from fastapi.testclient import TestClient

    from auction_search.api import install

    monkeypatch.setenv("DEVELOPAID_ADMIN_KEY", "секрет")
    monkeypatch.setenv("DEVELOPAID_ADMIN_IDS", "1")
    app = FastAPI()
    install(app)
    client = TestClient(app)
    assert client.get("/krt/nagatino/export.xlsx").status_code == 403
    answer = client.get("/krt/nagatino/export.xlsx", params={"key": "секрет"})
    assert answer.status_code == 200
    assert answer.content[:2] == b"PK", "это не книга Excel"
    assert "attachment" in answer.headers.get("content-disposition", "")


def test_the_registry_answer_and_our_inference_never_share_a_cell():
    """«И твоё про „распоряжается город“ там тоже должно быть отмечено»
    (владелец, 07.09.2026) — отмечено, но СВОЕЙ графой.

    В клетке «Статус земли» стоит ответ ЕГРН, в соседней — наше суждение с его
    основанием. Слитые в одну клетку, они читались бы как одна запись реестра.
    """
    import openpyxl
    from io import BytesIO

    from auction_search import nagatino_export

    book = openpyxl.load_workbook(BytesIO(
        nagatino_export.build(parcels.territory(), parcels.owners_summary())))
    sheet = book["ЗУ и объекты"]
    head = [cell.value for cell in sheet[1]]
    status = head.index("Статус земли (запись ЕГРН)")
    who = head.index("Кто распоряжается — вывод DevelopAid")
    ground = head.index("На чём этот вывод")
    number = head.index("Кадастровый номер")
    rows = {str(row[number]).split()[0]: row for row in sheet.iter_rows(values_only=True)
            if row[0] == "участок"}
    free = rows["77:05:0004001:2475"]
    assert free[status] == "собственность в ЕГРН не зарегистрирована"
    assert "город" not in free[status], "вывод затесался в графу реестра"
    assert free[who] == "город Москва"
    assert "М-05-061753" in free[ground] and "вывод DevelopAid" in free[ground]
    # Там, где документов нет, слабее и утверждение.
    bare = rows["77:05:0004001:16"]
    assert bare[who] == "вероятно город Москва"
    assert "подтверждения в документах нет" in bare[ground]
    # У собственника вывода нет вовсе — распоряжается он сам.
    owned = rows["77:05:0004001:7"]
    assert owned[status].startswith("собственность: ") and not owned[who]


def test_the_city_contract_is_recognised_by_its_number():
    """Арендодателя выписка не называет — его выдаёт номер договора."""
    assert parcels.disposal_note(
        {"owner": {}, "leases": [{"document_number": "М-05-061753"}]})["who"] == "город Москва"
    assert parcels.disposal_note(
        {"owner": {}, "leases": [{"document_number": "4827-05 ДГИ"}]})["who"] == "город Москва"
    # Чужой номер городским не считается.
    private = parcels.disposal_note({"owner": {}, "leases": [{"document_number": "1731/К-ЗН-1/24"}]})
    assert "сдан в аренду" in private["ground"], "чужой номер выдан за городской"
