"""ГПР-факт по пакетам: укрупнённый ГПР с колонкой «Факт <дата>».

Владелец (29.10.2026) прислал «ГПР до ввода» — около 40 пакетов («Монолит
надземной части», «Кладка стен и перегородок», проектные этапы) с кодом РСС и
колонкой «Факт 20.10.2026» — и спросил, можно ли считать по нему выполнение;
отдельных форм загрузки — не плодить. Монитор ведёт выполнение по работам
утверждённого ГПР, поэтому файл грузится той же кнопкой «ГПР-факт», а здесь
процент пакета разносится на его работы:

- по коду РСС: работа с кодом статьи пакета или её подстатьи («2.2.2.6.» →
  2.2.2.6.1 …), «2.2.1.*» — вся ветка, «2.2.1.10 + 2.2.2.8» — обе статьи;
  к работе — пакет с самым длинным совпавшим кодом, при равных — с бо́льшим
  числом общих слов в названии, а если и названия не различают — не сопоставлено;
- пакеты без кода — по названию работы (общие слова), только однозначно;
- процент пакета — у всех его работ (укрупнённо: так и сказано в сводке),
  но не ниже процента самой работы в утверждённом ГПР;
  пустой «Факт» — «нет числа», работа остаётся как в утверждённом ГПР.

Что не сопоставилось, названо списком, а не пропадает молча.
"""

from __future__ import annotations

import datetime
import io
import re
from typing import Any

_NAME = "раздел / пакет"
_FACT = re.compile(r"^\s*факт\b\s*(\d{1,2}\.\d{1,2}\.\d{4})?", re.I)
_WORD = re.compile(r"[a-zа-яё0-9]+", re.I)
# Слова, по которым работы не различаются.
_STOP = {"работы", "работ", "устройство", "монтаж", "корп", "корпус", "и", "в", "на",
         "по", "для", "с", "до", "из", "от", "часть", "части"}


def _norm(value: Any) -> str:
    return re.sub(r"\s+", " ", str(value or "")).strip().lower()


def _codes(value: Any) -> list[str]:
    """«2.2.1.10 + 2.2.2.8» → ['2.2.1.10', '2.2.2.8'], «2.2.1.*» → ['2.2.1'], «—» → []."""
    out = []
    for part in re.split(r"[+,;]", str(value or "")):
        code = part.strip().rstrip("*").strip().rstrip(".")
        if re.fullmatch(r"\d+(\.\d+)*", code):
            out.append(code)
    return out


def _share(value: Any) -> float | None:
    if value is None or value == "":
        return None
    try:
        share = float(str(value).replace("%", "").replace(",", ".").strip())
    except ValueError:
        return None
    return share / 100.0 if share > 1.0 else share


def read(data: bytes) -> dict[str, Any] | None:
    """Пакеты из книги; не та книга (нет «Раздел / пакет» и «Факт …») — None."""
    from openpyxl import load_workbook

    try:
        book = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    except Exception:  # noqa: BLE001 — не книга Excel: это не наш формат
        return None
    try:
        for sheet in book.worksheets:
            found = _read_sheet(sheet)
            if found:
                return found
    finally:
        book.close()
    return None


def _read_sheet(sheet: Any) -> dict[str, Any] | None:
    rows = list(sheet.iter_rows(values_only=True, max_row=400))
    head, columns, fact_day = None, {}, ""
    for index, row in enumerate(rows[:12]):
        cells = [_norm(c) for c in row]
        if _NAME not in cells:
            continue
        for pos, cell in enumerate(cells):
            if cell == _NAME:
                columns["name"] = pos
            elif cell.startswith("код рсс"):
                columns["code"] = pos
            elif cell.startswith("стадия"):
                columns["stage"] = pos
            elif cell.startswith("статус"):
                columns["status"] = pos
            elif cell.startswith("прогноз"):
                columns.setdefault("forecast", pos)
            elif (match := _FACT.match(cell)) and "fact" not in columns:
                columns["fact"] = pos
                if match.group(1):
                    fact_day = datetime.datetime.strptime(match.group(1), "%d.%m.%Y").date() \
                        .isoformat()
        head = index
        break
    if head is None or "fact" not in columns:
        return None

    def cell(row: tuple, key: str) -> Any:
        pos = columns.get(key)
        return row[pos] if pos is not None and pos < len(row) else None

    packages, group = [], ""
    for row in rows[head + 1:]:
        first, name = row[0] if row else None, str(cell(row, "name") or "").strip()
        if first not in (None, "") and not name and not str(first).strip().isdigit():
            group = str(first).strip()  # «Этап 1. …», «СМР» — заголовок группы
            continue
        if not name or first in (None, ""):
            continue
        packages.append({"group": group, "name": name,
                         "codes": _codes(cell(row, "code")),
                         "code_text": str(cell(row, "code") or "").strip(),
                         "stage": str(cell(row, "stage") or "").strip(),
                         "status": str(cell(row, "status") or "").strip(),
                         "fact": _share(cell(row, "fact")),
                         "forecast": _share(cell(row, "forecast"))})
    if not packages:
        return None
    return {"sheet": sheet.title, "fact_day": fact_day, "packages": packages}


def _words(text: str) -> set[str]:
    """Основы слов: первые шесть букв — «металлоконструкции» и «…конструкций» одно."""
    return {w[:6] for w in _WORD.findall(_norm(text)) if len(w) > 2 and w not in _STOP}


def _shared(package: dict[str, Any], work: dict[str, Any]) -> int:
    """Сколько основ слов пакета есть в названии работы."""
    return len(_words(package["name"])
               & _words(f"{work.get('name') or ''} {work.get('section') or ''}"))


def _likeness(package: dict[str, Any], work: dict[str, Any]) -> float:
    """Доля основ слов пакета, которые есть в названии работы."""
    words = _words(package["name"])
    return _shared(package, work) / len(words) if words else 0.0


def _by_code(work: dict[str, Any], packages: list[dict[str, Any]]) -> list[tuple[int, int]]:
    """(длина совпавшего кода, номер пакета) — пакеты, чей код покрывает код работы."""
    code = str(work.get("estimate_code") or "").strip().rstrip(".")
    if not code:
        return []
    out = []
    for index, package in enumerate(packages):
        hits = [len(c.split(".")) for c in package["codes"]
                if code == c or code.startswith(c + ".")]
        if hits:
            out.append((max(hits), index))
    return out


def apply(packages: list[dict[str, Any]], works: list[dict[str, Any]],
          status_word: str = "Завершено") -> dict[str, Any]:
    """Проценты пакетов → работы утверждённого ГПР.

    Возвращает строки работ с процентом пакета (поверх утверждённого ГПР, как
    обычный ГПР-факт), какому пакету отнесена каждая работа и что не
    сопоставилось — пакеты без работ, работы с двумя равными пакетами.
    """
    chosen: dict[int, int] = {}
    ambiguous: list[dict[str, Any]] = []
    disputed: set[int] = set()
    for wi, work in enumerate(works):
        hits = _by_code(work, packages)
        if hits:
            best = max(depth for depth, _ in hits)
            tied = [pi for depth, pi in hits if depth == best]
            if len(tied) > 1:
                scored = sorted(((_shared(packages[pi], work), pi) for pi in tied),
                                reverse=True)
                if scored[0][0] > scored[1][0]:
                    tied = [scored[0][1]]
            if len(tied) == 1:
                chosen[wi] = tied[0]
            else:
                disputed.update(tied)
                ambiguous.append({"work": work.get("name") or "",
                                  "packages": [packages[pi]["name"] for pi in tied]})
            continue
        # Пакеты без кода — только по названию и только однозначно.
        scored = sorted(((_likeness(p, work), pi) for pi, p in enumerate(packages)
                         if not p["codes"]), reverse=True)
        if scored and scored[0][0] >= 0.6 and (len(scored) == 1 or scored[1][0] < scored[0][0]):
            chosen[wi] = scored[0][1]

    rows, used = [], {}
    for wi, pi in chosen.items():
        package = packages[pi]
        used.setdefault(pi, 0)
        used[pi] += 1
        if package["fact"] is None:
            continue
        work = works[wi]
        # Выполнение назад не идёт: работа, уже закрытая в утверждённом ГПР,
        # не «раскрывается» средним процентом пакета, где она одна из многих.
        before = _share(work.get("progress")) or 0.0
        done = "заверш" in package["status"].lower() or package["fact"] >= 0.999
        rows.append(dict(work, progress=max(before, package["fact"]),
                         status=status_word if done else (work.get("status") or ""),
                         package=package["name"]))
    unmatched = [p["name"] for pi, p in enumerate(packages)
                 if pi not in used and pi not in disputed and p["fact"] is not None]
    no_fact = [p["name"] for pi, p in enumerate(packages) if pi in used and p["fact"] is None]
    return {"works": rows, "packages": len(packages),
            "matched_packages": len([pi for pi in used if packages[pi]["fact"] is not None]),
            "unmatched": unmatched, "ambiguous": ambiguous, "no_fact": no_fact}
