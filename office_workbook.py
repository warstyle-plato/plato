"""Книга нежилого проекта на каркасе жилой v4: у офиса три способа продажи рядом, месяцы вправо.

Образец утвердил владелец 22.10.2026 (`docs/samples/office_book/`): «модель
должна быть такой же по философии и архитектуре, как для жилья». Цепочка та
же, что у жилой v4: вводные → сроки → затраты → продажи или аренда по месяцам
→ кредит → налоги → поток проекта и собственника → отчёт → сверка с движком.

Состав: Дашборд, ОТЧЕТ, Вводные, ИНСТРУКЦИЯ, ТЭП, СРОКИ, CAPEX, Ставки,
Продажа ДДУ, Продажа, Аренда и Выход, листы прочих объектов проекта, ПФ,
CF_ДДУ, CF_Продажа и CF_Аренда (НДС, кредит, поток собственника, NPV и IRR),
Налоги, ПРОВЕРКИ, Источники.

Три способа продажи офиса считаются формулами одновременно (владелец,
24.10.2026: «Новая книга считает и ДДУ»):
- продажа по ДДУ — `sales_schedule` (доля до ввода и остаток после, рост
  цены), эскроу до РВЭ и `simulate_financing` (собственные деньги и БРИДЖ до
  разрешения на стройку, ПФ с покрытием эскроу и ступенями специальной ставки,
  плата за лимит, раскрытие эскроу в РВЭ, погашение продажами), НДС по ДДУ в
  РВЭ и LLCR — построчный пересказ движка;
- продажа блоками после ввода и аренда с продажей здания — `object_flows`.
Каждый способ — расчёт движка с этой стратегией у офисов (`office_paths`).
Прочие объекты проекта (ТЦ и т. п.) считаются своей стратегией во всех трёх
способах; объект по ДДУ входит в ПФ каждого способа. CAPEX по
статьям — формулами `build_operating_model`, налог на прибыль проекта —
правилом `_profit_tax_schedule` (маржа до РВЭ откладывается, перенос убытка не
больше половины базы года), NPV проекта — `monthly_npv` ((1+r)^(k/12)), NPV
собственника — `equity_npv` (r/12).

Деньги — в млн ₽. Формулы — только те, что считают и Excel, и вычислитель
репозитория (`xlsx_eval`): книга пересчитывается им при сборке, значения
ложатся рядом с формулами (`v4_value_cache`), а «ПРОВЕРКИ» сравнивают итоги с
авторитетным расчётом.

Нежилой проект без офиса (ТЦ, ФОК, наземный паркинг) получает эту же книгу с
одним способом — стратегией каждого объекта (`PROJECT_PATH`): листы объектов,
ПФ при объекте по ДДУ, лист «CF», сверка с основным расчётом движка; сравнения
путей и листов способов офиса у неё нет (`docs/nonres_workbook.md`).
"""

from __future__ import annotations

import io
import re
import sys
from datetime import date, datetime
from typing import Any

from openpyxl import Workbook, load_workbook
from openpyxl.chart import BarChart, Reference
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter as L
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.worksheet.pagebreak import Break
from openpyxl.worksheet.properties import PageSetupProperties

import developaid_nonres_strategy as ns
import site_massing_book as massing_book
from v4_value_cache import with_cached_values
from xlsx_eval import Evaluator

# --- вид жилой v4 -------------------------------------------------------------
FONT = "Carlito"
DARK, HEAD, TILE, LIGHT, INPUT, GREY = "0B1F33", "1F4E78", "17365D", "EAF2F8", "FFF2CC", "666666"
SALE_TAB, RENT_TAB, DDU_TAB = "2F6DB5", "C26A0A", "2E7D32"
ENGINE = "FFF3E0"  # значения движка там, где формулы нет («СЧИТАЕТ ДВИЖОК»)
NUM = "#,##0.0;[Red](#,##0.0);-"
NUM0 = "#,##0;[Red](#,##0);-"
NUM3 = "#,##0.000;[Red](#,##0.000);-"
PCT = "0.0%;[Red]-0.0%;-"
PCT2 = "0.00%;[Red]-0.00%;-"
RATIO = "0.00"
INT = "0"
DATE = "dd.mm.yyyy"
MONTH = "mmm yyyy"
FLAG = '0;-0;"-"'
_thin = Side(style="thin", color="C9D3DE")

FIRST_COL = 4  # D: месяц k = 0
DATE_ROW, K_ROW = 3, 4  # шапка помесячных листов
EXTRA_MONTHS = 12  # год после горизонта — NOI покупателя для выхода
TAX_TAIL = 15  # доплата налога года — в марте следующего

INPUTS_SHEET, TEP_SHEET, TIME_SHEET, CAPEX_SHEET, RATE_SHEET = "Вводные", "ТЭП", "СРОКИ", "CAPEX", "Ставки"
SALE_SHEET, RENT_SHEET, EXIT_SHEET, DDU_SHEET, PF_SHEET = "Продажа", "Аренда", "Выход", "Продажа ДДУ", "ПФ"
TAX_SHEET, REPORT_SHEET, DASH_SHEET = "Налоги", "ОТЧЕТ", "Дашборд"
GUIDE_SHEET, CHECK_SHEET, SOURCES_SHEET = "ИНСТРУКЦИЯ", "ПРОВЕРКИ", "Источники"
PASSED, FAILED = "ПРОЙДЕНО", "РАСХОЖДЕНИЯ"
AGREE, DIFF = "сходится", "РАСХОЖДЕНИЕ"

# Способы продажи офиса: код движка → подпись, листы и цвет ярлыка. Порядок —
# как у `economics.OFFICE_PATHS`.
PATHS: tuple[dict[str, str], ...] = (
    {"key": ns.STRATEGY_DDU, "label": "Продажа по ДДУ", "code": "d",
     "flows": DDU_SHEET, "cf": "CF_ДДУ", "tab": DDU_TAB},
    {"key": ns.STRATEGY_DIRECT, "label": "Продажа блоками", "code": "s",
     "flows": SALE_SHEET, "cf": "CF_Продажа", "tab": SALE_TAB},
    {"key": ns.STRATEGY_INCOME, "label": "Аренда и продажа здания", "code": "r",
     "flows": RENT_SHEET, "cf": "CF_Аренда", "tab": RENT_TAB},
)
# Нежилой проект без офиса продаётся одним способом — стратегией, выбранной у
# каждого объекта (поручение владельца 24.10.2026: книга ТЦ, ФОК и наземного
# паркинга — на каркасе v4, без сравнения путей). Его потоки — лист «CF»,
# сверка — с основным расчётом движка.
PROJECT_PATH: dict[str, str] = {"key": "project", "label": "Стратегия объектов проекта", "code": "p",
                                "flows": "", "cf": "CF", "tab": HEAD}
STRATEGY = {ns.STRATEGY_INCOME: "Аренда и выход", ns.STRATEGY_DIRECT: "Продажа блоками",
            ns.STRATEGY_DDU: "Продажа по ДДУ"}
INTEREST_MODES = ("Капитализация в ПФ", "Оплата капиталом")
EXIT = {ns.EXIT_SALE: "Продажа здания", ns.EXIT_HOLD: "Удержание с оценкой"}
REPAY = {ns.REPAY_ANNUITY: "Аннуитет с баллоном", ns.REPAY_SWEEP: "Из NOI по мере поступления",
         ns.REPAY_BULLET: "Одним платежом при выходе"}
CURVE = {"flat": "Равномерно", "bell": "Колоколом", "front_loaded": "Больше в начале",
         "back_loaded": "Больше в конце"}
SCENARIO = {"low": "Низкий", "base": "Базовый", "high": "Высокий"}
YES, NO = "Да", "Нет"
SPACES = "места"

# Шкала класса офиса на «Вводных»: поле шкалы → (ключ книги, подпись, ед., формат,
# множитель к числу шкалы). Ставка аренды у шкалы — в год, у движка — в месяц.
CLASS_ROWS: tuple[tuple[str, str, str, str, str, float], ...] = (
    ("price_th_per_sqm", "price", "Цена продажи блоками", "тыс ₽/м² с НДС", NUM0, 1.0),
    ("rent_th_per_sqm_year", "rent_year", "Ставка аренды «всё включено»", "тыс ₽/м² в год", NUM, 1.0),
    ("occupancy_stable_pct", "occ_stable", "Заполняемость после заполнения", "", PCT, 0.01),
    ("leaseup_months", "leaseup", "Срок заполнения", "мес.", INT, 1.0),
    ("opex_pct", "opex", "Расходы эксплуатации, доля дохода аренды", "", PCT, 0.01),
    ("exit_cap_pct", "cap", "Ставка капитализации при продаже здания", "", PCT, 0.01),
)


def _month(value: Any) -> date:
    if isinstance(value, datetime):
        return value.date().replace(day=1)
    if isinstance(value, date):
        return value.replace(day=1)
    return date.fromisoformat(str(value)[:10]).replace(day=1)


def _months_between(start: date, when: date) -> int:
    return (when.year - start.year) * 12 + when.month - start.month


def _date_or_none(raw: Any) -> date | None:
    try:
        return date.fromisoformat(str(raw or "").strip()[:10])
    except ValueError:
        return None


def ref(name: str, mod: str = "") -> str:
    """Ссылка шаблона на строку книги.

    `{имя}` — та же колонка месяца (у вводной — её клетка), `{имя@p}` —
    прошлый месяц, `{имя@-3}` — три месяца назад, `{имя:}` — весь ряд,
    `{имя:k}` — ряд по этот месяц, `{имя:p}` — ряд по прошлый месяц.
    """
    return "{" + name + mod + "}"


_TOKEN = re.compile(r"\{([^{}@:]+)(@p|@n|@[+-]\d+|:k|:p|:)?\}")


def _fnt(**kw: Any) -> Font:
    return Font(name=FONT, size=11, **kw)


def _fill(color: str) -> PatternFill:
    return PatternFill("solid", fgColor=color)


class _Book:
    """Книга: листы, реестр строк и отложенные формулы (ряды ссылаются вперёд)."""

    def __init__(self, n: int, start: date) -> None:
        self.wb = Workbook()
        self.wb.remove(self.wb.active)
        self.n = n
        self.start = start
        self.reg: dict[str, tuple[str, int, str | None]] = {}
        self.pending: list[tuple[Any, str, int | None, str]] = []
        self.last = L(FIRST_COL + n - 1)
        # Способы продажи, которые книга считает: три у офиса, один у проекта без офиса.
        self.paths: tuple[dict[str, str], ...] = PATHS

    @property
    def project_only(self) -> bool:
        """Книга без офиса: один способ — стратегия объектов проекта."""
        return self.paths == (PROJECT_PATH,)

    # --- листы ---------------------------------------------------------------
    def sheet(self, name: str, title: str, sub: str = "", monthly: bool = False,
              tab: str | None = None, width: int = 8):
        ws = self.wb.create_sheet(name)
        ws.sheet_view.showGridLines = False
        if tab:
            ws.sheet_properties.tabColor = tab
        ws["A1"] = title
        ws["A1"].font = _fnt(b=True, color="FFFFFF")
        for c in range(1, (FIRST_COL + self.n if monthly else width + 1)):
            ws.cell(1, c).fill = _fill(DARK)
        ws["A2"] = sub
        ws["A2"].font = _fnt(color=GREY)
        ws.page_setup.orientation = "landscape"
        if monthly:
            for c, text in enumerate(("Показатель", "Итого / вводная", "Ед."), 1):
                cell = ws.cell(DATE_ROW, c, text)
                cell.font = _fnt(b=True, color="FFFFFF")
                cell.fill = _fill(HEAD)
            ws.cell(K_ROW, 1, "Месяц проекта, №").font = _fnt(color=DARK)
            for c in range(1, FIRST_COL):
                ws.cell(K_ROW, c).fill = _fill(LIGHT)
            for k in range(self.n):
                col = L(FIRST_COL + k)
                if name == TIME_SHEET:
                    d = ws.cell(DATE_ROW, FIRST_COL + k,
                                f"=EDATE({self.cell_ref('in.start', name)},{col}{K_ROW})")
                    i = ws.cell(K_ROW, FIRST_COL + k, 0 if k == 0 else f"={L(FIRST_COL + k - 1)}{K_ROW}+1")
                else:
                    d = ws.cell(DATE_ROW, FIRST_COL + k, f"='{TIME_SHEET}'!{col}{DATE_ROW}")
                    i = ws.cell(K_ROW, FIRST_COL + k, f"='{TIME_SHEET}'!{col}{K_ROW}")
                d.number_format = MONTH
                d.font = _fnt(b=True, color="FFFFFF")
                d.fill = _fill(HEAD)
                i.number_format = INT
                i.font = _fnt(color=DARK)
                i.fill = _fill(LIGHT)
                ws.column_dimensions[col].width = 11
            ws.column_dimensions["A"].width = 52
            ws.column_dimensions["B"].width = 16
            ws.column_dimensions["C"].width = 12
            ws.freeze_panes = "D5"
            # Печать — по 12 месяцев на лист с повтором подписей и шапки, а не
            # сжатием всех месяцев в ширину одной страницы.
            ws.sheet_properties.pageSetUpPr = PageSetupProperties(fitToPage=False)
            ws.print_title_cols = "A:C"
            ws.print_title_rows = f"{DATE_ROW}:{K_ROW}"
            ws.page_setup.scale = 55
            for k in range(12, self.n, 12):
                ws.col_breaks.append(Break(id=FIRST_COL - 1 + k))
        else:
            ws.sheet_properties.pageSetUpPr = PageSetupProperties(fitToPage=True)
            ws.page_setup.fitToWidth = 1
            ws.page_setup.fitToHeight = 0
        return ws

    # --- ссылки ----------------------------------------------------------------
    def register(self, name: str, sheet: str, row: int, col: str | None = None) -> None:
        if name in self.reg:
            raise ValueError(f"строка «{name}» уже есть в книге")
        self.reg[name] = (sheet, row, col)

    def cell_ref(self, name: str, here: str | None = None) -> str:
        sheet, row, col = self.reg[name]
        prefix = "" if sheet == here else f"'{sheet}'!"
        return f"{prefix}${col or 'B'}${row}"

    def expand(self, tpl: str, k: int | None, here: str) -> str:
        def sub(m: re.Match) -> str:
            name, mod = m.group(1), m.group(2) or ""
            if name not in self.reg:
                raise KeyError(f"в книге нет строки «{name}» (формула «{tpl[:80]}»)")
            sheet, row, col = self.reg[name]
            prefix = "" if sheet == here else f"'{sheet}'!"
            if col is not None:
                if mod:
                    raise ValueError(f"«{name}» — вводная, у неё нет месяцев")
                return f"{prefix}${col}${row}"
            first = L(FIRST_COL)
            if mod == ":":
                return f"{prefix}${first}${row}:${self.last}${row}"
            if k is None:
                raise ValueError(f"«{name}» — ряд месяцев: из клетки итога нужен диапазон")
            if mod == ":k":
                return f"{prefix}${first}${row}:{L(FIRST_COL + k)}{row}"
            if mod == ":p":
                return "0" if k == 0 else f"{prefix}${first}${row}:{L(FIRST_COL + k - 1)}{row}"
            shift = {"": 0, "@p": -1, "@n": 1}.get(mod)
            if shift is None:
                shift = int(mod[1:])
            j = k + shift
            if not 0 <= j < self.n:
                return "0"
            return f"{prefix}{L(FIRST_COL + j)}{row}"
        return _TOKEN.sub(sub, tpl)

    def resolve(self) -> None:
        for cell, tpl, k, here in self.pending:
            cell.value = "=" + self.expand(tpl, k, here)

    def to_bytes(self) -> bytes:
        stream = io.BytesIO()
        self.wb.save(stream)
        return stream.getvalue()


class _Rows:
    """Строки листа сверху вниз: разделы, вводные-итоги (колонка B) и ряды месяцев."""

    def __init__(self, book: _Book, ws, start: int = 6) -> None:
        self.book, self.ws, self.r = book, ws, start

    def section(self, text: str, color: str = DARK) -> None:
        self.r += 1
        cell = self.ws.cell(self.r, 1, text)
        cell.font = _fnt(b=True, color=color)
        for j in range(1, FIRST_COL):
            self.ws.cell(self.r, j).fill = _fill(LIGHT)
        self.r += 1

    def note(self, text: str) -> None:
        self.ws.cell(self.r, 1, text).font = _fnt(i=True, color=GREY)
        self.r += 1

    def param(self, name: str, label: str, value: Any, unit: str = "", fmt: str = NUM,
              inp: bool = False, bold: bool = False, note: str | None = None, col: str = "B") -> None:
        ws = self.ws
        ws.cell(self.r, 1, label).font = _fnt(b=bold)
        cell = ws[f"{col}{self.r}"]
        if isinstance(value, str) and value.startswith("="):
            self.book.pending.append((cell, value[1:], None, ws.title))
        else:
            cell.value = value
        cell.number_format = fmt
        cell.font = _fnt(b=bold)
        if inp:
            cell.fill = _fill(INPUT)
            cell.font = _fnt(color="0000FF")
        ws.cell(self.r, 3, unit).font = _fnt(color=GREY)
        if note:
            ws.cell(self.r, 4, note).font = _fnt(i=True, color=GREY)
        self.book.register(name, ws.title, self.r, col)
        self.r += 1

    def month(self, name: str, label: str, tpl: str | None = None, first: str | None = None,
              values: list[float] | None = None, unit: str = "млн ₽", fmt: str = NUM,
              total: bool = True, bold: bool = False, amount: str | None = None) -> None:
        """Ряд месяцев. `amount` — формула суммы в колонке B (строка «{имя}.amt»),
        от которой ряд раскладывается; иначе B — сумма ряда."""
        ws, book = self.ws, self.book
        ws.cell(self.r, 1, label).font = _fnt(b=bold, i=values is not None,
                                             color=GREY if values is not None else None)
        ws.cell(self.r, 3, unit).font = _fnt(color=GREY)
        book.register(name, ws.title, self.r)
        for k in range(book.n):
            cell = ws.cell(self.r, FIRST_COL + k)
            if values is not None:
                cell.value = float(values[k]) if k < len(values) else 0.0
                cell.fill = _fill(ENGINE)
                cell.font = _fnt(i=True, color=GREY)
            else:
                book.pending.append((cell, first if (k == 0 and first is not None) else tpl, k, ws.title))
                if bold:
                    cell.font = _fnt(b=True)
            cell.number_format = fmt
        if amount is not None:
            t = ws.cell(self.r, 2)
            book.pending.append((t, amount, None, ws.title))
            book.register(f"{name}.amt", ws.title, self.r, "B")
            t.number_format = fmt
            t.font = _fnt(b=True)
        elif total:
            t = ws.cell(self.r, 2, f"=SUM(D{self.r}:{book.last}{self.r})")
            t.number_format = fmt
            t.font = _fnt(b=True)
        for j in range(1, FIRST_COL + book.n):
            ws.cell(self.r, j).border = Border(bottom=_thin)
        self.r += 1


def _k_of(book: _Book, date_ref: str) -> str:
    """Номер месяца даты от старта проекта — формулой."""
    s = ref("in.start")
    return f"((YEAR({date_ref})-YEAR({s}))*12+MONTH({date_ref})-MONTH({s}))"


def _s_weight(k: str, start: str, months: str) -> str:
    """Вес месяца S-кривой книги ПЛАТО (`build_curve.monthly_weights`)."""
    i = f"({k}-{start}+1)"
    return (f"IF(AND({k}>={start},{k}<{start}+{months}),"
            f"IF({i}/{months}<=0.2,0.6,IF({i}/{months}<=0.8,1.2,0.8)),0)")


def _even_weight(k: str, start: str, months: str) -> str:
    return f"IF(AND({k}>={start},{k}<{start}+{months}),1,0)"


def _sum(parts: list[str]) -> str:
    return "+".join(parts) if parts else "0"


def grouped(value: str) -> str:
    """Целое число текстом с пробелом между тысячами — как подпись движка
    «не показателен: своих денег 1 012 млн ₽» (`_equity_irr_number`)."""
    r = f"ROUND({value},0)"
    return f'IF({r}>=1000,INT({r}/1000)&" "&REPT("0",3-LEN(MOD({r},1000)))&MOD({r},1000),{r})'


# --- «Вводные» ---------------------------------------------------------------

# Проект: ключ → (подпись, ед., формат, список значений или None). Значение —
# колонка B. Проценты хранятся долями и показываются процентом.
PROJECT_INPUTS: tuple[tuple[str, str, str, str], ...] = (
    ("start", "Начало проекта", "", DATE),
    ("ird_months", "Документы до разрешения на стройку (ИРД)", "мес.", INT),
    ("ird_min", "Минимальный срок ИРД", "мес.", INT),
    ("construction_months", "Стройка основной части: от разрешения до ввода", "мес.", INT),
    ("residual_months", "Остаточные продажи после ввода (для горизонта)", "мес.", INT),
    ("vat", "НДС", "", PCT),
    ("profit_tax", "Налог на прибыль", "", PCT),
    ("loss_limit", "Убыток прошлых лет гасит не больше этой доли базы года", "", PCT),
    ("discount", "Ставка дисконтирования", "в год", PCT),
    ("revenue_mult", "Сценарий: множитель выручки", "×", "0.00"),
    ("cost_mult", "Сценарий: множитель затрат", "×", "0.00"),
    ("rate_scenario", "Сценарий ключевой ставки", "", "@"),
    ("marketing", "Маркетинг при продаже блоками", "доля выручки", PCT),
    ("selling", "Брокеры при продаже блоками", "доля выручки", PCT),
    ("parking_price_th", "Цена машино-места проекта по умолчанию", "тыс ₽", NUM0),
)
DEAL_INPUTS: tuple[tuple[str, str, str, str], ...] = (
    ("purchase_price", "Цена участка", "млн ₽", NUM),
    ("land_buyout", "Выкуп участков и зданий у третьих лиц", "млн ₽", NUM),
    ("land_rights_gross", "Плата за смену вида разрешённого использования", "млн ₽", NUM),
    ("land_rights_relief", "Льгота и зачёт по этой плате", "млн ₽", NUM),
    ("reservation_fee", "Плата банку за резервирование лимита кредита", "доля лимита", PCT2),
    ("limit_fee", "Плата банку за невыбранный лимит", "в год", PCT2),
)
COST_INPUTS: tuple[tuple[str, str, str, str], ...] = (
    ("core_above_gns", "Наземная площадь жилого ядра (у нежилого проекта — 0)", "м²", NUM0),
    ("core_under_gns", "Подземная площадь ядра: паркинг и кладовые", "м²", NUM0),
    ("ird_th", "ИРД", "тыс ₽/м² базы", NUM),
    ("design_p_th", "Проект (стадия П)", "тыс ₽/м² базы", NUM),
    ("design_rd_th", "Рабочая документация", "тыс ₽/м² базы", NUM),
    ("author_supervision", "Авторский надзор", "доля П + РД", PCT),
    ("demolition_area", "Площадь сноса", "м²", NUM0),
    ("demolition_th", "Снос", "тыс ₽/м² сноса", NUM),
    ("resettlement_mln", "Расселение", "млн ₽", NUM),
    ("preparation_th", "Подготовка территории", "тыс ₽/м² базы", NUM),
    ("main_above_th", "Стройка наземной части ядра", "тыс ₽/м²", NUM),
    ("main_under_th", "Стройка подземной части и гаражей объектов", "тыс ₽/м²", NUM),
    ("utilities_th", "Сети", "тыс ₽/м² базы", NUM),
    ("commissioning_th", "Ввод в эксплуатацию", "тыс ₽/м² базы", NUM),
    ("site_maintenance_th", "Содержание площадки", "тыс ₽/м² базы", NUM),
    ("project_management", "Управление проектом", "доля базы управления", PCT),
    ("technical_supervision", "Технический заказчик", "доля стройки", PCT),
    ("gc_fee", "Генподряд", "доля стройки", PCT),
    ("reserve", "Резерв", "доля статей до резерва", PCT),
)
# Объект: ключ → (подпись, ед., формат). Колонки B, C, … — объекты.
OBJECT_BLOCKS: tuple[tuple[str, tuple[tuple[str, str, str, str], ...]], ...] = (
    ("Объекты: площади и стройка", (
        ("measure", "Мера объекта", "м² / места", "@"),
        ("volume", "Наземная площадь здания (наземный паркинг — мест)", "м² / мест", NUM0),
        ("total_area", "Общая площадь", "м²", NUM0),
        ("saleable", "Продаваемая и арендопригодная площадь", "м²", NUM0),
        ("rate_cost", "Стоимость стройки здания (места — млн ₽ за место)", "тыс ₽/м²", NUM),
        ("turnkey", "Ставка «под ключ»: статьи проекта, генподряд и техзаказчик уже в ней", "", "@"),
        ("under_gns", "Подземный гараж объекта", "м²", NUM0),
        ("start", "Начало стройки", "", DATE),
        ("months", "Срок стройки", "мес.", INT),
        ("in_tax_pool", "Затраты объекта признаются его продажами", "", "@"),
        ("parking_units", "Машино-места в продаже или аренде (без гостевых)", "шт.", NUM0),
        ("parking_under_units", "Мест в подземном гараже", "шт.", NUM0),
        ("over_units", "Машино-места на первых этажах", "шт.", NUM0),
        ("over_area", "Площадь места на первом этаже", "м²", NUM0),
    )),
    ("Цена продажи и её рост (общие для продажи по ДДУ и продажи блоками)", (
        ("price", "Цена продажи (у офиса пусто — по классу)", "тыс ₽/м² с НДС", NUM0),
        ("sales_start", "Отсчёт роста цены", "", DATE),
        ("growth_pre", "Рост цены до ввода", "в месяц", PCT2),
        ("growth_post", "Рост цены после ввода", "в месяц", PCT2),
        ("parking_under_mln", "Цена подземного места (0 — цена проекта)", "млн ₽", NUM),
        ("parking_over_mln", "Цена места на первом этаже (0 — цена проекта)", "млн ₽", NUM),
    )),
    ("Продажа по ДДУ: деньги покупателей на эскроу до ввода проекта", (
        ("share_before", "Доля продаж до ввода объекта", "", PCT),
        ("residual", "Остаточные продажи после ввода объекта", "мес.", INT),
    )),
    ("Продажа блоками после ввода по договорам купли-продажи", (
        ("sale_offset", "Старт продаж от ввода (минус — до ввода)", "мес.", INT),
        ("sale_months", "Срок продаж", "мес.", INT),
        ("curve", "Профиль продаж", "", "@"),
    )),
    ("Аренда и продажа здания", (
        ("rent", "Ставка аренды с НДС (у офиса пусто — по классу)", "тыс ₽/м² в месяц", NUM3),
        ("rent_date", "Дата ставки аренды (пусто — на дату ввода)", "", DATE),
        ("rent_index", "Индексация аренды", "в год", PCT),
        ("parking_rent", "Аренда машино-места с НДС", "тыс ₽ в месяц", NUM),
        ("occ_start", "Заполняемость на открытии", "", PCT),
        ("occ_stable", "Заполняемость после заполнения (у офиса пусто — по классу)", "", PCT),
        ("leaseup", "Срок заполнения (у офиса пусто — по классу)", "мес.", INT),
        ("opex", "Расходы эксплуатации (у офиса пусто — по классу)", "доля дохода", PCT),
        ("property_tax", "Налог на имущество", "в год от стоимости без НДС", PCT),
        ("hold_years", "Срок удержания", "лет", INT),
        ("exit_mode", "Выход в конце удержания", "", "@"),
        ("cap", "Ставка капитализации (у офиса пусто — по классу)", "", PCT),
        ("exit_cost", "Расходы на сделку при продаже здания", "", PCT),
        ("dep_years", "Срок амортизации для налога", "лет", INT),
    )),
    ("Кредит объекта при продаже блоками и аренде", (
        ("loan_share", "Доля кредита в затратах", "", PCT),
        ("spread", "Надбавка к ключевой ставке", "", PCT2),
        ("fee", "Комиссия за выдачу", "доля выдачи", PCT2),
        ("repayment", "Погашение после ввода (при аренде)", "", "@"),
        ("term", "Срок кредита от первой выдачи (при аренде)", "лет", INT),
        ("balloon", "Остаток долга к сроку — баллон (при аренде)", "доля долга на ввод", PCT),
    )),
)
# Поля, которые у офиса задаёт класс: ключ книги → поле шкалы.
CLASS_FIELDS = {"price": "price_th_per_sqm", "rent": "rent_th_per_sqm_year",
                "occ_stable": "occupancy_stable_pct", "leaseup": "leaseup_months",
                "opex": "opex_pct", "cap": "exit_cap_pct"}
LISTS = {"measure": ("м²", SPACES), "turnkey": (YES, NO), "in_tax_pool": (YES, NO),
         "curve": tuple(CURVE.values()), "exit_mode": tuple(EXIT.values()),
         "repayment": tuple(REPAY.values())}


def _object_values(item: dict[str, Any]) -> dict[str, Any]:
    """Вводные объекта в единицах книги: проценты — долями."""
    params = item.get("params") or {}
    values: dict[str, Any] = {
        "measure": SPACES if item["measure"] == "spaces" else "м²",
        "volume": item["volume"], "total_area": item.get("total_area", 0.0),
        "rate_cost": item["rate_cost"], "turnkey": YES if item["turnkey"] else NO,
        "under_gns": item["under_gns"], "start": item["start"], "months": item["months"],
        "in_tax_pool": YES if item["in_tax_pool"] else NO,
    }
    if item.get("nonres") or item.get("ddu"):
        values.update({
            "saleable": item["saleable_sqm"], "over_units": item["over_units"],
            "over_area": item["over_area"], "parking_units": item["parking_saleable_units"],
            "parking_under_units": item["parking_under_units"],
            "price": item["price"], "sales_start": item["sales_start"],
            "growth_pre": float(item["growth_pre_pct"]) / 100,
            "growth_post": float(item["growth_post_pct"]) / 100,
            "parking_under_mln": item["parking_under_mln"], "parking_over_mln": item["parking_over_mln"],
        })
    if "share_before_pct" in item:
        values.update({"share_before": float(item["share_before_pct"]) / 100,
                       "residual": int(item["residual_months"])})
    if params:
        pct = lambda key: float(params[key]) / 100  # noqa: E731
        values.update({
            "sale_offset": params["direct_sale_offset_months"],
            "sale_months": params["direct_sale_months"], "curve": CURVE[params["direct_sale_curve"]],
            "rent": params["rent_th_per_sqm_month"], "rent_date": _date_or_none(params.get("rent_date")),
            "rent_index": pct("rent_index_pct"), "parking_rent": params["parking_rent_th_month"],
            "occ_start": pct("occupancy_start_pct"), "occ_stable": pct("occupancy_stable_pct"),
            "leaseup": params["leaseup_months"], "opex": pct("opex_pct"),
            "property_tax": pct("property_tax_pct"), "hold_years": params["hold_years"],
            "exit_mode": EXIT[params["exit_mode"]], "cap": pct("exit_cap_pct"),
            "exit_cost": pct("exit_cost_pct"), "dep_years": params["depreciation_years"],
            "loan_share": pct("loan_share_pct"), "spread": pct("loan_spread_pp"),
            "fee": pct("loan_fee_pct"), "repayment": REPAY[params["debt_repayment"]],
            "term": params["loan_term_years"], "balloon": pct("loan_balloon_pct"),
        })
    return values


def _class_value(office_class: str, field: str) -> float:
    """Число шкалы класса в единицах книги: аренда — в месяц, как у движка."""
    scale = ns.OFFICE_CLASS_SCALE[office_class]
    if field == "rent_th_per_sqm_year":
        return ns.office_class_fields(office_class)["rent_th_per_sqm_month"]
    multiplier = next(m for f, _, _, _, _, m in CLASS_ROWS if f == field)
    return float(scale[field]) * multiplier


def _same(a: Any, b: Any) -> bool:
    try:
        return abs(float(a) - float(b)) < 1e-9
    except (TypeError, ValueError):
        return False


def office_class_of(values: dict[str, Any]) -> tuple[str, dict[str, Any]]:
    """Класс офиса, к которому ближе вводные, и поля, заданные вручную.

    Класс в проекте пока не хранится (выбор класса появится вместе со
    страницей): берётся ступень, с которой совпало больше полей, при равенстве
    — та, чья цена совпала, затем младшая. Поле, отличное от ступени, уходит в
    «Вручную»: правка руками сильнее класса и видна.
    """
    def score(cls: str) -> tuple[int, int, int]:
        hits = sum(_same(values.get(key), _class_value(cls, field)) for key, field in CLASS_FIELDS.items())
        price = int(_same(values.get("price"), _class_value(cls, "price_th_per_sqm")))
        return hits, price, -ns.OFFICE_CLASSES.index(cls)
    best = max(ns.OFFICE_CLASSES, key=score)
    manual = {key: values.get(key) for key, field in CLASS_FIELDS.items()
              if not _same(values.get(key), _class_value(best, field))}
    return best, manual


class _Inputs:
    def __init__(self) -> None:
        self.columns: dict[str, str] = {}   # код объекта → колонка «Вводных»
        self.plan = ("", "")                # диапазоны графика оплаты участка: k и доля
        self.rates: dict[str, str] = {}     # «С даты» и сценарии → диапазоны
        self.class_head = ""                # шапка шкалы класса — для MATCH
        self.class_rows: dict[str, str] = {}  # поле шкалы → диапазон ступеней
        self.office_class: dict[str, str] = {}  # код офиса → выбранный класс
        self.steps: list[tuple[str, str]] = []  # ступени специальной ставки ПФ: (порог, ставка)


def _inputs_sheet(book: _Book, spec: dict[str, Any], objects: list[dict[str, Any]]) -> _Inputs:
    if book.project_only:
        ws = book.sheet(INPUTS_SHEET, "DEVELOPAID · ВВОДНЫЕ НЕЖИЛОГО ПРОЕКТА",
                        "Печатают только здесь: жёлтые клетки синим шрифтом. Каждый объект продаётся "
                        "стратегией, выбранной у него на странице проекта.", tab=DARK, width=8)
    else:
        ws = book.sheet(INPUTS_SHEET, "DEVELOPAID · ВВОДНЫЕ ОФИСНОГО ПРОЕКТА",
                        "Печатают только здесь: жёлтые клетки синим шрифтом. Класс офиса задаёт допущения "
                        "аренды и продажи; число, вписанное у офиса вручную, сильнее класса.",
                        tab=DARK, width=8)
    ws.column_dimensions["A"].width = 64
    for c in "BCDEFG":
        ws.column_dimensions[c].width = 17
    ws.column_dimensions["H"].width = 40
    info = _Inputs()
    rows = _Rows(book, ws, 3)
    cost = spec.get("cost") or {}
    project = {
        "start": spec["start"], "ird_months": spec["ird_months"], "ird_min": spec["ird_min"],
        "construction_months": spec["construction_months"], "residual_months": spec["residual_months"],
        "vat": spec["vat_pct"] / 100, "profit_tax": spec["profit_tax_pct"] / 100,
        "loss_limit": spec["loss_limit"], "discount": spec["discount_pct"] / 100,
        "revenue_mult": spec["revenue_mult"], "cost_mult": spec["cost_mult"],
        "rate_scenario": SCENARIO[spec["rate_scenario"]],
        "marketing": spec["marketing_pct"] / 100, "selling": spec["selling_pct"] / 100,
        "parking_price_th": spec["parking_price_th"],
    }
    rows.section("Проект")
    rows.param("in.name", "Проект", spec.get("name") or "", fmt="@", inp=True)
    for key, label, unit, fmt in PROJECT_INPUTS:
        rows.param(f"in.{key}", label, project[key], unit, fmt, inp=True)
    scenario = DataValidation(type="list", formula1='"' + ",".join(SCENARIO.values()) + '"')
    ws.add_data_validation(scenario)
    scenario.add(f"B{book.reg['in.rate_scenario'][1]}")

    deal = {"purchase_price": spec["purchase_price_mln"], "land_buyout": spec["land_buyout_mln"],
            "land_rights_gross": spec["land_rights_gross_mln"],
            "land_rights_relief": spec["land_rights_relief_mln"],
            "reservation_fee": spec.get("reservation_fee_pct", 0.0) / 100,
            "limit_fee": spec.get("limit_fee_pct", 0.0) / 100}
    rows.section("Сделка, участок и банк")
    for key, label, unit, fmt in DEAL_INPUTS:
        rows.param(f"in.{key}", label, deal[key], unit, fmt, inp=True)
    rows.r += 1
    ws.cell(rows.r, 1, "График оплаты участка и выкупа: месяц проекта и доля цены").font = _fnt(b=True)
    ws.cell(rows.r, 2, "Месяц, №").font = _fnt(b=True)
    ws.cell(rows.r, 3, "Доля").font = _fnt(b=True)
    rows.r += 1
    plan = list(spec.get("purchase_plan") or []) or [(0, 1.0)]
    first = rows.r
    for k, share in plan:
        for c, value, fmt in (("B", int(k), INT), ("C", float(share), PCT)):
            cell = ws[f"{c}{rows.r}"]
            cell.value, cell.number_format = value, fmt
            cell.fill, cell.font = _fill(INPUT), _fnt(color="0000FF")
        rows.r += 1
    info.plan = (f"'{INPUTS_SHEET}'!$B${first}:$B${rows.r - 1}",
                 f"'{INPUTS_SHEET}'!$C${first}:$C${rows.r - 1}")

    # Проектное финансирование продажи по ДДУ — те же вводные, что читает
    # `simulate_financing`.
    pf = spec.get("pf") or {}
    rows.section("Продажа по ДДУ: собственные деньги, БРИДЖ и проектное финансирование (ПФ)")
    for key, label, unit, fmt, value in (
            ("bridge_spread", "БРИДЖ до разрешения на стройку: надбавка к ключевой ставке", "", PCT2,
             pf.get("bridge_spread_pp", 0.0) / 100),
            ("bridge_cap_spread", "БРИДЖ: надбавка к ключевой ставке на проценты к уплате", "", PCT2,
             pf.get("bridge_cap_spread_pp", 0.0) / 100),
            ("pf_spread", "ПФ: надбавка к ключевой ставке на часть долга без покрытия эскроу", "", PCT2,
             pf.get("pf_spread_pp", 0.0) / 100),
            ("pf_special", "ПФ: льготная ставка на часть долга, покрытую эскроу", "в год", PCT2,
             pf.get("pf_special_pct", 0.0) / 100),
            ("own_funds", "Собственные деньги до ПФ (тратятся раньше БРИДЖа)", "млн ₽", NUM,
             pf.get("own_funds_mln", 0.0)),
            ("approved", "Одобренный банком лимит ПФ (0 — лимит по выдаче)", "млн ₽", NUM,
             pf.get("approved_limit_mln", 0.0))):
        rows.param(f"in.pf.{key}", label, value, unit, fmt, inp=True)
    rows.param("in.pf.mode", "Проценты БРИДЖа при открытии ПФ", pf.get("interest_mode", INTEREST_MODES[0]),
               fmt="@", inp=True)
    modes = DataValidation(type="list", formula1='"' + ",".join(INTEREST_MODES) + '"')
    ws.add_data_validation(modes)
    modes.add(f"B{book.reg['in.pf.mode'][1]}")
    rows.r += 1
    ws.cell(rows.r, 1, "ПФ: ступени льготной ставки по покрытию долга эскроу (порог 0 — ступени нет)").font = _fnt(
        b=True)
    ws.cell(rows.r, 2, "Покрытие от").font = _fnt(b=True)
    ws.cell(rows.r, 3, "Ставка").font = _fnt(b=True)
    rows.r += 1
    for edge, rate in pf.get("pf_special_steps") or []:
        for c, value, fmt in (("B", float(edge), PCT), ("C", float(rate), PCT2)):
            cell = ws[f"{c}{rows.r}"]
            cell.value, cell.number_format = value, fmt
            cell.fill, cell.font = _fill(INPUT), _fnt(color="0000FF")
        info.steps.append((f"'{INPUTS_SHEET}'!$B${rows.r}", f"'{INPUTS_SHEET}'!$C${rows.r}"))
        rows.r += 1
    if not info.steps:
        rows.note("Ступеней нет: на покрытую часть действует льготная ставка выше.")

    # Шкала класса офиса — стартовые допущения; правятся тоже здесь.
    offices = [item for item in objects if item.get("office")]
    if offices:
        rows.section("Шкала класса офиса (Москва, с НДС; аренда — на 12.2025, до ввода растёт на индексацию)")
        head = rows.r
        for j, text in enumerate(("Допущение", *ns.OFFICE_CLASSES, "Ед.")):
            cell = ws.cell(head, 1 + j, text)
            cell.font = _fnt(b=True, color="FFFFFF")
            cell.fill = _fill(HEAD)
        info.class_head = f"'{INPUTS_SHEET}'!$B${head}:$E${head}"
        rows.r += 1
        for field, key, label, unit, fmt, mult in CLASS_ROWS:
            ws.cell(rows.r, 1, label)
            for j, cls in enumerate(ns.OFFICE_CLASSES):
                cell = ws.cell(rows.r, 2 + j, float(ns.OFFICE_CLASS_SCALE[cls][field]) * mult)
                cell.number_format, cell.fill, cell.font = fmt, _fill(INPUT), _fnt(color="0000FF")
            ws.cell(rows.r, 6, unit).font = _fnt(color=GREY)
            info.class_rows[field] = f"'{INPUTS_SHEET}'!$B${rows.r}:$E${rows.r}"
            rows.r += 1
        rows.note("Ставка аренды в расчёте — в месяц: год / 12 с точностью до 0,001 тыс ₽, как в движке. "
                  "Источники — на листе «Источники».")

    # Объекты — колонками.
    for i, item in enumerate(objects):
        info.columns[item["code"]] = L(2 + i)
    validations = {}
    for key, options in LISTS.items():
        dv = DataValidation(type="list", formula1='"' + ",".join(options) + '"', allow_blank=True)
        ws.add_data_validation(dv)
        validations[key] = dv
    class_list = DataValidation(type="list", formula1='"' + ",".join(ns.OFFICE_CLASSES) + '"')
    ws.add_data_validation(class_list)
    values = {item["code"]: _object_values(item) for item in objects}
    for item in offices:
        cls, manual = office_class_of(values[item["code"]])
        info.office_class[item["code"]] = cls
        for key in CLASS_FIELDS:
            values[item["code"]][key] = manual.get(key)
    for title, block in OBJECT_BLOCKS:
        rows.section(title)
        ws.cell(rows.r, 1, "Объект").font = _fnt(b=True, color="FFFFFF")
        ws.cell(rows.r, 1).fill = _fill(HEAD)
        for item in objects:
            cell = ws[f"{info.columns[item['code']]}{rows.r}"]
            cell.value = item["title"]
            cell.font, cell.fill = _fnt(b=True, color="FFFFFF"), _fill(HEAD)
        rows.r += 1
        if title.startswith("Объекты"):
            ws.cell(rows.r, 1, "Как продаётся в книге")
            for item in objects:
                text = ("все три способа: продажа по ДДУ, продажа блоками и аренда" if item.get("office") else
                        STRATEGY.get(item.get("strategy") or "", "без продажи (только стройка)")
                        + " — меняется на странице проекта")
                ws[f"{info.columns[item['code']]}{rows.r}"] = text
                ws[f"{info.columns[item['code']]}{rows.r}"].alignment = Alignment(wrap_text=True)
            rows.r += 1
            if offices:
                ws.cell(rows.r, 1, "Класс офиса (шкала выше)")
                for item in offices:
                    cell = ws[f"{info.columns[item['code']]}{rows.r}"]
                    cell.value = info.office_class[item["code"]]
                    cell.fill, cell.font = _fill(INPUT), _fnt(color="0000FF")
                    class_list.add(cell.coordinate)
                    book.register(f"in.{item['code']}.class", INPUTS_SHEET, rows.r, info.columns[item["code"]])
                rows.r += 1
        for key, label, unit, fmt in block:
            ws.cell(rows.r, 1, label)
            ws.cell(rows.r, 2 + len(objects), unit).font = _fnt(color=GREY)
            for item in objects:
                column = info.columns[item["code"]]
                if key not in values[item["code"]]:
                    continue
                cell = ws[f"{column}{rows.r}"]
                cell.value = values[item["code"]][key]
                cell.number_format = fmt
                cell.fill, cell.font = _fill(INPUT), _fnt(color="0000FF")
                if key in validations:
                    validations[key].add(cell.coordinate)
                book.register(f"in.{item['code']}.{key}", INPUTS_SHEET, rows.r, column)
            rows.r += 1

    rows.section("Смета: ставки статей проекта")
    for key, label, unit, fmt in COST_INPUTS:
        value = spec.get(key) if key in ("core_above_gns", "core_under_gns", "main_under_th") else None
        if value is None:
            raw = cost.get(key if key.endswith(("_th", "_mln", "_area")) else f"{key}_pct", 0.0)
            value = float(raw) / 100 if fmt == PCT else raw
        rows.param(f"in.{key}", label, value, unit, fmt, inp=True)
    rows.section("Пятно застройки и подземная часть: множитель стройки за глубину")
    massing = spec.get("massing") or {}
    for key, label, kind in massing_book.ROWS:
        if kind == massing_book.INPUT and key != "underground_sqm":
            rows.param(f"ms.{key}", label, massing.get(key, 0.0), fmt="#,##0.##", inp=True)

    rows.section("Прогноз ключевой ставки ЦБ: действует с даты до следующей строки")
    head = rows.r
    for j, text in enumerate(("С даты", *SCENARIO.values())):
        cell = ws.cell(head, 2 + j, text)
        cell.font, cell.fill = _fnt(b=True, color="FFFFFF"), _fill(HEAD)
    rows.r += 1
    first = rows.r
    for rate in spec.get("rate_table") or []:
        ws[f"B{rows.r}"] = rate["date"]
        ws[f"B{rows.r}"].number_format = DATE
        for c, key in zip("CDE", ("low", "base", "high")):
            cell = ws[f"{c}{rows.r}"]
            cell.value, cell.number_format = float(rate[key]) / 100, PCT2
            cell.fill, cell.font = _fill(INPUT), _fnt(color="0000FF")
        rows.r += 1
    last = max(first, rows.r - 1)
    info.rates = {"dates": f"'{INPUTS_SHEET}'!$B${first}:$B${last}",
                  **{label: f"'{INPUTS_SHEET}'!${c}${first}:${c}${last}"
                     for label, c in zip(SCENARIO.values(), "CDE")}}
    ws.freeze_panes = "B3"
    return info


# --- «СРОКИ», «ТЭП», «Ставки» --------------------------------------------------

def _time_sheet(book: _Book, objects: list[dict[str, Any]]) -> None:
    ws = book.sheet(TIME_SHEET, "СРОКИ · КАЛЕНДАРЬ ПРОЕКТА",
                    "Номер месяца, разрешение на стройку, ввод объектов. Окна продаж и аренды — на листах способов продажи.",
                    monthly=True, tab=HEAD)
    book.register("date", TIME_SHEET, DATE_ROW)
    book.register("k", TIME_SHEET, K_ROW)
    rows = _Rows(book, ws, 5)
    rows.section("Этапы проекта")
    rows.param("t.ird_eff", "Срок документов до разрешения с учётом минимума",
               f"=MAX({ref('in.ird_min')},INT({ref('in.ird_months')}))", "мес.", INT)
    rows.param("t.permit", "Разрешение на стройку, № месяца", f"={ref('t.ird_eff')}", "№", INT)
    rows.param("t.build", "Срок стройки основной части",
               f"=MAX(1,INT({ref('in.construction_months')}))", "мес.", INT)
    rows.param("t.rve", "Ввод проекта (РВЭ), № месяца",
               f"={ref('t.permit')}+INT({ref('in.construction_months')})", "№", INT)
    rows.param("t.window", "Окно проектирования до разрешения", f"=MIN(6,{ref('t.ird_eff')})", "мес.", INT)
    rows.section("Объекты")
    for item in objects:
        o = item["code"]
        rows.param(f"t.{o}.start", f"{item['title']}: начало стройки, № месяца",
                   f"={_k_of(book, ref(f'in.{o}.start'))}", "№", INT)
        rows.param(f"t.{o}.comm", f"{item['title']}: ввод, № месяца",
                   f"={ref(f't.{o}.start')}+INT({ref(f'in.{o}.months')})", "№", INT)
    flows = [item for item in objects if item.get("nonres")]
    rows.param("t.comm_last", "Ввод последнего объекта с продажей или арендой, № месяца",
               "=MAX(" + ",".join(ref(f"t.{item['code']}.comm") for item in flows) + ")"
               if flows else "=0", "№", INT)
    rows.section("По месяцам")
    rows.month("year", "Календарный год", f"YEAR({ref('date')})", unit="год", fmt=INT, total=False)
    rows.month("moy", "Месяц года", f"MONTH({ref('date')})", unit="№", fmt=INT, total=False)
    rows.month("f.ird", "Документы до разрешения", f"IF({ref('k')}<{ref('t.permit')},1,0)",
               unit="0/1", fmt=FLAG, total=False)
    rows.month("f.build", "Стройка основной части",
               f"IF(AND({ref('k')}>={ref('t.permit')},{ref('k')}<{ref('t.rve')}),1,0)",
               unit="0/1", fmt=FLAG, total=False)
    rows.month("f.rve", "Месяц ввода проекта", f"IF({ref('k')}={ref('t.rve')},1,0)",
               unit="0/1", fmt=FLAG, total=False)
    for item in objects:
        o = item["code"]
        rows.month(f"f.{o}.build", f"{item['title']}: стройка",
                   f"IF(AND({ref('k')}>={ref(f't.{o}.start')},{ref('k')}<{ref(f't.{o}.comm')}),1,0)",
                   unit="0/1", fmt=FLAG, total=False)
        rows.month(f"f.{o}.comm", f"{item['title']}: месяц ввода",
                   f"IF({ref('k')}={ref(f't.{o}.comm')},1,0)", unit="0/1", fmt=FLAG, total=False)


def _tep_sheet(book: _Book, objects: list[dict[str, Any]], info: _Inputs) -> None:
    ws = book.sheet(TEP_SHEET, "ТЭП · ПЛОЩАДИ, ЭФФЕКТИВНОСТЬ, МАШИНО-МЕСТА",
                    "Считается от «Вводных». Продаваемая площадь в расчёте — без площади мест на первых этажах.",
                    tab=HEAD, width=2 + len(objects) + 1)
    ws.column_dimensions["A"].width = 64
    for item in objects:
        ws.column_dimensions[info.columns[item["code"]]].width = 18
    r = 4
    ws.cell(r, 1, "Показатель").font = _fnt(b=True, color="FFFFFF")
    ws.cell(r, 1).fill = _fill(HEAD)
    for item in objects:
        cell = ws[f"{info.columns[item['code']]}{r}"]
        cell.value, cell.font, cell.fill = item["title"], _fnt(b=True, color="FFFFFF"), _fill(HEAD)
    unit_col = 2 + len(objects)
    r += 1

    def line(key: str, label: str, formula, fmt: str, unit: str, only_flows: bool = False) -> None:
        nonlocal r
        ws.cell(r, 1, label)
        ws.cell(r, unit_col, unit).font = _fnt(color=GREY)
        for item in objects:
            if only_flows and not (item.get("nonres") or item.get("ddu")):
                continue
            column = info.columns[item["code"]]
            cell = ws[f"{column}{r}"]
            book.pending.append((cell, formula(item["code"]), None, TEP_SHEET))
            cell.number_format = fmt
            book.register(f"tep.{item['code']}.{key}", TEP_SHEET, r, column)
        r += 1

    i = lambda o, key: ref(f"in.{o}.{key}")  # noqa: E731
    spaces = lambda o: f'{i(o, "measure")}="{SPACES}"'  # noqa: E731
    line("gba", "Наземная площадь здания", lambda o: i(o, "volume"), NUM0, "м² / мест")
    line("total_area", "Общая площадь", lambda o: i(o, "total_area"), NUM0, "м²")
    line("under", "Подземный гараж", lambda o: i(o, "under_gns"), NUM0, "м²")
    line("saleable", "Продаваемая и арендопригодная площадь по вводным", lambda o: i(o, "saleable"),
         NUM0, "м²", True)
    line("over_sqm", "Под местами на первых этажах",
         lambda o: f"MAX(0,{i(o, 'over_units')})*{i(o, 'over_area')}", NUM0, "м²", True)
    line("area", "Продаваемая и арендопригодная площадь в расчёте",
         lambda o: (f"IF({spaces(o)},{i(o, 'volume')},IF({i(o, 'volume')}>0,"
                    f"MAX(0,{i(o, 'volume')}-MAX(0,{i(o, 'over_units')})*{i(o, 'over_area')})"
                    f"*MAX(0,{i(o, 'saleable')})/{i(o, 'volume')},0))"), NUM0, "м² / мест", True)
    line("efficiency", "Эффективность: продаваемая в расчёте / наземная",
         lambda o: f"IF({i(o, 'volume')}>0,{ref(f'tep.{o}.area')}/{i(o, 'volume')},0)", PCT, "", True)
    line("spaces_total", "Машино-места: подземные и на первых этажах",
         lambda o: f"MAX(0,{i(o, 'parking_under_units')})+MAX(0,{i(o, 'over_units')})", NUM0, "шт.", True)
    line("spaces", "Машино-места в продаже или аренде", lambda o: i(o, "parking_units"), NUM0, "шт.", True)
    line("guest", "Гостевые места (не продаются и не сдаются)",
         lambda o: f"MAX(0,{ref(f'tep.{o}.spaces_total')}-{ref(f'tep.{o}.spaces')})", NUM0, "шт.", True)

    r += 1
    ws.cell(r, 1, "Пятно застройки и подземные уровни").font = _fnt(b=True, color=DARK)
    r += 1
    formulas = massing_book.formulas(lambda key: ref(f"ms.{key}"))
    # Подземная площадь — от вводных книги: ядро и гаражи объектов.
    formulas["underground_sqm"] = _sum([ref("in.core_under_gns")]
                                       + [ref(f"in.{item['code']}.under_gns") for item in objects])
    massing_rows = {}
    for key, label, kind in massing_book.ROWS:
        if kind != massing_book.FORMULA and key != "underground_sqm":
            continue
        ws.cell(r, 1, label)
        massing_rows[key] = r
        book.register(f"ms.{key}", TEP_SHEET, r, "B")
        r += 1
    for key, row in massing_rows.items():
        cell = ws[f"B{row}"]
        book.pending.append((cell, formulas[key], None, TEP_SHEET))
        cell.number_format = "#,##0.00"
    ws.freeze_panes = "B5"


def _rate_sheet(book: _Book, info: _Inputs) -> None:
    ws = book.sheet(RATE_SHEET, "СТАВКИ", "Ключевая ставка — прогноз сценария с «Вводных»; "
                    "ставка кредита объекта — ключевая плюс его надбавка (листы CF).", monthly=True, tab=HEAD)
    rows = _Rows(book, ws, 5)
    position = f'MAX(1,COUNTIF({info.rates["dates"]},"<="&{ref("date")}))'
    scenario = ref("in.rate_scenario")
    pick = (f'IF({scenario}="{SCENARIO["low"]}",INDEX({info.rates[SCENARIO["low"]]},{position}),'
            f'IF({scenario}="{SCENARIO["high"]}",INDEX({info.rates[SCENARIO["high"]]},{position}),'
            f'INDEX({info.rates[SCENARIO["base"]]},{position})))')
    rows.month("key_rate", "Ключевая ставка ЦБ", pick, unit="в год", fmt=PCT2, total=False, bold=True)


# --- «CAPEX» -------------------------------------------------------------------

ENGINE_LABELS = {"landscaping": "Благоустройство", "social": "Соцобъекты",
                 "vri_interest": "Проценты рассрочки платы за смену ВРИ",
                 "vri_security": "Обеспечение рассрочки платы за смену ВРИ",
                 "land_rights": "Плата за смену ВРИ (рассрочка)"}


def _capex_sheet(book: _Book, spec: dict[str, Any], objects: list[dict[str, Any]], info: _Inputs,
                 months: int) -> dict[str, Any]:
    """Смета «ставка × база» и её график по месяцам (`build_operating_model`).

    Сумма статьи — в колонке B со сценарием затрат, ряд месяцев раскладывает
    её теми же кривыми, что движок: S-кривая книги ПЛАТО, равномерно, разово,
    профиль управления, график оплаты участка.
    """
    ws = book.sheet(CAPEX_SHEET, "CAPEX · ЗАТРАТЫ НА СОЗДАНИЕ",
                    "Смета «ставка × база» по статьям, млн ₽ с НДС, со сценарием затрат. Курсив на оранжевом — "
                    "статьи, которые считает движок (формулы в книге нет).", monthly=True, tab=HEAD)
    rows = _Rows(book, ws, 5)
    k = ref("k")
    mult = ref("in.cost_mult")
    i = lambda o, key: ref(f"in.{o}.{key}")  # noqa: E731
    smr_only = lambda o, value: f'IF({i(o, "turnkey")}="{NO}",{value},0)'  # noqa: E731
    gns = "(" + _sum([ref("in.core_above_gns"), ref("in.core_under_gns")]
                     + [smr_only(item["code"], f"{i(item['code'], 'volume')}+{i(item['code'], 'under_gns')}")
                        for item in objects]) + ")"
    depth = ref("ms.depth_factor")
    engine = spec.get("engine_articles") or {}

    def engine_values(key: str) -> list[float]:
        return [float((engine.get(key) or {}).get(j, 0.0)) / 1e6 for j in range(book.n)]

    rows.section("Веса графика")
    for item in objects:
        o = item["code"]
        rows.month(f"w.{o}", f"S-кривая стройки: {item['title']}",
                   _s_weight(k, ref(f"t.{o}.start"), f"INT({i(o, 'months')})"), unit="вес", fmt="0.0", total=False)
    rows.month("w.main", "S-кривая основной части", _s_weight(k, ref("t.permit"), ref("t.build")),
               unit="вес", fmt="0.0", total=False)
    rows.month("w.ird", "Равномерно: документы до разрешения", _even_weight(k, "0", ref("t.ird_eff")),
               unit="вес", fmt=FLAG, total=False)
    rows.month("w.design", "Равномерно: окно проектирования",
               _even_weight(k, f"({ref('t.permit')}-{ref('t.window')})", ref("t.window")),
               unit="вес", fmt=FLAG, total=False)
    rows.month("w.rve3", "Равномерно: три месяца до ввода",
               _even_weight(k, f"({ref('t.rve')}-3)", "3"), unit="вес", fmt=FLAG, total=False)

    labels: dict[str, str] = {}
    order: list[str] = []
    main = f"{ref('w.main')}/SUM({ref('w.main', ':')})"

    def article(key: str, label: str, amount: str | None, monthly: str | None = None) -> None:
        labels[key] = label
        order.append(key)
        if amount is None:  # статья движка
            rows.month(f"cx.{key}", label + " — СЧИТАЕТ ДВИЖОК", values=engine_values(key))
            book.register(f"cx.{key}.amt", CAPEX_SHEET, rows.r - 1, "B")
            return
        rows.month(f"cx.{key}", label, monthly or f"{ref(f'cx.{key}.amt')}*{main}", amount=amount)

    rows.section("Статьи сметы")
    for item in objects:
        o, title = item["code"], item["title"]
        w = f"{ref(f'w.{o}')}/SUM({ref(f'w.{o}', ':')})"
        article(f"{o}.building", f"{title}: здание",
                f'{i(o, "volume")}*{i(o, "rate_cost")}*IF({i(o, "measure")}="{SPACES}",1,0.001)*{mult}',
                f"{ref(f'cx.{o}.building.amt')}*{w}")
        article(f"{o}.garage", f"{title}: подземный гараж",
                f"{i(o, 'under_gns')}*{ref('in.main_under_th')}*{depth}*0.001*{mult}",
                f"{ref(f'cx.{o}.garage.amt')}*{w}")
    per_sqm = lambda rate: f"{gns}*{ref(rate)}*0.001*{mult}"  # noqa: E731
    design = lambda key: f"{ref(f'cx.{key}.amt')}/{ref('t.window')}*{ref('w.design')}"  # noqa: E731
    article("ird", "ИРД", per_sqm("in.ird_th"),
            f"{ref('cx.ird.amt')}/{ref('t.ird_eff')}*{ref('w.ird')}")
    article("design_p", "Проект (стадия П)", per_sqm("in.design_p_th"), design("design_p"))
    article("design_rd", "Рабочая документация", per_sqm("in.design_rd_th"), design("design_rd"))
    article("demolition", "Снос", f"{ref('in.demolition_area')}*{ref('in.demolition_th')}*0.001*{mult}",
            design("demolition"))
    article("resettlement", "Расселение", f"{ref('in.resettlement_mln')}*{mult}", design("resettlement"))
    article("preparation", "Подготовка территории", per_sqm("in.preparation_th"), design("preparation"))
    article("main_above", "Стройка наземной части ядра",
            f"{ref('in.core_above_gns')}*{ref('in.main_above_th')}*0.001*{mult}")
    article("main_under", "Стройка подземной части ядра",
            f"{ref('in.core_under_gns')}*{ref('in.main_under_th')}*{depth}*0.001*{mult}")
    article("utilities", "Сети", per_sqm("in.utilities_th"))
    article("site_maintenance", "Содержание площадки", per_sqm("in.site_maintenance_th"))
    article("commissioning", "Ввод в эксплуатацию", per_sqm("in.commissioning_th"),
            f"{ref('cx.commissioning.amt')}/3*{ref('w.rve3')}")
    for key in ("landscaping", "social"):
        article(key, ENGINE_LABELS[key], None)
    amt = lambda key: ref(f"cx.{key}.amt")  # noqa: E731
    article("author_supervision", "Авторский надзор",
            f"({amt('design_p')}+{amt('design_rd')})*{ref('in.author_supervision')}")
    # База генподряда и техзаказчика — стройка ядра и соцобъектов, как в
    # движке, плюс здание и гараж объекта, чья ставка — СМР: ставка «под ключ»
    # несёт их сама, процент сверху был бы двойным счётом.
    object_works = "".join("+" + smr_only(item["code"], f"{amt(item['code'] + '.building')}"
                                                         f"+{amt(item['code'] + '.garage')}")
                           for item in objects)
    works = f"({amt('main_above')}+{amt('main_under')}+{amt('social')}{object_works})"
    article("technical_supervision", "Технический заказчик", f"{works}*{ref('in.technical_supervision')}")
    profile = ("ird", "design_p", "design_rd", "author_supervision", "preparation", "main_above",
               "main_under", "utilities", "landscaping", "site_maintenance")
    base = "+".join(amt(key) for key in profile) + object_works
    totals = "+".join(f"SUM({ref(f'cx.{key}', ':')})" for key in profile)
    monthly = "+".join(ref(f"cx.{key}") for key in profile)
    rve = f"MAX(1,{ref('t.rve')})"
    article("project_management", "Управление проектом", f"({base})*{ref('in.project_management')}",
            f"IF(({totals})>0,{amt('project_management')}*({monthly})/({totals}),"
            f"{amt('project_management')}/{rve}*{_even_weight(k, '0', rve)})")
    article("gc_fee", "Генподряд", f"{works}*{ref('in.gc_fee')}")
    if spec.get("vri_enabled"):
        article("land_rights", ENGINE_LABELS["land_rights"], None)
    else:
        gross, relief = ref("in.land_rights_gross"), ref("in.land_rights_relief")
        article("land_rights", "Плата за смену ВРИ к оплате",
                f"{gross}*(1-IF({gross}>0,MIN(1,{relief}/{gross}),0))*{mult}",
                f"IF({k}={ref('t.permit')},{amt('land_rights')},0)")
    # Резерв не берётся с известных сумм (`RESERVE_EXCLUDED_ARTICLES` движка).
    excluded = set(spec["reserve_excluded"])
    reserve_base = [amt(key) for key in order if key.split(".")[-1] not in excluded]
    article("reserve", "Резерв", f"({_sum(reserve_base)})*{ref('in.reserve')}")
    article("purchase", "Участок и выкуп участков и зданий",
            f"({ref('in.purchase_price')}+{ref('in.land_buyout')})*{mult}",
            f"{amt('purchase')}*SUMIF({info.plan[0]},{k},{info.plan[1]})")
    for key in ("vri_interest", "vri_security"):
        if sum((engine.get(key) or {}).values()):
            article(key, ENGINE_LABELS[key], None)

    rows.section("Итоги по месяцам")
    for item in objects:
        o = item["code"]
        rows.month(f"cx.{o}", f"{item['title']}: затраты объекта (здание и гараж)",
                   f"{ref(f'cx.{o}.building')}+{ref(f'cx.{o}.garage')}")
    rows.month("cx.total", "Все затраты на создание (CAPEX)",
               "+".join(ref(f"cx.{key}") for key in order), bold=True)
    flows = [item for item in objects if item.get("nonres")]
    rows.month("cx.flow_obj", "в т.ч. объекты с продажей или арендой",
               _sum([ref(f"cx.{item['code']}") for item in flows]))
    rows.month("cx.other_obj", "в т.ч. объекты без продажи",
               _sum([ref(f"cx.{item['code']}") for item in objects if not item.get("nonres")]))
    vri_equity = spec.get("vri_equity") or {}
    rows.month("cx.vri_equity", "Плата за ВРИ за счёт собственных денег — СЧИТАЕТ ДВИЖОК",
               values=[float(vri_equity.get(j, 0.0)) / 1e6 for j in range(book.n)])
    # Нежилой проект: банк кредитует долю ВСЕЙ стоимости — общие затраты
    # уходят в кредит объектов долей их стройки.
    rows.month("cx.common", "Общие затраты проекта на кредит объектов",
               f"MAX(0,MAX(0,{ref('cx.total')}-{ref('cx.vri_equity')}-{ref('cx.flow_obj')})"
               f"-{ref('cx.other_obj')})")
    rows.section("Доли и итоги")
    rows.param("cx.capex_total", "Все затраты на создание, итого", f"=SUM({ref('cx.total', ':')})", "млн ₽",
               bold=True)
    every = "+".join("SUM(" + ref("cx." + item["code"], ":") + ")" for item in objects) or "0"
    for item in flows:
        o = item["code"]
        rows.param(f"cx.share.{o}", f"Доля «{item['title']}» в общих затратах (его стройка / стройка всех объектов)",
                   f"=IF(({every})>0,SUM({ref(f'cx.{o}', ':')})/({every}),0)", "", "0.00%")
    pool = _sum([f"{amt(item['code'] + '.building')}+{amt(item['code'] + '.garage')}"
                 for item in objects if item["in_tax_pool"]])
    rows.param("cx.tax_pool", "Затраты объектов, признаваемые их продажами", f"={pool}", "млн ₽")
    # Горизонт: месяцев в таблицах столько, сколько нужно самому длинному
    # способа продажи; расчёт, уводящий проект дальше, краснит «ПРОВЕРКИ», а не обрезает поток.
    rows.param("cx.plan_last", "Последний платёж за участок, № месяца", f"=MAX({info.plan[0]})", "№", INT)
    return {"labels": labels, "order": order}


# --- объект: продажа или аренда, НДС и кредит (`object_flows`) -----------------

class _Block:
    """Деньги одного объекта при одной стратегии: имена его строк в книге."""

    def __init__(self, item: dict[str, Any], strategy: str, pre: str) -> None:
        self.item, self.strategy, self.pre = item, strategy, pre
        self.code = item["code"]
        self.names: set[str] = set()

    @property
    def direct(self) -> bool:
        return self.strategy == ns.STRATEGY_DIRECT

    def n(self, key: str) -> str:
        return f"{self.pre}.{key}"

    def r(self, key: str, mod: str = "") -> str:
        return ref(self.n(key), mod)

    def z(self, key: str) -> str:
        """Строка блока или ноль, если у этой стратегии её нет (аренда у продажи)."""
        return self.r(key) if self.n(key) in self.names else "0"


def _class_pick(info: _Inputs, item: dict[str, Any], key: str) -> str:
    """Вводная офиса «в расчёт»: вписанное вручную, иначе ступень класса."""
    o = item["code"]
    manual = ref(f"in.{o}.{key}")
    if not item.get("office"):
        return manual
    field = CLASS_FIELDS[key]
    pick = f"INDEX({info.class_rows[field]},1,MATCH({ref(f'in.{o}.class')},{info.class_head},0))"
    if key == "rent":
        pick = f"ROUND({pick}/12,3)"
    return f'IF({manual}<>"",{manual},{pick})'


def _track(rows: _Rows, block: _Block):
    """Обёртка над строками листа: каждое имя блока помнится, чтобы `z` знал, что есть."""
    def param(key, *args, **kw):
        block.names.add(block.n(key))
        rows.param(block.n(key), *args, **kw)

    def month(key, *args, **kw):
        block.names.add(block.n(key))
        rows.month(block.n(key), *args, **kw)
    return param, month


def _flows(book: _Book, rows: _Rows, exit_rows: _Rows | None, block: _Block, info: _Inputs) -> None:
    """Выручка и расходы объекта: продажа блоками или аренда с выходом."""
    item, o = block.item, block.code
    b = block.r
    i = lambda key: ref(f"in.{o}.{key}")  # noqa: E731
    k = ref("k")
    param, month = _track(rows, block)
    spaces = f'{i("measure")}="{SPACES}"'
    rows.section(f"{item['title']} · вводные способа продажи")
    param("comm", "Ввод объекта, № месяца", f"={ref(f't.{o}.comm')}", "№", INT)
    param("area", "Продаваемая и арендопригодная площадь (места — шт.)", f"={ref(f'tep.{o}.area')}",
          "м² / шт.", NUM0)
    param("spaces", "Машино-места в продаже или аренде", f"={i('parking_units')}", "шт.", NUM0)
    param("rev_mult", "Сценарий: множитель выручки", f"={ref('in.revenue_mult')}", "×", "0.00")
    param("vat", "НДС", f"=MAX(0,{ref('in.vat')})", "", PCT)
    param("ptax_rate", "Налог на имущество", f"=MAX(0,{i('property_tax')})", "в год", PCT)
    if block.direct:
        under, over = i("parking_under_units"), i("over_units")
        default = f"{ref('in.parking_price_th')}*0.001"
        under_price = f"IF({i('parking_under_mln')}>0,{i('parking_under_mln')},{default})"
        over_price = f"IF({i('parking_over_mln')}>0,{i('parking_over_mln')},{default})"
        param("price", "Цена продажи «в расчёт» (места — млн ₽ за место)",
              f"={_class_pick(info, item, 'price')}", "тыс ₽/м² с НДС", NUM)
        param("parking_price", "Цена машино-места, средняя по гаражу (млн ₽)",
              f"=IF(MAX(0,{under})+MAX(0,{over})>0,(MAX(0,{under})*{under_price}+MAX(0,{over})*{over_price})"
              f"/(MAX(0,{under})+MAX(0,{over})),{default})", "млн ₽", NUM3)
        param("growth_pre", "Рост цены до ввода", f"={i('growth_pre')}", "в месяц", PCT2)
        param("growth_post", "Рост цены после ввода", f"={i('growth_post')}", "в месяц", PCT2)
        param("price_start", "Отсчёт роста цены, № месяца", f"={_k_of(book, i('sales_start'))}", "№", INT)
        param("selling_share", "Маркетинг и брокеры", f"={ref('in.marketing')}+{ref('in.selling')}", "доля выручки", PCT)
        param("cost_mult", "Сценарий: множитель затрат", f"={ref('in.cost_mult')}", "×", "0.00")
        param("sale_start", "Старт продаж, № месяца", f"={b('comm')}+MAX(-120,INT({i('sale_offset')}))", "№", INT)
        param("sale_months", "Срок продаж", f"=MAX(1,INT({i('sale_months')}))", "мес.", INT)
        param("curve", "Профиль продаж", f"={i('curve')}", "", "@")
        n_ = b("sale_months")
        curve = lambda code: f'{b("curve")}="{CURVE[code]}"'  # noqa: E731
        param("curve_sum", "Сумма весов профиля продаж",
              f"=IF({curve('bell')},{n_}*({n_}+1)*({n_}+2)/6,IF(OR({curve('front_loaded')},"
              f"{curve('back_loaded')}),{n_}*({n_}+1)/2,{n_}))", "", NUM0)
        param("horizon", "Последний месяц способа продажи, №",
              f"=MAX({b('comm')},{b('sale_start')}+{b('sale_months')}-1)", "№", INT)
    else:
        param("rent_rate", "Ставка аренды «в расчёт»", f"={_class_pick(info, item, 'rent')}",
              "тыс ₽/м² в месяц", NUM3)
        param("parking_rent", "Аренда машино-места", f"={i('parking_rent')}", "тыс ₽ в месяц", NUM)
        param("rent_index", "Индексация аренды", f"={i('rent_index')}", "в год", PCT)
        # Ставка дана на свою дату и до ввода индексируется (`rent_months_before`).
        param("rent_from", "Месяц ставки аренды, № (без даты — ввод)",
              f'=IF({i("rent_date")}="",{b("comm")},{_k_of(book, i("rent_date"))})', "№", INT)
        param("occ_start", "Заполняемость на открытии", f"=MIN(1,MAX(0,{i('occ_start')}))", "", PCT)
        param("occ_stable", "Заполняемость после заполнения",
              f"=MIN(1,MAX(0,{_class_pick(info, item, 'occ_stable')}))", "", PCT)
        param("leaseup", "Срок заполнения", f"=MAX(1,INT({_class_pick(info, item, 'leaseup')}))", "мес.", INT)
        param("opex_share", "Расходы эксплуатации", f"=MIN(0.95,MAX(0,{_class_pick(info, item, 'opex')}))",
              "доля дохода", PCT)
        param("hold", "Срок удержания", f"=MAX(1,INT({i('hold_years')}))*12", "мес.", INT)
        param("exit_mode", "Выход", f"={i('exit_mode')}", "", "@")
        param("horizon", "Последний месяц удержания, №", f"={b('comm')}+{b('hold')}-1", "№", INT)
        param("dep_years", "Срок амортизации для налога", f"=MAX(1,INT({i('dep_years')}))", "лет", INT)
    rows.section(f"{item['title']} · затраты объекта")
    month("capex", "Затраты стройки объекта", ref(f"cx.{o}"))
    month("common", "Доля общих затрат проекта", f"{ref('cx.common')}*{ref(f'cx.share.{o}')}")
    param("capex_total", "Затраты стройки объекта, итого", f"=SUM({b('capex', ':')})", "млн ₽")
    param("common_total", "Общие затраты проекта на объекте", f"=SUM({b('common', ':')})", "млн ₽")
    param("basis", "Стоимость объекта без НДС (база налога на имущество и амортизации)",
          f"={b('capex_total')}*(1-{b('vat')}/(1+{b('vat')}))", "млн ₽")
    if block.direct:
        rows.section(f"{item['title']} · продажа блоками")
        j = f"({k}-{b('sale_start')}+1)"
        n_ = b("sale_months")
        curve = lambda code: f'{b("curve")}="{CURVE[code]}"'  # noqa: E731
        month("weight", "Доля продаж месяца",
              f"IF(AND({j}>=1,{j}<={n_}),IF({curve('bell')},{j}*({n_}+1-{j}),IF({curve('front_loaded')},"
              f"{n_}+1-{j},IF({curve('back_loaded')},{j},1)))/{b('curve_sum')},0)", unit="доля", fmt=PCT)
        ps, comm = b("price_start"), b("comm")
        month("factor", "Индекс цены",
              f"(1+{b('growth_pre')})^MAX(0,MIN({k}-{ps},{comm}-{ps}))"
              f"*(1+{b('growth_post')})^MAX(0,{k}-MAX({ps},{comm}))", unit="×", fmt="0.000", total=False)
        price = f"{b('price')}*IF({spaces},1,0.001)"
        base = f"({b('area')}*{price}+{b('spaces')}*{b('parking_price')})*{b('factor')}*{b('weight')}"
        month("sale", "Выручка продаж", f"{base}*{b('rev_mult')}", bold=True)
        month("selling", "Маркетинг и брокеры", f"{base}*{b('selling_share')}*{b('cost_mult')}")
        month("unsold", "Не продано на начало месяца", f"MAX(0,1-SUM({b('weight', ':p')}))",
              unit="доля", fmt=PCT, total=False)
        month("ptax", "Налог на имущество с непроданной части",
              f"IF(AND({k}>={comm},{k}<={b('horizon')}),{b('basis')}*{b('ptax_rate')}/12*{b('unsold')},0)")
        rows.section(f"{item['title']} · налоговая маржа")
        month("recognized", "Затраты, признанные проданными м²",
              f"IF({k}>{b('horizon')},0,IF({k}<{comm},0,IF({k}={comm},SUM({b('weight', ':k')}),"
              f"{b('weight')}))*{b('basis')})")
        month("realized", "Выручка без НДС к признанию (до ввода копится)",
              f"IF({k}<{comm},0,IF({k}={comm},SUM({b('sale', ':k')})-SUM({b('vat_charged', ':k')}),"
              f"{b('sale')}-{b('vat_charged')}))")
        month("margin", "Налоговая маржа объекта",
              f"IF({k}>{b('horizon')},0,{b('realized')}-{b('recognized')}-{b('selling')}-{b('ptax')})")
        return
    rows.section(f"{item['title']} · аренда")
    comm, hold = b("comm"), b("hold")
    month("open", "Месяц эксплуатации", f"IF({k}>={comm},{k}-{comm}+1,0)", unit="№", fmt=INT, total=False)
    os_, ost = b("occ_start"), b("occ_stable")
    month("occ", "Заполняемость",
          f"IF({b('open')}=0,0,MIN({os_},{ost})+({ost}-MIN({os_},{ost}))*IF({b('leaseup')}=1,1,"
          f"MAX(0,MIN(1,({b('open')}-1)/({b('leaseup')}-1)))))", unit="", fmt=PCT, total=False)
    month("growth", "Индекс аренды",
          f"(1+MAX(-0.95,{b('rent_index')}))^((MAX(0,{b('open')}-1)+MAX(0,{comm}-{b('rent_from')}))/12)",
          unit="×", fmt="0.000", total=False)
    month("op_rev", "Доход аренды по всем месяцам эксплуатации",
          f"({b('area')}*{b('rent_rate')}+{b('spaces')}*{b('parking_rent')})*0.001*{b('growth')}*{b('occ')}*{b('rev_mult')}",
          total=False)
    month("hold_on", "Удержание", f"IF(AND({b('open')}>=1,{b('open')}<={hold}),1,0)",
          unit="0/1", fmt=FLAG, total=False)
    month("rent", "Доход аренды", f"{b('op_rev')}*{b('hold_on')}", bold=True)
    month("opex", "Расходы эксплуатации", f"{b('rent')}*{b('opex_share')}")
    month("ptax", "Налог на имущество", f"{b('basis')}*{b('ptax_rate')}/12*{b('hold_on')}")
    month("noi", "NOI", f"{b('rent')}-{b('opex')}-{b('ptax')}", bold=True)
    month("fwd", "NOI покупателя: год после срока удержания (до налога на имущество)",
          f"IF(AND({b('open')}>={hold}+1,{b('open')}<={hold}+12),{b('op_rev')}*(1-{b('opex_share')}),0)")
    month("stab", "NOI года после заполнения (до налога на имущество)",
          f"IF(AND({b('open')}>={b('leaseup')},{b('open')}<={b('leaseup')}+11),"
          f"{b('op_rev')}*(1-{b('opex_share')}),0)")
    month("year", "Год эксплуатации",
          f"IF({b('hold_on')}=1,INT(({b('open')}-1)/12)+1,0)", unit="№", fmt=INT, total=False)

    # Выход: стоимость здания = NOI следующего года / ставка капитализации.
    ep, em = _track(exit_rows, block)
    exit_rows.section(f"{item['title']} · продажа здания", RENT_TAB)
    ep("fwd_noi", "NOI следующих 12 месяцев после срока удержания",
       f"=SUM({b('fwd', ':')})-{b('basis')}*{b('ptax_rate')}", "млн ₽")
    ep("cap", "Ставка капитализации", f"={_class_pick(info, item, 'cap')}", "", PCT)
    ep("exit_value", "Стоимость здания", f"=IF(AND({b('cap')}>0,{b('fwd_noi')}>0),{b('fwd_noi')}/{b('cap')},0)",
       "млн ₽", bold=True)
    ep("exit_cost_share", "Расходы на сделку", f"=MAX(0,{i('exit_cost')})", "", PCT)
    ep("exit_net", "Получено от продажи здания", f"={b('exit_value')}*(1-{b('exit_cost_share')})", "млн ₽")
    ep("noi_stab", "NOI года после заполнения",
       f"=SUM({b('stab', ':')})-{b('basis')}*{b('ptax_rate')}", "млн ₽", bold=True)
    ep("cost_total", "Все затраты создания: стройка и доля общих затрат с участком",
       f"={b('capex_total')}+{b('common_total')}", "млн ₽")
    ep("yoc", "Доходность на затраты = NOI года после заполнения / все затраты",
       f"=IF({b('cost_total')}>0,{b('noi_stab')}/{b('cost_total')},0)", "", PCT, bold=True)
    ep("yoc_spread", "Доходность на затраты минус ставка капитализации", f"={b('yoc')}-{b('cap')}", "", PCT)
    exit_rows.note("Меньше нуля — готовое здание стоит дешевле, чем обходится его стройка: строить, "
                   "чтобы сдавать, невыгодно.")
    is_sale = f'{b("exit_mode")}="{EXIT[ns.EXIT_SALE]}"'
    is_hold = f'{b("exit_mode")}="{EXIT[ns.EXIT_HOLD]}"'
    rows.section(f"{item['title']} · выход")
    month("exit", "Продажа здания", f"IF(AND({is_sale},{k}={b('horizon')}),{b('exit_value')},0)", bold=True)
    month("exit_cost", "Расходы на сделку", f"{b('exit')}*{b('exit_cost_share')}")
    month("residual", "Оценка удержанного здания (не деньги)",
          f"IF(AND({is_hold},{k}={b('horizon')}),{b('exit_value')},0)")
    rows.section(f"{item['title']} · налоговая маржа")
    month("book_open", "Остаточная стоимость на начало месяца", f"{b('book_close', '@p')}",
          first=b("basis"), total=False)
    month("recognized", "Признано в расходах: амортизация, остаток при продаже",
          f"IF({k}>{b('horizon')},0,IF({k}>={comm},IF(AND({b('exit')}>0,{k}={b('horizon')}),{b('book_open')},"
          f"MIN({b('book_open')},{b('basis')}/({b('dep_years')}*12))),0))")
    month("book_close", "Остаточная стоимость на конец месяца", f"{b('book_open')}-{b('recognized')}",
          total=False)
    month("margin", "Налоговая маржа объекта",
          f"IF({k}>{b('horizon')},0,{b('rent')}+{b('exit')}-{b('vat_charged')}-{b('recognized')}"
          f"-{b('opex')}-{b('ptax')}-{b('exit_cost')})")


def _ddu_flows(book: _Book, rows: _Rows, block: _Block, info: _Inputs) -> None:
    """Продажа объекта по ДДУ — `sales_schedule` движка: доля продаж до ввода
    объекта ровными долями от старта продаж, остаток — ровными долями после
    ввода; цена растёт до ввода от старта продаж, после — от цены ввода.
    Профиль продаж и лестница цены, заданные на странице проекта, приходят
    рядами по месяцам. Деньги покупателей идут на эскроу (лист «ПФ»)."""
    item, o = block.item, block.code
    b = block.r
    i = lambda key: ref(f"in.{o}.{key}")  # noqa: E731
    k = ref("k")
    param, month = _track(rows, block)
    spaces = f'{i("measure")}="{SPACES}"'
    rows.section(f"{item['title']} · продажа по ДДУ: вводные", DDU_TAB)
    param("comm", "Ввод объекта, № месяца", f"={ref(f't.{o}.comm')}", "№", INT)
    param("area", "Продаваемая площадь (места — шт.)", f"={ref(f'tep.{o}.area')}", "м² / шт.", NUM0)
    param("spaces", "Машино-места в продаже", f"=MAX(0,{i('parking_units')})", "шт.", NUM0)
    param("rev_mult", "Сценарий: множитель выручки", f"={ref('in.revenue_mult')}", "×", "0.00")
    under, over = i("parking_under_units"), i("over_units")
    default = f"{ref('in.parking_price_th')}*0.001"
    under_price = f"IF({i('parking_under_mln')}>0,{i('parking_under_mln')},{default})"
    over_price = f"IF({i('parking_over_mln')}>0,{i('parking_over_mln')},{default})"
    param("price", "Цена продажи «в расчёт» (места — млн ₽ за место)",
          f"={_class_pick(info, item, 'price')}", "тыс ₽/м² с НДС", NUM)
    param("parking_price", "Цена машино-места, средняя по гаражу (млн ₽)",
          f"=IF(MAX(0,{under})+MAX(0,{over})>0,(MAX(0,{under})*{under_price}+MAX(0,{over})*{over_price})"
          f"/(MAX(0,{under})+MAX(0,{over})),{default})", "млн ₽", NUM3)
    param("growth_pre", "Рост цены до ввода", f"={i('growth_pre')}", "в месяц", PCT2)
    param("growth_post", "Рост цены после ввода", f"={i('growth_post')}", "в месяц", PCT2)
    param("price_start", "Старт продаж, № месяца", f"={_k_of(book, i('sales_start'))}", "№", INT)
    param("share", "Доля продаж до ввода объекта", f"=MAX(0,MIN(1,{i('share_before')}))", "", PCT)
    param("residual", "Остаточные продажи после ввода", f"=MAX(0,INT({i('residual')}))", "мес.", INT)
    param("pre_n", "Месяцев продаж до ввода", f"=MAX(1,{b('comm')}-{b('price_start')})", "мес.", INT)
    param("selling_share", "Маркетинг и брокеры", f"={ref('in.marketing')}+{ref('in.selling')}", "доля выручки",
          PCT)
    param("cost_mult", "Сценарий: множитель затрат", f"={ref('in.cost_mult')}", "×", "0.00")
    param("cost", "Затраты объекта, признаваемые его продажами (здание и гараж)",
          f"={ref(f'cx.{o}.building.amt')}+{ref(f'cx.{o}.garage.amt')}", "млн ₽")
    rows.section(f"{item['title']} · продажа по ДДУ по месяцам", DDU_TAB)
    s, comm, n_pre = b("price_start"), b("comm"), b("pre_n")
    weights, factor = item.get("sales_weights"), item.get("price_factor")
    if weights is not None:
        month("weight", "Доля продаж месяца — профиль задан на странице проекта", unit="доля", fmt=PCT,
              values=[float(weights.get(j, 0.0)) for j in range(book.n)])
    else:
        month("weight", "Доля продаж месяца",
              f"IF(AND({k}>={s},{k}<{s}+{n_pre}),{b('share')}/{n_pre},0)"
              f"+IF(AND({b('residual')}>0,{k}>={comm},{k}<{comm}+{b('residual')}),"
              f"(1-{b('share')})/{b('residual')},0)", unit="доля", fmt=PCT)
    if factor is not None:
        month("ladder", "Множитель цены — лестница задана на странице проекта", unit="×", fmt="0.000",
              total=False, values=[float(factor.get(j, 1.0)) for j in range(book.n)])
        growth = f"{b('ladder')}*IF({k}>={comm},(1+{b('growth_post')})^({k}-{comm}),1)"
    else:
        growth = (f"IF({k}<{comm},(1+{b('growth_pre')})^({k}-{s}),"
                  f"(1+{b('growth_pre')})^{n_pre}*(1+{b('growth_post')})^({k}-{comm}))")
    month("factor", "Индекс цены", growth, unit="×", fmt="0.000", total=False)
    month("qty", "Продано за месяц", f"{b('area')}*{b('weight')}", unit="м² / шт.", fmt=NUM0)
    price = f"{b('price')}*IF({spaces},1,0.001)*{b('factor')}"
    parking = f"{b('spaces')}*{b('weight')}*{b('parking_price')}*{b('factor')}"
    month("sale", "Выручка площади", f"{b('qty')}*{price}*{b('rev_mult')}")
    month("parking", "Выручка машино-мест", f"{parking}*{b('rev_mult')}")
    month("revenue", "Выручка по ДДУ — деньги покупателей на эскроу", f"{b('sale')}+{b('parking')}", bold=True)
    month("base", "Выручка без сценария — база расходов на продажу", f"{b('qty')}*{price}+{parking}")
    month("operating", "Маркетинг и брокеры", f"{b('base')}*{b('selling_share')}*{b('cost_mult')}")
    # Затраты объекта признаются по проданным метрам; места своих затрат не
    # несут — они в затратах объекта (`_profit_tax_schedule`).
    month("margin", "Налоговая маржа: выручка − затраты по проданным метрам",
          f"{b('revenue')}-IF({b('qty_total')}>0,{b('cost')}*{b('qty')}/{b('qty_total')},0)")
    month("last_k", "Служебная: месяц с продажами", f"IF({b('weight')}>0,{k},-1)", unit="№", fmt=INT, total=False)
    param("qty_total", "Продано за весь срок", f"=SUM({b('qty', ':')})", "м² / шт.", NUM0)
    param("horizon", "Последний месяц продаж, №", f"=MAX({b('last_k', ':')})", "№", INT)


def _pf(book: _Book, rows: _Rows, path: dict[str, str], ddu: list[_Block], credit: list[_Block],
        info: _Inputs) -> None:
    """Проектное финансирование объектов по ДДУ — `simulate_financing` движка
    по месяцам: до разрешения на стройку затраты платят собственные деньги,
    затем БРИДЖ; с разрешения ПФ выдаёт все затраты и рефинансирует БРИДЖ;
    ставка ПФ — смесь надбавки и льготной ставки по покрытию долга эскроу;
    в РВЭ эскроу раскрывается и гасит ПФ, после — продажи."""
    c = path["code"]
    P = lambda key, mod="": ref(f"pf.{c}.{key}", mod)  # noqa: E731
    param = lambda key, *a, **kw: rows.param(f"pf.{c}.{key}", *a, **kw)  # noqa: E731
    month = lambda key, *a, **kw: rows.month(f"pf.{c}.{key}", *a, **kw)  # noqa: E731
    k = ref("k")
    tab = path["tab"]
    credit_capex = _sum([f"SUM({ref(f'cx.{b.code}', ':')})" for b in credit])
    rows.section("Проектное финансирование: условия", tab)
    param("permit", "Разрешение на стройку — открытие ПФ, № месяца", f"={ref('t.permit')}", "№", INT)
    param("rve", "Ввод проекта (РВЭ) — раскрытие эскроу, № месяца", f"={ref('t.rve')}", "№", INT)
    param("own_funds", "Собственные деньги до ПФ", f"=MAX(0,{ref('in.pf.own_funds')})", "млн ₽")
    param("approved", "Одобренный лимит ПФ", f"=MAX(0,{ref('in.pf.approved')})", "млн ₽")
    param("limit", "Лимит ПФ: одобренный, иначе выдача с округлением вверх до 10 млн",
          f"=IF({P('approved')}>0,{P('approved')},ROUNDUP(SUM({P('p_draw', ':')})/10,0)*10)", "млн ₽", bold=True)
    # Объекты вне ДДУ берут общие затраты в свой кредит, и БРИДЖа нет
    # (`common_on_object_loans`); иначе лимит БРИДЖа — оплаты до разрешения.
    before = "+".join(f'SUMIF({ref("k", ":")},"<"&{P("permit")},{ref(f"cx.{key}", ":")})'
                      for key in ("purchase", "design_p", "design_rd"))
    param("bridge_limit", "Лимит БРИДЖа: участок и проект до разрешения на стройку",
          "=0" if credit else f"={before}", "млн ₽")
    param("bridge_fee", "Плата за резервирование БРИДЖа (в первый месяц)",
          f"={P('bridge_limit')}*MAX(0,{ref('in.reservation_fee')})", "млн ₽")
    param("reservation_fee", "Плата за резервирование лимита ПФ (при открытии)",
          f"={P('limit')}*MAX(0,{ref('in.reservation_fee')})", "млн ₽")
    shares = _sum([ref(f"cx.share.{b.code}") for b in credit])
    cap_mode = f'{ref("in.pf.mode")}="{INTEREST_MODES[0]}"'
    pay_mode = f'{ref("in.pf.mode")}<>"{INTEREST_MODES[0]}"'
    permit, rve = P("permit"), P("rve")
    rows.section("Деньги по ДДУ и затраты на ПФ", tab)
    month("key_rate", "Ключевая ставка", ref("key_rate"), unit="в год", fmt=PCT2, total=False)
    month("sales", "Выручка по ДДУ", _sum([b.r("revenue") for b in ddu]), bold=True)
    month("operating", "Маркетинг и брокеры продаж по ДДУ", _sum([b.r("operating") for b in ddu]))
    # Долговые затраты — `debt_capex` движка: CAPEX без платы за ВРИ за счёт
    # собственных денег, без объектов вне ДДУ и без их доли общих затрат.
    month("capex", "Затраты на ПФ: CAPEX без объектов вне ДДУ и их доли общих затрат",
          f"MAX(0,MAX(0,{ref('cx.total')}-{ref('cx.vri_equity')})"
          f"-({_sum([ref(f'cx.{b.code}') for b in credit])}))-{ref('cx.common')}*({shares})")
    month("costs", "Потребность месяца",
          f"{P('capex')}+{P('operating')}+IF(AND({k}={permit},{pay_mode}),{P('transfer')},0)")
    month("esc", "Эскроу на конец месяца (до раскрытия)",
          f"IF({k}<{rve},{P('esc', '@p')}+{P('sales')},IF({k}={rve},{P('esc', '@p')},0))", total=False)
    month("release", "Раскрытие эскроу", f"IF({k}={rve},{P('esc')},0)")
    # 214-ФЗ: до ввода деньги покупателей на эскроу, застройщику они приходят
    # раскрытием счёта, проданное после ввода — в месяц продажи. Это выручка
    # потока проекта, как `escrow_release` в `_calculate_economics`.
    month("cash", "Деньги по ДДУ застройщику: раскрытие эскроу и продажи после ввода",
          f"IF({k}<{rve},0,{P('sales')})+{P('release')}")
    rows.section("Собственные деньги и БРИДЖ до разрешения на стройку", tab)
    need = f"MAX({P('costs')},0)"
    month("own_avail", "Собственные деньги — остаток на начало месяца", P("own_left", "@p"),
          first=P("own_funds"), total=False)
    month("own_draw", "Собственные деньги", f"IF({k}<{permit},MIN({P('own_avail')},{need}),0)")
    month("own_left", "Собственные деньги — остаток", f"{P('own_avail')}-{P('own_draw')}", total=False)
    month("b_draw", "БРИДЖ — выдача", f"IF({k}<{permit},MAX(0,{need}-{P('own_draw')}),0)")
    month("b_bal", "БРИДЖ — долг", f"IF({k}<{permit},{P('b_bal', '@p')}+{P('b_draw')},0)", total=False)
    b_rate = f"({P('key_rate')}+{ref('in.pf.bridge_spread')})"
    b_cap_rate = f"({P('key_rate')}+{ref('in.pf.bridge_cap_spread')})"
    month("b_int", "БРИДЖ — проценты",
          f"IF(AND({k}<{permit},{P('b_bal')}>0),{P('b_bal')}*{b_rate}/12,0)")
    month("b_cap", "БРИДЖ — проценты на проценты к уплате",
          f"IF(AND({k}<{permit},{P('b_bal')}>0),{P('b_pay', '@p')}*{b_cap_rate}/12,0)")
    month("b_pay", "БРИДЖ — проценты к уплате",
          f"IF({k}<{permit},{P('b_pay', '@p')}+{P('b_int')}+{P('b_cap')},0)", total=False)
    month("b_repay", "БРИДЖ — погашение выдачей ПФ", f"IF({k}={permit},{P('b_bal', '@p')},0)")
    month("transfer", "Проценты БРИДЖа при открытии ПФ", f"IF({k}={permit},{P('b_pay', '@p')},0)")
    rows.section("ПФ: выдача, покрытие эскроу, ставка, проценты, погашение", tab)
    # ПФ выбирает все расходы с разрешения; линия, рассчитавшаяся в РВЭ, закрыта.
    draw = f"IF(OR({k}<={rve},{P('settled', '@p')}=0),{need},0)+{P('b_repay')}"
    month("p_draw", "ПФ — выдача",
          f"IF({k}>={permit},IF({P('approved')}>0,MIN({draw},MAX({P('approved')}-{P('p_bal', '@p')},0)),"
          f"{draw}),0)", bold=True)
    month("p_bal_after", "ПФ — долг после выдачи", f"{P('p_bal', '@p')}+{P('p_draw')}", total=False)
    month("cov", "Покрытие долга эскроу",
          f"IF({P('p_bal_after')}>0,{P('esc')}/{P('p_bal_after')},0)", unit="×", fmt=RATIO, total=False)
    ladder = ref("in.pf.pf_special")
    for edge, rate in info.steps:
        ladder = f"IF(AND({edge}>0,{P('cov')}+0.000000001>={edge}),{rate},{ladder})"
    month("special", "Льготная ставка ступени", ladder, unit="в год", fmt=PCT2, total=False)
    weight = f"MIN({P('cov')},1)"
    month("p_rate", "Ставка ПФ",
          f"IF({k}>={permit},({P('key_rate')}+{ref('in.pf.pf_spread')})*(1-{weight})+{P('special')}*{weight},0)",
          unit="в год", fmt=PCT2, total=False)
    # Расчёт движка кончается горизонтом проекта: дальше проценты не идут.
    live = f"AND({k}>={permit},{P('p_bal_after')}>0,{k}<{ref(f'{c}.horizon')})"
    month("p_pay_open", "ПФ — проценты к уплате на начало месяца",
          f"{P('p_pay', '@p')}+IF(AND({k}={permit},{cap_mode}),{P('transfer')},0)", total=False)
    month("p_int", "ПФ — проценты", f"IF({live},{P('p_bal_after')}*{P('p_rate')}/12,0)")
    month("p_cap", "ПФ — проценты на проценты к уплате", f"IF({live},{P('p_pay_open')}*{P('p_rate')}/12,0)")
    month("fee", "Плата за невыбранный лимит ПФ",
          f"IF(AND({live},{k}<{rve},{P('limit')}>0),MAX({P('limit')}-{P('p_bal_after')},0)"
          f"*MAX(0,{ref('in.limit_fee')})/12,0)")
    accrued = f"{P('p_pay_open')}+{P('p_int')}+{P('p_cap')}+{P('fee')}"
    month("pay_out", "Проценты и плата уплачены (с РВЭ)", f"IF(AND({k}>={permit},{k}>={rve}),{accrued},0)")
    month("p_pay", "ПФ — проценты к уплате на конец месяца",
          f"IF(AND({k}>={permit},{k}>={rve}),0,{accrued})", total=False)
    month("avail", "Деньги на погашение: раскрытый эскроу, затем продажи",
          f"IF({k}={rve},{P('release')},IF({k}>{rve},{P('sales')},0))")
    month("p_repay", "ПФ — погашение",
          f"IF(AND({k}>={permit},{P('p_bal_after')}>0),MIN({P('avail')},{P('p_bal_after')}),0)", bold=True)
    month("p_bal", "ПФ — долг на конец месяца", f"{P('p_bal_after')}-{P('p_repay')}", total=False, bold=True)
    month("settled", "ПФ закрыт в РВЭ", f"IF({k}={rve},IF(AND({k}>={permit},{P('p_bal')}<=0),1,0),"
          f"IF({k}>{rve},{P('settled', '@p')},0))", unit="0/1", fmt=FLAG, total=False)
    rows.section("Что ПФ даёт налогам и потокам", tab)
    month("interest", "Проценты начислены: БРИДЖ и ПФ",
          f"{P('b_int')}+{P('b_cap')}+{P('p_int')}+{P('p_cap')}")
    month("fees", "Комиссии: резервирование и невыбранный лимит",
          f"{P('fee')}+IF({k}=0,{P('bridge_fee')},0)+IF({k}={permit},{P('reservation_fee')},0)")
    month("deduct", "Вычет из базы налога: проценты и комиссии", f"{P('interest')}+{P('fees')}")
    month("vat", "НДС продаж по ДДУ к уплате (в РВЭ)", f"IF({k}={rve},{P('vat_due')},0)")
    # Поток капитала: деньги продаж — с раскрытия эскроу; выдачи долга
    # покрывают затраты, погашения уходят банку.
    month("equity", "Деньги ДДУ и ПФ собственнику (без затрат стройки и налога на прибыль)",
          f"IF({k}<{rve},0,{P('sales')}+{P('release')})-{P('operating')}-{P('pay_out')}+{P('b_draw')}"
          f"+{P('p_draw')}-{P('b_repay')}-{P('p_repay')}-IF({k}=0,{P('bridge_fee')},0)"
          f"-IF({k}={permit},{P('reservation_fee')},0)-{P('vat')}")
    rows.section("Итоги ПФ", tab)
    param("pf_draw_total", "Выдача ПФ всего", f"=SUM({P('p_draw', ':')})", "млн ₽", bold=True)
    param("transferred", "Проценты БРИДЖа, перенесённые в ПФ", f"=SUM({P('transfer', ':')})", "млн ₽")
    param("financing_cost", "Проценты и комиссии БРИДЖа и ПФ",
          f"=SUM({P('interest', ':')})+SUM({P('fees', ':')})", "млн ₽")
    param("peak_bridge", "БРИДЖ — пик долга", f"=MAX({P('b_bal', ':')})", "млн ₽")
    param("peak_pf", "ПФ — пик долга", f"=MAX({P('p_bal', ':')})", "млн ₽", bold=True)
    param("peak_escrow", "Эскроу — пик", f"=MAX({P('esc', ':')})", "млн ₽")
    param("ending_pf", "Долг ПФ на конец проекта",
          f"=SUMIF({ref('k', ':')},{ref(f'{c}.horizon')}-1,{P('p_bal', ':')})", "млн ₽", bold=True)
    param("rve_unpaid", "Не погашено раскрытым эскроу в РВЭ",
          f"=SUMIF({ref('k', ':')},{rve},{P('p_bal', ':')})", "млн ₽")
    # НДС продаж по ДДУ — как у движка: начислен в РВЭ со всей выручки по ДДУ,
    # вычет — входящий НДС затрат без участка, выкупа, платы за ВРИ и объектов
    # вне ДДУ (их НДС зачитывают они сами); к уплате — разница.
    gross = f"(MAX(0,{ref('in.vat')})/(1+MAX(0,{ref('in.vat')})))"
    revenue = f"SUM({P('sales', ':')})"
    vri = "".join(f"-{ref(f'cx.{key}.amt')}" for key in ("land_rights", "vri_interest", "vri_security")
                  if f"cx.{key}.amt" in book.reg)
    bearing = (f"{ref('cx.capex_total')}+SUM({P('operating', ':')})-{ref('in.purchase_price')}"
               f"-{ref('in.land_buyout')}{vri}-({credit_capex})")
    param("vat_due", "НДС продаж по ДДУ к уплате в РВЭ",
          f"=IF({ref('in.vat')}>0,MAX({revenue}*{gross}-({bearing})*{gross}*IF({revenue}>0,1,0),0),0)", "млн ₽")
    # LLCR — покрытие ПФ: выручка по ДДУ, расходы на продажу, налог на прибыль,
    # НДС по ДДУ, затраты без объектов вне ДДУ и выдача ПФ против выдачи,
    # процентов и комиссий всего проекта (`llcr` движка).
    credit_cost = _sum([f"SUM({b.r('interest', ':')})+SUM({b.r('fees', ':')})" for b in credit])
    param("llcr_num", "LLCR: числитель",
          f"={revenue}-SUM({P('operating', ':')})-SUM({ref(f'tx.{c}.tax', ':')})-{P('vat_due')}"
          f"-({ref('cx.capex_total')}-({credit_capex}))+{P('pf_draw_total')}", "млн ₽")
    param("llcr_den", "LLCR: знаменатель", f"={P('pf_draw_total')}+{P('financing_cost')}+{credit_cost}", "млн ₽")
    param("llcr", "LLCR проектного финансирования",
          f"=IF({P('llcr_den')}<>0,{P('llcr_num')}/{P('llcr_den')},0)", "×", RATIO, bold=True)


def _finance(book: _Book, rows: _Rows, block: _Block) -> None:
    """НДС и кредит объекта: выдача — доля затрат до ввода, проценты до ввода
    капитализируются, после ввода — погашение по схеме способа продажи."""
    item, o = block.item, block.code
    b, z = block.r, block.z
    i = lambda key: ref(f"in.{o}.{key}")  # noqa: E731
    k = ref("k")
    param, month = _track(rows, block)
    comm, horizon = b("comm"), b("horizon")
    rows.section(f"{item['title']} · условия кредита")
    param("loan_share", "Доля кредита в затратах", f"=MIN(0.95,MAX(0,{i('loan_share')}))", "", PCT)
    param("spread", "Надбавка к ключевой ставке", f"={i('spread')}", "", PCT2)
    param("fee", "Комиссия за выдачу", f"=MAX(0,{i('fee')})", "доля выдачи", PCT2)
    param("reservation", "Плата за резервирование лимита", f"=MAX(0,{ref('in.reservation_fee')})",
          "доля лимита", PCT2)
    param("commitment", "Плата за невыбранный лимит", f"=MAX(0,{ref('in.limit_fee')})", "в год", PCT2)
    param("permit", "Разрешение на стройку — граница двух кредитных линий, № месяца", f"={ref('t.permit')}",
          "№", INT)
    if not block.direct:
        param("repayment", "Погашение после ввода", f"={i('repayment')}", "", "@")
        param("term", "Срок кредита от первой выдачи", f"=MAX(1,INT({i('term')}))*12", "мес.", INT)
        param("balloon_share", "Баллон", f"=MIN(1,MAX(0,{i('balloon')}))", "доля долга на ввод", PCT)
    rows.section(f"{item['title']} · НДС")
    vs = f"{b('vat')}/(1+{b('vat')})"
    money_in = f"({z('sale')}+{z('rent')}+{z('exit')})"
    month("vat_charged", "НДС начислен", f"IF({k}>{horizon},0,{money_in}*{vs})")
    month("vat_pre", "НДС к вычету до зачёта",
          f"IF({k}>{horizon},0,{b('vat_credit', '@p')}+{b('capex')}*{vs}-{b('vat_charged')})", total=False)
    month("vat_paid", "НДС уплачен (минус — возмещён)",
          f"IF(AND({k}={comm},{b('vat_pre')}>0),-{b('vat_pre')},IF({b('vat_pre')}<0,-{b('vat_pre')},0))")
    month("vat_credit", "НДС к вычету на конец месяца", f"{b('vat_pre')}+{b('vat_paid')}", total=False)
    rows.section(f"{item['title']} · кредит")
    month("rate", "Ставка кредита", f"{ref('key_rate')}+{b('spread')}", unit="в год", fmt=PCT2, total=False)
    month("bal_open", "Долг на начало месяца", f"{b('bal_close', '@p')}", total=False)
    month("interest", "Проценты месяца", f"{b('bal_open')}*{b('rate')}/12")
    month("int_cap", "Проценты капитализированы (до ввода)", f"IF({k}<{comm},{b('interest')},0)")
    month("int_paid", "Проценты уплачены (после ввода)", f"IF({k}>={comm},{b('interest')},0)")
    month("draw", "Выдача кредита", f"IF({k}<={comm},({b('capex')}+{b('common')})*{b('loan_share')},0)")
    month("draw_obj", "в т.ч. на стройку объекта", f"IF({k}<={comm},{b('capex')}*{b('loan_share')},0)")
    month("first_k", "Служебная: месяц выдачи", f"IF({b('draw')}>0,{k},1000000)", unit="№", fmt=INT, total=False)
    month("first_build_k", "Служебная: месяц выдачи после разрешения",
          f"IF(AND({b('draw')}>0,{k}>={b('permit')}),{k},1000000)", unit="№", fmt=INT, total=False)
    param("loan_limit", "Лимит кредита = вся плановая выдача", f"=SUM({b('draw', ':')})", "млн ₽")
    param("limit_land", "Линия на участок и проект: выдача до разрешения",
          f'=SUMIF({ref("k", ":")},"<"&{b("permit")},{b("draw", ":")})', "млн ₽")
    param("limit_build", "Строительная линия: выдача с разрешения", f"={b('loan_limit')}-{b('limit_land')}", "млн ₽")
    param("first_draw", "Первая выдача, № месяца", f"=MIN({b('first_k', ':')})", "№", INT)
    param("first_land", "Открытие линии на участок, № месяца",
          f"=IF({b('first_draw')}<{b('permit')},{b('first_draw')},1000000)", "№", INT)
    param("first_build", "Открытие строительной линии, № месяца", f"=MIN({b('first_build_k', ':')})", "№", INT)
    param("reservation_fee", "Плата за резервирование лимитов",
          f"=({b('limit_land')}+{b('limit_build')})*{b('reservation')}", "млн ₽")
    # Невыбранный лимит — с первой выдачи до ввода. Две линии: до разрешения —
    # участок и проект, с разрешения — стройка; каждая платит за невыбранный
    # остаток своего лимита.
    month("commitment_fee", "Плата за невыбранный лимит",
          f"(IF(AND({k}>={b('first_land')},{k}<{b('permit')}),MAX(0,{b('limit_land')}-SUM({b('draw', ':k')})),0)"
          f"+IF({k}>={b('first_build')},MAX(0,{b('limit_build')}-(SUM({b('draw', ':k')})-{b('limit_land')})),0))"
          f"*{b('commitment')}/12")
    month("fees", "Комиссии банка",
          f"{b('draw')}*{b('fee')}+IF({k}={b('first_land')},{b('limit_land')}*{b('reservation')},0)"
          f"+IF({k}={b('first_build')},{b('limit_build')}*{b('reservation')},0)+{b('commitment_fee')}")
    month("bal_after", "Долг до погашения", f"{b('bal_open')}+{b('int_cap')}+{b('draw')}", total=False)
    # Долг стройки без погашений: до ввода погашений нет, и баллон берётся
    # отсюда — без ссылки на месяцы после ввода, которые сами читают баллон.
    month("bal_build", "Служебная: долг стройки до погашений",
          f"IF({k}<={comm},{b('bal_build', '@p')}+IF({k}<{comm},{b('bal_build', '@p')}*{b('rate')}/12,0)"
          f"+{b('draw')},{b('bal_build', '@p')})", total=False)
    month("cash", "Деньги объекта до погашения",
          f"{z('sale')}+{z('rent')}+{z('exit')}-{z('selling')}-{z('opex')}-{z('ptax')}-{z('exit_cost')}"
          f"-{b('vat_paid')}-{b('int_paid')}")
    if block.direct:
        month("repay", "Погашение из выручки",
              f"IF({b('bal_after')}<=0,0,IF({k}={horizon},{b('bal_after')},MIN({b('bal_after')},MAX(0,{b('cash')}))))")
    else:
        param("maturity", "Срок кредита, № месяца",
              f"=MAX(IF({b('first_draw')}<1000000,{b('first_draw')},{comm})+{b('term')},{comm})", "№", INT)
        param("comm_balance", "Долг на ввод", f"=MAX({b('bal_build', ':')})", "млн ₽")
        param("balloon", "Баллон", f"={b('comm_balance')}*{b('balloon_share')}", "млн ₽")
        r12, left = f"{b('rate')}/12", f"({b('maturity')}-{k})"
        mode = lambda code: f'{b("repayment")}="{REPAY[code]}"'  # noqa: E731
        month("annuity", "Тело по аннуитету к баллону",
              f"IF(AND({k}>{comm},{k}<{b('maturity')}),IF({b('rate')}>0,MIN({b('bal_after')},MAX(0,"
              f"({b('bal_after')}-{b('balloon')}/(1+{r12})^{left})*({r12})/(1-(1+{r12})^(-{left}))"
              f"-{b('interest')})),MIN({b('bal_after')},MAX(0,({b('bal_after')}-{b('balloon')})/{left}))),0)")
        month("repay", "Погашение",
              f"IF({b('bal_after')}<=0,0,IF({k}>={comm},IF({k}={horizon},{b('bal_after')},"
              f"IF({mode(ns.REPAY_SWEEP)},MIN({b('bal_after')},MAX(0,{b('cash')})),"
              f"IF({mode(ns.REPAY_ANNUITY)},IF({k}={comm},IF({b('maturity')}<={comm},{b('bal_after')},0),"
              f"IF({k}>={b('maturity')},{b('bal_after')},{b('annuity')})),0))),0))", bold=True)
        month("amort", "Плановое тело (для DSCR)",
              f"IF(AND({mode(ns.REPAY_ANNUITY)},{k}>{comm},{k}<{b('maturity')},{k}<>{horizon}),{b('repay')},0)")
        # Погашение по причине (`object_financing`): баллон в срок, остаток
        # при выходе, из денег объекта.
        month("part_balloon", "Погашение: баллон в срок",
              f"IF(AND({b('repay')}>0,{mode(ns.REPAY_ANNUITY)},{k}={b('maturity')}),{b('repay')}-{b('amort')},0)")
        month("part_exit", "Погашение: остаток при выходе",
              f"IF(AND({b('repay')}>0,{k}={horizon},{b('part_balloon')}=0),{b('repay')}-{b('amort')},0)")
    month("bal_close", "Долг на конец месяца", f"{b('bal_after')}-{b('repay')}", total=False, bold=True)
    param("peak", "Пик долга", f"=MAX({b('bal_close', ':')})", "млн ₽", bold=True)
    month("to_equity", "Деньги объекта собственнику (без затрат стройки)",
          f"{z('sale')}+{z('rent')}+{z('exit')}+{z('residual')}-{z('selling')}-{z('opex')}-{z('ptax')}"
          f"-{z('exit_cost')}-{b('vat_paid')}-{b('int_paid')}-{b('fees')}+{b('draw')}-{b('repay')}")
    if block.direct:
        return
    rows.section(f"{item['title']} · покрытие платежей по годам эксплуатации")
    year = b("year")
    month("year_end", "Последний месяц года эксплуатации",
          f"IF(AND({year}>0,{b('year', '@n')}<>{year}),1,0)", unit="0/1", fmt=FLAG, total=False)
    year_sum = lambda key: f"SUMIF({b('year', ':')},{year},{b(key, ':')})"  # noqa: E731
    month("dscr", "DSCR года: NOI / (проценты + плановое тело)",
          f'IF(AND({b("year_end")}=1,{year_sum("int_paid")}>0),{year_sum("noi")}'
          f'/({year_sum("int_paid")}+{year_sum("amort")}),"")', unit="×", fmt=RATIO, total=False)
    month("icr", "ICR года: NOI / проценты",
          f'IF(AND({b("year_end")}=1,{year_sum("int_paid")}>0),{year_sum("noi")}/{year_sum("int_paid")},"")',
          unit="×", fmt=RATIO, total=False)
    param("dscr_min", "DSCR — минимум по годам",
          f'=IF(COUNT({b("dscr", ":")})>0,MIN({b("dscr", ":")}),"")', "×", RATIO, bold=True)
    param("icr_min", "ICR — минимум по годам",
          f'=IF(COUNT({b("icr", ":")})>0,MIN({b("icr", ":")}),"")', "×", RATIO)


# --- способ продажи: налоги, поток проекта и собственника ---------------------------------

def section_of(path: dict[str, str]) -> str:
    """Раздел книги и «ПРОВЕРОК» у способа продажи."""
    return f"Вариант «{path['label']}»"


def _tax_path(book: _Book, rows: _Rows, path: dict[str, str], blocks: list[_Block],
              ddu: list[_Block]) -> None:
    """Налог на прибыль проекта и НДС способа (`_profit_tax_schedule`
    спецзастройщика): до РВЭ откладывается только маржа — она приходит в месяц
    РВЭ, — проценты и комиссии остаются расходом месяца и дают убыток своих
    лет; убыток прошлых лет гасит не больше доли базы года. Объекты по ДДУ
    дают маржу своими продажами, НДС по ДДУ уходит из базы, проценты и
    комиссии БРИДЖа и ПФ — вычет месяца начисления."""
    c = f"tx.{path['code']}"
    t = lambda key, mod="": ref(f"{c}.{key}", mod)  # noqa: E731
    pf = [f"pf.{path['code']}"] if ddu else []
    k = ref("k")
    rows.section(section_of(path), path["tab"])
    # Маркетинг и брокеры продаж по ДДУ — в общем пуле (`commercial_costs`).
    operating = "".join(f"+SUM({ref(name + '.operating', ':')})" for name in pf)
    rows.param(f"{c}.pool", "Общие затраты проекта — признаются в месяц ввода проекта",
               f"=MAX(0,{ref('cx.capex_total')}{operating}-{ref('cx.tax_pool')})", "млн ₽")
    rows.month(f"{c}.vat_charged", "НДС начислен", _sum([b.r("vat_charged") for b in blocks]))
    rows.month(f"{c}.vat_paid", "НДС к уплате (минус — возмещение)",
               _sum([b.r("vat_paid") for b in blocks] + [ref(name + ".vat") for name in pf]))
    rows.month(f"{c}.obj_margin", "Маржа объектов", _sum([b.r("margin") for b in blocks + ddu]))
    rows.month(f"{c}.common", "Общие затраты проекта (в месяц ввода)",
               f"IF({k}={ref('t.rve')},-{t('pool')},0)")
    # НДС продаж по ДДУ к уплате уходит из базы налога на прибыль.
    rows.month(f"{c}.margin", "Маржа проекта",
               f"{t('obj_margin')}+{t('common')}" + "".join(f"-{ref(name + '.vat')}" for name in pf))
    rows.month(f"{c}.financing", "Проценты и комиссии — вычет",
               _sum([f"{b.r('interest')}+{b.r('fees')}" for b in blocks] + [ref(name + ".deduct") for name in pf]))
    rows.month(f"{c}.net", "База месяца", f"{t('margin')}-{t('financing')}")
    rows.month(f"{c}.gate", "До ввода проекта: маржа откладывается", f"IF({k}<{ref('t.rve')},1,0)",
               unit="0/1", fmt=FLAG, total=False)
    rate, limit = ref("in.profit_tax"), ref("in.loss_limit")
    same_year = f"{ref('year')}={ref('year', '@p')}"
    rows.month(f"{c}.year_res", "Результат года нарастающим",
               f"IF({same_year},{t('year_res', '@p')},0)+IF({t('gate')}=1,-{t('financing')},"
               f"{t('net')}+IF({t('gate', '@p')}=1,SUM({t('margin', ':p')}),0))",
               first=f"IF({t('gate')}=1,-{t('financing')},{t('net')})", total=False)
    rows.month(f"{c}.prior", "Убыток прошлых лет",
               f"IF({ref('year')}<>{ref('year', '@p')},{t('prior', '@p')}+IF({t('year_res', '@p')}<0,"
               f"-{t('year_res', '@p')},-{t('used', '@p')}),{t('prior', '@p')})", first="0", total=False)
    rows.month(f"{c}.used", "Зачтено убытка",
               f"IF(AND({t('year_res')}>0,{t('prior')}>0),MIN({t('prior')},{t('year_res')}*{limit}),0)")
    rows.month(f"{c}.base", "База года", f"MAX({t('year_res')}-{t('used')},0)", total=False)
    rows.month(f"{c}.paid", "Налог года, начисленный ранее",
               f"IF({same_year},{t('paid', '@p')}+{t('tax', '@p')},0)", first="0", total=False)
    rows.month(f"{c}.tax", "Налог на прибыль начислен",
               f"IF({t('gate')}=1,0,MAX({t('base')}*{rate}-{t('paid')},0))", bold=True)
    # Уплата: авансы по итогам квартала (апрель, июль, октябрь) и доплата года
    # в марте следующего — `profit_tax_payments`.
    moy = ref("moy")
    rows.month(f"{c}.pay", "Налог на прибыль к уплате",
               f"IF({moy}=3,{t('tax', '@-5')}+{t('tax', '@-4')}+{t('tax', '@-3')},"
               f"IF(OR({moy}=4,{moy}=7,{moy}=10),{t('tax', '@-3')}+{t('tax', '@-2')}+{t('tax', '@-1')},0))",
               bold=True)


def _annual_irr(flows: str) -> str:
    """Годовой IRR месячного потока. Excel и LibreOffice ищут IRR Ньютоном от
    догадки, и у части потоков (собственник ТЦ по ДДУ) от догадки 0,01 не
    сходятся или уходят в ложный корень на бесконечности. Догадки по очереди — 0,01, −0,1, −0,3; ответ больше 100% в
    месяц — ложный корень, берётся следующая догадка; ни одна не сошлась — пусто."""
    formula = '""'
    for guess in ("-0.3", "-0.1", "0.01"):
        irr = f"IRR({flows},{guess})"
        formula = f"IFERROR(IF(ABS({irr})<1,(1+{irr})^12-1,1/0),{formula})"
    return formula


def _cash_path(book: _Book, rows: _Rows, path: dict[str, str], blocks: list[_Block],
               offices: list[_Block], ddu: list[_Block], spec: dict[str, Any], capex: dict[str, Any],
               months: int) -> None:
    """Поток проекта (NPV, IRR — до финансирования, как у движка) и
    собственника (NPV по ставке / 12, IRR капитала), раскладка вложений.

    `blocks` — объекты на своём кредите, `ddu` — объекты по ДДУ этого способа:
    их выручка идёт в поток проекта и собственнику с раскрытия эскроу
    (214-ФЗ), а проданное после ввода — в месяц продажи."""
    c = path["code"]
    p = lambda key, mod="": ref(f"{c}.{key}", mod)  # noqa: E731
    k = ref("k")
    tx = lambda key, mod="": ref(f"tx.{c}.{key}", mod)  # noqa: E731
    over = lambda key: _sum([b.z(key) for b in blocks])  # noqa: E731
    pf = (lambda key: [ref(f"pf.{c}.{key}")]) if ddu else (lambda key: [])  # noqa: E731
    pf_param = lambda key: ref(f"pf.{c}.{key}") if ddu else '""'  # noqa: E731
    rows.section("Проект: деньги по месяцам", path["tab"])
    rows.month(f"{c}.revenue", "Выручка", _sum([f"{b.z('sale')}+{b.z('rent')}+{b.z('exit')}+{b.z('residual')}"
                                              for b in blocks] + pf("sales")), bold=True)
    rows.month(f"{c}.capex", "Затраты на создание", ref("cx.total"))
    rows.month(f"{c}.costs", "Расходы продаж и эксплуатации, налог на имущество, сделка",
               _sum([f"{b.z('selling')}+{b.z('opex')}+{b.z('ptax')}+{b.z('exit_cost')}" for b in blocks]
                    + pf("operating")))
    rows.month(f"{c}.obj_vat", "НДС объектов на своём кредите (минус — возмещение)", over("vat_paid"))
    rows.month(f"{c}.vat", "НДС к уплате (минус — возмещение)", _sum([p("obj_vat")] + pf("vat")))
    rows.month(f"{c}.interest", "Проценты начислены", _sum([over("interest")] + pf("interest")))
    rows.month(f"{c}.fees", "Комиссии банка", _sum([over("fees")] + pf("fees")))
    rows.month(f"{c}.tax", "Налог на прибыль к уплате", tx("pay"))
    rows.month(f"{c}.cash_in", "Выручка деньгами: по ДДУ — с раскрытия эскроу",
               _sum([f"{b.z('sale')}+{b.z('rent')}+{b.z('exit')}+{b.z('residual')}"
                     for b in blocks] + pf("cash")))
    rows.month(f"{c}.project", "Поток проекта до кредита",
               f"{p('cash_in')}-{p('costs')}-{p('vat')}-{p('capex')}-{p('tax')}", bold=True)
    rows.section("Собственник: деньги по месяцам", path["tab"])
    rows.month(f"{c}.to_equity", "Деньги объектов, ДДУ и ПФ собственнику (без затрат стройки)",
               _sum([over("to_equity")] + pf("equity")))
    rows.month(f"{c}.equity", "Поток собственника", f"{p('to_equity')}-{p('capex')}-{p('tax')}", bold=True)
    rows.month(f"{c}.residual", "Оценка удержанного здания (не деньги)", over("residual"))
    rows.month(f"{c}.cash", "Поток собственника — только деньги", f"{p('equity')}-{p('residual')}")
    rows.month(f"{c}.cum", "Поток собственника нарастающим", f"{p('cum', '@p')}+{p('cash')}", total=False)
    disc = f"MAX({ref('in.discount')},-0.999999)"
    rows.month(f"{c}.df", "Множитель дисконтирования проекта: (1 + ставка)^(месяц / 12)",
               f"1/(1+{disc})^({k}/12)", unit="×", fmt="0.0000", total=False)
    rows.month(f"{c}.dfe", "Множитель дисконтирования собственника: ставка / 12",
               f"1/(1+{disc}/12)^({k})", unit="×", fmt="0.0000", total=False)
    rows.month(f"{c}.contributed", "Вложено собственником", f"MAX(0,-{p('cash')})")
    rows.month(f"{c}.returned", "Получено собственником", f"MAX(0,{p('cash')})")
    income = [b for b in blocks if not b.direct]
    exits = [f"{k}={b.r('horizon')}" for b in income]
    rows.month(f"{c}.exit", "Месяц выхода из удержания", f"IF(OR({','.join(exits)}),1,0)" if exits else "0",
               unit="0/1", fmt=FLAG, total=False)
    rows.month(f"{c}.in_build", "Вложено до ввода", f"IF({k}<={p('comm_last')},{p('contributed')},0)")
    rows.month(f"{c}.in_exit", "Вложено при выходе",
               f"IF(AND({k}>{p('comm_last')},{p('exit')}=1),{p('contributed')},0)")
    rows.month(f"{c}.out_vat", "Возмещение НДС стройки",
               f"IF({p('cash')}>0,MIN({p('cash')},MAX(0,-{p('obj_vat')})),0)")
    rows.month(f"{c}.out_exit", "Получено при выходе", f"IF({p('exit')}=1,{p('returned')}-{p('out_vat')},0)")
    rows.month(f"{c}.neg_k", "Служебная: накопленный поток меньше нуля", f"IF({p('cum')}<-0.0000005,{k},-1)",
               unit="№", fmt=INT, total=False)

    rows.section("Итоги варианта", path["tab"])
    param = lambda key, *a, **kw: rows.param(f"{c}.{key}", *a, **kw)  # noqa: E731
    total = lambda key: f"SUM({p(key, ':')})"  # noqa: E731
    articles = [ref(f"cx.{key}.amt") for key in capex["order"]]
    # Граница «до ввода» собственного участия — ввод последнего объекта на
    # своём кредите (`equity_participation_report`); без них — объектов по ДДУ.
    comm = [b.r("comm") for b in blocks] or [b.r("comm") for b in ddu] or [ref("t.comm_last")]
    param("comm_last", "Ввод последнего объекта на своём кредите, № месяца", f"=MAX({','.join(comm)})", "№", INT)
    param("t_revenue", "Выручка", f"={total('revenue')}", "млн ₽", bold=True)
    param("t_capex", "Затраты на создание (CAPEX, с НДС)", f"={total('capex')}", "млн ₽")
    param("t_costs", "Расходы продаж и эксплуатации, налог на имущество, сделка", f"={total('costs')}", "млн ₽")
    param("t_commercial", "в т.ч. маркетинг и брокеры продаж по ДДУ",
          "=" + (f"SUM({ref(f'pf.{c}.operating', ':')})" if ddu else "0"), "млн ₽")
    param("t_nonres_costs", "в т.ч. расходы объектов на своём кредите", f"={p('t_costs')}-{p('t_commercial')}",
          "млн ₽")
    param("t_ebitda", "EBITDA", f"={p('t_revenue')}-{p('t_capex')}-{p('t_costs')}", "млн ₽")
    param("t_financing", "Проценты и комиссии", f"={total('interest')}+{total('fees')}", "млн ₽")
    param("t_pbt", "Прибыль до налога", f"={p('t_ebitda')}-{p('t_financing')}", "млн ₽")
    param("t_tax", "Налог на прибыль", f"=SUM({tx('tax', ':')})", "млн ₽")
    param("t_vat", "НДС к уплате", f"=SUM({tx('vat_paid', ':')})", "млн ₽")
    param("t_net", "Чистая прибыль", f"={p('t_pbt')}-{p('t_tax')}-{p('t_vat')}", "млн ₽", bold=True)
    param("t_expenses", "Расходы всего: статьи сметы, расходы объектов, проценты, налог, НДС",
          f"={_sum(articles)}+{p('t_costs')}+{p('t_financing')}+{p('t_tax')}+{p('t_vat')}", "млн ₽")
    param("id_total", "Выручка − расходы всего − чистая прибыль (должно быть 0)",
          f"={p('t_revenue')}-{p('t_expenses')}-{p('t_net')}", "млн ₽")
    param("id_ebitda", "EBITDA − проценты − прибыль до налога (должно быть 0)",
          f"={p('t_ebitda')}-{p('t_financing')}-{p('t_pbt')}", "млн ₽")
    param("npv", "NPV проекта", f"=SUMPRODUCT({p('project', ':')},{p('df', ':')})", "млн ₽", bold=True)
    param("irr", "IRR проекта (до кредита)", "=" + _annual_irr(p("project", ":")),
          "в год", PCT, bold=True)
    # IRR и NPV собственника не числа, если ПФ к концу проекта не погашен
    # (решение владельца 2 от 29.09.2026): порог и подпись — движка.
    na = spec.get("equity_na") or {}

    def unrepaid(formula: str) -> str:
        if not ddu or not na:
            return formula
        return (f'IF({ref(f"pf.{c}.ending_pf")}>{float(na["min_ending_pf_rub"]) / _MLN:g},'
                f'"{na["label"]}",{formula})')
    param("npv_eq", "NPV собственника", "=" + unrepaid(f"SUMPRODUCT({p('equity', ':')},{p('dfe', ':')})"),
          "млн ₽", bold=True)
    # IRR капитала — числом, только если своих денег не меньше порога движка
    # от полных расходов (владелец, 16.10.2026).
    irr_eq = _annual_irr(p("equity", ":"))
    min_share = (spec.get("equity_irr") or {}).get("min_share")
    if min_share is not None:
        full = f"({p('t_capex')}+{p('t_costs')}+{p('t_financing')}+{p('t_tax')}+{p('t_vat')})"
        own = f'(-SUMIF({p("equity", ":")},"<0"))'
        irr_eq = (f'IF(AND({full}>0,{own}<{float(min_share):g}*{full}),"не показателен: своих денег "'
                  f'&{grouped(own)}&" млн ₽ ("&SUBSTITUTE(ROUND({own}/{full}*100,1),".",",")&"% затрат)",{irr_eq})')
    param("irr_eq", "IRR собственника", "=" + unrepaid(irr_eq), "в год", PCT)
    param("peak", "Пик собственных денег", f"=-MIN(0,MIN({p('cum', ':')}))", "млн ₽", bold=True)
    loan_peaks = [b.r("peak") for b in blocks]
    param("loan_peak", "Пик долга объектов на своём кредите (сумма пиков)", f"={_sum(loan_peaks)}", "млн ₽")
    dscr = [b.r("dscr_min") for b in income]
    param("dscr_min", "DSCR кредита — минимум по годам и объектам",
          f'=IF(COUNT({",".join(dscr)})>0,MIN({",".join(dscr)}),"")' if dscr else '=""', "×", RATIO)
    param("llcr", "LLCR проектного финансирования", "=" + pf_param("llcr"), "×", RATIO)
    param("pf_peak", "Пик долга проектного финансирования", "=" + pf_param("peak_pf"), "млн ₽")
    param("escrow_peak", "Пик денег на эскроу", "=" + pf_param("peak_escrow"), "млн ₽")
    horizons = "".join(f",{b.r('horizon')}" for b in blocks + ddu)
    param("horizon", f"Горизонт варианта, мес. (месяцев в таблицах: {book.n})",
          f"=MAX({ref('t.rve')}+MAX(INT({ref('in.residual_months')})+3,12){horizons},{ref('cx.plan_last')})+1",
          "мес.", INT)
    param("rows_ok", "Месяцев в таблицах хватает на вариант и год покупателя после него",
          f'=IF({p("horizon")}+{EXTRA_MONTHS}<={book.n},"хватает","НЕ ХВАТАЕТ")', "", "@")

    if not book.project_only:
        _office_totals(rows, path, offices, ddu)

    rows.section("Собственное участие", path["tab"])
    param("e_in_build", "Вложено до ввода", f"={total('in_build')}", "млн ₽")
    param("e_contributed", "Вложено всего", f"={total('contributed')}", "млн ₽", bold=True)
    param("e_in_exit", "в т.ч. при выходе — довнесение на погашение долга", f"={total('in_exit')}", "млн ₽")
    param("e_in_operation", "в т.ч. после ввода — дефицит",
          f"={p('e_contributed')}-{p('e_in_build')}-{p('e_in_exit')}", "млн ₽")
    param("e_returned", "Получено всего (деньгами)", f"={total('returned')}", "млн ₽", bold=True)
    param("e_out_vat", "в т.ч. возмещение НДС стройки", f"={total('out_vat')}", "млн ₽")
    param("e_out_exit", "в т.ч. при выходе — после погашения долга", f"={total('out_exit')}", "млн ₽")
    param("e_out_running", "в т.ч. по ходу — продажи, аренда после долга и налогов",
          f"={p('e_returned')}-{p('e_out_vat')}-{p('e_out_exit')}", "млн ₽")
    param("e_residual", "Оценка удержанного здания — не деньги", f"={total('residual')}", "млн ₽")
    param("e_net", "Чистый денежный результат (получено − вложено)",
          f"={p('e_returned')}-{p('e_contributed')}", "млн ₽")
    param("e_multiple", "Мультипликатор капитала (получено / вложено)",
          f'=IF({p("e_contributed")}>0,{p("e_returned")}/{p("e_contributed")},"")', "×", RATIO)
    param("e_irr_cash", "IRR собственника — только деньги",
          "=" + unrepaid(f'IF({p("e_residual")}>0,{_annual_irr(p("cash", ":"))},"")'),
          "в год", PCT)
    last_neg = f"MAX({p('neg_k', ':')})"
    param("e_payback", "Окупаемость собственных денег",
          f'=IF(AND({last_neg}>=0,{last_neg}+1<={months - 1}),EDATE({ref("in.start")},{last_neg}+1),'
          f'"не окупается за горизонт проекта")', "", MONTH)


def _office_totals(rows: _Rows, path: dict[str, str], offices: list[_Block], ddu: list[_Block]) -> None:
    """Офис в способе продажи: его выручка, расходы, аренда и кредит — для «ОТЧЕТА»
    и сверки с `office_paths`."""
    c = path["code"]
    p = lambda key, mod="": ref(f"{c}.{key}", mod)  # noqa: E731
    param = lambda key, *a, **kw: rows.param(f"{c}.{key}", *a, **kw)  # noqa: E731
    rows.section("Офис в этом варианте", path["tab"])
    office_sum = lambda key: _sum([f"SUM({b.r(key, ':')})" for b in offices if b.n(key) in b.names])  # noqa: E731
    ddu_offices = [b for b in ddu if b.item.get("office")]
    ddu_sum = lambda key: _sum([f"SUM({b.r(key, ':')})" for b in ddu_offices])  # noqa: E731
    param("o_sale", "Выручка продаж", f"={office_sum('sale')}+{ddu_sum('revenue')}", "млн ₽")
    param("o_rent", "Доход аренды", f"={office_sum('rent')}", "млн ₽")
    param("o_exit", "Продажа здания", f"={office_sum('exit')}", "млн ₽")
    param("o_selling", "Маркетинг и брокеры", f"={office_sum('selling')}+{ddu_sum('operating')}", "млн ₽")
    param("o_operating", "Эксплуатация, налог на имущество, расходы на сделку",
          f"={_sum([office_sum('opex'), office_sum('ptax'), office_sum('exit_cost')])}", "млн ₽")
    rent_blocks = [b for b in offices if not b.direct]
    param("o_noi", "NOI года после заполнения", "=" + _sum([b.r("noi_stab") for b in rent_blocks]), "млн ₽")
    noi_cost = _sum([f"IF({b.r('noi_stab')}<>0,{b.r('cost_total')},0)" for b in rent_blocks])
    param("o_yoc", "Доходность на затраты", f"=IF(AND({p('o_noi')}<>0,({noi_cost})>0),{p('o_noi')}/({noi_cost}),\"\")"
          if rent_blocks else '=""', "", PCT)
    param("o_cap", "Ставка капитализации", "=" + (rent_blocks[0].r("cap") if rent_blocks else '""'), "", PCT)
    param("o_exit_value", "Стоимость здания", "=" + (_sum([b.r("exit_value") for b in rent_blocks])
                                                     if rent_blocks else '""'), "млн ₽")
    o_dscr = [b.r("dscr_min") for b in rent_blocks]
    param("o_dscr", "DSCR кредита офиса — минимум",
          f'=IF(COUNT({",".join(o_dscr)})>0,MIN({",".join(o_dscr)}),"")' if o_dscr else '=""', "×", RATIO)
    param("o_loan_peak", "Пик долга офиса на своём кредите", "=" + _sum([b.r("peak") for b in offices]), "млн ₽")


def _block_totals(rows: _Rows, block: _Block) -> None:
    """Итоги объекта в способе продажи — те же величины, что `object_flows` отдаёт в
    `totals`, `kpi` и `financing`: по ним «ПРОВЕРКИ» сверяют объект с движком."""
    b = block.r
    param, _ = _track(rows, block)
    s = lambda key: f"SUM({b(key, ':')})" if block.n(key) in block.names else "0"  # noqa: E731
    k_all = ref("k", ":")
    rows.section(f"{block.item['title']} · итоги объекта")
    param("T.revenue", "Выручка объекта", f"={s('sale')}+{s('rent')}+{s('exit')}+{s('residual')}", "млн ₽")
    param("T.loan_draw", "Кредит — выдача", f"={s('draw')}", "млн ₽")
    param("T.loan_interest", "Кредит — проценты", f"={s('interest')}", "млн ₽")
    param("T.loan_fee", "Кредит — комиссии всего", f"={s('fees')}", "млн ₽")
    param("T.loan_repayment", "Кредит — погашено", f"={s('repay')}", "млн ₽")
    param("T.draw_object", "Выдача на стройку объекта", f"=MIN({s('draw_obj')},{s('draw')})", "млн ₽")
    param("T.draw_common", "Выдача на общие затраты", f"=MAX(0,{s('draw')}-{s('draw_obj')})", "млн ₽")
    param("T.issue_fee", "Комиссия за выдачу",
          f"={s('fees')}-{b('reservation_fee')}-{s('commitment_fee')}", "млн ₽")
    param("T.interest_capitalized", "Проценты до ввода — капитализированы", f"={s('int_cap')}", "млн ₽")
    param("T.interest_paid", "Проценты после ввода — уплачены", f"={s('int_paid')}", "млн ₽")
    param("T.debt_at_commissioning", "Долг на ввод",
          f"=SUMIF({k_all},{b('comm')},{b('bal_close', ':')})", "млн ₽")
    param("T.rate_at_commissioning", "Ставка на вводе", f"=SUMIF({k_all},{b('comm')},{b('rate', ':')})", "", PCT2)
    param("T.avg_rate", "Средняя ставка по остатку долга",
          f"=IF({s('bal_close')}>0,SUMPRODUCT({b('rate', ':')},{b('bal_close', ':')})/{s('bal_close')},0)", "", PCT2)
    if block.direct:
        for key, label in (("balloon_planned", "Баллон по графику"), ("balloon_paid", "Баллон погашен в срок"),
                           ("repaid_scheduled", "Погашено плановым телом"),
                           ("repaid_at_exit", "Погашено при выходе"), ("avg_payment", "Аннуитетный платёж — средний")):
            param(f"T.{key}", label, "=0", "млн ₽")
        param("T.repaid_from_cash", "Погашено из выручки", f"={s('repay')}", "млн ₽")
        return
    annuity = f'{b("repayment")}="{REPAY[ns.REPAY_ANNUITY]}"'
    count = f'COUNTIF({b("amort", ":")},">0")'
    param("T.balloon_planned", "Баллон по графику", f"=IF({annuity},{b('balloon')},0)", "млн ₽")
    param("T.balloon_paid", "Баллон погашен в срок", f"={s('part_balloon')}", "млн ₽")
    param("T.repaid_scheduled", "Погашено плановым телом", f"={s('amort')}", "млн ₽")
    param("T.repaid_at_exit", "Погашено при выходе", f"={s('part_exit')}", "млн ₽")
    param("T.repaid_from_cash", "Погашено из NOI",
          f"={s('repay')}-{s('amort')}-{s('part_balloon')}-{s('part_exit')}", "млн ₽")
    param("T.avg_payment", "Аннуитетный платёж — средний",
          f"=IF({count}>0,SUMPRODUCT(({b('amort', ':')}>0)*({b('int_paid', ':')}+{b('amort', ':')}))/{count},0)",
          "млн ₽")


# --- «ОТЧЕТ» ------------------------------------------------------------------

def _report_sheet(book: _Book, spec: dict[str, Any], capex: dict[str, Any], current: str | None,
                  years: list[int], current_label: str = "") -> dict[str, Any]:
    title = ("DEVELOPAID · ОТЧЁТ · НЕЖИЛОЙ ПРОЕКТ" if book.project_only
             else "DEVELOPAID · ОТЧЁТ · ТРИ СПОСОБА ПРОДАТЬ ОФИС")
    ws = book.sheet(REPORT_SHEET, title, "Все числа — ссылки на листы модели, млн ₽ с НДС.", tab=DARK, width=12)
    PATHS = book.paths  # noqa: N806 — способы этой книги
    ws.column_dimensions["A"].width = 60
    for c in "BCDEFGHIJKLM":
        ws.column_dimensions[c].width = 17
    out: dict[str, Any] = {"years": {}}
    r = 4
    width = 1 + len(PATHS)
    fills = (HEAD, *(path["tab"] for path in PATHS))
    labels = tuple(path["label"] for path in PATHS)

    def head(cells: tuple[str, ...], fills: tuple[str, ...] | None = None) -> None:
        nonlocal r
        for j, text in enumerate(cells):
            cell = ws.cell(r, 1 + j, text)
            cell.font = _fnt(b=True, color="FFFFFF")
            cell.fill = _fill((fills or ())[j] if fills and j < len(fills) else HEAD)
        r += 1

    def section(text: str) -> None:
        nonlocal r
        r += 1
        ws.cell(r, 1, text).font = _fnt(b=True, color=DARK)
        for j in range(1, width + 1):
            ws.cell(r, j).fill = _fill(LIGHT)
        r += 1

    def line(key: str, label: str, cells: tuple[str, ...], fmt: str = NUM, bold: bool = False) -> None:
        nonlocal r
        ws.cell(r, 1, label).font = _fnt(b=bold)
        for j, value in enumerate(cells):
            cell = ws.cell(r, 2 + j)
            if isinstance(value, str) and value.startswith("="):
                book.pending.append((cell, value[1:], None, REPORT_SHEET))
            else:
                cell.value = value
            cell.number_format = fmt
            cell.font = _fnt(b=bold)
            cell.alignment = Alignment(horizontal="right")
        for j in range(1, width + 1):
            ws.cell(r, j).border = Border(bottom=_thin)
            if bold:
                ws.cell(r, j).fill = _fill(LIGHT)
        for j, path in enumerate(PATHS):
            book.register(f"rep.{path['code']}.{key}", REPORT_SHEET, r, L(2 + j))
        r += 1

    every = lambda key: tuple("=" + ref(p["code"] + "." + key) for p in PATHS)  # noqa: E731
    head(("Показатель", *labels), fills)
    out["head_row"] = r - 1
    current_label = current_label or next((p["label"] for p in PATHS if p["key"] == current),
                                          "у офисов проекта разные способы")
    ws.cell(r, 1, "Способ, выбранный в проекте")
    ws.cell(r, 2, current_label).font = _fnt(b=True)
    r += 1
    if not book.project_only:
        ws.cell(r, 1, "Выгоднее по NPV проекта").font = _fnt(b=True)
        cell = ws.cell(r, 2)
        book.pending.append((cell, _better_formula(), None, REPORT_SHEET))
        cell.font = _fnt(b=True)
        book.register("rep.better", REPORT_SHEET, r, "B")
        r += 1
    section("Экономика")
    for key, label, fmt, bold in (
            ("t_revenue", "Выручка", NUM, False), ("t_capex", "Затраты на создание с участком", NUM, False),
            ("t_costs", "Расходы продаж и эксплуатации, налог на имущество, сделка", NUM, False),
            ("t_ebitda", "EBITDA", NUM, False), ("t_financing", "Проценты и комиссии банка", NUM, False),
            ("t_pbt", "Прибыль до налога", NUM, False), ("t_tax", "Налог на прибыль", NUM, False),
            ("t_vat", "НДС к уплате", NUM, False), ("t_net", "Чистая прибыль", NUM, True),
            ("npv", "NPV проекта", NUM, True), ("irr", "IRR проекта (до кредита)", PCT, False),
            ("npv_eq", "NPV собственника", NUM, False), ("irr_eq", "IRR собственника", PCT, False),
            ("peak", "Пик собственных денег", NUM, False),
            ("loan_peak", "Пик долга объектов на своём кредите", NUM, False),
            ("dscr_min", "DSCR кредита — минимум по годам", RATIO, False),
            ("llcr", "LLCR проектного финансирования", RATIO, False),
            ("pf_peak", "Пик долга проектного финансирования", NUM, False),
            ("escrow_peak", "Пик денег покупателей на эскроу", NUM, False)):
        line(key, label, every(key), fmt, bold)
    if not book.project_only:
        section("Офис")
        rent_only = ("o_rent", "o_exit", "o_noi", "o_yoc", "o_cap", "o_exit_value", "o_dscr")
        for key, label, fmt in (
                ("o_sale", "Выручка продаж (по ДДУ или блоками)", NUM), ("o_rent", "Доход аренды", NUM),
                ("o_exit", "Продажа здания", NUM), ("o_selling", "Маркетинг и брокеры", NUM),
                ("o_operating", "Эксплуатация, налог на имущество, расходы на сделку", NUM),
                ("o_noi", "NOI года после заполнения", NUM), ("o_yoc", "Доходность на затраты", PCT),
                ("o_cap", "Ставка капитализации", PCT), ("o_exit_value", "Стоимость здания при продаже", NUM),
                ("o_dscr", "DSCR кредита офиса — минимум", RATIO),
                ("o_loan_peak", "Пик долга офиса на своём кредите", NUM)):
            cells = list(every(key))
            for j, path in enumerate(PATHS):
                rent = path["key"] == ns.STRATEGY_INCOME
                if (key in rent_only and not rent) or (key in ("o_sale", "o_selling") and rent) \
                        or (key == "o_loan_peak" and path["key"] == ns.STRATEGY_DDU):
                    cells[j] = "не применяется"
            line(key, label, tuple(cells), fmt)
    section("Как сложены итоги")
    line("id_total", "Выручка − расходы всего − чистая прибыль (должно быть 0)", every("id_total"))
    line("id_ebitda", "EBITDA − проценты − прибыль до налога (должно быть 0)", every("id_ebitda"))
    ws.cell(r, 1, "EBITDA = выручка − затраты на создание − расходы объектов; без процентов, налога на прибыль "
                  "и НДС. Расходы всего = статьи сметы + расходы объектов + проценты + налог + НДС.").font = _fnt(
        i=True, color=GREY)
    r += 1
    section("Структура расходов")
    for key in capex["order"]:
        label = capex["labels"][key]
        if key.endswith(".garage"):
            continue
        if key.endswith(".building"):
            o = key.split(".")[0]
            formula = f"={ref(f'cx.{o}.building.amt')}+{ref(f'cx.{o}.garage.amt')}"
            label = label.replace(": здание", "")
        else:
            formula = f"={ref(f'cx.{key}.amt')}"
        line(f"art.{key.split('.')[0]}", label, (formula,) * len(PATHS))
    line("s_costs", "Расходы продаж и эксплуатации, налог на имущество, сделка", every("t_costs"))
    line("s_financing", "Проценты и комиссии", every("t_financing"))
    line("s_tax", "Налог на прибыль", every("t_tax"))
    line("s_vat", "НДС", every("t_vat"))
    line("s_total", "Расходы всего", every("t_expenses"), NUM, True)
    section("Собственное участие")
    for key, label in (("e_in_build", "Вложено до ввода"), ("e_contributed", "Вложено всего"),
                       ("e_in_exit", "в т.ч. при выходе — довнесение на погашение долга"),
                       ("e_in_operation", "в т.ч. после ввода — дефицит"),
                       ("peak", "Пик потребности в собственных деньгах"),
                       ("e_returned", "Получено всего (деньгами)"), ("e_out_vat", "в т.ч. возмещение НДС стройки"),
                       ("e_out_exit", "в т.ч. при выходе — после погашения долга"),
                       ("e_out_running", "в т.ч. по ходу — продажи, аренда после долга и налогов"),
                       ("e_residual", "Оценка удержанного здания — не деньги"),
                       ("e_net", "Чистый денежный результат (получено − вложено)"),
                       ("e_multiple", "Мультипликатор капитала (получено / вложено)"),
                       ("e_irr_cash", "IRR собственника — только деньги"),
                       ("e_payback", "Окупаемость собственных денег")):
        fmt = RATIO if key == "e_multiple" else PCT if key == "e_irr_cash" else MONTH if key == "e_payback" else NUM
        line(f"eq.{key}", label, every(key), fmt)

    year_row = ref("year", ":")
    by_year = lambda rng, y: f"SUMIF({year_row},{y},{rng})"  # noqa: E731
    section("По годам: поток собственника")
    head(("Год", *labels), fills)
    out["owner_years"] = (r - 1, r, r + len(years) - 1)
    for y in years:
        ws.cell(r, 1, y)
        for j, path in enumerate(PATHS):
            cell = ws.cell(r, 2 + j)
            book.pending.append((cell, by_year(ref(f"{path['code']}.equity", ":"), f"$A{r}"), None, REPORT_SHEET))
            cell.number_format = NUM
        r += 1
    for path in PATHS:
        c = path["code"]
        section(f"По годам · {path['label']}")
        columns = (("Год", None), ("Выручка", "revenue"), ("Затраты на создание", "capex"),
                   ("Расходы объектов", "costs"), ("Проценты и комиссии", "fin"), ("НДС", "vat"),
                   ("Налог на прибыль", "tax"), ("Поток проекта", "project"),
                   ("Вложено собственником", "contributed"), ("Получено собственником", "returned"),
                   ("Чистый поток собственника", "net"), ("Накопленный на конец года", "cum"))
        head(tuple(text for text, _ in columns), (path["tab"],) * len(columns))
        first = r
        for y in years:
            ws.cell(r, 1, y)
            for j, (_, key) in enumerate(columns[1:], start=2):
                if key == "fin":
                    formula = (f"{by_year(ref(f'{c}.interest', ':'), f'$A{r}')}"
                               f"+{by_year(ref(f'{c}.fees', ':'), f'$A{r}')}")
                elif key == "net":
                    formula = f"J{r}-I{r}"
                elif key == "cum":
                    formula = f"K{r}" if r == first else f"L{r - 1}+K{r}"
                else:
                    formula = by_year(ref(f"{c}.{key}", ":"), f"$A{r}")
                cell = ws.cell(r, j)
                book.pending.append((cell, formula, None, REPORT_SHEET))
                cell.number_format = NUM
            out["years"][(c, y)] = r
            r += 1
        section(f"ОПУ по годам · {path['label']} (налоговая база проекта)")
        columns = ("Год", "Маржа объектов", "Общие затраты (в месяц ввода)", "Проценты и комиссии",
                   "Прибыль до налога", "Налог на прибыль", "Прибыль после налога")
        head(columns, (path["tab"],) * len(columns))
        for y in years:
            ws.cell(r, 1, y)
            for j, key in enumerate(("obj_margin", "common", "financing"), start=2):
                cell = ws.cell(r, j)
                book.pending.append((cell, by_year(ref(f"tx.{c}.{key}", ":"), f"$A{r}"), None, REPORT_SHEET))
                cell.number_format = NUM
            ws.cell(r, 5, f"=B{r}+C{r}-D{r}").number_format = NUM
            cell = ws.cell(r, 6)
            book.pending.append((cell, by_year(ref(f"tx.{c}.tax", ":"), f"$A{r}"), None, REPORT_SHEET))
            cell.number_format = NUM
            ws.cell(r, 7, f"=E{r}-F{r}").number_format = NUM
            r += 1
    ws.cell(r + 1, 1, "Чувствительность: сценарные множители выручки и затрат — на «Вводных»; таблицы данных "
                      "Excel книга не строит: их не пересчитывает проверка книги.").font = _fnt(i=True, color=GREY)
    ws.freeze_panes = "B5"
    return out


def _better_formula() -> str:
    """Название способа с наибольшим NPV проекта; при равенстве — первый по
    порядку `PATHS`, как `max` в `office_paths` движка."""
    npv = [ref(f"{path['code']}.npv") for path in PATHS]
    formula = f'"{PATHS[-1]["label"]}"'
    for j in range(len(PATHS) - 2, -1, -1):
        rest = ",".join(npv[j + 1:])
        formula = f'IF({npv[j]}>=MAX({rest}),"{PATHS[j]["label"]}",{formula})'
    return formula


# --- «Дашборд», «ИНСТРУКЦИЯ», «Источники» --------------------------------------

def _dashboard_sheet(book: _Book, report: dict[str, Any], with_pf: bool = False) -> None:
    PATHS = book.paths  # noqa: N806 — способы этой книги
    if book.project_only:
        ws = book.sheet(DASH_SHEET, "DEVELOPAID · ИНВЕСТИЦИОННЫЙ ДАШБОРД · НЕЖИЛОЙ ПРОЕКТ", "", tab=DARK, width=12)
        book.pending.append((ws["A2"], f'{ref("in.name")}&" · каждый объект продаётся своей стратегией"',
                             None, DASH_SHEET))
    else:
        ws = book.sheet(DASH_SHEET, "DEVELOPAID · ИНВЕСТИЦИОННЫЙ ДАШБОРД · ОФИС", "", tab=DARK, width=12)
        book.pending.append((ws["A2"], f'{ref("in.name")}&" · три способа: продажа по ДДУ, продажа блоками, '
                                       f'аренда с продажей здания"', None, DASH_SHEET))
    for c in "ABCDEFGHIJKL":
        ws.column_dimensions[c].width = 13
    short = {"d": "по ДДУ", "s": "блоками", "r": "аренда"}
    tiles: list[tuple[str, str, str]] = []
    if book.project_only:
        debt = (("LLCR проектного финансирования", ref("p.llcr"), RATIO),
                ("Пик эскроу, млн ₽", ref("p.escrow_peak"), NUM0)) if with_pf else (
                ("DSCR кредита — минимум", ref("p.dscr_min"), RATIO),
                ("Пик долга объектов, млн ₽", ref("p.loan_peak"), NUM0))
        tiles += [("NPV проекта, млн ₽", ref("p.npv"), NUM0), ("IRR проекта", ref("p.irr"), PCT),
                  ("Чистая прибыль, млн ₽", ref("p.t_net"), NUM0), ("NPV собственника, млн ₽", ref("p.npv_eq"), NUM0),
                  ("IRR собственника", ref("p.irr_eq"), PCT), ("Пик своих денег, млн ₽", ref("p.peak"), NUM0),
                  ("Выручка, млн ₽", ref("p.t_revenue"), NUM0), *debt, ("Проверки книги", ref("chk.verdict"), "@")]
    else:
        for key, label, fmt in (("npv", "NPV проекта, млн ₽", NUM0), ("irr", "IRR проекта", PCT),
                                ("peak", "Пик своих денег, млн ₽", NUM0), ("t_net", "Чистая прибыль, млн ₽", NUM0)):
            for path in PATHS:
                tiles.append((f"{label} · {short.get(path['code'], path['label'])}", ref(f"{path['code']}.{key}"),
                              fmt))
        tiles += [("Выгоднее по NPV", ref("rep.better"), "@"), ("Проверки книги", ref("chk.verdict"), "@"),
                  ("LLCR · по ДДУ", ref("d.llcr"), RATIO), ("Доходность на затраты · аренда", ref("r.o_yoc"), PCT),
                  ("Ставка капитализации", ref("r.o_cap"), PCT),
                  ("Пик эскроу · по ДДУ, млн ₽", ref("d.escrow_peak"), NUM0)]
    for n, (label, formula, fmt) in enumerate(tiles):
        c0 = 1 + (n % 3) * 4
        r0 = 4 + (n // 3) * 4
        ws.merge_cells(start_row=r0, start_column=c0, end_row=r0, end_column=c0 + 2)
        ws.merge_cells(start_row=r0 + 1, start_column=c0, end_row=r0 + 2, end_column=c0 + 2)
        h = ws.cell(r0, c0, label)
        h.font, h.fill = _fnt(b=True, color="FFFFFF"), _fill(TILE)
        v = ws.cell(r0 + 1, c0)
        book.pending.append((v, formula, None, DASH_SHEET))
        v.number_format = fmt
        v.font = Font(name=FONT, size=18, b=True, color=DARK)
        v.fill = _fill(LIGHT)
        v.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    row = 4 + 4 * ((len(tiles) + 2) // 3)
    ws.merge_cells(start_row=row, start_column=1, end_row=row + 1, end_column=12)
    cell = ws.cell(row, 1)
    npv = [ref(f"{path['code']}.npv") for path in PATHS]
    if book.project_only:
        llcr = ref("p.llcr")
        book.pending.append((cell, (
            f'IF({ref("p.npv")}<0,"Проект не окупается: NPV проекта меньше нуля.","NPV проекта положителен.")'
            + (f'&IF({llcr}="","",IF({llcr}<1," Денег покупателей и продаж не хватает на погашение '
               f'проектного финансирования (LLCR меньше 1).",""))' if with_pf else "")), None, DASH_SHEET))
    else:
        book.pending.append((cell, (
            f'IF(MAX({",".join(npv)})<0,"Проект не окупается ни одним из трёх способов продажи. ","")'
            f'&"Выгоднее по NPV проекта: "&{ref("rep.better")}&"."'
            f'&IF({ref("d.llcr")}="","",IF({ref("d.llcr")}<1," При продаже по ДДУ денег покупателей и продаж не хватает '
            f'на погашение проектного финансирования (LLCR меньше 1).",""))'
            f'&IF({ref("r.o_yoc")}="","",IF({ref("r.o_yoc")}<{ref("r.o_cap")}," Доходность на затраты ниже ставки '
            f'капитализации: готовое здание стоит дешевле, чем обходится его стройка.",""))'), None, DASH_SHEET))
    cell.font = _fnt(b=True, color=DARK)
    cell.alignment = Alignment(wrap_text=True, vertical="top")
    rep = book.wb[REPORT_SHEET]
    head, first, last = report["owner_years"]
    if last >= first:
        chart = BarChart()
        chart.title = "Поток собственника по годам, млн ₽"
        chart.height, chart.width = 7.5, 22
        chart.add_data(Reference(rep, min_col=2, max_col=1 + len(PATHS), min_row=head, max_row=last),
                       titles_from_data=True)
        chart.set_categories(Reference(rep, min_col=1, min_row=first, max_row=last))
        for series, path in zip(chart.series, PATHS):
            series.graphicalProperties.solidFill = path["tab"]
        chart.x_axis.tickLblPos = "low"
        chart.x_axis.delete = False
        chart.y_axis.delete = False
        ws.add_chart(chart, f"A{row + 3}")


GUIDE = (
    (DASH_SHEET, "Главные числа трёх способов продажи, какой выгоднее и итог проверок книги."),
    (REPORT_SHEET, "Три способа рядом: экономика, офис, структура расходов, собственное участие, по годам. "
                   "Только ссылки."),
    (INPUTS_SHEET, "Единственное место правки — жёлтые клетки синим шрифтом. Класс офиса задаёт шкалу допущений; "
                   "число, вписанное у офиса вручную, сильнее класса."),
    (TEP_SHEET, "Площади, эффективность, машино-места; пятно и подземные уровни для надбавки за глубину."),
    (TIME_SHEET, "Номер месяца, разрешение на стройку, ввод проекта и объектов."),
    (CAPEX_SHEET, "Смета «ставка × база» по статьям и её график по месяцам."),
    (RATE_SHEET, "Ключевая ставка по выбранному сценарию."),
    (DDU_SHEET, "Продажа по ДДУ: доля продаж до ввода и после, рост цены, выручка, маркетинг и брокеры, "
                "налоговая маржа. Деньги покупателей лежат на эскроу до ввода проекта."),
    (SALE_SHEET, "Продажа блоками после ввода: цена, профиль продаж, выручка, маркетинг, налог на имущество с "
                 "непроданной части, налоговая маржа."),
    (RENT_SHEET, "Аренда: заполняемость, индексация, доход, эксплуатация, NOI, продажа здания, налоговая маржа."),
    (EXIT_SHEET, "Стоимость здания по ставке капитализации и доходность на затраты."),
    (PF_SHEET, "Проектное финансирование продажи по ДДУ: собственные деньги и БРИДЖ до разрешения на стройку, "
               "кредит банка с разрешения, ставка по покрытию долга эскроу, раскрытие эскроу при вводе, "
               "погашение, НДС по ДДУ и LLCR."),
    ("CF_ДДУ, CF_Продажа, CF_Аренда", "НДС, кредит, поток проекта и собственника, NPV и IRR каждого способа."),
    (TAX_SHEET, "Налог на прибыль проекта каждого способа: маржа до ввода откладывается, убыток прошлых лет гасит "
                "не больше половины базы года; уплата авансами по кварталам."),
    (CHECK_SHEET, "Каждый итог книги против расчёта движка: все три способа и подробно — способ, выбранный "
                  "в проекте."),
    (SOURCES_SHEET, "Откуда допущения шкалы класса офиса."),
    ("Правка", "Поменяйте жёлтую клетку — пересчитается вся книга. «ПРОВЕРКИ» покраснеют: движок считал по прежним "
               "вводным, и это ожидаемо."),
)


PROJECT_GUIDE = (
    (DASH_SHEET, "Главные числа проекта и итог проверок книги."),
    (REPORT_SHEET, "Экономика проекта, структура расходов, собственное участие, по годам. Только ссылки."),
    (INPUTS_SHEET, "Единственное место правки — жёлтые клетки синим шрифтом. Объекты стоят колонками; способ "
                   "продажи объекта меняется на странице проекта."),
    (TEP_SHEET, "Площади, эффективность, машино-места; пятно и подземные уровни для надбавки за глубину."),
    (TIME_SHEET, "Номер месяца, разрешение на стройку, ввод проекта и объектов."),
    (CAPEX_SHEET, "Смета «ставка × база» по статьям и её график по месяцам."),
    (RATE_SHEET, "Ключевая ставка по выбранному сценарию."),
    ("Листы объектов", "Каждый объект — своей стратегией: продажа по ДДУ (доля продаж до ввода и после, рост "
                       "цены, выручка, маркетинг, налоговая маржа), продажа блоками после ввода или аренда с "
                       "продажей здания (заполнение, NOI, выход, свой кредит)."),
    (PF_SHEET, "Проектное финансирование объектов по ДДУ: собственные деньги и БРИДЖ до разрешения на стройку, "
               "кредит банка с разрешения, ставка по покрытию долга эскроу, раскрытие эскроу при вводе, "
               "погашение, НДС по ДДУ и LLCR. Лист есть, только если в проекте есть объект по ДДУ."),
    (PROJECT_PATH["cf"], "НДС, кредит, поток проекта и собственника, NPV и IRR, собственное участие."),
    (TAX_SHEET, "Налог на прибыль проекта: маржа до ввода откладывается, убыток прошлых лет гасит не больше "
                "половины базы года; уплата авансами по кварталам."),
    (CHECK_SHEET, "Каждый итог книги против расчёта движка: проект, смета, ПФ, объекты, собственное участие, "
                  "по годам."),
    ("Правка", "Поменяйте жёлтую клетку — пересчитается вся книга. «ПРОВЕРКИ» покраснеют: движок считал по прежним "
               "вводным, и это ожидаемо."),
)


def _guide_sheet(book: _Book) -> None:
    ws = book.sheet(GUIDE_SHEET, "ИНСТРУКЦИЯ", "Как устроена книга.", tab=GREY, width=2)
    ws.column_dimensions["A"].width = 26
    ws.column_dimensions["B"].width = 110
    for r, (name, text) in enumerate(PROJECT_GUIDE if book.project_only else GUIDE, start=4):
        ws.cell(r, 1, name).font = _fnt(b=True)
        ws.cell(r, 2, text).alignment = Alignment(wrap_text=True)


def _sources_sheet(book: _Book) -> None:
    ws = book.sheet(SOURCES_SHEET, "ИСТОЧНИКИ ДОПУЩЕНИЙ КЛАССА ОФИСА", "", tab=GREY, width=7)
    ws.column_dimensions["A"].width = 50
    for c in "BCDEF":
        ws.column_dimensions[c].width = 14
    ws.column_dimensions["G"].width = 60
    r = 4
    for j, text in enumerate(("Допущение", *ns.OFFICE_CLASSES, "Ед.", "Откуда")):
        cell = ws.cell(r, 1 + j, text)
        cell.font, cell.fill = _fnt(b=True, color="FFFFFF"), _fill(HEAD)
    origin = {
        "price_th_per_sqm": "решение владельца 22.10.2026, только в нежилом проекте",
        "rent_th_per_sqm_year": "рыночная ставка Москвы на конец 2025 года «всё включено»: базовая ставка, "
                                "эксплуатация и НДС",
    }
    for field, _, label, unit, fmt, mult in CLASS_ROWS:
        r += 1
        ws.cell(r, 1, label)
        for j, cls in enumerate(ns.OFFICE_CLASSES):
            ws.cell(r, 2 + j, float(ns.OFFICE_CLASS_SCALE[cls][field]) * mult).number_format = fmt
        ws.cell(r, 6, unit).font = _fnt(color=GREY)
        ws.cell(r, 7, origin.get(field, "стартовое допущение класса; правится на «Вводных»")).alignment = Alignment(
            wrap_text=True)
    r += 2
    for text in ("Класс офиса по умолчанию идёт за классом проекта: комфорт — B, бизнес — B+, элитный — A.",
                 "Ставка аренды дана на 12.2025 и до ввода растёт на индексацию аренды, как цена продажи растёт до "
                 "ввода. Стоимость стройки в шкалу класса не входит: она идёт статьями сметы на листе CAPEX.",
                 "Здесь — справка: расчёт читает шкалу с «Вводных»."):
        ws.cell(r, 1, text).font = _fnt(i=True, color=GREY)
        r += 1


# --- «ПРОВЕРКИ» ----------------------------------------------------------------

TOLERANCE_SHARE = 1e-6
RATIO_TOLERANCE = 1e-6
_MLN = 1e6


def _engine_number(value: Any, ratio: bool) -> float | str:
    if value is None:
        return ""
    if isinstance(value, str):
        return value
    return float(value) if ratio else float(value) / _MLN


def _irr_label(summary: dict[str, Any]) -> str | None:
    """Подпись движка «IRR капитала не показателен» (`summary.equity_irr`)."""
    verdict = summary.get("equity_irr") or {}
    return str(verdict.get("label")) if verdict.get("status") == "thin_equity" else None


PATH_CHECKS: tuple[tuple[str, str, str, bool], ...] = (
    ("revenue", "t_revenue", "Выручка", False), ("capex", "t_capex", "Затраты на создание", False),
    ("selling_cost", "o_selling", "Офис: маркетинг и брокеры", False),
    ("operating_cost", "o_operating", "Офис: эксплуатация, налог на имущество, сделка", False),
    ("financing_cost", "t_financing", "Проценты и комиссии", False),
    ("profit_tax", "t_tax", "Налог на прибыль", False), ("net_profit", "t_net", "Чистая прибыль", False),
    ("npv", "npv", "NPV проекта", False), ("irr_project", "irr", "IRR проекта", True),
    ("npv_equity", "npv_eq", "NPV собственника", False), ("equity_peak", "peak", "Пик собственных денег", False),
    ("sale_revenue", "o_sale", "Офис: выручка продаж", False), ("rent_revenue", "o_rent", "Офис: доход аренды", False),
    ("exit_revenue", "o_exit", "Офис: продажа здания", False),
    ("stabilized_noi", "o_noi", "Офис: NOI года после заполнения", False),
    ("yield_on_cost", "o_yoc", "Офис: доходность на затраты", True),
    ("exit_value", "o_exit_value", "Офис: стоимость здания", False),
    ("dscr_min", "o_dscr", "Офис: DSCR — минимум", True), ("loan_peak", "o_loan_peak", "Офис: пик долга", False),
    ("commercial_costs", "t_commercial", "Маркетинг и брокеры продаж по ДДУ", False),
    ("llcr", "llcr", "LLCR проектного финансирования", True),
    ("pf_peak", "pf_peak", "Пик долга проектного финансирования", False),
)
# Чего у способа нет: у аренды — продажи, у продаж — аренды. Маркетинг офиса
# при продаже по ДДУ движок считает в расходах проекта (`commercial_costs`).
PATH_SKIP = {
    ns.STRATEGY_DIRECT: {"rent_revenue", "exit_revenue", "stabilized_noi", "yield_on_cost", "exit_value",
                         "dscr_min"},
    ns.STRATEGY_INCOME: {"sale_revenue", "selling_cost"},
    ns.STRATEGY_DDU: {"rent_revenue", "exit_revenue", "stabilized_noi", "yield_on_cost", "exit_value",
                      "dscr_min", "selling_cost"},
}
# Итоги ПФ продажи по ДДУ: строка книги → ключ `finance` движка.
PF_CHECKS: tuple[tuple[str, str, str], ...] = (
    ("pf_draw_total", "pf_draw_total", "Выдача ПФ всего"), ("peak_bridge", "peak_bridge", "БРИДЖ — пик долга"),
    ("peak_pf", "peak_pf", "ПФ — пик долга"), ("peak_escrow", "peak_escrow", "Эскроу — пик"),
    ("limit", "pf_limit", "Лимит ПФ"), ("bridge_fee", "bridge_fee", "Плата за резервирование БРИДЖа"),
    ("reservation_fee", "pf_reservation_fee", "Плата за резервирование лимита ПФ"),
    ("transferred", "transferred_bridge_interest", "Проценты БРИДЖа, перенесённые в ПФ"),
    ("ending_pf", "ending_pf", "Долг ПФ на конец проекта"),
    ("rve_unpaid", "rve_pf_shortfall", "Не погашено раскрытым эскроу в РВЭ"),
)
OBJECT_CHECKS: tuple[tuple[str, str, str, str, bool], ...] = (
    # (где у движка, ключ движка, строка книги, подпись, доля ли)
    ("totals", "revenue", "T.revenue", "Выручка объекта", False),
    ("totals", "rent_revenue", "rent", "Доход аренды", False),
    ("totals", "sale_revenue", "sale", "Выручка продаж", False),
    ("totals", "exit_revenue", "exit", "Продажа здания", False),
    ("totals", "residual_value", "residual", "Оценка удержания", False),
    ("totals", "opex", "opex", "Расходы эксплуатации", False),
    ("totals", "property_tax", "ptax", "Налог на имущество", False),
    ("totals", "selling_cost", "selling", "Маркетинг и брокеры", False),
    ("totals", "exit_cost", "exit_cost", "Расходы на сделку", False),
    ("totals", "vat_paid", "vat_paid", "НДС к уплате", False),
    ("totals", "vat_charged", "vat_charged", "НДС начислен", False),
    ("totals", "loan_draw", "T.loan_draw", "Кредит — выдача", False),
    ("totals", "loan_interest", "T.loan_interest", "Кредит — проценты", False),
    ("totals", "loan_fee", "T.loan_fee", "Кредит — комиссии", False),
    ("totals", "loan_repayment", "T.loan_repayment", "Кредит — погашено", False),
    ("totals", "loan_peak", "peak", "Кредит — пик долга", False),
    ("totals", "noi", "noi", "NOI за срок удержания", False),
    ("totals", "tax_margin", "margin", "Налоговая маржа", False),
    ("totals", "capex", "capex_total", "Затраты стройки объекта", False),
    ("kpi", "exit_value", "exit_value", "Стоимость здания", False),
    ("kpi", "stabilized_noi", "noi_stab", "NOI года после заполнения", False),
    ("kpi", "dscr_min", "dscr_min", "DSCR — минимум", True),
    ("kpi", "icr_min", "icr_min", "ICR — минимум", True),
    ("financing", "draw_object", "T.draw_object", "Выдача на стройку объекта", False),
    ("financing", "draw_common", "T.draw_common", "Выдача на общие затраты", False),
    ("financing", "fee", "T.issue_fee", "Комиссия за выдачу", False),
    ("financing", "reservation_fee", "reservation_fee", "Плата за резервирование лимита", False),
    ("financing", "limit", "loan_limit", "Лимит кредита", False),
    ("financing", "commitment_fee", "commitment_fee", "Плата за невыбранный лимит", False),
    ("financing", "limit_land", "limit_land", "Линия на участок и проект", False),
    ("financing", "limit_build", "limit_build", "Строительная линия", False),
    ("financing", "interest_capitalized", "T.interest_capitalized", "Проценты до ввода", False),
    ("financing", "interest_paid", "T.interest_paid", "Проценты после ввода", False),
    ("financing", "debt_at_commissioning", "T.debt_at_commissioning", "Долг на ввод", False),
    ("financing", "rate_at_commissioning", "T.rate_at_commissioning", "Ставка на вводе", True),
    ("financing", "avg_rate", "T.avg_rate", "Средняя ставка", True),
    ("financing", "balloon_planned", "T.balloon_planned", "Баллон по графику", False),
    ("financing", "balloon_paid", "T.balloon_paid", "Баллон погашен в срок", False),
    ("financing", "repaid_scheduled", "T.repaid_scheduled", "Погашено плановым телом", False),
    ("financing", "repaid_at_exit", "T.repaid_at_exit", "Погашено при выходе", False),
    ("financing", "repaid_from_cash", "T.repaid_from_cash", "Погашено из денег объекта", False),
    ("financing", "avg_payment", "T.avg_payment", "Аннуитетный платёж — средний", False),
    ("result", "common_capex", "common_total", "Общие затраты проекта на объекте", False),
)


def _check_rows(result: dict[str, Any], current: str | None, blocks: dict[str, list[_Block]],
                ddu: dict[str, list[_Block]], capex: dict[str, Any], report: dict[str, Any],
                objects: list[dict[str, Any]],
                ways: tuple[dict[str, str], ...] = PATHS) -> list[tuple[str, str, str, Any, bool]]:
    """(раздел, показатель, ссылка книги, число движка в единицах книги, доля ли).

    Книга офиса сверяет каждый способ с `office_paths`, а способ проекта —
    подробно; книга без офиса (`PROJECT_PATH`) — только подробно, с основным
    расчётом движка."""
    rows: list[tuple[str, str, str, Any, bool]] = []
    paths = {p.get("key"): p for p in (result.get("office_paths") or {}).get("paths") or []}
    project_only = ways == (PROJECT_PATH,)
    if project_only:
        rows.append(("Проект", "Месяцев в таблицах хватает", ref(f"{PROJECT_PATH['code']}.rows_ok"), "хватает",
                     False))
    for path in () if project_only else ways:
        engine = paths.get(path["key"]) or {}
        section = section_of(path)
        if engine.get("error"):
            rows.append((section, "Движок этот вариант не посчитал: " + str(engine["error"]), '""', "посчитан",
                         False))
            continue
        for key, book_key, label, ratio in PATH_CHECKS:
            if key in PATH_SKIP.get(path["key"], ()):
                continue
            value = engine.get(key)
            # NPV собственника снят при непогашенном ПФ — движок даёт подпись.
            if key == "npv_equity" and value is None and engine.get("equity_returns_label"):
                value = engine["equity_returns_label"]
            rows.append((section, label, ref(f"{path['code']}.{book_key}"), _engine_number(value, ratio), ratio))
        rows.append((section, "Месяцев в таблицах хватает", ref(f"{path['code']}.rows_ok"), "хватает", False))
    code = next((p["code"] for p in ways if p["key"] == current), None)
    if code is None:
        return rows
    summary = result.get("summary") or {}
    finance = result.get("finance") or {}
    equity = (result.get("report") or {}).get("equity_participation") or {}
    verdict = summary.get("equity_returns") or {}
    na_label = (verdict.get("label") or None) if verdict.get("status") == "debt_unrepaid" else None
    irr_equity = (summary.get("irr_equity") if na_label is None and _irr_label(summary) is None
                  else (na_label or _irr_label(summary)))
    debt_metric = (((result.get("report") or {}).get("layout") or {}).get("debt_metric") or {}).get("value")
    sec = "Проект"
    for key, book_key, label, value, ratio in (
            ("revenue", "t_revenue", "Выручка проекта", summary.get("revenue"), False),
            ("capex", "t_capex", "CAPEX", summary.get("capex"), False),
            ("nonres_costs", "t_nonres_costs", "Расходы объектов", finance.get("nonres_costs"), False),
            ("commercial_costs", "t_commercial", "Маркетинг и брокеры продаж по ДДУ",
             summary.get("commercial_costs"), False),
            ("ebitda", "t_ebitda", "EBITDA", summary.get("ebitda"), False),
            ("financing_cost", "t_financing", "Проценты и комиссии", summary.get("financing_cost"), False),
            ("profit_before_tax", "t_pbt", "Прибыль до налога", summary.get("profit_before_tax"), False),
            ("profit_tax", "t_tax", "Налог на прибыль", summary.get("profit_tax"), False),
            ("vat", "t_vat", "НДС", summary.get("vat"), False),
            ("net_profit", "t_net", "Чистая прибыль", summary.get("net_profit"), False),
            ("total_expenses", "t_expenses", "Расходы всего", summary.get("total_expenses"), False),
            ("npv", "npv", "NPV проекта", summary.get("npv"), False),
            ("irr_project", "irr", "IRR проекта", summary.get("irr_project"), True),
            ("npv_equity", "npv_eq", "NPV собственника", na_label or summary.get("npv_equity"), False),
            ("irr_equity", "irr_eq", "IRR собственного капитала", irr_equity, True),
            ("loan_peak", "loan_peak", "Кредит объектов — пик",
             sum(float((o.get("totals") or {}).get("loan_peak") or 0.0)
                 for o in finance.get("nonres", {}).get("objects") or []), False),
            ("id_total", "id_total", "Выручка − расходы всего − чистая прибыль", 0.0, False),
            ("id_ebitda", "id_ebitda", "EBITDA − проценты − прибыль до налога", 0.0, False),
            # Проект с ПФ меряет долг LLCR, без ПФ — DSCR кредита объектов.
            ("debt_metric", "llcr" if ddu.get(code) else "dscr_min",
             "LLCR" if ddu.get(code) else "DSCR кредита объектов — минимум", debt_metric, True)):
        rows.append((sec, label, ref(f"{code}.{book_key}"), _engine_number(value, ratio), ratio))
    # Горизонт — месяцы расчёта движка (`finance.rows`).
    rows.append((sec, "Горизонт модели, мес.", ref(f"{code}.horizon"), float(len(finance.get("rows") or [])), True))
    if ddu.get(code):
        financing = (result.get("report") or {}).get("financing") or {}
        for book_key, engine_key, label in PF_CHECKS:
            value = finance.get(engine_key)
            if engine_key == "rve_pf_shortfall":
                value = financing.get(engine_key, value)
            rows.append(("ПФ", label, ref(f"pf.{code}.{book_key}"), _engine_number(value, False), False))
        rows.append(("ПФ", "НДС продаж по ДДУ к уплате",  ref(f"pf.{code}.vat_due"),
                     _engine_number(sum(float(v or 0.0) for v in (finance.get("vat_ddu_schedule") or {}).values()),
                                    False), False))
        revenue = result.get("revenue") or {}
        margins = finance.get("tax_margin_by_product") or {}
        for block in ddu[code]:
            keys = (block.item["key"], block.item.get("parking_product_key") or "")
            title = block.item["title"]
            rows.append((title, "Выручка по ДДУ", f"SUM({block.r('revenue', ':')})",
                         _engine_number(sum(float(revenue.get(key) or 0.0) for key in keys), False), False))
            rows.append((title, "Налоговая маржа продаж по ДДУ", f"SUM({block.r('margin', ':')})",
                         _engine_number(sum(float(margins.get(key) or 0.0) for key in keys), False), False))
    engine_capex = result.get("capex") or {}
    keys = {item["code"]: item["key"] for item in objects}
    for key in capex["order"]:
        if key.endswith(".garage"):
            continue
        base = keys.get(key.split(".")[0], key) if key.endswith(".building") else key
        if base not in engine_capex:
            continue
        label = capex["labels"][key].replace(": здание", "")
        rows.append(("Смета", label, ref(f"rep.{code}.art.{key.split('.')[0]}"),
                     _engine_number(engine_capex.get(base), False), False))
    parts = equity.get("parts") or {}
    if equity:
        for book_key, label, value in (
                ("e_contributed", "Вложено всего", equity.get("contributed")),
                ("e_returned", "Получено всего", equity.get("returned")),
                ("e_in_build", "Вложено до ввода", parts.get("in_build")),
                ("e_in_operation", "Вложено после ввода", parts.get("in_operation")),
                ("e_in_exit", "Вложено при выходе", parts.get("in_exit")),
                ("e_out_vat", "Возмещение НДС стройки", parts.get("out_vat")),
                ("e_out_running", "Получено по ходу", parts.get("out_running")),
                ("e_out_exit", "Получено при выходе", parts.get("out_exit")),
                ("e_residual", "Оценка удержания", equity.get("residual"))):
            rows.append(("Собственное участие", label, ref(f"{code}.{book_key}"), _engine_number(value, False), False))
        for line in equity.get("rows") or []:
            label = str(line.get("label") or "")
            if label.startswith("Пик потребности"):
                rows.append(("Собственное участие", "Пик потребности", ref(f"{code}.peak"),
                             _engine_number(line.get("value"), False), False))
            elif label.startswith("NPV собственного капитала"):
                rows.append(("Собственное участие", "NPV собственного капитала", ref(f"{code}.npv_eq"),
                             _engine_number(line.get("value"), False), False))
            elif label.startswith("IRR собственного капитала — только деньги") and line.get("unit") == "pct":
                rows.append(("Собственное участие", "IRR — только деньги", ref(f"{code}.e_irr_cash"),
                             _engine_number(line.get("value"), True), True))
        for year in equity.get("years") or []:
            r = report["years"].get((code, year.get("year")))
            if r:
                rows.append(("По годам", f"{year['year']}: вложено собственником", f"'{REPORT_SHEET}'!$I${r}",
                             _engine_number(year.get("contributed"), False), False))
                rows.append(("По годам", f"{year['year']}: получено собственником", f"'{REPORT_SHEET}'!$J${r}",
                             _engine_number(year.get("returned"), False), False))
    taxes = (result.get("cashflow") or {}).get("profit_tax") or []
    by_year: dict[int, float] = {}
    for month, value in zip((result.get("cashflow") or {}).get("months") or [], taxes):
        by_year[_month(month).year] = by_year.get(_month(month).year, 0.0) + float(value or 0.0)
    for year, value in sorted(by_year.items()):
        r = report["years"].get((code, year))
        if r:
            rows.append(("По годам", f"{year}: налог на прибыль", f"'{REPORT_SHEET}'!$G${r}", value / _MLN, False))
    engine_objects = {o.get("key"): o for o in finance.get("nonres", {}).get("objects") or []}
    ratios_none = {"dscr_min", "icr_min"}
    for block in blocks[code]:
        engine = engine_objects.get(block.item["key"]) or {}
        for where, engine_key, book_key, label, ratio in OBJECT_CHECKS:
            if block.n(book_key) not in block.names:
                value = (engine.get(where) or {}).get(engine_key)
                # Строки нет у этой стратегии (аренда у продажи): у движка — ноль или пусто.
                if value in (None, 0, 0.0):
                    continue
            value = (engine.get(where) or {}).get(engine_key)
            book_ref = (f"SUM({block.r(book_key, ':')})"
                        if book_key in ("rent", "sale", "exit", "residual", "opex", "ptax", "selling", "exit_cost",
                                        "vat_paid", "vat_charged", "noi", "margin", "commitment_fee")
                        and block.n(book_key) in block.names else
                        block.r(book_key) if block.n(book_key) in block.names else "0")
            if engine_key in ratios_none and value is None and block.direct:
                continue
            rows.append((block.item["title"], label, book_ref, _engine_number(value, ratio), ratio))
    return rows


def _check_sheet(book: _Book, rows: list[tuple[str, str, str, Any, bool]], missing: list[str]) -> None:
    ws = book.sheet(CHECK_SHEET, "ПРОВЕРКИ · КНИГА ПРОТИВ ДВИЖКА",
                    "Каждый итог книги против авторитетного расчёта движка. Допуск — 0,0001% суммы (не меньше "
                    "1 ₽), у долей — 0,000001.", tab=GREY, width=10)
    for c, w in zip("ABCDEFGHIJ", (34, 46, 16, 16, 12, 12, 16, 2, 46, 12)):
        ws.column_dimensions[c].width = w
    head = 4
    for j, text in enumerate(("Раздел", "Показатель", "Книга", "Движок", "Разница", "Допуск", "Итог")):
        cell = ws.cell(head, 1 + j, text)
        cell.font, cell.fill = _fnt(b=True, color="FFFFFF"), _fill(HEAD)
    r = head
    for section, label, book_ref, engine, ratio in rows:
        r += 1
        ws.cell(r, 1, section)
        ws.cell(r, 2, label)
        cell = ws.cell(r, 3)
        book.pending.append((cell, book_ref, None, CHECK_SHEET))
        cell.number_format = "0.000000" if ratio else NUM3
        engine_cell = ws.cell(r, 4, engine)
        engine_cell.fill = _fill(ENGINE)
        engine_cell.number_format = "0.000000" if ratio else NUM3
        if isinstance(engine, str) and engine:
            # Движок дал не число, а вердикт или слово: книга обязана сказать то же.
            ws.cell(r, 7, f'=IF(C{r}=D{r},"{AGREE}","{DIFF}")')
            continue
        ws.cell(r, 5, f'=IF(OR(D{r}="",C{r}=""),"",C{r}-D{r})').number_format = (
            "0.000000" if ratio else NUM3)
        tolerance = f"{RATIO_TOLERANCE}" if ratio else f"MAX(0.000001,ABS(D{r})*{TOLERANCE_SHARE})"
        ws.cell(r, 6, f'=IF(D{r}="","",{tolerance})')
        # Пусто у обоих — «показателя нет» с обеих сторон; пусто у одного —
        # расхождение: книга посчитала то, чего движок не дал, или наоборот.
        ws.cell(r, 7, f'=IF(AND(D{r}="",C{r}=""),"—",IF(OR(D{r}="",C{r}=""),"{DIFF}",'
                      f'IF(ABS(E{r})<=F{r},"{AGREE}","{DIFF}")))')
    last = r
    ws["I4"] = "Вердикт"
    ws["I4"].font = _fnt(b=True)
    # Пройдено — когда КАЖДАЯ строка сказала «сходится» или «—»: строка с
    # ошибкой (#VALUE!, цикл) слова «РАСХОЖДЕНИЕ» не содержит.
    agreed = f'COUNTIF(G{head + 1}:G{last},"{AGREE}")+COUNTIF(G{head + 1}:G{last},"—")'
    ws["I5"] = f'=IF({agreed}={last - head},"{PASSED}","{FAILED}")'
    ws["I5"].font = _fnt(b=True, color=DARK)
    book.register("chk.verdict", CHECK_SHEET, 5, "I")
    ws["I6"] = "Расхождений и ошибок"
    ws["J6"] = f"={last - head}-({agreed})"
    ws["I8"] = "Формулы нет — значения движка"
    ws["I8"].font = _fnt(b=True)
    for n, text in enumerate(missing or ["нет"], start=9):
        ws[f"I{n}"] = text
        ws[f"I{n}"].alignment = Alignment(wrap_text=True)
    ws.freeze_panes = "A5"


# --- сборка ------------------------------------------------------------------

def _horizon_months(spec: dict[str, Any], objects: list[dict[str, Any]], start: date) -> int:
    """Месяцев в таблицах: самый длинный способ, хвост уплаты налога и год покупателя."""
    rve = max(int(spec["ird_min"]), int(spec["ird_months"])) + int(spec["construction_months"])
    ends = [rve + max(int(spec["residual_months"]) + 3, 12)]
    ends += [int(k) for k, _ in spec.get("purchase_plan") or []]
    for item in objects:
        if not (item.get("nonres") or item.get("ddu")):
            continue
        comm = _months_between(start, _month(item["start"])) + int(item["months"])
        if item.get("params"):
            params = item["params"]
            ends.append(max(comm, comm + max(-120, int(params.get("direct_sale_offset_months") or 0))
                            + max(1, int(params.get("direct_sale_months") or 12)) - 1))
            ends.append(comm + max(1, int(params.get("hold_years") or 5)) * 12 - 1)
        if "share_before_pct" in item:
            ends.append(comm + max(0, int(item.get("residual_months") or 0)))
            ends += [int(k) for k, w in (item.get("sales_weights") or {}).items() if w]
    return max(len(spec.get("months") or []), max(ends) + 1 + TAX_TAIL) + EXTRA_MONTHS


def evaluate(workbook) -> Evaluator:
    """Вычислитель книги, прогретый по месяцам слева направо: ряд долга читает
    прошлый месяц, и счёт подряд не упирается в глубину рекурсии."""
    evaluator = Evaluator(workbook)
    n = workbook[TIME_SHEET].max_column - FIRST_COL + 1
    monthly = [ws for ws in workbook.worksheets
               if ws.max_column >= FIRST_COL + n - 1 and ws.cell(K_ROW, FIRST_COL).value is not None]
    limit = sys.getrecursionlimit()
    sys.setrecursionlimit(max(limit, 20000))
    try:
        for col in range(FIRST_COL, FIRST_COL + n):
            letter = L(col)
            for ws in monthly:
                for r in range(1, ws.max_row + 1):
                    value = ws[f"{letter}{r}"].value
                    if isinstance(value, str) and value.startswith("="):
                        evaluator.cell(ws.title, f"{letter}{r}")
    finally:
        sys.setrecursionlimit(limit)
    return evaluator


def verdict(evaluator: Evaluator) -> tuple[str, list[tuple[str, str]]]:
    """Вердикт «ПРОВЕРОК» и строки, которые не сошлись: (раздел, показатель)."""
    sheet = evaluator.workbook[CHECK_SHEET]
    bad = [(sheet[f"A{r}"].value, sheet[f"B{r}"].value) for r in range(5, sheet.max_row + 1)
           if sheet[f"G{r}"].value and evaluator.cell(CHECK_SHEET, f"G{r}") not in (AGREE, "—")]
    return evaluator.cell(CHECK_SHEET, "I5"), bad


def build(result: dict[str, Any], spec: dict[str, Any], office_keys: set[str]) -> tuple[bytes, dict[str, Any]]:
    """Книга нежилого проекта: модель формулами от «Вводных» и сверка с движком.

    `result` — авторитетный расчёт (`consolidated`; у проекта с офисом — с
    `office_paths`), `spec` — вводные, которые движок читал (`nonres_book_spec`),
    `office_keys` — ключи офисов, которые книга считает тремя способами. Без
    офиса в проекте книга считает один способ — стратегию каждого объекта
    (`PROJECT_PATH`) — и сверяет его с основным расчётом. Отдаёт байты книги и
    отчёт сборки: вердикт «ПРОВЕРОК», расхождения и непосчитанные формулы.
    """
    start = _month(spec["start"])
    objects = []
    for n_, item in enumerate(spec.get("objects") or [], start=1):
        office = bool(item.get("nonres") or item.get("ddu")) and item["key"] in office_keys
        # Офис книга продаёт всеми способами: в книге он объект с продажей или
        # арендой (`nonres`) при любой стратегии проекта.
        objects.append({**item, "code": f"o{n_}", "office": office,
                        **({"nonres": True, "ddu": False} if office else {})})
    offices = [item for item in objects if item["office"]]
    if not any(item.get("nonres") or item.get("ddu") for item in objects):
        raise ValueError("книга нежилого проекта без объектов с продажей или арендой")
    months = len(spec.get("months") or [])
    book = _Book(_horizon_months(spec, objects, start), start)
    if not offices:
        book.paths = (PROJECT_PATH,)
    info = _inputs_sheet(book, spec, objects)
    _time_sheet(book, objects)
    _tep_sheet(book, objects, info)
    capex = _capex_sheet(book, spec, objects, info, months)
    _rate_sheet(book, info)

    flow_rows: dict[str, _Rows] = {}
    exit_rows: _Rows | None = None
    if offices:
        ddu_ws = book.sheet(DDU_SHEET, "СПОСОБ 1 · ПРОДАЖА ПО ДДУ",
                            "Договоры долевого участия: деньги покупателей лежат на эскроу до ввода проекта. "
                            "млн ₽ с НДС.", monthly=True, tab=DDU_TAB)
        sale_ws = book.sheet(SALE_SHEET, "СПОСОБ 2 · ПРОДАЖА БЛОКАМИ",
                             "Договоры купли-продажи после ввода, без ДДУ и эскроу. млн ₽ с НДС.", monthly=True,
                             tab=SALE_TAB)
        rent_ws = book.sheet(RENT_SHEET, "СПОСОБ 3 · АРЕНДА И ПРОДАЖА ЗДАНИЯ",
                             "Заполнение, индексация, эксплуатация, налог на имущество, NOI, продажа здания. млн ₽ с НДС.",
                             monthly=True, tab=RENT_TAB)
        exit_ws = book.sheet(EXIT_SHEET, "СПОСОБ 3 · ПРОДАЖА ЗДАНИЯ",
                             "Стоимость здания = NOI следующих 12 месяцев после срока удержания / ставка капитализации.",
                             tab=RENT_TAB, width=4)
        exit_ws.column_dimensions["A"].width = 64
        exit_ws.column_dimensions["B"].width = 16
        exit_ws.column_dimensions["C"].width = 14
        flow_rows = {"d": _Rows(book, ddu_ws, 5), "s": _Rows(book, sale_ws, 5), "r": _Rows(book, rent_ws, 5)}
        exit_rows = _Rows(book, exit_ws, 3)
    office_blocks: dict[str, list[_Block]] = {path["code"]: [] for path in book.paths}
    for item in offices:
        for path in PATHS:
            block = _Block(item, path["key"], f"{item['code']}{path['code']}")
            if path["key"] == ns.STRATEGY_DDU:
                _ddu_flows(book, flow_rows[path["code"]], block, info)
            else:
                _flows(book, flow_rows[path["code"]], exit_rows, block, info)
            office_blocks[path["code"]].append(block)
    other_blocks: list[_Block] = []
    other_ddu: list[_Block] = []
    used = {DDU_SHEET, SALE_SHEET, RENT_SHEET, EXIT_SHEET, PF_SHEET, INPUTS_SHEET, TEP_SHEET,
            TIME_SHEET, CAPEX_SHEET, RATE_SHEET, TAX_SHEET, REPORT_SHEET, DASH_SHEET, GUIDE_SHEET,
            CHECK_SHEET, SOURCES_SHEET, *(path["cf"] for path in book.paths)}
    other_sheets = []
    for item in objects:
        if item["office"] or not (item.get("nonres") or item.get("ddu")):
            continue
        # Excel не допускает в имени листа []:*?/\ — «ФОК / медцентр» → «ФОК - медцентр».
        name = re.sub(r"[\[\]:*?/\\]", "-", item["title"])[:31]
        while name in used:
            name = name[:28] + f" {len(used) % 10}"
        used.add(name)
        other_sheets.append(name)
        ws = book.sheet(name, f"{item['title'].upper()} · {STRATEGY[item['strategy']].upper()}",
                        ("Объект считается стратегией, выбранной у него в проекте. млн ₽ с НДС." if not offices else
                         "Объект проекта считается своей стратегией во всех трёх способах продажи офиса. "
                         "млн ₽ с НДС."), monthly=True,
                        tab=HEAD if offices else DDU_TAB if item.get("ddu")
                        else RENT_TAB if item["strategy"] == ns.STRATEGY_INCOME else SALE_TAB)
        rows = _Rows(book, ws, 5)
        block = _Block(item, item["strategy"], item["code"])
        if item.get("ddu"):
            _ddu_flows(book, rows, block, info)
            other_ddu.append(block)
            continue
        _flows(book, rows, rows, block, info)
        _finance(book, rows, block)
        _block_totals(rows, block)
        other_blocks.append(block)
    tax_ws = book.sheet(TAX_SHEET, "НАЛОГИ · НАЛОГ НА ПРИБЫЛЬ И НДС ПРОЕКТА",
                        "Налоговая база — маржа объектов; общие затраты признаются в месяц ввода проекта; вычет — "
                        "начисленные проценты и комиссии.", monthly=True, tab=GREY)
    tax_rows = _Rows(book, tax_ws, 5)
    tax_rows.param("tx.rate", "Ставка налога на прибыль", f"={ref('in.profit_tax')}", "", PCT)
    tax_rows.param("tx.limit", "Убыток прошлых лет гасит не больше доли базы года", f"={ref('in.loss_limit')}", "",
                   PCT)
    path_blocks: dict[str, list[_Block]] = {}
    path_ddu: dict[str, list[_Block]] = {}
    pf_sheets: list[str] = []
    for path in book.paths:
        c = path["code"]
        ddu_path = path["key"] == ns.STRATEGY_DDU
        credit_offices = [] if ddu_path else office_blocks[c]
        ddu = (office_blocks[c] if ddu_path else []) + other_ddu
        blocks = credit_offices + other_blocks
        path_blocks[c], path_ddu[c] = blocks, ddu
        # ПФ — свой лист у продажи офиса по ДДУ и у проекта без офиса с
        # объектом по ДДУ; у прочих способов офиса — блок на листе потока.
        pf_sheet = ddu_path or (book.project_only and bool(ddu))
        if pf_sheet:
            pf_ws = book.sheet(PF_SHEET, "ПФ · ПРОЕКТНОЕ ФИНАНСИРОВАНИЕ ПРОДАЖИ ПО ДДУ",
                               "Собственные деньги и БРИДЖ до разрешения на стройку, кредит банка (ПФ) с разрешения, "
                               "ставка по покрытию долга эскроу, раскрытие эскроу при вводе проекта, погашение.",
                               monthly=True, tab=DDU_TAB)
            pf_sheets.append(PF_SHEET)
        ws = book.sheet(path["cf"], f"{path['cf'].upper()} · НДС, КРЕДИТ, ДЕНЬГИ СОБСТВЕННИКА",
                        "Кредит: выдача — доля затрат до ввода, проценты до ввода капитализируются, после ввода — "
                        "погашение. Поток проекта — до кредита; поток собственника — после кредита и налогов.",
                        monthly=True, tab=path["tab"])
        rows = _Rows(book, ws, 5)
        if ddu:
            _pf(book, _Rows(book, pf_ws, 5) if pf_sheet else rows, path, ddu, blocks, info)
        for block in credit_offices:
            _finance(book, rows, block)
        _tax_path(book, tax_rows, path, blocks, ddu)
        _cash_path(book, rows, path, blocks, credit_offices, ddu, spec, capex, months)
        for block in credit_offices:
            _block_totals(rows, block)
    years = sorted({(start.year * 12 + start.month - 1 + k) // 12 for k in range(book.n)})
    if offices:
        current = (result.get("office_paths") or {}).get("current")
        current_label = ""
    else:
        current = PROJECT_PATH["key"]
        current_label = "; ".join(f"{item['title']} — {STRATEGY[item['strategy']][:1].lower()}{STRATEGY[item['strategy']][1:]}" for item in objects
                                  if item.get("nonres") or item.get("ddu"))
    report = _report_sheet(book, spec, capex, current, years, current_label)
    _dashboard_sheet(book, report, with_pf=bool(pf_sheets))
    _guide_sheet(book)
    if offices:
        _sources_sheet(book)
    missing = list(spec.get("missing") or [])
    if current not in {p["key"] for p in book.paths}:
        missing.append("Офисы проекта продаются разными способами: подробная сверка способа проекта не строится, "
                       "все три способа сверяются с движком итогами")
    checks = _check_rows(result, current, path_blocks, path_ddu, capex, report, objects, book.paths)
    _check_sheet(book, checks, missing)
    book.resolve()
    office_sheets = [DDU_SHEET, SALE_SHEET, RENT_SHEET, EXIT_SHEET] if offices else []
    order = [DASH_SHEET, REPORT_SHEET, INPUTS_SHEET, GUIDE_SHEET, TEP_SHEET, TIME_SHEET, CAPEX_SHEET, RATE_SHEET,
             *office_sheets, *other_sheets, *pf_sheets,
             *(path["cf"] for path in book.paths), TAX_SHEET, CHECK_SHEET, *([SOURCES_SHEET] if offices else [])]
    book.wb._sheets = [book.wb[name] for name in order]
    book.wb.active = 0
    content = book.to_bytes()
    evaluator = evaluate(load_workbook(io.BytesIO(content)))
    limit = sys.getrecursionlimit()
    sys.setrecursionlimit(max(limit, 20000))
    try:
        content, cache = with_cached_values(content, evaluator=evaluator)
        verdict_, bad_rows = verdict(evaluator)
    finally:
        sys.setrecursionlimit(limit)
    bad = [f"{section}: {label}" for section, label in bad_rows]
    return content, {"verdict": verdict_, "mismatches": bad, "missing": missing,
                     "unresolved": cache["unresolved"], "unresolved_examples": cache["examples"],
                     "months": book.n}
