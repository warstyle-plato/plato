"""Финансы книги v4: налог по годам, выручка по продуктам, БРИДЖ и потолок ПФ,
начисление процентов, cash sweep, строки дефолта и чистой прибыли ОТЧЁТА,
блок паритета с движком.

Кусок книги v4: файлы `engine/xlsx_book_v4*.py` выполняются подряд на
месте прежнего `_engine_part("xlsx_book_v4")`. Не самостоятельный модуль:
`main_legacy.py` выполняет этот файл в своём пространстве имён
(`_engine_part("xlsx_book_v4_finance")`), поэтому здесь видны все имена движка,
а подмены `core.<имя>` из `main.py` и тестов действуют и на этот код.
Отдельно не импортировать: второй импорт поднял бы вторую копию движка.
"""
# ruff: noqa

# Столбцы месяцев на листах CF: D..DS. Сетка у всех листов одна — строка 3
# каждого берёт даты с «Ставок», — поэтому налоговый год месяца объявляется
# ОДИН раз строкой на КОНСОЛИДАТОРЕ и читается всеми очередями.
_V4_TAX_YEARS_SPAN = 12          # горизонт 120 месяцев = 10 лет; с запасом
_V4_TAX_QUEUE_COLUMNS = ("C", "D", "E", "F")      # база очередей 1–4
_V4_TAX_SHARE_COLUMNS = ("H", "I", "J", "K")      # налог очередей 1–4
_V4_TAX_MONTHLY_COLUMNS = ("L", "M", "N", "O")    # запасная мера — месяцы с налогом
_V4_TAX_RESERVE_TOTAL = "P"


def _v4_tax_share_by_year(xml: str, missing: list[str]) -> str:
    """Сводный налог раздаётся очередям ПО ГОДАМ, а не одним весом за горизонт.

    Год — единственная мера, которую задаёт НК: налоговый период по налогу на
    прибыль это календарный год (ст. 285), база считается нарастающим итогом с
    начала периода (п. 7 ст. 274), а месяц объектом налогообложения не является
    вовсе. Книга делила налог долей `MAX(SUM(база за все месяцы),0)`, потом —
    долей ПОЛОЖИТЕЛЬНЫХ МЕСЯЦЕВ: обе меры приписывали очереди налог за месяцы и
    годы, в которых проект не был должен бюджету ничего. На четырёх очередях с
    шагом 24 месяца это 733 млн ₽ на строке очереди при совпадающем итоге.

    Свёртка до РВЭ обязательна и повторяет `_profit_tax_schedule`: всё, что
    признано до первого облагаемого месяца, приходит в ГОД этого месяца.
    Наивная раздача «по календарному году» без неё даёт на том же проекте
    770 млн вместо 1 847 — она выбрасывает отсрочку, которую делает сам расчёт
    налога.

    Блок дописывается в свободный низ листа: вставить строку в занятое место
    нельзя — поедут все ссылки. Налоговый год месяца считается один раз строкой
    `MAX(YEAR('CF'!D$3); первый облагаемый год)` — clamp и есть свёртка, и
    после него у каждого года простое равенство, без условий по датам.
    """
    rows = [int(number) for number in re.findall(r'<x:row r="(\d+)"', xml)]
    if not rows:
        missing.append("КОНСОЛИДАТОР: лист пуст — блок раздачи налога не дописан")
        return xml
    at = max(rows) + 2

    def label(coord: str, text: str) -> str:
        return (f'<x:c r="{coord}"{_v4_label_attr()} t="inlineStr">'
                f"<x:is><x:t>{xml_escape(text)}</x:t></x:is></x:c>")

    def formula(coord: str, text: str) -> str:
        return (f'<x:c r="{coord}"{_v4_formula_attr()}>'
                f"<x:f>{xml_escape(text)}</x:f></x:c>")

    parts: list[str] = []
    head = at
    parts.append(f'<x:row r="{head}">'
                 + _v4_head_cell(f"A{head}", "РАЗДАЧА СВОДНОГО НАЛОГА ПО ГОДАМ")
                 + "</x:row>")
    note = head + 1
    parts.append(f'<x:row r="{note}">'
                 + label(f"A{note}",
                         "Налог каждого года делится по вкладу очередей в базу "
                         "ЭТОГО года. Признанное до первого РВЭ отложено в год "
                         "РВЭ — так же считает строка 37 листа CF. Считается, "
                         "править нечего.")
                 + "</x:row>")
    gate = note + 1
    # Первый облагаемый месяц объявлен на листе CF строкой 31 — оттуда его и
    # берём. Второе объявление «с какого месяца облагается» разошлось бы с
    # первым молча: у книги тогда две даты гейта, и обе выглядят верными.
    # 99999 в строке 31 значит «включённых очередей нет» — год из него не
    # берём вовсе.
    parts.append(f'<x:row r="{gate}">'
                 + label(f"A{gate}", "Первый облагаемый год — год строки 31 листа CF")
                 + formula(f"B{gate}",
                           "IF('CF'!$B$31>50000,0,IFERROR(YEAR('CF'!$B$31),0))")
                 + "</x:row>")
    years = gate + 1
    year_cells = [label(f"A{years}", "Налоговый год месяца")]
    for column in _v4_cf_columns():
        year_cells.append(
            formula(f"{column}{years}", f"MAX(YEAR('CF'!{column}$3),$B${gate})"))
    parts.append(_v4_row_xml(years, year_cells))
    # Запасная мера на год, у которого налог есть, а годовая база
    # отрицательна: внутри года накопленным итогом была прибыль, налог
    # начислен и обратно не возвращается. Такой год делится по МЕСЯЦАМ, в
    # которых налог и начислен, — по ним же считает движок.
    monthly = years + 1
    for index in range(len(_V4_TAX_QUEUE_COLUMNS)):
        row = monthly + index
        cells = [label(f"A{row}",
                       f"База О{index + 1} в месяцы с начисленным налогом")]
        for column in _v4_cf_columns():
            cells.append(formula(
                f"{column}{row}",
                f"IF('CF'!{column}$37>0,MAX('CF_{index + 1}'!{column}$22,0),0)"))
        parts.append(_v4_row_xml(row, cells))
    header = monthly + len(_V4_TAX_QUEUE_COLUMNS)
    titles = ["Год", "Налог свода"]
    titles += [f"База О{index + 1}" for index in range(len(_V4_TAX_QUEUE_COLUMNS))]
    titles += ["Итого положительная"]
    titles += [f"Налог О{index + 1}" for index in range(len(_V4_TAX_SHARE_COLUMNS))]
    titles += [f"Запас О{index + 1}" for index in range(len(_V4_TAX_MONTHLY_COLUMNS))]
    titles += ["Итого запас"]
    parts.append(f'<x:row r="{header}">'
                 + "".join(_v4_head_cell(f"{_v4_column_letter(index + 1)}{header}", title)
                           for index, title in enumerate(titles))
                 + "</x:row>")
    grid = _v4_month_span(years)
    first = header + 1
    last = first + _V4_TAX_YEARS_SPAN - 1
    for offset in range(_V4_TAX_YEARS_SPAN):
        row = first + offset
        cells = [formula(f"A{row}", f"$B${gate}" if offset == 0 else f"A{row - 1}+1"),
                 formula(f"B{row}", f"SUMIF({grid},$A{row},{_v4_month_span(37, 'CF')})")]
        for index, column in enumerate(_V4_TAX_QUEUE_COLUMNS):
            cells.append(formula(
                f"{column}{row}",
                f"MAX(SUMIF({grid},$A{row},"
                f"{_v4_month_span(22, f'CF_{index + 1}')}),0)"))
        span = f"{_V4_TAX_QUEUE_COLUMNS[0]}{row}:{_V4_TAX_QUEUE_COLUMNS[-1]}{row}"
        cells.append(formula(f"G{row}", f"SUM({span})"))
        for index, column in enumerate(_V4_TAX_MONTHLY_COLUMNS):
            cells.append(formula(
                f"{column}{row}",
                f"SUMIF({grid},$A{row},{_v4_month_span(monthly + index)})"))
        reserve = (f"{_V4_TAX_MONTHLY_COLUMNS[0]}{row}:"
                   f"{_V4_TAX_MONTHLY_COLUMNS[-1]}{row}")
        cells.append(formula(f"{_V4_TAX_RESERVE_TOTAL}{row}", f"SUM({reserve})"))
        for index, column in enumerate(_V4_TAX_SHARE_COLUMNS):
            annual = f"{_V4_TAX_QUEUE_COLUMNS[index]}{row}"
            monthly_cell = f"{_V4_TAX_MONTHLY_COLUMNS[index]}{row}"
            total_reserve = f"${_V4_TAX_RESERVE_TOTAL}{row}"
            # Год не дал ответа — делим по месяцам с налогом; не дали и они —
            # поровну между ВКЛЮЧЁННЫМИ очередями: терять налог нельзя, а
            # основания у деления в этом случае нет, и так и написано.
            cells.append(formula(
                f"{column}{row}",
                f"IF($G{row}>0,$B{row}*{annual}/$G{row},"
                f"IF({total_reserve}>0,$B{row}*{monthly_cell}/{total_reserve},"
                f"IF($B${4 + index}=\"Да\",$B{row}/MAX(COUNTIF($B$4:$B$7,\"Да\"),1),0)))"))
        parts.append(_v4_row_xml(row, cells))
    total = last + 1
    total_cells = [label(f"A{total}", "Итого"),
                   formula(f"B{total}", f"SUM(B{first}:B{last})")]
    for column in _V4_TAX_SHARE_COLUMNS:
        total_cells.append(
            formula(f"{column}{total}", f"SUM({column}{first}:{column}{last})"))
    parts.append(_v4_row_xml(total, total_cells))
    check = total + 1
    parts.append(f'<x:row r="{check}">'
                 + label(f"A{check}",
                         "Расхождение со сводным налогом — обязано быть 0")
                 + formula(f"B{check}", f"B{total}-$K$8")
                 + "</x:row>")

    tail = "</x:sheetData>"
    if tail not in xml:
        missing.append("КОНСОЛИДАТОР: конец sheetData не найден — блок не дописан")
        return xml
    xml = xml.replace(tail, "".join(parts) + tail, 1)

    # Доля очереди теперь считается от её налога по годам, а не от базы за
    # горизонт. Подпись столбца обязана переехать вместе с величиной: «доля»
    # под двумя разными мерами — это два ответа на один вопрос.
    for index, column in enumerate(_V4_TAX_SHARE_COLUMNS, start=1):
        row = 29 + index
        span = _v4_month_span(22, f"CF_{index}")
        old = f"MAX(SUM({span}),0)"
        alt = f"SUMIF({span},\">0\")"
        source = f"{column}{total}"
        for candidate in (old, alt):
            encoded = xml_escape(candidate)
            if encoded in xml:
                xml = xml.replace(encoded, xml_escape(source), 1)
                break
        else:
            missing.append(f"КОНСОЛИДАТОР: доля налога очереди {index} не опознана")
    was = "<x:t>Положительная база</x:t>"
    if was in xml:
        xml = xml.replace(was, "<x:t>Налог очереди по годам</x:t>", 1)
    else:
        missing.append("КОНСОЛИДАТОР: подпись «Положительная база» не опознана")
    return xml


def _v4_revenue_by_product(xml: str, products: list[dict[str, Any]],
                           missing: list[str]) -> str:
    """Колонки «Выручка · продукт» по очередям и строка итога.

    Колонка заводится под продукт, у которого выручка есть хоть в одной
    очереди: семь нулевых колонок — шум, а не полнота (то же правило, что на
    экране). Продукт, которого книга посчитать не умеет, в `missing`, а не
    молча пропущен: пустая колонка и отсутствующая выглядят одинаково.
    """
    if not products:
        return xml
    # Своим импортом: прежде имя приходило в пространство движка из удалённой
    # выгрузки ПЛАТО, и книга падала NameError без неё.
    from openpyxl.utils import get_column_letter
    labels = {str(item.get("key")): str(item.get("label") or item.get("key"))
              for item in products}
    column = _V4_CONSOLIDATOR_FIRST_COL
    for key in labels:
        letter = get_column_letter(column)
        if key in _V4_SALES_PRODUCT_ROW:
            base = _V4_SALES_PRODUCT_ROW[key]
            per_queue = [f"'Продажи'!$B${base + _V4_SALES_PHASE_STRIDE * index}"
                         for index in range(len(_V4_CONSOLIDATOR_ROWS))]
        elif key in _V4_OBJECT_PRODUCT_CELLS:
            phase_cell, revenue_cell = _V4_OBJECT_PRODUCT_CELLS[key]
            per_queue = [f"IF('ОБЪЕКТЫ'!$B${phase_cell}={index + 1},"
                         f"'ОБЪЕКТЫ'!$B${revenue_cell},0)"
                         for index in range(len(_V4_CONSOLIDATOR_ROWS))]
        elif key in _V4_OBJECT_PARKING_KEYS:
            # Паркинг — продукт СВОЕГО объекта: колонка на объект, строка его
            # выручки на листе ОБЪЕКТЫ и очередь та же, что у самого объекта
            # (владелец, 27.09.2026: «Паркинг — МФОЦ / офисы», «Паркинг — ТЦ»…).
            parking = _V4_OBJECT_PARKING_BY_KEY.get(_V4_OBJECT_PARKING_KEYS[key])
            if parking is None:
                missing.append(f"КОНСОЛИДАТОР: у книги нет строки паркинга для «{labels[key]}»")
                continue
            enabled_row, revenue_row = parking[1], parking[3]
            per_queue = [f"IF('ОБЪЕКТЫ'!$B${enabled_row + 1}={index + 1},"
                         f"'ОБЪЕКТЫ'!$B${revenue_row},0)"
                         for index in range(len(_V4_CONSOLIDATOR_ROWS))]
        else:
            missing.append(f"КОНСОЛИДАТОР: книга не умеет считать выручку «{labels[key]}»")
            continue
        xml, done = _v4_set_or_insert_cell(
            xml, f"{letter}3", text=f"Выручка · {labels[key]}, млн ₽")
        if not done:
            missing.append(f"КОНСОЛИДАТОР: заголовок {letter}3 не поставлен")
            continue
        for row, formula in zip(_V4_CONSOLIDATOR_ROWS, per_queue):
            xml, done = _v4_set_or_insert_cell(xml, f"{letter}{row}", formula=formula)
            if not done:
                missing.append(f"КОНСОЛИДАТОР: ячейка {letter}{row} не поставлена")
        first, last = _V4_CONSOLIDATOR_ROWS[0], _V4_CONSOLIDATOR_ROWS[-1]
        xml, done = _v4_set_or_insert_cell(
            xml, f"{letter}{_V4_CONSOLIDATOR_TOTAL_ROW}",
            formula=f"SUM({letter}{first}:{letter}{last})")
        if not done:
            missing.append(f"КОНСОЛИДАТОР: итог {letter}{_V4_CONSOLIDATOR_TOTAL_ROW}")
        column += 1
    return xml


# Кэш-свип в книге.
#
# Очередь передала долг, но продавать не перестала: строка 12 её листа —
# «Прямые поступления после РВЭ», и долю от них банк забирает в погашение
# долга, который теперь несёт следующая. Строка 26 листов CF в шаблоне пуста —
# ряд встаёт, не сдвигая ссылок; доля лежит на «Вводных» рядом с признаком
# переноса.
# Ячейка одобренного лимита ПФ на «Вводных» и строка непокрытой потребности,
# дописываемая в конец листа CF: свободных строк внутри блока ПФ (38–47) в
# шаблоне не осталось — 26, 29, 64 и 65 уже заняты свипом, нехваткой в РВЭ и
# переносом долга. Вставить строку в занятое место нельзя: поедут все ссылки.
_V4_PF_APPROVED_CELL = "F26"
# Блок очередей «Вводных»: лимит НКЛ очереди (ввод) и её потолок (формула).
_V4_PF_QUEUE_INDIVIDUAL_COL = "AS"
_V4_PF_QUEUE_CAP_COL = "AT"
# Количество квартир очереди. Колонки под него в шаблоне нет вовсе: книга
# знала метры квартир и не знала, сколько их, — «Единицы» строки «Квартиры»
# стояли пустыми и на ТЭП, и на Дашборде. База (штуки на исходных метрах) и
# живое число, которое идёт за плотностью так же, как за ней идёт продаваемая
# площадь (L = Z×K15): числом оно стояло бы мёртвым — правка плотности
# двигала бы метры и не двигала квартиры.
#
# Буквы колонок объявлены у читателя (дашборд), а не здесь второй раз:
# писатель и читатель одной ячейки, назвавшие её порознь, однажды назовут
# разные.
_V4_APARTMENT_UNITS_BASE_COL = v4_dashboard.QUEUE_APARTMENT_UNITS_BASE_COL
_V4_APARTMENT_UNITS_COL = v4_dashboard.QUEUE_APARTMENT_UNITS_COL
_V4_PF_UNCOVERED_ROW = 87
_V4_PF_UNCOVERED_PARITY_ROW = 87

_V4_SWEEP_ROW = 26
_V4_SWEEP_SHARE_CELL = "F92"
# Единая касса проекта: две строки в конце листа CF (87 занята непокрытой
# потребностью). 88 — сколько расходов до РнС покрыто свободной кассой
# предыдущих очередей, 89 — сколько её доступно на месяц.
_V4_PROJECT_CASH_ROW = 88
_V4_PROJECT_CASH_POOL_ROW = 89


# Служебная строка блока CAPEX: база лимита БРИДЖа — те же статьи, но только
# в месяцах ДО РнС. 36 и 37 заняты сносом и расселением, 38 — процентами ВРИ,
# 39 в шаблоне пуста. В итог блока и в базу резерва она не входит по
# построению: обе формулы читают диапазон 14:31 и перечисленные строки.
_V4_CAPEX_BRIDGE_BASE_ROW = 39


def _v4_apply_bridge_limit_before_permit(xml: str, missing: list[str]) -> str:
    """Лимит БРИДЖа в книге — по платежам ДО РнС, как у движка.

    «Бридж по определению это только то, что ДО РнС» (владелец, 04.09.2026).
    База комиссии выдачи (строка 57 листов CF) брала итоги статей целиком —
    покупку, П, РД и денежную соцкомпенсацию, — и рассрочка входила в неё вся,
    включая платежи после открытия ПФ. С РнС эти платежи финансирует ПФ, и
    лимит под них — просьба о деньгах дважды.

    Считается по помесячным строкам самих статей, а не по их итогам: у покупки
    график задаётся вводной, у соцкомпенсации своя дата, а РД у длинного
    проекта переваливает за РнС. При совмещённом режиме соцнагрузки денежная
    часть идёт профилем строки 31 в той доле, в какой она денежная, — сумма за
    месяцы до РнС выходит той же, что у движка.
    """
    from openpyxl.utils import get_column_letter as _col

    def style_of(coord: str) -> str | None:
        found = re.search(r'<x:c r="%s"[^>]*?\ss="(\d+)"' % coord, xml)
        return found.group(1) if found else None

    label_style, amount_style, month_style = style_of("A20"), style_of("B20"), style_of("D20")
    for phase in range(_V4_CAPEX_PHASES):
        base = _V4_CAPEX_BLOCK_STRIDE * phase
        row = _V4_CAPEX_BRIDGE_BASE_ROW + base
        permit = f"$B${7 + base}"
        social_total = f"$B${31 + base}"
        queue_row = _V4_CF_QUEUE_ENABLED_ROW + phase
        cash = f"'Вводные'!$B$56*'Вводные'!$R${queue_row}*'Вводные'!$H$6"
        xml = _v4_ensure_row(xml, row)
        cells = {f"A{row}": dict(text="Справочно: база лимита БРИДЖа (платежи до РнС)"),
                 f"B{row}": dict(formula=f"SUM(D{row}:{_V4_LAST_COLUMN}{row})"),
                 f"C{row}": dict(text="млн ₽")}
        styles = {f"A{row}": label_style, f"B{row}": amount_style, f"C{row}": None}
        for index in range(_V4_CAPEX_MONTH_COLUMNS):
            column = _col(4 + index)
            social = (
                f"IF('Вводные'!$B$37=\"Денежная компенсация\",{column}{31 + base},"
                f"IF('Вводные'!$B$37=\"{SOCIAL_MODE_BOTH}\","
                f"{column}{31 + base}*IFERROR(({cash})/{social_total},0),0))")
            cells[f"{column}{row}"] = dict(formula=(
                f"IF({column}$3<{permit},"
                f"{column}{14 + base}+{column}{17 + base}+{column}{18 + base}+{social},0)"))
            styles[f"{column}{row}"] = month_style
        xml, done = _v4_set_cells(xml, row, cells, styles)
        if not done:
            missing.append(f"CAPEX · база лимита БРИДЖа: строка {row} не записана")
    return xml


def _v4_use_bridge_base_row(xml: str, phase: int, missing: list[str]) -> str:
    """Строка 57 листа CF берёт базу лимита из служебной строки CAPEX."""
    base = _V4_CAPEX_BLOCK_STRIDE * (phase - 1)
    row = _V4_CAPEX_BRIDGE_BASE_ROW + base
    purchase, design_p, design_rd = 14 + base, 17 + base, 18 + base
    social = 31 + base
    queue_row = _V4_CF_QUEUE_ENABLED_ROW + (phase - 1)
    cash = f"'Вводные'!$B$56*'Вводные'!$R${queue_row}*'Вводные'!$H$6"
    old = (f"('CAPEX'!$B${purchase}+'CAPEX'!$B${design_p}+'CAPEX'!$B${design_rd}"
           f"+IF('Вводные'!$B$37=\"Денежная компенсация\",'CAPEX'!$B${social},"
           f"IF('Вводные'!$B$37=\"{SOCIAL_MODE_BOTH}\",{cash},0)))")
    new = f"SUM({_v4_month_span(row, 'CAPEX')})"

    def build(_column: str, body: str) -> str:
        return body.replace(old, new)

    xml, count = _v4_rewrite_row_formulas(xml, _V4_BRIDGE_FEE_ROW, old, build)
    if not count:
        missing.append(
            f"CF_{phase}: строка {_V4_BRIDGE_FEE_ROW} не опознана — лимит БРИДЖа "
            "остался от полной цены, а не от платежей до РнС")
    return xml


def _v4_apply_extra_objects_tax_row(xml: str, phase: int, missing: list[str]) -> str:
    """Признание стоимости дописанных объектов в налоговой базе очереди — как в движке.

    Проданный объект признаёт СВОЙ CAPEX по своему проданному объёму: он и в
    движке стоит отдельным пулом (`krt_products`) — тем же слагаемым, что
    шаблон пишет трём своим объектам. У объекта с признаком «что дальше»
    (ФОК) переданный городу не признаёт ничего — выручки у него нет, — и его
    стоимость возвращается в общий пул очереди, откуда её вычитала строка
    аллокации объектов. Забыть вторую половину значит оставить стоимость
    непризнанной до конца проекта: налог вырос бы на четверть её, и обе книги
    выглядели бы верными.
    """
    if not _v4_extras():
        return xml
    queue = phase
    alloc_row = 96 + 8 * (phase - 1)
    pool_old = f"(SUM({_v4_month_span(20)})-'ОБЪЕКТЫ'!$B${alloc_row})"
    returned = ""
    for lay in _v4_extras():
        if lay.disposition_row:
            sold = (f"'Вводные'!$K${lay.disposition_row}=\"{_V4_OBJECT_SALE_WORD}\"")
            capex = f"'ОБЪЕКТЫ'!$B${lay.object_head + 22}"
            returned += (f"+IF({sold},0,IF('ОБЪЕКТЫ'!$B${lay.object_head + 2}={queue},"
                         f"{capex},0))")
    pool_new = pool_old[:-1] + returned + ")"

    def build(column: str, body: str) -> str:
        # Хвост формулы — колонко-зависимый: «-D53-D57-D21» в колонке D и
        # «-E53-E57-E21» в соседней. Искать его строкой одной колонки значит
        # не найти его ни в одной, кроме первой.
        tail = f"-{column}53-{column}57-{column}21"
        if pool_old not in body or tail not in body:
            return body
        body = body.replace(pool_old, pool_new)
        own = ""
        for lay in _v4_extras():
            head = lay.object_head
            mine = f"'ОБЪЕКТЫ'!$B${head + 2}={queue}"
            volume = _V4_OBJECT_PRODUCT_CELLS[lay.obj.key][1] - 2
            share = (f"IF({mine},IFERROR('ОБЪЕКТЫ'!$B${head + 22}*'ОБЪЕКТЫ'!{column}{volume}"
                     f"/'ОБЪЕКТЫ'!$B${volume},0),0)")
            if lay.disposition_row:
                sold = (f"'Вводные'!$K${lay.disposition_row}=\"{_V4_OBJECT_SALE_WORD}\"")
                own += f"-IF({sold},{share},0)"
            else:
                own += f"-{share}"
        # Слагаемые встают ПЕРЕД процентами и НДС — там же, где стоят три
        # объекта шаблона: порядок слагаемых читается как порядок признания.
        at = body.rindex(tail)
        return body[:at] + own + body[at:]

    xml, count = _v4_rewrite_row_formulas(xml, 22, "'ОБЪЕКТЫ'!$B$64=", build)
    if count < 100:
        missing.append(f"CF_{phase}: строка налоговой базы не знает о дописанных объектах ({count})")
    return xml


def _v4_apply_core_pool_without_sales(xml: str, phase: int, missing: list[str]) -> str:
    """Общий пул расходов очереди без проданных штук признаётся в РВЭ — как в движке.

    Строка налоговой базы делит пул (все расходы очереди минус CAPEX
    объектов) на проданные квартиры, коммерцию, паркинг и кладовые. У проекта
    без них — нежилого, офисника — делитель ноль, IFERROR отдавал 0, и пул не
    вычитался НИ В ОДНОМ месяце. Движок (`core_cost and not
    core_quantity_total`) вычитает его целиком в месяц РВЭ. На проекте
    владельца «Вавилов» (офисы 19 110 м²) это 1 819,65 млн ₽ мимо базы:
    налог книги 763,9 млн против 308,6 у движка, и паритет налога, чистой
    прибыли и LLCR падал.

    Правится только запасная ветка IFERROR, и после прохода дописанных
    объектов: пул к этому моменту уже несёт возвращённую стоимость ФОКа.
    """
    head = f"IFERROR((SUM({_v4_month_span(20)})-'ОБЪЕКТЫ'!$B${96 + 8 * (phase - 1)}"

    def build(column: str, body: str) -> str:
        at = body.find(head)
        if at < 0:
            return body
        # Пул — первая скобка внутри IFERROR; запасная ветка — «,0)» в конце
        # вызова. Скобки считаются, а не ищутся строкой: внутри пула стоят
        # вложенные IF дописанных объектов.
        start = at + len("IFERROR(")
        depth, pool_end, call_end = 0, -1, -1
        for index in range(start, len(body)):
            char = body[index]
            if char == "(":
                depth += 1
            elif char == ")":
                depth -= 1
                if depth == 0 and pool_end < 0:
                    pool_end = index + 1
                if depth < 0:
                    call_end = index
                    break
        if pool_end < 0 or call_end < 0 or body[call_end - 2:call_end] != ",0":
            return body
        pool = body[start:pool_end]
        in_rve = f"{_v4_month_offset(column, '$B$8')}=0"
        return (body[:call_end - 1] + f"IF({in_rve},{pool},0)" + body[call_end:])

    xml, count = _v4_rewrite_row_formulas(xml, 22, head, build)
    if count < 100:
        missing.append(f"CF_{phase}: строка налоговой базы не признаёт пул без продаж ({count})")
    return xml


_V4_ACCRUAL_ROW = 58                 # Финансовые расходы месяца (начисление)


def _v4_accrual_row_xml(row: int, label: str, formulas: dict[str, str]) -> str:
    """Строка листа CF деньгами — стилями строк 53 и 57, ближайших соседей той
    же природы: у них месяц стоит стилем 117, а подпись и итог идут без своего.
    Свой формат посреди блока Excel показал бы чужим."""
    cells = [
        f'<x:c r="A{row}" t="inlineStr"><x:is><x:t>{xml_escape(label)}</x:t></x:is></x:c>',
        f'<x:c r="B{row}"><x:f>SUM(D{row}:{_V4_LAST_COLUMN}{row})</x:f></x:c>',
        f'<x:c r="C{row}" t="str"><x:v>млн ₽</x:v></x:c>',
    ]
    for column in _v4_cf_columns():
        cells.append(
            f'<x:c r="{column}{row}" s="117">'
            f"<x:f>{xml_escape(formulas[column])}</x:f></x:c>")
    return f'<x:row r="{row}">' + "".join(cells) + "</x:row>"


def _v4_apply_interest_accrual(xml: str, phase: int, missing: list[str]) -> str:
    """Проценты уменьшают базу налога в месяце НАЧИСЛЕНИЯ (п. 8 ст. 272 НК).

    Строка 53 «ПФ — уплата процентов и комиссий» отвечает на вопрос кассы, и
    там она и остаётся: вклад капитала (49), распределение собственнику (50) и
    проверка фондирования (63) считают деньги. А база налога (строка 22)
    вычитала её же — то есть признавала расход в месяце ВЫПЛАТЫ, а платится в
    этой модели всё разом в РВЭ. Пока очереди прибыльны, разницы нет; она
    появляется на убыточной очереди, где половинное ограничение ст. 283
    считает по годам и год признания решает, сколько убытка зачтётся.

    Начисление собирается из строк, которые уже есть, а не считается второй
    раз: проценты БРИДЖа (32) и их капитализация — та же величина, что растит
    строку 37, — проценты ПФ (42), плата за лимит (43) и капитализация ПФ,
    растящая строку 48. Гейты повторяют гейты этих строк слово в слово, иначе
    начисленное за горизонт не сойдётся с уплаченным.

    Капитализацию в базе не путать с капитализацией в ДОЛГЕ: та лимит ПФ не
    выбирает (решение владельца 04.08.2026) и живёт отдельным обязательством —
    здесь речь только о том, в каком месяце расход признан.
    """
    sheet = f"CF_{phase}"
    columns = _v4_cf_columns()
    row = _V4_ACCRUAL_ROW

    # Перенесённые в РнС проценты БРИДЖа капитализируются по ставке ПФ уже в
    # месяце переноса — так считает движок (`pf_cap` от остатка к уплате после
    # переноса). Шаблон прибавлял их после роста на (1+ставка/12), и первый
    # месяц капитализации на них пропадал: 2,5 млн ₽ стоимости финансирования
    # на умолчаниях. Условие переноса берётся из самой строки 48 — то же слово
    # в слово, что у неё.
    from xml.sax.saxutils import unescape as _xml_unescape
    _entities = {"&quot;": '"', "&apos;": "'"}
    transfers: dict[str, str] = {}

    def compound_transfer(column: str, body: str) -> str:
        found = re.search(
            r"(?P<prev>[A-Z]{1,3})48\*\(1\+%s41/12\)\+%s42\+%s43\+(IF\(AND\(AND\(YEAR\(%s\$3\)"
            r"=YEAR\(\$B\$7\),MONTH\(%s\$3\)=MONTH\(\$B\$7\)\),[^,]+,(?P=prev)37,0\))"
            % ((column,) * 5), body)
        if not found:
            return body
        previous, transfer = found.group(1), found.group(2)
        transfers[column] = transfer
        return body.replace(
            found.group(0),
            f"({previous}48+{transfer})*(1+{column}41/12)+{column}42+{column}43", 1)

    xml, rewritten = _v4_rewrite_row_formulas(xml, 48, "37,0)", compound_transfer)
    if rewritten != len(transfers) or not transfers:
        missing.append(f"{sheet}: строка 48 — перенос процентов БРИДЖа не опознан")

    formulas: dict[str, str] = {}
    for index, column in enumerate(columns):
        previous = columns[index - 1] if index else "C"
        # Капитализация БРИДЖа — тем же условием, что у строки 37, и по ставке
        # БРИДЖа: строка 37 берёт её оттуда же.
        bridge_cap = (f"IF(AND({column}$3<$B$7,$B$5=1),"
                      f"{previous}37*'Ставки'!{column}$6/12,0)")
        # Капитализация ПФ живёт, пока копится строка 48: с РнС и по РВЭ
        # включительно — в сам месяц РВЭ её начисляет строка 53, грубя
        # накопленное на (1+ставка/12).
        transfer = transfers.get(column)
        base = (f"({previous}48+{_xml_unescape(transfer, _entities)})"
                if transfer else f"{previous}48")
        pf_cap = (f"IF(AND({column}$3>=$B$7,{column}$3<=$B$8),"
                  f"{base}*{column}41/12,0)")
        formulas[column] = (f"{column}32+{bridge_cap}"
                            f"+{column}42+{column}43+{pf_cap}")

    body = _v4_accrual_row_xml(row, "Финансовые расходы месяца (начисление)", formulas)
    existing = re.search(r'<x:row r="%d"[^>]*>(.*?)</x:row>' % row, xml, re.S)
    if not existing:
        missing.append(f"{sheet}: строка {row} не найдена")
        return xml
    # «Занята» — это формула или число в МЕСЯЧНЫХ ячейках. Подпись и единица
    # строки-разделителя наши: в шаблоне там пустая строка и знак «×».
    taken = [column for column, cell in re.findall(
        r'<x:c r="([A-Z]{1,3})%d"[^>]*>(.*?)</x:c>' % row, existing.group(1), re.S)
        if column in set(columns) and ("<x:f>" in cell or "<x:v>" in cell)]
    if taken:
        missing.append(f"{sheet}: строка {row} занята ({taken[0]}{row})")
        return xml
    xml = xml[:existing.start()] + body + xml[existing.end():]

    # База налога вычитает начисленное вместо уплаченного. Хвост формулы
    # колонко-зависимый — «-D53-D57-D21» в колонке D и «-E53-E57-E21» в
    # соседней: искать его строкой одной колонки значит не найти ни в одной,
    # кроме первой.
    changed = 0

    def swap(column: str, formula: str) -> str:
        nonlocal changed
        tail = f"-{column}53-{column}57-{column}21"
        if tail not in formula:
            return formula
        changed += 1
        return formula.replace(tail, f"-{column}{row}-{column}57-{column}21")

    xml, seen = _v4_rewrite_row_formulas(xml, 22, "'ОБЪЕКТЫ'!", swap)
    if changed < len(columns):
        missing.append(
            f"{sheet}: база налога не отдала выплату процентов начислению "
            f"({changed} из {len(columns)}, узнано {seen})")
    return xml


def _v4_mark_as_entry(xml: str, coord: str, missing: list[str]) -> str:
    """Ставит ячейке стиль ввода, не трогая её содержимого.

    Стиль здесь — утверждение о ячейке: «это правят руками». Лист ввода
    отбирает по нему строки, поэтому вводная без стиля остаётся на расчётном
    листе — то есть там, где её никто не ищет.
    """
    found = re.search(r'<x:c r="%s"(?P<attrs>[^>]*?)(/>|>)' % coord, xml)
    if not found:
        missing.append(f"стиль ввода: ячейки {coord} нет")
        return xml
    attrs = found.group("attrs")
    updated = (re.sub(r'\ss="\d+"', "", attrs) + ' s="12"').strip()
    return (xml[:found.start()] + f'<x:c r="{coord}" {updated}'
            + found.group(2) + xml[found.end():])


def _v4_apply_pf_ceiling(xml: str, phase: int, missing: list[str]) -> str:
    """Одобренный лимит ПФ становится потолком выборки — как у движка.

    Поле `pf_limit_approved_mln` работало только в движке: там заданный лимит
    становится настоящим потолком, выборка режется по свободному остатку, а
    непокрытая потребность уходит в `pf_shortfall` с месяцем первой нехватки.
    Книга же считала лимит СВОЙ (B26 — округление выборки вверх до 10 млн) и
    потолком его не ограничивала: поле стояло на «Вводных» справочно, и это
    было прямо записано в причине `V4_INPUTS_SHOWN_ONLY`.

    Цена: «требуется 30 млрд, банк одобрил 20» — движок обрежет выборку и
    покажет дыру, а книга посчитает так, будто все 30 доступны. Это не допуск
    в полпроцента, а разные проекты в двух файлах одного расчёта.

    Потолок очереди — клетка блока очередей (AT «Вводных»): её лимит НКЛ, а
    у одиночного проекта — одобренный лимит F26. Общий лимит у очередей
    потолком не служит (владелец, 20.09.2026: он ожидание на все очереди,
    а НКЛ каждой считается заново при её открытии). Считает это ровно так же
    движок (`_pf_limit_cap_mln`): разойдись книга с ним, расхождение было бы
    нашим, а не методики.

    Ноль в поле — прежнее поведение до последней формулы: `MIN` и `MAX` из
    формулы исчезают вовсе, и книга без одобренного лимита считает ровно как
    раньше.
    """
    columns = _v4_cf_columns()
    sheet = f"CF_{phase}"
    queue_row = _V4_CF_QUEUE_ENABLED_ROW + phase - 1
    # Потолок очереди — её же клетка блока очередей; ноль в ней значит
    # «потолка нет», и есть ли потолок, решает та же клетка, что и его
    # величину: второй ответ на «задан ли» разошёлся бы с первым.
    approved = f"'Вводные'!${_V4_PF_QUEUE_CAP_COL}${queue_row}"
    has_cap = f"{approved}>0"
    cells: dict[str, dict[str, Any]] = {}
    for index, column in enumerate(columns):
        # Своё тело на НАЧАЛО месяца: потолок держит только его. Принятый долг
        # лимита не выбирает, поэтому из остатка вычитается всё, что принято и
        # ещё не сметено свипом, — по ПРЕДЫДУЩИЕ месяцы: долг этого месяца
        # ложится на баланс ниже по циклу, уже после проверки потолка.
        if index == 0:
            own = f"MAX(0,{column}38)"
        else:
            previous = columns[index - 1]
            own = (f"MAX(0,{column}38-MAX(0,"
                   f"SUM($D${_V4_CARRY_ACCEPTED_ROW}:{previous}{_V4_CARRY_ACCEPTED_ROW})"
                   f"-SUM($D${_V4_SWEEP_ROW}:{previous}{_V4_SWEEP_ROW})))")
        # Формула идёт в `_v4_set_cells` СЫРОЙ: экранирует она сама, и «&gt;»
        # здесь превратилось бы в «&amp;gt;» — Excel открыл бы книгу, а
        # вычислитель не понял 18 592 формулы разом.
        cells[f"{column}45"] = dict(formula=(
            f"IF({has_cap},"
            f"MIN({column}44,MAX(0,{approved}-{own})),{column}44)"))
    xml, done = _v4_set_cells(xml, 45, cells)
    if not done:
        missing.append(f"{sheet}: строка 45 (выборка ПФ) не найдена")
        return xml

    # Непокрытая потребность отдельной величиной: без неё дыра видна только
    # как разбухший «вклад собственного капитала» (строка 49) — там она
    # неотличима от денег, которые инвестор действительно собирался внести.
    body = (
        f'<x:c r="A{_V4_PF_UNCOVERED_ROW}" s="113" t="str"><x:v>'
        f"ПФ — непокрытая потребность (одобренного лимита не хватило)"
        f'</x:v></x:c>'
        f'<x:c r="B{_V4_PF_UNCOVERED_ROW}" s="113" t="n"><x:f>'
        f"IF($B$5=1,SUM(D44:{_V4_LAST_COLUMN}44)"
        f"-SUM(D45:{_V4_LAST_COLUMN}45),0)"
        f'</x:f></x:c>'
        f'<x:c r="C{_V4_PF_UNCOVERED_ROW}" s="113" t="str"><x:v>млн ₽</x:v></x:c>')
    xml, done = _v4_insert_row(xml, _V4_PF_UNCOVERED_ROW, body, None)
    if not done:
        missing.append(
            f"{sheet}: строка {_V4_PF_UNCOVERED_ROW} занята — "
            "непокрытую потребность писать некуда")
    return xml


_V4_FEES_ROW = 57                    # комиссии выдачи (БРИДЖ и резервирование ПФ)
_V4_FEES_EQUITY_NOTE = (
    " — платятся собственным капиталом (вклад в строке 49), "
    "в потребность и лимит ПФ не входят")


def _v4_name_fees_paid_by_equity(xml: str, phase: int, missing: list[str]) -> str:
    """Комиссии выдачи книга платит капиталом — и говорит это в подписи.

    Строка 57 листов CF (комиссия БРИДЖа в первом месяце, резервирование ПФ
    в месяц РнС) входит во вклад собственного капитала (49) и не входит ни
    в потребность ПФ, ни в лимит — так же считает движок (`project_cf` и
    `equity_cf` вычитают `bridge_fee` и `pf_reservation_fee` напрямую).
    Это допустимая банковская конструкция, а не ошибка (решение владельца,
    19.09.2026: «оставляем капиталом и пишем»), но до сих пор она нигде не
    была объявлена, и человек вправе был считать, что комиссия сидит в
    лимите. Банк бывает требует оплаты и за счёт кредитных средств —
    вводная «источник оплаты комиссий» стоит в очереди задач, а пока
    подпись строки говорит, как посчитано.
    """
    found = re.search(
        r'<x:c r="A%d"([^>]*)>\s*(?:<x:v>[^<]*</x:v>|<x:is>.*?</x:is>)\s*</x:c>'
        % _V4_FEES_ROW, xml, re.S)
    if not found:
        missing.append(f"CF_{phase}: подпись строки {_V4_FEES_ROW} не найдена")
        return xml
    attrs = re.sub(r'\s+t="[^"]*"', "", found.group(1))
    label = "Комиссии выдачи (БРИДЖ и резервирование ПФ)" + _V4_FEES_EQUITY_NOTE
    cell = (f'<x:c r="A{_V4_FEES_ROW}"{attrs} t="inlineStr">'
            f"<x:is><x:t>{xml_escape(label)}</x:t></x:is></x:c>")
    return xml[:found.start()] + cell + xml[found.end():]


# НДС к уплате (строка 21) — живые деньги в бюджет в месяц РВЭ. Шаблон брал
# его только в прибыль, налог и LLCR, а вклад и распределение капитала и поток
# проекта его не видели: IRR и NPV книги, как и у движка, были завышены на
# НДС продаж. Решение владельца (13.10.2026) — вычитать; движок вычитает тот
# же НДС по ДДУ (`vat_ddu_schedule`) в `project_cf` и `equity_cf`.
_V4_VAT_CASH_ROWS = (
    (49, "{c}20+{c}24+", "{c}20+{c}21+{c}24+"),
    (50, "-{c}20-{c}24-", "-{c}20-{c}21-{c}24-"),
    (54, "{c}27-{c}20-{c}24", "{c}27-{c}20-{c}21-{c}24"),
    (63, "-{c}20-{c}24-", "-{c}20-{c}21-{c}24-"),
)


def _v4_apply_single_queue_tax(xml: str, phase: int, missing: list[str]) -> str:
    """Налог в потоках единственной очереди — налог свода (CF строка 37).

    Строка 24 шаблона считает налог накопленной базой за весь горизонт:
    убыток прошлых лет гасит прибыль целиком. Движок и свод книги считают
    по годам с переносом убытка не больше половины базы года (ст. 283): на
    проекте с убыточным стартом и прибылью позже свод платит налог, а строка
    24 — ноль. Строка 24 кормит вклад и распределение капитала (49, 50), CFADS
    и потребность в БРИДЖе, и у проекта с одной очередью капитал выходил не
    тем, что у движка (12 мес. остаточных продаж: 27 млн ₽).

    Только при одной включённой очереди: тогда свод и есть её налог. Раздача
    свода по очередям годами (блок КОНСОЛИДАТОРа) тянет весь год вперёд, и
    строка 24 через потребность БРИДЖа замкнулась бы сама на себя; там
    остаётся формула шаблона — так считает и движок очередей.
    """
    def build(column: str, body: str) -> str:
        sheet = re.search(r"('[^']+')!\$B\$21", body)
        if sheet is None:
            return body
        return (f'IF(COUNTIF({sheet.group(1)}!$B$88:$B$91,"Да")=1,'
                f"IF($B$5=1,'CF'!{column}37,0),{body})")

    xml, count = _v4_rewrite_row_formulas(xml, 24, "$B$21", build)
    if count == 0:
        missing.append(f"CF_{phase}: строка 24 (налог) не опознана")
    return xml


def _v4_apply_vat_cash(xml: str, phase: int, missing: list[str]) -> str:
    """Строки 49, 50, 54 и 63 видят НДС к уплате (строка 21)."""
    sheet = f"CF_{phase}"
    columns = _v4_cf_columns()
    for row, pattern, replacement in _V4_VAT_CASH_ROWS:
        formulas = _v4_row_formulas(xml, row)
        cells: dict[str, dict[str, Any]] = {}
        for column in columns:
            coord = f"{column}{row}"
            formula = formulas.get(coord)
            if formula is None:
                missing.append(f"{sheet}: формулы {coord} нет")
                break
            old = pattern.format(c=column)
            if old not in formula:
                missing.append(f"{sheet}: формула {coord} не опознана")
                break
            cells[coord] = dict(formula=formula.replace(old, replacement.format(c=column), 1))
        if cells:
            xml, done = _v4_set_cells(xml, row, cells)
            if not done:
                missing.append(f"{sheet}: строка {row} не найдена")
    return xml


def _v4_apply_project_cash(xml: str, phase: int, missing: list[str]) -> str:
    """Единая касса проекта: расходы очереди до РнС сначала из свободной кассы.

    Стратегия `unified_project_cash` жила только в расчёте: книга вела БРИДЖ
    каждой очереди от нуля, и на четырёх очередях Нагатино четвёртая брала
    964 млн БРИДЖа там, где движок покрыл 930,5 млн кассой первой (разбор
    16–17.10.2026: +401,6 млн стоимости финансирования и +930 млн пика ПФ
    против движка). Методика — движковая (`financing.py`, ветка до РнС):
    касса идёт первой, затем заявленные собственные средства, остаток —
    БРИДЖ; касса предыдущей очереди — её положительный поток капитала
    (строка 50 её листа), доступный с месяца выплаты включительно; очереди
    считаются по порядку, и касса, взятая более ранней очередью, следующей
    уже не достаётся — списывается с самых ранних поступлений (FIFO), как
    `_consume_project_cash_sources`.

    Строка 88 — покрытие в месяц, 89 — доступная касса на месяц. Выборка
    БРИДЖа (34) и расход собственных средств (56) считают потребность уже
    после покрытия. Признак выключен — обе строки нули, и правленые формулы
    сходятся к прежним.
    """
    columns = _v4_cf_columns()
    sheet = f"CF_{phase}"
    row, pool_row = _V4_PROJECT_CASH_ROW, _V4_PROJECT_CASH_POOL_ROW
    flag = f"'Вводные'!${_V4_UNIFIED_CASH_CELL[0]}${_V4_UNIFIED_CASH_CELL[1:]}=\"Да\""
    draw: dict[str, str] = {}
    pool: dict[str, str] = {}
    for index, column in enumerate(columns):
        if phase <= 1:
            draw[column] = pool[column] = "0"
            continue
        # Касса предыдущих очередей к этому месяцу включительно за вычетом
        # того, что из неё уже взяли очереди между ними (FIFO — с ранних).
        available = f"SUM('CF_1'!$D$50:'CF_1'!{column}50)"
        for earlier in range(2, phase):
            available = (f"MAX(0,{available}-'CF_{earlier}'!$B${row})"
                         f"+SUM('CF_{earlier}'!$D$50:'CF_{earlier}'!{column}50)")
        spent = f"-SUM($D${row}:{columns[index - 1]}{row})" if index else ""
        pool[column] = f"MAX(0,{available}{spent})"
        draw[column] = (f"IF(AND({flag},$B$5=1,{column}$3<$B$7),"
                        f"MIN({column}{pool_row},{column}33),0)")
    for number, label, total, formulas in (
            (row, "Касса проекта — покрытие расходов до РнС",
             f"SUM(D{row}:{_V4_LAST_COLUMN}{row})", draw),
            (pool_row, "Касса проекта — доступно на месяц (свободная касса предыдущих очередей)",
             f"MAX(D{pool_row}:{_V4_LAST_COLUMN}{pool_row})", pool)):
        body = (f'<x:c r="A{number}" s="113" t="str"><x:v>{xml_escape(label)}</x:v></x:c>'
                f'<x:c r="B{number}" s="113" t="n"><x:f>{total}</x:f></x:c>'
                f'<x:c r="C{number}" s="113" t="str"><x:v>млн ₽</x:v></x:c>'
                + "".join(f'<x:c r="{column}{number}" s="120"><x:f>'
                          f"{xml_escape(formulas[column])}</x:f></x:c>"
                          for column in columns))
        xml, done = _v4_insert_row(xml, number, body, None)
        if not done:
            missing.append(f"{sheet}: строка {number} занята — касса проекта не записана")
            return xml
    # Потребность после покрытия кассой: БРИДЖ (34) и собственные средства (56).
    # Строки 43 и 57 (плата за невыбранный лимит и комиссия резервирования) брали лимит очереди как
    # «общий авто-лимит B26 × доля очереди V»: B26 суммирует выборку ПФ ВСЕХ
    # очередей, и через кассу (89 → строка 50 предыдущей → её 57 → B26 →
    # выборка следующей) это замыкалось в круг. Лимит очереди — та же
    # величина без обхода через свод: ROUNDUP(своя B45/10)×10 (V = он / B26).
    queue_row = _V4_CF_QUEUE_ENABLED_ROW + (phase - 1)
    for target, was, now in (
            (34, "MAX(0,{c}33-{c}56)", "MAX(0,{c}33-{c}" + str(row) + "-{c}56)"),
            (56, ",{c}33)", ",MAX(0,{c}33-{c}" + str(row) + "))"),
            # 43 — плата за невыбранный лимит, 57 — комиссия резервирования:
            # обе брали лимит очереди через свод.
            (43, f"'Вводные'!$B$26*'Вводные'!$V${queue_row}", "ROUNDUP($B$45/10,0)*10"),
            (57, f"'Вводные'!$B$26*'Вводные'!$V${queue_row}", "ROUNDUP($B$45/10,0)*10")):
        cells: dict[str, dict[str, Any]] = {}
        for coord, formula in _v4_row_formulas(xml, target).items():
            column = re.match(r"[A-Z]+", coord).group(0)
            if column in ("A", "B", "C"):
                continue
            if was.format(c=column) not in formula:
                missing.append(f"{sheet}: строка {target} не опознана — касса проекта не учтена")
                return xml
            cells[coord] = dict(formula=formula.replace(
                was.format(c=column), now.format(c=column), 1))
        xml, done = _v4_set_cells(xml, target, cells)
        if not done:
            missing.append(f"{sheet}: строка {target} не записана — касса проекта не учтена")
            return xml
    return xml


def _v4_apply_cash_sweep(xml: str, phase: int, missing: list[str]) -> str:
    """Строка 26: сколько принято кэш-свипом с продаж предыдущей очереди.

    Формула — та же, что у движка: доля от прямых поступлений предыдущей
    очереди, но не больше её долга на этот месяц. Больше, чем должны, банк не
    берёт; излишек остаётся проекту, и в книге он просто не списывается.
    """
    columns = _v4_cf_columns()
    sheet = f"CF_{phase}"
    if phase <= 1:
        formulas = {column: "0" for column in columns}
    else:
        previous = f"CF_{phase - 1}"
        share = f"'Вводные'!${_V4_SWEEP_SHARE_CELL[0]}${_V4_SWEEP_SHARE_CELL[1:]}"
        flag = f"'Вводные'!${_V4_CARRY_FLAG_CELL[0]}${_V4_CARRY_FLAG_CELL[1:]}=\"Да\""
        formulas = {}
        for index, column in enumerate(columns):
            if index == 0:
                # В первом месяце передавать ещё нечего.
                formulas[column] = "0"
                continue
            # Свип идёт, только когда предыдущая очередь уже передала долг:
            # до этого её поступления гасят её собственную линию.
            handed = f"SUM('{previous}'!$D$65:'{previous}'!${columns[index - 1]}$65)>0"
            owed = f"MAX(0,{column}38+{column}45+{column}{_V4_CARRY_ACCEPTED_ROW}-{column}46)"
            formulas[column] = (
                f"IF(AND({flag},$B$5=1,{handed}),"
                f"MIN({owed},{share}*'{previous}'!{column}12),0)")
    body = _v4_carry_row_xml(_V4_SWEEP_ROW, "ПФ — погашение кэш-свипом предыдущей очереди",
                             formulas)
    existing = re.search(r'<x:row r="%d"[^>]*>(.*?)</x:row>' % _V4_SWEEP_ROW, xml, re.S)
    if not existing:
        missing.append(f"{sheet}: строка {_V4_SWEEP_ROW} не найдена")
        return xml
    if "<x:v>" in existing.group(1) or "<x:f>" in existing.group(1):
        missing.append(f"{sheet}: строка {_V4_SWEEP_ROW} занята")
        return xml
    xml = xml[:existing.start()] + body + xml[existing.end():]
    # Долг на конец и его самопроверка обязаны увидеть свип, иначе книга
    # покажет погашение и тот же долг.
    for row, pattern, replacement in (
            (47, "{c}38+{c}45+{c}64-{c}46-{c}65", "{c}38+{c}45+{c}64-{c}46-{c}65-{c}26"),
            (61, "{c}38+{c}45+{c}64-{c}46-{c}65", "{c}38+{c}45+{c}64-{c}46-{c}65-{c}26")):
        # Строка читается и пишется ОДНИМ проходом на каждую: правка по ячейке
        # копировала мегабайтный лист на каждый из 180 месяцев — 1440 копий на
        # сборку, 1,4 секунды. Цена прохода растёт вместе с сеткой, и сторож,
        # считающий ПРОХОДЫ, этого подорожания не видит.
        formulas = _v4_row_formulas(xml, row)
        cells: dict[str, dict[str, Any]] = {}
        for column in columns:
            coord = f"{column}{row}"
            formula = formulas.get(coord)
            if formula is None:
                # Ячейка может и быть — без формулы. Читатель об этом узнаёт.
                missing.append(f"{sheet}: формулы {coord} нет")
                break
            old = pattern.format(c=column)
            if old not in formula:
                missing.append(f"{sheet}: формула {coord} не опознана")
                break
            cells[coord] = dict(formula=formula.replace(old, replacement.format(c=column), 1))
        if cells:
            xml, done = _v4_set_cells(xml, row, cells)
            if not done:
                missing.append(f"{sheet}: строка {row} не найдена")
    return xml


_V4_CARRY_PARITY_ROW = 85
_V4_VAT_PARITY_ROW = 86


_V4_REPORT_DEFAULT_ROW = 21          # ОТЧЕТ: строка под «Статусом модели»


def _v4_add_report_default_row(xml: str, missing: list[str]) -> str:
    """Строка ОТЧЁТа: сколько не погашено в РВЭ и не передано дальше.

    Строка 22 у книги уже есть — «Непогашенный долг на конец проекта (дефолт,
    если > 0)». У неё тот же изъян, который на экране чинился отдельно: она
    смотрит на КОНЕЦ горизонта. Банк ждёт погашения в дату раскрытия эскроу, и
    если остаточные продажи закрывают долг годом позже, строка 22 показывает
    ноль — книга о дефолте молчит (владелец, 30.08.2026: «в экселе мы где-то
    увидим что проект дефолт»).

    Нехватку в РВЭ книга считает по очередям (строка 29 листов CF), но одной
    суммой её показывать нельзя: та же величина возникает и при законном
    переносе долга следующей очереди — банк его принял, дефолта нет. Поэтому
    из суммы вычитается переданное (строка 65), и остаётся ровно то, что никто
    не принял.

    Строка 21 в шаблоне пуста в колонках A–C (справа, в F–H, живёт своя
    панель темпов продаж — её не трогаем). Значит ряд встаёт, не сдвигая ни
    одной ссылки.
    """
    row = _V4_REPORT_DEFAULT_ROW
    unpaid = "+".join(f"'CF_{phase}'!$B${_V4_RVE_UNPAID_ROW}" for phase in range(1, 5))
    passed = "+".join(f"'CF_{phase}'!$B${_V4_CARRY_PASSED_ROW}" for phase in range(1, 5))
    for coord, kwargs in (
            (f"A{row}", {"text": "Не погашено в дату раскрытия эскроу и не передано "
                                 "следующей очереди (дефолт по линии, если > 0)"}),
            (f"B{row}", {"formula": f"MAX(0,({unpaid})-({passed}))"})):
        # Занятая ячейка значит, что шаблон изменился: писать поверх молча
        # нельзя — так теряется то, что там стояло.
        existing = re.search(
            r'<x:c r="%s"[^>]*?(?:/>|>(.*?)</x:c>)' % re.escape(coord), xml, re.S)
        if existing and (existing.group(1) or ""):
            missing.append(f"ОТЧЁТ: ячейка {coord} занята")
            continue
        xml, done = _v4_set_cell(xml, coord, **kwargs)
        if not done:
            missing.append(f"ОТЧЁТ: ячейка {coord} не найдена")
    return xml


_V4_REPORT_NET_PROFIT_ROW = 12       # ОТЧЕТ: «Чистая прибыль»
_V4_CF_VAT_ROW = 21                  # CF очереди: НДС к уплате


def _v4_report_net_profit_from_its_own_rows(xml: str, missing: list[str]) -> str:
    """Чистая прибыль считается из строк ТОГО ЖЕ листа, а не из чужого блока.

    В шаблоне `ОТЧЕТ!B12` читает `КОНСОЛИДАТОР!L8` — там чистая собирается по
    очередям, и из каждой вычитается её доля налога (`K4:K7 = K8 × доля`). У
    владельца Excel посчитал эти доли нулями, и налог не вычелся ни у одной
    очереди: на экране стояло 56 598,6 вместо 42 128,6 — ровно на налог
    14 470,0 больше (14.09.2026). Наш вычислитель на том же файле давал верное
    число, то есть расхождение живёт между Excel и цепочкой раздачи, а не в
    методике, и искать в ней виноватую ячейку нечем: формулы там законные.

    Лечится тем, что **величину, которую читает человек, собирают из величин
    рядом с ней**: прибыль до налога (B10) минус налог (B11) минус НДС. Обе
    строки Excel считает верно — это видно на том же экране, — и потерять
    налог по дороге больше негде.

    НДС стоит внутри формулы, а не своей строкой: свободной строки на листе
    нет, а вставить её значит сдвинуть все ссылки. Поэтому лист по-прежнему не
    складывается глазами (B10 − B11 ≠ B12) — это отдельная правка, и она про
    вёрстку, а не про число.

    Раздача налога по очередям остаётся там, где ей место, — в таблице
    сравнения очередей.
    """
    row = _V4_REPORT_NET_PROFIT_ROW
    before = re.search(r'<x:c r="B%d"[^>]*>(.*?)</x:c>' % row, xml, re.S)
    if not before or "КОНСОЛИДАТОР" not in (before.group(1) or ""):
        # Формула шаблона не опознана — значит лист другой, и писать поверх
        # нельзя: расхождение уходит в missing, а не проходит молча.
        missing.append(f"ОТЧЁТ: формула чистой прибыли в B{row} не опознана")
        return xml
    vat = "+".join(f"'CF_{phase}'!$B${_V4_CF_VAT_ROW}" for phase in range(1, 5))
    xml, done = _v4_set_cell(
        xml, f"B{row}", formula=f"B{row - 2}-B{row - 1}-({vat})")
    if not done:
        missing.append(f"ОТЧЁТ: ячейка B{row} не найдена")
    return xml


def _v4_relax_limit_check_for_carried_debt(xml: str, missing: list[str]) -> str:
    """Проверка 72 книги не знает, что принятый долг лимита не выбирает.

    Лимит покрывает только тело, выбранное этой очередью; долг, переоформленный
    с предыдущей, приходит на линию, не выбирая лимита (решение владельца,
    27.08.2026). Проверка сравнивала пик ПФ с лимитом как есть и на любом
    проекте с переносом кричала «превышение», которого по методике нет.
    Кричащая зря проверка хуже отсутствующей: её перестают читать.
    """
    row = 72
    found = re.search(r'<x:c r="B%d"[^>]*>(.*?)</x:c>' % row, xml, re.S)
    if not found:
        missing.append(f"ПРОВЕРКИ: ячейка B{row} не найдена")
        return xml
    body = found.group(1)
    changed = body
    for phase in range(1, 5):
        old = f"'CF_{phase}'!B83-"
        new = f"'CF_{phase}'!B83-'CF_{phase}'!$B${_V4_CARRY_ACCEPTED_ROW}-"
        if old not in changed:
            missing.append(f"ПРОВЕРКИ: формула B{row} не опознана (очередь {phase})")
            return xml
        changed = changed.replace(old, new, 1)
    return xml[:found.start(1)] + changed + xml[found.end(1):]


# Строка-гейт паритета: правил ли человек вводные после сборки книги.
_V4_PARITY_GATE_ROW = 75
# Ссылка на ЛИСТ ВВОДА из чужого листа едет меткой. При записи в архив каждый
# лист проходит переименование «Вводные» → «Параметры модели», и написанная
# прямо ссылка уехала бы на расчётный лист — туда, где формулы читают CAPEX и
# CF, то есть отпечаток вводных двигался бы от расхождения, которое он обязан
# показывать.
_V4_ENTRY_SHEET_TOKEN = "'__ЛИСТ-ВВОДА__'!"


def _v4_parity_verdict(row: int) -> str:
    """Вердикт строки паритета — объявлен один раз на все строки блока.

    Ответов три, и третий появился не для красоты. Цель в колонке C — это
    значение движка на дату сборки, а книга считает сама, и правка вводной в
    ней законна: «эксель должен работать почти как движок, если что-то
    меняешь где-то, всё должно меняться так же» (владелец, 03.09.2026).
    Измерено: подними цену квартир на 10% прямо в книге — она пересчитается
    верно, а ВОСЕМЬ строк паритета из одиннадцати покраснеют, потому что цель
    осталась прежней. Кричащая зря проверка хуже отсутствующей: её перестают
    читать, — а читать её надо, именно эти строки нашли и забытую статью
    сноса, и базу НДС, и паркинг офисника.

    Гейт снимает вердикт только явным «WARN»: не собрался лист ввода, не
    записался отпечаток — и строка ведёт себя как прежде. Отказ гейта обязан
    возвращать проверку, а не выключать её.
    """
    gate = _V4_PARITY_GATE_ROW
    return (f'IF(C{row}="","",IF($F${gate}="WARN",'
            f'"вводные правлены — сверять не с чем",'
            f'IF(ABS(D{row})<=E{row},"OK","FAIL")))')


def _v4_entry_fingerprint(entry_xml: str) -> float:
    """Отпечаток листа ввода: сумма всех его чисел.

    Считается ровно то, что считает `SUM` по всему листу в самой книге:
    числовые ячейки, без подписей и без формул — формул на листе ввода нет
    ни одной, он чистый ввод человека. Поэтому отпечаток не зависит от
    расчёта книги и не может замолчать расхождение, ради показа которого
    строки паритета и стоят.

    Отбрасывается ТЕКСТ, а не всякий объявленный тип: у числа тип бывает
    написан явно (`t="n"`), и запрет по одному наличию `t=` выкидывал
    четырнадцать живых чисел, а даты (они без типа) при этом считал только
    отпечаток — три разных суммы на одну книгу, и гейт кричал на
    свежесобранной. Текст, логическое и ошибку Excel в диапазоне
    пропускает — их и пропускаем.
    """
    total = 0.0
    for match in re.finditer(
            r"<x:c(?P<attrs>\s[^>]*?)?>(?P<body>(?:(?!</x:c>).)*)</x:c>",
            entry_xml, re.S):
        attrs = match.group("attrs") or ""
        body = match.group("body") or ""
        kind = re.search(r'\st="([^"]*)"', attrs)
        if "<x:f" in body or (kind and kind.group(1) != "n"):
            continue
        value = re.search(r"<x:v>([^<]*)</x:v>", body)
        if not value:
            continue
        try:
            total += float(value.group(1))
        except ValueError:
            continue
    return total


def _v4_add_parity_gate_row(xml: str, missing: list[str]) -> str:
    """Строка «вводные не менялись с даты сборки» над блоком паритета.

    Отпечаток дописывается позже — когда лист ввода собран, — поэтому цель
    здесь пустая: пустая цель гасит и саму строку, и гейт, то есть паритет
    остаётся таким, каким был до этой правки.
    """
    row = _V4_PARITY_GATE_ROW
    body = (
        f'<x:c r="A{row}" t="inlineStr"><x:is><x:t>'
        + xml_escape("Лист ввода не менялся с даты сборки")
        + "</x:t></x:is></x:c>"
        + f'<x:c r="B{row}"><x:f>'
        + xml_escape(f"ROUND(SUM({_V4_ENTRY_SHEET_TOKEN}$A$1:$BZ$1000),6)")
        + "</x:f></x:c>"
        + f'<x:c r="C{row}"/>'
        + f'<x:c r="D{row}"><x:f>IF(C{row}="","",B{row}-C{row})</x:f></x:c>'
        + f'<x:c r="E{row}"><x:v>0.001</x:v></x:c>'
        + f'<x:c r="F{row}"><x:f>'
        + xml_escape(f'IF(C{row}="","",IF(ABS(D{row})<=E{row},"OK","WARN"))')
        + "</x:f></x:c>")
    updated, done = _v4_insert_row(xml, row, body, row + 1)
    if not done:
        missing.append(f"гейт паритета: строка {row} ПРОВЕРОК занята")
        return xml
    return updated


# Доля допуска строк паритета. Была одна на все — полпроцента, — и это
# измерено, а не выбрано: на десяти формах проекта (одиночный, дорогой вход,
# рассрочка ВРИ, ступени ставки, лестница цены, одобренный лимит, объекты, три
# и четыре очереди, перенос долга) выручка, CAPEX, EBITDA, НДС, пики БРИДЖа и
# ПФ, переданный долг и непокрытая потребность расходятся не более чем на
# 0,003% — это округление. Полпроцента на них — запас в полтораста раз: на
# проекте в сорок миллиардов допуск CAPEX выходил 188 млн ₽ при фактическом
# расхождении в 50 тыс.
#
# Исключение ровно одно и названо: стоимость финансирования. У неё расхождение
# НАСТОЯЩЕЕ — 0,21–0,42% на тех же десяти формах, то есть до 84% собственного
# допуска, — и сидит оно в капитализированных процентах (302,1 млн у книги
# против 311,3 у движка на проверочном проекте, при совпадении всех четырёх
# слагаемых до шести знаков). Ужать её допуск значит покрасить книгу красным у
# каждого, кто её скачает; причина не найдена и стоит в задачах. Налог и чистая
# прибыль — производные от неё, и их допуск считается от слагаемых ниже.
_V4_PARITY_SHARE_DEFAULT = 0.0005
_V4_PARITY_SHARE = {"financing_mln": 0.005}
# LLCR — абсолютный допуск: доля от 1,1x ничего не значит. Худшее измеренное
# расхождение на тех же формах — 0,00056, запас девятикратный.
_V4_PARITY_LLCR_TOLERANCE = 0.005


def _v4_add_parity_row(xml: str, row: int, label: str, total: str,
                       target_mln: float, missing: list[str], what: str) -> str:
    """Ещё одна строка паритета в свободный низ листа ПРОВЕРКИ.

    Две таких строки были написаны копией, и различались они только подписью и
    формулой: вторая копия однажды забыла бы расширить вердикт B3 на свою же
    строку — она красная, а лист «ПРОЙДЕНО». Допуск считается той же мерой,
    что у строк блока выше.
    """
    tolerance = max(1.0, abs(target_mln) * _V4_PARITY_SHARE_DEFAULT)
    body = (
        f'<x:c r="A{row}" t="inlineStr"><x:is><x:t>{xml_escape(label)}'
        f'</x:t></x:is></x:c>'
        f'<x:c r="B{row}"><x:f>{total}</x:f></x:c>'
        f'<x:c r="C{row}"><x:v>{_v4_number(round(target_mln, 4))}</x:v></x:c>'
        f'<x:c r="D{row}"><x:f>IF(C{row}="","",B{row}-C{row})</x:f></x:c>'
        f'<x:c r="E{row}"><x:v>{_v4_number(round(tolerance, 4))}</x:v></x:c>'
        f'<x:c r="F{row}"><x:f>{xml_escape(_v4_parity_verdict(row))}</x:f>'
        f'</x:c>')
    updated, done = _v4_insert_row(xml, row, body, None)
    if not done:
        missing.append(f"{what}: строка {row} ПРОВЕРОК занята")
        return xml
    # Вердикт листа считает по диапазону — новая строка обязана в него войти,
    # иначе она красная, а лист «ПРОЙДЕНО».
    before = updated
    updated = re.sub(r"COUNTIF\(F6:F\d+,", f"COUNTIF(F6:F{row},", updated)
    if updated == before:
        missing.append(f"{what}: вердикт B3 не расширен на строку {row}")
    return updated


def _v4_object_phase_by_result(phases: list[dict[str, Any]]) -> dict[str, int]:
    """В какой очереди движок ПОСТРОИЛ каждый отдельно стоящий объект.

    Читается результат, а не правила: очередь, в чьих вводных объект ВКЛЮЧЁН,
    и есть его очередь — это то, что движок применил, разложив проект. Повторять
    здесь лестницу приоритета нельзя: это был бы третий ответ на вопрос, у
    которого уже два.

    Первая версия смотрела на строку ТЭП очереди и ошибалась молча: площадь
    объекта движок читает из `offices_gba_sqm`, а строка ТЭП при этом может
    быть нулевой — так устроены умолчания, и на них объект «нигде не построен»
    при живых 10 000 м² и 2 519 млн ₽ CAPEX в четвёртой очереди. Признак
    включения верен и для ПЕРЕДАННОГО городу объекта, у которого ни выручки,
    ни своего пула расходов нет вовсе.

    Объект, не включённый ни в одной очереди, в ответ не попадает: пустое
    значит «движок его не строил», и книга останется на своём умолчании, а не
    получит выдуманную единицу.
    """
    out: dict[str, int] = {}
    for index, phase in enumerate(phases or []):
        phase_inputs = (phase or {}).get("inputs") or {}
        for obj in standalone_objects():
            if obj.key in out:
                continue
            if bool(phase_inputs.get(f"{obj.prefix}_enabled")):
                out[obj.key] = index + 1
    return out


def _v4_add_carry_parity_row(xml: str, target_mln: float, missing: list[str]) -> str:
    """Строка паритета «долг, переданный между очередями» на листе ПРОВЕРКИ.

    Без неё книга могла бы переносить не то и не тогда, а вердикт листа остался
    бы «ПРОЙДЕНО»: остальные строки паритета смотрят на итоги, а перенос между
    очередями итоги проекта почти не двигает — он меняет, КТО платит.
    """
    return _v4_add_parity_row(
        xml, _V4_CARRY_PARITY_ROW, "Паритет: долг, переданный между очередями, млн",
        "+".join(f"'CF_{phase}'!$B${_V4_CARRY_PASSED_ROW}" for phase in range(1, 5)),
        target_mln, missing, "паритет переноса долга")


def _v4_add_vat_parity_row(xml: str, target_mln: float, missing: list[str]) -> str:
    """Строка паритета НДС.

    Своей строки у НДС не было, и проверялся он только внутри чистой прибыли —
    вместе с EBITDA и стоимостью финансирования, у которых допуск считается
    долей от их собственной величины. На проекте с околонулевой прибылью это
    поймало настоящую ошибку (движок не облагал выручку паркинга отдельно
    стоящих объектов — 166,0 млн выручки и 4,7 млн налога), но узнать, что дело
    в НДС, было нельзя: строка называла прибыль. Теперь величина проверяется
    там, где она считается, а прибыль отвечает за себя.
    """
    return _v4_add_parity_row(
        xml, _V4_VAT_PARITY_ROW, "Паритет: НДС к уплате, млн",
        "+".join(f"'CF_{phase}'!$B$21" for phase in range(1, 5)),
        target_mln, missing, "паритет НДС")


def _v4_teach_funding_check_about_the_limit(
    xml: str, missing: list[str]) -> str:
    """Строка 32 ПРОВЕРОК: «потребность профинансирована банком полностью».

    Проверка была написана тогда, когда выборка равнялась потребности ВСЕГДА
    — потолка у книги не существовало, — и «не сошлось» могло означать ровно
    одно: сломалась арифметика. Отсюда FAIL и «СБОЙ» на весь лист.

    С одобренным лимитом разрыв стал законным: банк дал меньше, чем нужно, и
    это факт о ПРОЕКТЕ, а не о книге. FAIL здесь — та самая кричащая зря
    проверка, которую перестают читать: она загоралась бы у каждого, кто
    впишет настоящий лимит из term sheet. Разрыв без заданного лимита
    по-прежнему FAIL: объяснить его нечем.

    Заодно факт берётся из строк, которые эту дыру и считают
    (`_V4_PF_UNCOVERED_ROW` листов CF), а не вторым таким же вычитанием:
    двух ответов на «сколько не покрыто» в одной книге быть не должно.
    """
    row = 32
    facts = ",".join(f"'CF_{phase}'!$B${_V4_PF_UNCOVERED_ROW}"
                     for phase in range(1, 5))
    edits = (
        ("B", "SUM('CF_1'!D44:", f"SUM({facts})"),
        # Метка — начало формулы шаблона: полный текст с прежним вердиктом
        # здесь не пишется, вердикт паритета собирает один _v4_parity_verdict.
        ("F", 'IF(B32&lt;=C32+E32,"OK",',
         'IF(B32&lt;=C32+E32,"OK",'
         "IF(OR(" + ",".join(
             [f"'Вводные'!${_V4_PF_APPROVED_CELL[0]}${_V4_PF_APPROVED_CELL[1:]}&gt;0"]
             + [f"'Вводные'!${_V4_PF_QUEUE_INDIVIDUAL_COL}${_V4_CF_QUEUE_ENABLED_ROW + k}&gt;0"
                for k in range(4)]) + '),"WARN","FAIL"))'),
    )
    for column, mark, replacement in edits:
        found = re.search(r'<x:c r="%s%d"[^>]*>\s*<x:f>(.*?)</x:f>' % (column, row),
                          xml, re.S)
        if not found:
            missing.append(f"ПРОВЕРКИ: формулы {column}{row} нет")
            return xml
        if mark not in found.group(1):
            missing.append(f"ПРОВЕРКИ: формула {column}{row} не опознана")
            return xml
        xml = xml[:found.start(1)] + replacement + xml[found.end(1):]
    # Подпись обязана объяснить WARN: «не профинансирована полностью» без
    # причины читается как поломка книги. Ячейка подписи живёт ОБЩЕЙ СТРОКОЙ
    # (t="s" с номером в таблице строк), поэтому правится не её тело, а она
    # сама — на inlineStr, как у дописанных строк паритета. Стиль остаётся
    # свой: чужая краска в середине блока читалась бы как чужая вставка.
    label = re.search(r'<x:c r="A%d"(?P<attrs>[^>]*?)(?:/>|>.*?</x:c>)' % row,
                      xml, re.S)
    if not label:
        missing.append(f"ПРОВЕРКИ: подписи A{row} нет")
        return xml
    style = re.search(r'\ss="\d+"', label.group("attrs"))
    text = xml_escape(
        "Потребность в долге профинансирована банком полностью "
        "(WARN — не хватило ОДОБРЕННОГО лимита ПФ, см. строку "
        f"{_V4_PF_UNCOVERED_ROW} листов CF)")
    cell = (f'<x:c r="A{row}"{style.group(0) if style else ""} t="inlineStr">'
            f"<x:is><x:t>{text}</x:t></x:is></x:c>")
    return xml[:label.start()] + cell + xml[label.end():]


def _v4_add_pf_uncovered_parity_row(
    xml: str, target_mln: float, missing: list[str]) -> str:
    """Строка паритета: непокрытая потребность в ПФ.

    Пока книга не знала об одобренном лимите, сверять было нечего — у неё
    дыры не бывало по построению. Теперь бывает, и она обязана совпасть с
    движковой: иначе «потолок в книге есть» значит только, что он есть.
    """
    return _v4_add_parity_row(
        xml, _V4_PF_UNCOVERED_PARITY_ROW,
        "Паритет: непокрытая потребность в ПФ, млн",
        "+".join(f"'CF_{phase}'!$B${_V4_PF_UNCOVERED_ROW}" for phase in range(1, 5)),
        target_mln, missing, "паритет непокрытой потребности")


def _v4_parity_targets(consolidated: dict[str, Any]) -> dict[str, float]:
    """Контрольные числа движка для parity-блока листа ПРОВЕРКИ."""
    summary = consolidated.get("summary") or {}
    finance = consolidated.get("finance") or {}
    return {
        "revenue_mln": float(summary.get("revenue") or 0) / 1e6,
        "capex_mln": float(summary.get("capex") or 0) / 1e6,
        "ebitda_mln": float(summary.get("ebitda") or 0) / 1e6,
        "financing_mln": float(summary.get("financing_cost") or 0) / 1e6,
        "tax_mln": float(summary.get("profit_tax") or 0) / 1e6,
        "net_profit_mln": float(summary.get("net_profit") or 0) / 1e6,
        "vat_mln": float(finance.get("vat") or 0) / 1e6,
        "llcr": float(summary.get("llcr") or 0),
        "peak_bridge_mln": float(finance.get("peak_bridge") or 0) / 1e6,
        "peak_pf_mln": float(finance.get("peak_pf") or 0) / 1e6,
        "pf_shortfall_mln": float(finance.get("pf_shortfall") or 0) / 1e6,
    }


# Продукты пула МКД на листе «Продажи»: строка объёма и строка выручки в
# блоке очереди 1 (шаг блока — `_V4_SALES_PHASE_STRIDE`). Порядок — порядок
# `COST_POOL_PRODUCTS`, тот же, что у слагаемых строки налоговой базы.
_V4_POOL_SALES_ROWS: dict[str, tuple[int, int]] = {
    "apartments": (14, 16), "ground_commercial": (17, 19),
    "underground_parking": (20, 22), "storage": (23, 25),
}
_V4_POOL_LABELS: dict[str, str] = {
    "apartments": "квартиры", "ground_commercial": "коммерция 1 этажа",
    "underground_parking": "подземный паркинг", "storage": "кладовые",
}


def _v4_balanced(text: str, start: int) -> int:
    """Индекс за закрывающей скобкой группы, открытой в `start`, или -1."""
    depth = 0
    for index in range(start, len(text)):
        if text[index] == "(":
            depth += 1
        elif text[index] == ")":
            depth -= 1
            if depth == 0:
                return index + 1
    return -1


def _v4_apply_core_pool_by_product(xml: str, phase: int, missing: list[str]) -> str:
    """Себестоимость продуктов пула МКД — «прямые по месту + общие по выручке».

    Решение владельца 15.10.2026, методика движка `cost_pool_product_costs`.
    Строка налоговой базы делила весь пул очереди долей СУММЫ проданного
    количества — метры квартир и коммерции складывались со ШТУКАМИ паркинга и
    кладовых. Теперь пул раскладывается по продуктам блоком внизу листа:
      · прямые — СМР наземной и подземной части с надбавками (генподряд,
        техзаказчик, управление проектом, резерв — тот же процент с СМР
        части), продуктам своей части по их ГНС; наземные кладовые лежат в
        ГНС квартир и идут с наземной частью (`storage_placement`);
      · общие — остаток пула, по выручке продуктов (без выручки — по ГНС);
    и каждый продукт признаёт СВОЮ себестоимость долей своего проданного
    объёма в своих единицах. Продукт без продаж и остаток, который делить не
    по чему, признаются в РВЭ. Проход идёт после признания дописанных
    объектов и пула без продаж: пул к этому моменту уже окончательный.
    """
    sheet = f"CF_{phase}"
    head = f"IFERROR((SUM({_v4_month_span(20)})-'ОБЪЕКТЫ'!$B${96 + 8 * (phase - 1)}"
    numbers = [int(number) for number in re.findall(r'<x:row r="(\d+)"', xml)]
    base = (max(numbers) if numbers else 0) + 2
    rows = {"pool": base + 1, "direct_above": base + 2, "direct_under": base + 3,
            "scale": base + 4, "common": base + 5}
    for index, key in enumerate(COST_POOL_PRODUCTS):
        rows[key] = base + 6 + index
    rows["unallocated"] = base + 6 + len(COST_POOL_PRODUCTS)
    shift = _V4_SALES_PHASE_STRIDE * (phase - 1)
    want = [_V4_POOL_SALES_ROWS[key][0] + shift for key in COST_POOL_PRODUCTS]
    pools: list[str] = []

    def build(column: str, body: str) -> str:
        at = body.find(head)
        if at < 0:
            return body
        start = at + len("IFERROR(")
        pool_end = _v4_balanced(body, start)
        if pool_end < 0 or not body.startswith("*(", pool_end):
            return body
        sold_end = _v4_balanced(body, pool_end + 1)
        if sold_end < 0 or not body.startswith("/(", sold_end):
            return body
        total_end = _v4_balanced(body, sold_end + 1)
        if total_end < 0 or not body.startswith(",IF(", total_end):
            return body
        call_end = _v4_balanced(body, at + len("IFERROR"))
        if call_end < 0:
            return body
        sold = body[pool_end + 1:sold_end]
        got = [int(r) for r in re.findall(r"'Продажи'!" + column + r"(\d+)", sold)]
        if got != want:
            return body
        pools.append(body[start:pool_end])
        in_rve = f"{_v4_month_offset(column, '$B$8')}=0"
        terms = []
        for key, volume in zip(COST_POOL_PRODUCTS, want):
            cost = f"$B${rows[key]}"
            terms.append(f"IFERROR({cost}*'Продажи'!{column}{volume}"
                         f"/SUM({_v4_month_span(volume, 'Продажи')}),IF({in_rve},{cost},0))")
        terms.append(f"IF({in_rve},$B${rows['unallocated']},0)")
        return body[:at] + "(" + "+".join(terms) + ")" + body[call_end:]

    xml, count = _v4_rewrite_row_formulas(xml, 22, head, build)
    columns = len(_v4_cf_columns())
    if count < columns or len(pools) < columns or len(set(pools)) != 1:
        missing.append(f"{sheet}: строка налоговой базы не разложила пул МКД по продуктам "
                       f"({len(pools)} из {columns})")
        return xml

    queue = _V4_CF_QUEUE_ENABLED_ROW + phase - 1
    capex = _V4_CAPEX_BLOCK_STRIDE * (phase - 1)
    flats, shops, garage, store = (f"'Вводные'!${col}${queue}" for col in ("I", "J", "K", "AR"))
    under, above = _v4_storage_under_flag(), _v4_storage_above_flag()
    inside = f"MIN({store},{flats})*{above}"
    part_gns = {
        "apartments": (f"({flats}-{inside})", "0"),
        "ground_commercial": (shops, "0"),
        "underground_parking": ("0", garage),
        "storage": (inside, f"{store}*{under}"),
    }
    above_total, under_total = f"({flats}+{shops})", f"({garage}+{store}*{under})"
    uplift = ("(1+'Вводные'!$B$50+'Вводные'!$B$53+'Вводные'!$B$52)"
              "*(1+'Вводные'!$B$54)")
    revenue = {key: f"'Продажи'!$B${_V4_POOL_SALES_ROWS[key][1] + shift}"
               for key in COST_POOL_PRODUCTS}
    revenue_total = "(" + "+".join(revenue.values()) + ")"
    gns_total = f"({above_total}+{under_total})"
    b = {key: f"$B${row}" for key, row in rows.items()}
    formulas = {
        "pool": pools[0],
        "direct_above": f"'CAPEX'!$B${_V4_CAPEX_ARTICLE_ROW['main_above'] + capex}*{uplift}",
        "direct_under": f"'CAPEX'!$B${_V4_CAPEX_ARTICLE_ROW['main_under'] + capex}*{uplift}",
        "scale": (f"IF(AND({b['direct_above']}+{b['direct_under']}>0,{b['pool']}>0),"
                  f"MIN(1,{b['pool']}/({b['direct_above']}+{b['direct_under']})),0)"),
        "common": (f"{b['pool']}-{b['scale']}*(IF({above_total}>0,{b['direct_above']},0)"
                   f"+IF({under_total}>0,{b['direct_under']},0))"),
    }
    for key in COST_POOL_PRODUCTS:
        own_above, own_under = part_gns[key]
        weight = (f"IF({revenue_total}>0,{revenue[key]}/{revenue_total},"
                  f"IFERROR(({own_above}+{own_under})/{gns_total},0))")
        formulas[key] = (f"{b['scale']}*(IFERROR({b['direct_above']}*{own_above}/{above_total},0)"
                         f"+IFERROR({b['direct_under']}*{own_under}/{under_total},0))"
                         f"+{b['common']}*{weight}")
    formulas["unallocated"] = (f"{b['pool']}-SUM({b[COST_POOL_PRODUCTS[0]]}"
                               f":{b[COST_POOL_PRODUCTS[-1]]})")
    labels = {
        "pool": "Пул МКД очереди: расходы без объектов",
        "direct_above": "Прямые: СМР наземной части с надбавками",
        "direct_under": "Прямые: СМР подземной части с надбавками",
        "scale": "Доля прямых, вошедшая в пул (не больше 1)",
        "common": "Общие затраты — делятся по выручке продуктов",
        **{key: f"Себестоимость: {_V4_POOL_LABELS[key]}" for key in COST_POOL_PRODUCTS},
        "unallocated": "Не распределено — признаётся в РВЭ",
    }
    parts = [f'<x:row r="{base}">' + _v4_head_cell(
        f"A{base}", f"СЕБЕСТОИМОСТЬ ПРОДУКТОВ ПУЛА МКД · ОЧЕРЕДЬ {phase}: "
        "прямые — по ГНС своей части, общие — по выручке") + "</x:row>"]
    for key, row in rows.items():
        unit = "доля" if key == "scale" else "млн ₽"
        parts.append(
            f'<x:row r="{row}">'
            f'<x:c r="A{row}" t="inlineStr"><x:is><x:t>{xml_escape(labels[key])}</x:t></x:is></x:c>'
            f'<x:c r="B{row}"><x:f>{xml_escape(formulas[key])}</x:f></x:c>'
            f'<x:c r="C{row}" t="str"><x:v>{unit}</x:v></x:c>'
            "</x:row>")
    tail = "</x:sheetData>"
    if tail not in xml:
        missing.append(f"{sheet}: блок себестоимости продуктов не записан")
        return xml
    return xml.replace(tail, "".join(parts) + tail, 1)
