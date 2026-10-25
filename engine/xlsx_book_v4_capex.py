"""Лист CAPEX книги v4: строки статей, общепроектные кассовые доли, снос и
расселение, рассрочка ВРИ, база НДС, «под ключ», нежилая смета, профиль
управления проектом.

Кусок книги v4: файлы `engine/xlsx_book_v4*.py` выполняются подряд на
месте прежнего `_engine_part("xlsx_book_v4")`. Не самостоятельный модуль:
`main_legacy.py` выполняет этот файл в своём пространстве имён
(`_engine_part("xlsx_book_v4_capex")`), поэтому здесь видны все имена движка,
а подмены `core.<имя>` из `main.py` и тестов действуют и на этот код.
Отдельно не импортировать: второй импорт поднял бы вторую копию движка.
"""
# ruff: noqa

MANAGEMENT_PROFILE_ARTICLES: tuple[str, ...] = (
    "ird", "design_p", "design_rd", "author_supervision", "preparation",
    "main_above", "main_under", "utilities", "landscaping", "site_maintenance",
)

# Строки листа CAPEX книги v4: блок очереди повторяется каждые 34 строки.
_V4_CAPEX_ARTICLE_ROW: dict[str, int] = {
    "ird": 16, "design_p": 17, "design_rd": 18, "author_supervision": 19,
    "preparation": 20, "main_above": 21, "main_under": 22, "utilities": 23,
    "landscaping": 24, "commissioning": 25, "site_maintenance": 26,
    "technical_supervision": 27, "project_management": 28, "gc_fee": 29,
    "reserve": 30,
}
_V4_CAPEX_BLOCK_STRIDE = 34
_V4_CAPEX_PHASES = 4


# Общепроектные статьи, которые движок раскладывает по очередям КАССОВЫМИ
# долями, а книга считала от площадей самой очереди. Пока доли близки, разница
# прячется в допуске; на проекте с длинным ИРД и крупной покупкой это 686 млн ₽
# пика БРИДЖа при совпадающих выручке и CAPEX. Ключ — как в `shared_cash`
# движка, строка — в блоке CAPEX первой очереди, колонка — где книга берёт долю.
_V4_SHARED_CASH_ARTICLES = (
    ("ird", 16, "AJ"),
    ("design", 17, "AK"),
    ("design", 18, "AK"),
    ("preparation", 20, "AL"),
    ("utilities", 23, "AM"),
)
# Площади очереди в формуле статьи: их место занимает площадь всего проекта,
# а долю очереди даёт кассовый вес.
_V4_QUEUE_AREA = "('Вводные'!$I${row}+'Вводные'!$J${row}+'Вводные'!$K${row})"
_V4_PROJECT_AREA = "SUM('Вводные'!$I$88:$K$91)"


def _v4_apply_shared_cash_articles(xml: str, missing: list[str]) -> str:
    """Приводит ИРД, проектирование, подготовку и наружные сети к кассовым долям.

    Движок берёт эти статьи по проекту целиком и делит их между очередями
    кассовыми долями (`shared_cash`; умолчание — веса очередей, со страницы
    приходит front-loaded пресет). Книга считала каждую от площадей своей
    очереди — то есть по долям ТЭП. Отношение выборок на проверочном проекте
    вышло 1,571, а это ровно 55/35: кассовая доля против доли ТЭП.

    Формула не переписывается вслепую: не опознали — в `missing`.
    """
    for phase in range(_V4_CAPEX_PHASES):
        base = _V4_CAPEX_BLOCK_STRIDE * phase
        queue_row = _V4_CF_QUEUE_ENABLED_ROW + phase
        area = _V4_QUEUE_AREA.format(row=queue_row)
        for _key, article_row, column in _V4_SHARED_CASH_ARTICLES:
            row = article_row + base
            pattern = re.compile(r'(<x:c r="B%d"[^>]*>)(.*?)(</x:c>)' % row, re.S)
            found = pattern.search(xml)
            if not found:
                missing.append(f"CAPEX · кассовая доля: ячейка B{row} не найдена")
                return xml
            body = found.group(2)
            if area not in body:
                missing.append(
                    f"CAPEX · кассовая доля: формула B{row} не опознана "
                    f"(нет площадей очереди {phase + 1})")
                return xml
            share = f"'Вводные'!${column}${queue_row}"
            body = body.replace(area, f"({_V4_PROJECT_AREA}*{share})", 1)
            xml = (xml[:found.start()] + found.group(1) + body + found.group(3)
                   + xml[found.end():])
    return xml


# Снос и расселение в блоке CAPEX. У шаблона для них нет строки: движок тратил
# 270,9 млн ₽ на снос, книга — ничего, и ровно на эту сумму расходились пик
# БРИДЖа, а следом CAPEX, проценты, налог и LLCR (владелец, 04.09.2026:
# «расхождения опять»). Вводные для них с 0.21.84 стоят в F39–F41 «Вводных»,
# но их не читала ни одна формула. Строки 36–37 каждого блока в шаблоне пусты —
# ряд заводится туда, не сдвигая ни одной ссылки; итог блока (строка 32) и
# база резерва (строка 30) переписываются, чтобы видеть новые строки.
_V4_CAPEX_EXTRA_ROWS = (
    (36, "Снос и демонтаж",
     "('Вводные'!$F$39*'Вводные'!$F$40/1000)"),
    (37, "Расселение", "'Вводные'!$F$41"),
)
_V4_CAPEX_TOTAL_ROW = 32
_V4_CAPEX_RESERVE_ROW = 30
_V4_CAPEX_MONTH_COLUMNS = _V4_MONTH_COLUMNS


def _v4_apply_demolition_rows(xml: str, missing: list[str]) -> str:
    """Строки сноса и расселения в каждом блоке CAPEX по образцу подготовки.

    Сумма — как у строки 20: вводная × сценарий × множитель затрат очереди ×
    признак очереди × коэффициент затрат к старту × кассовая доля (та же
    колонка AL, что у подготовки: движок делит снос её долей). Помесячно —
    окно перед РнС, как у подготовки. Формула не опознана — `missing`.
    """
    from openpyxl.utils import get_column_letter as _col

    def styled(cell_xml: str, coord: str, style: str | None) -> str:
        return cell_xml.replace(f'<x:c r="{coord}"', f'<x:c r="{coord}" s="{style}"', 1) if style else cell_xml

    def style_of(coord: str) -> str | None:
        found = re.search(r'<x:c r="%s"[^>]*?\ss="(\d+)"' % coord, xml)
        return found.group(1) if found else None

    label_style, amount_style, month_style = style_of("A20"), style_of("B20"), style_of("D20")
    for phase in range(_V4_CAPEX_PHASES):
        base = _V4_CAPEX_BLOCK_STRIDE * phase
        queue_row = _V4_CF_QUEUE_ENABLED_ROW + phase
        permit = f"$B${7 + base}"
        for row_at, label, amount in _V4_CAPEX_EXTRA_ROWS:
            row = row_at + base
            xml = _v4_ensure_row(xml, row)
            formula = (f"{amount}*'Вводные'!$H$6*'Вводные'!$T${queue_row}"
                       f"*--('Вводные'!$B${queue_row}=\"Да\")*'Вводные'!$AH${queue_row}"
                       f"*'Вводные'!$AL${queue_row}")
            cells = {f"A{row}": dict(text=label),
                     f"B{row}": dict(formula=formula),
                     f"C{row}": dict(text="млн ₽")}
            styles = {f"A{row}": label_style, f"B{row}": amount_style, f"C{row}": None}
            for index in range(_V4_CAPEX_MONTH_COLUMNS):
                column = _col(4 + index)
                window = f"MIN(6,'Вводные'!$E${queue_row})"
                cells[f"{column}{row}"] = dict(formula=(
                    f"IF(AND({column}$3>=EDATE({permit},-{window}),{column}$3<{permit}),"
                    f"$B${row}/MAX(1,{window}),0)"))
                styles[f"{column}{row}"] = month_style
            xml, done = _v4_set_cells(xml, row, cells, styles)
            if not done:
                missing.append(f"CAPEX · снос и расселение: строка {row} не записана")
                return xml
        # Итог блока видит новые строки: SUM(X14:X31)+ОБЪЕКТЫ → +X36+X37.
        total_row = _V4_CAPEX_TOTAL_ROW + base
        extra = [row_at + base for row_at, _label, _amount in _V4_CAPEX_EXTRA_ROWS]
        pattern = re.compile(
            r'(<x:c r="(?P<col>[A-Z]{1,3})%d"[^>]*><x:f>)SUM\((?P=col)%d:(?P=col)%d\)\+'
            % (total_row, 14 + base, 31 + base))
        xml, count = pattern.subn(
            lambda m: m.group(1) + "SUM(%s%d:%s%d)+%s+" % (
                m.group("col"), 14 + base, m.group("col"), 31 + base,
                "+".join(f"{m.group('col')}{r}" for r in extra)), xml)
        if count < _V4_CAPEX_MONTH_COLUMNS:
            missing.append(f"CAPEX · итог очереди {phase + 1}: формула строки {total_row} не опознана "
                           f"({count} колонок из {_V4_CAPEX_MONTH_COLUMNS})")
        # База резерва — все статьи блока: (SUM(B15:B29)+B31+…) → +B36+B37.
        reserve_row = _V4_CAPEX_RESERVE_ROW + base
        pattern = re.compile(r'(<x:c r="B%d"[^>]*><x:f>\(SUM\(B%d:B%d\)\+B%d)' % (
            reserve_row, 15 + base, 29 + base, 31 + base))
        xml, count = pattern.subn(lambda m: m.group(1) + "".join(f"+B{r}" for r in extra), xml)
        if count != 1:
            missing.append(f"CAPEX · резерв очереди {phase + 1}: формула B{reserve_row} не опознана")
    return xml


# Лист «ВРИ»: четыре блока по 13 строк, первый начинается на строке 6.
_V4_VRI_BLOCK_STRIDE = 13
_V4_VRI_PRINCIPAL_ROW = 12     # погашение основного долга
_V4_VRI_INTEREST_ROW = 13      # проценты по рассрочке
_V4_VRI_SECURITY_ROW = 14      # расходы на обеспечение
# Строка CAPEX, куда уходят проценты и обеспечение: 36 и 37 заняты сносом и
# расселением, 38 в шаблоне пуста.
_V4_CAPEX_VRI_EXTRA_ROW = 38
_V4_CF_VAT_ROW = 21                  # CF очереди: «НДС к уплате»
_V4_CAPEX_VRI_ROW = 15


def _v4_apply_vri_installment_start(xml: str, missing: list[str]) -> str:
    """Первый платёж рассрочки ВРИ — в дату обязательства, а не периодом позже.

    Правило движка с 08.2026: рассрочка начинается в дату обязательства.
    Шаблон остался на прежней методике — условие платежа требовало, чтобы от
    даты обязательства прошло БОЛЬШЕ нуля месяцев, и первый взнос уезжал на
    период вперёд. Число платежей при этом сохранялось, поэтому расходилось не
    тело, а проценты: на трёхлетней квартальной рассрочке книга начисляла их
    двенадцать раз вместо одиннадцати и на полную сумму — 317,9 млн ₽ против
    269,1 у движка при одинаковой плате.

    Замена точечная: «прошло > 0 месяцев» → «прошло >= 0», а верхняя граница
    становится строгой, иначе платежей стало бы тринадцать. Формула шаблона не
    опознана — `missing`, а не молча посчитанный по-своему график.
    """
    expected = _V4_CAPEX_MONTH_COLUMNS * 2 * _V4_CAPEX_PHASES
    start_from, start_count = ")&gt;0,MOD(", xml.count(")&gt;0,MOD(")
    upper_from, upper_count = "&lt;=&#39;Вводные&#39;!$B$77*12", 0
    # Апостроф в XML может быть и целым, и мнемоникой — считаем оба написания.
    for candidate in ("&lt;='Вводные'!$B$77*12", "&lt;=&#39;Вводные&#39;!$B$77*12"):
        if xml.count(candidate):
            upper_from, upper_count = candidate, xml.count(candidate)
            break
    if start_count != expected or upper_count != expected:
        missing.append(
            f"ВРИ · начало рассрочки: условие графика не опознано "
            f"({start_count} и {upper_count} вместо {expected})")
        return xml
    xml = xml.replace(start_from, ")&gt;=0,MOD(")
    xml = xml.replace(upper_from, upper_from.replace("&lt;=", "&lt;", 1))
    return xml


_V4_VRI_BALANCE_END_ROW = 16   # остаток на конец месяца
_V4_VRI_RATE_ROW = 17          # ставка по рассрочке
_V4_VRI_ACCRUAL_ROW = 18       # накопленные проценты — служебная строка блока


def _v4_apply_vri_monthly_accrual(xml: str, missing: list[str]) -> str:
    """Проценты по рассрочке ВРИ копятся помесячно, как в движке.

    Шаблон начислял их в месяц платежа: остаток × ставка ЭТОГО месяца ×
    периодичность / 12. Движок идёт по месяцам между платежами и берёт ставку
    каждого месяца (`build_vri_schedule`), а ключевая ставка у нас снижается —
    на трёхлетней рассрочке книга недосчитывала около двух процентов.

    Свободная строка блока (18, 31, 44, 57 — в шаблоне пусты) копит начисление
    на остаток КОНЦА месяца: в месяц платежа это остаток после погашения, то
    есть ровно та база, на которой считает движок. Платёж забирает накопленное
    за прошлые месяцы, поэтому на первом взносе процентов нет.
    """
    from openpyxl.utils import get_column_letter as _col

    def style_of(coord: str) -> str | None:
        found = re.search(r'<x:c r="%s"[^>]*?\ss="(\d+)"' % coord, xml)
        return found.group(1) if found else None

    label_style, month_style = style_of("A13"), style_of("D13")
    for phase in range(_V4_CAPEX_PHASES):
        base = _V4_VRI_BLOCK_STRIDE * phase
        principal = _V4_VRI_PRINCIPAL_ROW + base
        interest = _V4_VRI_INTEREST_ROW + base
        balance_end = _V4_VRI_BALANCE_END_ROW + base
        rate = _V4_VRI_RATE_ROW + base
        accrual = _V4_VRI_ACCRUAL_ROW + base
        xml = _v4_ensure_row(xml, accrual)
        carried_cells = {f"A{accrual}": dict(text="Накопленные проценты (служебная строка)"),
                         f"B{accrual}": dict(formula=f"{_V4_LAST_COLUMN}{accrual}")}
        styles = {f"A{accrual}": label_style, f"B{accrual}": None}
        paid_cells: dict[str, dict[str, Any]] = {}
        for index in range(_V4_CAPEX_MONTH_COLUMNS):
            column = _col(4 + index)
            previous = _col(3 + index)
            carried = "0" if index == 0 else f"{previous}{accrual}"
            month = (f"IF(AND('Вводные'!$B$75<>\"Единовременно\",'Вводные'!$B$80=\"Да\","
                     f"{column}{balance_end}>0),{column}{balance_end}*{column}{rate}/12,0)")
            carried_cells[f"{column}{accrual}"] = dict(
                formula=f"{carried}+{month}-{column}{interest}")
            styles[f"{column}{accrual}"] = month_style
            paid_cells[f"{column}{interest}"] = dict(formula=(
                "0" if index == 0 else f"IF({column}{principal}>0,{previous}{accrual},0)"))
        for target, payload, style_map in ((accrual, carried_cells, styles),
                                           (interest, paid_cells, None)):
            xml, done = _v4_set_cells(xml, target, payload, style_map)
            if not done:
                missing.append(f"ВРИ · накопление процентов: строка {target} не записана")
                return xml
    return xml


def _v4_vat_exempt_ddu(formula: str) -> str | None:
    """Нежилое дома, проданное до РВЭ по ДДУ, выходит из облагаемой выручки.

    Пп. 23.1 п. 3 ст. 149 НК освобождает услуги застройщика по ДДУ на
    «входящие в состав многоквартирных домов жилые и (или) нежилые
    помещения, машино-места» (решение владельца 25.10.2026, ревизия
    docs/vat_ddu_review.md). Шаблон облагал «всю выручку минус квартиры» —
    в обоих местах формулы (начисление и доля вычета) к квартирам теперь
    прибавляется проданное ДО РВЭ очереди из строк коммерции, паркинга и
    кладовых её блока листа «Продажи» (квартиры + 3, + 6, + 9). Движок
    вычитает то же (`VAT_MKD_NONRES_PRODUCTS`, `vat_exempt_ddu`).
    """
    found = re.search(r"SUM\('Продажи'!\$D\$(\d+):\$GA\$\1\)", formula)
    if not found:
        return None
    apartments = int(found.group(1))
    old = found.group(0)
    exempt = "+".join(
        f"SUMPRODUCT(--('Продажи'!$D$3:$GA$3<$B$8),'Продажи'!$D${row}:$GA${row})"
        for row in (apartments + 3, apartments + 6, apartments + 9))
    if formula.count(old) != 2:
        return None
    # В шаблоне строка квартир ВЫЧИТАЕТСЯ: «вся выручка − SUM(квартиры)».
    # Освобождённое нежилое уходит вместе с ней — в одних скобках, иначе
    # минус перед квартирами его бы прибавил.
    return formula.replace(old, f"({old}+{exempt})")


def _v4_apply_vat_base(xml: str, phase: int, missing: list[str]) -> str:
    """Проценты по рассрочке ВРИ НДС не облагаются — значит и вычета не дают.

    Строка 21 листа CF берёт базу вычета как «все расходы очереди минус покупка
    и минус тело платы за ВРИ». Проценты и обеспечение рассрочки приезжают в
    расходы строкой 38 CAPEX (0.23.х) и в эту базу входили, то есть книга
    возвращала НДС с процентов. Движок их из базы вычитает
    (`vri_interest`, `vri_security` в `vat_bearing_costs`), и на проекте с
    трёхлетней рассрочкой расхождение выходило 5,9 млн ₽ — при совпадающих
    CAPEX, выручке и EBITDA. Нашла его строка паритета НДС первым же прогоном.

    Формула шаблона не опознана — `missing`, а не тихий вычет по прежней базе.
    """
    row = _V4_CF_VAT_ROW
    # `phase` здесь — НОМЕР листа CF (1..4), как в остальных проходах этого
    # цикла: блок CAPEX у первой очереди идёт без смещения.
    extra = (f"'CAPEX'!$B$"
             f"{_V4_CAPEX_VRI_EXTRA_ROW + _V4_CAPEX_BLOCK_STRIDE * (phase - 1)}")
    was = "(SUM($D$20:$GA$20)-"
    now = f"(SUM($D$20:$GA$20)-{extra}-"
    # Из базы исключается плата за ВРИ, которой в расходах нет. Шаблон вычитал
    # НАЧИСЛЕННУЮ (B16 × доля очереди): при льготе 100 % уплачено 0, а база
    # уменьшалась на 10 166,6 млн — и книга платила на 324,7 млн НДС больше
    # движка (Нагатино КРТ, 17.10.2026). Движок вычитает уплаченное тело —
    # оно же стоит строкой 15 блока CAPEX очереди.
    vri_was = (f"-'Вводные'!$B$16*'Вводные'!$Q$"
               f"{_V4_CF_QUEUE_ENABLED_ROW + phase - 1}")
    vri_now = f"-'CAPEX'!$B${_V4_CAPEX_VRI_ROW + _V4_CAPEX_BLOCK_STRIDE * (phase - 1)}"
    formulas = _v4_row_formulas(xml, row)
    cells: dict[str, dict[str, Any]] = {}
    for coord, formula in formulas.items():
        # Колонка берётся целиком: «AC21» начинается на «A», и проверка по
        # первой букве выбрасывала весь блок AA..CZ — 81 ячейку из 183, ровно
        # ту середину строки, где и стоит живой месяц. Снаружи это выглядело
        # как «правка не сработала»: D21 с правкой, а число прежнее.
        column = re.match(r"[A-Z]+", coord).group(0)
        if column in ("A", "B", "C"):
            continue
        if was not in formula or vri_was not in formula:
            continue
        formula = formula.replace(was, now, 1).replace(vri_was, vri_now, 1)
        exempt = _v4_vat_exempt_ddu(formula)
        if exempt is None:
            continue
        cells[coord] = dict(formula=exempt)
    if not cells:
        missing.append(f"CF очереди {phase}: база НДС в строке {row} не опознана")
        return xml
    xml, done = _v4_set_cells(xml, row, cells)
    if not done:
        missing.append(f"CF очереди {phase}: база НДС строки {row} не записана")
    return xml


def _v4_apply_vri_interest_row(xml: str, missing: list[str]) -> str:
    """Проценты и обеспечение по рассрочке ВРИ — своей строкой, вне базы резерва.

    Строка 15 CAPEX брала ВЕСЬ денежный платёж по ВРИ (тело, проценты и
    обеспечение), а база резерва — это `SUM(B15:B29)`: книга начисляла резерв и
    на проценты. Движок берёт резерв до того, как проценты и обеспечение
    становятся статьями (они прибавляются после `base_for_overheads`), и на
    317,9 млн ₽ процентов это давало лишние 15,9 млн ₽ резерва.

    Поэтому строка 15 несёт только тело, а проценты с обеспечением встают
    строкой 38 своего блока: она входит в итог блока и не входит в базу резерва.
    """
    from openpyxl.utils import get_column_letter as _col

    def style_of(coord: str) -> str | None:
        found = re.search(r'<x:c r="%s"[^>]*?\ss="(\d+)"' % coord, xml)
        return found.group(1) if found else None

    label_style, amount_style, month_style = style_of("A15"), style_of("B15"), style_of("D15")
    for phase in range(_V4_CAPEX_PHASES):
        base = _V4_CAPEX_BLOCK_STRIDE * phase
        vri_base = _V4_VRI_BLOCK_STRIDE * phase
        principal = _V4_VRI_PRINCIPAL_ROW + vri_base
        interest = _V4_VRI_INTEREST_ROW + vri_base
        security = _V4_VRI_SECURITY_ROW + vri_base
        row = _V4_CAPEX_VRI_EXTRA_ROW + base
        vri_row = _V4_CAPEX_VRI_ROW + base
        xml = _v4_ensure_row(xml, row)
        extra = {f"A{row}": dict(text="Проценты и обеспечение по рассрочке ВРИ"),
                 f"B{row}": dict(formula=(
                     f"SUM('ВРИ'!D{interest}:{_V4_LAST_COLUMN}{interest})"
                     f"+SUM('ВРИ'!D{security}:{_V4_LAST_COLUMN}{security})")),
                 f"C{row}": dict(text="млн ₽")}
        styles = {f"A{row}": label_style, f"B{row}": amount_style, f"C{row}": None}
        body = {f"B{vri_row}": dict(
            formula=f"SUM('ВРИ'!D{principal}:{_V4_LAST_COLUMN}{principal})")}
        for index in range(_V4_CAPEX_MONTH_COLUMNS):
            column = _col(4 + index)
            extra[f"{column}{row}"] = dict(
                formula=f"'ВРИ'!{column}{interest}+'ВРИ'!{column}{security}")
            styles[f"{column}{row}"] = month_style
            body[f"{column}{vri_row}"] = dict(formula=f"'ВРИ'!{column}{principal}")
        for target, payload, style_map in ((row, extra, styles), (vri_row, body, None)):
            xml, done = _v4_set_cells(xml, target, payload, style_map)
            if not done:
                missing.append(f"CAPEX · проценты ВРИ: строка {target} не записана")
                return xml
        # Итог блока видит новую строку; база резерва (B30) её не видит намеренно.
        total_row = _V4_CAPEX_TOTAL_ROW + base
        pattern = re.compile(
            r'(<x:c r="(?P<col>[A-Z]{1,3})%d"[^>]*><x:f>SUM\((?P=col)%d:(?P=col)%d\)'
            r'(?:\+(?P=col)\d+)*)' % (total_row, 14 + base, 31 + base))
        xml, count = pattern.subn(
            lambda m: m.group(1) + f"+{m.group('col')}{row}", xml)
        if count < _V4_CAPEX_MONTH_COLUMNS:
            missing.append(f"CAPEX · итог очереди {phase + 1}: строка процентов ВРИ не добавлена "
                           f"({count} колонок из {_V4_CAPEX_MONTH_COLUMNS})")
    return xml


# Строки статей блока CAPEX, чья база меняется методикой «объект под ключ».
_V4_CAPEX_TECH_SUPERVISION_ROW = 27
_V4_CAPEX_GC_FEE_ROW = 29


def _v4_apply_turnkey_overheads(xml: str, missing: list[str]) -> str:
    """Генподряд, техзаказчик и резерв — на базах движка (решение 06.10.2026).

    Шаблон брал техзаказчика (B27) и генподряд (B29) от СМР ядра, соцстройки
    И «CAPEX объектов» с листа «ОБЪЕКТЫ», а резерв (B30) — от `SUM(B15:B29)`,
    то есть и от платы за смену ВРИ (B15). Методика теперь иная: ставка
    объекта — «под ключ», генподряд и техзаказчик в ней уже есть; плата за ВРИ
    — известная сумма, резерв на неё не нужен. Правка точечная, поверх формулы
    шаблона: из B27/B29 уходит слагаемое объектов, у резерва сумма начинается
    со строки 16. Не опознанная формула — в `missing`, а не молчаливое
    расхождение с движком.
    """
    objects_term = re.compile(r"\+'ОБЪЕКТЫ'!\$B\$\d+")
    for phase in range(_V4_CAPEX_PHASES):
        base = _V4_CAPEX_BLOCK_STRIDE * phase
        for row_at, label in ((_V4_CAPEX_TECH_SUPERVISION_ROW, "техзаказчик"),
                              (_V4_CAPEX_GC_FEE_ROW, "генподряд")):
            row = row_at + base
            formula = _v4_cell_formula(xml, f"B{row}")
            fixed, count = objects_term.subn("", formula or "", count=1)
            if count != 1:
                missing.append(f"CAPEX · {label} очереди {phase + 1}: формула B{row} не опознана")
                continue
            xml, done = _v4_set_cells(xml, row, {f"B{row}": dict(formula=fixed)})
            if not done:
                missing.append(f"CAPEX · {label} очереди {phase + 1}: B{row} не записана")
        row = _V4_CAPEX_RESERVE_ROW + base
        formula = _v4_cell_formula(xml, f"B{row}") or ""
        was = f"SUM(B{_V4_CAPEX_VRI_ROW + base}:B{29 + base})"
        if was not in formula:
            missing.append(f"CAPEX · резерв очереди {phase + 1}: формула B{row} не опознана")
            continue
        fixed = formula.replace(was, f"SUM(B{_V4_CAPEX_VRI_ROW + 1 + base}:B{29 + base})", 1)
        xml, done = _v4_set_cells(xml, row, {f"B{row}": dict(formula=fixed)})
        if not done:
            missing.append(f"CAPEX · резерв очереди {phase + 1}: B{row} не записана")
    return xml


# Статьи проекта в блоке CAPEX очереди: разносимые долями движка по очередям
# (база — сумма площадей всех очередей × доля) и считаемые своей очередью.
_V4_SHARED_AREA_ARTICLE_ROWS = (16, 17, 18, 20, 23)
_V4_PHASE_AREA_ARTICLE_ROWS = (25, 26)
_V4_CAPEX_MANAGEMENT_ROW = 28


def _v4_apply_nonres_full_estimate(xml: str, inputs: dict[str, Any], missing: list[str]) -> str:
    """Чисто нежилой проект: полная смета формулами, как в движке.

    Решение владельца 06.10.2026 (`object_rate_is_turnkey`): ставка метрового
    объекта нежилого проекта — СМР здания. Статьи проекта (ИРД, П, РД,
    подготовка, сети, ввод, содержание) идут по своим ставкам от общей
    площади объектов — здание (GBA) плюс свой гараж (места × площадь
    места), — а генподряд, техзаказчик и управление начисляются и на СМР
    объекта. Шаблон считал эти статьи от площади МКД (у нежилого — ноль) и без
    объектов. Правка точечная, поверх формул: к площади добавляются слагаемые
    объектов, к базам процентов — CAPEX объектов своей очереди. Объект «под
    ключ» (наземный паркинг) в эти слагаемые не входит. Не опознанная формула —
    в `missing`.
    """
    def cond(lay: "_V4ObjectLayout") -> str:
        head = lay.object_head
        return (f"AND('ОБЪЕКТЫ'!$B${head + 1}=\"Да\",IFERROR(INDEX('Вводные'!$B$88:$B$91,"
                f"'ОБЪЕКТЫ'!$B${head + 2}),\"Нет\")=\"Да\")")

    def area(lay: "_V4ObjectLayout") -> str:
        head = lay.object_head
        garage = ""
        if lay.parking_under:
            col, row = re.match(r"([A-Z]+)(\d+)$", lay.parking_under).groups()
            garage = f"+'Вводные'!${col}${row}*'Вводные'!$K$158"
        return f"IF({cond(lay)},'ОБЪЕКТЫ'!$B${head + 6}{garage},0)"

    smr = [lay for lay in _v4_layouts() if not object_rate_is_turnkey(inputs, lay.obj)]
    if not smr:
        return xml
    total_area = "+".join(area(lay) for lay in smr)

    def phase_area(phase: int) -> str:
        return "+".join(f"IF('ОБЪЕКТЫ'!$B${lay.object_head + 2}={phase + 1},{area(lay)},0)"
                        for lay in smr)

    def phase_works(phase: int) -> str:
        return "+".join(f"IF('ОБЪЕКТЫ'!$B${lay.object_head + 2}={phase + 1},"
                        f"'ОБЪЕКТЫ'!$B${lay.object_head + 22},0)" for lay in smr)

    def patch(row: int, was: str, now: str, label: str) -> None:
        nonlocal xml
        formula = _v4_cell_formula(xml, f"B{row}") or ""
        if was not in formula:
            missing.append(f"CAPEX · нежилой проект, {label}: формула B{row} не опознана")
            return
        xml, done = _v4_set_cells(xml, row, {f"B{row}": dict(formula=formula.replace(was, now, 1))})
        if not done:
            missing.append(f"CAPEX · нежилой проект, {label}: B{row} не записана")

    # Площадь кладовых в базах уже стоит с признаком размещения
    # (`_v4_apply_storage_area_rows`): хвост базы — он, а не голая колонка AR.
    flag = _v4_storage_under_flag()
    shared = f"SUM('Вводные'!$AR$88:$AR$91)*{flag}"
    for phase in range(_V4_CAPEX_PHASES):
        base = _V4_CAPEX_BLOCK_STRIDE * phase
        for row in _V4_SHARED_AREA_ARTICLE_ROWS:
            patch(row + base, shared + ")", f"{shared}+{total_area})",
                  f"статья проекта очереди {phase + 1}")
        own = f"'Вводные'!$AR${88 + phase}*{flag})"
        for row in _V4_PHASE_AREA_ARTICLE_ROWS:
            patch(row + base, own, f"'Вводные'!$AR${88 + phase}*{flag}+{phase_area(phase)})",
                  f"статья проекта очереди {phase + 1}")
        works = phase_works(phase)
        for row in (_V4_CAPEX_TECH_SUPERVISION_ROW, _V4_CAPEX_GC_FEE_ROW):
            patch(row + base, f"B{31 + base})*", f"B{31 + base}+{works})*",
                  f"генподряд и техзаказчик очереди {phase + 1}")
        patch(_V4_CAPEX_MANAGEMENT_ROW + base, f"B{26 + base})*", f"B{26 + base}+{works})*",
              f"управление очереди {phase + 1}")
    return xml


def _v4_management_profile_ranges() -> tuple[int, int, list[int]]:
    """Первая и последняя строка профиля и строки, которые в него не входят."""
    rows = sorted(_V4_CAPEX_ARTICLE_ROW[key] for key in MANAGEMENT_PROFILE_ARTICLES)
    first, last = rows[0], rows[-1]
    excluded = [row for row in range(first, last + 1) if row not in rows]
    return first, last, excluded


# Подписи, называющие подземную площадь наземной. ГНС — наземная площадь
# здания (внутренний термин DevelopAid), и у гаража с кладовыми её нет вовсе:
# в их строке стоит площадь подземного этажа. Итог складывает обе величины и
# потому зовётся строительным объёмом — на нём считаются общие статьи и все
# удельные «на метр». Формулы это не двигает: `ТЭП!C36` читает единственная
# ячейка `ОТЧЕТ!G6`, и её саму не читает никто.
_V4_GNS_LABELS: tuple[tuple[str, str, str, str], ...] = (
    ("ТЭП", "C3", "ГНС, м²", ROW_AREA_HEADER_FULL),
    ("ТЭП", "A36", "ИТОГО ПРОЕКТ", f"ИТОГО ПРОЕКТ — {TOTAL_AREA.lower()}"),
    ("ОТЧЕТ", "F6", "ГНС", TOTAL_AREA.name),
    ("ОТЧЕТ", "A33", "Покупка", "Покупка: участок + выкуп ЗУ/ОКС"),
)


# Выкуп ЗУ/ОКС у третьих лиц — строка стоимости сделки (решение владельца,
# 29.09.2026). В книге он стоит своей ячейкой блока сделки, а там, где книга
# берёт цену покупки, — покупка идёт вместе с выкупом: строка «Приобретение»
# каждой очереди (та же доля очереди, тот же график), база вычета НДС листов
# CF и «Покупка» структуры расходов ОТЧЕТа. Движок делает то же самое: выкуп
# платится статьёй покупки и из базы НДС вычитается.
_V4_CAPEX_PURCHASE_ROW = 14
# Апостроф в XML листа бывает и целым, и мнемоникой — опознаются все написания.
_V4_DEAL_PRICE_CELLS = ("'Вводные'!$B$15", "&apos;Вводные&apos;!$B$15",
                        "&#39;Вводные&#39;!$B$15")


def _v4_deal_includes_buyout(xml: str, sheet: str, rows: tuple[int, ...],
                             missing: list[str]) -> str:
    """Цена покупки в названных строках → цена покупки плюс выкуп ЗУ/ОКС."""
    for row in rows:
        cells: dict[str, dict[str, Any]] = {}
        for coord, formula in _v4_row_formulas(xml, row).items():
            for price in _V4_DEAL_PRICE_CELLS:
                buyout = price.replace("$B$15", "$F$16")
                if price in formula and buyout not in formula:
                    cells[coord] = dict(formula=formula.replace(
                        price, f"({price}+{buyout})"))
                    break
        if not cells:
            missing.append(f"{sheet} · строка {row}: цена покупки не опознана — "
                           "выкуп ЗУ/ОКС в стоимость сделки не добавлен")
            continue
        xml, done = _v4_set_cells(xml, row, cells)
        if not done:
            missing.append(f"{sheet} · строка {row}: стоимость сделки с выкупом не записана")
    return xml


def _v4_rename_labels(xml: str, sheet: str, missing: list[str]) -> str:
    """Переименовать подписи листа — по опознанному тексту, а не вслепую.

    Не нашли прежнюю подпись — это `missing`, а не молчание: подпись, которую
    не удалось заменить, продолжает называть подземную площадь наземной, и
    снаружи это неотличимо от переименованной.
    """
    for name, coord, was, becomes in _V4_GNS_LABELS:
        if name != sheet:
            continue
        pattern = re.compile(
            r'(<(?:x:)?c r="%s"[^>]*>\s*<(?:x:)?v>)%s(</(?:x:)?v>)' % (coord, re.escape(was))
        )
        xml, count = pattern.subn(lambda m: m.group(1) + becomes + m.group(2), xml, count=1)
        if not count:
            missing.append(
                f"{sheet}!{coord}: подпись «{was}» не опознана — "
                f"переименование в «{becomes}» не применено")
    return xml


def _v4_apply_management_profile(xml: str, missing: list[str]) -> str:
    """Приводит формулу «Управление проектом» книги v4 к списку движка.

    Шаблон брал сплошной диапазон строк, а в нём лежит и статья, которой в
    профиле движка нет. Формула не переписывается вслепую: не опознали — в
    `missing`, а не молча посчитали по-своему.
    """
    first, last, excluded = _v4_management_profile_ranges()
    if not excluded:
        return xml
    for phase in range(_V4_CAPEX_PHASES):
        base = _V4_CAPEX_BLOCK_STRIDE * phase
        top, bottom = first + base, last + base
        drop = [row + base for row in excluded]
        pattern = re.compile(
            r"SUM\((?P<col>[A-Z]{1,2})%d:(?P=col)%d\)/SUM\(\$D\$%d:\$%s\$%d\)"
            % (top, bottom, top, _V4_LAST_COLUMN, bottom))
        def replace(match: "re.Match[str]") -> str:
            column = match.group("col")
            month = (f"(SUM({column}{top}:{column}{bottom})-"
                     + "-".join(f"{column}{row}" for row in drop) + ")")
            total = (f"(SUM($D${top}:${_V4_LAST_COLUMN}${bottom})-"
                     + "-".join(f"SUM({_v4_month_span(row)})" for row in drop) + ")")
            return f"{month}/{total}"
        xml, count = pattern.subn(replace, xml)
        if not count:
            missing.append(
                f"CAPEX · профиль управления очереди {phase + 1}: формула строки "
                f"{_V4_CAPEX_ARTICLE_ROW['project_management'] + base} не опознана")
    return xml
