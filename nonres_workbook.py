"""Книга нежилого проекта — финансовая модель формулами от одного листа вводных.

Владелец (05.10.2026): «Excel в данном случае не модель, а убогий свод».
Прежняя книга несла восемь готовых чисел движка и формулы только на деньги
объекта; CAPEX, общие затраты, ключевая ставка, налог проекта, NPV и IRR
приходили значениями. Здесь всё считает книга:

* «Вводные» — все исходные данные одним листом: площади, цены, стратегии,
  ставки затрат, кредит, налоги, дисконт, график оплаты участка, прогноз
  ключевой ставки. Варианты — выпадающими списками по-русски, даты — датами.
* «Затраты» — смета «ставка × база» по статьям и её график по месяцам теми же
  кривыми, что у движка (`build_operating_model`): S-кривая книги ПЛАТО
  (`build_curve`), равномерный разнос, разовый платёж, профиль управления.
* «Объект N» — построчный пересказ `developaid_nonres_strategy.object_flows`:
  выручка, эксплуатация, налог на имущество, НДС, кредит; и итог объекта как
  отдельного проекта (`main_legacy.object_result`).
* «Кредит», «Налоги», «Денежный поток» — проект: выборка и погашение,
  НДС и налог на прибыль с переносом убытка (`_profit_tax_schedule`), поток
  проекта и собственника, NPV, IRR, окупаемость.
* «Отчёт» — экономика проекта, структура расходов, финансирование объектов,
  собственное участие, по годам — ссылками на листы.
* «Сверка» — каждый итог книги против авторитетного расчёта движка.

Методика книги и движка обязана совпадать в обе стороны: где книга формулой
повторить движок не может (благоустройство, соцобъекты, рассрочка ВРИ), число
приходит значением движка, подписано «СЧИТАЕТ ДВИЖОК» и названо в `missing`.

Формулы — только те, что считают и Excel, и вычислитель репозитория
(`xlsx_eval.Evaluator`).
"""

from __future__ import annotations

import io
from datetime import date, datetime
from typing import Any

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

import developaid_nonres_strategy as ns
import site_massing_book as massing_book

ENGINE_FILL = PatternFill("solid", fgColor="FFF3E0")
INPUT_FILL = PatternFill("solid", fgColor="FFFDE7")
BOLD = Font(bold=True)
MONEY = "#,##0"
MONTH = "mm.yyyy"

FIRST_ROW = 80  # строка месяца k = 0 на каждом помесячном листе
EXTRA_MONTHS = 12  # объектам — ещё год после горизонта: NOI покупателя

INPUTS_SHEET = "Вводные"
COSTS_SHEET = "Затраты"
CREDIT_SHEET = "Кредит"
TAX_SHEET = "Налоги"
CASH_SHEET = "Денежный поток"
REPORT_SHEET = "Отчёт"
CHECK_SHEET = "Сверка"
PASSED, FAILED = "ПРОЙДЕНО", "РАСХОЖДЕНИЯ"

# Выпадающие списки: код движка → подпись в книге. Формулы сравнивают с
# подписью, движку подпись не нужна — код к ней переводит `nonres_book_spec`.
STRATEGY = {ns.STRATEGY_INCOME: "Аренда и выход", ns.STRATEGY_DIRECT: "Прямая продажа (ДКП)",
            ns.STRATEGY_DDU: "Продажа по ДДУ (эскроу)"}
INTEREST_MODES = ("Капитализация в ПФ", "Оплата капиталом")
EXIT = {ns.EXIT_SALE: "Продажа объекта", ns.EXIT_HOLD: "Удержание с оценкой"}
REPAY = {ns.REPAY_ANNUITY: "Аннуитет с баллоном", ns.REPAY_SWEEP: "Из NOI по мере поступления",
         ns.REPAY_BULLET: "Одним платежом при выходе"}
CURVE = {"flat": "Равномерно", "bell": "Колоколом", "front_loaded": "Больше в начале",
         "back_loaded": "Больше в конце"}
SCENARIO = {"low": "Низкий", "base": "Базовый", "high": "Высокий"}
YES, NO = "Да", "Нет"
SPACES = "места"


def _month(value: Any) -> date:
    if isinstance(value, datetime):
        return value.date().replace(day=1)
    if isinstance(value, date):
        return value.replace(day=1)
    return date.fromisoformat(str(value)[:10]).replace(day=1)


def _irr_label(summary: dict[str, Any]) -> str | None:
    """Подпись движка «IRR капитала не показателен» (`summary.equity_irr`)."""
    verdict = summary.get("equity_irr") or {}
    return str(verdict.get("label")) if verdict.get("status") == "thin_equity" else None


def _sum(parts: list[str]) -> str:
    return "+".join(parts) if parts else "0"


# --- «Вводные» ---------------------------------------------------------------

# Проект: ключ → подпись. Значение — столбец B.
PROJECT_INPUTS: tuple[tuple[str, str], ...] = (
    ("start", "Старт проекта"),
    ("ird_months", "Срок ИРД до РнС, мес."),
    ("ird_min", "Минимальный срок ИРД, мес."),
    ("construction_months", "Срок строительства основной части (РнС → РВЭ), мес."),
    ("residual_months", "Остаточные продажи после РВЭ, мес. (для горизонта)"),
    ("vat", "Ставка НДС, %"),
    ("profit_tax", "Налог на прибыль, %"),
    ("loss_limit", "Зачёт убытка прошлых лет — не более доли базы года"),
    ("discount", "Ставка дисконтирования, % годовых"),
    ("revenue_mult", "Сценарий: множитель выручки"),
    ("cost_mult", "Сценарий: множитель затрат"),
    ("marketing", "Маркетинг, % выручки ДКП"),
    ("selling", "Продажи (агенты), % выручки ДКП"),
    ("rate_scenario", "Сценарий ключевой ставки"),
    ("core_above_gns", "ГНС наземная жилого ядра, м² (у нежилого проекта — 0)"),
    ("core_under_gns", "ГНС подземная ядра (паркинг и подземные кладовые), м²"),
    ("ird_th", "ИРД, тыс ₽/м² базы статей проекта"),
    ("design_p_th", "Проект (П), тыс ₽/м² базы статей проекта"),
    ("design_rd_th", "Рабочая документация, тыс ₽/м² базы статей проекта"),
    ("author_supervision_pct", "Авторский надзор, % от П + РД"),
    ("demolition_area", "Площадь сноса, м²"),
    ("demolition_th", "Снос, тыс ₽/м²"),
    ("resettlement_mln", "Расселение, млн ₽"),
    ("preparation_th", "Подготовка территории, тыс ₽/м² базы статей проекта"),
    ("main_above_th", "СМР наземной части ядра, тыс ₽/м²"),
    ("main_under_th", "СМР подземной части, тыс ₽/м² (и гаражи объектов; надбавка за глубину — "
                      "в блоке пятна застройки ниже)"),
    ("utilities_th", "Сети, тыс ₽/м² базы статей проекта"),
    ("commissioning_th", "Ввод в эксплуатацию, тыс ₽/м² базы статей проекта"),
    ("site_maintenance_th", "Содержание площадки, тыс ₽/м² базы статей проекта"),
    ("project_management_pct", "Управление проектом, % базы управления"),
    ("technical_supervision_pct", "Технический заказчик, % СМР"),
    ("gc_fee_pct", "Генподряд, % СМР"),
    ("reserve_pct", "Резерв, % всех статей до резерва"),
    ("purchase_price", "Цена участка, млн ₽"),
    ("land_buyout", "Выкуп ЗУ/ОКС, млн ₽"),
    ("land_rights_gross", "Плата за смену ВРИ, млн ₽"),
    ("land_rights_relief", "Льгота и зачёт по плате за ВРИ, млн ₽"),
    ("parking_price_th", "Цена машино-места проекта по умолчанию, тыс ₽"),
    ("reservation_fee_pct", "Плата за резервирование лимита НКЛ объекта, % лимита"),
    ("limit_fee_pct", "Плата за невыбранный лимит НКЛ, % годовых (0 — нет)"),
    # Проектное финансирование объектов по ДДУ — `simulate_financing`.
    ("bridge_spread", "БРИДЖ: спред к ключевой ставке, п.п."),
    ("bridge_cap_spread", "БРИДЖ: спред капитализации процентов, п.п."),
    ("pf_spread", "ПФ: спред к ключевой ставке на непокрытую эскроу часть, п.п."),
    ("pf_special", "ПФ: специальная ставка на покрытую эскроу часть, %"),
    ("own_funds", "Собственные средства до ПФ (идут раньше БРИДЖа), млн ₽"),
    ("pf_limit_approved", "Одобренный лимит ПФ, млн ₽ (0 — лимит по выборке)"),
    ("interest_mode", "Проценты БРИДЖа при открытии ПФ"),
    # Пятно застройки, этажность и подземные уровни (`site_massing_book`):
    # вводные — числа движка, расчётные строки — формулы, общие с книгой v4.
    *((f"massing_{key}", label) for key, label, _ in massing_book.ROWS),
)
# Расчётные строки блока пятна — формулы, а не жёлтые вводные.
MASSING_FORMULAS = {f"massing_{key}" for key, _, kind in massing_book.ROWS
                    if kind == massing_book.FORMULA}

# Объект: ключ → подпись. Значения — столбцы C, D, … по объектам.
OBJECT_INPUTS: tuple[tuple[str, str], ...] = (
    ("title", "Объект"),
    ("strategy", "Стратегия реализации"),
    ("measure", "Мера объекта (м² / места)"),
    ("volume", "ГНС объекта, м² (наземный паркинг — мест)"),
    ("rate_cost", "Стоимость строительства, тыс ₽/м² ГНС (места — млн ₽/место)"),
    ("turnkey", "Ставка «под ключ»: статьи проекта, генподряд и техзаказчик — в ней"),
    ("under_gns", "Подземный гараж объекта, м²"),
    ("start", "Начало строительства"),
    ("months", "Срок строительства, мес."),
    ("in_tax_pool", "Затраты признаются продажами объекта"),
    ("saleable", "Продаваемая площадь по вводным, м²"),
    ("over_units", "Машино-места на первых этажах, шт."),
    ("over_area", "Площадь места на первом этаже, м²"),
    ("price", "Цена продажи, тыс ₽/м² с НДС (места — млн ₽)"),
    ("sales_start", "Отсчёт цены (старт продаж)"),
    ("growth_pre", "Рост цены до ввода, % в месяц"),
    ("growth_pre_year", "  справочно: рост до ввода, % в год"),
    ("growth_pre_source", "  откуда рост до ввода"),
    ("growth_post", "Рост цены после ввода, % в месяц"),
    ("share_before", "Доля продаж до ввода объекта, % (ДДУ)"),
    ("residual", "Остаточные продажи после ввода, мес. (ДДУ)"),
    ("parking_units", "Машино-места объекта в продаже или аренде, шт."),
    ("parking_under_units", "Мест в гараже подземных, шт. (вес средней цены)"),
    ("parking_under_mln", "Цена подземного места, млн ₽ (0 — цена проекта)"),
    ("parking_over_mln", "Цена места на первом этаже, млн ₽ (0 — цена проекта)"),
    ("rent", "Ставка аренды, тыс ₽/м²/мес. с НДС"),
    ("parking_rent", "Аренда машино-места, тыс ₽/мес. с НДС"),
    ("rent_index", "Индексация аренды, % в год"),
    ("rent_date", "Дата ставки аренды (пусто — на дату ввода)"),
    ("occ_start", "Загрузка на открытии, %"),
    ("occ_stable", "Стабильная загрузка, %"),
    ("leaseup", "Срок заполнения, мес."),
    ("opex", "Операционные расходы, % выручки аренды"),
    ("property_tax", "Налог на имущество, % в год"),
    ("hold_years", "Срок удержания, лет"),
    ("exit_mode", "Выход в конце удержания"),
    ("cap", "Ставка капитализации выхода, %"),
    ("exit_cost", "Затраты на выход, %"),
    ("sale_offset", "Старт прямых продаж от ввода, мес. (минус — до ввода)"),
    ("sale_months", "Срок прямых продаж, мес."),
    ("curve", "Профиль прямых продаж"),
    ("loan_share", "Доля кредита в затратах, %"),
    ("spread", "Спред к ключевой ставке, п.п."),
    ("fee", "Комиссия за выдачу, % выборки"),
    ("repayment", "Погашение кредита"),
    ("term", "Срок кредита от первой выдачи, лет"),
    ("balloon", "Баллон, % долга на ввод"),
    ("dep_years", "Срок амортизации для налога, лет"),
)

P_ROW: dict[str, int] = {key: 3 + i for i, (key, _) in enumerate(PROJECT_INPUTS)}
O_HEAD = 3 + len(PROJECT_INPUTS) + 2
O_ROW: dict[str, int] = {key: O_HEAD + i for i, (key, _) in enumerate(OBJECT_INPUTS)}
TABLES_ROW = O_HEAD + len(OBJECT_INPUTS) + 2


def _p(key: str) -> str:
    """Абсолютная ссылка на вводную проекта."""
    return f"'{INPUTS_SHEET}'!$B${P_ROW[key]}"


def _o(key: str, column: str) -> str:
    return f"'{INPUTS_SHEET}'!${column}${O_ROW[key]}"


def _date_or_none(raw: Any) -> date | None:
    """Дата вводной или пусто: пустая дата ставки — «на дату ввода»."""
    try:
        return date.fromisoformat(str(raw or "").strip()[:10])
    except ValueError:
        return None


def _k_of(date_ref: str) -> str:
    """Номер месяца даты от старта проекта — формулой."""
    return (f"((YEAR({date_ref})-YEAR({_p('start')}))*12"
            f"+MONTH({date_ref})-MONTH({_p('start')}))")


class _Inputs:
    """Где на листе «Вводные» лежат таблицы."""

    def __init__(self) -> None:
        self.plan_first = self.plan_last = 0
        self.rate_first = self.rate_last = 0
        self.steps: list[tuple[str, str]] = []
        self.schedule: dict[tuple[str, str], str] = {}


def _inputs_sheet(book: Workbook, spec: dict[str, Any], objects: list[dict[str, Any]]) -> _Inputs:
    ws = book.active
    ws.title = INPUTS_SHEET
    ws["A1"] = f"Вводные нежилого проекта{(' · ' + spec['name']) if spec.get('name') else ''}"
    ws["A1"].font = BOLD
    ws["A2"] = "Все исходные данные модели — здесь. Жёлтые клетки правятся, остальное считают листы."
    cost = spec.get("cost") or {}
    pf = spec.get("pf") or {}
    values = {
        "start": spec["start"], "ird_months": spec["ird_months"], "ird_min": spec["ird_min"],
        "construction_months": spec["construction_months"], "residual_months": spec["residual_months"],
        "vat": spec["vat_pct"], "profit_tax": spec["profit_tax_pct"], "loss_limit": spec["loss_limit"],
        "discount": spec["discount_pct"], "revenue_mult": spec["revenue_mult"],
        "cost_mult": spec["cost_mult"], "marketing": spec["marketing_pct"],
        "selling": spec["selling_pct"], "rate_scenario": SCENARIO[spec["rate_scenario"]],
        "core_above_gns": spec["core_above_gns"], "core_under_gns": spec["core_under_gns"],
        "main_under_th": spec["main_under_th"], "parking_price_th": spec["parking_price_th"],
        "purchase_price": spec["purchase_price_mln"], "land_buyout": spec["land_buyout_mln"],
        "land_rights_gross": spec["land_rights_gross_mln"],
        "land_rights_relief": spec["land_rights_relief_mln"],
        "reservation_fee_pct": spec.get("reservation_fee_pct", 0.0),
        "limit_fee_pct": spec.get("limit_fee_pct", 0.0),
        "bridge_spread": pf.get("bridge_spread_pp", 0.0),
        "bridge_cap_spread": pf.get("bridge_cap_spread_pp", 0.0),
        "pf_spread": pf.get("pf_spread_pp", 0.0), "pf_special": pf.get("pf_special_pct", 0.0),
        "own_funds": pf.get("own_funds_mln", 0.0),
        "pf_limit_approved": pf.get("approved_limit_mln", 0.0),
        "interest_mode": pf.get("interest_mode", INTEREST_MODES[0]),
        **{key: cost.get(key, 0.0) for key in (
            "ird_th", "design_p_th", "design_rd_th", "author_supervision_pct", "demolition_area",
            "demolition_th", "resettlement_mln", "preparation_th", "main_above_th", "utilities_th",
            "commissioning_th", "site_maintenance_th", "project_management_pct",
            "technical_supervision_pct", "gc_fee_pct", "reserve_pct")},
    }
    massing = spec.get("massing") or {}
    values.update({f"massing_{key}": massing.get(key, 0.0)
                   for key, _, kind in massing_book.ROWS if kind == massing_book.INPUT})
    values.update({f"massing_{key}": "=" + formula for key, formula in
                   massing_book.formulas(lambda key: f"$B${P_ROW['massing_' + key]}").items()})
    # Подземная площадь — формулой от вводных книги: ядро и гаражи объектов.
    if objects:
        garages = "+".join(_o("under_gns", item["column"]) for item in objects)
        values["massing_underground_sqm"] = f"={_p('core_under_gns')}+{garages}"
    formula_keys = MASSING_FORMULAS | ({"massing_underground_sqm"} if objects else set())
    for key, label in PROJECT_INPUTS:
        row = P_ROW[key]
        ws[f"A{row}"] = label
        ws[f"B{row}"] = values[key]
        if key not in formula_keys:
            ws[f"B{row}"].fill = INPUT_FILL
    ws[f"B{P_ROW['start']}"].number_format = "dd.mm.yyyy"
    scenario_list = DataValidation(type="list", formula1='"' + ",".join(SCENARIO.values()) + '"',
                                   allow_blank=False)
    ws.add_data_validation(scenario_list)
    scenario_list.add(f"B{P_ROW['rate_scenario']}")
    mode_list = DataValidation(type="list", formula1='"' + ",".join(INTEREST_MODES) + '"',
                               allow_blank=False)
    ws.add_data_validation(mode_list)
    mode_list.add(f"B{P_ROW['interest_mode']}")

    ws[f"A{O_HEAD - 1}"] = "Объекты"
    ws[f"A{O_HEAD - 1}"].font = BOLD
    lists = {
        "strategy": STRATEGY.values(), "exit_mode": EXIT.values(), "repayment": REPAY.values(),
        "curve": CURVE.values(), "in_tax_pool": (YES, NO), "measure": ("м²", SPACES),
        "turnkey": (YES, NO),
    }
    validations = {}
    for key, options in lists.items():
        dv = DataValidation(type="list", formula1='"' + ",".join(options) + '"', allow_blank=True)
        ws.add_data_validation(dv)
        validations[key] = dv
    for key, label in OBJECT_INPUTS:
        ws[f"A{O_ROW[key]}"] = label
    for item in objects:
        column = item["column"]
        params = item.get("params") or {}
        nonres = bool(item.get("nonres"))
        ddu = bool(item.get("ddu"))
        cells = {
            "title": item["title"],
            # Объект без продажи (переданный городу) — только стройка.
            "strategy": STRATEGY.get(item.get("strategy") or "", "Без продажи (стройка)"),
            "measure": SPACES if item["measure"] == "spaces" else "м²",
            "volume": item["volume"], "rate_cost": item["rate_cost"],
            "under_gns": item["under_gns"], "start": item["start"], "months": item["months"],
            "in_tax_pool": YES if item["in_tax_pool"] else NO,
            "turnkey": YES if item["turnkey"] else NO,
        }
        if nonres:
            cells.update({
                "saleable": item["saleable_sqm"], "over_units": item["over_units"],
                "over_area": item["over_area"], "price": item["price"],
                "sales_start": item["sales_start"], "growth_pre": item["growth_pre_pct"],
                "growth_post": item["growth_post_pct"], "parking_units": item["parking_saleable_units"],
                "parking_under_units": item["parking_under_units"],
                "parking_under_mln": item["parking_under_mln"],
                "parking_over_mln": item["parking_over_mln"],
                "rent": params["rent_th_per_sqm_month"], "parking_rent": params["parking_rent_th_month"],
                "rent_index": params["rent_index_pct"], "occ_start": params["occupancy_start_pct"],
                "rent_date": _date_or_none(params.get("rent_date")),
                "occ_stable": params["occupancy_stable_pct"], "leaseup": params["leaseup_months"],
                "opex": params["opex_pct"], "property_tax": params["property_tax_pct"],
                "hold_years": params["hold_years"], "exit_mode": EXIT[params["exit_mode"]],
                "cap": params["exit_cap_pct"], "exit_cost": params["exit_cost_pct"],
                "sale_offset": params["direct_sale_offset_months"],
                "sale_months": params["direct_sale_months"], "curve": CURVE[params["direct_sale_curve"]],
                "loan_share": params["loan_share_pct"], "spread": params["loan_spread_pp"],
                "fee": params["loan_fee_pct"], "repayment": REPAY[params["debt_repayment"]],
                "term": params["loan_term_years"], "balloon": params["loan_balloon_pct"],
                "dep_years": params["depreciation_years"],
            })
        if ddu:
            cells.update({
                "saleable": item["saleable_sqm"], "over_units": item["over_units"],
                "over_area": item["over_area"], "price": item["price"],
                "sales_start": item["sales_start"], "growth_pre": item["growth_pre_pct"],
                "growth_post": item["growth_post_pct"],
                "share_before": item["share_before_pct"], "residual": item["residual_months"],
                "parking_units": item["parking_saleable_units"],
                "parking_under_units": item["parking_under_units"],
                "parking_under_mln": item["parking_under_mln"],
                "parking_over_mln": item["parking_over_mln"],
            })
        if nonres or ddu:
            default = float(item.get("growth_pre_default_pct") or 0.0)
            same = abs(float(item["growth_pre_pct"]) - default) < 1e-9
            ws[f"{column}{O_ROW['growth_pre_year']}"] = (
                f"=((1+{column}{O_ROW['growth_pre']}/100)^12-1)*100")
            ws[f"{column}{O_ROW['growth_pre_year']}"].number_format = "0.0"
            ws[f"{column}{O_ROW['growth_pre_source']}"] = (
                f"умолчание объекта {default:g} %/мес. (реестр объектов движка); у жилья своё"
                if same else "задано в проекте")
        for key, value in cells.items():
            cell = ws[f"{column}{O_ROW[key]}"]
            cell.value = value
            cell.fill = INPUT_FILL
            if key in ("start", "sales_start", "rent_date"):
                cell.number_format = "dd.mm.yyyy"
            if key in validations and (nonres or key in ("measure", "in_tax_pool", "turnkey")):
                validations[key].add(f"{column}{O_ROW[key]}")
        ws[f"{column}{O_ROW['title']}"].font = BOLD

    info = _Inputs()
    row = TABLES_ROW
    ws[f"A{row}"] = "График оплаты участка и выкупа: месяц от старта проекта и доля цены"
    ws[f"A{row}"].font = BOLD
    ws[f"B{row}"] = "Месяц (k)"
    ws[f"C{row}"] = "Доля"
    plan = list(spec.get("purchase_plan") or []) or [(0, 1.0)]
    info.plan_first = row + 1
    for i, (k, share) in enumerate(plan):
        ws[f"B{row + 1 + i}"] = int(k)
        ws[f"C{row + 1 + i}"] = float(share)
        ws[f"B{row + 1 + i}"].fill = ws[f"C{row + 1 + i}"].fill = INPUT_FILL
    info.plan_last = row + len(plan)
    row = info.plan_last + 2
    ws[f"A{row}"] = "Прогноз ключевой ставки ЦБ, % годовых (действует с даты до следующей строки)"
    ws[f"A{row}"].font = BOLD
    for column, label in zip("BCDE", ("С даты", "Низкий", "Базовый", "Высокий")):
        ws[f"{column}{row}"] = label
    rates = list(spec.get("rate_table") or [])
    info.rate_first = row + 1
    for i, rate in enumerate(rates):
        r = row + 1 + i
        ws[f"B{r}"] = rate["date"]
        ws[f"B{r}"].number_format = "dd.mm.yyyy"
        for column, key in zip("CDE", ("low", "base", "high")):
            ws[f"{column}{r}"] = float(rate[key])
            ws[f"{column}{r}"].fill = INPUT_FILL
    info.rate_last = row + max(1, len(rates))
    # Ступени специальной ставки ПФ по покрытию эскроу — вводная, как в
    # договоре (`pf_special_steps`): порог покрытия и ставка ступени.
    row = info.rate_last + 2
    ws[f"A{row}"] = "ПФ: ступени специальной ставки по покрытию эскроу (пустой порог — ступени нет)"
    ws[f"A{row}"].font = BOLD
    ws[f"B{row}"], ws[f"C{row}"] = "Покрытие от, %", "Ставка, %"
    steps = list((spec.get("pf") or {}).get("pf_special_steps") or [])
    info.steps = []
    for i, (edge, rate) in enumerate(steps):
        r = row + 1 + i
        ws[f"B{r}"], ws[f"C{r}"] = edge * 100, rate * 100
        ws[f"B{r}"].fill = ws[f"C{r}"].fill = INPUT_FILL
        info.steps.append((f"'{INPUTS_SHEET}'!$B${r}", f"'{INPUTS_SHEET}'!$C${r}"))
    # Профиль продаж и лестница цены объекта по ДДУ — вводные: их задают
    # текстом на странице («60% в сделку, 40% через год»), книга получает
    # разобранными по месяцам. Пустой столбец — распределение «доля до ввода
    # + остаток» и ежемесячный рост формулами листа объекта.
    info.schedule = {}
    shaped = [item for item in objects if item.get("ddu")
              and (item.get("sales_weights") is not None or item.get("price_factor") is not None)]
    if shaped:
        row = row + 2 + len(steps)
        ws[f"A{row}"] = "Объекты по ДДУ: профиль продаж (доля месяца) и лестница цены (множитель) — по месяцам"
        ws[f"A{row}"].font = BOLD
        ws[f"B{row}"] = "Месяц (k)"
        months = len(spec.get("months") or [])
        column = 3
        for item in shaped:
            for kind in ("sales_weights", "price_factor"):
                values = item.get(kind)
                if values is None:
                    continue
                letter = get_column_letter(column)
                ws[f"{letter}{row}"] = (item["title"] + (": доля продаж" if kind == "sales_weights"
                                                         else ": множитель цены"))
                for k in range(months):
                    ws[f"B{row + 1 + k}"] = k
                    ws[f"{letter}{row + 1 + k}"] = float(values.get(k, 0.0))
                    ws[f"{letter}{row + 1 + k}"].fill = INPUT_FILL
                info.schedule[(item["key"], kind)] = (
                    f"'{INPUTS_SHEET}'!${letter}${row + 1}:${letter}${row + months}")
                column += 1
    ws.column_dimensions["A"].width = 66
    for i in range(2, 3 + len(objects)):
        ws.column_dimensions[get_column_letter(i)].width = 22
    return info


# --- «Затраты» ---------------------------------------------------------------

S_CURVE = "S-кривая"
EVEN = "равномерно"


def _s_weight(k: str, start: str, months: str) -> str:
    """Вес месяца S-кривой книги ПЛАТО (`build_curve.monthly_weights`)."""
    i = f"({k}-{start}+1)"
    return (f"IF(AND({k}>={start},{k}<{start}+{months}),"
            f"IF({i}/{months}<=0.2,0.6,IF({i}/{months}<=0.8,1.2,0.8)),0)")


def _even_weight(k: str, start: str, months: str) -> str:
    return f"IF(AND({k}>={start},{k}<{start}+{months}),1,0)"


class _Costs:
    def __init__(self) -> None:
        self.cells: dict[str, str] = {}      # имя → абсолютная ссылка
        self.amounts: dict[str, str] = {}    # статья → сумма со сценарием
        self.monthly: dict[str, str] = {}    # статья → буква помесячного столбца
        self.total_col = ""
        self.common_col = ""
        self.share: dict[str, str] = {}      # ключ объекта → доля общих затрат
        self.labels: dict[str, str] = {}


def _costs_sheet(book: Workbook, spec: dict[str, Any], objects: list[dict[str, Any]],
                 inputs: _Inputs, months: int) -> _Costs:
    ws = book.create_sheet(COSTS_SHEET)
    out = _Costs()
    sheet = f"'{COSTS_SHEET}'"
    ws["A1"] = "Затраты: смета «ставка × база» и график по месяцам"
    ws["A1"].font = BOLD
    mult = _p("cost_mult")
    # Ключевые даты — номером месяца от старта проекта.
    ird = f"MAX({_p('ird_min')},INT({_p('ird_months')}))"
    dates = (
        ("ird_eff", "Срок ИРД с минимумом, мес.", f"={ird}"),
        ("permit", "РнС — месяц от старта (k)", f"={ird}"),
        ("build", "Срок строительства основной части, мес.", f"=MAX(1,INT({_p('construction_months')}))"),
        ("rve", "РВЭ — месяц от старта (k)", f"=$B$4+INT({_p('construction_months')})"),
        ("design_window", "Окно проектирования до РнС, мес.", "=MIN(6,$B$3)"),
    )
    for i, (name, label, formula) in enumerate(dates, start=3):
        ws[f"A{i}"] = label
        ws[f"B{i}"] = formula
        out.cells[name] = f"{sheet}!$B${i}"
    head = 10
    for column, label in zip("ABCDEFGHI", ("Статья", "База", "Мера базы", "Ставка",
                                           "Сумма, ₽", "Со сценарием, ₽", "Начало (k)",
                                           "Срок, мес.", "График")):
        ws[f"{column}{head}"] = label
        ws[f"{column}{head}"].font = BOLD
    # База статей проекта (ИРД, П, РД, подготовка, сети, ввод, содержание) —
    # как в движке (`object_rate_is_turnkey`): ГНС ядра плюс ГНС здания и
    # подземный гараж каждого объекта, чья ставка — СМР («под ключ» = «Нет»).
    # У объекта «под ключ» статьи уже в ставке.
    def smr_only(c: str, value: str) -> str:
        return f'IF({_o("turnkey", c)}="{NO}",{value},0)'
    objects_area = "+".join(
        smr_only(item["column"], f"{_o('volume', item['column'])}+{_o('under_gns', item['column'])}")
        for item in objects)
    gns = f"({_p('core_above_gns')}+{_p('core_under_gns')}" + (
        f"+{objects_area})" if objects_area else ")")
    base_measure = "м² базы статей"
    rows: list[tuple[str, str, str, str, str, str, str, str, str]] = []
    # (статья, подпись, база, мера, ставка, сумма, начало, срок, график)

    def article(key, label, base, measure, rate, amount, start="", length="", schedule=""):
        rows.append((key, label, base, measure, rate, amount, start, length, schedule))

    permit, build, rve = out.cells["permit"], out.cells["build"], out.cells["rve"]
    window = out.cells["design_window"]
    # Объекты: здание и гараж — две строки, одна статья движка.
    for item in objects:
        c = item["column"]
        key = item["key"]
        article(f"{key}:building", f"{item['title']}: здание", f"={_o('volume', c)}",
                f'=IF({_o("measure", c)}="{SPACES}","мест","м² ГНС")', f"={_o('rate_cost', c)}",
                f'=B{{r}}*D{{r}}*IF({_o("measure", c)}="{SPACES}",1000000,1000)',
                f"={_k_of(_o('start', c))}", f"=INT({_o('months', c)})", S_CURVE)
        article(f"{key}:garage", f"{item['title']}: подземный гараж", f"={_o('under_gns', c)}",
                "м²", f"={_p('main_under_th')}*{_p('massing_depth_factor')}", "=B{r}*D{r}*1000",
                f"={_k_of(_o('start', c))}", f"=INT({_o('months', c)})", S_CURVE)
    article("ird", "ИРД", f"={gns}", base_measure, f"={_p('ird_th')}", "=B{r}*D{r}*1000",
            "=0", f"={out.cells['ird_eff']}", EVEN)
    article("design_p", "Проект (П)", f"={gns}", base_measure, f"={_p('design_p_th')}",
            "=B{r}*D{r}*1000", f"={permit}-{window}", f"={window}", EVEN)
    article("design_rd", "Рабочая документация", f"={gns}", base_measure, f"={_p('design_rd_th')}",
            "=B{r}*D{r}*1000", f"={permit}-{window}", f"={window}", EVEN)
    article("demolition", "Снос", f"={_p('demolition_area')}", "м² сноса", f"={_p('demolition_th')}",
            "=B{r}*D{r}*1000", f"={permit}-{window}", f"={window}", EVEN)
    article("resettlement", "Расселение", "=1", "млн ₽", f"={_p('resettlement_mln')}",
            "=B{r}*D{r}*1000000", f"={permit}-{window}", f"={window}", EVEN)
    article("preparation", "Подготовка территории", f"={gns}", base_measure,
            f"={_p('preparation_th')}", "=B{r}*D{r}*1000", f"={permit}-{window}", f"={window}", EVEN)
    article("main_above", "СМР наземной части ядра", f"={_p('core_above_gns')}", "м²",
            f"={_p('main_above_th')}", "=B{r}*D{r}*1000", f"={permit}", f"={build}", S_CURVE)
    article("main_under", "СМР подземной части ядра", f"={_p('core_under_gns')}", "м²",
            f"={_p('main_under_th')}*{_p('massing_depth_factor')}", "=B{r}*D{r}*1000", f"={permit}", f"={build}", S_CURVE)
    article("utilities", "Сети", f"={gns}", base_measure, f"={_p('utilities_th')}",
            "=B{r}*D{r}*1000", f"={permit}", f"={build}", S_CURVE)
    article("site_maintenance", "Содержание площадки", f"={gns}", base_measure,
            f"={_p('site_maintenance_th')}", "=B{r}*D{r}*1000", f"={permit}", f"={build}", S_CURVE)
    article("commissioning", "Ввод в эксплуатацию", f"={gns}", base_measure,
            f"={_p('commissioning_th')}", "=B{r}*D{r}*1000", f"={rve}-3", "=3", EVEN)
    engine = spec.get("engine_articles") or {}
    cost_mult = float(spec.get("cost_mult") or 1.0) or 1.0
    engine_labels = {"landscaping": "Благоустройство", "social": "Соцобъекты",
                     "vri_interest": "Проценты рассрочки ВРИ", "vri_security": "Обеспечение рассрочки ВРИ",
                     "land_rights": "Плата за смену ВРИ (рассрочка)"}
    for key in ("landscaping", "social"):
        total = sum((engine.get(key) or {}).values()) / cost_mult
        article(key, engine_labels[key] + " — СЧИТАЕТ ДВИЖОК", "", "", "", f"={total!r}", "", "",
                "движок")
    article("author_supervision", "Авторский надзор", "=E{design_p}+E{design_rd}", "П + РД, ₽",
            f"={_p('author_supervision_pct')}/100", "=B{r}*D{r}", f"={permit}", f"={build}", S_CURVE)
    # База генподряда и техзаказчика — СМР ядра и соцобъектов, как в движке,
    # плюс здание и гараж объекта, чья ставка — СМР. Ставка «под ключ» несёт
    # генподряд и техзаказчика сама, процент сверху был бы двойным счётом.
    object_works = "".join(
        "+" + smr_only(item["column"], f"E{{{item['key']}__building}}+E{{{item['key']}__garage}}")
        for item in objects)
    works = "=E{main_above}+E{main_under}+E{social}" + object_works
    article("technical_supervision", "Технический заказчик", works, "СМР, ₽",
            f"={_p('technical_supervision_pct')}/100", "=B{r}*D{r}", f"={permit}", f"={build}", S_CURVE)
    article("project_management", "Управление проектом",
            "=E{ird}+E{design_p}+E{design_rd}+E{author_supervision}+E{preparation}+E{main_above}"
            "+E{main_under}+E{utilities}+E{landscaping}+E{site_maintenance}" + object_works,
            "база управления, ₽", f"={_p('project_management_pct')}/100", "=B{r}*D{r}",
            "=0", f"={rve}", "профиль расходов")
    article("gc_fee", "Генподряд", works, "СМР, ₽", f"={_p('gc_fee_pct')}/100", "=B{r}*D{r}",
            f"={permit}", f"={build}", S_CURVE)
    if spec.get("vri_enabled"):
        total = sum((engine.get("land_rights") or {}).values()) / cost_mult
        article("land_rights", engine_labels["land_rights"] + " — СЧИТАЕТ ДВИЖОК", "", "", "",
                f"={total!r}", "", "", "движок")
    else:
        article("land_rights", "Плата за смену ВРИ к оплате", f"={_p('land_rights_gross')}", "млн ₽",
                f"=1-IF({_p('land_rights_gross')}>0,MIN(1,{_p('land_rights_relief')}/{_p('land_rights_gross')}),0)",
                "=B{r}*D{r}*1000000", f"={permit}", "=1", "разово в РнС")
    # Плата за смену ВРИ — известная сумма, резерв на неё не начисляется
    # (список движка `RESERVE_EXCLUDED_ARTICLES` приходит в спецификации).
    reserve_excluded = set(spec["reserve_excluded"])
    reserve_parts = [key for key, *_ in rows if key not in reserve_excluded]
    article("reserve", "Резерв", "=" + _sum([f"E{{{key.replace(':', '__')}}}" for key in reserve_parts]),
            "все статьи до резерва, ₽", f"={_p('reserve_pct')}/100", "=B{r}*D{r}",
            f"={permit}", f"={build}", S_CURVE)
    article("purchase", "Участок и выкуп ЗУ/ОКС", f"={_p('purchase_price')}+{_p('land_buyout')}",
            "млн ₽", "=1", "=B{r}*D{r}*1000000", "", "", "по графику оплаты")
    for key in ("vri_interest", "vri_security"):
        total = sum((engine.get(key) or {}).values()) / cost_mult
        if total:
            article(key, engine_labels[key] + " — СЧИТАЕТ ДВИЖОК", "", "", "", f"={total!r}", "", "",
                    "движок")
    row_of = {key: head + 1 + i for i, (key, *_) in enumerate(rows)}
    for key, label, base, measure, rate, amount, start, length, schedule in rows:
        r = row_of[key]
        fill = {name.replace(":", "__"): str(row_of[name]) for name in row_of}
        fill["r"] = str(r)

        def put(column, text):
            if text == "":
                return
            value = text.format(**fill) if isinstance(text, str) else text
            ws[f"{column}{r}"] = value

        ws[f"A{r}"] = label
        put("B", base)
        put("C", measure)
        put("D", rate)
        put("E", amount)
        ws[f"F{r}"] = f"=E{r}*{mult}"
        put("G", start)
        put("H", length)
        ws[f"I{r}"] = schedule
        if schedule == "движок":
            ws[f"E{r}"].fill = ENGINE_FILL
        ws[f"E{r}"].number_format = ws[f"F{r}"].number_format = MONEY
        out.labels[key] = label
    total_row = head + len(rows) + 1
    ws[f"A{total_row}"] = "Итого затраты проекта (CAPEX)"
    ws[f"A{total_row}"].font = BOLD
    ws[f"F{total_row}"] = f"=SUM(F{head + 1}:F{total_row - 1})"
    ws[f"F{total_row}"].number_format = MONEY
    out.cells["capex_total"] = f"{sheet}!$F${total_row}"
    for key in row_of:
        out.amounts[key] = f"{sheet}!$F${row_of[key]}"

    # Помесячный график: столбец на статью движка (объект — здание + гараж).
    first = FIRST_ROW
    last = FIRST_ROW + months - 1
    columns: list[tuple[str, str]] = [("k", "k"), ("month", "Месяц")]
    for item in objects:
        columns.append((f"w:{item['key']}", f"Вес S-кривой: {item['title']}"))
    columns += [("w:main", "Вес S-кривой основной части"), ("w:ird", "Вес: ИРД"),
                ("w:design", "Вес: окно проектирования"), ("w:rve3", "Вес: 3 мес. до РВЭ")]
    for item in objects:
        columns.append((item["key"], item["title"]))
    articles_monthly = [key for key, *_ in rows if ":" not in key]
    columns += [(key, out.labels[key]) for key in articles_monthly]
    columns += [("total", "Итого CAPEX"), ("nonres_obj", "Объекты вне ДДУ"),
                ("ddu_obj", "Объекты по ДДУ"), ("vri_equity", "ВРИ за счёт капитала (движок)"),
                ("common", "Общие затраты на кредит объектов")]
    letter = {key: get_column_letter(i + 1) for i, (key, _) in enumerate(columns)}
    for key, label in columns:
        ws[f"{letter[key]}{first - 1}"] = label
        ws[f"{letter[key]}{first - 1}"].font = BOLD
    for key in engine:
        if key in letter:
            ws[f"{letter[key]}{first - 2}"] = "СЧИТАЕТ ДВИЖОК"
    plan = (f"'{INPUTS_SHEET}'!$B${inputs.plan_first}:$B${inputs.plan_last}",
            f"'{INPUTS_SHEET}'!$C${inputs.plan_first}:$C${inputs.plan_last}")
    vri_equity = spec.get("vri_equity") or {}

    def col(key):
        return f"${letter[key]}${first}:${letter[key]}${last}"

    def amount_of(key):
        return f"$F${row_of[key]}"

    def length_of(key):
        return f"$H${row_of[key]}"

    def start_of(key):
        return f"$G${row_of[key]}"

    profile = [key for key in ("ird", "design_p", "design_rd", "author_supervision", "preparation",
                               "main_above", "main_under", "utilities", "landscaping",
                               "site_maintenance") if key in letter]
    for i in range(months):
        r = first + i
        k = f"$A{r}"
        ws[f"A{r}"] = i
        ws[f"B{r}"] = f"=EDATE({_p('start')},A{r})"
        ws[f"B{r}"].number_format = MONTH
        for item in objects:
            key = f"{item['key']}:building"
            ws[f"{letter['w:' + item['key']]}{r}"] = "=" + _s_weight(k, start_of(key), length_of(key))
            ws[f"{letter[item['key']]}{r}"] = (
                f"=({amount_of(key)}+{amount_of(item['key'] + ':garage')})"
                f"*{letter['w:' + item['key']]}{r}/SUM({col('w:' + item['key'])})")
        ws[f"{letter['w:main']}{r}"] = "=" + _s_weight(k, permit, build)
        ws[f"{letter['w:ird']}{r}"] = "=" + _even_weight(k, "0", out.cells["ird_eff"])
        ws[f"{letter['w:design']}{r}"] = "=" + _even_weight(k, f"({permit}-{window})", window)
        ws[f"{letter['w:rve3']}{r}"] = "=" + _even_weight(k, f"({rve}-3)", "3")
        main = f"{letter['w:main']}{r}/SUM({col('w:main')})"
        for key in articles_monthly:
            cell = f"{letter[key]}{r}"
            if key in engine:
                ws[cell] = float((engine.get(key) or {}).get(i, 0.0))
                ws[cell].fill = ENGINE_FILL
            elif key == "ird":
                ws[cell] = f"={amount_of(key)}/{length_of(key)}*{letter['w:ird']}{r}"
            elif key in ("design_p", "design_rd", "demolition", "resettlement", "preparation"):
                ws[cell] = f"={amount_of(key)}/{window}*{letter['w:design']}{r}"
            elif key == "commissioning":
                ws[cell] = f"={amount_of(key)}/3*{letter['w:rve3']}{r}"
            elif key == "land_rights":
                ws[cell] = f"=IF({k}={permit},{amount_of(key)},0)"
            elif key == "purchase":
                ws[cell] = f"={amount_of(key)}*SUMIF({plan[0]},{k},{plan[1]})"
            elif key == "project_management":
                base = "+".join(f"{letter[p]}{r}" for p in profile)
                total = "+".join(f"SUM({col(p)})" for p in profile)
                ws[cell] = (f"=IF(({total})>0,{amount_of(key)}*({base})/({total}),"
                            f"{amount_of(key)}/MAX(1,{rve})*{_even_weight(k, '0', f'MAX(1,{rve})')})")
            elif key in ("landscaping", "social", "vri_interest", "vri_security"):
                ws[cell] = 0
            else:
                ws[cell] = f"={amount_of(key)}*{main}"
            ws[cell].number_format = MONEY
        spend = [letter[item["key"]] for item in objects] + [letter[key] for key in articles_monthly]
        ws[f"{letter['total']}{r}"] = "=" + "+".join(f"{c}{r}" for c in spend)
        ws[f"{letter['nonres_obj']}{r}"] = "=" + _sum(
            [f"{letter[item['key']]}{r}" for item in objects if item.get("nonres")])
        ws[f"{letter['ddu_obj']}{r}"] = "=" + _sum(
            [f"{letter[item['key']]}{r}" for item in objects if not item.get("nonres")])
        ws[f"{letter['vri_equity']}{r}"] = float(vri_equity.get(i, 0.0))
        ws[f"{letter['vri_equity']}{r}"].fill = ENGINE_FILL
        # Нежилой проект: банк кредитует долю ВСЕЙ стоимости — общие затраты
        # уходят в кредит объектов вне ДДУ долей их стройки.
        ws[f"{letter['common']}{r}"] = (
            f"=MAX(0,MAX(0,{letter['total']}{r}-{letter['vri_equity']}{r}-{letter['nonres_obj']}{r})"
            f"-{letter['ddu_obj']}{r})")
        for key in ("total", "nonres_obj", "ddu_obj", "vri_equity", "common"):
            ws[f"{letter[key]}{r}"].number_format = MONEY
    out.monthly = {key: letter[key] for key, _ in columns}
    out.total_col, out.common_col = letter["total"], letter["common"]
    # Доля общих затрат объекта — его стройка в стройке всех объектов.
    share_row = total_row + 2
    ws[f"A{share_row}"] = "Доля объекта в общих затратах (его стройка / стройка всех объектов)"
    ws[f"A{share_row}"].font = BOLD
    all_objects = "+".join(f"SUM({col(item['key'])})" for item in objects) or "0"
    for i, item in enumerate(objects, start=1):
        r = share_row + i
        ws[f"A{r}"] = item["title"]
        if item.get("nonres"):
            ws[f"B{r}"] = f"=IF(({all_objects})>0,SUM({col(item['key'])})/({all_objects}),0)"
            ws[f"B{r}"].number_format = "0.0000%"
            out.share[item["key"]] = f"{sheet}!$B${r}"
    out.cells["months_first"] = str(first)
    out.cells["plan_last_k"] = f"MAX({plan[0]})"
    ws.column_dimensions["A"].width = 44
    for i in range(2, len(columns) + 1):
        ws.column_dimensions[get_column_letter(i)].width = 16
    return out


# --- «Объект N» --------------------------------------------------------------

# Вводные объекта на его листе: ключ → (строка, подпись). Клетки — ссылки на
# «Вводные» с теми же ограничениями, что у `object_flows`.
INPUTS: dict[str, tuple[int, str]] = {key: (2 + i, label) for i, (key, label) in enumerate((
    ("strategy", "Стратегия"),
    ("area", "Арендопригодная / продаваемая площадь, м² (места — шт.)"),
    ("spaces", "Машино-места объекта (без гостевых), шт."),
    ("rent", "Ставка аренды, ₽/м²/мес. с НДС"),
    ("parking_rent", "Аренда машино-места, ₽/мес. с НДС"),
    ("rent_index", "Индексация аренды, доля в год"),
    ("rent_from", "Месяц ставки аренды (k); без даты — месяц ввода"),
    ("occ_start", "Загрузка на открытии, доля"),
    ("occ_stable", "Стабильная загрузка, доля"),
    ("leaseup", "Срок заполнения, мес."),
    ("opex", "Операционные расходы, доля выручки"),
    ("rev_mult", "Множитель выручки сценария"),
    ("property_tax", "Налог на имущество, доля в год"),
    ("vat", "Ставка НДС"),
    ("loan_share", "Доля кредита в затратах"),
    ("spread", "Спред кредита к ключевой ставке"),
    ("fee", "Комиссия за выдачу, доля выборки"),
    ("reservation", "Плата за резервирование лимита НКЛ, доля лимита"),
    ("commitment", "Плата за невыбранный лимит НКЛ, доля в год"),
    ("permit", "РнС — граница двух НКЛ (k)"),
    ("term", "Срок кредита от первой выдачи, мес."),
    ("balloon_share", "Баллон, доля долга на ввод"),
    ("hold", "Срок удержания, мес."),
    ("cap", "Ставка капитализации выхода"),
    ("exit_cost", "Затраты на выход, доля"),
    ("exit_mode", "Выход"),
    ("repayment", "Погашение кредита"),
    ("comm", "Ввод объекта — месяц от старта проекта (k)"),
    ("horizon", "Конец горизонта объекта (k)"),
    ("price", "Цена продажи, ₽/м² с НДС (места — ₽/место)"),
    ("parking_price", "Цена машино-места, ₽ (средняя по гаражу)"),
    ("growth_pre", "Рост цены до ввода, доля в месяц"),
    ("growth_post", "Рост цены после ввода, доля в месяц"),
    ("price_start", "Отсчёт цены (k)"),
    ("selling", "Маркетинг и продажи, доля выручки"),
    ("cost_mult", "Множитель затрат сценария"),
    ("sale_start", "Старт прямых продаж (k)"),
    ("sale_months", "Срок прямых продаж, мес."),
    ("curve", "Профиль продаж"),
    ("dep_years", "Срок амортизации для налога, лет"),
    ("share", "Доля объекта в общих затратах"),
    ("tax_rate", "Ставка налога на прибыль"),
    ("loss_limit", "Зачёт убытка — не более доли базы года"),
    ("discount", "Ставка дисконтирования, годовая"),
))}
_D0 = 2 + len(INPUTS) + 1
DERIVED: dict[str, tuple[int, str]] = {key: (_D0 + i, label) for i, (key, label) in enumerate((
    ("capex_total", "CAPEX объекта, ₽ с НДС"),
    ("basis", "Стоимость объекта без НДС (база налога на имущество и амортизации)"),
    ("first_draw", "Первая выдача кредита (k)"),
    ("maturity", "Срок погашения (k)"),
    ("comm_balance", "Долг на ввод до погашения, ₽"),
    ("balloon", "Баллон, ₽"),
    ("forward_noi", "NOI следующих 12 мес. после срока удержания, ₽"),
    ("exit_value", "Стоимость выхода = NOI / ставка, ₽"),
    ("stabilized_noi", "Стабилизированный NOI, год, ₽"),
    ("curve_sum", "Сумма весов профиля продаж"),
    ("peak", "Пик долга, ₽"),
    ("common_total", "Общие затраты проекта на объекте, ₽"),
    ("monthly_rate", "Ставка дисконтирования, месячная"),
    ("loan_limit", "Лимит НКЛ объекта = вся плановая выборка, ₽"),
    ("limit_land", "НКЛ на участок и проект: выборка до РнС, ₽"),
    ("limit_build", "Строительная НКЛ: выборка с РнС, ₽"),
    ("first_land", "Открытие НКЛ на участок (k)"),
    ("first_build", "Открытие строительной НКЛ (k)"),
    ("reservation_fee", "Плата за резервирование лимитов, ₽ (при открытии каждой НКЛ)"),
))}
assert _D0 + len(DERIVED) < FIRST_ROW - 2

_COLUMN_LIST: tuple[tuple[str, str], ...] = (
    ("k", "k"), ("month", "Месяц"),
    ("capex", "CAPEX объекта"), ("common", "Общие затраты на объекте"),
    ("key_rate", "Ключевая ставка"),
    ("open", "Месяц эксплуатации"), ("occ", "Загрузка"), ("growth", "Индексация"),
    ("rent", "Арендная выручка"), ("opex", "OPEX"), ("ptax", "Налог на имущество"),
    ("weight", "Доля продаж месяца"), ("factor", "Рост цены"),
    ("sale", "Выручка прямых продаж"), ("selling", "Расходы на продажу"),
    ("exit", "Выход — продажа"), ("exit_cost", "Затраты на выход"),
    ("residual", "Удержание — оценка"), ("op_rev", "Выручка эксплуатации (без срока)"),
    ("fwd", "NOI после срока удержания"), ("stab", "NOI стабилизированного года"),
    ("vat_charged", "НДС начисленный"), ("vat_pre", "НДС к зачёту до расчёта"),
    ("vat_paid", "НДС к уплате (− возмещение)"), ("vat_credit", "НДС к зачёту на конец"),
    ("rate", "Ставка кредита"), ("bal_open", "Долг на начало"),
    ("interest", "Проценты"), ("int_cap", "Проценты капитализированные"),
    ("int_paid", "Проценты уплаченные"), ("draw", "Выборка"), ("fee", "Комиссия"),
    ("first_k", "Служебная: выборка"), ("first_build_k", "Служебная: выборка с РнС"),
    ("commitment_fee", "Плата за невыбранный лимит"), ("bal_after", "Долг после выборки"),
    ("bal_build", "Служебная: долг стройки до погашений"), ("cash", "Деньги до погашения"),
    ("repay", "Погашение"), ("annuity", "Тело по аннуитету"),
    ("amort", "Плановое тело (для DSCR)"), ("bal_close", "Долг на конец"),
    ("to_equity", "Деньги объекта собственнику (без CAPEX)"),
    ("year", "Год эксплуатации"), ("noi", "NOI"),
    ("book_open", "Остаточная стоимость на начало"), ("recognized", "Признано в расходах"),
    ("book_close", "Остаточная стоимость на конец"), ("realized", "Выручка ДКП к признанию"),
    ("margin", "Налоговая маржа объекта"),
    ("draw_obj", "Выборка на стройку объекта"),
    ("part_balloon", "Погашение: баллон в срок"), ("part_exit", "Погашение: остаток при выходе"),
    ("part_cash", "Погашение: из выручки / NOI"),
    ("cal_year", "Календарный год"),
    ("common_rec", "Общие затраты — признано"),
    ("obj_margin", "База объекта до процентов"), ("obj_fin", "Проценты и комиссия — вычет"),
    ("tx_net", "База месяца"), ("tx_gate", "До первого облагаемого месяца"),
    ("tx_year", "Результат года нарастающим"), ("tx_prior", "Убыток прошлых лет"),
    ("tx_used", "Зачтено убытка"), ("tx_base", "База года"),
    ("tx_paid", "Налог года уплачен ранее"), ("tx_tax", "Налог на прибыль объекта"),
    ("eq_flow", "Поток капитала объекта после налога"),
    ("eq_cash", "То же без оценки удержания"), ("eq_cum", "Накопленный (деньги)"),
    ("neg_k", "Служебная: накопленный < 0"), ("first_neg", "Служебная: первое вложение"),
    ("df", "Дисконт-множитель"), ("eq_disc", "Дисконтированный поток капитала"),
)
COLUMNS: dict[str, tuple[str, str]] = {
    key: (get_column_letter(i + 1), label) for i, (key, label) in enumerate(_COLUMN_LIST)}


def _c(name: str) -> str:
    return COLUMNS[name][0]


def _in(name: str) -> str:
    row = (INPUTS.get(name) or DERIVED[name])[0]
    return f"$B${row}"


def _is(name: str, labels: dict[str, str], code: str) -> str:
    return f'{_in(name)}="{labels[code]}"'


def _tax_formulas(v: dict[str, str], p: dict[str, str], r: int, first: int,
                  gate: str, rate: str, limit: str, net: str) -> dict[str, str]:
    """`_profit_tax_schedule` строкой Excel: год закрывается сменой
    календарного года, убыток закрытых лет зачитывается не более `limit`
    базы года, всё признанное до первого облагаемого месяца приходит в его год.
    """
    if r == first:
        return {
            "tx_gate": f"={gate}",
            "tx_year": f"=IF({v['tx_gate']}=1,0,{v['tx_net']})",
            "tx_prior": "=0",
            "tx_used": f"=IF(AND({v['tx_year']}>0,{v['tx_prior']}>0),MIN({v['tx_prior']},{v['tx_year']}*{limit}),0)",
            "tx_base": f"=MAX({v['tx_year']}-{v['tx_used']},0)",
            "tx_paid": "=0",
            "tx_tax": f"=IF({v['tx_gate']}=1,0,MAX({v['tx_base']}*{rate}-{v['tx_paid']},0))",
        }
    same_year = f"AND({p['tx_gate']}=0,{v['cal_year']}={p['cal_year']})"
    return {
        "tx_gate": f"={gate}",
        "tx_year": (f"=IF({v['tx_gate']}=1,0,IF({same_year},{p['tx_year']},0)+{v['tx_net']}"
                    f"+IF({p['tx_gate']}=1,SUM({net}${first}:{net}{r - 1}),0))"),
        "tx_prior": (f"=IF(AND({p['tx_gate']}=0,{v['cal_year']}<>{p['cal_year']}),"
                     f"{p['tx_prior']}+IF({p['tx_year']}<0,-{p['tx_year']},-{p['tx_used']}),{p['tx_prior']})"),
        "tx_used": f"=IF(AND({v['tx_year']}>0,{v['tx_prior']}>0),MIN({v['tx_prior']},{v['tx_year']}*{limit}),0)",
        "tx_base": f"=MAX({v['tx_year']}-{v['tx_used']},0)",
        "tx_paid": f"=IF({same_year},{p['tx_paid']}+{p['tx_tax']},0)",
        "tx_tax": f"=IF({v['tx_gate']}=1,0,MAX({v['tx_base']}*{rate}-{v['tx_paid']},0))",
    }


def _project_tax_formulas(v: dict[str, str], p: dict[str, str], r: int, first: int,
                          gate: str, rate: str, limit: str, margin: str) -> dict[str, str]:
    """Налог проекта спецзастройщика строкой Excel (`profit_tax_schedule` с
    `defer_financing=False`): до РВЭ откладывается только маржа — она приходит
    в месяц РВЭ, — а проценты и комиссии остаются расходом месяца начисления и
    дают убыток своих лет, который гасит прибыль не больше чем `limit` базы
    года (владелец, 16.10.2026)."""
    if r == first:
        return {
            "tx_gate": f"={gate}",
            "tx_year": f"=IF({v['tx_gate']}=1,-{v['financing']},{v['tx_net']})",
            "tx_prior": "=0",
            "tx_used": f"=IF(AND({v['tx_year']}>0,{v['tx_prior']}>0),MIN({v['tx_prior']},{v['tx_year']}*{limit}),0)",
            "tx_base": f"=MAX({v['tx_year']}-{v['tx_used']},0)",
            "tx_paid": "=0",
            "tx_tax": f"=IF({v['tx_gate']}=1,0,MAX({v['tx_base']}*{rate}-{v['tx_paid']},0))",
        }
    same_year = f"{v['cal_year']}={p['cal_year']}"
    return {
        "tx_gate": f"={gate}",
        "tx_year": (f"=IF({same_year},{p['tx_year']},0)+IF({v['tx_gate']}=1,-{v['financing']},"
                    f"{v['tx_net']}+IF({p['tx_gate']}=1,SUM({margin}${first}:{margin}{r - 1}),0))"),
        "tx_prior": (f"=IF({v['cal_year']}<>{p['cal_year']},"
                     f"{p['tx_prior']}+IF({p['tx_year']}<0,-{p['tx_year']},-{p['tx_used']}),{p['tx_prior']})"),
        "tx_used": f"=IF(AND({v['tx_year']}>0,{v['tx_prior']}>0),MIN({v['tx_prior']},{v['tx_year']}*{limit}),0)",
        "tx_base": f"=MAX({v['tx_year']}-{v['tx_used']},0)",
        "tx_paid": f"=IF({same_year},{p['tx_paid']}+{p['tx_tax']},0)",
        "tx_tax": f"=IF({v['tx_gate']}=1,0,MAX({v['tx_base']}*{rate}-{v['tx_paid']},0))",
    }


def _tax_payment_formula(v: dict[str, str], r: int, first: int, tax: str) -> str:
    """Уплата налога месяца: авансы по итогам квартала (апрель, июль, октябрь)
    и доплата года в марте следующего — `profit_tax_payments` строкой Excel."""
    def span(start: int, stop: int) -> str:
        cells = [f"{tax}{row}" for row in range(r - start, r - stop + 1) if row >= first]
        return "+".join(cells) or "0"
    month = f"MONTH({v['month']})"
    return (f"=IF({month}=3,{span(5, 3)},IF(OR({month}=4,{month}=7,{month}=10),"
            f"{span(3, 1)},0))")


def _row_formulas(r: int, refs: dict[str, str]) -> dict[str, str]:
    """Формулы строки r — построчный пересказ `object_flows` и `object_result`
    языком Excel."""
    A = f"{_c('k')}{r}"
    prev = r - 1
    first = FIRST_ROW
    v = {name: f"{_c(name)}{r}" for name in COLUMNS}
    p = {name: f"{_c(name)}{prev}" for name in COLUMNS}
    vs = f"{_in('vat')}/(1+{_in('vat')})"
    income = _is("strategy", STRATEGY, ns.STRATEGY_INCOME)
    direct = _is("strategy", STRATEGY, ns.STRATEGY_DIRECT)
    opex_share = f"MIN(0.95,{_in('opex')})"
    base_price = f"({_in('area')}*{_in('price')}+{_in('spaces')}*{_in('parking_price')})"
    j = f"({A}-{_in('sale_start')}+1)"
    n = _in("sale_months")
    prev_weights = f"SUM({_c('weight')}${first}:{_c('weight')}{prev})" if r > first else "0"
    r12 = f"{v['rate']}/12"
    left = f"({_in('maturity')}-{A})"
    annuity_mode = _is("repayment", REPAY, ns.REPAY_ANNUITY)
    sweep_mode = _is("repayment", REPAY, ns.REPAY_SWEEP)
    out = {
        "capex": f"={refs['capex']}",
        "common": f"={refs['common']}*{_in('share')}",
        "key_rate": f"={refs['key_rate']}",
        "open": f"=IF({A}>={_in('comm')},{A}-{_in('comm')}+1,0)",
        "occ": (f"=IF({v['open']}=0,0,MIN({_in('occ_start')},{_in('occ_stable')})"
                f"+({_in('occ_stable')}-MIN({_in('occ_start')},{_in('occ_stable')}))"
                f"*IF({_in('leaseup')}=1,1,MAX(0,MIN(1,({v['open']}-1)/({_in('leaseup')}-1)))))"),
        "growth": (f"=(1+MAX(-0.95,{_in('rent_index')}))^((MAX(0,{v['open']}-1)"
                   f"+MAX(0,{_in('comm')}-{_in('rent_from')}))/12)"),
        "op_rev": (f"=({_in('area')}*{_in('rent')}+{_in('spaces')}*{_in('parking_rent')})"
                   f"*{v['growth']}*{v['occ']}*{_in('rev_mult')}"),
        "rent": f"=IF(AND({income},{v['open']}>=1,{v['open']}<={_in('hold')}),{v['op_rev']},0)",
        "opex": f"={v['rent']}*{opex_share}",
        "weight": (f"=IF(AND({direct},{j}>=1,{j}<={n}),IF({_is('curve', CURVE, 'bell')},{j}*({n}+1-{j}),"
                   f"IF({_is('curve', CURVE, 'front_loaded')},{n}+1-{j},"
                   f"IF({_is('curve', CURVE, 'back_loaded')},{j},1)))/{_in('curve_sum')},0)"),
        "ptax": (f"=IF({income},IF(AND({v['open']}>=1,{v['open']}<={_in('hold')}),"
                 f"{_in('basis')}*{_in('property_tax')}/12,0),"
                 f"IF(AND({A}>={_in('comm')},{A}<={_in('horizon')}),"
                 f"{_in('basis')}*{_in('property_tax')}/12*MAX(0,1-{prev_weights}),0))"),
        "factor": (f"=(1+{_in('growth_pre')})^MAX(0,MIN({A}-{_in('price_start')},"
                   f"{_in('comm')}-{_in('price_start')}))"
                   f"*(1+{_in('growth_post')})^MAX(0,{A}-MAX({_in('price_start')},{_in('comm')}))"),
        "sale": f"={base_price}*{v['factor']}*{v['weight']}*{_in('rev_mult')}",
        "selling": f"={base_price}*{v['factor']}*{v['weight']}*{_in('selling')}*{_in('cost_mult')}",
        "exit": (f'=IF(AND({income},{_is("exit_mode", EXIT, ns.EXIT_SALE)},{A}={_in("horizon")}),'
                 f'{_in("exit_value")},0)'),
        "exit_cost": f"={v['exit']}*{_in('exit_cost')}",
        "residual": (f'=IF(AND({income},{_is("exit_mode", EXIT, ns.EXIT_HOLD)},{A}={_in("horizon")}),'
                     f'{_in("exit_value")},0)'),
        "fwd": (f"=IF(AND({v['open']}>={_in('hold')}+1,{v['open']}<={_in('hold')}+12),"
                f"{v['op_rev']}*(1-{opex_share}),0)"),
        "stab": (f"=IF(AND({v['open']}>={_in('leaseup')},{v['open']}<={_in('leaseup')}+11),"
                 f"{v['op_rev']}*(1-{opex_share}),0)"),
        "vat_charged": f"=IF({A}>{_in('horizon')},0,({v['sale']}+{v['rent']}+{v['exit']})*{vs})",
        "vat_pre": (f"=IF({A}>{_in('horizon')},0,{p['vat_credit'] if r > first else '0'}"
                    f"+{v['capex']}*{vs}-{v['vat_charged']})"),
        "vat_paid": (f"=IF(AND({A}={_in('comm')},{v['vat_pre']}>0),-{v['vat_pre']},"
                     f"IF({v['vat_pre']}<0,-{v['vat_pre']},0))"),
        "vat_credit": f"={v['vat_pre']}+{v['vat_paid']}",
        "rate": f"={v['key_rate']}+{_in('spread')}",
        "bal_open": f"={p['bal_close'] if r > first else '0'}",
        "interest": f"={v['bal_open']}*{v['rate']}/12",
        "int_cap": f"=IF({A}<{_in('comm')},{v['interest']},0)",
        "int_paid": f"=IF({A}>={_in('comm')},{v['interest']},0)",
        "draw": f"=IF({A}<={_in('comm')},({v['capex']}+{v['common']})*{_in('loan_share')},0)",
        "fee": (f"={v['draw']}*{_in('fee')}"
                f"+IF({A}={_in('first_land')},{_in('limit_land')}*{_in('reservation')},0)"
                f"+IF({A}={_in('first_build')},{_in('limit_build')}*{_in('reservation')},0)"
                f"+{v['commitment_fee']}"),
        "first_build_k": f"=IF(AND({v['draw']}>0,{A}>={_in('permit')}),{A},1000000)",
        # Невыбранный лимит — с первой выдачи до ввода, как у ПФ до РВЭ.
        # Две НКЛ: до РнС — участок и проект, с РнС — стройка. Каждая платит
        # за невыбранный остаток своего лимита с открытия; после её последней
        # выдачи остаток нулевой.
        "commitment_fee": (f"=(IF(AND({A}>={_in('first_land')},{A}<{_in('permit')}),"
                           f"MAX(0,{_in('limit_land')}-SUM({_c('draw')}${first}:{_c('draw')}{r})),0)"
                           f"+IF({A}>={_in('first_build')},MAX(0,{_in('limit_build')}"
                           f"-(SUM({_c('draw')}${first}:{_c('draw')}{r})-{_in('limit_land')})),0))"
                           f"*{_in('commitment')}/12"),
        "first_k": f"=IF({v['draw']}>0,{A},1000000)",
        "bal_after": f"={v['bal_open']}+{v['int_cap']}+{v['draw']}",
        # Долг стройки без погашений: до ввода у доходного объекта погашений
        # нет, и баллон берётся отсюда — без ссылки на строки после ввода,
        # которые сами читают баллон (Excel видел бы цикл).
        "bal_build": (f"=IF({A}<={_in('comm')},{p['bal_build'] if r > first else '0'}"
                      f"+IF({A}<{_in('comm')},{p['bal_build'] if r > first else '0'}*{v['rate']}/12,0)"
                      f"+{v['draw']},{p['bal_build'] if r > first else '0'})"),
        "cash": (f"={v['sale']}+{v['rent']}+{v['exit']}-{v['selling']}-{v['opex']}-{v['ptax']}"
                 f"-{v['exit_cost']}-{v['vat_paid']}-{v['int_paid']}"),
        "annuity": (f"=IF(AND({A}>{_in('comm')},{A}<{_in('maturity')}),"
                    f"IF({v['rate']}>0,MIN({v['bal_after']},MAX(0,({v['bal_after']}-{_in('balloon')}"
                    f"/(1+{r12})^{left})*({r12})/(1-(1+{r12})^(-{left}))-{v['interest']})),"
                    f"MIN({v['bal_after']},MAX(0,({v['bal_after']}-{_in('balloon')})/{left}))),0)"),
        "repay": (f"=IF({v['bal_after']}<=0,0,IF(OR({A}>={_in('comm')},{direct}),"
                  f"IF({A}={_in('horizon')},{v['bal_after']},"
                  f"IF(OR({sweep_mode},{direct}),MIN({v['bal_after']},MAX(0,{v['cash']})),"
                  f"IF({annuity_mode},IF({A}={_in('comm')},"
                  f"IF({_in('maturity')}<={_in('comm')},{v['bal_after']},0),"
                  f"IF({A}>={_in('maturity')},{v['bal_after']},{v['annuity']})),0))),0))"),
        "amort": (f"=IF(AND({income},{annuity_mode},{A}>{_in('comm')},"
                  f"{A}<{_in('maturity')},{A}<>{_in('horizon')}),{v['repay']},0)"),
        "bal_close": f"={v['bal_after']}-{v['repay']}",
        "to_equity": (f"={v['sale']}+{v['rent']}+{v['exit']}+{v['residual']}-{v['selling']}"
                      f"-{v['opex']}-{v['ptax']}-{v['exit_cost']}-{v['vat_paid']}-{v['int_paid']}"
                      f"-{v['fee']}+{v['draw']}-{v['repay']}"),
        "year": f"=IF(AND({v['open']}>=1,{v['open']}<={_in('hold')}),INT(({v['open']}-1)/12)+1,0)",
        "noi": f"={v['rent']}-{v['opex']}-{v['ptax']}",
        "book_open": f"={p['book_close'] if r > first else _in('basis')}",
        "recognized": (f"=IF({A}>{_in('horizon')},0,IF({direct},"
                       f"IF({A}<{_in('comm')},0,IF({A}={_in('comm')},"
                       f"SUM({_c('weight')}${first}:{_c('weight')}{r}),{v['weight']}))*{_in('basis')},"
                       f"IF({A}>={_in('comm')},IF(AND({v['exit']}>0,{A}={_in('horizon')}),{v['book_open']},"
                       f"MIN({v['book_open']},{_in('basis')}/({_in('dep_years')}*12))),0)))"),
        "book_close": f"={v['book_open']}-{v['recognized']}",
        "realized": (f"=IF({A}<{_in('comm')},0,IF({A}={_in('comm')},"
                     f"SUM({_c('sale')}${first}:{_c('sale')}{r})-SUM({_c('vat_charged')}${first}:{_c('vat_charged')}{r}),"
                     f"{v['sale']}-{v['vat_charged']}))"),
        "margin": (f"=IF({A}>{_in('horizon')},0,IF({direct},{v['realized']}-{v['recognized']}"
                   f"-{v['selling']}-{v['ptax']},{v['sale']}+{v['rent']}+{v['exit']}-{v['vat_charged']}"
                   f"-{v['recognized']}-{v['selling']}-{v['opex']}-{v['ptax']}-{v['exit_cost']}))"),
        "draw_obj": f"=IF({A}<={_in('comm')},{v['capex']}*{_in('loan_share')},0)",
        # Погашение по причине (`object_financing`): плановое тело, баллон в
        # срок, остаток при выходе, из денег объекта.
        "part_balloon": (f"=IF(AND({v['repay']}>0,{income},{annuity_mode},{A}={_in('maturity')}),"
                         f"{v['repay']}-{v['amort']},0)"),
        "part_exit": (f"=IF(AND({v['repay']}>0,{income},{A}={_in('horizon')},{v['part_balloon']}=0),"
                      f"{v['repay']}-{v['amort']},0)"),
        "part_cash": f"={v['repay']}-{v['amort']}-{v['part_balloon']}-{v['part_exit']}",
        "cal_year": f"=YEAR({v['month']})",
        # Итог объекта как отдельного плательщика (`object_result`): доля общих
        # затрат признаётся графиком его стройки.
        "common_rec": (f"=IF({_in('basis')}<>0,{_in('common_total')}*{v['recognized']}/{_in('basis')},"
                       f"IF({A}={_in('comm')},{_in('common_total')},0))"),
        "obj_margin": f"={v['margin']}-{v['common_rec']}",
        "obj_fin": f"={v['interest']}+{v['fee']}",
        "tx_net": f"={v['obj_margin']}-{v['obj_fin']}",
        "eq_flow": f"={v['to_equity']}-{v['capex']}-{v['common']}-{v['tx_tax']}",
        "eq_cash": f"={v['eq_flow']}-{v['residual']}",
        "eq_cum": f"={p['eq_cum'] + '+' if r > first else ''}{v['eq_cash']}",
        "neg_k": f"=IF({v['eq_cum']}<-0.5,{A},-1)",
        "first_neg": f"=IF({v['eq_cash']}<0,{A},1000000)",
        "df": f"=1/(1+{_in('monthly_rate')})^{A}",
        "eq_disc": f"={v['eq_flow']}*{v['df']}",
    }
    out.update(_tax_formulas(v, p, r, first, "0", _in("tax_rate"), _in("loss_limit"), _c("tx_net")))
    return out


def _object_sheet(book: Workbook, item: dict[str, Any], title: str, costs: _Costs,
                  months: int) -> dict[str, str]:
    """Лист объекта: вводные-ссылки, помесячная таблица, итоги. Адреса итогов."""
    ws = book.create_sheet(title[:31])
    c = item["column"]
    ws["A1"] = f"{item['title']} — построчно `object_flows`; вводные — ссылки на «{INPUTS_SHEET}»"
    ws["A1"].font = BOLD
    strategy = _o("strategy", c)
    spaces_measure = f'{_o("measure", c)}="{SPACES}"'
    default_parking = f"{_p('parking_price_th')}*1000"
    under, over = _o("parking_under_units", c), _o("over_units", c)
    under_price = f"IF({_o('parking_under_mln', c)}>0,{_o('parking_under_mln', c)}*1000000,{default_parking})"
    over_price = f"IF({_o('parking_over_mln', c)}>0,{_o('parking_over_mln', c)}*1000000,{default_parking})"
    comm = f"({_k_of(_o('start', c))}+INT({_o('months', c)}))"
    links = {
        "strategy": f"={strategy}",
        "area": (f"=IF({spaces_measure},{_o('volume', c)},IF({_o('volume', c)}>0,"
                 f"MAX(0,{_o('volume', c)}-MAX(0,{over})*{_o('over_area', c)})"
                 f"*MAX(0,{_o('saleable', c)})/{_o('volume', c)},0))"),
        "spaces": f"={_o('parking_units', c)}",
        "rent": f"={_o('rent', c)}*1000",
        "parking_rent": f"={_o('parking_rent', c)}*1000",
        "rent_index": f"={_o('rent_index', c)}/100",
        # Ставка дана на свою дату и до ввода индексируется (`rent_months_before`).
        "rent_from": f'=IF({_o("rent_date", c)}="",{comm},{_k_of(_o("rent_date", c))})',
        "occ_start": f"=MIN(1,MAX(0,{_o('occ_start', c)}/100))",
        "occ_stable": f"=MIN(1,MAX(0,{_o('occ_stable', c)}/100))",
        "leaseup": f"=MAX(1,INT({_o('leaseup', c)}))",
        "opex": f"=MAX(0,{_o('opex', c)}/100)",
        "rev_mult": f"={_p('revenue_mult')}",
        "property_tax": f"=MAX(0,{_o('property_tax', c)}/100)",
        "vat": f"=MAX(0,{_p('vat')})/100",
        "loan_share": f"=MIN(0.95,MAX(0,{_o('loan_share', c)}/100))",
        "spread": f"={_o('spread', c)}/100",
        "fee": f"=MAX(0,{_o('fee', c)}/100)",
        "reservation": f"=MAX(0,{_p('reservation_fee_pct')}/100)",
        "commitment": f"=MAX(0,{_p('limit_fee_pct')}/100)",
        "permit": f"={costs.cells['permit']}",
        "term": f"=MAX(1,INT({_o('term', c)}))*12",
        "balloon_share": f"=MIN(1,MAX(0,{_o('balloon', c)}/100))",
        "hold": f"=MAX(1,INT({_o('hold_years', c)}))*12",
        "cap": f"={_o('cap', c)}/100",
        "exit_cost": f"=MAX(0,{_o('exit_cost', c)}/100)",
        "exit_mode": f"={_o('exit_mode', c)}",
        "repayment": f"={_o('repayment', c)}",
        "comm": f"={comm}",
        "horizon": (f'=IF({_in("strategy")}="{STRATEGY[ns.STRATEGY_DIRECT]}",'
                    f"MAX({_in('comm')},{_in('sale_start')}+{_in('sale_months')}-1),"
                    f"{_in('comm')}+{_in('hold')}-1)"),
        "price": f"={_o('price', c)}*IF({spaces_measure},1000000,1000)",
        "parking_price": (f"=IF(MAX(0,{under})+MAX(0,{over})>0,(MAX(0,{under})*{under_price}"
                          f"+MAX(0,{over})*{over_price})/(MAX(0,{under})+MAX(0,{over})),{default_parking})"),
        "growth_pre": f"={_o('growth_pre', c)}/100",
        "growth_post": f"={_o('growth_post', c)}/100",
        "price_start": f"={_k_of(_o('sales_start', c))}",
        "selling": f"=({_p('marketing')}+{_p('selling')})/100",
        "cost_mult": f"={_p('cost_mult')}",
        "sale_start": f"={_in('comm')}+MAX(-120,INT({_o('sale_offset', c)}))",
        "sale_months": f"=MAX(1,INT({_o('sale_months', c)}))",
        "curve": f"={_o('curve', c)}",
        "dep_years": f"=MAX(1,INT({_o('dep_years', c)}))",
        "share": f"={costs.share[item['key']]}",
        "tax_rate": f"={_p('profit_tax')}/100",
        "loss_limit": f"={_p('loss_limit')}",
        "discount": f"={_p('discount')}/100",
    }
    for name, (row, label) in INPUTS.items():
        ws[f"A{row}"] = label
        ws[f"B{row}"] = links[name]
    rows = months + EXTRA_MONTHS
    last = FIRST_ROW + rows - 1
    col = lambda name: f"{_c(name)}${FIRST_ROW}:{_c(name)}${last}"  # noqa: E731
    n = _in("sale_months")
    derived = {
        "capex_total": f"=SUM({col('capex')})",
        "basis": f"={_in('capex_total')}*(1-{_in('vat')}/(1+{_in('vat')}))",
        "first_draw": f"=MIN({col('first_k')})",
        "maturity": (f"=MAX(IF({_in('first_draw')}<1000000,{_in('first_draw')},{_in('comm')})"
                     f"+{_in('term')},{_in('comm')})"),
        "comm_balance": f"=MAX({col('bal_build')})",
        "balloon": f"={_in('comm_balance')}*{_in('balloon_share')}",
        "forward_noi": (f'=IF({_is("strategy", STRATEGY, ns.STRATEGY_INCOME)},'
                        f'SUM({col("fwd")})-{_in("basis")}*{_in("property_tax")},0)'),
        "exit_value": f"=IF(AND({_in('cap')}>0,{_in('forward_noi')}>0),{_in('forward_noi')}/{_in('cap')},0)",
        "stabilized_noi": (f'=IF({_is("strategy", STRATEGY, ns.STRATEGY_INCOME)},'
                           f'SUM({col("stab")})-{_in("basis")}*{_in("property_tax")},0)'),
        "curve_sum": (f'=IF({_is("curve", CURVE, "bell")},{n}*({n}+1)*({n}+2)/6,'
                      f'IF(OR({_is("curve", CURVE, "front_loaded")},{_is("curve", CURVE, "back_loaded")}),'
                      f'{n}*({n}+1)/2,{n}))'),
        "peak": f"=MAX({col('bal_close')})",
        "common_total": f"=SUM({col('common')})",
        "monthly_rate": f"=(1+MAX({_in('discount')},-0.999999))^(1/12)-1",
        "loan_limit": f"=SUM({col('draw')})",
        "limit_land": f'=SUMIF({col("k")},"<"&{_in("permit")},{col("draw")})',
        "limit_build": f"={_in('loan_limit')}-{_in('limit_land')}",
        "first_land": f"=IF({_in('first_draw')}<{_in('permit')},{_in('first_draw')},1000000)",
        "first_build": f"=MIN({col('first_build_k')})",
        "reservation_fee": f"=({_in('limit_land')}+{_in('limit_build')})*{_in('reservation')}",
    }
    for name, (row, label) in DERIVED.items():
        ws[f"A{row}"] = label
        ws[f"B{row}"] = derived[name]
    head = FIRST_ROW - 1
    for name, (letter, label) in COLUMNS.items():
        ws[f"{letter}{head}"] = label
        ws[f"{letter}{head}"].font = BOLD
    for i in range(rows):
        r = FIRST_ROW + i
        ws[f"A{r}"] = i
        ws[f"B{r}"] = f"=EDATE({_p('start')},A{r})"
        ws[f"B{r}"].number_format = MONTH
        source = r if i < months else None
        refs = {
            "capex": f"'{COSTS_SHEET}'!{costs.monthly[item['key']]}{source}" if source else "0",
            "common": f"'{COSTS_SHEET}'!{costs.common_col}{source}" if source else "0",
            "key_rate": f"'{CREDIT_SHEET}'!$C{source}" if source else "0",
        }
        for name, formula in _row_formulas(r, refs).items():
            ws[f"{_c(name)}{r}"] = formula
    for i in range(1, len(COLUMNS) + 1):
        ws.column_dimensions[get_column_letter(i)].width = 16
    ws.column_dimensions["A"].width = 58
    # Блок другой стратегии свёрнут группой: у прямой продажи нет аренды и
    # выхода, у аренды — продаж ДКП. Столбцы на месте, их формулы дают нули и
    # разворачиваются кнопкой группы, если стратегию сменить на «Вводных».
    hidden = (("open", "occ", "growth", "rent", "opex", "exit", "exit_cost", "residual", "op_rev",
               "fwd", "stab", "year", "noi")
              if item.get("strategy") == ns.STRATEGY_DIRECT
              else ("weight", "factor", "sale", "selling", "realized"))
    for name in hidden:
        ws.column_dimensions[_c(name)].outlineLevel = 1
        ws.column_dimensions[_c(name)].hidden = True

    # Годы эксплуатации: NOI, проценты, плановое тело, DSCR и ICR.
    years_row = last + 3
    for letter, label in zip("ABCDEFG", ("Год эксплуатации", "NOI", "Проценты", "Плановое тело",
                                         "DSCR", "ICR", "")):
        ws[f"{letter}{years_row}"] = label
    year_count = (rows + 11) // 12
    income = _is("strategy", STRATEGY, ns.STRATEGY_INCOME)
    for y in range(1, year_count + 1):
        r = years_row + y
        ws[f"A{r}"] = y
        ws[f"B{r}"] = f"=SUMIF({col('year')},A{r},{col('noi')})"
        ws[f"C{r}"] = f"=SUMIF({col('year')},A{r},{col('int_paid')})"
        ws[f"D{r}"] = f"=SUMIF({col('year')},A{r},{col('amort')})"
        # Покрытие считается у доходного объекта: у прямой продажи NOI нет.
        ws[f"E{r}"] = f'=IF(AND({income},C{r}>0),B{r}/(C{r}+D{r}),"")'
        ws[f"F{r}"] = f'=IF(AND({income},C{r}>0),B{r}/C{r},"")'
    first_y, last_y = years_row + 1, years_row + year_count
    income = _is("strategy", STRATEGY, ns.STRATEGY_INCOME)
    payback_last = f"MAX({col('neg_k')})"
    payback_first = f"MIN({col('first_neg')})"
    annuity_scheme = f"AND({income},{_is('repayment', REPAY, ns.REPAY_ANNUITY)})"
    draw_object = f"SUM({col('draw_obj')})"
    amort_count = f'COUNTIF({col("amort")},">0")'
    totals = {
        "revenue": f"=SUM({col('sale')})+SUM({col('rent')})+SUM({col('exit')})+SUM({col('residual')})",
        "rent_revenue": f"=SUM({col('rent')})", "sale_revenue": f"=SUM({col('sale')})",
        "exit_revenue": f"=SUM({col('exit')})", "residual_value": f"=SUM({col('residual')})",
        "opex": f"=SUM({col('opex')})", "property_tax": f"=SUM({col('ptax')})",
        "selling_cost": f"=SUM({col('selling')})", "exit_cost": f"=SUM({col('exit_cost')})",
        "vat_paid": f"=SUM({col('vat_paid')})", "vat_charged": f"=SUM({col('vat_charged')})",
        "loan_draw": f"=SUM({col('draw')})",
        "loan_interest": f"=SUM({col('int_cap')})+SUM({col('int_paid')})",
        "loan_fee": f"=SUM({col('fee')})", "loan_repayment": f"=SUM({col('repay')})",
        "issue_fee": f"=SUM({col('fee')})-{_in('reservation_fee')}-SUM({col('commitment_fee')})",
        "commitment_fee": f"=SUM({col('commitment_fee')})",
        "reservation_fee": f"={_in('reservation_fee')}", "limit": f"={_in('loan_limit')}",
        "limit_land": f"={_in('limit_land')}", "limit_build": f"={_in('limit_build')}",
        "loan_peak": f"={_in('peak')}", "tax_margin": f"=SUM({col('margin')})",
        "noi": f"=IF({income},SUM({col('noi')}),0)",
        "exit_value": f"={_in('exit_value')}", "stabilized_noi": f"={_in('stabilized_noi')}",
        "dscr_min": f'=IF(COUNT(E{first_y}:E{last_y})>0,MIN(E{first_y}:E{last_y}),"")',
        "icr_min": f'=IF(COUNT(F{first_y}:F{last_y})>0,MIN(F{first_y}:F{last_y}),"")',
        "capex": f"={_in('capex_total')}",
        "common_capex": f"={_in('common_total')}",
        # Кредит объекта (`object_financing`).
        "draw_object": f"=MIN({draw_object},SUM({col('draw')}))",
        "draw_common": f"=MAX(0,SUM({col('draw')})-{draw_object})",
        "interest_capitalized": f"=SUM({col('int_cap')})",
        "interest_paid": f"=SUM({col('int_paid')})",
        "debt_at_commissioning": f"=SUMIF({col('k')},{_in('comm')},{col('bal_close')})",
        "rate_at_commissioning": f"=SUMIF({col('k')},{_in('comm')},{col('rate')})",
        "avg_rate": (f"=IF(SUM({col('bal_close')})>0,SUMPRODUCT({col('rate')},{col('bal_close')})"
                     f"/SUM({col('bal_close')}),0)"),
        "balloon_planned": f"=IF({annuity_scheme},{_in('balloon')},0)",
        "balloon_paid": f"=SUM({col('part_balloon')})",
        "repaid_scheduled": f"=SUM({col('amort')})",
        "repaid_at_exit": f"=SUM({col('part_exit')})",
        "repaid_from_cash": f"=SUM({col('part_cash')})",
        "avg_payment": (f"=IF({amort_count}>0,SUMPRODUCT(({col('amort')}>0)*({col('int_paid')}"
                        f"+{col('amort')}))/{amort_count},0)"),
        "maturity": f"=IF({annuity_scheme},EDATE({_p('start')},{_in('maturity')}),\"\")",
        # Итог объекта как отдельного проекта (`object_result`).
        "cost_total": f"={_in('capex_total')}+{_in('common_total')}",
        "equity_invested": f"={_in('capex_total')}+{_in('common_total')}-SUM({col('draw')})",
        "equity_peak": f"=-MIN(0,MIN({col('eq_cum')}))",
        "common_recognized": f"=SUM({col('common_rec')})",
        "financing_cost": f"=SUM({col('obj_fin')})",
        "profit_before_tax": f"=SUM({col('obj_margin')})-SUM({col('obj_fin')})",
        "profit_tax": f"=SUM({col('tx_tax')})",
        "equity_cash": f"=SUM({col('eq_cash')})",
        "irr": f'=IFERROR((1+IRR({col("eq_flow")},0.01))^12-1,"")',
        "irr_cash": f'=IF(SUM({col("residual")})>0,IFERROR((1+IRR({col("eq_cash")},0.01))^12-1,""),"")',
        "npv": f"=SUM({col('eq_disc')})",
        "payback_months": (f'=IF(OR({payback_first}>=1000000,{payback_last}<0,'
                           f'{payback_last}+1>{_in("horizon")}),"",{payback_last}+1-{payback_first})'),
    }
    out: dict[str, str] = {}
    totals_row = years_row + year_count + 3
    ws[f"A{totals_row}"] = "Итоги объекта (формулы книги)"
    ws[f"A{totals_row}"].font = BOLD
    for i, (name, formula) in enumerate(totals.items(), start=1):
        r = totals_row + i
        ws[f"A{r}"] = TOTAL_LABELS.get(name, name)
        ws[f"B{r}"] = formula
        ws[f"B{r}"].number_format = MONTH if name == "maturity" else MONEY
        out[name] = f"'{ws.title}'!$B${r}"
    out["sheet"] = ws.title
    out["comm"] = f"'{ws.title}'!{_in('comm')}"
    out["horizon"] = f"'{ws.title}'!{_in('horizon')}"
    out["strategy"] = f"'{ws.title}'!{_in('strategy')}"
    return out


TOTAL_LABELS = {
    "revenue": "Выручка объекта", "rent_revenue": "Арендная выручка",
    "sale_revenue": "Выручка прямых продаж", "exit_revenue": "Выход — продажа",
    "residual_value": "Удержание — оценка", "opex": "OPEX", "property_tax": "Налог на имущество",
    "selling_cost": "Расходы на продажу", "exit_cost": "Затраты на выход",
    "vat_paid": "НДС к уплате", "vat_charged": "НДС начисленный",
    "loan_draw": "Кредит — выборка", "loan_interest": "Кредит — проценты",
    "loan_fee": "Кредит — комиссии всего", "loan_repayment": "Кредит — погашение",
    "issue_fee": "Комиссия за выдачу", "reservation_fee": "Плата за резервирование лимита НКЛ",
    "limit": "Лимит НКЛ объекта", "commitment_fee": "Плата за невыбранный лимит НКЛ",
    "limit_land": "Лимит НКЛ на участок и проект (до РнС)", "limit_build": "Лимит строительной НКЛ (с РнС)",
    "loan_peak": "Кредит — пик долга", "noi": "NOI за срок удержания",
    "tax_margin": "Налоговая маржа объекта", "exit_value": "Стоимость выхода",
    "stabilized_noi": "Стабилизированный NOI", "dscr_min": "DSCR — минимум",
    "icr_min": "ICR — минимум", "capex": "Затраты стройки объекта",
    "common_capex": "Общие затраты проекта на объекте",
    "draw_object": "Выборка на стройку объекта", "draw_common": "Выборка на общие затраты",
    "interest_capitalized": "Проценты до ввода (капитализированы)",
    "interest_paid": "Проценты после ввода (уплачены)",
    "debt_at_commissioning": "Долг на ввод", "rate_at_commissioning": "Ставка на вводе",
    "avg_rate": "Средняя ставка по остатку долга", "balloon_planned": "Баллон по графику",
    "balloon_paid": "Баллон погашен в срок", "repaid_scheduled": "Погашено плановым телом",
    "repaid_at_exit": "Погашено при выходе", "repaid_from_cash": "Погашено из выручки / NOI",
    "avg_payment": "Аннуитетный платёж — средний", "maturity": "Срок кредита",
    "cost_total": "Затраты объекта всего", "equity_invested": "Собственный капитал объекта",
    "equity_peak": "Пик капитала объекта", "common_recognized": "Общие затраты — признано",
    "financing_cost": "Проценты и комиссии объекта", "profit_before_tax": "Прибыль объекта до налога",
    "profit_tax": "Налог на прибыль объекта", "equity_cash": "Денежный итог объекта",
    "irr": "IRR капитала объекта", "irr_cash": "IRR капитала объекта — только деньги",
    "npv": "NPV капитала объекта", "payback_months": "Окупаемость капитала объекта, мес.",
}


# --- проект: «Кредит», «Налоги», «Денежный поток» ----------------------------

def _object_sum(objects_refs: list[dict[str, str]], column: str, r: int) -> str:
    return _sum([f"'{ref['sheet']}'!{_c(column)}{r}" for ref in objects_refs])


def _credit_sheet(book: Workbook, objects_refs: list[dict[str, str]], inputs: _Inputs,
                  months: int) -> dict[str, str]:
    ws = book.create_sheet(CREDIT_SHEET)
    ws["A1"] = "Кредит объектов: ключевая ставка, выборка, проценты, погашение, долг"
    ws["A1"].font = BOLD
    ws["A2"] = "Помесячно — сумма по объектам; условия и итог каждого кредита — на листе объекта и в «Отчёте»."
    columns = (("k", "k"), ("month", "Месяц"), ("key_rate", "Ключевая ставка"),
               ("draw", "Выборка"), ("int_cap", "Проценты капитализированные"),
               ("int_paid", "Проценты уплаченные"), ("fee", "Комиссия"),
               ("repay", "Погашение"), ("bal_close", "Долг на конец"))
    letters = {key: get_column_letter(i + 1) for i, (key, _) in enumerate(columns)}
    for key, label in columns:
        ws[f"{letters[key]}{FIRST_ROW - 1}"] = label
        ws[f"{letters[key]}{FIRST_ROW - 1}"].font = BOLD
    dates = f"'{INPUTS_SHEET}'!$B${inputs.rate_first}:$B${inputs.rate_last}"
    scenario = _p("rate_scenario")
    by_scenario = {label: f"'{INPUTS_SHEET}'!${column}${inputs.rate_first}:${column}${inputs.rate_last}"
                   for label, column in zip(SCENARIO.values(), "CDE")}
    for i in range(months):
        r = FIRST_ROW + i
        ws[f"A{r}"] = i
        ws[f"B{r}"] = f"=EDATE({_p('start')},A{r})"
        ws[f"B{r}"].number_format = MONTH
        position = f'MAX(1,COUNTIF({dates},"<="&B{r}))'
        pick = (f'IF({scenario}="{SCENARIO["low"]}",INDEX({by_scenario[SCENARIO["low"]]},{position}),'
                f'IF({scenario}="{SCENARIO["high"]}",INDEX({by_scenario[SCENARIO["high"]]},{position}),'
                f'INDEX({by_scenario[SCENARIO["base"]]},{position})))')
        ws[f"C{r}"] = f"={pick}/100"
        ws[f"C{r}"].number_format = "0.00%"
        for key in ("draw", "int_cap", "int_paid", "fee", "repay", "bal_close"):
            ws[f"{letters[key]}{r}"] = "=" + _object_sum(objects_refs, key, r)
            ws[f"{letters[key]}{r}"].number_format = MONEY
    for i in range(1, len(columns) + 1):
        ws.column_dimensions[get_column_letter(i)].width = 18
    last = FIRST_ROW + months - 1
    return {key: f"'{CREDIT_SHEET}'!${letters[key]}${FIRST_ROW}:${letters[key]}${last}" for key in letters}


def _tax_sheet(book: Workbook, objects_refs: list[dict[str, str]], objects: list[dict[str, Any]],
               costs: _Costs, months: int, ddu_refs: list[dict[str, str]] | None = None,
               pf: dict[str, str] | None = None) -> dict[str, str]:
    """Налог на прибыль проекта и НДС: база — маржа объектов, общие затраты
    проекта признаются в РВЭ (продаж ядра у нежилого проекта нет), вычет —
    начисленные проценты и комиссии; правило переноса убытка — то же, что у
    объектов и у движка."""
    ws = book.create_sheet(TAX_SHEET)
    ws["A1"] = "Налоги проекта: НДС и налог на прибыль с переносом убытка"
    ws["A1"].font = BOLD
    pool = "+".join(f"{costs.amounts[item['key'] + ':building']}+{costs.amounts[item['key'] + ':garage']}"
                    for item in objects if item["in_tax_pool"]) or "0"
    ws["A3"] = "Общие затраты проекта — признаются в РВЭ, ₽ (CAPEX минус статьи объектов)"
    # Маркетинг и продажи объектов по ДДУ — в общем пуле (`commercial_costs`).
    ws["B3"] = (f"=MAX(0,{costs.cells['capex_total']}"
                + (f"+SUM({pf['operating']})" if pf else "") + f"-({pool}))")
    ws["A4"] = "РВЭ — маржа объектов признаётся в этот месяц (k)"
    ws["B4"] = f"={costs.cells['rve']}"
    ws["A5"] = "Ставка налога на прибыль"
    ws["B5"] = f"={_p('profit_tax')}/100"
    ws["A6"] = "Зачёт убытка — не более доли базы года"
    ws["B6"] = f"={_p('loss_limit')}"
    for r in range(3, 7):
        ws[f"B{r}"].number_format = MONEY if r == 3 else "0.00"
    columns = (("k", "k"), ("month", "Месяц"), ("cal_year", "Год"),
               ("vat_charged", "НДС начисленный"), ("vat_paid", "НДС к уплате (− возмещение)"),
               ("objects_margin", "Маржа объектов"), ("common", "Общие затраты (в РВЭ)"),
               ("margin", "Маржа проекта"), ("financing", "Проценты и комиссии — вычет"),
               ("tx_net", "База месяца"), ("tx_gate", "До РВЭ"), ("tx_year", "Результат года"),
               ("tx_prior", "Убыток прошлых лет"), ("tx_used", "Зачтено убытка"),
               ("tx_base", "База года"), ("tx_paid", "Уплачено ранее в году"),
               ("tx_tax", "Налог на прибыль"),
               ("tx_pay", "Налог к уплате (авансы по кварталам, доплата года в марте)"))
    letters = {key: get_column_letter(i + 1) for i, (key, _) in enumerate(columns)}
    for key, label in columns:
        ws[f"{letters[key]}{FIRST_ROW - 1}"] = label
        ws[f"{letters[key]}{FIRST_ROW - 1}"].font = BOLD
    for i in range(months):
        r = FIRST_ROW + i
        v = {key: f"{letters[key]}{r}" for key in letters}
        p = {key: f"{letters[key]}{r - 1}" for key in letters}
        ws[f"A{r}"] = i
        ws[f"B{r}"] = f"=EDATE({_p('start')},A{r})"
        ws[f"B{r}"].number_format = MONTH
        ws[v["cal_year"]] = f"=YEAR({v['month']})"
        ws[v["vat_charged"]] = "=" + _object_sum(objects_refs, "vat_charged", r)
        ddu_vat = f"+INDEX({pf['vat']},{i + 1})" if pf else ""
        ws[v["vat_paid"]] = "=" + _object_sum(objects_refs, "vat_paid", r) + ddu_vat
        ws[v["objects_margin"]] = "=" + _object_sum(objects_refs, "margin", r) + "".join(
            f"+INDEX({ref['margin']},{i + 1})" for ref in ddu_refs or [])
        ws[v["common"]] = f"=IF(A{r}=$B$4,-$B$3,0)"
        # НДС объектов по ДДУ к уплате уходит из базы налога на прибыль.
        ws[v["margin"]] = f"={v['objects_margin']}+{v['common']}" + (
            f"-INDEX({pf['vat']},{i + 1})" if pf else "")
        ws[v["financing"]] = "=" + _sum([f"'{ref['sheet']}'!{_c('interest')}{r}+'{ref['sheet']}'!{_c('fee')}{r}"
                                         for ref in objects_refs]) + (
            f"+INDEX({pf['deduct']},{i + 1})" if pf else "")
        ws[v["tx_net"]] = f"={v['margin']}-{v['financing']}"
        # Пересказ `_profit_tax_schedule` спецзастройщика: до РВЭ откладывается
        # только маржа, проценты — в месяце начисления.
        tax = _project_tax_formulas(v, p, r, FIRST_ROW, f"IF(A{r}<$B$4,1,0)", "$B$5", "$B$6",
                                    letters["margin"])
        for key, formula in tax.items():
            ws[v[key]] = formula
        ws[v["tx_pay"]] = _tax_payment_formula(v, r, FIRST_ROW, letters["tx_tax"])
        for key in letters:
            if key not in ("k", "month", "cal_year", "tx_gate"):
                ws[v[key]].number_format = MONEY
    for i in range(1, len(columns) + 1):
        ws.column_dimensions[get_column_letter(i)].width = 17
    ws.column_dimensions["A"].width = 60
    last = FIRST_ROW + months - 1
    return {key: f"'{TAX_SHEET}'!${letters[key]}${FIRST_ROW}:${letters[key]}${last}" for key in letters}


def _cash_sheet(book: Workbook, objects_refs: list[dict[str, str]], costs: _Costs,
                taxes: dict[str, str], months: int, pf: dict[str, str] | None = None) -> dict[str, str]:
    ws = book.create_sheet(CASH_SHEET)
    ws["A1"] = "Денежный поток проекта и собственника; NPV, IRR, окупаемость"
    ws["A1"].font = BOLD
    columns = (("k", "k"), ("month", "Месяц"), ("cal_year", "Год"),
               ("revenue", "Выручка объектов"), ("capex", "CAPEX проекта"),
               ("costs", "Расходы объектов: продажи, эксплуатация, налог на имущество, выход"),
               ("vat", "НДС к уплате"), ("interest", "Проценты начисленные"),
               ("fee", "Комиссии (выдача, резервирование лимита НКЛ)"),
               ("tax", "Налог на прибыль"), ("project", "Поток проекта"),
               ("to_equity", "Деньги объектов собственнику"), ("equity", "Поток собственного капитала"),
               ("residual", "Оценка удержания (не деньги)"), ("cash", "Поток капитала — деньги"),
               ("cum", "Накопленный — деньги"), ("df", "Дисконт-множитель"),
               ("project_disc", "Поток проекта дисконтированный"),
               ("equity_disc", "Поток капитала дисконтированный"),
               ("contributed", "Вложено собственником"), ("returned", "Получено собственником"),
               ("exit", "Месяц выхода"), ("in_build", "Вложено до ввода"),
               ("in_exit", "Вложено при выходе"), ("out_vat", "Возмещение НДС стройки"),
               ("out_exit", "Получено при выходе"), ("neg_k", "Служебная: накопленный < 0"))
    letters = {key: get_column_letter(i + 1) for i, (key, _) in enumerate(columns)}
    for key, label in columns:
        ws[f"{letters[key]}{FIRST_ROW - 1}"] = label
        ws[f"{letters[key]}{FIRST_ROW - 1}"].font = BOLD
    ws["A3"] = "Ставка дисконтирования, месячная"
    ws["B3"] = f"=(1+MAX({_p('discount')}/100,-0.999999))^(1/12)-1"
    ws["A4"] = "Ввод последнего объекта (k)"
    ws["B4"] = "=MAX(" + ",".join(ref["comm"] for ref in objects_refs) + ")" if objects_refs else "=0"
    income_label = STRATEGY[ns.STRATEGY_INCOME]
    for i in range(months):
        r = FIRST_ROW + i
        v = {key: f"{letters[key]}{r}" for key in letters}
        ws[f"A{r}"] = i
        ws[f"B{r}"] = f"=EDATE({_p('start')},A{r})"
        ws[f"B{r}"].number_format = MONTH
        ws[v["cal_year"]] = f"=YEAR({v['month']})"
        at = lambda name: f"+INDEX({pf[name]},{i + 1})" if pf else ""  # noqa: E731
        ws[v["revenue"]] = "=" + _sum([
            f"'{ref['sheet']}'!{_c(name)}{r}" for ref in objects_refs
            for name in ("sale", "rent", "exit", "residual")]) + at("sales")
        ws[v["capex"]] = f"='{COSTS_SHEET}'!{costs.total_col}{r}"
        ws[v["costs"]] = "=" + _sum([
            f"'{ref['sheet']}'!{_c(name)}{r}" for ref in objects_refs
            for name in ("selling", "opex", "ptax", "exit_cost")]) + at("operating")
        # НДС к уплате: объектов вне ДДУ — их листы, по ДДУ — лист «ПФ» в РВЭ.
        ws[v["vat"]] = "=" + _object_sum(objects_refs, "vat_paid", r) + at("vat")
        ws[v["interest"]] = "=" + _object_sum(objects_refs, "interest", r) + at("pay_out")
        ws[v["fee"]] = "=" + _object_sum(objects_refs, "fee", r) + (
            f"+IF(A{r}=0,{_pi('bridge_fee')},0)+IF(A{r}={_pi('permit')},{_pi('reservation_fee')},0)"
            if pf else "")
        # Поток видит уплату, а не начисление: деньги уходят по кварталам.
        ws[v["tax"]] = f"=INDEX({taxes['tx_pay']},{i + 1})"
        # Поток проекта — до финансирования, как у движка (владелец,
        # 16.10.2026): без процентов и комиссий; по нему NPV и IRR проекта.
        ws[v["project"]] = f"={v['revenue']}-{v['costs']}-{v['vat']}-{v['capex']}-{v['tax']}"
        ws[v["to_equity"]] = "=" + _object_sum(objects_refs, "to_equity", r)
        ws[v["equity"]] = f"={v['to_equity']}-{v['capex']}-{v['tax']}" + (
            f"+INDEX({pf['equity']},{i + 1})" if pf else "")
        ws[v["residual"]] = "=" + _object_sum(objects_refs, "residual", r)
        ws[v["cash"]] = f"={v['equity']}-{v['residual']}"
        ws[v["cum"]] = (f"={letters['cum']}{r - 1}+{v['cash']}" if i else f"={v['cash']}")
        ws[v["df"]] = f"=1/(1+$B$3)^A{r}"
        ws[v["project_disc"]] = f"={v['project']}*{v['df']}"
        ws[v["equity_disc"]] = f"={v['equity']}*{v['df']}"
        ws[v["contributed"]] = f"=MAX(0,-{v['cash']})"
        ws[v["returned"]] = f"=MAX(0,{v['cash']})"
        exits = [f'AND({ref["strategy"]}="{income_label}",A{r}={ref["horizon"]})' for ref in objects_refs]
        ws[v["exit"]] = f"=IF(OR({','.join(exits)}),1,0)" if exits else "=0"
        ws[v["in_build"]] = f"=IF(A{r}<=$B$4,{v['contributed']},0)"
        ws[v["in_exit"]] = f"=IF(AND(A{r}>$B$4,{v['exit']}=1),{v['contributed']},0)"
        ws[v["out_vat"]] = f"=IF({v['cash']}>0,MIN({v['cash']},MAX(0,-{v['vat']})),0)"
        ws[v["out_exit"]] = f"=IF({v['exit']}=1,{v['returned']}-{v['out_vat']},0)"
        ws[v["neg_k"]] = f"=IF({v['cum']}<-0.5,A{r},-1)"
        for key in letters:
            if key not in ("k", "month", "cal_year", "df", "exit", "neg_k"):
                ws[v[key]].number_format = MONEY
    for i in range(1, len(columns) + 1):
        ws.column_dimensions[get_column_letter(i)].width = 17
    ws.column_dimensions["A"].width = 40
    last = FIRST_ROW + months - 1
    return {key: f"'{CASH_SHEET}'!${letters[key]}${FIRST_ROW}:${letters[key]}${last}" for key in letters}


# --- «Объект по ДДУ» и «ПФ» ----------------------------------------------------
#
# Решение владельца (13.10.2026): «рабочая книга с живыми формулами, а не
# выгрузка хардов». Объект нежилого проекта, проданный по ДДУ, прежде уводил
# весь проект в жилую книгу v4. Теперь его продажи и проектное финансирование
# считает книга нежилого проекта — построчным пересказом движка:
# `sales_schedule` (доля до ввода объекта + остаток, рост цены), эскроу до РВЭ
# проекта и `simulate_financing` (собственные средства, БРИДЖ до РнС, ПФ с
# покрытием эскроу и ступенями специальной ставки, плата за лимит,
# раскрытие эскроу в РВЭ, погашение продажами после него).

DDU_INPUTS: dict[str, tuple[int, str]] = {key: (2 + i, label) for i, (key, label) in enumerate((
    ("area", "Продаваемая площадь, м² (места — шт.)"),
    ("spaces", "Машино-места объекта в продаже, шт."),
    ("price", "Цена продажи, ₽/м² с НДС (места — ₽/место)"),
    ("parking_price", "Цена машино-места, ₽ (средняя по гаражу)"),
    ("growth_pre", "Рост цены до ввода, доля в месяц"),
    ("growth_post", "Рост цены после ввода, доля в месяц"),
    ("price_start", "Старт продаж (k)"),
    ("comm", "Ввод объекта (k)"),
    ("share", "Доля продаж до ввода"),
    ("residual", "Остаточные продажи после ввода, мес."),
    ("pre_n", "Месяцев продаж до ввода"),
    ("rev_mult", "Множитель выручки сценария"),
    ("cost", "Затраты объекта к признанию продажами, ₽ (здание и гараж)"),
    ("qty_total", "Продано за горизонт, м² (места — шт.)"),
))}
DDU_COLUMNS: dict[str, tuple[str, str]] = {key: (get_column_letter(i + 1), label) for i, (key, label) in enumerate((
    ("k", "k"), ("month", "Месяц"), ("weight", "Доля продаж месяца"),
    ("price", "Цена, ₽/м²"), ("parking_price", "Цена места, ₽"),
    ("qty", "Продано, м²"), ("sale", "Выручка площади"), ("parking", "Выручка мест"),
    ("revenue", "Выручка объекта"), ("base", "Выручка без сценария (база расходов на продажу)"),
    ("margin", "Налоговая маржа: выручка − затраты по проданным м²"),
))}


def _dc(name: str) -> str:
    return DDU_COLUMNS[name][0]


def _di(name: str) -> str:
    return f"$B${DDU_INPUTS[name][0]}"


def _ddu_sheet(book: Workbook, item: dict[str, Any], title: str, costs: _Costs,
               inputs: _Inputs, months: int) -> dict[str, str]:
    """Лист объекта по ДДУ: продажи формулами `sales_schedule`."""
    ws = book.create_sheet(title[:31])
    c = item["column"]
    ws["A1"] = (f"{item['title']} — продажа по ДДУ: `sales_schedule` движка; вводные — ссылки на "
                f"«{INPUTS_SHEET}», деньги идут на эскроу (лист «{PF_SHEET}»)")
    ws["A1"].font = BOLD
    spaces_measure = f'{_o("measure", c)}="{SPACES}"'
    default_parking = f"{_p('parking_price_th')}*1000"
    under, over = _o("parking_under_units", c), _o("over_units", c)
    under_price = f"IF({_o('parking_under_mln', c)}>0,{_o('parking_under_mln', c)}*1000000,{default_parking})"
    over_price = f"IF({_o('parking_over_mln', c)}>0,{_o('parking_over_mln', c)}*1000000,{default_parking})"
    last = FIRST_ROW + months - 1
    col = lambda name: f"${_dc(name)}${FIRST_ROW}:${_dc(name)}${last}"  # noqa: E731
    links = {
        "area": (f"=IF({spaces_measure},{_o('volume', c)},IF({_o('volume', c)}>0,"
                 f"MAX(0,{_o('volume', c)}-MAX(0,{over})*{_o('over_area', c)})"
                 f"*MAX(0,{_o('saleable', c)})/{_o('volume', c)},0))"),
        "spaces": f"=MAX(0,{_o('parking_units', c)})",
        "price": f"={_o('price', c)}*IF({spaces_measure},1000000,1000)",
        "parking_price": (f"=IF(MAX(0,{under})+MAX(0,{over})>0,(MAX(0,{under})*{under_price}"
                          f"+MAX(0,{over})*{over_price})/(MAX(0,{under})+MAX(0,{over})),{default_parking})"),
        "growth_pre": f"={_o('growth_pre', c)}/100",
        "growth_post": f"={_o('growth_post', c)}/100",
        "price_start": f"={_k_of(_o('sales_start', c))}",
        "comm": f"={_k_of(_o('start', c))}+INT({_o('months', c)})",
        "share": f"=MAX(0,MIN(1,{_o('share_before', c)}/100))",
        "residual": f"=MAX(0,INT({_o('residual', c)}))",
        "pre_n": f"=MAX(1,{_di('comm')}-{_di('price_start')})",
        "rev_mult": f"={_p('revenue_mult')}",
        "cost": (f"={costs.amounts[item['key'] + ':building']}"
                 f"+{costs.amounts[item['key'] + ':garage']}"),
        "qty_total": f"=SUM({col('qty')})",
    }
    for name, (row, label) in DDU_INPUTS.items():
        ws[f"A{row}"] = label
        ws[f"B{row}"] = links[name]
        ws[f"B{row}"].number_format = MONEY if name in ("price", "parking_price", "cost") else "0.0000"
    weights = inputs.schedule.get((item["key"], "sales_weights"))
    factor = inputs.schedule.get((item["key"], "price_factor"))
    for key, (letter, label) in DDU_COLUMNS.items():
        ws[f"{letter}{FIRST_ROW - 1}"] = label
        ws[f"{letter}{FIRST_ROW - 1}"].font = BOLD
    s, comm, n_pre = _di("price_start"), _di("comm"), _di("pre_n")
    for i in range(months):
        r = FIRST_ROW + i
        v = {name: f"{_dc(name)}{r}" for name in DDU_COLUMNS}
        k = v["k"]
        ws[k] = i
        ws[v["month"]] = f"=EDATE({_p('start')},{k})"
        ws[v["month"]].number_format = MONTH
        # Доли `sales_weights`: доля до ввода ровно на месяцы до ввода, остаток
        # — равными долями после; профиль продаж, если задан, их замещает.
        ws[v["weight"]] = (
            f"=INDEX({weights},{k}+1)" if weights else
            (f"=IF(AND({k}>={s},{k}<{s}+{n_pre}),{_di('share')}/{n_pre},0)"
             f"+IF(AND({_di('residual')}>0,{k}>={comm},{k}<{comm}+{_di('residual')}),"
             f"(1-{_di('share')})/{_di('residual')},0)"))
        # Цена: до ввода — рост от старта продаж, после — от цены ввода; при
        # лестнице до ввода её ступень, после — рост от достигнутой ступени.
        growth = (f"IF({k}<{comm},(1+{_di('growth_pre')})^({k}-{s}),"
                  f"(1+{_di('growth_pre')})^{n_pre}*(1+{_di('growth_post')})^({k}-{comm}))")
        if factor:
            growth = (f"INDEX({factor},{k}+1)*IF({k}>={comm},"
                      f"(1+{_di('growth_post')})^({k}-{comm}),1)")
        ws[v["price"]] = f"={_di('price')}*{growth}"
        ws[v["parking_price"]] = f"={_di('parking_price')}*{growth}"
        ws[v["qty"]] = f"={_di('area')}*{v['weight']}"
        ws[v["base"]] = (f"={_di('area')}*{v['weight']}*{v['price']}"
                         f"+{_di('spaces')}*{v['weight']}*{v['parking_price']}")
        ws[v["sale"]] = f"={_di('area')}*{v['weight']}*{v['price']}*{_di('rev_mult')}"
        ws[v["parking"]] = f"={_di('spaces')}*{v['weight']}*{v['parking_price']}*{_di('rev_mult')}"
        ws[v["revenue"]] = f"={v['sale']}+{v['parking']}"
        # Затраты объекта признаются по проданным м²; места своих затрат не
        # несут — они в затратах объекта (`_profit_tax_schedule`).
        ws[v["margin"]] = (f"={v['revenue']}-IF({_di('qty_total')}>0,"
                           f"{_di('cost')}*{v['qty']}/{_di('qty_total')},0)")
        for name in ("price", "parking_price", "sale", "parking", "revenue", "base", "margin"):
            ws[v[name]].number_format = MONEY
    ws.column_dimensions["A"].width = 60
    for i in range(2, len(DDU_COLUMNS) + 1):
        ws.column_dimensions[get_column_letter(i)].width = 17
    return {name: f"'{ws.title}'!{col(name)}" for name in DDU_COLUMNS}


PF_SHEET = "ПФ"
PF_CELLS: dict[str, tuple[int, str]] = {key: (2 + i, label) for i, (key, label) in enumerate((
    ("permit", "РнС — открытие ПФ (k)"),
    ("rve", "РВЭ — раскрытие эскроу (k)"),
    ("own_funds", "Собственные средства до ПФ, ₽"),
    ("approved", "Одобренный лимит ПФ, ₽ (0 — нет)"),
    ("limit", "Лимит ПФ, ₽ (одобренный, иначе выборка с округлением до 10 млн вверх)"),
    ("common_share", "Доля общих затрат на ПФ (1 − доли объектов вне ДДУ)"),
    ("bridge_limit", "Расчётный лимит БРИДЖа: участок и проект до РнС, ₽"),
    ("bridge_fee", "Плата за резервирование БРИДЖа, ₽ (в первом месяце)"),
    ("reservation_fee", "Плата за резервирование лимита ПФ, ₽ (в РнС)"),
    ("pf_draw_total", "Выборка ПФ всего, ₽"),
    ("transferred", "Проценты БРИДЖа, перенесённые в ПФ, ₽"),
    ("financing_cost", "Проценты и комиссии БРИДЖа и ПФ, ₽"),
    ("llcr_num", "LLCR: числитель, ₽"),
    ("llcr_den", "LLCR: знаменатель, ₽"),
    ("llcr", "LLCR"),
    ("peak_bridge", "БРИДЖ — пик, ₽"),
    ("peak_pf", "ПФ — пик, ₽"),
    ("peak_escrow", "Эскроу — пик, ₽"),
    ("ending_pf", "Долг ПФ на конец проекта, ₽"),
    ("rve_unpaid", "Не погашено раскрытым эскроу в РВЭ, ₽"),
    ("vat_due", "НДС объектов по ДДУ к уплате в РВЭ, ₽"),
))}
PF_COLUMNS: dict[str, tuple[str, str]] = {key: (get_column_letter(i + 1), label) for i, (key, label) in enumerate((
    ("k", "k"), ("month", "Месяц"), ("key_rate", "Ключевая ставка"),
    ("sales", "Продажи объектов по ДДУ"), ("capex", "Затраты на ПФ (стройка ДДУ и доля общих)"),
    ("operating", "Маркетинг и продажи"), ("costs", "Расходы проекта на финансирование"),
    ("esc", "Эскроу (до раскрытия)"), ("release", "Раскрытие эскроу"),
    ("own_draw", "Собственные средства"), ("own_left", "Собственные средства — остаток"),
    ("b_draw", "БРИДЖ — выборка"), ("b_bal", "БРИДЖ — долг"), ("b_int", "БРИДЖ — проценты"),
    ("b_cap", "БРИДЖ — капитализация процентов"), ("b_pay", "БРИДЖ — проценты к уплате"),
    ("b_repay", "БРИДЖ — погашение (рефинансирование в ПФ)"),
    ("transfer", "Проценты БРИДЖа при открытии ПФ"),
    ("p_draw", "ПФ — выборка"), ("p_bal_after", "ПФ — долг после выборки"),
    ("cov", "Покрытие эскроу"), ("special", "Специальная ставка ступени"),
    ("p_rate", "Ставка ПФ"), ("p_int", "ПФ — проценты"), ("p_cap", "ПФ — капитализация процентов"),
    ("fee", "Плата за невыбранный лимит"), ("p_pay_open", "ПФ — проценты к уплате на начало"),
    ("pay_out", "Проценты и комиссии уплачены"), ("p_pay", "ПФ — проценты к уплате"),
    ("avail", "Деньги на погашение"), ("p_repay", "ПФ — погашение"), ("p_bal", "ПФ — долг"),
    ("settled", "ПФ закрыт в РВЭ"), ("deduct", "Вычет из базы налога: проценты и комиссии"),
    ("vat", "НДС к уплате"), ("project", "Поток проекта: ДДУ"),
    ("equity", "Поток капитала: ДДУ и ПФ"),
))}


def _pf_link_taxes(book: Workbook, pf: dict[str, str], ddu_refs: list[dict[str, str]],
                   objects: list[dict[str, Any]], costs: _Costs, taxes: dict[str, str],
                   objects_refs: list[dict[str, str]], months: int) -> None:
    """НДС объектов по ДДУ и числитель LLCR — после листа «Налоги».

    НДС — как у движка: начислен в РВЭ со всей выручки по ДДУ, вычет — входящий
    НДС затрат проекта без участка, платы за ВРИ и объектов вне ДДУ; к уплате
    разница. LLCR — покрытие ПФ: выручка по ДДУ, расходы на продажу, налог на
    прибыль, НДС по ДДУ, затраты без объектов вне ДДУ и выборка ПФ против
    выборки, процентов и комиссий.
    """
    ws = book[PF_SHEET]
    last = FIRST_ROW + months - 1
    revenue = "+".join(f"SUM({ref['revenue']})" for ref in ddu_refs)
    gross = f"(MAX(0,{_p('vat')})/100)/(1+MAX(0,{_p('vat')})/100)"
    vat_cols = [key for key in ("vri_interest", "vri_security") if key in costs.monthly]
    column = lambda key: (f"'{COSTS_SHEET}'!${costs.monthly[key]}${FIRST_ROW}:"  # noqa: E731
                          f"${costs.monthly[key]}${last}")
    bearing = (f"{costs.cells['capex_total']}+SUM({pf['operating']})-{_p('purchase_price')}*1000000"
               f"-{_p('land_buyout')}*1000000-{_p('land_rights_gross')}*1000000"
               + "".join(f"-SUM({column(key)})" for key in vat_cols)
               + f"-SUM({column('nonres_obj')})")
    ws[f"B{PF_CELLS['vat_due'][0]}"] = (f"=IF({_p('vat')}>0,MAX(({revenue})*{gross}"
                                       f"-({bearing})*{gross}*IF(({revenue})>0,1,0),0),0)")
    # Кредит объектов вне ДДУ — их проценты и комиссии в знаменателе, как у
    # движка (`financing_cost` целиком).
    interest = "+".join(
        f"SUM('{ref['sheet']}'!${_c(name)}${FIRST_ROW}:${_c(name)}${FIRST_ROW + months + EXTRA_MONTHS - 1})"
        for ref in objects_refs for name in ("interest", "fee")) or "0"
    ws[f"B{PF_CELLS['llcr_num'][0]}"] = (
        f"=({revenue})-SUM({pf['operating']})-SUM({taxes['tx_tax']})-{_pi('vat_due')}"
        f"-({costs.cells['capex_total']}-SUM({column('nonres_obj')}))+{_pi('pf_draw_total')}")
    ws[f"B{PF_CELLS['llcr_den'][0]}"] = f"={_pi('pf_draw_total')}+{_pi('financing_cost')}+{interest}"


def _pc(name: str) -> str:
    return PF_COLUMNS[name][0]


def _pi(name: str) -> str:
    return f"'{PF_SHEET}'!$B${PF_CELLS[name][0]}"


def _pf_sheet(book: Workbook, ddu_refs: list[dict[str, str]], objects: list[dict[str, Any]],
              costs: _Costs, credit: dict[str, str], inputs: _Inputs, months: int) -> dict[str, str]:
    """Проектное финансирование объектов по ДДУ — `simulate_financing` строкой Excel."""
    ws = book.create_sheet(PF_SHEET)
    ws["A1"] = ("ПФ объектов по ДДУ: собственные средства и БРИДЖ до РнС, ПФ с покрытием эскроу, "
                "раскрытие эскроу в РВЭ — построчно `simulate_financing`")
    ws["A1"].font = BOLD
    last = FIRST_ROW + months - 1
    col = lambda name: f"${_pc(name)}${FIRST_ROW}:${_pc(name)}${last}"  # noqa: E731
    cost_col = lambda key: (f"'{COSTS_SHEET}'!${costs.monthly[key]}${FIRST_ROW}:"  # noqa: E731
                            f"${costs.monthly[key]}${last}")
    nonres_shares = [costs.share[item["key"]] for item in objects if item.get("nonres")]
    # Лимит БРИДЖа считают оплаты до РнС; при объектах вне ДДУ общие затраты
    # берёт их кредит, и БРИДЖа нет (`common_on_object_loans`).
    before_permit = "+".join(
        f'SUMIF({cost_col("k")},"<"&{_pi("permit")},{cost_col(key)})'
        for key in ("purchase", "design_p", "design_rd") if key in costs.monthly) or "0"
    cells = {
        "permit": f"={costs.cells['permit']}",
        "rve": f"={costs.cells['rve']}",
        "own_funds": f"=MAX(0,{_p('own_funds')})*1000000",
        "approved": f"=MAX(0,{_p('pf_limit_approved')})*1000000",
        "limit": (f"=IF({_pi('approved')}>0,{_pi('approved')},"
                  f"ROUNDUP(SUM({col('p_draw')})/10000000,0)*10000000)"),
        "common_share": "=1-(" + _sum(nonres_shares) + ")",
        "bridge_limit": (f"=0" if nonres_shares else f"={before_permit}"),
        "bridge_fee": f"={_pi('bridge_limit')}*{_p('reservation_fee_pct')}/100",
        "reservation_fee": f"={_pi('limit')}*{_p('reservation_fee_pct')}/100",
        "pf_draw_total": f"=SUM({col('p_draw')})",
        "transferred": f"=SUM({col('transfer')})",
        "financing_cost": (f"=SUM({col('b_int')})+SUM({col('b_cap')})+{_pi('bridge_fee')}"
                           f"+SUM({col('p_int')})+SUM({col('p_cap')})+SUM({col('fee')})"
                           f"+{_pi('reservation_fee')}"),
        "llcr_num": "=0",
        "llcr_den": f"={_pi('pf_draw_total')}+{_pi('financing_cost')}",
        "llcr": f"=IF({_pi('llcr_den')}<>0,{_pi('llcr_num')}/{_pi('llcr_den')},0)",
        "peak_bridge": f"=MAX({col('b_bal')})",
        "peak_pf": f"=MAX({col('p_bal')})",
        "peak_escrow": f"=MAX({col('esc')})",
        "ending_pf": f"=INDEX({col('p_bal')},{months})",
        "rve_unpaid": (f"=IFERROR(INDEX({col('p_bal')},{_pi('rve')}+1),0)"),
        "vat_due": "=0",
    }
    for name, (row, label) in PF_CELLS.items():
        ws[f"A{row}"] = label
        ws[f"B{row}"] = cells[name]
        ws[f"B{row}"].number_format = "0.00" if name == "llcr" else MONEY
    for key, (letter, label) in PF_COLUMNS.items():
        ws[f"{letter}{FIRST_ROW - 1}"] = label
        ws[f"{letter}{FIRST_ROW - 1}"].font = BOLD
    permit, rve = _pi("permit"), _pi("rve")
    cap_mode = f'{_p("interest_mode")}="{INTEREST_MODES[0]}"'
    selling = f"({_p('marketing')}+{_p('selling')})/100*{_p('cost_mult')}"
    steps = inputs.steps
    for i in range(months):
        r = FIRST_ROW + i
        v = {name: f"{_pc(name)}{r}" for name in PF_COLUMNS}
        p = {name: f"{_pc(name)}{r - 1}" for name in PF_COLUMNS}
        first = i == 0
        prev = (lambda name: "0") if first else (lambda name: p[name])  # noqa: E731
        k = v["k"]
        cm = lambda key: f"'{COSTS_SHEET}'!{costs.monthly[key]}{r}"  # noqa: E731
        ws[k] = i
        ws[v["month"]] = f"=EDATE({_p('start')},{k})"
        ws[v["month"]].number_format = MONTH
        ws[v["key_rate"]] = f"=INDEX({credit['key_rate']},{i + 1})"
        ws[v["sales"]] = "=" + _sum([f"INDEX({ref['revenue']},{i + 1})" for ref in ddu_refs])
        ws[v["operating"]] = "=(" + _sum([f"INDEX({ref['base']},{i + 1})" for ref in ddu_refs]) + f")*{selling}"
        # Долговые затраты — `debt_capex` движка: CAPEX без ВРИ за счёт
        # капитала, без объектов вне ДДУ и без их доли общих затрат.
        ws[v["capex"]] = (f"=MAX(0,MAX(0,{cm('total')}-{cm('vri_equity')})-{cm('nonres_obj')})"
                          f"-{cm('common')}*(1-{_pi('common_share')})")
        ws[v["costs"]] = (f"={v['capex']}+{v['operating']}"
                          f"+IF(AND({k}={permit},{_p('interest_mode')}<>\"{INTEREST_MODES[0]}\"),{v['transfer']},0)")
        ws[v["esc"]] = (f"=IF({k}<{rve},{prev('esc')}+{v['sales']},IF({k}={rve},{prev('esc')},0))")
        ws[v["release"]] = f"=IF({k}={rve},{v['esc']},0)"
        need = f"MAX({v['costs']},0)"
        own_left = _pi("own_funds") if first else p["own_left"]
        ws[v["own_draw"]] = f"=IF({k}<{permit},MIN({own_left},{need}),0)"
        ws[v["own_left"]] = f"={own_left}-{v['own_draw']}"
        ws[v["b_draw"]] = f"=IF({k}<{permit},MAX(0,{need}-{v['own_draw']}),0)"
        ws[v["b_bal"]] = f"=IF({k}<{permit},{prev('b_bal')}+{v['b_draw']},0)"
        b_rate = f"({v['key_rate']}+{_p('bridge_spread')}/100)"
        b_cap_rate = f"({v['key_rate']}+{_p('bridge_cap_spread')}/100)"
        ws[v["b_int"]] = f"=IF(AND({k}<{permit},{v['b_bal']}>0),{v['b_bal']}*{b_rate}/12,0)"
        ws[v["b_cap"]] = f"=IF(AND({k}<{permit},{v['b_bal']}>0),{prev('b_pay')}*{b_cap_rate}/12,0)"
        ws[v["b_pay"]] = f"=IF({k}<{permit},{prev('b_pay')}+{v['b_int']}+{v['b_cap']},0)"
        ws[v["b_repay"]] = f"=IF({k}={permit},{prev('b_bal')},0)"
        ws[v["transfer"]] = f"=IF({k}={permit},{prev('b_pay')},0)"
        # ПФ: с РнС выбирает все расходы; линия, рассчитавшаяся в РВЭ, закрыта.
        open_line = f"OR({k}<={rve},{prev('settled')}=0)"
        draw = f"IF({open_line},{need},0)+{v['b_repay']}"
        ws[v["p_draw"]] = (f"=IF({k}>={permit},IF({_pi('approved')}>0,"
                           f"MIN({draw},MAX({_pi('approved')}-{prev('p_bal')},0)),{draw}),0)")
        ws[v["p_bal_after"]] = f"={prev('p_bal')}+{v['p_draw']}"
        ws[v["cov"]] = f"=IF({v['p_bal_after']}>0,{v['esc']}/{v['p_bal_after']},0)"
        base_special = f"{_p('pf_special')}/100"
        ladder = base_special
        for edge_ref, rate_ref in steps:
            ladder = (f"IF(AND({edge_ref}>0,{v['cov']}+0.000000001>={edge_ref}/100),"
                      f"{rate_ref}/100,{ladder})")
        ws[v["special"]] = f"={ladder}"
        weight = f"MIN({v['cov']},1)"
        ws[v["p_rate"]] = (f"=IF({k}>={permit},({v['key_rate']}+{_p('pf_spread')}/100)*(1-{weight})"
                           f"+{v['special']}*{weight},0)")
        live = f"AND({k}>={permit},{v['p_bal_after']}>0)"
        ws[v["p_pay_open"]] = (f"={prev('p_pay')}+IF(AND({k}={permit},{cap_mode}),{v['transfer']},0)")
        ws[v["p_int"]] = f"=IF({live},{v['p_bal_after']}*{v['p_rate']}/12,0)"
        ws[v["p_cap"]] = f"=IF({live},{v['p_pay_open']}*{v['p_rate']}/12,0)"
        ws[v["fee"]] = (f"=IF(AND({live},{k}<{rve},{_pi('limit')}>0),"
                        f"MAX({_pi('limit')}-{v['p_bal_after']},0)*{_p('limit_fee_pct')}/100/12,0)")
        accrued = f"{v['p_pay_open']}+{v['p_int']}+{v['p_cap']}+{v['fee']}"
        ws[v["pay_out"]] = f"=IF(AND({k}>={permit},{k}>={rve}),{accrued},0)"
        ws[v["p_pay"]] = f"=IF(AND({k}>={permit},{k}>={rve}),0,{accrued})"
        ws[v["avail"]] = f"=IF({k}={rve},{v['release']},IF({k}>{rve},{v['sales']},0))"
        ws[v["p_repay"]] = (f"=IF(AND({k}>={permit},{v['p_bal_after']}>0),"
                            f"MIN({v['avail']},{v['p_bal_after']}),0)")
        ws[v["p_bal"]] = f"={v['p_bal_after']}-{v['p_repay']}"
        ws[v["settled"]] = (f"=IF({k}={rve},IF(AND({k}>={permit},{v['p_bal']}<=0),1,0),"
                            f"IF({k}>{rve},{prev('settled')},0))")
        ws[v["deduct"]] = (f"={v['b_int']}+{v['b_cap']}+{v['p_int']}+{v['p_cap']}+{v['fee']}"
                           f"+IF({k}=0,{_pi('bridge_fee')},0)+IF({k}={permit},{_pi('reservation_fee')},0)")
        ws[v["vat"]] = f"=IF({k}={rve},{_pi('vat_due')},0)"
        fees = f"(IF({k}=0,{_pi('bridge_fee')},0)+IF({k}={permit},{_pi('reservation_fee')},0))"
        # Поток проекта движка: выручка по ДДУ в месяц продажи и расходы, до
        # финансирования — без процентов и комиссий (CAPEX — у листа
        # «Денежный поток»).
        ws[v["project"]] = f"={v['sales']}-{v['operating']}-{v['vat']}"
        # Поток капитала: деньги продаж — с раскрытия эскроу; выборки долга
        # покрывают затраты, погашения уходят банку.
        ws[v["equity"]] = (f"=IF({k}<{rve},0,{v['sales']}+{v['release']})-{v['operating']}"
                           f"-{v['pay_out']}+{v['b_draw']}+{v['p_draw']}-{v['b_repay']}"
                           f"-{v['p_repay']}-{fees}-{v['vat']}")
        for name in PF_COLUMNS:
            if name in ("cov", "special", "p_rate", "key_rate"):
                ws[v[name]].number_format = "0.0000"
            elif name not in ("k", "month", "settled"):
                ws[v[name]].number_format = MONEY
    ws.column_dimensions["A"].width = 64
    for i in range(2, len(PF_COLUMNS) + 1):
        ws.column_dimensions[get_column_letter(i)].width = 16
    return {name: f"'{PF_SHEET}'!{col(name)}" for name in PF_COLUMNS}


# --- «Отчёт» -----------------------------------------------------------------

def _report_sheet(book: Workbook, spec: dict[str, Any], objects: list[dict[str, Any]],
                  objects_refs: list[dict[str, str]], costs: _Costs, cash: dict[str, str],
                  taxes: dict[str, str], months: int, pf: dict[str, str] | None = None
                  ) -> dict[str, str]:
    ws = book.create_sheet(REPORT_SHEET)
    ws["A1"] = f"Отчёт нежилого проекта{(' · ' + spec['name']) if spec.get('name') else ''}"
    ws["A1"].font = BOLD
    ws["A2"] = "Все числа — ссылки на листы модели."
    out: dict[str, str] = {}
    row = 4

    def section(title: str) -> None:
        nonlocal row
        row += 1
        ws[f"A{row}"] = title
        ws[f"A{row}"].font = BOLD
        row += 1

    def line(key: str, label: str, formula: str, fmt: str = MONEY) -> None:
        nonlocal row
        ws[f"A{row}"] = label
        ws[f"B{row}"] = formula
        ws[f"B{row}"].number_format = fmt
        out[key] = f"'{REPORT_SHEET}'!$B${row}"
        row += 1

    def ref(key: str) -> str:
        return out[key].split("!")[1]

    # IRR и NPV КАПИТАЛА не числа, если ПФ к концу проекта не погашен (решение
    # владельца 2 от 29.09.2026): порог и подпись — движка, из `spec`. Книга
    # без ПФ (объекты на своих кредитах) долга ПФ не несёт — правила нет.
    # NPV проекта (поток до финансирования) — числом всегда (владелец, 17.10.2026).
    na = spec.get("equity_na") or {}

    def na_when_unrepaid(formula: str) -> str:
        if not pf or not na:
            return f"={formula}"
        return (f'=IF({_pi("ending_pf")}>{float(na["min_ending_pf_rub"]):g},'
                f'"{na["label"]}",{formula})')

    section("Экономика проекта")
    line("revenue", "Выручка проекта", f"=SUM({cash['revenue']})")
    line("capex", "Затраты (CAPEX, с НДС)", f"=SUM({cash['capex']})")
    line("nonres_costs", "Объекты вне ДДУ: продажи ДКП, эксплуатация, налог на имущество, выход",
         f"=SUM({cash['costs']})" + (f"-SUM({pf['operating']})" if pf else ""))
    if pf:
        # Маркетинг и продажи объектов по ДДУ — `commercial_costs` движка.
        line("commercial_costs", "Объекты по ДДУ: маркетинг и продажи", f"=SUM({pf['operating']})")
    line("ebitda", "EBITDA", f"={ref('revenue')}-{ref('capex')}-{ref('nonres_costs')}"
         + (f"-{ref('commercial_costs')}" if pf else ""))
    # Проценты ПФ в потоке — уплаченные; в расходах — начисленные
    # (`financing_cost` движка): их итог на листе «ПФ».
    line("financing_cost", "Проценты и комиссии", f"=SUM({cash['interest']})+SUM({cash['fee']})" + (
        f"-SUM({pf['pay_out']})-{_pi('bridge_fee')}-{_pi('reservation_fee')}+{_pi('financing_cost')}"
        if pf else ""))
    line("profit_before_tax", "Прибыль до налога", f"={ref('ebitda')}-{ref('financing_cost')}")
    line("profit_tax", "Налог на прибыль", f"=SUM({taxes['tx_tax']})")
    line("vat", "НДС к уплате", f"=SUM({taxes['vat_paid']})")
    line("net_profit", "Чистая прибыль", f"={ref('profit_before_tax')}-{ref('profit_tax')}-{ref('vat')}")
    line("npv", "NPV проекта", f"=SUM({cash['project_disc']})")
    # IRR капитала — числом, только если своих денег не меньше порога движка
    # от полных расходов (владелец, 16.10.2026): при почти нулевых своих
    # деньгах он выходит сотнями процентов и доходностью не является. Свои
    # деньги — оттоки потока капитала, как `equity_contributed` движка.
    irr_formula = f'IFERROR((1+IRR({cash["equity"]},0.01))^12-1,"")'
    min_share = (spec.get("equity_irr") or {}).get("min_share")
    if min_share is not None:
        full_cost = "+".join([ref("capex"), ref("nonres_costs")]
                             + ([ref("commercial_costs")] if pf else [])
                             + [ref("financing_cost"), ref("profit_tax"), ref("vat")])
        own = f'(-SUMIF({cash["equity"]},"<0"))'
        irr_formula = (f'IF(AND(({full_cost})>0,{own}<{float(min_share):g}*({full_cost})),'
                       f'"не показателен: своих денег "&ROUND({own}/1000000,0)&" млн ₽ ("'
                       f'&SUBSTITUTE(ROUND({own}/({full_cost})*100,1),".",",")&"% затрат)",'
                       f'{irr_formula})')
    line("irr_equity", "IRR собственного капитала", na_when_unrepaid(irr_formula), "0.00%")

    # Горизонт: строк в помесячных таблицах столько, сколько месяцев у движка.
    # Правка вводных, уводящая проект дальше, должна краснить сверку, а не
    # обрезать поток молча.
    horizons = [ref["horizon"] for ref in objects_refs]
    line("horizon_months", f"Горизонт модели, мес. (строк в помесячных таблицах: {months})",
         f"=MAX({costs.cells['rve']}+MAX(INT({_p('residual_months')})+3,12)"
         + "".join(f",{h}" for h in horizons) + f",{costs.cells['plan_last_k']})+1", "0")

    section("Как сложены итоги")
    line("identity_ebitda", "EBITDA − проценты − прибыль до налога (должно быть 0)",
         f"={ref('ebitda')}-{ref('financing_cost')}-{ref('profit_before_tax')}")
    ws[f"C{row - 1}"] = "EBITDA = выручка − CAPEX − расходы объектов; без процентов, налога на прибыль и НДС"
    identity_row = row
    row += 0

    section("Структура расходов")
    article_rows = []
    for key, label in costs.labels.items():
        if key.endswith(":garage"):
            continue
        if key.endswith(":building"):
            base = key.split(":")[0]
            formula = f"={costs.amounts[key]}+{costs.amounts[base + ':garage']}"
            label = label.replace(": здание", "")
        else:
            formula = f"={costs.amounts[key]}"
        article_rows.append(row)
        line(f"article:{key.split(':')[0]}", label, formula)
    line("structure_nonres", "Объекты вне ДДУ: продажи ДКП, эксплуатация, налог на имущество, выход",
         f"={ref('nonres_costs')}")
    if pf:
        line("structure_commercial", "Объекты по ДДУ: маркетинг и продажи", f"={ref('commercial_costs')}")
    line("structure_financing", "Проценты и комиссии", f"={ref('financing_cost')}")
    line("structure_tax", "Налог на прибыль", f"={ref('profit_tax')}")
    line("structure_vat", "НДС", f"={ref('vat')}")
    first_article = article_rows[0] if article_rows else row
    line("total_expenses", "Расходы всего", f"=SUM(B{first_article}:B{row - 1})")
    ws[f"C{row - 1}"] = ("CAPEX + расходы объектов + проценты и комиссии + налог на прибыль + НДС; "
                         "поэтому выручка − расходы всего = чистая прибыль, а не EBITDA")
    ws[f"A{identity_row}"] = "Выручка − расходы всего − чистая прибыль (должно быть 0)"
    ws[f"B{identity_row}"] = f"={ref('revenue')}-{ref('total_expenses')}-{ref('net_profit')}"
    ws[f"B{identity_row}"].number_format = MONEY
    out["identity_total"] = f"'{REPORT_SHEET}'!$B${identity_row}"

    section("Финансирование объектов")
    ws[f"A{row}"] = "Показатель"
    for i, item in enumerate(objects_refs):
        ws.cell(row, 2 + i, item["title"]).font = BOLD
    row += 1
    financing_rows = (
        ("loan_draw", "Выборка кредита"), ("draw_object", "в т.ч. на стройку объекта"),
        ("draw_common", "в т.ч. на общие затраты проекта"), ("limit", "Лимит НКЛ объекта"),
        ("limit_land", "в т.ч. НКЛ на участок и проект (до РнС)"),
        ("limit_build", "в т.ч. строительная НКЛ (с РнС)"),
        ("reservation_fee", "Плата за резервирование лимита НКЛ"),
        ("commitment_fee", "Плата за невыбранный лимит НКЛ"), ("issue_fee", "Комиссия за выдачу"),
        ("interest_capitalized", "Проценты до ввода — капитализированы"),
        ("debt_at_commissioning", "Долг на ввод"), ("loan_peak", "Пик долга"),
        ("balloon_planned", "Баллон по графику"), ("repaid_scheduled", "Погашено плановым телом"),
        ("balloon_paid", "Баллон погашен в срок"), ("repaid_from_cash", "Погашено из выручки ДКП / NOI"),
        ("repaid_at_exit", "Остаток погашен при выходе"), ("loan_repayment", "Погашено всего"),
        ("interest_paid", "Проценты после ввода — уплачены"),
        ("dscr_min", "DSCR — минимум по годам"), ("icr_min", "ICR — минимум"),
    )
    for key, label in financing_rows:
        ws[f"A{row}"] = label
        for i, item in enumerate(objects_refs):
            ws.cell(row, 2 + i, f"={item[key]}").number_format = (
                "0.00" if key in ("dscr_min", "icr_min") else MONEY)
        row += 1
    line("loan_peak", "Кредит объектов — пик (сумма пиков объектов)",
         "=" + _sum([item["loan_peak"] for item in objects_refs]))
    dscr = [item["dscr_min"] for item in objects_refs]
    if pf:
        # Проект с ПФ меряет долг LLCR (`debt_metric` движка).
        line("debt_metric", "LLCR проектного финансирования", f"={_pi('llcr')}", "0.00")
    else:
        line("debt_metric", "DSCR кредита объектов — минимум по годам и объектам",
             (f'=IF(COUNT({",".join(dscr)})>0,MIN({",".join(dscr)}),"")'
              if dscr else '=""'), "0.00")
        ws[f"C{row - 1}"] = "пусто — кредит гасится выручкой ДКП, платежа из NOI нет"
    if pf:
        section("Проектное финансирование объектов по ДДУ")
        for key in ("pf_draw_total", "peak_bridge", "peak_pf", "peak_escrow", "limit",
                    "bridge_fee", "reservation_fee", "transferred", "financing_cost", "llcr",
                    "ending_pf", "rve_unpaid", "vat_due"):
            line("pf:" + key, PF_CELLS[key][1], f"={_pi(key)}", "0.00" if key == "llcr" else MONEY)

    section("Собственное участие")
    line("equity_capex", "Затраты проекта (CAPEX, с НДС)", f"={ref('capex')}")
    line("equity_draw", "Кредит объектов — выборка", "=" + _sum([item["loan_draw"] for item in objects_refs]))
    line("in_build", "Собственные средства до ввода", f"=SUM({cash['in_build']})")
    line("contributed", "Вложено всего", f"=SUM({cash['contributed']})")
    line("in_exit", "в т.ч. при выходе — довнесение на погашение долга", f"=SUM({cash['in_exit']})")
    line("in_operation", "в т.ч. после ввода — дефицит",
         f"={ref('contributed')}-{ref('in_build')}-{ref('in_exit')}")
    line("equity_peak", "Пик потребности в собственных средствах", f"=-MIN(0,MIN({cash['cum']}))")
    line("returned", "Получено всего (деньгами)", f"=SUM({cash['returned']})")
    line("out_vat", "в т.ч. возмещение НДС стройки", f"=SUM({cash['out_vat']})")
    line("out_exit", "в т.ч. при выходе — после погашения долга", f"=SUM({cash['out_exit']})")
    line("out_running", "в т.ч. по ходу — продажи ДКП, аренда после долга и налогов",
         f"={ref('returned')}-{ref('out_vat')}-{ref('out_exit')}")
    line("residual", "Оценка удержанного объекта — не деньги", f"=SUM({cash['residual']})")
    line("equity_net", "Чистый денежный результат (получено − вложено)",
         f"={ref('returned')}-{ref('contributed')}")
    line("equity_multiple", "Мультипликатор капитала (получено / вложено)",
         f'=IF({ref("contributed")}>0,{ref("returned")}/{ref("contributed")},"")', "0.00")
    line("equity_npv", "NPV собственного капитала", na_when_unrepaid(f"SUM({cash['equity_disc']})"))
    line("equity_irr_cash", "IRR собственного капитала — только деньги",
         na_when_unrepaid(f'IF({ref("residual")}>0,IFERROR((1+IRR({cash["cash"]},0.01))^12-1,""),"")'),
         "0.00%")
    last_neg = f"MAX({cash['neg_k']})"
    line("payback", "Окупаемость собственных средств",
         f'=IF(AND({last_neg}>=0,{last_neg}+1<={months - 1}),EDATE({_p("start")},{last_neg}+1),'
         f'"не окупается за горизонт проекта")', MONTH)

    section("ОПУ по годам (налоговая база проекта)")
    head = ("Год", "Маржа объектов", "Общие затраты (в РВЭ)", "Проценты и комиссии",
            "Прибыль до налога", "Налог на прибыль", "Прибыль после налога")
    for i, label in enumerate(head):
        ws.cell(row, 1 + i, label).font = BOLD
    row += 1
    for year in sorted({_month(m).year for m in spec["months"]}):
        ws.cell(row, 1, year)
        for i, key in enumerate(("objects_margin", "common", "financing"), start=2):
            ws.cell(row, i, f"=SUMIF({taxes['cal_year']},A{row},{taxes[key]})").number_format = MONEY
        ws.cell(row, 5, f"=B{row}+C{row}-D{row}").number_format = MONEY
        ws.cell(row, 6, f"=SUMIF({taxes['cal_year']},A{row},{taxes['tx_tax']})").number_format = MONEY
        ws.cell(row, 7, f"=E{row}-F{row}").number_format = MONEY
        row += 1
    ws.cell(row, 1, "Чувствительность: сценарные множители выручки и затрат — на «Вводных»; "
                    "таблицы данных Excel книга не строит").font = Font(italic=True)
    row += 1

    section("По годам")
    years = sorted({_month(m).year for m in spec["months"]})
    head = ("Год", "Выручка", "CAPEX", "Расходы объектов", "Проценты и комиссии",
            "НДС", "Налог на прибыль", "Поток проекта", "Вложено собственником",
            "Получено собственником", "Чистый поток собственника", "Накопленный на конец года")
    for i, label in enumerate(head):
        ws.cell(row, 1 + i, label).font = BOLD
    row += 1
    out["years_first"] = str(row)
    for year in years:
        ws.cell(row, 1, year)
        series = (cash["revenue"], cash["capex"], cash["costs"], None, cash["vat"], cash["tax"],
                  cash["project"], cash["contributed"], cash["returned"])
        for i, rng in enumerate(series, start=2):
            if rng is None:
                formula = (f"=SUMIF({cash['cal_year']},A{row},{cash['interest']})"
                           f"+SUMIF({cash['cal_year']},A{row},{cash['fee']})")
            else:
                formula = f"=SUMIF({cash['cal_year']},A{row},{rng})"
            ws.cell(row, i, formula).number_format = MONEY
        ws.cell(row, 11, f"=J{row}-I{row}").number_format = MONEY
        ws.cell(row, 12, f"=K{row}" if year == years[0] else f"=L{row - 1}+K{row}").number_format = MONEY
        out[f"year:{year}"] = str(row)
        row += 1
    ws.column_dimensions["A"].width = 70
    for letter in "BCDEFGHIJKL":
        ws.column_dimensions[letter].width = 18
    return out


# --- «Сверка» ----------------------------------------------------------------

TOLERANCE_SHARE = 1e-6
RATIO_TOLERANCE = 1e-6


def _engine_number(value: Any) -> float | str:
    if value is None or isinstance(value, str):
        return ""
    return float(value)


def _check_rows(result: dict[str, Any], spec: dict[str, Any], objects_refs: list[dict[str, str]],
                engine_objects: list[dict[str, Any]], costs: _Costs, report: dict[str, str],
                ddu_refs: list[dict[str, str]] | None = None
                ) -> list[tuple[str, str, str, Any, bool]]:
    """(раздел, показатель, ссылка книги, число движка, доля ли это)."""
    summary = result.get("summary") or {}
    finance = result.get("finance") or {}
    equity = (result.get("report") or {}).get("equity_participation") or {}
    verdict = summary.get("equity_returns") or {}
    na_label = (verdict.get("label") or None) if verdict.get("status") == "debt_unrepaid" else None
    rows: list[tuple[str, str, str, Any, bool]] = []
    project = (
        ("revenue", "Выручка проекта", summary.get("revenue"), False),
        ("capex", "CAPEX", summary.get("capex"), False),
        ("nonres_costs", "Расходы объектов вне ДДУ", finance.get("nonres_costs"), False),
        ("ebitda", "EBITDA", summary.get("ebitda"), False),
        ("financing_cost", "Проценты и комиссии", summary.get("financing_cost"), False),
        ("profit_before_tax", "Прибыль до налога", summary.get("profit_before_tax"), False),
        ("profit_tax", "Налог на прибыль", summary.get("profit_tax"), False),
        ("vat", "НДС", summary.get("vat"), False),
        ("net_profit", "Чистая прибыль", summary.get("net_profit"), False),
        ("total_expenses", "Расходы всего", summary.get("total_expenses"), False),
        ("npv", "NPV проекта", summary.get("npv"), False),
        # Подпись движка вместо IRR капитала — «N/A — долг не погашен» или
        # «не показателен: своих денег …» (`equity_irr`).
        ("irr_equity", "IRR собственного капитала",
         summary.get("irr_equity") if na_label is None and _irr_label(summary) is None
         else (na_label or _irr_label(summary)), True),
        ("loan_peak", "Кредит объектов — пик",
         sum(float((o.get("totals") or {}).get("loan_peak") or 0.0) for o in engine_objects), False),
        ("identity_total", "Выручка − расходы всего − чистая прибыль", 0.0, False),
        ("identity_ebitda", "EBITDA − проценты − прибыль до налога", 0.0, False),
        # Горизонт — месяцы расчёта движка (`finance.rows`); ряд потоков бывает
        # длиннее на хвост уплаты налога (доплата года в марте следующего).
        ("horizon_months", "Горизонт модели, мес.", len(finance.get("rows") or []),
         True),
        ("debt_metric", "LLCR" if ddu_refs else "DSCR кредита объектов — минимум",
         (((result.get("report") or {}).get("layout") or {}).get("debt_metric") or {}).get("value"), True),
    )
    for key, label, value, ratio in project:
        rows.append(("Проект", label, report[key], value, ratio))
    if ddu_refs:
        financing = (result.get("report") or {}).get("financing") or {}
        for key, value in (
                ("pf_draw_total", finance.get("pf_draw_total")),
                ("peak_bridge", finance.get("peak_bridge")),
                ("peak_pf", finance.get("peak_pf")),
                ("peak_escrow", finance.get("peak_escrow")),
                ("limit", finance.get("pf_limit")),
                ("bridge_fee", finance.get("bridge_fee")),
                ("reservation_fee", finance.get("pf_reservation_fee")),
                ("transferred", finance.get("transferred_bridge_interest")),
                ("ending_pf", finance.get("ending_pf")),
                ("rve_unpaid", financing.get("rve_pf_shortfall", finance.get("rve_pf_shortfall")))):
            rows.append(("ПФ", PF_CELLS[key][1], report["pf:" + key], value, False))
    capex = result.get("capex") or {}
    for key in costs.labels:
        base = key.split(":")[0]
        # Статьи, которой в смете движка нет (покупка участка идёт прямо в
        # месячный CAPEX), сверяет итог CAPEX.
        if key.endswith(":garage") or base not in capex:
            continue
        rows.append(("Смета", costs.labels[key].replace(": здание", ""), report[f"article:{base}"],
                     capex.get(base), False))
    parts = equity.get("parts") or {}
    equity_rows = (
        ("contributed", "Вложено всего", equity.get("contributed")),
        ("returned", "Получено всего", equity.get("returned")),
        ("in_build", "Вложено до ввода", parts.get("in_build")),
        ("in_operation", "Вложено после ввода", parts.get("in_operation")),
        ("in_exit", "Вложено при выходе", parts.get("in_exit")),
        ("out_vat", "Возмещение НДС стройки", parts.get("out_vat")),
        ("out_running", "Получено по ходу", parts.get("out_running")),
        ("out_exit", "Получено при выходе", parts.get("out_exit")),
        ("residual", "Оценка удержания", equity.get("residual")),
    )
    # Отчёт о собственном участии движок строит у проекта с объектами вне
    # ДДУ; у проекта только с ДДУ его нет — сверять не с чем.
    for key, label, value in equity_rows if equity else ():
        rows.append(("Собственное участие", label, report[key], value, False))
    for line in equity.get("rows") or []:
        label = str(line.get("label") or "")
        if label.startswith("Пик потребности"):
            rows.append(("Собственное участие", "Пик потребности", report["equity_peak"],
                         line.get("value"), False))
        elif label.startswith("NPV собственного капитала"):
            rows.append(("Собственное участие", "NPV собственного капитала", report["equity_npv"],
                         line.get("value"), False))
        elif label.startswith("IRR собственного капитала — только деньги") and line.get("unit") == "pct":
            rows.append(("Собственное участие", "IRR — только деньги", report["equity_irr_cash"],
                         line.get("value"), True))
    for year in equity.get("years") or []:
        r = report.get(f"year:{year.get('year')}")
        if r:
            rows.append(("По годам", f"{year['year']}: вложено собственником", f"'{REPORT_SHEET}'!$I${r}",
                         year.get("contributed"), False))
            rows.append(("По годам", f"{year['year']}: получено собственником", f"'{REPORT_SHEET}'!$J${r}",
                         year.get("returned"), False))
    taxes = (result.get("cashflow") or {}).get("profit_tax") or []
    tax_months = (result.get("cashflow") or {}).get("months") or []
    by_year: dict[int, float] = {}
    for month, value in zip(tax_months, taxes):
        by_year[_month(month).year] = by_year.get(_month(month).year, 0.0) + float(value or 0.0)
    for year, value in sorted(by_year.items()):
        r = report.get(f"year:{year}")
        if r:
            rows.append(("По годам", f"{year}: налог на прибыль", f"'{REPORT_SHEET}'!$G${r}", value, False))
    object_checks = (
        ("revenue", "totals"), ("rent_revenue", "totals"), ("sale_revenue", "totals"),
        ("exit_revenue", "totals"), ("residual_value", "totals"), ("opex", "totals"),
        ("property_tax", "totals"), ("selling_cost", "totals"), ("exit_cost", "totals"),
        ("vat_paid", "totals"), ("vat_charged", "totals"), ("loan_draw", "totals"),
        ("loan_interest", "totals"), ("loan_fee", "totals"), ("loan_repayment", "totals"),
        ("loan_peak", "totals"), ("noi", "totals"), ("tax_margin", "totals"), ("capex", "totals"),
        ("exit_value", "kpi"), ("stabilized_noi", "kpi"), ("dscr_min", "kpi"), ("icr_min", "kpi"),
        ("draw_object", "financing"), ("draw_common", "financing"),
        ("issue_fee", "financing:fee"), ("reservation_fee", "financing"), ("limit", "financing"),
        ("commitment_fee", "financing"), ("limit_land", "financing"), ("limit_build", "financing"),
        ("interest_capitalized", "financing"), ("interest_paid", "financing"),
        ("debt_at_commissioning", "financing"), ("rate_at_commissioning", "financing"),
        ("avg_rate", "financing"), ("balloon_planned", "financing"), ("balloon_paid", "financing"),
        ("repaid_scheduled", "financing"), ("repaid_at_exit", "financing"),
        ("repaid_from_cash", "financing"), ("avg_payment", "financing"),
        ("common_capex", "result"), ("equity_invested", "result"), ("equity_peak", "result"),
        ("common_recognized", "result"), ("financing_cost", "result"),
        ("profit_before_tax", "result"), ("profit_tax", "result"), ("equity_cash", "result"),
        ("irr", "result"), ("irr_cash", "result"), ("npv", "result"), ("payback_months", "result"),
    )
    ratios = {"dscr_min", "icr_min", "rate_at_commissioning", "avg_rate", "irr", "irr_cash"}
    for ref, item in zip(objects_refs, engine_objects):
        for key, where in object_checks:
            where, _, engine_key = where.partition(":")
            rows.append((item.get("title") or item.get("key") or "", TOTAL_LABELS.get(key, key), ref[key],
                         (item.get(where) or {}).get(engine_key or key), key in ratios))
    return rows


def _check_sheet(book: Workbook, rows: list[tuple[str, str, str, Any, bool]],
                 missing: list[str]) -> None:
    ws = book.create_sheet(CHECK_SHEET)
    ws["A1"] = "Сверка формул книги с авторитетным расчётом движка"
    ws["A1"].font = BOLD
    ws.append(["Раздел", "Показатель", "Книга", "Движок", "Разница", "Допуск", "Итог"])
    for section, label, ref, engine, ratio in rows:
        r = ws.max_row + 1
        ws[f"A{r}"] = section
        ws[f"B{r}"] = label
        ws[f"C{r}"] = f"={ref}"
        ws[f"D{r}"] = _engine_number(engine)
        ws[f"D{r}"].fill = ENGINE_FILL
        if isinstance(engine, str) and engine:
            # Движок дал не число, а вердикт («N/A — долг не погашен»): книга
            # обязана сказать то же слово, сравнивается текст.
            ws[f"D{r}"] = engine
            ws[f"G{r}"] = f'=IF(C{r}=D{r},"сходится","РАСХОЖДЕНИЕ")'
            continue
        ws[f"E{r}"] = f'=IF(OR(D{r}="",C{r}=""),"",C{r}-D{r})'
        tolerance = (f"{RATIO_TOLERANCE}" if ratio else f"MAX(1,ABS(D{r})*{TOLERANCE_SHARE})")
        ws[f"F{r}"] = f'=IF(D{r}="","",{tolerance})'
        # Пусто у обоих — «показателя нет» с обеих сторон; пусто у одного —
        # расхождение: книга посчитала то, чего движок не дал, или наоборот.
        ws[f"G{r}"] = (f'=IF(AND(D{r}="",C{r}=""),"—",IF(OR(D{r}="",C{r}=""),"РАСХОЖДЕНИЕ",'
                       f'IF(ABS(E{r})<=F{r},"сходится","РАСХОЖДЕНИЕ")))')
        for column in "CDE":
            ws[f"{column}{r}"].number_format = "0.000000" if ratio else MONEY
    last = ws.max_row
    ws["I2"] = "Вердикт"
    ws["I2"].font = BOLD
    ws["I3"] = f'=IF(COUNTIF(G3:G{last},"РАСХОЖДЕНИЕ")=0,"{PASSED}","{FAILED}")'
    ws["I4"] = "Расхождений"
    ws["J4"] = f'=COUNTIF(G3:G{last},"РАСХОЖДЕНИЕ")'
    ws["I6"] = "Формулы нет — значения движка (missing)"
    ws["I6"].font = BOLD
    for i, text in enumerate(missing or ["нет"], start=7):
        ws[f"I{i}"] = text
    for letter, width in (("A", 26), ("B", 44), ("C", 20), ("D", 20), ("E", 14), ("F", 12),
                          ("G", 14), ("I", 60)):
        ws.column_dimensions[letter].width = width


# --- сборка ------------------------------------------------------------------

def build(result: dict[str, Any], spec: dict[str, Any]) -> bytes:
    """Книга нежилого проекта: модель формулами от «Вводных» и сверка с движком.

    `result` — авторитетный расчёт (`consolidated`), `spec` — вводные, которые
    движок читал (`main_legacy.nonres_book_spec`).
    """
    finance = result.get("finance") or {}
    engine_objects = [o for o in (finance.get("nonres") or {}).get("objects") or []]
    by_key = {o.get("key"): o for o in engine_objects}
    months = len(spec["months"])
    objects = []
    for i, item in enumerate(spec.get("objects") or []):
        objects.append({**item, "column": get_column_letter(3 + i)})
    book = Workbook()
    inputs = _inputs_sheet(book, spec, objects)
    costs = _costs_sheet(book, spec, objects, inputs, months)
    nonres = [item for item in objects if item.get("nonres")]
    objects_refs: list[dict[str, str]] = []
    ordered_engine: list[dict[str, Any]] = []
    for number, item in enumerate(nonres, start=1):
        refs = _object_sheet(book, item, f"Объект {number}", costs, months)
        refs["title"] = item["title"]
        objects_refs.append(refs)
        ordered_engine.append(by_key.get(item["key"]) or {})
    ddu = [item for item in objects if item.get("ddu")]
    ddu_refs = [_ddu_sheet(book, item, f"ДДУ {number}", costs, inputs, months)
                for number, item in enumerate(ddu, start=1)]
    credit = _credit_sheet(book, objects_refs, inputs, months)
    pf = _pf_sheet(book, ddu_refs, objects, costs, credit, inputs, months) if ddu_refs else None
    taxes = _tax_sheet(book, objects_refs, objects, costs, months, ddu_refs, pf)
    if pf:
        _pf_link_taxes(book, pf, ddu_refs, objects, costs, taxes, objects_refs, months)
    cash = _cash_sheet(book, objects_refs, costs, taxes, months, pf)
    report = _report_sheet(book, spec, objects, objects_refs, costs, cash, taxes, months, pf)
    rows = _check_rows(result, spec, objects_refs, ordered_engine, costs, report, ddu_refs)
    _check_sheet(book, rows, list(spec.get("missing") or []))
    stream = io.BytesIO()
    book.save(stream)
    return stream.getvalue()
