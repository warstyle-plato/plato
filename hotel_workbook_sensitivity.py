"""Чувствительность: те же формулы месячной модели, связанные с вводными.

Без таблиц данных Excel и без линейной аппроксимации: каждый вариант
пересчитывает кредит и налог, а не умножает готовый IRR.
"""
from __future__ import annotations

from openpyxl.styles import Alignment
from openpyxl.utils.cell import coordinate_from_string

import hotel_workbook_capex as capex

SHEET = "Чувствительность"
FACTORS = ("ADR", "Целевая загрузка", "Ставка СМР", "FF&E на номер",
           "Стоимость выхода", "Ставка кредита")
# Факторы, которые меняют смету (их вводные читает лист «Затраты»).
CAPEX_FACTORS = (2, 3)


def build(book, rows, hw):
    ws = book.create_sheet(SHEET, book.index(book[hw.CHECK_SHEET]))
    ws["A1"] = "Чувствительность гостиницы"
    ws["A2"] = "Один фактор за раз; кредит, налог и потоки пересчитываются формулами. Вводные — общие с базой."
    for col, label in zip("ABCDE", ("Фактор", "Изменение", "NPV проекта, ₽", "IRR проекта", "IRR капитала")):
        ws[f"{col}3"] = label
    ws["A4"] = "Базовый проект"
    ws["B4"] = 0
    ws["C4"] = f"='{hw.TOTALS_SHEET}'!B4"
    ws["D4"] = f"='{hw.TOTALS_SHEET}'!B5"
    ws["E4"] = f"='{hw.TOTALS_SHEET}'!B7"
    clones = []
    # Вводные и снимок движка копия берёт ссылкой на базу; узнаются по заливке.
    marked = {book._fills.index(fill) for fill in (hw.INPUT_FILL, hw.ENGINE_FILL)
              if fill in book._fills}
    for factor_index, label in enumerate(FACTORS):
        for delta_index, delta in enumerate((-0.1, 0.1)):
            row = 5 + factor_index * 2 + delta_index
            factor = f"(1+'{SHEET}'!$B${row})"
            ws[f"A{row}"], ws[f"B{row}"] = label, delta
            ws[f"B{row}"].fill = hw.INPUT_FILL
            names = {hw.INPUTS_SHEET: f"С{row}_Вводные", hw.CALC_SHEET: f"С{row}_Расчёт"}
            # Своя смета — только варианту, который её меняет (ставка СМР,
            # FF&E). Остальные читают смету базы: её вводные у них те же, а
            # копия на каждый вариант удваивала выгрузку.
            if factor_index in CAPEX_FACTORS:
                names[capex.SHEET] = f"С{row}_Затраты"
            for original, name in names.items():
                # Копия скрыта и считается в Excel: ей нужны значения и
                # формулы, а не оформление — копирование стилей стоило секунды.
                clone = book.create_sheet(name)
                clone.sheet_state = "hidden"
                clones.append(name)
                for cells in book[original].iter_rows():
                    for cell in cells:
                        value = cell.value
                        if value is None:
                            continue
                        if cell.data_type == "f":
                            for old, new in names.items():
                                value = value.replace(f"'{old}'!", f"'{new}'!")
                        elif cell.has_style and cell._style.fillId in marked:
                            value = f"='{original}'!{cell.coordinate}"
                        clone.cell(row=cell.row, column=cell.column, value=value)
            inp = book[names[hw.INPUTS_SHEET]]
            if factor_index in (0, 1):
                key = "adr" if factor_index == 0 else "occ_target"
                target = f"C{hw._ROWS[key]}"
                expr = f"'{hw.INPUTS_SHEET}'!{target}*{factor}"
                inp[target] = "=" + (f"MIN(1,MAX(0,{expr}))" if key == "occ_target" else expr)
            elif factor_index in (2, 3):
                key = "hotel_cost_th_per_sqm" if factor_index == 2 else "hotel_ffe_th_per_key"
                target = f"C{capex.INPUT_ROWS[key]}"
                inp[target] = f"='{hw.INPUTS_SHEET}'!{target}*{factor}"
            elif factor_index == 4:
                cell = inp[f"C{hw._ROWS['exit_value']}"]
                cell.value = f"=({cell.value[1:]})*{factor}"
            else:
                for key in ("spread", "pref_margin"):
                    target = f"C{hw._ROWS[key]}"
                    inp[target] = f"='{hw.INPUTS_SHEET}'!{target}*{factor}"
                for r in range(120, 120 + rows):
                    inp[f"F{r}"] = f"='{hw.INPUTS_SHEET}'!F{r}*{factor}"
            calc = names[hw.CALC_SHEET]
            def col(key):
                return f"'{calc}'!{hw._c(key)}{hw.FIRST_ROW}:{hw._c(key)}{hw.FIRST_ROW+rows-1}"
            ws[f"C{row}"] = f"=SUM({col('disc_p')})"
            ws[f"D{row}"] = f'=IFERROR((1+IRR({col("project_cf")},0.01))^12-1,"")'
            ws[f"E{row}"] = f'=IFERROR((1+IRR({col("equity_cf")},0.01))^12-1,"")'
            # Эти служебные колонки не участвуют в результате сценария.
            for r in range(hw.FIRST_ROW, hw.FIRST_ROW + rows):
                for key in ("irr_p", "irr_e"):
                    book[calc][f"{hw._c(key)}{r}"] = "=0"
    for row in range(4, 5 + len(FACTORS) * 2):
        ws[f"B{row}"].number_format = '+0%;-0%;"база"'
        ws[f"C{row}"].number_format = hw.MONEY
        ws[f"D{row}"].number_format = ws[f"E{row}"].number_format = hw.PERCENT
    ws.column_dimensions["A"].width = 30
    for col in "BCDE":
        ws.column_dimensions[col].width = 22
    ws.freeze_panes = "C4"
    ws["A18"] = "Изменяйте отклонения в жёлтых ячейках. Здесь сравнение с базой, а не со старым снимком движка."
    ws["A19"] = EXCEL_ONLY_NOTE
    for row in (18, 19):
        ws.merge_cells(f"A{row}:E{row}")
        ws[f"A{row}"].alignment = Alignment(wrap_text=True)
        ws.row_dimensions[row].height = 36
    return clones


# Двенадцать полных копий месячной модели — сотни тысяч формул. Выгрузка
# их не вычисляет (минута на сервере на каждое скачивание), а оставляет
# Excel: книга пересчитывается при открытии. Вне Excel клетки вариантов
# пусты, и лист говорит это сам.
EXCEL_ONLY_NOTE = ("Варианты считает Excel при открытии файла. В просмотре без пересчёта "
                   "(телефон, браузер) клетки вариантов пусты — это не ошибка расчёта.")


def excel_only(clones):
    """Клетки, которые считает только Excel: копии модели и итоги вариантов."""
    clones = set(clones)
    last = 4 + len(FACTORS) * 2

    def test(sheet: str, coordinate: str) -> bool:
        if sheet in clones:
            return True
        if sheet != SHEET:
            return False
        column, row = coordinate_from_string(coordinate)
        return column in ("C", "D", "E") and 5 <= row <= last
    return test
