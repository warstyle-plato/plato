#!/usr/bin/env python3
"""Книга Excel против движка на наборе эталонных проектов — формулами книги.

Книга считается `xlsx_eval` (тем же вычислителем, что кладёт в неё сохранённые
значения), а не LibreOffice: тот в песочнице ненадёжен (`excel_book_parity.py`
остаётся для ручной сверки пересчётом офиса). Сверяются не только итоги, но и
то, из чего они складываются:

* жилая книга v4 и смешанный проект — строки ОТЧЕТ; по годам — выручка по
  договорам, раскрытие и остаток эскроу, расходы, налог (начислен и уплачен),
  проценты, выборка и погашение БРИДЖа и ПФ, долг на конец года, поток
  собственника и проектный поток; CAPEX по главам; NPV и IRR проекта,
  посчитанные функцией движка из потока книги; пик собственных денег;
* книга нежилого проекта (`office_workbook`: офис по ДДУ, офис без ДДУ, ТЦ,
  ФОК, паркинг) — каждая строка её листа «ПРОВЕРКИ»: итоги, смета, ПФ, налог
  и деньги собственника по годам, NPV, IRR, пик.

    python3 scripts/book_engine_parity.py                  # все эталоны
    python3 scripts/book_engine_parity.py --only default office_ddu
    python3 scripts/book_engine_parity.py --tamper         # подделка обязана краснеть
    python3 scripts/book_engine_parity.py --summary out.md # таблица для CI

Код выхода 1 — есть расхождения. Каждое пишется одной строкой:
«проект X, строка Y: книга A, движок B, разница C». Протокол — docs/book_engine_parity.md.
"""

from __future__ import annotations

import argparse
import copy
import io
import sys
import time
from collections import defaultdict
from dataclasses import dataclass, field
from datetime import date, datetime
from pathlib import Path
from typing import Any, Callable

ROOT = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(ROOT))
sys.path.insert(0, str(ROOT / "scripts"))

import golden_snapshot as gs  # noqa: E402
from excel_book_parity import NONRES_BOOKS, REPORT_ROWS, nonres_scenario  # noqa: E402

MLN = 1e6
# Допуск — тот же, что у сверки ОТЧЕТ: 0,01% числа, не меньше 1 тыс. ₽ (0,001 млн);
# у долей — 1e-6 абсолютно.
REL_TOL = 1e-4
ABS_TOL_MLN = 1e-3
ABS_TOL_RATIO = 1e-6


# --- эталонные проекты ---------------------------------------------------------
# Вид книги каждого — для подписи; какую книгу собрать, решает сам движок
# (`build_project_workbook`), и несовпадение вида — тоже расхождение.

def _office(strategy: str) -> Callable[[Any], dict]:
    def build(core: Any) -> dict:
        spec = gs.SCENARIOS["nonresidential"](core, ROOT)
        spec["inputs"]["offices_strategy"] = strategy
        return spec
    return build


def _golden(name: str) -> Callable[[Any], dict]:
    return lambda core: gs.SCENARIOS[name](core, ROOT)


def _nonres(name: str) -> Callable[[Any], dict]:
    return lambda core: nonres_scenario(core, NONRES_BOOKS[name])


PROJECTS: dict[str, tuple[str, str, Callable[[Any], dict]]] = {
    # имя: (вид книги, что это, вводные)
    "default": ("v4", "Жильё одной очередью (эталон default)", _golden("default")),
    "nagatino_phased": ("v4", "Жильё очередями, ТЦ в четвёртой (эталон Нагатино)", _golden("nagatino_phased")),
    "mixed_osz_parking": ("v4", "Смешанный: жильё и офис ОСЗ с гаражом", _golden("mixed_osz_parking")),
    "retail_sports": ("v4", "Смешанный: жильё, ТЦ и ФОК", _golden("retail_sports")),
    "office_ddu": ("office", "Офис по ДДУ", _office("ddu")),
    "office_direct": ("office", "Офис без ДДУ: продажа блоками", _office("direct")),
    "office_rent": ("office", "Офис без ДДУ: аренда с выходом", _golden("office_rent")),
    "nonres_retail_sports_ddu": ("office", "ТЦ и ФОК по ДДУ", _nonres("nonres_retail_sports_ddu")),
    "nonres_parking_ddu": ("office", "Наземный паркинг по ДДУ", _nonres("nonres_parking_ddu")),
}

# Известные расхождения методики, по которым есть решение и PR, но он ещё не
# слит. Строка из этого списка печатается отдельно и CI не валит, а тест
# требует, чтобы расхождение по ней ещё было: исчезло — убрать строку отсюда.
KNOWN_GAPS: dict[str, str] = {
    "Проектный поток до финансирования":
        "#892 (214-ФЗ): движок берёт выручку ДДУ в поток проекта в месяц договора, книга — раскрытием эскроу",
    "NPV проекта из потока книги": "#892: тот же поток проекта",
    "IRR проекта из потока книги": "#892: тот же поток проекта",
}


@dataclass
class Line:
    row: str
    book: Any
    engine: Any
    ratio: bool = False

    @property
    def delta(self) -> float | None:
        if _num(self.book) and _num(self.engine):
            return float(self.book) - float(self.engine)
        return None

    def agrees(self) -> bool:
        if not (_num(self.book) and _num(self.engine)):
            # Слово против слова («N/A — долг не погашен»), пусто против пусто.
            return _blank(self.book) == _blank(self.engine) and (
                _blank(self.book) or str(self.book).strip() == str(self.engine).strip())
        tol = ABS_TOL_RATIO if self.ratio else ABS_TOL_MLN
        return abs(self.delta) <= max(tol, abs(float(self.engine)) * REL_TOL)

    def text(self, project: str) -> str:
        delta = self.delta
        return (f"проект {project}, строка «{self.row}»: книга {_fmt(self.book, self.ratio)}, "
                f"движок {_fmt(self.engine, self.ratio)}, разница "
                + (_fmt(delta, self.ratio) if delta is not None else "—"))


@dataclass
class Result:
    project: str
    title: str
    kind: str
    checked: int = 0
    mismatches: list[Line] = field(default_factory=list)
    known: list[tuple[Line, str]] = field(default_factory=list)
    seconds: float = 0.0


def _num(value: Any) -> bool:
    return isinstance(value, (int, float)) and not isinstance(value, bool)


def _blank(value: Any) -> bool:
    return value is None or (isinstance(value, str) and not value.strip())


def _fmt(value: Any, ratio: bool) -> str:
    if not _num(value):
        return repr(value)
    return f"{value:.6f}" if ratio else f"{value:,.3f} млн".replace(",", " ")


def _year(value: Any) -> int | None:
    if isinstance(value, (date, datetime)):
        return value.year
    text = str(value or "")
    return int(text[:4]) if text[:4].isdigit() else None


# --- жилая книга v4 ------------------------------------------------------------

# Годовые ряды: (строка, где в книге, откуда у движка, остаток ли на конец года).
# Где в книге — ('CF', строка) или ('CF_n', строка) — сумма по очередям.
# Ключи строки движка перечислены вариантами: берётся первый, который есть в
# каждой строке; ни одного — пробел карты, а не ноль движка.
V4_SERIES: tuple[tuple[str, tuple[str, int], tuple[str, tuple[tuple[str, ...], ...]], bool], ...] = (
    ("Выручка по договорам", ("CF", 6), ("rows", (("sales",),)), False),
    ("Раскрытие эскроу", ("CF", 8), ("rows", (("escrow_release",),)), False),
    ("Эскроу на конец года", ("CF", 9), ("rows", (("escrow",),)), True),
    ("Расходы проекта: CAPEX и коммерческие", ("CF", 13),
     ("rows", (("project_costs",),)), False),
    # CF!14 — налог в потоке, то есть уплата; начисление — свод CF!37.
    ("Налог на прибыль начислен", ("CF", 37), ("rows", (("profit_tax",),)), False),
    ("Налог на прибыль уплачен", ("CF", 14), ("cashflow", (("profit_tax",),)), False),
    ("Проценты БРИДЖа и ПФ", ("CF", 20), ("rows", (("bridge_interest", "pf_interest"),)), False),
    ("БРИДЖ — выборка", ("CF_n", 34), ("rows", (("bridge_draw",),)), False),
    ("БРИДЖ — погашение", ("CF_n", 35), ("rows", (("bridge_repayment",),)), False),
    ("ПФ — выборка", ("CF_n", 45), ("rows", (("pf_draw",),)), False),
    ("ПФ — погашение", ("CF_n", 46), ("rows", (("pf_repayment",),)), False),
    ("Долг БРИДЖ на конец года", ("CF", 17), ("rows", (("bridge_balance",),)), True),
    ("Долг ПФ на конец года", ("CF", 18), ("rows", (("pf_balance",),)), True),
    ("Поток собственника", ("CF", 25), ("cashflow", (("equity",),)), False),
    ("Проектный поток до финансирования", ("CF", 28), ("cashflow", (("project",),)), False),
)

# Главы CAPEX: подпись строки листа CAPEX (сумма по очередям) → ключи сметы движка.
V4_CAPEX: tuple[tuple[str, tuple[str, ...]], ...] = (
    ("Покупка / цена входа", ("purchase",)),
    ("Плата за ВРИ", ("land_rights",)),
    ("ИРД и согласования", ("ird",)),
    ("Проектирование, стадия П", ("design_p",)),
    ("Проектирование, стадия РД", ("design_rd",)),
    ("Авторский надзор", ("author_supervision",)),
    ("Подготовительные работы", ("preparation",)),
    ("Основное строительство — наземная часть", ("main_above",)),
    ("Основное строительство — подземная часть", ("main_under",)),
    ("Наружные инженерные сети", ("utilities",)),
    ("Благоустройство", ("landscaping",)),
    ("Сдача и ввод", ("commissioning",)),
    ("Содержание стройплощадки", ("site_maintenance",)),
    ("Техзаказчик / стройконтроль", ("technical_supervision",)),
    ("Управление проектом", ("project_management",)),
    ("Вознаграждение генподрядчика", ("gc_fee",)),
    ("Резерв", ("reserve",)),
    ("Социальная нагрузка", ("social",)),
    ("Снос и демонтаж", ("demolition",)),
    ("Расселение", ("resettlement",)),
    ("Проценты и обеспечение по рассрочке ВРИ", ("vri_interest", "vri_security")),
)


def _v4_lines(core: Any, book: Any, consolidated: dict) -> list[Line]:
    from openpyxl.utils import get_column_letter
    from xlsx_eval import Evaluator

    import excel_book_parity as parity

    sys.setrecursionlimit(max(sys.getrecursionlimit(), 400000))
    ev = Evaluator(book)
    lines: list[Line] = []

    # 1. ОТЧЕТ — итоги (та же мера, что у excel_book_parity.compare_report).
    report = book["ОТЧЕТ"]
    labels = {}
    for r in range(1, 40):
        label = report.cell(r, 1).value
        if isinstance(label, str) and not label.startswith("="):
            labels.setdefault(label.strip(), r)
    for label, path, scale in REPORT_ROWS:
        r = next((row for key, row in labels.items() if key.startswith(label)), None)
        engine = parity._path(consolidated, path)
        if r is None:
            lines.append(Line(f"ОТЧЕТ: {label}", "нет строки", engine))
            continue
        value = ev.cell("ОТЧЕТ", f"B{r}")
        if isinstance(value, str) and engine is None:
            engine = value  # «N/A …» книги там, где движок числа не дал
        lines.append(Line(f"ОТЧЕТ: {label}", value,
                          engine / scale if _num(engine) else engine, ratio=scale == 1))

    # 2. Годовые ряды CF.
    cf = book["CF"]
    cols = [c for c in range(4, cf.max_column + 1) if cf.cell(3, c).value is not None]
    years = [_year(ev.cell("CF", f"{get_column_letter(c)}3")) for c in cols]
    queues = [name for name in book.sheetnames if name.startswith("CF_")]
    finance = consolidated.get("finance") or {}
    rows = finance.get("rows") or []
    cashflow = consolidated.get("cashflow") or {}
    engine_years = sorted({_year(r.get("month")) for r in rows})
    last_engine_year = engine_years[-1] if engine_years else None

    def book_series(where: tuple[str, int]) -> list[float]:
        sheets = queues if where[0] == "CF_n" else [where[0]]
        out = []
        for c in cols:
            letter = get_column_letter(c)
            out.append(sum(float(ev.cell(s, f"{letter}{where[1]}") or 0.0) for s in sheets))
        return out

    def engine_series(source: str, variants: tuple[tuple[str, ...], ...]) -> tuple[list[int | None], list[float]]:
        if source == "rows":
            present = set.intersection(*(set(r) for r in rows)) if rows else set()
            keys = next((v for v in variants if set(v) <= present), None)
            if keys is None:
                # Ни одного варианта: это пробел карты, а не ноль движка.
                raise KeyError(f"в строках движка нет ни одного из {variants}")
            return ([_year(r.get("month")) for r in rows],
                    [sum(float(r.get(k) or 0.0) for k in keys) / MLN for r in rows])
        months = cashflow.get("months") or []
        return [_year(m) for m in months], [float(v or 0.0) / MLN for v in cashflow[variants[0][0]]]

    book_cache: dict[str, list[float]] = {}
    for label, where, (source, variants), balance in V4_SERIES:
        values = book_series(where)
        book_cache[label] = values
        e_years, e_values = engine_series(source, variants)
        by_book: dict[int, float] = defaultdict(float)
        by_engine: dict[int, float] = defaultdict(float)
        for y, v in zip(years, values):
            by_book[y] = v if balance else by_book[y] + v
        for y, v in zip(e_years, e_values):
            by_engine[y] = v if balance else by_engine[y] + v
        for y in sorted(set(by_book) | set(by_engine)):
            if balance and last_engine_year is not None and y > last_engine_year:
                continue  # остаток за горизонтом движка — его последний остаток, не новый год
            lines.append(Line(f"{label}, {y}", by_book.get(y, 0.0), by_engine.get(y, 0.0)))

    # 3. Пик собственных денег — худший накопленный поток собственника.
    cumulative = worst = 0.0
    for v in book_cache["Поток собственника"]:
        cumulative += v
        worst = min(worst, cumulative)
    engine_peak = core.owner_cash_peak(finance, cashflow.get("equity") or []) / MLN
    lines.append(Line("Пик собственных денег", -worst, engine_peak))

    # 4. NPV и IRR проекта — функцией движка из потока книги.
    summary = consolidated.get("summary") or {}
    flow = [v * MLN for v in book_cache["Проектный поток до финансирования"]]
    while flow and abs(flow[-1]) < 1e-6:
        flow.pop()
    rate = float(consolidated.get("_discount_rate") or 0.0)
    lines.append(Line("NPV проекта из потока книги", core._monthly_npv(flow, rate) / MLN,
                      float(summary.get("npv") or 0.0) / MLN))
    irr = core._monthly_irr(flow)
    lines.append(Line("IRR проекта из потока книги", irr, summary.get("irr_project"), ratio=True))

    # 5. CAPEX по главам: сумма блоков очередей листа CAPEX.
    capex_sheet = book["CAPEX"]
    totals: dict[str, float] = defaultdict(float)
    for r in range(1, capex_sheet.max_row + 1):
        label = capex_sheet.cell(r, 1).value
        if isinstance(label, str) and capex_sheet.cell(r, 2).value is not None:
            value = ev.cell("CAPEX", f"B{r}")
            if _num(value):
                totals[label.strip()] += float(value)
    capex = consolidated.get("capex") or {}
    for label, keys in V4_CAPEX:
        lines.append(Line(f"CAPEX: {label}", totals.get(label, 0.0),
                          sum(float(capex.get(k) or 0.0) for k in keys) / MLN))
    return lines


# --- книга нежилого проекта ------------------------------------------------------

def _office_lines(book: Any) -> list[Line]:
    import office_workbook as ow

    ev = ow.evaluate(book)
    sheet = book[ow.CHECK_SHEET]
    lines = []
    for r in range(5, sheet.max_row + 1):
        if not sheet[f"G{r}"].value:
            continue
        ratio = sheet[f"C{r}"].number_format == "0.000000"
        line = Line(f"{sheet[f'A{r}'].value}: {sheet[f'B{r}'].value}",
                    ev.cell(ow.CHECK_SHEET, f"C{r}"), sheet[f"D{r}"].value, ratio=ratio)
        verdict = ev.cell(ow.CHECK_SHEET, f"G{r}")
        # Вердикт решает формула самой книги (её допуск строже нашего); строка
        # с ошибкой вместо слова — расхождение.
        line.verdict_ok = verdict in (ow.AGREE, "—")  # type: ignore[attr-defined]
        lines.append(line)
    return lines


# --- сборка и сверка ------------------------------------------------------------

def build(core: Any, name: str) -> tuple[bytes, dict, dict]:
    """Книга эталона, отчёт её сборки и свод движка."""
    spec = PROJECTS[name][2](core)
    content, _file, meta = core.build_project_workbook(
        copy.deepcopy(spec["inputs"]), copy.deepcopy(spec["tep"]), [],
        copy.deepcopy(spec["phasing"]), project_name=name)
    bundle = core._run_authoritative_model(
        copy.deepcopy(spec["inputs"]), copy.deepcopy(spec["tep"]), [], copy.deepcopy(spec["phasing"]))
    consolidated = dict(bundle["consolidated"])
    consolidated["_discount_rate"] = float(spec["inputs"].get("discount_rate_pct", 20) or 20) / 100
    return content, meta, consolidated


def check(core: Any, name: str, tamper: Callable[[Any], None] | None = None,
          built: tuple[bytes, dict, dict] | None = None) -> Result:
    import openpyxl

    started = time.time()
    kind, title, _ = PROJECTS[name]
    content, meta, consolidated = built or build(core, name)
    result = Result(name, title, kind)
    book = openpyxl.load_workbook(io.BytesIO(content))
    actual = "office" if meta.get("nonres_book") else "v4"
    if actual != kind:
        result.mismatches.append(Line("вид книги", actual, kind))
        return result
    if tamper:
        tamper(book)
    lines = _v4_lines(core, book, consolidated) if kind == "v4" else _office_lines(book)
    result.checked = len(lines)
    for line in lines:
        ok = getattr(line, "verdict_ok", None)
        agrees = line.agrees() if ok is None else ok
        if agrees:
            continue
        reason = next((why for key, why in KNOWN_GAPS.items() if line.row.startswith(key)), None)
        if reason:
            result.known.append((line, reason))
        else:
            result.mismatches.append(line)
    result.seconds = time.time() - started
    return result


# --- подделки: проверка обязана краснеть ---------------------------------------

def tamper_v4_reserve(book: Any) -> None:
    """Подделка методики: резерв очереди 1 на 1% больше, чем у движка."""
    sheet = book["CAPEX"]
    row = next(r for r in range(1, sheet.max_row + 1) if sheet.cell(r, 1).value == "Резерв")
    sheet.cell(row, 2).value = f"{sheet.cell(row, 2).value}*1.01"


def tamper_office_revenue(book: Any) -> None:
    """Подделка методики: выручка проекта в книге нежилого проекта на 1% выше."""
    import office_workbook as ow

    sheet = book[ow.CHECK_SHEET]
    row = next(r for r in range(5, sheet.max_row + 1) if sheet[f"B{r}"].value == "Выручка проекта")
    sheet[f"C{row}"] = f"={str(sheet[f'C{row}'].value).lstrip('=')}*1.01"


TAMPERS = {"v4": tamper_v4_reserve, "office": tamper_office_revenue}


# --- вывод ------------------------------------------------------------------------

def markdown(results: list[Result], tampered: bool = False) -> str:
    out = ["## Книга Excel ↔ движок" + (" — подделка" if tampered else ""), "",
           "| Проект | Книга | Строк | Расхождений | Известных | с |", "|---|---|---:|---:|---:|---:|"]
    for r in results:
        out.append(f"| {r.project} — {r.title} | {r.kind} | {r.checked} | "
                   f"{'**' + str(len(r.mismatches)) + '**' if r.mismatches else 0} | {len(r.known)} | {r.seconds:.0f} |")
    bad = [line.text(r.project) for r in results for line in r.mismatches]
    if bad:
        out += ["", "### Расхождения", ""] + [f"- {text}" for text in bad[:200]]
    known = [(line.text(r.project), why) for r in results for line, why in r.known]
    if known:
        out += ["", "### Известные расхождения (решение принято, PR не слит)", ""]
        out += [f"- {text} — {why}" for text, why in known[:60]]
    return "\n".join(out) + "\n"


def main() -> int:
    parser = argparse.ArgumentParser(description=__doc__.split("\n\n")[0])
    parser.add_argument("--only", nargs="*", choices=list(PROJECTS))
    parser.add_argument("--tamper", action="store_true",
                        help="испортить формулу книги в каждом эталоне; код 0 — только если все покраснели")
    parser.add_argument("--summary", type=Path, help="куда дописать таблицу Markdown (GITHUB_STEP_SUMMARY)")
    args = parser.parse_args()
    core = gs.load_core(ROOT)
    results = []
    for name in args.only or list(PROJECTS):
        tamper = TAMPERS[PROJECTS[name][0]] if args.tamper else None
        result = check(core, name, tamper)
        results.append(result)
        state = "расхождений нет" if not result.mismatches else f"РАСХОЖДЕНИЙ: {len(result.mismatches)}"
        print(f"===== {name} — {result.title}: строк {result.checked}, {state} ({result.seconds:.0f} с)",
              flush=True)
        for line in result.mismatches:
            print("  " + line.text(name), flush=True)
        for line, why in result.known:
            print(f"  известное: {line.text(name)} — {why}", flush=True)
    if args.summary:
        with args.summary.open("a", encoding="utf-8") as fh:
            fh.write(markdown(results, args.tamper))
    if args.tamper:
        missed = [r.project for r in results if not r.mismatches]
        if missed:
            print(f"ПОДДЕЛКА НЕ ПОЙМАНА: {', '.join(missed)}")
            return 1
        print("Подделка поймана во всех эталонах.")
        return 0
    return 1 if any(r.mismatches for r in results) else 0


if __name__ == "__main__":
    sys.exit(main())
