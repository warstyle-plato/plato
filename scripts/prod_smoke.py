#!/usr/bin/env python3
"""Проверка прода после выката: готовые поверхности на эталонном проекте.

Тест PR проверяет код, а не то, что видит пользователь. За два дня на прод
уехали тизер КРТ Нагатино с 10 участками из 20 и выгрузка торгов без
комментариев — при зелёном CI. Здесь проверяется то, что человек получает на
руки с живого прода, на одном эталоне (`scripts/prod_smoke_reference.json`).

Только чтение и по одному запросу на поверхность: проект собирается в браузере
из пресета (как это делает человек кнопкой «Импорт проекта / пресета»), на
сервере ничего не сохраняется. Единственный побочный эффект — выгрузка торгов:
её лоты встают в очередь комментария Платона ровно так же, как при нажатии
кнопки человеком; очередь повторно известный лот не ставит.

Запуск:  python3 scripts/prod_smoke.py --base https://developaid.ru
Ключ владельца (тизер и экономика «Итога» — за входом): PROD_SMOKE_ADMIN_KEY.
Нет ключа — проверка ПРОПУЩЕНА с причиной, а не зелёная.
"""

from __future__ import annotations

import argparse
import io
import json
import os
import re
import socket
import sys
import time
import urllib.error
import urllib.parse
import urllib.request
from dataclasses import asdict, dataclass, field
from http.client import HTTPException
from pathlib import Path
from typing import Any, Callable

ROOT = Path(__file__).resolve().parent.parent
REFERENCE = Path(__file__).resolve().parent / "prod_smoke_reference.json"

OK, FAIL, SKIP = "ok", "fail", "skip"
MARK = {OK: "✅", FAIL: "❌", SKIP: "⚪ пропущено"}

# Следы сырого исключения вёрстки вместо документа или вместо честного отказа.
RAW_DUMP_MARKERS = ("Traceback", "LayoutError", "too large on page", "Flowable",
                    "reportlab", "ReportLab", "<Paragraph", "<Table@")


@dataclass
class Check:
    name: str
    status: str
    expected: str
    got: str
    detail: str = ""
    evidence: dict[str, Any] = field(default_factory=dict)


def load_reference(path: Path = REFERENCE) -> dict[str, Any]:
    return json.loads(path.read_text(encoding="utf-8"))


# --- судьи: чистые функции над тем, что отдал прод --------------------------
# Их проверяет юнит-тест на подделках (tests/test_prod_smoke_judges.py): зелёный
# ответ на тизере с 10 участками из 20 значил бы, что проверка ничего не ловит.

def judge_autoload(data: dict[str, Any], ref: dict[str, Any]) -> Check:
    name = "1. Автозагрузка участков"
    want = list(ref["cadastral_numbers"])
    area_want = float(ref["land_area_ha"])
    tol = float(ref.get("land_area_tolerance_ha", 0.05))
    expected = f"найдено {len(want)} из {len(want)}, ≈ {area_want:.2f} га (±{tol})"
    results = data.get("results") or []
    found = [r for r in results if r.get("found")]
    area = sum(float(r.get("area_ha") or 0) for r in found)
    got_numbers = {str(r.get("cadastral_number")) for r in found}
    missing = [n for n in want if n not in got_numbers]
    got = f"найдено {len(found)} из {len(results)}, {area:.4f} га"
    problems = []
    if len(results) != len(want):
        problems.append(f"в ответе {len(results)} номеров вместо {len(want)}")
    if missing:
        problems.append("не найдены: " + ", ".join(missing))
    if abs(area - area_want) > tol:
        problems.append(f"площадь {area:.4f} га вместо ≈ {area_want:.4f}")
    if data.get("reason"):
        problems.append(f"причина от сервера: {data['reason']}")
    return Check(name, FAIL if problems else OK, expected, got, "; ".join(problems))


def pdf_text(pdf: bytes) -> tuple[str, int]:
    from pypdf import PdfReader

    reader = PdfReader(io.BytesIO(pdf))
    return "\n".join(page.extract_text() or "" for page in reader.pages), len(reader.pages)


def judge_teaser(text: str, ref: dict[str, Any]) -> Check:
    """Тизер: сколько участков он называет и сколько контуров нарисовал.

    Число — из строки «Кадастровые номера (N)» паспорта участка; подпись
    карты «Территория из N участков» и «контур нарисован для X из N» —
    сверяются, если тизер их печатает. Карта — картинка, контуры из неё не
    извлекаются: о недорисованных говорит только подпись."""
    name = "2. Тизер PDF"
    want = len(ref["cadastral_numbers"])
    expected = f"PDF, «Кадастровые номера ({want})», контуров {want} из {want}, без дампа ReportLab"
    flat = re.sub(r"\s+", " ", text)
    problems: list[str] = []
    listed = re.findall(r"Кадастровые номера \((\d+)\)", flat)
    territory = re.findall(r"Территория из (\d+) участк", flat)
    drawn = re.findall(r"контур нарисован для (\d+) из (\d+)", flat)
    got_parts = []
    if listed:
        got_parts.append(f"«Кадастровые номера ({listed[0]})»")
        if int(listed[0]) != want:
            problems.append(f"тизер называет {listed[0]} участков из {want}")
    else:
        single = "Кадастровые номера" in flat
        got_parts.append("одиночный номер" if single else "паспорта участка нет")
        problems.append(f"тизер не называет числа участков (ждали {want})")
    if territory:
        got_parts.append(f"подпись карты: {territory[0]} участков")
        if int(territory[0]) != want:
            problems.append(f"подпись карты — территория из {territory[0]} участков, а не {want}")
    if drawn:
        x, n = drawn[0]
        got_parts.append(f"контур нарисован для {x} из {n}")
        problems.append(f"на карте {x} контуров из {n}")
    if "Карта участка не построена" in flat:
        got_parts.append("карты нет")
        problems.append("карта участка не построена")
    dumps = [m for m in RAW_DUMP_MARKERS if m in text]
    if dumps:
        problems.append("в тексте сырой дамп ошибки: " + ", ".join(dumps))
    return Check(name, FAIL if problems else OK, expected, "; ".join(got_parts), "; ".join(problems))


def judge_teaser_response(status: int, content_type: str, body: bytes,
                          ref: dict[str, Any]) -> Check:
    name = "2. Тизер PDF"
    want = len(ref["cadastral_numbers"])
    expected = f"HTTP 200, PDF, «Кадастровые номера ({want})»"
    if status != 200 or not body.startswith(b"%PDF"):
        text = body[:2000].decode("utf-8", "replace")
        dumps = [m for m in RAW_DUMP_MARKERS if m in text]
        detail = "сырой дамп ошибки: " + ", ".join(dumps) if dumps else ""
        return Check(name, FAIL, expected, f"HTTP {status} {content_type}: {text[:300]}", detail)
    text, pages = pdf_text(body)
    check = judge_teaser(text, ref)
    check.got = f"PDF {pages} стр.; " + check.got
    return check


def _header_index(header: list[Any], prefix: str) -> int | None:
    for index, value in enumerate(header):
        if isinstance(value, str) and value.strip().startswith(prefix):
            return index
    return None


# Сырое исключение в клетке — не «честная причина», а поломка разбора.
RAW_ERROR = re.compile(r"HTTPException|Traceback|\w+Error\b|\w+Exception\b")


def classify_note(value: Any) -> str:
    """Клетка комментария: `text` (комментарий или названная причина),
    `raw` (сырой текст исключения), `empty`."""
    text = str(value or "").strip()
    if not text:
        return "empty"
    return "raw" if RAW_ERROR.search(text) else "text"


def nspd_link_problem(url: str, host: str) -> str:
    """Пусто — ссылка ведёт на объект в НСПД; иначе — что с ней не так.

    Общая карта (`map?thematic=PKK`) участок не открывает: НСПД номер из адреса
    не читает, объект открывают только координаты или `selectedCard`."""
    parsed = urllib.parse.urlparse(url)
    if parsed.scheme != "https" or not parsed.netloc or " " in url:
        return "не URL"
    if not (parsed.netloc == host or parsed.netloc.endswith("." + host)):
        return f"не {host}"
    query = urllib.parse.parse_qs(parsed.query)
    if "selectedCard" in query or ("coordinate_x" in query and "coordinate_y" in query):
        return ""
    return "общая карта, а не участок"


def judge_auction_export(xlsx: bytes, ref: dict[str, Any]) -> Check:
    """Выгрузка торгов: комментарий заполнен, ссылка НСПД ведёт на участок.

    «Заполнен» — в клетке комментарий либо названная причина, почему его нет.
    Пустая клетка и сырой текст исключения — поломка, а не ответ."""
    import openpyxl

    name = "3. Excel-выгрузка торгов"
    cfg = ref["auction_export"]
    expected = (f"«{cfg['comment_prefix']}…» у большинства строк — текст или честная причина "
                f"(не пусто и не сырое исключение); у строки с кадастровым номером — ссылка "
                f"{cfg['nspd_host']} на участок или причина, почему её нет")
    wb = openpyxl.load_workbook(io.BytesIO(xlsx))
    ws = wb.active
    header = [cell.value for cell in ws[1]]
    rows = [r for r in range(2, ws.max_row + 1)
            if any(ws.cell(r, c).value not in (None, "") for c in range(1, ws.max_column + 1))]
    problems: list[str] = []
    got: list[str] = [f"лист «{ws.title}», строк {len(rows)}"]
    if not rows:
        problems.append("в книге нет строк")

    for label, prefix, required in (("Комментарий", cfg["comment_prefix"], True),
                                    ("Примечание", cfg["note_prefix"], False)):
        column = _header_index(header, prefix)
        if column is None:
            if required:
                problems.append(f"колонки «{prefix}…» нет; заголовки: "
                                + ", ".join(str(h) for h in header if h))
            else:
                got.append(f"колонки «{prefix}» в выгрузке нет")
            continue
        kinds = [classify_note(ws.cell(r, column + 1).value) for r in rows]
        text, raw, empty = kinds.count("text"), kinds.count("raw"), kinds.count("empty")
        got.append(f"«{label}»: текст {text}, сырое исключение {raw}, пусто {empty} из {len(rows)}")
        sample = next((str(ws.cell(r, column + 1).value)[:90] for r, k in zip(rows, kinds)
                       if k == "raw"), "")
        if rows and text * 2 <= len(rows):
            problems.append(f"«{label}» содержателен лишь у {text} строк из {len(rows)}"
                            + (f"; например «{sample}»" if sample else ""))
        elif raw:
            # Сырой текст исключения человек читает как ответ — это поломка,
            # даже когда большинство строк в порядке.
            problems.append(f"«{label}»: сырое исключение вместо ответа у {raw} строк из "
                            f"{len(rows)}, например «{sample}»")

    cadastre = _header_index(header, cfg["cadastre_header"])
    reason_column = _header_index(header, cfg["nspd_header"])
    with_numbers = bad_links = unexplained = good = 0
    bad_sample = ""
    for r in rows:
        targets = [ws.cell(r, c).hyperlink.target for c in range(1, ws.max_column + 1)
                   if ws.cell(r, c).hyperlink and cfg["nspd_host"] in (ws.cell(r, c).hyperlink.target or "")]
        verdicts = [nspd_link_problem(t, cfg["nspd_host"]) for t in targets]
        for t, v in zip(targets, verdicts):
            if v:
                bad_links += 1
                bad_sample = bad_sample or f"{t} — {v}"
        has_number = cadastre is not None and str(ws.cell(r, cadastre + 1).value or "").strip()
        if not has_number:
            continue
        with_numbers += 1
        if any(not v for v in verdicts):
            good += 1
        elif not (reason_column is not None and str(ws.cell(r, reason_column + 1).value or "").strip()):
            unexplained += 1
    got.append(f"строк с кадастровым номером {with_numbers}: ссылка на участок у {good}, "
               f"без ссылки и без причины {unexplained}")
    if bad_links:
        problems.append(f"ссылок НСПД не на участок: {bad_links}, например {bad_sample}")
    if unexplained:
        problems.append(f"у {unexplained} строк с кадастровым номером нет ни ссылки НСПД на участок, "
                        "ни причины её отсутствия")
    return Check(name, FAIL if problems else OK, expected, "; ".join(got), "; ".join(problems))


def judge_workbook(xlsx: bytes, tep: dict[str, Any], ref: dict[str, Any]) -> Check:
    """Книга v4: лист «Вводные» есть; у каждого ОСЗ проекта — строка мест его гаража."""
    import openpyxl

    name = "4. Книга Excel v4"
    parking = {k: v for k, v in ref["osz_parking"].items() if not k.startswith("_")}
    present = [k for k in parking if float((tep.get(k) or {}).get("gns") or 0) > 0]
    expected = "лист «Вводные»" + ("; строки: " + ", ".join(f"«{parking[k]}»" for k in present)
                                   if present else "; ОСЗ в проекте нет — строки паркинга ОСЗ не нужны")
    wb = openpyxl.load_workbook(io.BytesIO(xlsx), read_only=True)
    if "Вводные" not in wb.sheetnames:
        return Check(name, FAIL, expected, "листы: " + ", ".join(wb.sheetnames[:12]),
                     "листа «Вводные» нет")
    labels = set()
    for row in wb["Вводные"].iter_rows(values_only=True):
        labels.update(v.strip() for v in row if isinstance(v, str))
    missing = [parking[k] for k in present if parking[k] not in labels]
    got = f"листов {len(wb.sheetnames)}, «Вводные» есть; строк паркинга ОСЗ {len(present) - len(missing)} из {len(present)}"
    detail = ("нет строк: " + ", ".join(f"«{m}»" for m in missing)) if missing else ""
    return Check(name, FAIL if missing else OK, expected, got, detail)


def judge_itog(page: dict[str, Any], ref: dict[str, Any]) -> Check:
    """«Итог»: удельные числа подписаны единицей с делителем.

    `page` — то, что отрисовано: строки «Ключевых параметров» (подпись →
    значение), заголовки «Удельной экономики» и число её строк, текст раздела."""
    name = "5. Страница «Итог»"
    cfg = ref["itog"]
    expected = ("удельные строки «Ключевых параметров» — «тыс. ₽/м² <база>» в обеих базах; "
                "заголовки «Удельной экономики»: " + ", ".join(f"«{h}»" for h in cfg["unit_headers"]))
    params: dict[str, str] = page.get("params") or {}
    problems: list[str] = []
    if not params:
        return Check(name, FAIL, expected, "«Ключевые параметры» пусты", "расчёт не отрисован")
    for label in cfg["per_metre_rows"]:
        value = params.get(label)
        if value is None:
            problems.append(f"нет строки «{label}»")
            continue
        units = re.findall(r"(?:тыс\.|млн|млрд)\s*₽\s*/\s*м²\s*([^\d·]*)", value)
        bases = [u.strip() for u in units if u.strip()]
        if len(bases) < 2:
            problems.append(f"«{label}: {value}» — нет единицы с делителем в обеих базах")
    headers = [re.sub(r"\s+", " ", h).strip() for h in page.get("unit_headers") or []]
    for want in cfg["unit_headers"]:
        if not any(h.lower() == want.lower() for h in headers):
            problems.append(f"в «Удельной экономике» нет колонки «{want}» (есть: {', '.join(headers)})")
    if not page.get("unit_rows"):
        problems.append("«Удельная экономика» без строк")
    text = page.get("text") or ""
    for junk in ("NaN", "undefined", "Infinity"):
        if junk in text:
            problems.append(f"в разделе напечатано «{junk}»")
    sample = "; ".join(f"{k}: {params[k]}" for k in cfg["per_metre_rows"][:2] if k in params)
    got = f"строк параметров {len(params)}, удельной экономики {page.get('unit_rows', 0)}; {sample}"
    return Check(name, FAIL if problems else OK, expected, got, "; ".join(problems))


# --- сеть -------------------------------------------------------------------

def http(method: str, url: str, body: Any = None, timeout: int = 300,
         headers: dict[str, str] | None = None) -> tuple[int, str, bytes]:
    data = None
    hdrs = {"User-Agent": "developaid-prod-smoke/1", **(headers or {})}
    if body is not None:
        data = json.dumps(body, ensure_ascii=False).encode("utf-8")
        hdrs["Content-Type"] = "application/json"
    request = urllib.request.Request(url, data=data, method=method, headers=hdrs)
    try:
        with urllib.request.urlopen(request, timeout=timeout) as response:
            return response.status, response.headers.get("Content-Type", ""), response.read()
    except urllib.error.HTTPError as exc:
        return exc.code, exc.headers.get("Content-Type", ""), exc.read()


# Ответа нет вовсе: прод оборвал соединение (RemoteDisconnected), сбросил его,
# не ответил за таймаут, не нашёлся. HTTP 4xx/5xx сюда не входят — это ответ.
NETWORK_ERRORS = (OSError, HTTPException, socket.timeout)
RETRY_PAUSE = 20  # с; обрыв после выката — обычно холодный старт воркера


class NetworkDrop(Exception):
    """Запрос остался без ответа и после повтора; текст — причина и место."""


def describe_drop(exc: BaseException) -> str:
    kind = type(exc).__name__
    if isinstance(exc, urllib.error.URLError) and not isinstance(exc, urllib.error.HTTPError):
        reason = exc.reason
        if isinstance(reason, (TimeoutError, socket.timeout)):
            return f"таймаут ({reason})"
        return f"нет соединения ({type(reason).__name__}: {reason})"
    if isinstance(exc, (TimeoutError, socket.timeout)):
        return f"таймаут ({kind})"
    if isinstance(exc, ConnectionResetError) and kind != "RemoteDisconnected":
        return f"сброс соединения ({kind}: {exc})"
    return f"обрыв соединения ({kind}: {exc})" if str(exc) else f"обрыв соединения ({kind})"


def fetch(method: str, url: str, body: Any = None, timeout: int = 300,
          headers: dict[str, str] | None = None,
          log: Callable[[str], None] | None = None) -> tuple[int, str, bytes, str]:
    """`http` с одной повторной попыткой на обрыв без ответа.

    Четвёртое значение — пометка «повтор после обрыва …» для строки итога
    (пусто, если хватило первой попытки). Любой HTTP-статус — ответ, его не
    повторяем. Второй обрыв — `NetworkDrop` с причиной, методом, путём и
    временем: провал этой проверки, а не всего прогона."""
    path = urllib.parse.urlparse(url).path or "/"
    first = ""
    for attempt in (1, 2):
        started = time.monotonic()
        try:
            status, ctype, data = http(method, url, body, timeout=timeout, headers=headers)
            return status, ctype, data, (f"повтор после обрыва: {first}" if first else "")
        except NETWORK_ERRORS as exc:
            where = f"{describe_drop(exc)} на {method} {path} через {time.monotonic() - started:.0f} с"
            if attempt == 2:
                raise NetworkDrop(f"{where}; повтор после обрыва ({first}) тоже без ответа") from exc
            first = where
            if log:
                log(f"{where}; повтор через {RETRY_PAUSE} с")
            time.sleep(RETRY_PAUSE)
    raise AssertionError("unreachable")


def with_note(check: Check, note: str) -> Check:
    if note:
        check.got = f"{check.got} ({note})"
    return check


def health(base: str) -> dict[str, Any]:
    try:
        status, _, body = http("GET", base + "/health", timeout=30)
        return json.loads(body) if status == 200 else {"status": f"HTTP {status}"}
    except Exception as exc:  # noqa: BLE001 — сеть: причина уходит в итог
        return {"status": f"нет ответа: {exc}"}


def deployed(seen: dict[str, Any], sha: str, version: str) -> tuple[bool, str]:
    """Прод отдаёт собранное: коммит совпал; а пуст коммит — совпала версия.

    Пустой `commit` — изъян наблюдаемости (образ собран без APP_COMMIT), и
    итог называет его, но версия из того же коммита остаётся свидетельством."""
    commit = str(seen.get("commit") or "")
    if commit:
        return commit.startswith(sha) or sha.startswith(commit), ""
    if version and str(seen.get("version") or "") == version:
        return True, "/health отдаёт пустой commit — сверено по версии"
    return False, ""


def wait_for_commit(base: str, sha: str, timeout: int, log: Callable[[str], None],
                    version: str = "") -> Check:
    name = "0. Выкат дошёл"
    deadline = time.time() + timeout
    want = f"/health commit {sha[:12]}" + (f" (версия {version})" if version else "")
    while True:
        seen = health(base)
        commit = str(seen.get("commit") or "")
        ok, note = deployed(seen, sha, version)
        got = f"{seen.get('version')} ({commit[:12] or 'commit пуст'})"
        if ok:
            return Check(name, OK, want, got, note)
        if time.time() >= deadline:
            return Check(name, FAIL, want + f" за {timeout // 60} мин", got,
                         "прод не поднял собранный образ — проверять старую версию бессмысленно")
        log(f"прод отдаёт {got}, ждём {sha[:12]}…")
        time.sleep(30)


# --- браузер: проект собирается так же, как у человека ----------------------

def _chromium_launch(p):
    exe = os.environ.get("PROD_SMOKE_CHROMIUM") or ""
    if not exe and Path("/opt/pw-browsers/chromium").exists():
        exe = "/opt/pw-browsers/chromium"
    return p.chromium.launch(executable_path=exe or None)


def _route_via_python(page) -> None:
    """Песочница с перехватывающим TLS-прокси: браузер его CA не знает, а Python
    знает. Трафик страницы идёт через urllib — проверка сертификата остаётся."""
    def handle(route):
        rq = route.request
        headers = {k: v for k, v in rq.headers.items()
                   if k.lower() not in ("host", "content-length", "accept-encoding")}
        request = urllib.request.Request(rq.url, data=rq.post_data_buffer, method=rq.method,
                                         headers=headers)
        try:
            with urllib.request.urlopen(request, timeout=600) as response:
                status, hdrs, body = response.status, dict(response.headers), response.read()
        except urllib.error.HTTPError as exc:
            status, hdrs, body = exc.code, dict(exc.headers), exc.read()
        except Exception:  # noqa: BLE001
            route.abort()
            return
        for key in ("Transfer-Encoding", "Content-Encoding", "Content-Length"):
            hdrs.pop(key, None)
        route.fulfill(status=status, headers=hdrs, body=body)
    page.route("**/*", handle)


def browser_project(base: str, ref: dict[str, Any], admin_key: str, out: Path,
                    log: Callable[[str], None]) -> dict[str, Any]:
    """Главная страница: импорт пресета эталона, расчёт, «Итог».

    Возвращает то, что увидел человек («Итог»), и груз страницы для тизера и
    книги — тот же `currentPdfReportPayload()`, что шлют её кнопки."""
    from playwright.sync_api import sync_playwright

    preset = ROOT / ref["preset"]
    with sync_playwright() as p:
        browser = _chromium_launch(p)
        page = browser.new_page(viewport={"width": 1440, "height": 1000})
        if os.environ.get("HTTPS_PROXY") and os.environ.get("PROD_SMOKE_ROUTE_VIA_PYTHON", "1") == "1":
            _route_via_python(page)
        if admin_key:
            page.add_init_script(
                f"try{{localStorage.setItem('plato_projects_key',{json.dumps(admin_key)})}}catch(e){{}}")
        alerts: list[str] = []
        page.on("dialog", lambda d: (alerts.append(d.message), d.dismiss()))
        page.goto(base + "/", wait_until="domcontentloaded", timeout=120_000)
        page.wait_for_function("typeof uploadPreset==='function'&&typeof applyPreset==='function'",
                               timeout=120_000)
        page.set_input_files("#presetFile", str(preset))
        page.evaluate("uploadPreset()")
        page.wait_for_function("typeof presetPreview!=='undefined'&&!!presetPreview", timeout=120_000)
        page.evaluate("applyPreset()")
        try:
            page.wait_for_function("inputs._land_lookup&&inputs._land_lookup.found_count",
                                   timeout=180_000)
        except Exception:  # noqa: BLE001 — автозагрузку проверяет пункт 1
            log("страница не дождалась автозагрузки участка за 3 минуты")
        locked = page.evaluate("async()=>{const r=await calculate();openTab('report');"
                               "return r===null&&typeof calcNeedsLogin==='function'&&calcNeedsLogin()}")
        page.wait_for_timeout(3000)
        itog = page.evaluate("""()=>{
          const params={};
          document.querySelectorAll('#projectParamsTable tr').forEach(tr=>{
            const c=tr.querySelectorAll('td,th'); if(c.length>=2)
              params[c[0].innerText.trim()]=c[c.length-1].innerText.replace(/\\s+/g,' ').trim()});
          const table=document.getElementById('unitEconomicsTable');
          const heads=table?[...table.closest('table').querySelectorAll('thead th')].map(t=>t.innerText):[];
          const sec=document.getElementById('rsSummary');
          return {params, unit_headers:heads, unit_rows:table?table.querySelectorAll('tr').length:0,
                  text:sec?sec.innerText:''};
        }""")
        itog["locked"] = bool(locked)
        itog["alerts"] = alerts
        payload = page.evaluate("JSON.parse(JSON.stringify(currentPdfReportPayload()))")
        try:
            page.screenshot(path=str(out / "itog.png"), full_page=False)
        except Exception:  # noqa: BLE001
            pass
        browser.close()
    return {"itog": itog, "payload": payload}


def browser_auction_export(base: str, out: Path, log: Callable[[str], None]) -> tuple[int, bytes, int]:
    """Страница торгов: выборка, которую видит человек, и её кнопка выгрузки."""
    from playwright.sync_api import sync_playwright

    with sync_playwright() as p:
        browser = _chromium_launch(p)
        page = browser.new_page(viewport={"width": 1440, "height": 1000}, accept_downloads=True)
        if os.environ.get("HTTPS_PROXY") and os.environ.get("PROD_SMOKE_ROUTE_VIA_PYTHON", "1") == "1":
            _route_via_python(page)
        alerts: list[str] = []
        page.on("dialog", lambda d: (alerts.append(d.message), d.dismiss()))
        page.goto(base + "/auctions", wait_until="domcontentloaded", timeout=120_000)
        # Лоты страница сама не тянет: человек жмёт «Обновить» — и мы тоже.
        page.wait_for_function("typeof discover==='function'", timeout=120_000)
        page.click("#refresh")
        page.wait_for_function("typeof state!=='undefined'&&(state.filtered||[]).length>0",
                               timeout=240_000)
        count = page.evaluate("state.filtered.length")
        # Проверяется файл, который человек получает на диск, — а не ответ
        # маршрута: между ними ещё страница.
        statuses: list[int] = []
        page.on("response", lambda r: statuses.append(r.status)
                if "/auctions/export.xlsx" in r.url else None)
        with page.expect_download(timeout=300_000) as info:
            page.click("#auctionExport")
        target = out / "auctions.xlsx"
        info.value.save_as(str(target))
        browser.close()
    return (statuses[-1] if statuses else 200), target.read_bytes(), count


# --- прогон -----------------------------------------------------------------

def guarded(name: str, expected: str, body: Callable[[], Check]) -> Check:
    """Исключение внутри проверки — провал этой проверки с причиной, а не
    конец прогона: остальные проверки выполняются, итог пишется."""
    try:
        return body()
    except NetworkDrop as exc:
        return Check(name, FAIL, expected, "нет ответа", str(exc))
    except Exception as exc:  # noqa: BLE001 — причина уходит в итог
        first = (str(exc).splitlines() or [""])[0][:300]
        return Check(name, FAIL, expected, "ошибка проверки", f"{type(exc).__name__}: {first}")


def run(base: str, ref: dict[str, Any], admin_key: str, out: Path,
        log: Callable[[str], None]) -> list[Check]:
    checks: list[Check] = []
    out.mkdir(parents=True, exist_ok=True)

    # 1. Автозагрузка — маршрут кнопки «Найти участок»: один запрос.
    def autoload() -> Check:
        status, _, body, note = fetch("POST", base + "/land/lookup",
                                      {"query": ", ".join(ref["cadastral_numbers"]), "limit": 30},
                                      log=log)
        if status != 200:
            return with_note(Check("1. Автозагрузка участков", FAIL, "HTTP 200",
                                   f"HTTP {status}: {body[:200]!r}"), note)
        return with_note(judge_autoload(json.loads(body), ref), note)

    checks.append(guarded("1. Автозагрузка участков", "ответ маршрута", autoload))

    # Главная страница: импорт эталона, «Итог», груз для тизера и книги.
    project: dict[str, Any] | None = None
    browser_error = ""
    try:
        project = browser_project(base, ref, admin_key, out, log)
        (out / "payload.json").write_text(json.dumps(project["payload"], ensure_ascii=False),
                                          encoding="utf-8")
    except ImportError as exc:
        browser_error = f"нет playwright ({exc}); pip install playwright && playwright install chromium"
    except Exception as exc:  # noqa: BLE001
        browser_error = f"страница не собрала проект: {str(exc).splitlines()[0][:300]}"
    payload = (project or {}).get("payload") or {}

    # 2. Тизер.
    no_key = ("нужен ключ владельца: секрет DEVELOPAID_ADMIN_KEY в репозитории "
              "(в прогон — PROD_SMOKE_ADMIN_KEY)")
    if not payload:
        checks.append(Check("2. Тизер PDF", FAIL if browser_error.startswith("страница") else SKIP,
                            "PDF", "—", browser_error))
    else:
        def teaser() -> Check:
            status, ctype, body, note = fetch(
                "POST", base + "/report/teaser",
                {"session": "", "access_key": admin_key, **payload}, timeout=600, log=log)
            notes = [note] if note else []
            # Холодный тизер сервер собирает в фоне: 202 с билетом, PDF — опросом.
            deadline = time.monotonic() + 1200
            while status == 202 and time.monotonic() < deadline:
                ticket = str((json.loads(body or b"{}") or {}).get("ticket") or "")
                if not ticket:
                    break
                log(f"тизер готовится: {json.loads(body).get('detail', '')}")
                time.sleep(5)
                status, ctype, body, note = fetch("GET", base + "/report/teaser/" + ticket,
                                                  timeout=120, log=log)
                if note:
                    notes.append(note)
            note = "; ".join(notes)
            if status == 401:
                return with_note(Check("2. Тизер PDF", SKIP, "PDF", "HTTP 401 — тизер только после входа",
                                       no_key if not admin_key else "ключ не принят прод-сервером"), note)
            if body.startswith(b"%PDF"):
                (out / "teaser.pdf").write_bytes(body)
            return with_note(judge_teaser_response(status, ctype, body, ref), note)

        checks.append(guarded("2. Тизер PDF", "PDF", teaser))

    # 3. Выгрузка торгов.
    try:
        status, body, count = browser_auction_export(base, out, log)
        if status != 200:
            checks.append(Check("3. Excel-выгрузка торгов", FAIL, "HTTP 200, xlsx",
                                f"HTTP {status}: {body[:200]!r}"))
        else:
            check = judge_auction_export(body, ref)
            check.got = f"выборка на странице {count} лотов; " + check.got
            checks.append(check)
    except ImportError as exc:
        checks.append(Check("3. Excel-выгрузка торгов", SKIP, "xlsx", "—", f"нет playwright ({exc})"))
    except Exception as exc:  # noqa: BLE001
        checks.append(Check("3. Excel-выгрузка торгов", FAIL, "выборка и xlsx",
                            f"ошибка: {str(exc).splitlines()[0][:300]}"))

    # 4. Книга v4 — маршрут открыт без ключа.
    if not payload:
        checks.append(Check("4. Книга Excel v4", FAIL if browser_error.startswith("страница") else SKIP,
                            "xlsx", "—", browser_error))
    else:
        def workbook() -> Check:
            status, ctype, body, note = fetch("POST", base + "/report/workbook", {
                key: payload.get(key) for key in ("inputs", "tep", "rates", "phasing", "scenario")
            } | {"project_name": ref["project"]}, timeout=600, log=log)
            if status == 401:
                return with_note(Check("4. Книга Excel v4", SKIP, "xlsx", "HTTP 401", no_key), note)
            if status != 200 or not body.startswith(b"PK"):
                return with_note(Check("4. Книга Excel v4", FAIL, "HTTP 200, xlsx",
                                       f"HTTP {status} {ctype}: {body[:300]!r}"), note)
            (out / "workbook.xlsx").write_bytes(body)
            return with_note(judge_workbook(body, payload.get("tep") or {}, ref), note)

        checks.append(guarded("4. Книга Excel v4", "xlsx", workbook))

    # 5. «Итог».
    if not project:
        checks.append(Check("5. Страница «Итог»", FAIL if browser_error.startswith("страница") else SKIP,
                            "отрисованный «Итог»", "—", browser_error))
    elif project["itog"].get("locked"):
        checks.append(Check("5. Страница «Итог»", SKIP, "отрисованный «Итог»",
                            "экономика за входом через Telegram", no_key if not admin_key
                            else "ключ не принят страницей"))
    else:
        checks.append(guarded("5. Страница «Итог»", "отрисованный «Итог»",
                              lambda: judge_itog(project["itog"], ref)))
    return checks


def summary_markdown(checks: list[Check], base: str, prod: dict[str, Any], last_pr: str) -> str:
    version = f"{prod.get('version', '?')} ({str(prod.get('commit') or '')[:12] or 'commit пуст'})"
    failed = [c for c in checks if c.status == FAIL]
    skipped = [c for c in checks if c.status == SKIP]
    head = ("## ❌ Прод: сломано — " + ", ".join(c.name for c in failed)) if failed else \
           ("## ✅ Прод: поверхности эталона в порядке" + (f" (пропущено {len(skipped)})" if skipped else ""))
    lines = [head, "", f"Прод `{base}`, эталон — {load_reference()['project']}.", "",
             "| Проверка | Итог | Ожидалось | Получено | Версия прода | Последний слитый PR |",
             "|---|---|---|---|---|---|"]

    def cell(text: str) -> str:
        return str(text).replace("|", "\\|").replace("\n", " ")

    for c in checks:
        got = c.got + (f" — **{c.detail}**" if c.detail else "")
        lines.append(f"| {cell(c.name)} | {MARK[c.status]} | {cell(c.expected)} | {cell(got)} "
                     f"| {cell(version)} | {cell(last_pr)} |")
    if skipped:
        lines += ["", "Пропущенные проверки не зелёные: причина названа в строке."]
    return "\n".join(lines) + "\n"


def main(argv: list[str] | None = None) -> int:
    parser = argparse.ArgumentParser(description=__doc__.splitlines()[0])
    parser.add_argument("--base", default="https://developaid.ru")
    parser.add_argument("--wait-commit", default="", help="ждать, пока /health отдаст этот коммит")
    parser.add_argument("--wait-version", default="",
                        help="VERSION собранного коммита: запасная сверка, если /health не отдаёт commit")
    parser.add_argument("--wait-timeout", type=int, default=2400)
    parser.add_argument("--last-pr", default="")
    parser.add_argument("--summary", default="", help="куда записать таблицу итога (markdown)")
    parser.add_argument("--json", default="", help="куда записать результаты (json)")
    parser.add_argument("--out", default="smoke-out", help="папка для скачанных файлов")
    args = parser.parse_args(argv)

    base = args.base.rstrip("/")
    ref = load_reference()
    admin_key = os.environ.get("PROD_SMOKE_ADMIN_KEY", "").strip()
    log = lambda message: print(message, file=sys.stderr, flush=True)  # noqa: E731

    checks: list[Check] = []
    if args.wait_commit:
        arrived = wait_for_commit(base, args.wait_commit, args.wait_timeout, log, args.wait_version)
        checks.append(arrived)
    if not checks or checks[-1].status == OK:
        try:
            checks += run(base, ref, admin_key, Path(args.out), log)
        except Exception as exc:  # noqa: BLE001 — итог и json пишутся всегда
            first = (str(exc).splitlines() or [""])[0][:300]
            checks.append(Check("Прогон", FAIL, "все проверки выполнены", "прогон прерван",
                                f"{type(exc).__name__}: {first}"))
    prod = health(base)
    last_pr = args.last_pr or "не указан"
    text = summary_markdown(checks, base, prod, last_pr)
    print(text)
    for c in checks:
        if c.status == SKIP and os.environ.get("GITHUB_ACTIONS"):
            print(f"::warning title={c.name} пропущена::{c.detail}")
        if c.status == FAIL and os.environ.get("GITHUB_ACTIONS"):
            print(f"::error title={c.name}::{c.detail or c.got}")
    if args.summary:
        Path(args.summary).write_text(text, encoding="utf-8")
    if args.json:
        Path(args.json).write_text(json.dumps({
            "base": base, "prod": {k: prod.get(k) for k in ("version", "commit")},
            "last_pr": last_pr, "checks": [asdict(c) for c in checks]}, ensure_ascii=False, indent=2),
            encoding="utf-8")
    return 1 if any(c.status == FAIL for c in checks) else 0


if __name__ == "__main__":
    sys.exit(main())
