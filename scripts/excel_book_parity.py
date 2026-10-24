#!/usr/bin/env python3
"""Сверка посчитанной книги v4 с движком на эталонах и «Донском».

Формулы openpyxl не считает, поэтому книга пересчитывается LibreOffice
(headless, `calculateAll`) и читается сохранённым результатом. LibreOffice —
не Excel: расхождение, которого не видно по самой формуле, сначала проверить в
Excel, и только потом чинить книгу.

    python3 scripts/excel_book_parity.py                   # все эталоны и Донской
    python3 scripts/excel_book_parity.py --only donskoy default
    python3 scripts/excel_book_parity.py --out /tmp/books  # где оставить книги
    python3 scripts/excel_book_parity.py --only nonres_retail_sports_ddu nonres_parking_ddu

Печатает по каждой книге: статус модели, строки ПРОВЕРКИ с FAIL/WARN, ячейки
с ошибками (#REF!, #DIV/0! …) и строки ОТЧЕТа, расходящиеся с движком больше
чем на 0,01% (абсолютно — больше 1 тыс. ₽), и строки «ПРОВЕРОК» книги
нежилого проекта (`office_workbook`), кроме «сходится» и «—». Разбор последней сверки —
docs/excel_book_audit_2026-10-23.md.
"""

from __future__ import annotations

import argparse
import copy
import json
import subprocess
import sys
import tempfile
import time
from pathlib import Path
from typing import Any

ROOT = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(ROOT))
sys.path.insert(0, str(ROOT / "scripts"))

import golden_snapshot as gs  # noqa: E402

DONSKOY = ROOT / "tests" / "fixtures" / "project_settings_donskoy_0_25_78.json"
ERRORS = ("#REF!", "#DIV/0!", "#VALUE!", "#NAME?", "#N/A", "#NUM!", "#NULL!", "Err:")
MLN = 1e6

# Подпись строки ОТЧЕТ (начало) → путь в сводке движка и делитель.
REPORT_ROWS = (
    ("Выручка", "summary.revenue", MLN),
    ("CAPEX", "summary.capex", MLN),
    ("Коммерческие расходы", "summary.commercial_costs", MLN),
    ("EBITDA", "summary.ebitda", MLN),
    ("Стоимость финансирования", "summary.financing_cost", MLN),
    ("Прибыль до налога", "summary.profit_before_tax", MLN),
    ("Налог на прибыль", "summary.profit_tax", MLN),
    ("Чистая прибыль", "summary.net_profit", MLN),
    ("Маржинальность", "summary.margin", 1),
    ("IRR собственного", "summary.irr_equity", 1),
    ("NPV собственного", "summary.npv_equity", MLN),
    ("Пиковый долг", "report.financing.pf_peak", MLN),
    ("Пиковый эскроу", "report.financing.escrow_cover.peak_escrow", MLN),
    ("LLCR проекта", "summary.llcr", 1),
    ("Не погашено в дату", "summary.rve_pf_shortfall", MLN),
    ("Непогашенный долг на конец", "summary.ending_pf", MLN),
)


# Нежилой проект без офиса: ФОК по ДДУ рядом с ТЦ и наземный паркинг по ДДУ —
# книга нежилого проекта (`office_workbook`, один способ — стратегия объектов)
# против движка.
# Эталонами движка не являются: только сверка книги.
NONRES_BOOKS: dict[str, dict[str, Any]] = {
    # Нежилого проекта из одного ФОК не бывает (владелец, 24.10.2026): ФОК — рядом с ТЦ.
    "nonres_retail_sports_ddu": dict(retail_enabled=True, retail_gba_sqm=15000.0, retail_saleable_sqm=11000.0,
                                     retail_strategy="direct", sports_enabled=True, sports_gba_sqm=5000.0,
                                     sports_saleable_sqm=3500.0, sports_disposition="sale"),
    "nonres_parking_ddu": dict(above_parking_enabled=True, above_parking_spaces=300),
}


def nonres_scenario(core: Any, over: dict[str, Any]) -> dict:
    """Нежилой проект без офиса и жилья: объекты из `over`, ТЭП — от их площади."""
    x = {**copy.deepcopy(core.DEFAULT_INPUTS), "offices_enabled": False, **over,
         "project_kind": core.PROJECT_KIND_NONRESIDENTIAL}
    t = copy.deepcopy(core.TEP_DEFAULT)
    for obj in core.STANDALONE_OBJECTS:
        gba = float(x.get(f"{obj.prefix}_gba_sqm") or 0)
        if x.get(f"{obj.prefix}_enabled") and gba > 0:
            saleable = float(x.get(f"{obj.prefix}_saleable_sqm") or 0)
            t.setdefault(obj.key, {}).update(gns=gba, total_area=round(gba * 0.94, 2),
                                             useful=saleable, saleable=saleable)
    for key in core.MKD_PRODUCTS:
        for col in ("gns", "total_area", "useful", "saleable", "transfer", "units"):
            if key in t:
                t[key][col] = 0
    # Ставка метрового объекта нежилого проекта — СМР (`object_smr_rate`), как на странице.
    for key, turnkey in core.OBJECT_SMR_RATE_DEFAULTS.items():
        if key not in over and float(x.get(key) or 0) == turnkey:
            x[key] = core.object_smr_rate(turnkey, x)
    return {"inputs": x, "tep": t, "phasing": {}}


def scenarios(core: Any) -> dict[str, dict]:
    out = {name: build(core, ROOT) for name, build in gs.SCENARIOS.items()}
    saved = json.loads(DONSKOY.read_text(encoding="utf-8"))
    out["donskoy"] = {"inputs": {**copy.deepcopy(core.DEFAULT_INPUTS), **saved["inputs"]},
                      "tep": saved["tep"], "phasing": {}}
    out.update({name: nonres_scenario(core, over) for name, over in NONRES_BOOKS.items()})
    return out


def recalc(paths: list[Path]) -> list[Path]:
    """Пересчитать книги LibreOffice и сохранить рядом как *.lo.xlsx."""
    import uno  # noqa: PLC0415 — есть только там, где стоит LibreOffice
    from com.sun.star.beans import PropertyValue  # noqa: PLC0415

    def prop(name: str, value: Any) -> Any:
        item = PropertyValue()
        item.Name, item.Value = name, value
        return item

    office = subprocess.Popen(
        ["soffice", "--headless", "--invisible", "--norestore",
         "--accept=socket,host=localhost,port=2002;urp;"],
        stdout=subprocess.DEVNULL, stderr=subprocess.DEVNULL)
    try:
        local = uno.getComponentContext()
        resolver = local.ServiceManager.createInstanceWithContext(
            "com.sun.star.bridge.UnoUrlResolver", local)
        for _ in range(60):
            try:
                ctx = resolver.resolve(
                    "uno:socket,host=localhost,port=2002;urp;StarOffice.ComponentContext")
                break
            except Exception:  # noqa: BLE001 — офис ещё поднимается
                time.sleep(1)
        desktop = ctx.ServiceManager.createInstanceWithContext("com.sun.star.frame.Desktop", ctx)
        out = []
        for path in paths:
            target = path.with_name(path.stem + ".lo.xlsx")
            doc = desktop.loadComponentFromURL(
                uno.systemPathToFileUrl(str(path)), "_blank", 0, (prop("Hidden", True),))
            doc.calculateAll()
            doc.storeToURL(uno.systemPathToFileUrl(str(target)),
                           (prop("FilterName", "Calc MS Excel 2007 XML"),))
            doc.close(True)
            out.append(target)
        try:
            desktop.terminate()
        except Exception:  # noqa: BLE001 — офис закрывается сам
            pass
        return out
    finally:
        office.wait(timeout=30)


def _path(data: Any, path: str) -> Any:
    for part in path.split("."):
        data = data.get(part) if isinstance(data, dict) else None
    return data


def report(name: str, book_path: Path, consolidated: dict) -> list[str]:
    import openpyxl  # noqa: PLC0415

    book = openpyxl.load_workbook(book_path, data_only=True)
    lines = []
    errors = [f"{sheet.title}!{cell.coordinate}={cell.value}"
              for sheet in book.worksheets for row in sheet.iter_rows() for cell in row
              if isinstance(cell.value, str) and cell.value.startswith(ERRORS)]
    if errors:
        lines.append(f"  ошибки в ячейках: {len(errors)}: {errors[:6]}")
    if "ПРОВЕРКИ" in book.sheetnames and book["ПРОВЕРКИ"]["I4"].value == "Вердикт":  # книга нежилого проекта
        checks = book["ПРОВЕРКИ"]
        lines.append(f"  проверки книги нежилого проекта: {checks['I5'].value}")
        for row in range(5, checks.max_row + 1):
            verdict = checks.cell(row, 7).value
            if verdict not in ("сходится", "—", None):
                lines.append(f"  ПРОВЕРКИ!G{row} {checks.cell(row, 1).value}: {checks.cell(row, 2).value}: "
                             f"книга {checks.cell(row, 3).value!r}, движок {checks.cell(row, 4).value!r}")
        return lines
    if "ПРОВЕРКИ" in book.sheetnames:
        checks = book["ПРОВЕРКИ"]
        lines.append(f"  статус модели: {checks['B3'].value}")
        for row in range(4, checks.max_row + 1):
            verdict = checks.cell(row, 6).value
            if verdict in ("FAIL", "WARN"):
                lines.append(f"  {verdict}: {checks.cell(row, 1).value}: "
                             f"{checks.cell(row, 2).value} против {checks.cell(row, 3).value}")
    sheet = book["ОТЧЕТ"]
    rows = {}
    for row in range(1, 120):
        label = sheet.cell(row, 1).value
        if isinstance(label, str):
            rows.setdefault(label.strip(), (row, sheet.cell(row, 2).value))
    return lines + compare_report(rows, consolidated)


def compare_report(rows: dict[str, tuple[int, Any]], consolidated: dict) -> list[str]:
    """Строки ОТЧЕТ (подпись → (строка, значение B)) против сводки движка в допуске.

    Одна мера и для пересчёта LibreOffice, и для `xlsx_eval` в тестах.
    """
    lines = []
    for label, path, scale in REPORT_ROWS:
        found = next((value for key, value in rows.items() if key.startswith(label)), None)
        engine = _path(consolidated, path)
        if found is None:
            lines.append(f"  нет строки ОТЧЕТ «{label}»")
            continue
        row, value = found
        if not isinstance(value, (int, float)) or not isinstance(engine, (int, float)):
            if (value is None) != (engine is None) and not (isinstance(value, str) and engine is None):
                lines.append(f"  ОТЧЕТ!B{row} {label}: книга {value!r}, движок {engine!r}")
            continue
        expected = engine / scale
        delta = value - expected
        if abs(delta) > max(1e-3 if scale == MLN else 1e-6, abs(expected) * 1e-4):
            lines.append(f"  ОТЧЕТ!B{row} {label}: книга {value:,.3f}, движок {expected:,.3f}, "
                         f"Δ {delta:,.3f}")
    return lines


def main() -> int:
    parser = argparse.ArgumentParser(description=__doc__.split("\n\n")[0])
    parser.add_argument("--only", nargs="*")
    parser.add_argument("--out", type=Path)
    args = parser.parse_args()
    core = gs.load_core(ROOT)
    specs = scenarios(core)
    names = args.only or list(specs)
    out = args.out or Path(tempfile.mkdtemp(prefix="excel_book_parity_"))
    out.mkdir(parents=True, exist_ok=True)
    built = {}
    for name in names:
        spec = specs[name]
        content, _file, meta = core.build_project_workbook(
            copy.deepcopy(spec["inputs"]), copy.deepcopy(spec["tep"]), [],
            copy.deepcopy(spec["phasing"]), project_name=name)
        path = out / f"{name}.xlsx"
        path.write_bytes(content)
        built[name] = path
        if meta.get("missing"):
            print(f"{name}: не записано в книгу: {meta['missing']}")
    recalculated = dict(zip(built, recalc(list(built.values()))))
    for name, path in recalculated.items():
        spec = specs[name]
        bundle = core._run_authoritative_model(
            copy.deepcopy(spec["inputs"]), copy.deepcopy(spec["tep"]), [],
            copy.deepcopy(spec["phasing"]))
        print(f"===== {name} ({path})")
        for line in report(name, path, bundle["consolidated"]) or ["  расхождений нет"]:
            print(line)
    return 0


if __name__ == "__main__":
    sys.exit(main())
