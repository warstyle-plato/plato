#!/usr/bin/env python3
"""Мост IRR гостиницы: расчёт движка на ориентире «Отель 1» → модель Отеля 1.

Задача владельца #654: разложить разницу IRR проекта 20,5 → 17,68 % и IRR
капитала 57,8 → 34,6 % по причинам, с адресами ячеек книги. Мост идёт шагами,
каждый шаг меняет одно — вводную или правило расчёта — на то, что в книге
`docs/reference/hotel/dombai_ski_hotel_5star_2025.xlsm`; Δ шага — вклад
причины в п.п. Шаги накопительные: вклад причины зависит от порядка, порядок
назван здесь и в `docs/hotel_irr_vs_hotel1.md`.

Методика движка не меняется: правила книги (Гордон, FCFF без процентов,
капитал первым, годовой IRR) применяются к потокам движка здесь, в мосте.

    python3 scripts/hotel_irr_bridge.py     # напечатать таблицу моста
"""

from __future__ import annotations

import copy
import datetime as dt
import sys
import warnings
from functools import lru_cache
from pathlib import Path
from typing import Any, Callable

ROOT = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(ROOT))

import developaid_finance_math as fm  # noqa: E402
import developaid_hotel_strategy as hs  # noqa: E402
import hotel_reference as ref  # noqa: E402

# Гордон книги: Расчеты!E1634 = Инвестор!I19 (темп роста после прогноза) и
# Required_Rate_of_Return = Расчеты!E1623 (стоимость капитала в рублях).
GORDON_GROWTH_CELL = "Расчеты!E1634"
REQUIRED_RETURN_CELL = "Расчеты!E1623"
# Строки «Расчетов» книги (квартал в столбце, даты — строка 7, тыс. руб.).
ROWS = {
    "hotel_capex": 212,     # CAPEX отеля без НДС (Контр_панель!G66 = I212)
    "infra_capex": 170,     # коммунальная инфраструктура (Контр_панель!G71 = I170)
    "ffe_capex": 207,       # мебель и оборудование
    "fcff": 1640,           # FCFF = OCF + проценты × (1 − t) + ICF (1637:1639)
    "fcff_tv": 1641,        # терминальная стоимость проекта (Гордон)
    "fcfe": 1677,           # FCFE: капитал, займ, дивиденды (1673:1676)
    "fcfe_tv": 1678,        # терминальная стоимость для акционеров
}
BOOK_IRR_PROJECT = "hotel1:Контр_панель!G14"   # = Расчеты!E1651, IRR(L1645:AY1645)
BOOK_IRR_EQUITY = "Контр_панель!G20"           # = Расчеты!E1689, IRR(L1682:AY1682)


@lru_cache(maxsize=None)
def book() -> dict[str, Any]:
    """Квартальные ряды книги (руб.), ставки Гордона и ключевая ставка."""
    import openpyxl
    with warnings.catch_warnings():
        warnings.simplefilter("ignore")
        wb = openpyxl.load_workbook(ref.workbook_path(ref.DOMBAI), data_only=True,
                                    read_only=True)
    try:
        calc = wb["Расчеты"]
        head = next(calc.iter_rows(min_row=7, max_row=7, values_only=True))
        rows: dict[str, dict[dt.date, float]] = {}
        for name, n in ROWS.items():
            values = next(calc.iter_rows(min_row=n, max_row=n, values_only=True))
            rows[name] = {head[c].date(): float(values[c] or 0.0) * 1000.0
                          for c in range(11, len(head)) if isinstance(head[c], dt.datetime)}
        growth = float(calc["E1634"].value)
        required = float(calc["E1623"].value)
        irr_project = float(calc["E1651"].value)
        irr_equity = float(calc["E1689"].value)
        pre = list(wb["Предпосылки"].iter_rows(min_row=56, max_row=60, max_col=40,
                                                values_only=True))
    finally:
        wb.close()
    key_rate = {pre[0][c]: float(pre[4][c]) for c in range(len(pre[0]))
                if isinstance(pre[0][c], int) and pre[4][c] is not None}
    return {"rows": rows, "growth": growth, "required": required, "key_rate": key_rate,
            "irr_project": irr_project, "irr_equity": irr_equity}


def annual_irr(flows: list[float]) -> float | None:
    """IRR годового ряда, как IRR() Excel над годовыми суммами."""
    def npv(rate: float) -> float:
        return sum(cf / (1.0 + rate) ** i for i, cf in enumerate(flows))
    lo, hi = -0.9, 10.0
    if npv(lo) * npv(hi) > 0:
        return None
    for _ in range(200):
        mid = (lo + hi) / 2
        if npv(lo) * npv(mid) <= 0:
            hi = mid
        else:
            lo = mid
    return (lo + hi) / 2


def by_year(months: list[dt.date], flows: list[float]) -> list[float]:
    out: dict[int, float] = {}
    for month, value in zip(months, flows):
        out[month.year] = out.get(month.year, 0.0) + value
    return [out[y] for y in sorted(out)]


def _monthly(quarters: dict[dt.date, float], scale: float = 1.0) -> dict[dt.date, float]:
    """Квартал книги → три месяца поровну."""
    out: dict[dt.date, float] = {}
    for start, value in quarters.items():
        for k in range(3):
            month = hs.add_months(start, k)
            out[month] = out.get(month, 0.0) + value * scale / 3.0
    return out


def engine_plan() -> dict[str, Any]:
    """План, который движок отдаёт гостинице на ориентире «Отель 1» с умолчаниями
    проекта: CAPEX по месяцам, даты, ключевая ставка, ставки налога и дисконта."""
    sys.path.insert(0, str(ROOT))
    import hotel_presets
    import main as wrapper
    core = wrapper.core
    x = copy.deepcopy(core.DEFAULT_INPUTS)
    t = copy.deepcopy(core.TEP_DEFAULT)
    x["project_kind"] = core.PROJECT_KIND_HOTEL
    x = hotel_presets.apply_preset(x, "hotel1")
    seen: dict[str, Any] = {}
    original = core.hotel_strategy.hotel_flows

    def spy(plan, key_rate, **kw):
        result = original(plan, key_rate, **kw)
        seen.update(plan=plan, kw=kw, rates={m: key_rate(m) for m in result["months"]},
                    kpi=result["kpi"])
        return result

    core.hotel_strategy.hotel_flows = spy
    try:
        core.calculate(core.CalcRequest(inputs=x, tep=t, rates=[]))
    finally:
        core.hotel_strategy.hotel_flows = original
    return seen


def bridge(plan: dict[str, Any] | None = None, skip: frozenset[str] = frozenset()
           ) -> list[dict[str, Any]]:
    """Шаги моста: подпись, ячейки книги, IRR проекта и капитала после шага.

    `skip` — пропустить причину («infra», «equity_first»): для проверки, что
    мост без неё до книги не сходится."""
    seen = plan or engine_plan()
    b = book()
    vat = 1.0 + ref.value("hotel1:Предпосылки!E775")
    kw = seen["kw"]
    rates = seen["rates"]
    first, last = min(rates), max(rates)

    def engine_rate(month: dt.date) -> float:
        return rates.get(month, rates[first] if month < first else rates[last])

    def book_rate(month: dt.date) -> float:
        return b["key_rate"].get(month.year, b["key_rate"][max(b["key_rate"])])

    def flows(p: dict[str, Any], rate: Callable[[dt.date], float], **over: Any) -> dict:
        p = dict(p)
        p["params"] = {**p["params"], **over}
        return hs.hotel_flows(p, rate, **kw)

    steps: list[dict[str, Any]] = []

    def step(key: str, label: str, cells: str, project: float | None,
             equity: float | None) -> None:
        steps.append({"key": key, "label": label, "cells": cells,
                      "irr_project": project, "irr_equity": equity})

    plan0 = seen["plan"]
    r = flows(plan0, engine_rate)
    step("engine", "Движок: ориентир «Отель 1», даты и смета проекта движка", "—",
         r["kpi"]["irr_project"], r["kpi"]["irr_equity"])

    # Даты книги: ввод 01.01.2029 (Предпосылки!E24) вместо 01.07.2030 движка.
    shift = hs.month_index(plan0["commissioning"], dt.date(2029, 1, 1))
    moved = {name: {hs.add_months(m, shift): v for m, v in plan0[name].items()}
             for name in ("capex", "capex_ffe", "capex_land")}
    plan1 = {**plan0, **moved, "start": hs.add_months(plan0["start"], shift),
             "commissioning": dt.date(2029, 1, 1)}
    r = flows(plan1, engine_rate)
    step("dates", "Даты книги: ввод 01.2029, график стройки сдвинут целиком",
         "Предпосылки!E24", r["kpi"]["irr_project"], r["kpi"]["irr_equity"])

    r = flows(plan1, book_rate)
    step("key_rate", "Ключевая ставка книги по годам", "Предпосылки!M60:AD60",
         r["kpi"]["irr_project"], r["kpi"]["irr_equity"])

    rows = b["rows"]
    hotel = _monthly(rows["hotel_capex"], vat)
    plan3 = {**plan1, "start": dt.date(2025, 10, 1), "capex": hotel,
             "capex_ffe": _monthly(rows["ffe_capex"], vat), "capex_land": {}}
    r = flows(plan3, book_rate)
    step("capex", "CAPEX отеля книги по кварталам вместо сметы движка",
         "Расчеты!L212:CO212 (Контр_панель!G66), мебель — Расчеты!207",
         r["kpi"]["irr_project"], r["kpi"]["irr_equity"])

    capex = dict(hotel)
    for month, value in _monthly(rows["infra_capex"], vat).items():
        if "infra" not in skip:
            capex[month] = capex.get(month, 0.0) + value
    plan4 = {**plan3, "capex": capex}
    r = flows(plan4, book_rate)
    step("infra", "Коммунальная инфраструктура 1 371 млн ₽ без НДС",
         "Расчеты!L170:CO170 (Контр_панель!G71)",
         r["kpi"]["irr_project"], r["kpi"]["irr_equity"])

    # Выход книги — Гордон: поток последних 4 кварталов × (1 + g) ÷ (r − g)
    # в последнем квартале эксплуатации; продажи и её затрат нет.
    r = flows(plan4, book_rate, exit_mode=hs.EXIT_HOLD)
    months = r["months"]
    end = months.index(r["horizon_end"])
    residual = r["kpi"]["exit_value"]
    factor = (1.0 + b["growth"]) / (b["required"] - b["growth"])

    def gordon(series: list[float]) -> list[float]:
        out = list(series)
        out[end] -= residual
        out[end] += sum(out[end - 11:end + 1]) * factor
        return out

    project = gordon(r["project_cf"])
    equity = gordon(r["equity_cf"])
    step("exit", "Выход по Гордону вместо продажи по EV/EBITDA 4×",
         f"Расчеты!1641, 1678; g = {GORDON_GROWTH_CELL}, r = {REQUIRED_RETURN_CELL}",
         fm.monthly_irr(project), fm.monthly_irr(equity))

    # Поток проекта книги — FCFF: OCF + проценты × (1 − t) + ICF. У движка
    # поток проекта — после процентов и комиссий кредита.
    fcff = gordon([r["monthly"]["fcff"].get(m, 0.0) for m in months])
    step("fcff", "Поток проекта — FCFF без процентов и комиссий кредита",
         "Расчеты!1637:1640", fm.monthly_irr(fcff), fm.monthly_irr(equity))
    tax = kw["profit_tax_rate"]
    fcff = [v - tax * r["monthly"]["loan_interest_paid"].get(m, 0.0)
            for v, m in zip(fcff, months)]
    step("shield", "FCFF без налогового щита процентов: проценты × (1 − t)",
         "Расчеты!1638", fm.monthly_irr(fcff), fm.monthly_irr(equity))

    # Финансирование стройки: книга вкладывает капитал первым (L1381), дальше
    # платит кредит (1382:1383); у движка капитал — доля каждого платежа.
    before = _monthly(rows["fcfe"])
    if "equity_first" not in skip:
        equity = [before.get(m, 0.0) if m < r["commissioning"] else v
                  for m, v in zip(months, equity)]
    step("equity_first", "Капитал первым, затем кредит — поток капитала книги до ввода",
         "Расчеты!L1381:L1383, 1491, 1677", fm.monthly_irr(fcff), fm.monthly_irr(equity))

    step("annual", "IRR по годовым суммам вместо месячного ряда",
         "Расчеты!1645, 1651; 1682, 1689",
         annual_irr(by_year(months, fcff)), annual_irr(by_year(months, equity)))
    step("book", "Модель Отеля 1", "Контр_панель!G14, G20",
         b["irr_project"], b["irr_equity"])
    return steps


def table(steps: list[dict[str, Any]]) -> str:
    lines = ["| Шаг | Ячейки книги | IRR проекта | Δ, п.п. | IRR капитала | Δ, п.п. |",
             "|---|---|---:|---:|---:|---:|"]
    prev = None
    for s in steps:
        p, e = s["irr_project"] * 100, s["irr_equity"] * 100
        dp = "" if prev is None else f"{p - prev[0]:+.2f}".replace(".", ",")
        de = "" if prev is None else f"{e - prev[1]:+.2f}".replace(".", ",")
        pp, ee = (f"{v:.2f} %".replace(".", ",") for v in (p, e))
        lines.append(f"| {s['label']} | {s['cells']} | {pp} | {dp} | {ee} | {de} |")
        prev = (p, e)
    return "\n".join(lines)


if __name__ == "__main__":
    print(table(bridge()))
