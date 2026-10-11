"""Смета гостиницы: базы и графики движка, выраженные формулами Excel."""
from __future__ import annotations

from openpyxl.utils import get_column_letter

INPUT_DEFAULTS = {
    "hotel_gba_sqm": 0.0, "hotel_cost_th_per_sqm": 0.0,
    "hotel_ffe_th_per_key": 0.0, "ird_months": 18.0, "construction_months": 24.0,
    "ird_th_per_sqm": 0.0, "design_p_th_per_sqm": 0.0, "design_rd_th_per_sqm": 0.0,
    "preparation_th_per_sqm": 0.0, "utilities_th_per_sqm": 0.0,
    "commissioning_th_per_sqm": 0.0, "site_maintenance_th_per_sqm": 0.0,
    "demolition_area_sqm": 0.0, "demolition_cost_th_per_sqm": 0.0,
    "resettlement_cost_mln": 0.0, "author_supervision_pct": 0.0,
    "technical_supervision_pct": 0.0, "project_management_pct": 5.0,
    "gc_fee_pct": 0.0, "reserve_pct": 0.0, "purchase_price_mln": 0.0,
    "land_buyout_mln": 0.0, "scenario_cost_multiplier": 1.0,
}
LABELS = {
    "hotel_gba_sqm": "ГНС гостиницы, м²", "hotel_cost_th_per_sqm": "СМР и отделка, тыс. ₽/м² ГНС",
    "hotel_ffe_th_per_key": "FF&E, тыс. ₽/номер", "ird_months": "ИРД до РнС, мес.",
    "construction_months": "Строительство, мес.", "ird_th_per_sqm": "ИРД, тыс. ₽/м² ГНС",
    "design_p_th_per_sqm": "Проект (П), тыс. ₽/м² ГНС",
    "design_rd_th_per_sqm": "Рабочая документация, тыс. ₽/м² ГНС",
    "preparation_th_per_sqm": "Подготовка территории, тыс. ₽/м² ГНС",
    "utilities_th_per_sqm": "Сети, тыс. ₽/м² ГНС",
    "commissioning_th_per_sqm": "Ввод, тыс. ₽/м² ГНС",
    "site_maintenance_th_per_sqm": "Содержание площадки, тыс. ₽/м² ГНС",
    "demolition_area_sqm": "Площадь сноса, м²",
    "demolition_cost_th_per_sqm": "Снос, тыс. ₽/м²",
    "resettlement_cost_mln": "Расселение, млн ₽",
    "author_supervision_pct": "Авторский надзор, % от П + РД",
    "technical_supervision_pct": "Техзаказчик, % от СМР",
    "project_management_pct": "Управление проектом, % от базы движка",
    "gc_fee_pct": "Генподрядчик, % от СМР",
    "reserve_pct": "Резерв, % от базы движка",
    "purchase_price_mln": "Покупка участка, млн ₽",
    "land_buyout_mln": "Выкуп ЗУ/ОКС, млн ₽",
    "scenario_cost_multiplier": "Множитель затрат сценария",
}
ARTICLES = (
    "purchase", "land_rights", "ird", "design_p", "design_rd", "demolition",
    "resettlement", "preparation", "hotel", "hotel_ffe", "utilities",
    "site_maintenance", "commissioning", "author_supervision",
    "technical_supervision", "project_management", "gc_fee", "reserve",
)
SHEET = "Затраты"
FIRST = 32
INPUT_FIRST = 76
INPUT_ROWS = {key: INPUT_FIRST + i for i, key in enumerate(INPUT_DEFAULTS)}


def ref(key: str) -> str:
    return f"'Вводные'!$C${INPUT_ROWS[key]}"


def inputs(book, spec, hw) -> None:
    ws = book[hw.INPUTS_SHEET]
    ws[f"A{INPUT_FIRST - 1}"] = "СМЕТА И КАЛЕНДАРЬ — редактируемые вводные"
    ws[f"A{INPUT_FIRST - 1}"].font = hw.BOLD
    for key, row in INPUT_ROWS.items():
        ws[f"A{row}"] = LABELS[key]
        ws[f"C{row}"] = spec["inputs"][key]
        ws[f"C{row}"].fill = hw.INPUT_FILL
    ws["E76"] = "Платежи за участок: месяц от старта и доля"
    plan = spec.get("purchase_plan") or [(0, 1.0)]
    editable_plan = plan + [(0, 0.0)] * max(0, 12 - len(plan))
    for row, (month, share) in enumerate(editable_plan, 77):
        ws[f"F{row}"], ws[f"G{row}"] = month, share
        ws[f"F{row}"].fill = ws[f"G{row}"].fill = hw.INPUT_FILL
        ws[f"G{row}"].number_format = "0.0%"
    ws["E74"] = "Число строк графика участка"
    ws["G74"] = f'=COUNTIF(G77:G{76+len(editable_plan)},">0")'


def build(book, spec, rows, hw):
    """Постатейная смета и денежный график; возвращает ссылки для расчёта."""
    ws = book.create_sheet(SHEET)
    ws["A1"] = "Затраты гостиницы — смета и график оплаты"
    ws["A2"] = "Все суммы с НДС. Жёлтые ячейки — вводные; смета и график связаны формулами."
    for col, label in zip("ABCDEFGH", ("Статья", "База", "Ед.", "Ставка", "Сумма, ₽",
                                           "Сценарий, ₽", "Начало (k)", "Срок, мес.")):
        ws[f"{col}4"] = label
    p = ref
    i = hw._in
    permit = f"MAX({spec['ird_min']},INT({p('ird_months')}))"
    construction = f"MAX(1,INT({p('construction_months')}))"
    rve = f"({permit}+INT({p('construction_months')}))"
    window = f"MIN(6,{permit})"
    ffe = f"MAX(1,MIN({spec['ffe_months']},{construction}))"
    # (база, единица, ставка, сумма; начало, срок, профиль)
    definitions = {}
    def article(key, label, base, unit, rate, amount, start, length, profile):
        definitions[key] = (label, base, unit, rate, amount, start, length, profile)
    article("purchase", "Участок и выкуп ЗУ/ОКС", "1", "млн ₽", f"{p('purchase_price_mln')}+{p('land_buyout_mln')}",
            "B{r}*D{r}*1000000", "0", "1", "purchase")
    # Статутная льгота уже определена движком. Её сумма — вводная, а не новая
    # формула права в Excel. Единовременный денежный график книги — формула.
    ws["K2"] = "Плата за ВРИ после льгот, ₽ (исходная сумма)"
    ws["L2"] = spec["vri_net"]
    ws["L2"].fill = hw.INPUT_FILL
    article("land_rights", "Плата за ВРИ после льгот", "1", "₽", "$L$2", "B{r}*D{r}", permit, "1", "single")
    area = p("hotel_gba_sqm")
    for key, label, rate, start, length, profile in (
        ("ird", "ИРД", "ird_th_per_sqm", "0", permit, "even"),
        ("design_p", "Проект (П)", "design_p_th_per_sqm", f"{permit}-{window}", window, "even"),
        ("design_rd", "Рабочая документация", "design_rd_th_per_sqm", f"{permit}-{window}", window, "even"),
        ("preparation", "Подготовка территории", "preparation_th_per_sqm", f"{permit}-{window}", window, "even"),
        ("hotel", "СМР и отделка здания", "hotel_cost_th_per_sqm", permit, construction, "curve"),
        ("utilities", "Сети", "utilities_th_per_sqm", permit, construction, "curve"),
        ("site_maintenance", "Содержание площадки", "site_maintenance_th_per_sqm", permit, construction, "curve"),
        ("commissioning", "Ввод в эксплуатацию", "commissioning_th_per_sqm", f"{rve}-3", "3", "even"),
    ):
        article(key, label, area, "м² ГНС", p(rate), "B{r}*D{r}*1000", start, length, profile)
    article("demolition", "Снос", p("demolition_area_sqm"), "м² сноса", p("demolition_cost_th_per_sqm"),
            "B{r}*D{r}*1000", f"{permit}-{window}", window, "even")
    article("resettlement", "Расселение", "1", "млн ₽", p("resettlement_cost_mln"),
            "B{r}*D{r}*1000000", f"{permit}-{window}", window, "even")
    article("hotel_ffe", "Мебель и оборудование (FF&E)", i("keys"), "номеров", p("hotel_ffe_th_per_key"),
            "B{r}*D{r}*1000", f"{rve}-{ffe}", ffe, "even")
    positions = {key: 5 + index for index, key in enumerate(ARTICLES)}
    def amount(key):
        return f"$E${positions[key]}"
    design = f"{amount('design_p')}+{amount('design_rd')}"
    article("author_supervision", "Авторский надзор", design, "П + РД, ₽", f"{p('author_supervision_pct')}/100",
            "B{r}*D{r}", permit, construction, "curve")
    for key, label, rate in (("technical_supervision", "Технический заказчик", "technical_supervision_pct"),
                              ("gc_fee", "Генподрядчик", "gc_fee_pct")):
        article(key, label, amount("hotel"), "СМР, ₽", f"{p(rate)}/100", "B{r}*D{r}", permit, construction, "curve")
    management = [key for key in spec["management_profile"] if key in positions]
    article("project_management", "Управление проектом", "+".join(amount(key) for key in management),
            "База движка, ₽", f"{p('project_management_pct')}/100", "B{r}*D{r}", "0", rve, "management")
    reserve_base = [key for key in ARTICLES if key not in
                    {*spec["reserve_excluded"], "purchase", "reserve"}]
    article("reserve", "Резерв", "+".join(amount(key) for key in reserve_base), "База движка, ₽",
            f"{p('reserve_pct')}/100", "B{r}*D{r}", permit, construction, "curve")
    for key in ARTICLES:
        r = positions[key]
        label, base, unit, rate, total, start, length, _ = definitions[key]
        ws[f"A{r}"], ws[f"C{r}"] = label, unit
        for col, expr in (("B", base), ("D", rate), ("E", total.format(r=r)),
                          ("F", f"E{r}*{p('scenario_cost_multiplier')}"), ("G", start), ("H", length)):
            ws[f"{col}{r}"] = "=" + expr
    ws["A24"] = "Итого CAPEX"
    ws["F24"] = "=SUM(F5:F22)"
    # Дополнительные статьи не скрываются за «сошлось»: они поименованы,
    # лежат в отдельном исходном графике и перечислены в ограничениях книги.
    engine = {key: schedule for key, schedule in spec.get("engine_articles", {}).items() if schedule}
    if engine:
        ws["A26"] = "Дополнительные статьи: исходный график движка (не пересчитывается сметой)"
    monthly_keys = list(ARTICLES) + [key for key in engine if key not in ARTICLES]
    letters = {key: get_column_letter(3 + index) for index, key in enumerate(monthly_keys)}
    total_col = get_column_letter(3 + len(monthly_keys))
    ws[f"A{FIRST-1}"], ws[f"B{FIRST-1}"] = "k", "Месяц"
    for key, col in letters.items():
        ws[f"{col}{FIRST-1}"] = definitions[key][0] if key in definitions else spec.get("article_labels", {}).get(key, key)
    ws[f"{total_col}{FIRST-1}"] = "Итого CAPEX"
    plan_last = 76 + max(12, len(spec.get("purchase_plan") or [(0, 1.0)]))
    plan_k = f"'Вводные'!$F$77:$F${plan_last}"
    plan_share = f"'Вводные'!$G$77:$G${plan_last}"
    def col_range(key):
        return f"${letters[key]}${FIRST}:${letters[key]}${FIRST + rows - 1}"
    # База профиля управления — одна сумма графика на всю сетку, а не своя
    # в каждой строке: 7 SUM по всей сетке в каждом месяце стоили выгрузке
    # десятки секунд и Excel — столько же лишнего счёта.
    management_row = positions["project_management"]
    management_total = f"$I${management_row}"
    ws["I4"] = "График базы, ₽"
    ws[management_total.replace("$", "")] = "=" + ("+".join(f"SUM({col_range(p)})" for p in management) or "0")
    # Профиль управления исключает здание гостиницы: так задан движок.
    # Его исправление было бы изменением методики, а не ремонтом выгрузки.
    for k in range(rows):
        r = FIRST + k
        ws[f"A{r}"] = k
        ws[f"B{r}"] = f"='{hw.CALC_SHEET}'!{hw._c('month')}{hw.FIRST_ROW+k}"
        ws[f"B{r}"].number_format = "mmm yyyy"
        for key in monthly_keys:
            cell = f"{letters[key]}{r}"
            if key in engine:
                ws[cell] = float(engine[key].get(k, 0.0))
                ws[cell].fill = hw.ENGINE_FILL
                continue
            pos = positions[key]
            amount_ref, start_ref, length_ref = f"$F${pos}", f"$G${pos}", f"$H${pos}"
            kk = f"$A{r}"
            inside = f"AND({kk}>={start_ref},{kk}<{start_ref}+{length_ref})"
            profile = definitions[key][-1]
            if profile == "purchase":
                formula = f"{amount_ref}*SUMIF({plan_k},{kk},{plan_share})"
            elif profile == "single":
                formula = f"IF({kk}={start_ref},{amount_ref},0)"
            elif profile == "even":
                formula = f"IF({inside},{amount_ref}/MAX(1,{length_ref}),0)"
            elif profile == "curve":
                # Те же ступени build_curve.monthly_weights, что у движка.
                idx = f"({kk}-{start_ref}+1)"
                weight = f"IF({idx}/{length_ref}<=0.2,0.6,IF({idx}/{length_ref}<=0.8,1.2,0.8))"
                denominator = (f"(0.6*INT({length_ref}*0.2)+1.2*(INT({length_ref}*0.8)"
                               f"-INT({length_ref}*0.2))+0.8*({length_ref}-INT({length_ref}*0.8)))")
                formula = f"IF({inside},{amount_ref}*{weight}/{denominator},0)"
            else:
                base = "+".join(f"{letters[p]}{r}" for p in management)
                total = management_total
                formula = f"IF({total}>0,{amount_ref}*({base})/{total},IF({inside},{amount_ref}/MAX(1,{length_ref}),0))"
            ws[cell] = "=" + formula
        ws[f"{total_col}{r}"] = "=" + "+".join(f"{col}{r}" for col in letters.values())
    ws["F24"] = f"=SUM({total_col}{FIRST}:{total_col}{FIRST+rows-1})"
    ws.freeze_panes = "B5"
    ws.column_dimensions["A"].width = 44
    for col in range(2, 9):
        ws.column_dimensions[get_column_letter(col)].width = 18
    def month_ref(key, index):
        if key not in letters and key != "total":
            return "0"
        return f"'{SHEET}'!{letters[key] if key in letters else total_col}{FIRST+index}"
    return month_ref, list(engine)
