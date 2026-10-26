"""Образец книги офиса на каркасе жилой v4: месяцы вправо, расчёт формулами.

Формулы повторяют `developaid_nonres_strategy.object_flows` (продажа, аренда,
выход, НДС, кредит, деньги собственника). Из движка значениями приходят только
график затрат, ключевая ставка и налог на прибыль — в рабочей книге у них свои
листы формул, как в жилой v4.
"""
import re, sys
from datetime import date
sys.path.insert(0, "/tmp/claude-0/-home-user-plato/c936f76f-bdf5-5950-84e0-841f7b792428/scratchpad")
import run_office as ro, collect
from openpyxl import Workbook
from openpyxl.chart import BarChart, LineChart, Reference
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter as L
from openpyxl.worksheet.properties import PageSetupProperties

OUT = sys.argv[1]
# --- движок ------------------------------------------------------------------
RES = {p: ro.run(p, **collect.BASE) for p in ("direct", "income")}
OBJ = {p: RES[p]["finance"]["nonres"]["objects"][0] for p in RES}
START = date(2027, 1, 1); N = 114  # 01.2027 … 06.2036: год после продажи здания — для NOI покупателя


def mi(iso):
    d = date.fromisoformat(str(iso)[:10]); return (d.year - START.year) * 12 + d.month - START.month


def series(path, key, src="obj"):
    out = [0.0] * N
    if src == "cf":
        cf = RES[path]["cashflow"]
        for when, v in zip(cf["months"], cf[key]):
            k = mi(when)
            if 0 <= k < N:
                out[k] = float(v or 0.0)
        return out
    rows = OBJ[path]["rows"] if src == "obj" else RES[path]["finance"]["rows"]
    for r in rows:
        k = mi(r["month"])
        if 0 <= k < N:
            out[k] = float(r.get(key) or 0.0)
    return out


KEY = series("income", "key_rate", "fin")
last = max(i for i, v in enumerate(KEY) if v)
KEY = [v if i <= last else KEY[last] for i, v in enumerate(KEY)]

# --- стиль жилой v4 ---------------------------------------------------------------
F = "Carlito"
DARK, HEAD, TILE, LIGHT, INPUT, GREY = "0B1F33", "1F4E78", "17365D", "EAF2F8", "FFF2CC", "666666"
SALE, RENT = "2F6DB5", "C26A0A"
NUM = '#,##0.0;[Red](#,##0.0);-'; NUM0 = '#,##0;[Red](#,##0);-'; PCT = '0.0%;[Red]-0.0%;-'; PCT2 = '0.00%'
thin = Side(style="thin", color="C9D3DE")
fnt = lambda **k: Font(name=F, size=11, **k)
fill = lambda c: PatternFill("solid", fgColor=c)

wb = Workbook(); wb.remove(wb.active)
REG = {}  # имя ряда → (лист, строка)
PENDING = []  # формулы пишутся после разметки: ряды ссылаются вперёд


def sheet(name, title, sub, monthly=False, tab=None):
    ws = wb.create_sheet(name)
    ws.sheet_view.showGridLines = False
    if tab:
        ws.sheet_properties.tabColor = tab
    ws["A1"] = title; ws["A1"].font = fnt(b=True, color="FFFFFF")
    for c in range(1, (N + 4 if monthly else 8) + 1):
        ws.cell(1, c).fill = fill(DARK)
    ws["A2"] = sub; ws["A2"].font = fnt(color=GREY)
    ws.sheet_properties.pageSetUpPr = PageSetupProperties(fitToPage=True)
    ws.page_setup.orientation = "landscape"; ws.page_setup.fitToWidth = 1; ws.page_setup.fitToHeight = 0
    if monthly:
        for c, v in enumerate(["Показатель", "Итого / вводная", "Ед."], 1):
            x = ws.cell(3, c, v); x.font = fnt(b=True, color="FFFFFF"); x.fill = fill(HEAD)
        ws.cell(4, 1, "Период"); ws.cell(4, 3, "№")
        for k in range(N):
            col = 4 + k
            d = ws.cell(3, col, "=СРОКИ!$D$3" if name != "СРОКИ" else None)
            if name == "СРОКИ":
                d.value = f"=DATE({START.year},{START.month},1)" if k == 0 else f"=EDATE({L(col - 1)}3,1)"
            else:
                d.value = f"=СРОКИ!{L(col)}3"
            d.number_format = "mmm yyyy"; d.font = fnt(b=True, color="FFFFFF"); d.fill = fill(HEAD)
            i = ws.cell(4, col, k); i.number_format = "0"; i.font = fnt(color=DARK); i.fill = fill(LIGHT)
            ws.column_dimensions[L(col)].width = 11
        ws.column_dimensions["A"].width = 46; ws.column_dimensions["B"].width = 16; ws.column_dimensions["C"].width = 10
        ws.freeze_panes = "D5"
        # Помесячный лист печатается по 12 месяцев на лист с повтором подписей,
        # а не сжатием всех 114 столбцов в ширину одной страницы.
        ws.sheet_properties.pageSetUpPr = PageSetupProperties(fitToPage=False)
        ws.print_title_cols = "A:C"; ws.print_title_rows = "3:4"
        ws.page_setup.scale = 52
        for k in range(12, N, 12):
            from openpyxl.worksheet.pagebreak import Break
            ws.col_breaks.append(Break(id=3 + k))
    return ws


def ref(name, col=None, absolute=False):
    sh, row = REG[name]
    c = "$B$" if col is None else col
    return f"'{sh}'!{c}{row}" if col is None or not absolute else f"'{sh}'!${col}${row}"


def expand(tpl, col, prev):
    """{имя} — та же колонка, {имя@p} — прошлый месяц, {имя$} — вводная (колонка B)."""
    def sub(m):
        name, mod = m.group(1), m.group(2)
        sh, row = REG[name]
        if mod == "$":
            return f"'{sh}'!$B${row}"
        c = prev if mod == "@p" else col
        return f"'{sh}'!{c}{row}"
    return re.sub(r"\{([^{}@$]+)(@p|\$)?\}", sub, tpl)


class Rows:
    def __init__(self, ws, start=5):
        self.ws, self.r = ws, start

    def section(self, text):
        self.r += 1
        c = self.ws.cell(self.r, 1, text); c.font = fnt(b=True, color=DARK)
        for j in range(1, 4):
            self.ws.cell(self.r, j).fill = fill(LIGHT)
        self.r += 1

    def param(self, name, label, value, unit="", fmt=NUM, inp=False, note=None):
        ws = self.ws; ws.cell(self.r, 1, label)
        c = ws.cell(self.r, 2, value)
        if isinstance(value, str) and ("{" in value or "!" in value or "(" in value):
            PENDING.append((c, value, "B", "B"))
        c.number_format = fmt
        if inp:
            c.fill = fill(INPUT); c.font = fnt(color="0000FF")
        ws.cell(self.r, 3, unit).font = fnt(color=GREY)
        if note:
            ws.cell(self.r, 4, note).font = fnt(i=True, color=GREY)
        REG[name] = (ws.title, self.r); self.r += 1

    def month(self, name, label, tpl=None, values=None, unit="млн ₽", fmt=NUM, total=True, bold=False,
              first=None, engine=False):
        ws = self.ws; ws.cell(self.r, 1, label).font = fnt(b=bold, i=engine, color=GREY if engine else None)
        ws.cell(self.r, 3, unit).font = fnt(color=GREY)
        REG[name] = (ws.title, self.r)
        for k in range(N):
            col, prev = L(4 + k), L(3 + k)
            if values is not None:
                v = values[k]
            else:
                v = None
            c = ws.cell(self.r, 4 + k, v); c.number_format = fmt
            if values is None:
                PENDING.append((c, first if (k == 0 and first is not None) else tpl, col, prev))
            if bold:
                c.font = fnt(b=True)
            if engine:
                c.font = fnt(i=True, color=GREY)
        if total:
            t = ws.cell(self.r, 2, f"=SUM(D{self.r}:{L(3 + N)}{self.r})"); t.number_format = fmt; t.font = fnt(b=True)
        for j in range(1, 4 + N):
            ws.cell(self.r, j).border = Border(bottom=thin)
        self.r += 1


# =============================================================================
V = sheet("Вводные", "DEVELOPAID · ВВОДНЫЕ ОФИСНОГО ПРОЕКТА",
          "Печатают только здесь: жёлтые клетки. Класс офиса задаёт допущения аренды и продажи; число в колонке «Вручную» сильнее класса.", tab=DARK)
for c, w in zip("ABCDEFGH", (50, 16, 12, 14, 14, 14, 16, 30)):
    V.column_dimensions[c].width = w
R = Rows(V, 4)
R.section("Проект")
R.param("name", "Проект", "Москва, Пролетарский проспект, вл. 20", inp=True, fmt="@")
R.param("class", "Класс офиса", "B+ и A−", inp=True, fmt="@", note="B и ниже · B+ и A− · A и выше")
R.param("start", "Начало проекта", date(2027, 1, 1), fmt="mm.yyyy", inp=True)
R.param("ird", "Документы до разрешения на стройку", 18, "мес.", "0", inp=True)
R.param("build", "Стройка", 24, "мес.", "0", inp=True)
R.param("vat", "НДС", 0.22, "", PCT, inp=True)
R.param("disc", "Ставка дисконтирования", 0.20, "в год", PCT, inp=True)
R.section("Площади (лист ТЭП)")
R.param("gba", "Наземная площадь здания", 22842, "м²", NUM0, inp=True)
R.param("total_area", "Общая площадь", 21472, "м²", NUM0, inp=True)
R.param("gla", "Арендопригодная = продаваемая", 15030, "м²", NUM0, inp=True)
R.param("spaces_all", "Подземный паркинг, мест", 68, "шт.", NUM0, inp=True)
R.param("guest", "Гостевые места (не продаются и не сдаются)", 0.10, "", PCT, inp=True)
R.param("spaces", "Машино-места в продаже или аренде", "{spaces_all$}-ROUNDUP({spaces_all$}*{guest$},0)", "шт.", NUM0)
R.section("Сделка")
R.param("land", "Цена участка", 5, "млн ₽", NUM, inp=True)
R.param("buyout", "Выкуп участков и зданий у третьих лиц", 675, "млн ₽", NUM, inp=True)
R.param("entry", "Цена входа", "{land$}+{buyout$}", "млн ₽", NUM)

# шкала классов
R.section("Шкала класса офиса (стартовые допущения, Москва, с НДС)")
HR = R.r
hdr = ["Допущение", "B и ниже", "B+ и A−", "A и выше", "Вручную", "В расчёт", "Ед."]
for j, h in enumerate(hdr):
    c = V.cell(R.r, 1 + j if j == 0 else 1 + j, h); c.font = fnt(b=True, color="FFFFFF"); c.fill = fill(HEAD)
R.r += 1
CLASS = [("c_cost", "Стоимость стройки здания", (175, 200, 225), "тыс. ₽/м² наземной", NUM0, 110),
         ("c_price", "Стартовая цена продажи", (230, 300, 420), "тыс. ₽/м²", NUM0, 273),
         ("c_rent_y", "Ставка аренды", (24, 30, 42), "тыс. ₽/м² в год", NUM0, None),
         ("c_occ", "Заполняемость после заполнения", (0.90, 0.92, 0.94), "", PCT, None),
         ("c_leaseup", "Срок заполнения", (18, 15, 12), "мес.", "0", 18),
         ("c_opex", "Расходы эксплуатации, доля дохода", (0.20, 0.18, 0.16), "", PCT, None),
         ("c_cap", "Ставка капитализации при продаже здания", (0.115, 0.105, 0.095), "", PCT, None)]
for key, label, vals, unit, fmt, manual in CLASS:
    r = R.r; V.cell(r, 1, label)
    for j, v in enumerate(vals):
        c = V.cell(r, 2 + j, v); c.number_format = fmt
    m = V.cell(r, 5, manual); m.fill = fill(INPUT); m.font = fnt(color="0000FF"); m.number_format = fmt
    cls = REG["class"][1]
    f = V.cell(r, 6, f'=IF(E{r}<>"",E{r},INDEX(B{r}:D{r},1,MATCH($B${cls},$B${HR}:$D${HR},0)))')
    f.number_format = fmt; f.font = fnt(b=True)
    V.cell(r, 7, unit).font = fnt(color=GREY)
    REG[key] = ("Вводные", r); R.r += 1
# ссылки «В расчёт» — колонка F; перерегистрируем как колонку B-подобную через отдельный блок
R.r += 1
V.cell(R.r, 1, "Клетка «Вручную» заполнена — класс для этой строки не действует. В образце стоимость стройки и цена "
       "продажи заданы вручную по присланному отчёту.").font = fnt(i=True, color=GREY)
R.r += 1
R.section("Путь 1. Продажа блоками")
R.param("growth_pre", "Рост цены до ввода", 0.015, "в месяц", PCT, inp=True)
R.param("growth_post", "Рост цены после ввода", 0.0025, "в месяц", PCT2, inp=True)
R.param("sale_months", "Срок продаж после ввода", 12, "мес.", "0", inp=True)
R.param("sell", "Маркетинг и брокеры", 0.07, "доля выручки", PCT, inp=True)
R.param("pprice", "Цена машино-места", 1.5, "млн ₽", NUM, inp=True)
R.section("Путь 2. Аренда и продажа здания")
R.param("occ0", "Заполняемость на открытии", 0.55, "", PCT, inp=True)
R.param("rent_idx", "Индексация аренды", 0.05, "в год", PCT, inp=True)
R.param("prent", "Аренда машино-места", 8, "тыс. ₽ в месяц", NUM, inp=True)
R.param("hold", "Срок удержания", 5, "лет", "0", inp=True)
R.param("exit_cost", "Расходы на сделку при продаже здания", 0.01, "", PCT, inp=True)
R.param("ptax", "Налог на имущество", 0.022, "в год от стоимости без НДС", PCT, inp=True)
R.section("Кредит")
R.param("loan_share", "Доля кредита в затратах", 0.60, "", PCT, inp=True)
R.param("spread", "Надбавка к ключевой ставке", 0.04, "", PCT, inp=True)
R.param("fee", "Комиссия за выдачу", 0.01, "от выдачи", PCT, inp=True)
R.param("res_fee", "Плата за резервирование лимита", 0.001, "от лимита", PCT2, inp=True)
R.param("lim_fee", "Плата за невыбранный лимит", 0.007, "в год", PCT2, inp=True)
R.param("term", "Срок кредита при аренде", 10, "лет", "0", inp=True)
R.param("balloon", "Остаток долга к сроку (баллон)", 0.20, "доля долга на ввод", PCT, inp=True)
V.freeze_panes = "A4"


def cv(key):  # значение «в расчёт» из шкалы класса
    return f"'Вводные'!$F${REG[key][1]}"


# =============================================================================
S = sheet("СРОКИ", "СРОКИ · КАЛЕНДАРЬ ПРОЕКТА", "Номер месяца, этапы и окна продаж и аренды.", monthly=True, tab=HEAD)
R = Rows(S)
R.param("k_permit", "Разрешение на стройку, № месяца", "{ird$}", "№", "0")
R.param("k_comm", "Ввод, № месяца", "{ird$}+{build$}", "№", "0")
R.param("k_sale_end", "Последний месяц продаж (путь 1)", "{k_comm$}+{sale_months$}-1", "№", "0")
R.param("hold_m", "Месяцев удержания (путь 2)", "{hold$}*12", "мес.", "0")
R.param("k_exit", "Продажа здания, № месяца (путь 2)", "{k_comm$}+{hold_m$}-1", "№", "0")
R.month("k", "№ месяца", "{k@p}+1", first="0", unit="№", fmt="0", total=False)
R.month("f_build", "Идёт стройка", "IF(AND({k}>={k_permit$},{k}<{k_comm$}),1,0)", unit="0/1", fmt="0", total=False)
R.month("f_comm", "Месяц ввода", "IF({k}={k_comm$},1,0)", unit="0/1", fmt="0", total=False)
R.month("f_sale", "Окно продаж (путь 1)", "IF(AND({k}>={k_comm$},{k}<={k_sale_end$}),1,0)", unit="0/1", fmt="0", total=False)
R.month("m_open", "Месяц эксплуатации (путь 2)", "IF({k}>={k_comm$},{k}-{k_comm$}+1,0)", unit="№", fmt="0", total=False)
R.month("f_hold", "Удержание (путь 2)", "IF(AND({m_open}>=1,{m_open}<={hold_m$}),1,0)", unit="0/1", fmt="0", total=False)
R.month("f_after", "Год покупателя после продажи здания", "IF(AND({m_open}>{hold_m$},{m_open}<={hold_m$}+12),1,0)", unit="0/1", fmt="0", total=False)

# =============================================================================
C = sheet("CAPEX", "CAPEX · ЗАТРАТЫ НА СОЗДАНИЕ", "Курсив — график затрат движка. В рабочей книге — статьи сметы формулами, как лист CAPEX жилой v4.", monthly=True, tab=HEAD)
R = Rows(C)
R.month("capex_obj", "Стройка здания, с НДС", values=[v / 1e6 for v in series("income", "nonres_capex")], engine=True)
R.month("capex_common", "Общие затраты проекта с участком, с НДС", values=[v / 1e6 for v in series("income", "nonres_common_capex")], engine=True)
R.month("capex_all", "Все затраты на создание", "{capex_obj}+{capex_common}", bold=True)
R.param("cost_basis", "Стоимость здания без НДС (база налога на имущество)", "{capex_obj}*(1-{vat$}/(1+{vat$}))".replace("{capex_obj}", f"'CAPEX'!B{REG['capex_obj'][1]}"), "млн ₽")
R.param("cost_total", "Все затраты на создание, итого", f"'CAPEX'!B{REG['capex_all'][1]}", "млн ₽")

# =============================================================================
K = sheet("Ставки", "СТАВКИ", "Ключевая ставка — сценарий движка (значения). Ставка кредита = ключевая + надбавка.", monthly=True, tab=HEAD)
R = Rows(K)
R.month("key", "Ключевая ставка", values=KEY, unit="", fmt=PCT, total=False, engine=True)
R.month("rate", "Ставка кредита", "{key}+{spread$}", unit="", fmt=PCT, total=False, bold=True)

# =============================================================================
P = sheet("Продажа", "ПУТЬ 1 · ПРОДАЖА БЛОКАМИ", "Договоры купли-продажи после ввода, без ДДУ и эскроу. млн ₽ с НДС.", monthly=True, tab=SALE)
R = Rows(P)
R.param("k_ps", "Отсчёт цены, № месяца (старт продаж ДДУ по вводным)", 18, "№", "0", inp=True)
R.month("pfactor", "Индекс цены", f"(1+{{growth_pre$}})^MAX(0,MIN({{k}}-{{k_ps$}},{{k_comm$}}-{{k_ps$}}))*(1+{{growth_post$}})^MAX(0,{{k}}-MAX({{k_ps$}},{{k_comm$}}))",
        unit="×", fmt="0.000", total=False)
R.month("weight", "Доля продаж месяца", "{f_sale}/{sale_months$}", unit="", fmt=PCT, )
R.month("sale_rev", "Выручка", f"({{gla$}}*{cv('c_price')}/1000+{{spaces$}}*{{pprice$}})*{{pfactor}}*{{weight}}", bold=True)
R.month("sale_cost", "Маркетинг и брокеры", "{sale_rev}*{sell$}")
R.month("unsold", "Не продано на начало месяца", "IF({k}<{k_comm$},1,{unsold@p}-{weight@p}*IF({k}-1>={k_comm$},1,0))", first="1", unit="доля", fmt=PCT, total=False)
R.month("sale_ptax", "Налог на имущество с непроданной части", "IF(AND({k}>={k_comm$},{k}<={k_sale_end$}),{cost_basis$}*{ptax$}/12*{unsold},0)")

# =============================================================================
A = sheet("Аренда", "ПУТЬ 2 · АРЕНДА", "Заполнение, индексация, расходы эксплуатации, налог на имущество, NOI. млн ₽ с НДС.", monthly=True, tab=RENT)
R = Rows(A)
R.month("occ_m", "Заполняемость", f"IF({{m_open}}<=0,0,MIN({{occ0$}},{cv('c_occ')})+({cv('c_occ')}-MIN({{occ0$}},{cv('c_occ')}))*IF({cv('c_leaseup')}<=1,1,MIN(1,MAX(0,({{m_open}}-1)/({cv('c_leaseup')}-1)))))",
        unit="", fmt=PCT, total=False)
R.month("grow", "Индекс аренды", "IF({m_open}<=0,0,(1+{rent_idx$})^(({m_open}-1)/12))", unit="×", fmt="0.000", total=False)
R.month("rent_all", "Доход аренды (по всем месяцам эксплуатации)",
        f"({{gla$}}*{cv('c_rent_y')}/12/1000+{{spaces$}}*{{prent$}}/1000)*{{grow}}*{{occ_m}}", total=False)
R.month("rent_rev", "Доход аренды", "{rent_all}*{f_hold}", bold=True)
R.month("opex", "Расходы эксплуатации", f"{{rent_rev}}*{cv('c_opex')}")
R.month("rent_ptax", "Налог на имущество", "{cost_basis$}*{ptax$}/12*{f_hold}")
R.month("noi", "NOI", "{rent_rev}-{opex}-{rent_ptax}", bold=True)
R.month("noi_buyer", "NOI покупателя: год после продажи здания (до налога на имущество)", f"{{rent_all}}*(1-{cv('c_opex')})*{{f_after}}")

# =============================================================================
X = sheet("Выход", "ПУТЬ 2 · ПРОДАЖА ЗДАНИЯ", "Стоимость = NOI следующего года / ставка капитализации.", tab=RENT)
X.column_dimensions["A"].width = 58; X.column_dimensions["B"].width = 16; X.column_dimensions["C"].width = 14
R = Rows(X, 4)
R.param("fwd_noi", "NOI следующих 12 месяцев после срока удержания", f"'Аренда'!B{REG['noi_buyer'][1]}-{{cost_basis$}}*{{ptax$}}", "млн ₽")
R.param("cap_x", "Ставка капитализации", f"{cv('c_cap')}", "", PCT)
R.param("exit_value", "Стоимость здания", "{fwd_noi$}/{cap_x$}", "млн ₽")
R.param("exit_costv", "Расходы на сделку", "{exit_value$}*{exit_cost$}", "млн ₽")
R.param("exit_net", "Получено от продажи здания", "{exit_value$}-{exit_costv$}", "млн ₽")
R.param("noi_stab", "NOI года после заполнения", f"SUMPRODUCT(('СРОКИ'!D{REG['m_open'][1]}:{L(3 + N)}{REG['m_open'][1]}>={cv('c_leaseup')})*('СРОКИ'!D{REG['m_open'][1]}:{L(3 + N)}{REG['m_open'][1]}<={cv('c_leaseup')}+11)*'Аренда'!D{REG['rent_all'][1]}:{L(3 + N)}{REG['rent_all'][1]})*(1-{cv('c_opex')})-{{cost_basis$}}*{{ptax$}}", "млн ₽")
R.param("yoc", "Доходность на затраты = NOI года после заполнения / все затраты", "{noi_stab$}/{cost_total$}", "", PCT)
R.param("spread_yoc", "Доходность на затраты минус ставка капитализации", "{yoc$}-{cap_x$}", "", PCT)
X.cell(R.r + 1, 1, "Меньше нуля — готовое здание стоит дешевле, чем его стройка: строить, чтобы сдавать, невыгодно.").font = fnt(i=True, color=GREY)
A_ = Rows(A, A.max_row + 1)
A_.month("exit_rev", "Продажа здания", f"IF({{k}}={{k_exit$}},'Выход'!$B${REG['exit_value'][1]},0)", bold=True)
A_.month("exit_cst", "Расходы на сделку", f"IF({{k}}={{k_exit$}},'Выход'!$B${REG['exit_costv'][1]},0)")

# =============================================================================
def finance_sheet(path, title, tab, rev, costs, ptax_row, horizon):
    """НДС, кредит и деньги собственника пути — формулами object_flows."""
    p = "s" if path == "direct" else "r"
    ws = sheet(title, f"{title.upper()} · НДС, КРЕДИТ, ДЕНЬГИ СОБСТВЕННИКА",
               "Кредит: выборка — доля затрат до ввода; проценты до ввода капитализируются; после ввода — погашение.", monthly=True, tab=tab)
    R = Rows(ws)
    R.param(p + "hz", "Последний месяц пути", horizon, "№", "0")
    R.month(p + "on", "Месяц пути", f"IF({{k}}<={{{p}hz$}},1,0)", unit="0/1", fmt="0", total=False)
    R.month(p + "rev", "Выручка пути", rev + f"*{{{p}on}}", bold=True)
    R.month(p + "costs", "Расходы продаж и эксплуатации", costs + f"*{{{p}on}}")
    R.section("НДС")
    R.month(p + "vch", "НДС начислен", f"{{{p}rev}}*{{vat$}}/(1+{{vat$}})")
    R.month(p + "vpre", "НДС к вычету накоплено до зачёта",
            f"{{{p}vcar@p}}+{{capex_obj}}*{{vat$}}/(1+{{vat$}})*{{{p}on}}-{{{p}vch}}", first=f"{{capex_obj}}*{{vat$}}/(1+{{vat$}})-{{{p}vch}}", total=False)
    R.month(p + "vpaid", "НДС уплачен (минус — возмещён)", f"IF(AND({{f_comm}}=1,{{{p}vpre}}>0),-{{{p}vpre}},IF({{{p}vpre}}<0,-{{{p}vpre}},0))*{{{p}on}}")
    R.month(p + "vcar", "НДС к вычету на конец месяца", f"{{{p}vpre}}+{{{p}vpaid}}", total=False)
    R.section("Кредит")
    R.month(p + "draw", "Выдача", f"IF({{k}}<={{k_comm$}},{{capex_all}}*{{loan_share$}},0)")
    R.month(p + "int", "Проценты месяца", f"{{{p}bal@p}}*{{rate}}/12", first="0")
    R.month(p + "icap", "Проценты капитализированы (до ввода)", f"IF({{k}}<{{k_comm$}},{{{p}int}},0)")
    R.month(p + "ipaid", "Проценты уплачены", f"IF({{k}}>={{k_comm$}},{{{p}int}},0)*{{{p}on}}")
    R.month(p + "pre", "Долг до погашения", f"{{{p}bal@p}}+{{{p}icap}}+{{{p}draw}}", first=f"{{{p}icap}}+{{{p}draw}}", total=False)
    R.month(p + "cash", "Деньги на погашение", f"{{{p}rev}}-{{{p}costs}}-{{{p}vpaid}}-{{{p}ipaid}}")
    if path == "direct":
        R.month(p + "rep", "Погашение",
                f"IF({{{p}pre}}<=0,0,IF({{k}}={{{p}hz$}},{{{p}pre}},MIN({{{p}pre}},MAX(0,{{{p}cash}}))))*{{{p}on}}")
    else:
        R.month(p + "cumd", "Выдано нарастающим", f"{{{p}cumd@p}}+{{{p}draw}}", first=f"{{{p}draw}}", total=False)
        R.param(p + "first", "Первая выдача, № месяца", f"COUNTIF('{title}'!D{REG[p + 'cumd'][1]}:{L(3 + N)}{REG[p + 'cumd'][1]},0)", "№", "0")
        R.param(p + "mat", "Срок кредита, № месяца", f"MAX({{{p}first$}}+{{term$}}*12,{{k_comm$}})", "№", "0")
        R.month(p + "blnm", "Баллон к сроку (фиксируется в месяц ввода)", f"IF({{k}}={{k_comm$}},{{{p}pre}}*{{balloon$}},{{{p}blnm@p}})", first=f"IF({{k}}={{k_comm$}},{{{p}pre}}*{{balloon$}},0)", total=False)
        R.month(p + "left", "Месяцев до срока", f"{{{p}mat$}}-{{k}}", unit="мес.", fmt="0", total=False)
        R.month(p + "pay", "Плановый платёж (аннуитет к баллону)",
                f"IF(OR({{k}}<={{k_comm$}},{{{p}left}}<=0),0,({{{p}pre}}-{{{p}blnm}}/(1+{{rate}}/12)^({{{p}left}}))*({{rate}}/12)/(1-(1+{{rate}}/12)^(-{{{p}left}})))")
        R.month(p + "rep", "Погашение",
                f"IF({{{p}pre}}<=0,0,IF(OR({{k}}={{{p}hz$}},AND({{k}}>{{k_comm$}},{{{p}left}}<=0)),{{{p}pre}},IF({{k}}>{{k_comm$}},MIN({{{p}pre}},MAX(0,{{{p}pay}}-{{{p}int}})),0)))*{{{p}on}}")
    R.month(p + "bal", "Долг на конец месяца", f"{{{p}pre}}-{{{p}rep}}", total=False, bold=True)
    # две линии: до разрешения и после — резервирование и плата за невыбранный лимит
    R.month(p + "dl", "Выдача линии 1 (до разрешения)", f"IF({{k}}<{{k_permit$}},{{{p}draw}},0)")
    R.month(p + "db", "Выдача линии 2 (стройка)", f"IF({{k}}>={{k_permit$}},{{{p}draw}},0)")
    R.month(p + "cl", "Выбрано линии 1 нарастающим", f"{{{p}cl@p}}+{{{p}dl}}", first=f"{{{p}dl}}", total=False)
    R.month(p + "cb", "Выбрано линии 2 нарастающим", f"{{{p}cb@p}}+{{{p}db}}", first=f"{{{p}db}}", total=False)
    lim_l = f"'{title}'!$B${REG[p + 'dl'][1]}"; lim_b = f"'{title}'!$B${REG[p + 'db'][1]}"
    R.month(p + "fee", "Комиссии банка",
            f"{{{p}draw}}*{{fee$}}"
            f"+IF(AND({{{p}dl}}>0,{{{p}cl@p}}=0),{lim_l}*{{res_fee$}},0)+IF(AND({{{p}db}}>0,{{{p}cb@p}}=0),{lim_b}*{{res_fee$}},0)"
            f"+IF(AND({{{p}cl}}>0,{{k}}<{{k_permit$}}),({lim_l}-{{{p}cl}})*{{lim_fee$}}/12,0)"
            f"+IF(AND({{{p}cb}}>0,{{k}}<={{k_comm$}}),({lim_b}-{{{p}cb}})*{{lim_fee$}}/12,0)",
            first=f"{{{p}draw}}*{{fee$}}+IF({{{p}dl}}>0,{lim_l}*{{res_fee$}}+({lim_l}-{{{p}cl}})*{{lim_fee$}}/12,0)")
    R.section("Деньги собственника")
    R.month(p + "tax", "Налог на прибыль проекта к уплате (движок; в рабочей книге — лист «Налоги»)",
            values=[v / 1e6 for v in series(path, "profit_tax", "cf")], engine=True)
    R.month(p + "own", "Поток собственника",
            f"{{{p}rev}}-{{{p}costs}}-{{{p}vpaid}}-{{{p}ipaid}}-{{{p}fee}}+{{{p}draw}}-{{{p}rep}}-{{capex_all}}*IF({{k}}<={{k_comm$}},1,0)-{{{p}tax}}", bold=True)
    R.month(p + "cum", "Поток собственника нарастающим", f"{{{p}cum@p}}+{{{p}own}}", first=f"{{{p}own}}", total=False)
    R.month(p + "proj", "Поток проекта до кредита",
            f"{{{p}rev}}-{{{p}costs}}-{{{p}vpaid}}-{{capex_all}}*IF({{k}}<={{k_comm$}},1,0)-{{{p}tax}}")
    R.month(p + "dfac", "Множитель дисконтирования", "1/(1+{disc$})^({k}/12)", unit="×", fmt="0.000", total=False)
    R.param(p + "npv", "NPV проекта", f"SUMPRODUCT('{title}'!D{REG[p + 'proj'][1]}:{L(3 + N)}{REG[p + 'proj'][1]},'{title}'!D{REG[p + 'dfac'][1]}:{L(3 + N)}{REG[p + 'dfac'][1]})", "млн ₽")
    R.param(p + "irr", "IRR проекта", f"(1+IRR('{title}'!D{REG[p + 'proj'][1]}:{L(3 + N)}{REG[p + 'proj'][1]},-0.01))^12-1", "в год", PCT)
    R.month(p + "dfe", "Множитель дисконтирования собственника (ставка / 12)", "1/(1+{disc$}/12)^({k})", unit="×", fmt="0.000", total=False)
    R.param(p + "npv_eq", "NPV собственника", f"SUMPRODUCT('{title}'!D{REG[p + 'own'][1]}:{L(3 + N)}{REG[p + 'own'][1]},'{title}'!D{REG[p + 'dfe'][1]}:{L(3 + N)}{REG[p + 'dfe'][1]})", "млн ₽")
    R.param(p + "peak", "Пик собственных денег", f"-MIN('{title}'!D{REG[p + 'cum'][1]}:{L(3 + N)}{REG[p + 'cum'][1]})", "млн ₽")
    R.param(p + "net", "Чистая прибыль (сумма потока собственника)", f"'{title}'!B{REG[p + 'own'][1]}", "млн ₽")
    return ws


finance_sheet("direct", "CF_Продажа", SALE, "{sale_rev}", "({sale_cost}+{sale_ptax})", None, "{k_sale_end$}")
finance_sheet("income", "CF_Аренда", RENT, "({rent_rev}+{exit_rev})", "({opex}+{rent_ptax}+{exit_cst})", None, "{k_exit$}")

# =============================================================================
O = sheet("ОТЧЕТ", "DEVELOPAID · ОТЧЁТ · ДВА ПУТИ", "Все числа — ссылки на листы модели. млн ₽ с НДС.", tab=DARK)
for c, w in zip("ABCD", (52, 20, 24, 40)):
    O.column_dimensions[c].width = w
for j, h in enumerate(["Показатель", "Продажа блоками", "Аренда и продажа здания"]):
    c = O.cell(4, 1 + j, h); c.font = fnt(b=True, color="FFFFFF"); c.fill = fill([HEAD, SALE, RENT][j])
b = lambda n: f"='{REG[n][0]}'!$B${REG[n][1]}"
lines = [("Выручка", b("srev"), b("rrev"), NUM0), ("Все затраты на создание, с участком", b("cost_total"), b("cost_total"), NUM0),
         ("Расходы на продажу или эксплуатацию", b("scosts"), b("rcosts"), NUM0),
         ("Проценты уплаченные и капитализированные", f"='CF_Продажа'!$B${REG['sint'][1]}", f"='CF_Аренда'!$B${REG['rint'][1]}", NUM0),
         ("Комиссии банка", b("sfee"), b("rfee"), NUM0), ("Налог на прибыль", b("stax"), b("rtax"), NUM0),
         ("Чистая прибыль", b("snet"), b("rnet"), NUM0), ("NPV проекта", b("snpv"), b("rnpv"), NUM0),
         ("IRR проекта (до кредита)", b("sirr"), b("rirr"), PCT), ("NPV собственника", b("snpv_eq"), b("rnpv_eq"), NUM0),
         ("Пик собственных денег", b("speak"), b("rpeak"), NUM0),
         ("NOI года после заполнения", "не применяется", b("noi_stab"), NUM0), ("Доходность на затраты", "не применяется", b("yoc"), PCT),
         ("Ставка капитализации", "не применяется", b("cap_x"), PCT), ("Стоимость здания при продаже", "не применяется", b("exit_value"), NUM0),
         ("Пик долга", f"=MAX('CF_Продажа'!D{REG['sbal'][1]}:{L(3 + N)}{REG['sbal'][1]})", f"=MAX('CF_Аренда'!D{REG['rbal'][1]}:{L(3 + N)}{REG['rbal'][1]})", NUM0)]
for i, (lab, a, c2, fmt) in enumerate(lines):
    r = 5 + i
    O.cell(r, 1, lab)
    for j, v in enumerate((a, c2)):
        x = O.cell(r, 2 + j, v); x.number_format = fmt
        x.alignment = Alignment(horizontal="right")
    for j in range(1, 4):
        O.cell(r, j).border = Border(bottom=thin)
    if lab in ("Чистая прибыль", "NPV проекта"):
        for j in range(1, 4):
            O.cell(r, j).font = fnt(b=True); O.cell(r, j).fill = fill(LIGHT)
REG["rep_npv"] = ("ОТЧЕТ", 12)
# по годам
yr0 = 5 + len(lines) + 2
O.cell(yr0, 1, "По годам: поток собственника").font = fnt(b=True, color=DARK)
for j, h in enumerate(["Год", "Продажа блоками", "Аренда и продажа здания"]):
    c = O.cell(yr0 + 1, 1 + j, h); c.font = fnt(b=True, color="FFFFFF"); c.fill = fill(HEAD)
for i, y in enumerate(range(2027, 2037)):
    r = yr0 + 2 + i
    O.cell(r, 1, y)
    for j, (sh, key) in enumerate((("CF_Продажа", "sown"), ("CF_Аренда", "rown"))):
        rr = REG[key][1]
        x = O.cell(r, 2 + j, f"=SUMPRODUCT(('СРОКИ'!$D$3:${L(3 + N)}$3>=DATE({y},1,1))*('СРОКИ'!$D$3:${L(3 + N)}$3<DATE({y + 1},1,1))*'{sh}'!$D${rr}:${L(3 + N)}${rr})")
        x.number_format = NUM0
O.freeze_panes = "A5"

# =============================================================================
DSH = sheet("Дашборд", "DEVELOPAID · ИНВЕСТИЦИОННЫЙ ДАШБОРД · ОФИС", "", tab=DARK)
DSH["A2"] = "='Вводные'!$B$" + str(REG["name"][1]) + "&\" · класс \"&'Вводные'!$B$" + str(REG["class"][1])
for c in "ABCDEFGHIJKL":
    DSH.column_dimensions[c].width = 13
tiles = [("NPV · продажа, млн ₽", b("snpv"), NUM0, SALE), ("NPV · аренда, млн ₽", b("rnpv"), NUM0, RENT),
         ("Доходность на затраты", b("yoc"), PCT, RENT), ("Ставка капитализации", b("cap_x"), PCT, RENT),
         ("Пик своих денег · продажа", b("speak"), NUM0, SALE), ("Пик своих денег · аренда", b("rpeak"), NUM0, RENT)]
for i, (lab, f_, fmt, col) in enumerate(tiles):
    c0 = 1 + (i % 3) * 4; r0 = 4 + (i // 3) * 4
    DSH.merge_cells(start_row=r0, start_column=c0, end_row=r0, end_column=c0 + 2)
    DSH.merge_cells(start_row=r0 + 1, start_column=c0, end_row=r0 + 2, end_column=c0 + 2)
    h = DSH.cell(r0, c0, lab); h.font = fnt(b=True, color="FFFFFF"); h.fill = fill(col)
    v = DSH.cell(r0 + 1, c0, f_); v.number_format = fmt; v.font = Font(name=F, size=18, b=True, color=DARK)
    v.fill = fill(LIGHT); v.alignment = Alignment(horizontal="center", vertical="center")
DSH["A12"] = ("=IF(AND('ОТЧЕТ'!B12<0,'ОТЧЕТ'!C12<0),\"Проект не окупается ни продажей, ни арендой. \",\"\")"
              "&IF('ОТЧЕТ'!B12>'ОТЧЕТ'!C12,\"Продажа выгоднее аренды.\",\"Аренда выгоднее продажи.\")"
              f"&IF('Выход'!$B${REG['spread_yoc'][1]}<0,\" Доходность на затраты ниже ставки капитализации: готовое здание стоит дешевле его стройки.\",\"\")")
DSH["A12"].font = fnt(b=True, color=DARK); DSH.merge_cells("A12:L13"); DSH["A12"].alignment = Alignment(wrap_text=True, vertical="top")
ch = BarChart(); ch.title = "Поток собственника по годам, млн ₽"; ch.height = 7.5; ch.width = 22
ch.add_data(Reference(O, min_col=2, max_col=3, min_row=yr0 + 1, max_row=yr0 + 11), titles_from_data=True)
ch.set_categories(Reference(O, min_col=1, min_row=yr0 + 2, max_row=yr0 + 11))
ch.series[0].graphicalProperties.solidFill = SALE; ch.series[1].graphicalProperties.solidFill = RENT
ch.x_axis.tickLblPos = "low"; ch.x_axis.delete = False; ch.y_axis.delete = False
DSH.add_chart(ch, "A15")

# =============================================================================
I = sheet("ИНСТРУКЦИЯ", "ИНСТРУКЦИЯ", "Как устроена книга.", tab=GREY)
I.column_dimensions["A"].width = 26; I.column_dimensions["B"].width = 100
for i, (a, t) in enumerate([
        ("Дашборд", "Главные числа обоих путей и вывод."), ("ОТЧЕТ", "Два пути рядом и поток собственника по годам. Только ссылки."),
        ("Вводные", "Единственное место правки: жёлтые клетки. Класс офиса задаёт шкалу; «Вручную» сильнее класса."),
        ("СРОКИ", "Номер месяца, разрешение, ввод, окна продаж и аренды."), ("CAPEX", "График затрат на создание."),
        ("Ставки", "Ключевая ставка и ставка кредита по месяцам."), ("Продажа", "Путь 1: цена, продажи, маркетинг, налог на имущество."),
        ("Аренда", "Путь 2: заполнение, индексация, доход, эксплуатация, NOI, продажа здания."),
        ("Выход", "Стоимость здания по ставке капитализации, доходность на затраты."),
        ("CF_Продажа, CF_Аренда", "НДС, кредит двумя линиями, комиссии, поток собственника и проекта, NPV и IRR."),
        ("ПРОВЕРКИ", "Каждый итог книги против расчёта движка.")], start=4):
    I.cell(i, 1, a).font = fnt(b=True); I.cell(i, 2, t)

# =============================================================================
Q = sheet("ПРОВЕРКИ", "ПРОВЕРКИ · КНИГА ПРОТИВ ДВИЖКА", "Расхождение больше 0,5 млн ₽ (у IRR — 0,5 п.п.) — красное.", tab=GREY)
for c, w in zip("ABCDE", (52, 16, 16, 14, 12)):
    Q.column_dimensions[c].width = w
for j, h in enumerate(["Итог", "Книга", "Движок", "Разница", "Статус"]):
    c = Q.cell(4, 1 + j, h); c.font = fnt(b=True, color="FFFFFF"); c.fill = fill(HEAD)
D_, I_ = OBJ["direct"], OBJ["income"]
checks = [("Продажа: выручка", b("srev")[1:], D_["totals"]["sale_revenue"] / 1e6),
          ("Продажа: маркетинг и брокеры", f"'Продажа'!B{REG['sale_cost'][1]}", D_["totals"]["selling_cost"] / 1e6),
          ("Продажа: налог на имущество", f"'Продажа'!B{REG['sale_ptax'][1]}", D_["totals"]["property_tax"] / 1e6),
          ("Продажа: проценты", f"'CF_Продажа'!B{REG['sint'][1]}", D_["totals"]["loan_interest"] / 1e6),
          ("Продажа: комиссии", f"'CF_Продажа'!B{REG['sfee'][1]}", D_["totals"]["loan_fee"] / 1e6),
          ("Продажа: чистая прибыль", f"'CF_Продажа'!B{REG['sown'][1]}", RES["direct"]["summary"]["net_profit"] / 1e6),
          ("Продажа: NPV проекта", b("snpv")[1:], RES["direct"]["summary"]["npv"] / 1e6),
          ("Продажа: NPV собственника", b("snpv_eq")[1:], RES["direct"]["summary"]["npv_equity"] / 1e6),
          ("Продажа: IRR проекта, %", f"{b('sirr')[1:]}*100", RES["direct"]["summary"]["irr_project"] * 100),
          ("Аренда: доход", f"'Аренда'!B{REG['rent_rev'][1]}", I_["totals"]["rent_revenue"] / 1e6),
          ("Аренда: эксплуатация", f"'Аренда'!B{REG['opex'][1]}", I_["totals"]["opex"] / 1e6),
          ("Аренда: налог на имущество", f"'Аренда'!B{REG['rent_ptax'][1]}", I_["totals"]["property_tax"] / 1e6),
          ("Аренда: стоимость здания", b("exit_value")[1:], I_["kpi"]["exit_value"] / 1e6),
          ("Аренда: NOI года после заполнения", b("noi_stab")[1:], I_["kpi"]["stabilized_noi"] / 1e6),
          ("Аренда: проценты", f"'CF_Аренда'!B{REG['rint'][1]}", I_["totals"]["loan_interest"] / 1e6),
          ("Аренда: комиссии", f"'CF_Аренда'!B{REG['rfee'][1]}", I_["totals"]["loan_fee"] / 1e6),
          ("Аренда: чистая прибыль", f"'CF_Аренда'!B{REG['rown'][1]}", RES["income"]["summary"]["net_profit"] / 1e6),
          ("Аренда: NPV проекта", b("rnpv")[1:], RES["income"]["summary"]["npv"] / 1e6),
          ("Аренда: NPV собственника", b("rnpv_eq")[1:], RES["income"]["summary"]["npv_equity"] / 1e6),
          ("Аренда: IRR проекта, %", f"{b('rirr')[1:]}*100", RES["income"]["summary"]["irr_project"] * 100)]
for i, (lab, f_, eng) in enumerate(checks):
    r = 5 + i
    Q.cell(r, 1, lab); Q.cell(r, 2, "=" + f_).number_format = NUM; Q.cell(r, 3, eng).number_format = NUM
    Q.cell(r, 4, f"=B{r}-C{r}").number_format = NUM
    Q.cell(r, 5, f'=IF(ABS(D{r})<=0.5,"ПРОЙДЕНО","РАСХОЖДЕНИЕ")')
Q.freeze_panes = "A5"

for cell, tpl, col, prev in PENDING:
    cell.value = "=" + expand(tpl, col, prev)
order = ["Дашборд", "ОТЧЕТ", "Вводные", "ИНСТРУКЦИЯ", "СРОКИ", "CAPEX", "Ставки", "Продажа", "Аренда", "Выход",
         "CF_Продажа", "CF_Аренда", "ПРОВЕРКИ"]
wb._sheets = [wb[n] for n in order]
wb.save(OUT)
print("saved", OUT)
