"""Книга гостиничного проекта: живые формулы, годовой USALI, итоги и сверка.

Шаг — месяц, как у движка; смета от площадей и ставок, годовой USALI,
кредит, налоги, денежный поток и итоги связаны формулами. NPV, IRR, PBP и
DPBP пересчитываются в скачанном файле. Значения движка — только снимок
для сверки, а не источник итогов книги. Изменение вводных снимает сверку
с прежним снимком, но внутренние балансы продолжают проверяться.

Методика совпадает с `developaid_hotel_strategy.hotel_flows`: USALI,
амортизация, НДС, кредит, удержание денег до погашения, налог на прибыль
с переносом убытка. Дополнительные статьи без формулы поименованы отдельно.
Формулы проверяет `xlsx_eval`; рассчитанные значения сохраняются рядом с
формулами, чтобы книгу можно было читать и без пересчёта в Excel.
"""

from __future__ import annotations

import io
import sys
from copy import copy
from datetime import date
from typing import Any

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.workbook.properties import CalcProperties
from openpyxl.utils import get_column_letter

import developaid_hotel_strategy as hs
from developaid_finance_math import LOSS_CARRY_USE_LIMIT
import hotel_workbook_capex as capex_book
import hotel_workbook_sensitivity as sensitivity_book
from xlsx_eval import Evaluator
from v4_value_cache import with_cached_values

CAPEX_INPUT_DEFAULTS = capex_book.INPUT_DEFAULTS
FORMULA_CAPEX_ARTICLES = frozenset(capex_book.ARTICLES)

ENGINE_FILL = PatternFill("solid", fgColor="FFF3E0")
INPUT_FILL = PatternFill("solid", fgColor="FFFDE7")
BOLD = Font(bold=True)
MONEY = '#,##0;[Red](#,##0);"—"'
PERCENT = '0.0%;[Red](0.0%);"—"'
# Шрифт и цвета образца Отеля 1: Контр_панель!B2, Предпосылки!B2.
SECTION_FILL = PatternFill("solid", fgColor="A3D6C0")
TITLE_FILL = PatternFill("solid", fgColor="00B050")

INPUTS_SHEET = "Вводные"
CALC_SHEET = "Расчёт"
USALI_SHEET = "USALI по годам"
TOTALS_SHEET = "Итоги"
CHECK_SHEET = "Сверка"
FIRST_ROW = 6  # первая строка помесячной таблицы на листе «Расчёт»
NO_DRAW = 1_000_000  # «выдачи не было» — номер строки, которого нет
CHOICES = {
    "financing": {"commercial": "Обычный кредит", "preferential": "Льготный кредит", "none": "Без кредита"},
    "repayment": {"annuity": "Аннуитет", "sculpted": "Под DSCR", "sweep": "Весь свободный поток"},
    "exit_mode": {"sale": "Продажа", "hold": "Удержание"},
    "valuation": {"cap_rate": "Ставка капитализации", "ev_ebitda": "EV/EBITDA"},
}

# Модельные величины листа «Вводные» (столбец C): ключ → подпись.
MODEL: tuple[tuple[str, str], ...] = (
    ("keys", "Номерной фонд, номеров"),
    ("adr", "ADR на дату цены, ₽ без НДС"),
    ("price_k", "Дата цены ADR — номер строки расчёта (k)"),
    ("index", "Индексация ADR, доля в год"),
    ("occ_start", "Стартовая загрузка, доля"),
    ("occ_target", "Целевая загрузка, доля"),
    ("ramp", "Выход на целевую загрузку, мес."),
    ("fnb", "F&B, доля выручки номеров"),
    ("other", "Прочие департаменты, доля выручки номеров и F&B"),
    ("rooms_cost", "Расходы номерного фонда, доля"),
    ("fnb_cost", "Расходы F&B, доля"),
    ("other_cost", "Расходы прочих департаментов, доля"),
    ("ag", "A&G, доля выручки"),
    ("sm", "S&M, доля выручки"),
    ("pom", "POM, доля выручки"),
    ("util", "Utilities, доля выручки"),
    ("base_fee", "Базовое вознаграждение оператора, доля выручки"),
    ("incentive", "Поощрительное вознаграждение, доля GOP"),
    ("reserve", "Резерв FF&E, доля выручки"),
    ("relief_months", "Льгота НДС на проживание, мес. с ввода"),
    ("rev_mult", "Множитель выручки сценария"),
    ("vat", "Ставка НДС"),
    ("property_rate", "Налог на имущество, доля в год"),
    ("insurance_rate", "Страхование, доля CAPEX без НДС в год"),
    ("building_months", "Амортизация здания, мес."),
    ("ffe_months", "Амортизация мебели и оборудования, мес."),
    ("financing", "Финансирование (commercial / preferential / none)"),
    ("loan_share", "Доля кредита в затратах"),
    ("spread", "Спред обычного кредита к КС"),
    ("pref_share", "Льготный: доля ключевой ставки"),
    ("pref_margin", "Льготный: маржа"),
    ("limit_on", "Лимит кредита задан (1 / 0)"),
    ("limit", "Лимит кредита, ₽"),
    ("fee", "Комиссия за выдачу, доля"),
    ("term", "Срок кредита от первой выдачи, мес."),
    ("grace", "Отсрочка тела от первой выдачи, мес."),
    ("repayment", "Погашение (annuity / sculpted / sweep)"),
    ("dscr", "Целевой DSCR (погашение «под DSCR»)"),
    ("exit_mode", "Конец срока (sale / hold)"),
    ("valuation", "Оценка выхода (cap_rate / ev_ebitda)"),
    ("cap", "Ставка капитализации"),
    ("multiple", "Мультипликатор EV/EBITDA"),
    ("exit_cost", "Затраты на продажу, доля"),
    ("tax_rate", "Налог на прибыль"),
    ("loss_limit", "Зачёт убытка прошлых лет, не более доли базы"),
    ("discount", "Ставка дисконтирования"),
    ("comm", "Ввод — номер строки расчёта (k)"),
    ("horizon", "Конец срока — номер строки расчёта (k)"),
    ("hold", "Срок эксплуатации, лет"),
    ("limit_th", "Лимит кредита, тыс. ₽/номер (0 при отсутствии лимита)"),
    ("start", "Начало проекта"),
    ("price_date", "Дата цены ADR"),
)
DERIVED: tuple[tuple[str, str], ...] = (
    ("vs", "Доля НДС в сумме с НДС"),
    ("capex_total", "Затраты проекта за срок, ₽ с НДС"),
    ("land_total", "Участок и ВРИ, ₽ (без НДС, без амортизации)"),
    ("ffe_total", "Мебель и оборудование, ₽ с НДС"),
    ("capex_vat", "НДС в затратах, ₽"),
    ("capex_net", "Затраты без НДС, ₽"),
    ("ffe_net", "Мебель и оборудование без НДС, ₽"),
    ("building_net", "Здание без НДС, ₽ (база налога на имущество)"),
    ("insurance_month", "Страхование в месяц, ₽"),
    ("first_draw", "Первая выдача кредита — k"),
    ("maturity", "Срок кредита — k"),
    ("amort_start", "Начало погашения тела — k"),
    ("forward_ebitda", "EBITDA 12 месяцев после срока, ₽"),
    ("exit_value", "Стоимость гостиницы на выходе, ₽"),
    ("peak", "Пик долга, ₽"),
    ("rate_m", "Месячная ставка дисконтирования проекта (как NPV движка)"),
)
_ROWS: dict[str, int] = {key: 4 + i for i, (key, _) in enumerate(MODEL)}
_ROWS.update({key: 4 + len(MODEL) + 2 + i for i, (key, _) in enumerate(DERIVED)})

# Колонки листа «Расчёт»: ключ → подпись. Буквы — по порядку.
_COLUMN_LIST: tuple[tuple[str, str], ...] = (
    ("k", "k"), ("month", "Месяц"), ("year", "Год (календарь)"), ("days", "Дней (календарь)"),
    ("capex", "Затраты, ₽ с НДС (движок)"), ("capex_ffe", "в т.ч. FF&E (движок)"),
    ("capex_land", "в т.ч. участок и ВРИ (движок)"), ("key_rate", "Ключевая ставка (движок)"),
    ("in_h", "В сроке расчёта"), ("year_h", "Год в сроке"),
    ("open", "Месяц работы"), ("occ", "Загрузка"), ("growth", "Индекс ADR"), ("adr", "ADR"),
    ("avail", "Номеро-ночи в продаже"), ("sold", "Проданные номеро-ночи"),
    ("rooms", "Выручка номеров"), ("fnb", "Выручка F&B"), ("other", "Выручка прочих"),
    ("dept", "Выручка департаментов"), ("relief", "Льгота НДС на проживание"),
    ("rooms_exp", "Расходы НФ"), ("fnb_exp", "Расходы F&B"), ("other_exp", "Расходы прочих"),
    ("ag", "A&G"), ("sm", "S&M"), ("pom", "POM"), ("util", "Utilities"), ("gop", "GOP"),
    ("base_fee", "Базовое вознаграждение"), ("incentive", "Поощрительное вознаграждение"),
    ("reserve", "Резерв FF&E"),
    ("book_b_open", "Здание: стоимость на начало"), ("dep_b", "Здание: износ"),
    ("book_b_close", "Здание: стоимость на конец"), ("book_f_open", "FF&E: стоимость на начало"),
    ("dep_f", "FF&E: износ"), ("book_f_close", "FF&E: стоимость на конец"),
    ("ptax", "Налог на имущество"), ("ins", "Страхование"),
    ("revenue", "Выручка всего"), ("opex", "Расходы USALI"), ("ebitda", "EBITDA"),
    ("exit", "Выход — продажа"), ("exit_cost", "Затраты на продажу"),
    ("residual", "Удержание — оценка"),
    ("vatin", "НДС в затратах месяца"), ("refund", "Возмещение НДС стройки"),
    ("rate", "Ставка кредита"), ("bal_open", "Долг на начало"), ("interest", "Проценты"),
    ("int_cap", "Проценты капитализированные"), ("int_paid", "Проценты уплаченные"),
    ("want", "Нужно кредита"), ("cum_prev", "Выдано до месяца"), ("draw", "Выдача"),
    ("cum_draw", "Выдано всего"), ("fee", "Комиссия"), ("first_k", "Служебная: выдача"),
    ("bal_after", "Долг после выдачи"), ("prepay", "Возмещение НДС в погашение"),
    ("bal_pre", "Долг до погашения"), ("annuity", "Тело по аннуитету"),
    ("sched", "Плановое погашение тела"),
    ("free", "Свободные деньги после процентов и налога"),
    ("res_open", "Удержано в проекте на начало"),
    ("sweep", "Погашение удержанными деньгами"),
    ("repay", "Погашение тела"), ("repay_total", "Погашение всего"),
    ("bal_close", "Долг на конец"),
    ("res_close", "Удержано в проекте на конец"), ("retained", "Удержано за месяц"),
    ("amort", "Плановое тело (для DSCR)"),
    ("dep", "Амортизация"), ("margin", "Налоговая маржа"), ("fin_ded", "Проценты и комиссии к вычету"),
    ("def_acc", "Расходы до ввода, к году ввода"), ("newyear", "Новый налоговый год"),
    ("prior", "Убыток прошлых лет"), ("yres", "Результат года нарастающим"),
    ("used", "Зачтено убытка"), ("base", "База налога"), ("ypaid_prev", "Налог года до месяца"),
    ("tax", "Налог на прибыль"), ("ypaid", "Налог года с месяцем"),
    ("vat_paid", "НДС к уплате (− возмещение)"), ("operating", "Деньги эксплуатации"),
    ("to_equity", "Деньги собственнику без CAPEX"), ("project_cf", "Поток проекта"),
    ("equity_cf", "Поток капитала"), ("fcff", "Поток до финансирования"),
    ("disc_p", "Поток проекта, дисконт."), ("disc_e", "Поток капитала, дисконт."),
    ("irr_p", "Поток проекта при IRR движка"), ("irr_e", "Поток капитала при IRR движка"),
    ("cum_p", "Поток проекта нарастающим"), ("cross_p", "Окупаемость: месяц перехода"),
    ("cum_dp", "Дисконт. поток нарастающим"), ("cross_dp", "DPBP: месяц перехода"),
    ("cum_e", "Денежный поток капитала нарастающим"),
    ("profit_after_tax", "Финансовый результат после налога"),
)
COLUMNS: dict[str, tuple[str, str]] = {
    key: (get_column_letter(i), label) for i, (key, label) in enumerate(_COLUMN_LIST, start=1)}


def _c(name: str) -> str:
    return COLUMNS[name][0]


def _in(name: str) -> str:
    column = "I" if name in CHOICES else "C"
    return f"'{INPUTS_SHEET}'!${column}${_ROWS[name]}"


def _month(text: str) -> date:
    return date.fromisoformat(str(text)[:10]).replace(day=1)


def _index(start: date, month: date) -> int:
    return (month.year - start.year) * 12 + month.month - start.month


def _num(params: dict[str, Any], key: str) -> float:
    value = params.get(key)
    try:
        return 0.0 if value in (None, "") else float(value)
    except (TypeError, ValueError):
        return 0.0


def _row_formulas(r: int) -> dict[str, str]:
    """Формулы строки r — построчный пересказ `hotel_flows` языком Excel."""
    first = FIRST_ROW
    A = f"{_c('k')}{r}"
    v = {name: f"{_c(name)}{r}" for name in COLUMNS}
    p = {name: f"{_c(name)}{r - 1}" for name in COLUMNS}
    has_prev = r > first

    def prev(name: str, default: str = "0") -> str:
        return p[name] if has_prev else default

    i = _in
    open_ = v["open"]
    r12 = f"{v['rate']}/12"
    left = f"({i('maturity')}-{A})"
    refund_window = f"{_c('vatin')}{max(first, r - 3)}:{_c('vatin')}{r}"
    refund_lag = f"{_c('vatin')}{r - 3}" if r - 3 >= first else "0"
    irr_p = f'IF(\'{TOTALS_SHEET}\'!$B$5="",0,(1+\'{TOTALS_SHEET}\'!$B$5)^(1/12)-1)'
    irr_e = f'IF(\'{TOTALS_SHEET}\'!$B$7="",0,(1+\'{TOTALS_SHEET}\'!$B$7)^(1/12)-1)'
    return {
        "in_h": f"=IF({A}<={i('horizon')},1,IF({A}<={i('horizon')}+12,0,-1))",
        "profit_after_tax": f"={v['margin']}-{v['fin_ded']}-{v['tax']}",
        "year_h": f"=IF({v['in_h']}=1,{v['year']},0)",
        "open": f"=IF({A}>={i('comm')},{A}-{i('comm')}+1,0)",
        "occ": (f"=IF({open_}=0,0,IF({i('ramp')}=0,{i('occ_target')},{i('occ_start')}"
                f"+({i('occ_target')}-{i('occ_start')})*MAX(0,MIN(1,({open_}-1)/{i('ramp')}))))"),
        "growth": f"=(1+MAX(-0.95,{i('index')}))^(({A}-{i('price_k')})/12)",
        "adr": f"={i('adr')}*{v['growth']}",
        "avail": f"=IF({open_}>0,{i('keys')}*{v['days']},0)",
        "sold": f"={v['avail']}*{v['occ']}",
        "rooms": f"={v['sold']}*{v['adr']}*{i('rev_mult')}",
        "fnb": f"={v['rooms']}*{i('fnb')}",
        "other": f"=({v['rooms']}+{v['fnb']})*{i('other')}",
        "dept": f"={v['rooms']}+{v['fnb']}+{v['other']}",
        "relief": f"=IF(AND({open_}>0,{open_}<={i('relief_months')}),{v['rooms']}*{i('vat')},0)",
        "rooms_exp": f"={v['rooms']}*{i('rooms_cost')}",
        "fnb_exp": f"={v['fnb']}*{i('fnb_cost')}",
        "other_exp": f"={v['other']}*{i('other_cost')}",
        "ag": f"={v['dept']}*{i('ag')}",
        "sm": f"={v['dept']}*{i('sm')}",
        "pom": f"={v['dept']}*{i('pom')}",
        "util": f"={v['dept']}*{i('util')}",
        "gop": (f"={v['dept']}-{v['rooms_exp']}-{v['fnb_exp']}-{v['other_exp']}"
                f"-{v['ag']}-{v['sm']}-{v['pom']}-{v['util']}"),
        "base_fee": f"={v['dept']}*{i('base_fee')}",
        "incentive": f"=MAX(0,{v['gop']})*{i('incentive')}",
        "reserve": f"={v['dept']}*{i('reserve')}",
        "book_b_open": f"={prev('book_b_close', i('building_net'))}",
        "dep_b": f"=IF({open_}>0,MIN({v['book_b_open']},{i('building_net')}/{i('building_months')}),0)",
        "book_b_close": f"={v['book_b_open']}-{v['dep_b']}",
        "book_f_open": f"={prev('book_f_close', i('ffe_net'))}",
        "dep_f": f"=IF({open_}>0,MIN({v['book_f_open']},{i('ffe_net')}/{i('ffe_months')}),0)",
        "book_f_close": f"={v['book_f_open']}-{v['dep_f']}",
        "ptax": f"=IF({open_}>0,({v['book_b_open']}-{v['dep_b']}/2)*{i('property_rate')}/12,0)",
        "ins": f"=IF({open_}>0,{i('insurance_month')},0)",
        "revenue": f"={v['dept']}+{v['relief']}",
        "opex": f"={v['dept']}-{v['gop']}+{v['base_fee']}+{v['incentive']}+{v['reserve']}",
        "ebitda": f"={v['revenue']}-{v['opex']}-{v['ptax']}-{v['ins']}",
        "exit": f'=IF(AND({i("exit_mode")}="sale",{A}={i("horizon")}),{i("exit_value")},0)',
        "exit_cost": f"={v['exit']}*{i('exit_cost')}",
        "residual": f'=IF(AND({i("exit_mode")}="hold",{A}={i("horizon")}),{i("exit_value")},0)',
        "vatin": f"=IF({v['in_h']}=1,MAX(0,{v['capex']}-{v['capex_land']})*{i('vs')},0)",
        "refund": (f"=IF({A}>{i('horizon')},0,IF({A}={i('horizon')},SUM({refund_window}),"
                   f"{refund_lag}))"),
        "rate": (f'=IF({i("financing")}="none",0,IF({i("financing")}="preferential",'
                 f"{i('pref_share')}*{v['key_rate']}+{i('pref_margin')},"
                 f"{v['key_rate']}+{i('spread')}))"),
        "bal_open": f"={prev('bal_close')}",
        "interest": f"={v['bal_open']}*{v['rate']}/12",
        "int_cap": f"=IF({A}<{i('comm')},{v['interest']},0)",
        "int_paid": f"=IF(AND({A}>={i('comm')},{v['in_h']}=1),{v['interest']},0)",
        "want": f'=IF(AND({A}<={i("comm")},{v["in_h"]}=1,{i("financing")}<>"none"),{v["capex"]}*{i("loan_share")},0)',
        "cum_prev": f"={prev('cum_draw')}",
        "draw": (f"=IF({i('limit_on')}=1,MIN({v['want']},MAX(0,{i('limit')}-{v['cum_prev']})),"
                 f"{v['want']})"),
        "cum_draw": f"={v['cum_prev']}+{v['draw']}",
        "fee": f"={v['draw']}*{i('fee')}",
        "first_k": f"=IF({v['draw']}>0,{A},{NO_DRAW})",
        "bal_after": f"={v['bal_open']}+{v['int_cap']}+{v['draw']}",
        "prepay": f"=IF({v['bal_after']}>0,MIN({v['bal_after']},{v['refund']}),0)",
        "bal_pre": f"={v['bal_after']}-{v['prepay']}",
        "annuity": (f"=IF(AND({A}>={i('amort_start')},{A}<{i('maturity')},{v['bal_pre']}>0),"
                    f"IF({v['rate']}>0,MIN({v['bal_pre']},MAX(0,{v['bal_pre']}*{r12}"
                    f"/(1-(1+{r12})^(-{left}))-{v['interest']})),"
                    f"MIN({v['bal_pre']},MAX(0,{v['bal_pre']}/{left}-{v['interest']}))),0)"),
        "sched": (f"=IF(OR({v['bal_pre']}<=0,{A}<{i('comm')},{A}>{i('horizon')}),0,"
                  f"IF(OR({A}={i('horizon')},{A}>={i('maturity')}),{v['bal_pre']},"
                  f"IF({A}>={i('amort_start')},"
                  f'IF({i("repayment")}="sweep",MIN({v["bal_pre"]},MAX(0,{v["ebitda"]}-{v["interest"]})),'
                  f'IF({i("repayment")}="sculpted",MIN({v["bal_pre"]},MAX(0,MAX(0,{v["ebitda"]})'
                  f"/{i('dscr')}-{v['interest']})),{v['annuity']})),0)))"),
        # Удержание свободного потока (как в `hotel_flows`): пока долг не
        # погашен, деньги после процентов и налога копятся в проекте и гасят
        # кредит сверх планового платежа, когда погашение разрешено.
        "free": (f"=IF(AND({A}>={i('comm')},{A}<{i('horizon')}),"
                 f"{v['ebitda']}-{v['int_paid']}-{v['fee']}-{v['tax']},0)"),
        "res_open": f"={prev('res_close')}",
        "sweep": (f"=IF(AND({A}>={i('comm')},{A}<{i('horizon')},{v['bal_pre']}>{v['sched']},"
                  f"OR({A}>={i('amort_start')},{A}>={i('maturity')})),"
                  f"MIN({v['bal_pre']}-{v['sched']},MAX(0,{v['res_open']}+{v['free']}-{v['sched']})),0)"),
        "repay": f"={v['sched']}+{v['sweep']}",
        "repay_total": f"={v['repay']}+{v['prepay']}",
        "bal_close": f"={v['bal_pre']}-{v['repay']}",
        "res_close": (f"=IF(OR({A}<{i('comm')},{A}>={i('horizon')},{v['bal_close']}<=0.000001),0,"
                      f"MAX(0,{v['res_open']}+{v['free']}-{v['repay']}))"),
        "retained": f"={v['res_close']}-{v['res_open']}",
        "amort": f"=IF({A}={i('horizon')},0,{v['sched']})",
        "dep": f"=IF({v['in_h']}=1,{v['dep_b']}+{v['dep_f']},0)",
        "margin": (f"=IF({v['in_h']}<>1,0,{v['ebitda']}+{v['exit']}-{v['exit_cost']}-{v['dep']}"
                   f'-IF(AND({A}={i("horizon")},{i("exit_mode")}="sale"),'
                   f"{v['book_b_close']}+{v['book_f_close']}+{i('land_total')},0))"),
        "fin_ded": f"=IF({v['in_h']}=1,{v['int_cap']}+{v['int_paid']}+{v['fee']},0)",
        "def_acc": f"=IF({A}<{i('comm')},{prev('def_acc')}+{v['margin']}-{v['fin_ded']},0)",
        "newyear": (f"=IF({A}<={i('comm')},0,IF({v['year']}<>{prev('year', v['year'])},1,0))"),
        "prior": (f"=IF({A}<={i('comm')},0,IF({v['newyear']}=1,{prev('prior')}"
                  f"+IF({prev('yres')}<0,-{prev('yres')},-{prev('used')}),{prev('prior')}))"),
        "yres": (f"=IF({A}<{i('comm')},0,IF(OR({A}={i('comm')},{v['newyear']}=1),0,{prev('yres')})"
                 f"+{v['margin']}-{v['fin_ded']}+IF({A}={i('comm')},{prev('def_acc')},0))"),
        "used": (f"=IF(AND({A}>={i('comm')},{v['yres']}>0,{v['prior']}>0),"
                 f"MIN({v['prior']},{v['yres']}*{i('loss_limit')}),0)"),
        "base": f"=IF({A}<{i('comm')},0,MAX({v['yres']}-{v['used']},0))",
        "ypaid_prev": f"=IF(OR({A}<={i('comm')},{v['newyear']}=1),0,{prev('ypaid')})",
        "tax": (f"=IF(OR({A}<{i('comm')},{v['in_h']}<>1),0,"
                f"MAX({v['base']}*{i('tax_rate')}-{v['ypaid_prev']},0))"),
        "ypaid": f"={v['ypaid_prev']}+{v['tax']}",
        "vat_paid": f"=-{v['refund']}",
        "operating": (f"=IF({v['in_h']}<>1,0,{v['ebitda']}+{v['exit']}+{v['residual']}"
                      f"-{v['exit_cost']}-{v['vat_paid']})"),
        "to_equity": (f"={v['operating']}-{v['int_paid']}-{v['fee']}+{v['draw']}"
                      f"-{v['repay_total']}-{v['retained']}"),
        "project_cf": (f"=IF({v['in_h']}<>1,0,{v['operating']}-{v['capex']}-{v['int_cap']}"
                       f"-{v['int_paid']}-{v['fee']}-{v['tax']})"),
        "equity_cf": f"=IF({v['in_h']}<>1,0,{v['to_equity']}-{v['capex']}-{v['tax']})",
        "fcff": f"=IF({v['in_h']}<>1,0,{v['operating']}-{v['capex']}-{v['tax']})",
        "disc_p": f"={v['project_cf']}/(1+{i('rate_m')})^{A}",
        "disc_e": f"={v['equity_cf']}/(1+{i('discount')}/12)^{A}",
        "irr_p": f"={v['project_cf']}/(1+{irr_p})^{A}",
        "irr_e": f"={v['equity_cf']}/(1+{irr_e})^{A}",
        "cum_p": f"={prev('cum_p')}+{v['project_cf']}",
        "cum_e": f"={prev('cum_e')}+{v['equity_cf']}",
        "cross_p": (f"=IF(AND({prev('cum_p')}<0,{v['cum_p']}>=0),"
                    f"{A}+(-{prev('cum_p')})/{v['project_cf']},{NO_DRAW})"),
        "cum_dp": f"={prev('cum_dp')}+{v['disc_p']}",
        "cross_dp": (f"=IF(AND({prev('cum_dp')}<0,{v['cum_dp']}>=0),"
                     f"{A}+(-{prev('cum_dp')})/{v['disc_p']},{NO_DRAW})"),
    }


def _model_values(data: dict[str, Any], start: date) -> dict[str, Any]:
    params = hs.effective_params(data.get("params") or {})
    pct = lambda key: _num(params, key) / 100.0  # noqa: E731
    commissioning = _month(data["commissioning"])
    financing = str(params.get("financing") or hs.FINANCING_COMMERCIAL)
    limit_blank = params.get("loan_limit_th_per_key") in (None, "")
    keys = max(0.0, _num(params, "keys"))
    repayment = str(params.get("repayment") or hs.REPAY_ANNUITY)
    dscr = _num(params, "dscr_target")
    if repayment == hs.REPAY_SCULPTED and dscr <= 0:
        dscr = 1.0
    price_date = params.get("adr_price_date") or data.get("start")
    hold = max(1, int(_num(params, "hold_years") or 1))
    return {
        "keys": keys, "adr": max(0.0, _num(params, "adr_rub")),
        "price_k": _index(start, _month(price_date)),
        "index": pct("index_pct"),
        "occ_start": min(1.0, max(0.0, pct("occ_start_pct"))),
        "occ_target": min(1.0, max(0.0, pct("occ_target_pct"))),
        "ramp": max(0, int(_num(params, "ramp_months"))),
        "fnb": max(0.0, pct("fnb_pct")), "other": max(0.0, pct("other_pct")),
        "rooms_cost": max(0.0, pct("rooms_cost_pct")), "fnb_cost": max(0.0, pct("fnb_cost_pct")),
        "other_cost": max(0.0, pct("other_cost_pct")),
        "ag": max(0.0, pct("ag_pct")), "sm": max(0.0, pct("sm_pct")),
        "pom": max(0.0, pct("pom_pct")), "util": max(0.0, pct("utilities_pct")),
        "base_fee": max(0.0, pct("base_fee_pct")), "incentive": max(0.0, pct("incentive_fee_pct")),
        "reserve": max(0.0, pct("ffe_reserve_pct")),
        "relief_months": max(0, int(round(_num(params, "vat_relief_years") * 12))),
        "rev_mult": float(data.get("revenue_multiplier") or 1.0),
        "vat": float(data["vat_rate"]),
        "property_rate": max(0.0, pct("property_tax_pct")),
        "insurance_rate": max(0.0, pct("insurance_pct")),
        "building_months": max(1, int(_num(params, "building_years") or 30)) * 12,
        "ffe_months": max(1, int(_num(params, "ffe_years") or 8)) * 12,
        "financing": financing if financing in dict(hs.FINANCINGS) else hs.FINANCING_COMMERCIAL,
        "loan_share": min(1.0, max(0.0, pct("loan_share_pct"))),
        "spread": pct("loan_spread_pp"),
        "pref_share": min(1.0, max(0.0, pct("pref_key_share_pct"))),
        "pref_margin": pct("pref_margin_pp"),
        "limit_on": 0 if limit_blank else 1,
        "limit": 0.0 if limit_blank else _num(params, "loan_limit_th_per_key") * 1000.0 * keys,
        "fee": max(0.0, pct("loan_fee_pct")),
        "term": max(1, int(_num(params, "loan_term_years") or 1)) * 12,
        "grace": max(0, int(_num(params, "grace_months"))),
        "repayment": repayment, "dscr": dscr,
        "exit_mode": str(params.get("exit_mode") or hs.EXIT_SALE),
        "valuation": str(params.get("valuation") or hs.VALUATION_CAP),
        "cap": pct("exit_cap_pct"), "multiple": _num(params, "exit_multiple"),
        "exit_cost": max(0.0, pct("exit_cost_pct")),
        "tax_rate": float(data["profit_tax_rate"]), "loss_limit": LOSS_CARRY_USE_LIMIT,
        "discount": float(data["discount_rate"]),
        "comm": _index(start, commissioning),
        "horizon": _index(start, hs.add_months(commissioning, hold * 12 - 1)),
        "hold": hold,
        "limit_th": 0.0 if limit_blank else _num(params, "loan_limit_th_per_key"),
        "start": start, "price_date": _month(price_date),
    }


def _origin_text(field: hs.HotelField, origins: dict[str, Any], raw: Any) -> str:
    origin = origins.get(field.key)
    if origin and raw not in (None, ""):
        return str(origin.get("text") or "")
    if raw in (None, "") and field.default is not None:
        return "умолчание: " + field.origin
    if raw in (None, ""):
        return "не задано"
    return "введено вручную"


def _inputs_sheet(book: Workbook, data: dict[str, Any], values: dict[str, Any],
                  rows: int) -> None:
    ws = book.create_sheet(INPUTS_SHEET, 0)
    ws["A1"] = "Вводные гостиницы — как в проекте, с происхождением"
    ws["A1"].font = BOLD
    ws["A3"] = "Модельная величина"
    ws["C3"] = "Значение"
    for cell in ("A3", "C3"):
        ws[cell].font = BOLD
    for key, label in MODEL:
        row = _ROWS[key]
        ws[f"A{row}"] = label
        ws[f"C{row}"] = CHOICES[key].get(values[key], values[key]) if key in CHOICES else values[key]
        ws[f"C{row}"].fill = INPUT_FILL
        if key in CHOICES:
            formula = '""'
            for code, label in reversed(list(CHOICES[key].items())):
                formula = f'IF(C{row}="{label}","{code}",{formula})'
            ws[f"I{row}"] = "=" + formula
            validation = DataValidation(type="list", formula1='"' + ",".join(CHOICES[key].values()) + '"')
            validation.errorTitle = "Выберите вариант"
            validation.error = "Используйте один из вариантов списка."
            validation.showErrorMessage = True
            validation.errorStyle = "stop"
            ws.add_data_validation(validation)
            validation.add(f"C{row}")
        if "доля" in label.lower() or "спред" in label.lower() or "маржа" in label.lower() or key in ("vat", "tax_rate", "discount"):
            ws[f"C{row}"].number_format = PERCENT
    ws[f"C{_ROWS['limit']}"] = f"={_in('limit_th')}*1000*{_in('keys')}"
    ws[f"C{_ROWS['price_k']}"] = (f"=(YEAR({_in('price_date')})-YEAR({_in('start')}))*12"
                                   f"+MONTH({_in('price_date')})-MONTH({_in('start')})")
    if not (data.get("params") or {}).get("adr_price_date"):
        ws[f"C{_ROWS['price_date']}"] = "=" + _in("start")
        ws[f"C{_ROWS['price_date']}"].fill = PatternFill()
    for key in ("start", "price_date"):
        ws[f"C{_ROWS[key]}"].number_format = "dd.mm.yyyy"
    if data.get("capex_spec"):
        spec = data["capex_spec"]
        ws[f"C{_ROWS['comm']}"] = (f"=MAX({spec['ird_min']},INT({capex_book.ref('ird_months')}))"
                                   f"+INT({capex_book.ref('construction_months')})")
        ws[f"C{_ROWS['horizon']}"] = f"={_in('comm')}+MAX(1,INT({_in('hold')}))*12-1"
    for key in ("limit", "comm", "horizon", "price_k"):
        ws[f"C{_ROWS[key]}"].fill = PatternFill()
    last = f"{CALC_SHEET}"
    col = lambda name: f"'{last}'!${_c(name)}${FIRST_ROW}:${_c(name)}${FIRST_ROW + rows - 1}"  # noqa: E731
    i = _in
    derived = {
        "vs": f"={i('vat')}/(1+{i('vat')})",
        "capex_total": f"=SUMIF({col('in_h')},1,{col('capex')})",
        "land_total": f"=SUMIF({col('in_h')},1,{col('capex_land')})",
        "ffe_total": f"=SUMIF({col('in_h')},1,{col('capex_ffe')})",
        "capex_vat": f"=MAX(0,{i('capex_total')}-{i('land_total')})*{i('vs')}",
        "capex_net": f"={i('capex_total')}-{i('capex_vat')}",
        "ffe_net": f"={i('ffe_total')}*(1-{i('vs')})",
        "building_net": f"=MAX(0,{i('capex_net')}-{i('ffe_net')}-{i('land_total')})",
        "insurance_month": f"=({i('capex_net')}-{i('land_total')})*{i('insurance_rate')}/12",
        "first_draw": f"=MIN({col('first_k')})",
        "maturity": (f"=IF({i('first_draw')}<{NO_DRAW},{i('first_draw')}+{i('term')},{NO_DRAW})"),
        "amort_start": (f"=IF({i('first_draw')}<{NO_DRAW},MAX({i('comm')},{i('first_draw')}"
                        f"+{i('grace')}),{NO_DRAW})"),
        "forward_ebitda": f"=SUMIF({col('in_h')},0,{col('ebitda')})",
        "exit_value": (f'=IF({i("forward_ebitda")}<=0,0,IF({i("valuation")}="ev_ebitda",'
                       f"{i('forward_ebitda')}*{i('multiple')},"
                       f"IF({i('cap')}>0,{i('forward_ebitda')}/{i('cap')},0)))"),
        "peak": f"=MAX({col('bal_close')})",
        "rate_m": f"=(1+{i('discount')})^(1/12)-1",
    }
    head = _ROWS[DERIVED[0][0]] - 1
    ws[f"A{head}"] = "Производные (формулы)"
    ws[f"A{head}"].font = BOLD
    for key, label in DERIVED:
        row = _ROWS[key]
        ws[f"A{row}"] = label
        ws[f"C{row}"] = derived[key]

    # Поля проекта как их видит человек: значение и откуда оно.
    params = data.get("params") or {}
    origins = data.get("origins") or {}
    ws["E3"] = "Поле проекта"
    ws["F3"] = "Ед."
    ws["G3"] = "Исходное значение проекта (для происхождения)"
    ws["H3"] = "Происхождение"
    for cell in ("E3", "F3", "G3", "H3"):
        ws[cell].font = BOLD
    for n, field in enumerate(hs.FIELDS, start=4):
        raw = (data.get("raw_params") or {}).get(field.key, params.get(field.key))
        ws[f"E{n}"] = field.label
        ws[f"F{n}"] = field.unit
        shown = params.get(field.key)
        if field.kind == "choice":
            shown = dict(field.choices).get(str(shown), shown)
        ws[f"G{n}"] = "" if shown is None else shown
        ws[f"G{n}"].fill = ENGINE_FILL
        ws[f"H{n}"] = _origin_text(field, origins, raw)
    for letter, width in (("A", 54), ("B", 2), ("C", 20), ("D", 2), ("E", 46), ("F", 22),
                          ("G", 22), ("H", 90)):
        ws.column_dimensions[letter].width = width
    ws.column_dimensions["I"].hidden = True
    ws["A2"] = "Редактируйте жёлтые ячейки. Доли показаны в %. Оранжевое справа — происхождение исходного проекта."
    ws["E119"] = "Прогноз ключевой ставки, выбранный в проекте"
    horizon_validation = DataValidation(type="custom", formula1=f'$C${_ROWS["horizon"]}<={rows - 13}')
    horizon_validation.showErrorMessage = True
    horizon_validation.errorStyle = "stop"
    horizon_validation.error = f"Горизонт и 12 месяцев оценки выхода должны помещаться в {rows} месяцев сетки."
    ws.add_data_validation(horizon_validation)
    horizon_validation.add(f"C{_ROWS['hold']}")


def _calc_sheet(book: Workbook, data: dict[str, Any], start: date, values: dict[str, Any],
                rows: int, capex_ref=None) -> int:
    ws = book.create_sheet(CALC_SHEET)
    ws["A1"] = "Помесячный расчёт гостиницы — связанные формулы книги"
    ws["A1"].font = BOLD
    ws["A2"] = ("Строки после конца срока (В сроке = 0) — 12 месяцев для оценки выхода: "
                "в итоги они не входят.")
    head = FIRST_ROW - 1
    for name, (letter, label) in COLUMNS.items():
        ws[f"{letter}{head}"] = label
        ws[f"{letter}{head}"].font = BOLD
    for name in ("capex", "capex_ffe", "capex_land"):
        ws[f"{_c(name)}{head - 1}"] = "СМЕТА И ГРАФИК" if capex_ref else "СЧИТАЕТ ДВИЖОК"
    for name in ("year", "days"):
        ws[f"{_c(name)}{head - 1}"] = "календарь"
    months = data["months"]
    total = rows
    for idx in range(total):
        r = FIRST_ROW + idx
        month = hs.add_months(start, idx)
        source = months[idx] if idx < len(months) else {}
        ws[f"{_c('k')}{r}"] = idx
        month_ref = f"{_c('month')}{r}"
        ws[month_ref] = f"=EDATE({_in('start')},{_c('k')}{r})"
        ws[month_ref].number_format = "mmm yyyy"
        ws[f"{_c('year')}{r}"] = f"=YEAR({month_ref})"
        ws[f"{_c('days')}{r}"] = f"=EDATE({month_ref},1)-{month_ref}"
        for name, key in (("capex", "capex"), ("capex_ffe", "capex_ffe"),
                          ("capex_land", "capex_land"), ("key_rate", "key_rate")):
            if capex_ref and name != "key_rate":
                parts = {"capex": ["total"], "capex_ffe": ["hotel_ffe"],
                         "capex_land": ["purchase", "land_rights", "vri_interest", "vri_security"]}[name]
                ws[f"{_c(name)}{r}"] = "=" + "+".join(capex_ref(part, idx) for part in parts)
                continue
            value = float(source.get(key, 0.0) or 0.0)
            if name == "key_rate" and idx >= len(months) and months:
                value = float(months[-1].get("key_rate", 0.0) or 0.0)
            if name == "key_rate":
                input_row = 120 + idx
                inp = book[INPUTS_SHEET]
                forecast = (data.get("capex_spec") or {}).get("forecast") or []
                if idx < len(forecast):
                    value = float(forecast[idx][1])
                inp[f"E{input_row}"] = month
                inp[f"E{input_row}"].number_format = "mmm yyyy"
                inp[f"F{input_row}"] = value
                inp[f"F{input_row}"].fill = INPUT_FILL
                inp[f"F{input_row}"].number_format = PERCENT
                end = 119 + total
                ws[f"{_c(name)}{r}"] = (f"=INDEX('{INPUTS_SHEET}'!$F$120:$F${end},"
                                         f'MAX(1,COUNTIF(\'{INPUTS_SHEET}\'!$E$120:$E${end},"<="&{month_ref})))')
            else:
                ws[f"{_c(name)}{r}"] = value
                ws[f"{_c(name)}{r}"].fill = ENGINE_FILL
        for name, formula in _row_formulas(r).items():
            ws[f"{_c(name)}{r}"] = formula
    for idx in range(1, len(COLUMNS) + 1):
        ws.column_dimensions[get_column_letter(idx)].width = 15
    ws.freeze_panes = f"C{FIRST_ROW}"
    return total


# Годовой USALI: (ключ, подпись, колонка расчёта или формула от строк года).
USALI_LINES: tuple[tuple[str, str, str], ...] = (
    ("avail", "Номеро-ночи в продаже", "sum"),
    ("sold", "Проданные номеро-ночи", "sum"),
    ("occupancy", "Загрузка", "ratio:sold:avail"),
    ("adr", "ADR, ₽ без НДС", "ratio:rooms:sold"),
    ("revpar", "RevPAR, ₽ без НДС", "ratio:rooms:avail"),
    ("rooms", "Выручка номерного фонда", "sum"),
    ("fnb", "Выручка F&B", "sum"),
    ("other", "Выручка прочих департаментов", "sum"),
    ("dept", "Выручка департаментов", "sum"),
    ("relief", "Льгота НДС на проживание", "sum"),
    ("revenue", "Выручка всего", "sum"),
    ("rooms_exp", "Расходы номерного фонда", "sum"),
    ("fnb_exp", "Расходы F&B", "sum"),
    ("other_exp", "Расходы прочих департаментов", "sum"),
    ("ag", "A&G", "sum"),
    ("sm", "S&M", "sum"),
    ("pom", "POM", "sum"),
    ("util", "Utilities", "sum"),
    ("gop", "GOP", "sum"),
    ("base_fee", "Базовое вознаграждение оператора", "sum"),
    ("incentive", "Поощрительное вознаграждение оператора", "sum"),
    ("reserve", "Резерв FF&E", "sum"),
    ("ptax", "Налог на имущество", "sum"),
    ("ins", "Страхование", "sum"),
    ("ebitda", "EBITDA", "sum"),
    ("dep", "Амортизация", "sum"),
    ("interest", "Проценты по кредиту", "sum:int_cap+int_paid"),
    ("tax", "Налог на прибыль", "sum"),
    ("int_paid", "Проценты уплаченные", "sum"),
    ("amort", "Плановое тело", "sum"),
    ("dscr", "DSCR (EBITDA ÷ проценты и плановое тело)", "dscr"),
    ("capex", "CAPEX с НДС", "sum"),
    ("fcff", "Поток до финансирования (FCFF)", "sum"),
    ("equity_cf", "Поток капитала (FCFE)", "sum"),
    ("profit_after_tax", "Финансовый результат после налога", "sum"),
)


def _usali_sheet(book: Workbook, years: list[int], rows: int) -> dict[tuple[str, int], str]:
    ws = book.create_sheet(USALI_SHEET)
    ws["A1"] = "USALI по годам — свод помесячного расчёта (формулы)"
    ws["A1"].font = BOLD
    ws["A3"] = "Статья"
    ws["A3"].font = BOLD
    col = lambda name: f"'{CALC_SHEET}'!${_c(name)}${FIRST_ROW}:${_c(name)}${FIRST_ROW + rows - 1}"  # noqa: E731
    line_row = {key: 4 + n for n, (key, _, _) in enumerate(USALI_LINES)}
    refs: dict[tuple[str, int], str] = {}
    for j, year in enumerate(years):
        letter = get_column_letter(2 + j)
        ws[f"{letter}3"] = f"=YEAR({_in('start')})+{j}"
        ws[f"{letter}3"].font = BOLD
        for key, label, how in USALI_LINES:
            r = line_row[key]
            ws[f"A{r}"] = label
            if how == "sum":
                formula = f"=SUMIF({col('year_h')},{letter}$3,{col(key)})"
            elif how.startswith("sum:"):
                parts = how[4:].split("+")
                formula = "=" + "+".join(f"SUMIF({col('year_h')},{letter}$3,{col(part)})"
                                         for part in parts)
            elif how.startswith("ratio:"):
                _, num, den = how.split(":")
                formula = (f"=IF({letter}{line_row[den]}=0,0,"
                           f"{letter}{line_row[num]}/{letter}{line_row[den]})")
            else:  # dscr
                ds = f"({letter}{line_row['int_paid']}+{letter}{line_row['amort']})"
                formula = (f'=IF(AND({ds}>1,{letter}{line_row["avail"]}>0),'
                           f'{letter}{line_row["ebitda"]}/{ds},"")')
            ws[f"{letter}{r}"] = formula
            refs[(key, year)] = f"'{USALI_SHEET}'!${letter}${r}"
    ws.column_dimensions["A"].width = 46
    for j in range(len(years)):
        ws.column_dimensions[get_column_letter(2 + j)].width = 16
    ws.freeze_panes = "B4"
    refs[("dscr_range", 0)] = (f"'{USALI_SHEET}'!$B${line_row['dscr']}:"
                               f"${get_column_letter(1 + max(1, len(years)))}${line_row['dscr']}")
    return refs


def _totals(rows: int) -> dict[str, str]:
    col = lambda name: f"'{CALC_SHEET}'!${_c(name)}${FIRST_ROW}:${_c(name)}${FIRST_ROW + rows - 1}"  # noqa: E731

    def in_h(name: str) -> str:
        return f"=SUMIF({col('in_h')},1,{col(name)})"

    return {
        "revenue": in_h("revenue"), "department_revenue": in_h("dept"),
        "rooms_revenue": in_h("rooms"), "fnb_revenue": in_h("fnb"),
        "other_revenue": in_h("other"), "vat_relief": in_h("relief"), "gop": in_h("gop"),
        "opex": in_h("opex"), "property_tax": in_h("ptax"), "insurance": in_h("ins"),
        "ebitda": in_h("ebitda"), "depreciation": in_h("dep"),
        "exit_revenue": f"=SUM({col('exit')})", "exit_cost": f"=SUM({col('exit_cost')})",
        "residual_value": f"=SUM({col('residual')})", "vat_refund": f"=SUM({col('refund')})",
        "loan_draw": f"=SUM({col('draw')})",
        "loan_interest": f"=SUM({col('int_cap')})+SUM({col('int_paid')})",
        "loan_fee": f"=SUM({col('fee')})", "loan_repayment": f"=SUM({col('repay_total')})",
        "profit_tax": f"=SUM({col('tax')})", "tax_margin": f"=SUM({col('margin')})",
        "cash_to_equity": in_h("to_equity"),
        "forward_ebitda": f"={_in('forward_ebitda')}", "exit_value": f"={_in('exit_value')}",
        "loan_peak": f"={_in('peak')}",
        "capex_total": f"={_in('capex_total')}",
        "equity_peak": f"=MAX(0,-MIN({col('cum_e')}))",
        "npv_project": f"=SUM({col('disc_p')})", "npv_equity": f"=SUM({col('disc_e')})",
        "npv_at_irr_project": f"=SUM({col('irr_p')})",
        "npv_at_irr_equity": f"=SUM({col('irr_e')})",
        "irr_project": f'=IFERROR((1+IRR({col("project_cf")},0.01))^12-1,"")',
        "irr_equity": f'=IFERROR((1+IRR({col("equity_cf")},0.01))^12-1,"")',
        "pbp_years": f"=IF(MIN({col('cross_p')})<{NO_DRAW},MIN({col('cross_p')})/12,\"\")",
        "dpbp_years": f"=IF(MIN({col('cross_dp')})<{NO_DRAW},MIN({col('cross_dp')})/12,\"\")",
    }


# Сверка: итог книги → (подпись, где лежит у движка).
CHECKS: tuple[tuple[str, str, str], ...] = (
    ("revenue", "Выручка всего", "totals"),
    ("department_revenue", "Выручка департаментов", "totals"),
    ("rooms_revenue", "Выручка номеров", "totals"),
    ("vat_relief", "Льгота НДС на проживание", "totals"),
    ("gop", "GOP", "totals"),
    ("opex", "Расходы USALI", "totals"),
    ("property_tax", "Налог на имущество", "totals"),
    ("insurance", "Страхование", "totals"),
    ("ebitda", "EBITDA", "totals"),
    ("depreciation", "Амортизация", "totals"),
    ("exit_revenue", "Выход — продажа", "totals"),
    ("exit_cost", "Затраты на продажу", "totals"),
    ("residual_value", "Удержание — оценка", "totals"),
    ("vat_refund", "Возмещение НДС стройки", "totals"),
    ("loan_draw", "Кредит — выдача", "totals"),
    ("loan_interest", "Кредит — проценты", "totals"),
    ("loan_fee", "Кредит — комиссия", "totals"),
    ("loan_repayment", "Кредит — погашение", "totals"),
    ("tax_margin", "Налоговая маржа", "totals"),
    ("profit_tax", "Налог на прибыль", "totals"),
    ("cash_to_equity", "Деньги собственнику без CAPEX", "totals"),
    ("forward_ebitda", "EBITDA 12 мес. после срока", "kpi"),
    ("exit_value", "Стоимость на выходе", "kpi"),
    ("loan_peak", "Кредит — пик долга", "kpi"),
    ("npv_project", "NPV проекта", "kpi"),
    ("npv_equity", "NPV собственного капитала", "kpi"),
    ("irr_project", "IRR проекта", "kpi"),
    ("irr_equity", "IRR собственного капитала", "kpi"),
    ("pbp_years", "Срок окупаемости, лет", "kpi"),
    ("dpbp_years", "Дисконтированный срок окупаемости, лет", "kpi"),
    ("dscr_min", "DSCR — минимум по годам", "kpi"),
)
TOLERANCE_SHARE = 1e-6


def build(result: dict[str, Any], project_name: str = "") -> bytes:
    """Книга гостиничного проекта из посчитанного движком результата."""
    finance = result.get("finance") or {}
    summary = result.get("summary") or {}
    hotel = finance.get("hotel") or {}
    book = Workbook()
    book.remove(book.active)
    if not hotel.get("computed"):
        ws = book.create_sheet("Гостиница")
        ws["A1"] = "Гостиница не считается: не заданы " + ", ".join(hotel.get("missing") or [])
        stream = io.BytesIO()
        book.save(stream)
        return stream.getvalue()
    data = hotel["book"]
    kpi = hotel.get("kpi") or {}
    start = _month(data["months"][0]["month"])
    values = _model_values(data, start)
    # Резерв сетки позволяет менять срок эксплуатации в скачанной книге.
    # Выход всегда ровно 12 месяцев после её текущего горизонта.
    rows = max(360, values["horizon"] + 1 + 12)
    _inputs_sheet(book, data, values, rows)
    capex_ref = None
    missing = []
    if data.get("capex_spec"):
        capex_book.inputs(book, data["capex_spec"], sys.modules[__name__])
        capex_ref, missing = capex_book.build(book, data["capex_spec"], rows, sys.modules[__name__])
    _calc_sheet(book, data, start, values, rows, capex_ref)
    years = list(range(start.year, hs.add_months(start, rows - 13).year + 1))
    refs = _usali_sheet(book, years, rows)
    totals = _totals(rows)

    sheet = book.create_sheet(TOTALS_SHEET, 0)
    sheet["A1"] = f"DevelopAid · гостиница{(' · ' + project_name) if project_name else ''}"
    sheet["A1"].font = BOLD
    sheet["A2"] = "Финансовая модель гостиницы. Изменяйте жёлтые ячейки на «Вводных»: итоги пересчитываются."
    lines = (
        ("NPV проекта", "npv_project"),
        ("IRR проекта", "irr_project"),
        ("NPV собственного капитала", "npv_equity"),
        ("IRR собственного капитала", "irr_equity"),
        ("Срок окупаемости (PBP), лет от начала проекта", "pbp_years"),
        ("Дисконтированный срок окупаемости (DPBP), лет", "dpbp_years"),
        ("Выручка за срок", "revenue"),
        ("EBITDA за срок", "ebitda"),
        ("Налог на прибыль", "profit_tax"),
        ("Стоимость на выходе", "exit_value"),
        ("Кредит — пик долга", "loan_peak"),
        ("Кредит — проценты", "loan_interest"),
        ("DSCR — минимум по годам", "dscr_min"),
        ("Затраты проекта (CAPEX), ₽ с НДС", "capex_total"),
        ("Максимальное вложение собственного капитала, ₽", "equity_peak"),
    )
    total_refs: dict[str, str] = {}
    for n, line in enumerate(lines, start=4):
        label, key = line[0], line[1]
        sheet[f"A{n}"] = label
        if key is None:
            sheet[f"B{n}"] = "" if line[2] is None else float(line[2])
            sheet[f"B{n}"].fill = ENGINE_FILL
            sheet[f"C{n}"] = "СЧИТАЕТ ДВИЖОК"
            continue
        if key == "dscr_min":
            rng = refs[("dscr_range", 0)]
            sheet[f"B{n}"] = f'=IF(COUNT({rng})>0,MIN({rng}),"")'
        else:
            sheet[f"B{n}"] = totals[key]
        total_refs[key] = f"'{TOTALS_SHEET}'!$B${n}"
    sheet["A19"] = "Модель соответствует методике движка, а не заменяет методику исходного Отеля 1."
    sheet["A20"] = "Оранжевое — исходный снимок для сверки. Изменение вводных требует новой сверки с движком."
    labels = (data.get("capex_spec") or {}).get("article_labels") or {}
    sheet["A21"] = ("Не пересчитываются формулами: " + ", ".join(labels.get(key, key) for key in missing)) if missing else "Смета гостиницы считается формулами. Благоустройство отдельно не учтено."
    sheet.column_dimensions["A"].width = 56
    sheet.column_dimensions["B"].width = 22

    checks = book.create_sheet(CHECK_SHEET)
    checks["A1"] = "Сверка формул книги с движком"
    checks["A1"].font = BOLD
    checks.append(["Показатель", "Книга", "Движок", "Разница", "Допуск", "Итог"])
    hidden: dict[str, str] = {}
    for key, label, where in CHECKS:
        engine = (hotel.get(where) or {}).get(key)
        row = checks.max_row + 1
        ref = total_refs.get(key)
        if ref is None:
            hidden[key] = totals[key]
        checks[f"A{row}"] = label
        checks[f"B{row}"] = f"={ref}" if ref else totals[key]
        checks[f"C{row}"] = float(engine) if isinstance(engine, (int, float)) else ""
        checks[f"C{row}"].fill = ENGINE_FILL
        checks[f"D{row}"] = f'=IF(OR(C{row}="",B{row}=""),"",B{row}-C{row})'
        checks[f"E{row}"] = f'=IF(C{row}="","",MAX(1,ABS(C{row})*{TOLERANCE_SHARE}))'
        checks[f"F{row}"] = (f'=IF(AND(C{row}="",B{row}=""),"—",IF(OR(C{row}="",B{row}=""),'
                             f'"РАСХОЖДЕНИЕ",IF(ABS(D{row})<=E{row},"сходится","РАСХОЖДЕНИЕ")))')
    # Годовой USALI — по годам: выручка и EBITDA книги против движка.
    for year_row in hotel.get("annual") or []:
        if not year_row.get("rooms_available"):
            continue
        for key, label, engine_key in (("revenue", "Выручка", "revenue"),
                                       ("ebitda", "EBITDA", "ebitda")):
            row = checks.max_row + 1
            checks[f"A{row}"] = f"{label} {year_row['year']}"
            checks[f"B{row}"] = f"={refs[(key, year_row['year'])]}"
            checks[f"C{row}"] = float(year_row.get(engine_key) or 0.0)
            checks[f"C{row}"].fill = ENGINE_FILL
            checks[f"D{row}"] = f"=B{row}-C{row}"
            checks[f"E{row}"] = f"=MAX(1,ABS(C{row})*{TOLERANCE_SHARE})"
            checks[f"F{row}"] = f'=IF(ABS(D{row})<=E{row},"сходится","РАСХОЖДЕНИЕ")'
    last = checks.max_row
    checks["H2"] = "Вердикт"
    checks["H3"] = f'=IF(COUNTIF(F3:F{last},"РАСХОЖДЕНИЕ")=0,"ПРОЙДЕНО","ЕСТЬ РАСХОЖДЕНИЯ")'
    for letter, width in (("A", 40), ("B", 20), ("C", 20), ("D", 14), ("E", 12), ("F", 14)):
        checks.column_dimensions[letter].width = width
    _snapshot(book)
    _internal_checks(book, rows)
    _model_sheets(book, rows)
    clones = sensitivity_book.build(book, rows, sys.modules[__name__]) if capex_ref else []
    _style(book)
    book.calculation = CalcProperties(calcId=0, fullCalcOnLoad=True, forceFullCalc=True)
    stream = io.BytesIO()
    book.save(stream)
    # Долг и налог читают предыдущий месяц: предварительный расчёт сверху
    # вниз избегает глубокой рекурсии при сохранении значений для телефона.
    evaluator = Evaluator(book)
    for row in range(FIRST_ROW, FIRST_ROW + rows):
        for key, (letter, _) in COLUMNS.items():
            if key not in ("irr_p", "irr_e"):
                evaluator.cell(CALC_SHEET, f"{letter}{row}")
    # Варианты чувствительности считает Excel при открытии (`excel_only`).
    content, report = with_cached_values(stream.getvalue(), evaluator=evaluator,
                                         left_to_excel=sensitivity_book.excel_only(clones))
    if report["unresolved"]:
        raise ValueError("Формулы гостиничной книги не пересчитаны: " + "; ".join(report["examples"]))
    return content


def _snapshot(book: Workbook) -> None:
    """Сверка с исходным движком действительна только на исходных вводных."""
    checks = book[CHECK_SHEET]
    snapshot = book.create_sheet("Исходный снимок")
    snapshot.sheet_state = "hidden"
    inputs = [(ws.title, cell.coordinate, cell.value)
              for ws in book.worksheets for row in ws.iter_rows() for cell in row
              if cell.fill == INPUT_FILL and cell.data_type != "f"]
    for r, (sheet, address, value) in enumerate(inputs, 3):
        snapshot[f"A{r}"] = f"='{sheet}'!{address}"
        snapshot[f"B{r}"] = value
        snapshot[f"C{r}"] = f"=IF(A{r}=B{r},0,1)"
    checks["H4"] = f"=SUM('Исходный снимок'!C3:C{2+len(inputs)})"
    checks["H5"] = "Число изменённых вводных"
    for r in range(3, 3 + len(CHECKS)):
        original = checks[f"F{r}"].value[1:]
        checks[f"F{r}"] = f'=IF($H$4>0,"ВВОДНЫЕ ИЗМЕНЕНЫ",{original})'
    # И годовые строки сверяются с тем же исходным снимком.
    last = max(3 + len(CHECKS), max(cell.row for cell in checks["A"] if cell.value))
    for r in range(3 + len(CHECKS), last + 1):
        if checks[f"F{r}"].data_type == "f":
            original = checks[f"F{r}"].value[1:]
            checks[f"F{r}"] = f'=IF($H$4>0,"ВВОДНЫЕ ИЗМЕНЕНЫ",{original})'
    checks["H3"] = (f'=IF(COUNTIF(F3:F{last},"РАСХОЖДЕНИЕ")>0,"ЕСТЬ РАСХОЖДЕНИЯ",'
                     'IF(H4>0,"ВВОДНЫЕ ИЗМЕНЕНЫ — НУЖНА НОВАЯ СВЕРКА","ПРОЙДЕНО"))')


def _model_sheets(book: Workbook, rows: int) -> None:
    """Читаемые денежный поток, кредит и налоги — ссылки на общий расчёт."""
    groups = (
        ("ОПУ", (("revenue", "Выручка эксплуатации"), ("opex", "Расходы по USALI"),
                 ("ptax", "Налог на имущество"), ("ins", "Страхование"), ("ebitda", "EBITDA"),
                 ("dep", "Амортизация"), ("int_cap", "Проценты капитализированные"),
                 ("int_paid", "Проценты уплаченные"), ("fee", "Комиссия за выдачу"),
                 ("margin", "Налоговая маржа с выходом"), ("tax", "Налог на прибыль"),
                 ("profit_after_tax", "Финансовый результат после налога"))),
        ("Денежный поток", (("capex", "Затраты проекта"), ("revenue", "Выручка эксплуатации"),
                           ("ebitda", "EBITDA"), ("exit", "Продажа гостиницы"),
                           ("residual", "Оценка при удержании"), ("tax", "Налог на прибыль"),
                           ("project_cf", "Денежный поток проекта"), ("equity_cf", "Денежный поток собственника"),
                           ("cum_e", "Деньги собственника нарастающим"))),
        ("Кредит", (("key_rate", "Ключевая ставка"), ("rate", "Ставка кредита"),
                    ("bal_open", "Долг на начало месяца"), ("draw", "Выдача кредита"),
                    ("int_cap", "Проценты капитализированные"), ("int_paid", "Проценты уплаченные"),
                    ("fee", "Комиссия за выдачу"), ("repay_total", "Погашение кредита"),
                    ("bal_close", "Долг на конец месяца"))),
        ("Налоги", (("vatin", "Входящий НДС"), ("refund", "Возмещение НДС"),
                    ("ptax", "Налог на имущество"), ("dep", "Амортизация"),
                    ("margin", "Налоговая маржа"), ("prior", "Убытки прошлых лет"),
                    ("used", "Зачёт убытка"), ("base", "База налога на прибыль"), ("tax", "Налог на прибыль"))),
    )
    for name, lines in groups:
        ws = book.create_sheet(name)
        ws["A1"] = f"{name} гостиницы"
        ws["A2"] = "Связанные формулы помесячного расчёта. Суммы в ₽; ставки — % годовых."
        ws["A3"] = "Месяц"
        for index, (_, label) in enumerate(lines, 2):
            ws.cell(3, index, label)
        for index in range(rows):
            r = 4 + index
            calc_r = FIRST_ROW + index
            ws[f"A{r}"] = f"='{CALC_SHEET}'!{_c('month')}{calc_r}"
            ws[f"A{r}"].number_format = "mmm yyyy"
            for col, (key, _) in enumerate(lines, 2):
                cell = ws.cell(r, col, f"='{CALC_SHEET}'!{_c(key)}{calc_r}")
                cell.number_format = PERCENT if key in ("key_rate", "rate") else MONEY
        ws.freeze_panes = "B4"
        ws.column_dimensions["A"].width = 16
        for index in range(2, len(lines) + 2):
            ws.column_dimensions[get_column_letter(index)].width = 23
        ws.auto_filter.ref = f"A3:{get_column_letter(1 + len(lines))}{3 + rows}"
    order = (TOTALS_SHEET, INPUTS_SHEET, "Затраты", USALI_SHEET, "ОПУ",
             "Денежный поток", "Кредит", "Налоги", CHECK_SHEET, CALC_SHEET, "Исходный снимок")
    for index, name in enumerate(name for name in order if name in book.sheetnames):
        book.move_sheet(name, offset=index - book.index(book[name]))


def _internal_checks(book: Workbook, rows: int) -> None:
    checks = book[CHECK_SHEET]
    col = lambda key: f"'{CALC_SHEET}'!{_c(key)}{FIRST_ROW}:{_c(key)}{FIRST_ROW+rows-1}"
    lines = [
        ("Баланс долга: выдача + капитализированные проценты − погашение − остаток",
         f"SUM({col('draw')})+SUM({col('int_cap')})-SUM({col('repay_total')})"
         f"-'{CALC_SHEET}'!{_c('bal_close')}{FIRST_ROW+rows-1}"),
        ("Баланс потока проекта: поступления − затраты − финансирование − налог",
         f"SUM({col('project_cf')})-SUM({col('operating')})+SUM({col('capex')})"
         f"+SUM({col('int_cap')})+SUM({col('int_paid')})+SUM({col('fee')})+SUM({col('tax')})"),
    ]
    if capex_book.SHEET in book.sheetnames:
        lines.append(("Баланс сметы и денежного потока",
                      f"SUM({col('capex')})-'{capex_book.SHEET}'!F24"))
        inputs = book[INPUTS_SHEET]
        last = max(r for r in range(77, 120) if inputs[f"G{r}"].fill == INPUT_FILL)
        lines.append(("График покупки: сумма долей равна 100 %",
                      f"SUM('{INPUTS_SHEET}'!G77:G{last})-1"))
    for label, formula in lines:
        row = checks.max_row + 1
        checks[f"A{row}"] = label
        checks[f"B{row}"] = "=" + formula
        checks[f"C{row}"] = 0
        checks[f"D{row}"] = f"=B{row}"
        tolerance = 1e-9 if label.startswith("График покупки") else 1
        checks[f"E{row}"] = tolerance
        checks[f"F{row}"] = f'=IF(ABS(B{row})<=E{row},"сходится","РАСХОЖДЕНИЕ")'
    row = checks.max_row + 1
    checks[f"A{row}"] = "Календарь помещается в сетку, варианты выбраны из списка"
    checks[f"F{row}"] = (f'=IF(AND({_in("horizon")}<={rows-13},{_in("hold")}>0,'
                          + ",".join(f'{_in(key)}<>""' for key in CHOICES)
                          + '),"сходится","РАСХОЖДЕНИЕ")')
    checks["H3"] = (f'=IF(COUNTIF(F3:F{row},"РАСХОЖДЕНИЕ")>0,"ЕСТЬ РАСХОЖДЕНИЯ",'
                     'IF(H4>0,"ВВОДНЫЕ ИЗМЕНЕНЫ — НУЖНА НОВАЯ СВЕРКА","ПРОЙДЕНО"))')


def _style(book: Workbook) -> None:
    """Шрифт/палитра образца, закрепления, единицы и печать своих листов."""
    for ws in book.worksheets:
        if ws.sheet_state == "hidden":
            continue
        ws.sheet_view.showGridLines = False
        ws.sheet_view.zoomScale = 85
        ws.sheet_properties.pageSetUpPr.fitToPage = True
        ws.page_setup.orientation = "landscape"
        ws.page_setup.paperSize = ws.PAPERSIZE_A3
        ws.page_setup.fitToWidth, ws.page_setup.fitToHeight = 1, 0
        header = {INPUTS_SHEET: 3, CALC_SHEET: FIRST_ROW - 1, "Затраты": 4,
                  CHECK_SHEET: 2}.get(ws.title, 3)
        ws.print_title_rows = f"1:{header}"
        ws.freeze_panes = ws.freeze_panes or ("B4" if ws.title != INPUTS_SHEET else "D4")
        # Итог оформления клетки зависит только от её прежнего стиля, строки и
        # того, формула ли в ней: он считается один раз на сочетание и
        # переносится остальным. Новый Font на каждую из десятков тысяч клеток
        # стоил выгрузке секунды.
        styled: dict[tuple[Any, ...], Any] = {}
        for row in ws.iter_rows():
            for cell in row:
                if cell.value is None:
                    continue
                # Клетка без стиля хранит None, а не пустой набор.
                key = (tuple(cell._style or ()), cell.row <= header, cell.row == header,
                       cell.data_type == "f")
                done = styled.get(key)
                if done is not None:
                    cell._style = copy(done)
                    continue
                bold = cell.font.bold
                cell.font = Font(name="Montserrat", size=10, bold=bold, color="203C32")
                cell.alignment = Alignment(vertical="center", wrap_text=cell.row <= header)
                if cell.data_type == "f" and cell.number_format == "General":
                    cell.number_format = MONEY
                if cell.row == header:
                    cell.fill = SECTION_FILL
                    cell.font = Font(name="Montserrat", size=10, bold=True, color="203C32")
                    cell.border = Border(bottom=Side(style="thin", color="00B050"))
                styled[key] = copy(cell._style)
        ws.row_dimensions[header].height = 42
        ws["A1"].font = Font(name="Montserrat", size=16, bold=True, color="FFFFFF")
        ws["A1"].fill = TITLE_FILL
        ws.row_dimensions[1].height = 30
        ws.row_dimensions[2].height = 35
        ws["A2"].alignment = Alignment(wrap_text=True, vertical="center")
        # Титул и пояснение не раздувают печать до ширины технической сетки.
        end = 3 if ws.title == TOTALS_SHEET else min(8, ws.max_column)
        ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=end)
        if ws.title != CHECK_SHEET:
            ws.merge_cells(start_row=2, start_column=1, end_row=2, end_column=end)
    totals = book[TOTALS_SHEET]
    for row in (5, 7):
        totals[f"B{row}"].number_format = PERCENT
    for row in (8, 9, 16):
        totals[f"B{row}"].number_format = "0.00"
    for row in range(4, 19):
        totals[f"A{row}"].alignment = Alignment(wrap_text=True, vertical="center")
        totals.row_dimensions[row].height = 30
        totals[f"B{row}"].font = Font(name="Montserrat", size=12, bold=True, color="203C32")
    for row in (19, 20, 21):
        totals.merge_cells(start_row=row, start_column=1, end_row=row, end_column=3)
        totals[f"A{row}"].alignment = Alignment(wrap_text=True, vertical="center")
        totals.row_dimensions[row].height = 32
    annual = book[USALI_SHEET]
    for cell in annual[3]:
        if cell.column > 1:
            cell.number_format = "0"
    for index, (key, _, _) in enumerate(USALI_LINES, 4):
        for row in annual.iter_rows(min_row=index, max_row=index, min_col=2):
            for cell in row:
                cell.number_format = PERCENT if key == "occupancy" else "0.00" if key == "dscr" else MONEY
    calc = book[CALC_SHEET]
    for key, (letter, _) in COLUMNS.items():
        if key in ("occ", "key_rate", "rate"):
            for row in range(FIRST_ROW, calc.max_row + 1):
                calc[f"{letter}{row}"].number_format = PERCENT
        elif key in ("k", "year", "days"):
            for row in range(FIRST_ROW, calc.max_row + 1):
                calc[f"{letter}{row}"].number_format = "0"
