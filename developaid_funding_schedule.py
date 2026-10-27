"""График финансирования (ДДС) и его исполнение по реестру платежей РСС.

Владелец (26.10.2026): «Можем сделать проверку и отдельный свод, что реально
исполнилось по этому плану ДДС». Решения владельца:

* план — лист «ГФ_до конца стр» файла «График финансирования <проект> <дата>»;
  новая версия приходит раз в месяц, хранятся все, месяц сверяется с той
  версией, что на него действовала;
* факт — реестр платежей РСС, уже загруженный в Монитор (`read_payments`):
  оплаты с кодом РСС и датой; сшивка по коду РСС без хвостовой точки;
* вид — раздел в сводке руководству.

Что важно в листе (все места взяты из настоящего файла):

* месяцы — по заголовку (год в объединённой ячейке над месяцами), не по
  позиции; тройки «Лимит после … / Банк … / Резерв …» справа — не план;
* итог — объявленная строка ИТОГО, строки после неё (резервы, «Всего:») —
  не статьи;
* пустая ячейка месяца — «плана нет», а не ноль; текст в числовой ячейке
  («гар. 5 %») — «нет числа» с примечанием;
* один код может стоять двумя строками — сумма по коду берёт обе и говорит об
  этом; подстроки (2.2.1.7.1 и т. п.) — разбивка графика: факт по родителю
  сравнивается с родителем и его подстроками вместе, а подстрока со своим
  кодом в реестре — отдельно. Иерархия — по порядку строк и глубине кода, как
  в `developaid_actuals`, а не по префиксу.

Числа сверки считает одна функция `execution`: её читают и сводка, и Монитор.
"""

from __future__ import annotations

import datetime
import io
import json
import re
from pathlib import Path
from typing import Any

import developaid_actuals as actuals
import developaid_monitor as monitor

SHEET = "ГФ_до конца стр"
KIND = "funding_schedule"
HEADER_ROWS = (2, 3, 4)
FIRST_ROW = 5
# Расхождение «Выплачено» с реестром меньше этого — округление, а не разница.
PAID_TOLERANCE = 1_000.0
_MONTH_WORDS = ("янв", "фев", "мар", "апр", "май", "июн",
                "июл", "авг", "сен", "окт", "ноя", "дек")
MONTH_NAMES = ("январь", "февраль", "март", "апрель", "май", "июнь", "июль",
               "август", "сентябрь", "октябрь", "ноябрь", "декабрь")
# Колонки — по словам заголовка. Ключ: что ищем, затем чего в заголовке быть
# не должно («нач./макс. сумма договора» — не сумма договора).
_COLUMNS = {
    "code": (("код",), ()),
    "work": (("работ",), ("логик",)),
    "status": (("статус",), ()),
    "contract": (("сумма договора",), ("нач", "макс")),
    "paid": (("выплачено",), ()),
    "remain": (("остаток к выплате",), ()),
}
_REQUIRED = ("code", "contract", "paid")
_LABELS = {"contract": "Сумма договора", "paid": "Выплачено",
           "remain": "Остаток к выплате"}


class ScheduleUnreadable(ValueError):
    """Лист не тот или в нём нет того, без чего сверка не имеет смысла."""


def _norm(value: Any) -> str:
    return re.sub(r"\s+", " ", str(value or "").replace("ё", "е")).strip().lower()


def _month_of(value: Any, year: int | None) -> str:
    """«сентябрь» при годе 2026 → «2026-09»; дата в ячейке — её месяц."""
    if isinstance(value, (datetime.date, datetime.datetime)):
        return f"{value.year:04d}-{value.month:02d}"
    text = _norm(value)
    match = re.fullmatch(r"([а-я]+)\.?(?:\s+(\d{4})(?:\s*г(?:од|\.)?)?)?", text)
    if not match or match.group(1)[:3] not in _MONTH_WORDS:
        return ""
    found = int(match.group(2)) if match.group(2) else year
    if not found:
        return ""
    return f"{found:04d}-{_MONTH_WORDS.index(match.group(1)[:3]) + 1:02d}"


def _number(value: Any, scale: float) -> tuple[float | None, str]:
    """Число ячейки и, если вместо числа текст, — сам текст."""
    if value is None or (isinstance(value, str) and not value.strip()):
        return None, ""
    if isinstance(value, bool):
        return None, str(value)
    if isinstance(value, (int, float)):
        return float(value) * scale, ""
    text = str(value).replace("\xa0", "").replace(" ", "").replace(",", ".")
    try:
        return float(text) * scale, ""
    except ValueError:
        return None, str(value).strip()


def _scale(texts: list[str], contract_total: float | None) -> tuple[float, str]:
    """Единицы сумм: из шапки, а если там их нет — по величине итога.

    Второе — допущение, и оно называется: на проекте в сотни миллионов рублей
    итог договоров меньше ста тысяч бывает только в миллионах.
    """
    joined = " ".join(texts)
    if re.search(r"млн", joined):
        return 1e6, "в шапке: млн ₽"
    if re.search(r"тыс", joined):
        return 1e3, "в шапке: тыс. ₽"
    total = abs(contract_total or 0.0)
    if total and total < 1e5:
        return 1e6, "в шапке единиц нет; итог договоров меньше 100 тысяч — считаю в млн ₽"
    if total and total < 1e8:
        return 1e3, "в шапке единиц нет; итог договоров меньше 100 млн — считаю в тыс. ₽"
    return 1.0, "в шапке единиц нет; считаю в рублях"


def read_schedule(data: bytes | Path) -> dict[str, Any]:
    """Строки графика, месяцы и объявленный итог."""
    from openpyxl import load_workbook

    source = io.BytesIO(data) if isinstance(data, (bytes, bytearray)) else data
    try:
        book = load_workbook(source, data_only=True)
    except Exception as exc:  # noqa: BLE001 — причина уходит загружающему
        raise ScheduleUnreadable(f"файл не открылся как книга Excel: {exc}") from exc
    if SHEET not in book.sheetnames:
        raise ScheduleUnreadable(
            f"в книге нет листа «{SHEET}»; листы книги: {', '.join(book.sheetnames)}")
    sheet = book[SHEET]
    # Объединённая ячейка хранит значение только в левой верхней — год над
    # пятью месяцами стоит один раз.
    merged: dict[tuple[int, int], Any] = {}
    for area in sheet.merged_cells.ranges:
        top = sheet.cell(area.min_row, area.min_col).value
        for row in range(area.min_row, area.max_row + 1):
            for col in range(area.min_col, area.max_col + 1):
                merged[(row, col)] = top

    def value(row: int, col: int) -> Any:
        return merged.get((row, col), sheet.cell(row, col).value)

    width = sheet.max_column
    columns: dict[str, int] = {}
    months: dict[int, str] = {}
    head_texts: list[str] = []
    for col in range(1, width + 1):
        cells = [value(row, col) for row in HEADER_ROWS]
        texts = [_norm(cell) for cell in cells]
        head_texts += [text for text in texts if text]
        year = next((int(m.group(1)) for m in (re.search(r"(\d{4})\s*г", t) for t in texts) if m),
                    None)
        month = next((m for m in (_month_of(cell, year) for cell in cells) if m), "")
        if month and (year or any(isinstance(c, (datetime.date, datetime.datetime))
                                  for c in cells) or re.search(r"\d{4}", " ".join(texts))):
            months[col] = month
            continue
        joined = " ".join(texts)
        for key, (wanted, banned) in _COLUMNS.items():
            if key in columns:
                continue
            if all(word in joined for word in wanted) and not any(b in joined for b in banned):
                columns[key] = col
                break
    absent = [key for key in _REQUIRED if key not in columns]
    if absent:
        names = {"code": "Код", "contract": "Сумма договора", "paid": "Выплачено"}
        raise ScheduleUnreadable(
            "в шапке листа (строки 2–4) не нашлись колонки: "
            + ", ".join(f"«{names[key]}»" for key in absent))
    if not months:
        raise ScheduleUnreadable("в шапке листа нет месяцев плана (месяц под годом «2026 год»)")
    labels = {key: (" ".join(_norm(value(r, col)) for r in HEADER_ROWS if value(r, col)) or key)
              for key, col in columns.items()}

    total_row = None
    raw: list[tuple[int, list[Any]]] = []
    for row in range(FIRST_ROW, sheet.max_row + 1):
        cells = [value(row, col) for col in range(1, width + 1)]
        if any(_norm(cell).startswith("итого") for cell in cells[:4]):
            total_row = (row, cells)
            break
        raw.append((row, cells))
    if total_row is None:
        raise ScheduleUnreadable("в листе нет строки «ИТОГО» — итог графика взять неоткуда")

    def cell(cells: list[Any], key: str) -> Any:
        col = columns.get(key)
        return cells[col - 1] if col and col <= len(cells) else None

    contract_raw, _ = _number(cell(total_row[1], "contract"), 1.0)
    # Единицы пишут и в заглавии над шапкой («…, млн руб.»).
    title = [_norm(value(1, col)) for col in range(1, width + 1) if value(1, col)]
    scale, unit = _scale(title + head_texts, contract_raw)

    rows: list[dict[str, Any]] = []
    uncoded: list[dict[str, Any]] = []
    for number, cells in raw:
        raw_code = cell(cells, "code")
        code = actuals._code(raw_code)
        work = str(cell(cells, "work") or "").strip()
        if not code:
            if work:
                uncoded.append({"row": number, "work": work})
            continue
        item: dict[str, Any] = {"row": number, "raw_code": str(raw_code).strip(), "code": code,
                                "depth": code.count(".") + 1, "work": work,
                                "status": str(cell(cells, "status") or "").strip(),
                                "plan": {}, "notes": []}
        for key in ("contract", "paid", "remain"):
            amount, text = _number(cell(cells, key), scale)
            item[key] = amount
            if text:
                item["notes"].append(f"в колонке «{_LABELS[key]}» вместо числа: «{text}»")
        for col, month in months.items():
            amount, text = _number(cells[col - 1] if col <= len(cells) else None, scale)
            item["plan"][month] = amount
            if text:
                item["notes"].append(f"в плане на {_month_words(month)} вместо числа: «{text}»")
        rows.append(item)
    _link(rows)
    total = {key: _number(cell(total_row[1], key), scale)[0] for key in ("contract", "paid", "remain")}
    total["plan"] = {month: _number(total_row[1][col - 1], scale)[0] for col, month in months.items()}
    return {"sheet": SHEET, "rows": rows, "uncoded": uncoded,
            "months": sorted(set(months.values())), "total": total, "total_row": total_row[0],
            "unit": unit, "columns": labels}


def _link(rows: list[dict[str, Any]]) -> None:
    """Родитель — ближайшая строка выше с меньшей глубиной кода (как в РСС)."""
    stack: list[dict[str, Any]] = []
    for row in rows:
        while stack and stack[-1]["depth"] >= row["depth"]:
            stack.pop()
        row["parent"] = stack[-1]["code"] if stack else ""
        stack.append(row)


def _month_words(month: str) -> str:
    year, number = month.split("-")
    return f"{MONTH_NAMES[int(number) - 1]} {year}"


# --------------------------------------------------------------------------
# Версии графика в Мониторе
# --------------------------------------------------------------------------

def store(project: str, data: bytes, taken_at: Any, filename: str = "") -> dict[str, Any]:
    """Положить версию графика на её дату. Та же дата — прежняя уходит в архив."""
    day = monitor._iso(taken_at)
    if not re.fullmatch(r"\d{4}-\d{2}-\d{2}", day):
        raise ValueError("дата графика финансирования нужна в виде ГГГГ-ММ-ДД")
    parsed = read_schedule(data)
    folder = monitor._project_dir(project) / KIND
    folder.mkdir(parents=True, exist_ok=True)
    path = monitor._inside(folder, f"{day}.xlsx")
    replaced = path.exists()
    if replaced:
        archive = folder / ".replaced"
        archive.mkdir(exist_ok=True)
        stamp = datetime.datetime.now(datetime.timezone.utc).strftime("%Y%m%dT%H%M%S")
        path.rename(monitor._inside(archive, f"{day}-{stamp}.xlsx"))
    path.write_bytes(data)
    monitor._inside(folder, f"{day}.json").write_text(json.dumps({
        "taken_at": day, "filename": filename,
        "loaded_at": datetime.datetime.now(datetime.timezone.utc).isoformat(),
    }, ensure_ascii=False), encoding="utf-8")
    return {"taken_at": day, "filename": filename, "replaced": replaced,
            "rows": len(parsed["rows"]), "months": parsed["months"],
            "contract": parsed["total"]["contract"], "paid": parsed["total"]["paid"],
            "unit": parsed["unit"]}


# Книги, которые Монитор уже хранит, — где лист графика может лежать и без
# отдельной загрузки (владелец, 27.10.2026: «Откуда я возьму эти файлы и
# данные? Они все есть или в НКЛ, или в РСС, или в модели»). Лист ищется по
# имени в оглавлении книги, а не по имени файла.
BOOKS = (
    ("baseline/finance.xlsx", "финмодель (финансовая база Монитора)"),
    ("finance_history/*.xlsx", "прежние версии финмодели"),
    ("estimate/*.xlsx", "отчёты РСС"),
    ("budget/*.xlsx", "файлы БДДС «Бюджет месяца»"),
    ("sales/*.xlsx", "отчёты о продажах"),
    ("credit/*.xlsx", "файлы кредита"),
)
_DAY = re.compile(r"\d{4}-\d{2}-\d{2}")


def _sheet_names(path: Path) -> list[str]:
    """Листы книги по её оглавлению (xl/workbook.xml), не открывая листы."""
    import html
    import zipfile

    try:
        with zipfile.ZipFile(path) as book:
            index = book.read("xl/workbook.xml").decode("utf-8", "replace")
    except (OSError, KeyError, zipfile.BadZipFile):
        return []
    return [html.unescape(name) for name in re.findall(r'<sheet\b[^>]*\bname="([^"]*)"', index)]


def _book_day(path: Path) -> str:
    """Дата версии: из имени файла, у файлов без даты — день загрузки."""
    found = _DAY.match(path.stem)
    if found:
        return found.group(0)
    return datetime.date.fromtimestamp(path.stat().st_mtime).isoformat()


def sources(project: str, upto: str = "") -> dict[str, Any]:
    """Где лежит график: дата версии → файл, и что просмотрено без находки.

    Отдельно загруженный график на ту же дату главнее книги, где лист нашёлся
    попутно: его загрузили именно как график.
    """
    root = monitor._project_dir(project)
    found: dict[str, dict[str, Any]] = {}
    looked: list[dict[str, Any]] = []
    for pattern, where in BOOKS:
        paths = sorted(root.glob(pattern))
        looked.append({"where": where, "files": len(paths)})
        for path in paths:
            day = _book_day(path)
            if (upto and day > upto) or SHEET not in _sheet_names(path):
                continue
            found[day] = {"path": path, "where": where, "file": path.name}
    folder = root / KIND
    for path in sorted(folder.glob("*.xlsx")) if folder.exists() else []:
        if not upto or path.stem <= upto:
            found[path.stem] = {"path": path, "where": "загруженный график финансирования",
                                "file": path.name}
    return {"versions": dict(sorted(found.items())), "looked": looked}


def versions(project: str, upto: str = "") -> list[str]:
    return list(sources(project, upto)["versions"])


def _missing_words(looked: list[dict[str, Any]]) -> str:
    seen = [f"{item['where']} ({item['files']})" for item in looked if item["files"]]
    empty = [item["where"] for item in looked if not item["files"]]
    text = ("Графика финансирования нет ни в одной книге, которую хранит Монитор, поэтому "
            "сравнить план оплат с оплаченным не с чем.")
    if seen:
        text += " Искали в книгах: " + ", ".join(seen) + "."
    if empty:
        text += " Не загружены: " + ", ".join(empty) + "."
    return (text + " Если график лежит отдельным файлом, его можно загрузить в Мониторе: блок "
            "«Бюджет месяца», кнопка «Загрузить график».")


def version_for(month: str, days: list[str]) -> str:
    """Какая версия действовала на месяц: последняя, выпущенная до его конца.

    Месяцы раньше первого графика (август при графике от 23.09) сверяются с
    первым — другой версии на них нет, и это говорится в подписи.
    """
    year, number = (int(part) for part in month.split("-"))
    last = (datetime.date(year + number // 12, number % 12 + 1, 1)
            - datetime.timedelta(days=1)).isoformat()
    earlier = [day for day in days if day <= last]
    return earlier[-1] if earlier else days[0]


def _read_version(project: str, day: str, found: dict[str, Any] | None = None) -> dict[str, Any]:
    place = (found or sources(project)["versions"]).get(day)
    if place is None:
        raise ScheduleUnreadable(f"версии графика от {day} нет")
    return read_schedule(place["path"])


# --------------------------------------------------------------------------
# Сверка с реестром платежей
# --------------------------------------------------------------------------

def _service(code: str) -> bool:
    return code == "2.6" or code.startswith("2.6.")


def _lines(schedule: dict[str, Any], own: set[str]) -> tuple[list[dict[str, Any]], dict[str, str]]:
    """Строки сверки: статья графика с подстроками; подстрока со своим кодом в
    реестре — отдельной строкой. Возвращает строки и карту «код графика → строка»."""
    rows = schedule["rows"]
    parent = {}
    for row in rows:
        parent.setdefault(row["code"], row["parent"])

    def home(code: str) -> str:
        cursor, seen = code, set()
        while cursor and cursor not in seen:
            seen.add(cursor)
            up = parent.get(cursor, "")
            if cursor in own or not up:
                return cursor
            cursor = up
        return code

    lines: dict[str, dict[str, Any]] = {}
    target: dict[str, str] = {}
    for row in rows:
        key = home(row["code"])
        target[row["code"]] = key
        line = lines.setdefault(key, {"code": key, "work": "", "rows": [], "notes": []})
        if row["code"] == key and not line["work"]:
            line["work"] = row["work"]
        line["rows"].append(row)
    for line in lines.values():
        line["work"] = line["work"] or line["rows"][0]["work"]
        same = [row for row in line["rows"] if row["code"] == line["code"]]
        if len(same) > 1:
            line["work"] = " / ".join(row["work"] for row in same)
            line["notes"].append(
                f"в графике {len(same)} строки с этим кодом (строки "
                + ", ".join(str(row["row"]) for row in same) + ") — в сумме все")
        subs = [row for row in line["rows"] if row["code"] != line["code"]]
        if subs:
            line["notes"].append(("вместе с подстрокой графика " if len(subs) == 1
                                  else "вместе с подстроками графика ")
                                 + ", ".join(row["code"] for row in subs))
        for row in line["rows"]:
            line["notes"] += row["notes"]
        if _service(line["code"]):
            import developaid_monthly_budget as monthly_budget
            note = next(item["note"] for item in monthly_budget.LINES if item["key"] == "service")
            line["notes"].append(note[:1].upper() + note[1:])
    return list(lines.values()), target


def _attribution(schedule: dict[str, Any], estimate: dict[str, Any]) -> dict[str, str]:
    """Код платежа → код графика: сам код, иначе ближайший предок по дереву РСС."""
    codes = {row["code"] for row in schedule["rows"]}
    parent = {row["code"]: row["parent"] for row in estimate.get("rows") or []}

    def find(code: str) -> str:
        cursor, seen = code, set()
        while cursor and cursor not in seen:
            seen.add(cursor)
            if cursor in codes:
                return cursor
            cursor = parent.get(cursor, "")
        return ""
    return {code: find(code) for code in set(parent) | codes}


def _sum(values: list[float | None]) -> float | None:
    known = [value for value in values if value is not None]
    return sum(known) if known else None


def _compare(plan: float | None, fact: float) -> dict[str, Any]:
    return {"plan": plan, "fact": fact,
            "gap": (fact - plan) if plan is not None else None,
            "share": (fact / plan) if plan else None}


def execution(project: str, end: datetime.date) -> dict[str, Any]:
    """Исполнение графика по месяцам на дату ``end`` — один ответ для всех."""
    where = sources(project, end.isoformat())
    found = where["versions"]
    days = list(found)
    if not days:
        return {"state": "missing", "reason": _missing_words(where["looked"])}
    rss = monitor._latest(project, "estimate", ".xlsx", end.isoformat())
    if rss is None:
        return {"state": "missing",
                "reason": f"Не загружен отчёт РСС с реестром платежей на {end:%d.%m.%Y}: "
                          "оплаты взять неоткуда."}
    payments = actuals.read_payments(rss)
    estimate = actuals.read_estimate(rss)
    last = payments.get("last")
    fact_at = min(end, last) if last else end
    current = f"{end.year:04d}-{end.month:02d}"
    cache: dict[str, dict[str, Any]] = {}

    def schedule(day: str) -> dict[str, Any]:
        if day not in cache:
            cache[day] = _read_version(project, day, found)
        return cache[day]

    months_seen = sorted({m for day in days for m in schedule(day)["months"] if m <= current})
    months = []
    for month in months_seen:
        day = version_for(month, days)
        plan = schedule(day)
        attribution = _attribution(plan, estimate)
        dated = [row for row in payments["rows"] if row.get("date") and row["date"] <= end]
        # Подстрока «своя», если реестр хоть раз платил по её коду.
        subs = {row["code"] for row in plan["rows"] if row["parent"]}
        own = {attribution.get(row.get("estimate_code") or "", "") for row in dated} & subs
        lines, target = _lines(plan, own)
        facts: dict[str, float] = {}
        outside = 0.0
        for row in dated:
            if row["date"].strftime("%Y-%m") != month:
                continue
            code = attribution.get(row.get("estimate_code") or "", "")
            if code:
                key = target[code]
                facts[key] = facts.get(key, 0.0) + float(row["amount"] or 0.0)
            else:
                outside += float(row["amount"] or 0.0)
        out_lines = []
        for line in lines:
            planned = _sum([row["plan"].get(month) for row in line["rows"]])
            fact = facts.get(line["code"], 0.0)
            if planned is None and not fact:
                continue
            out_lines.append(dict(_compare(planned, fact), code=line["code"], work=line["work"],
                                  notes=list(line["notes"]), service=_service(line["code"])))
        declared = plan["total"]["plan"].get(month)
        lines_plan = _sum([line["plan"] for line in out_lines])
        months.append(dict(
            _compare(declared, sum(facts.values())),
            month=month, version=day, closed=month < current,
            fact_at=fact_at.isoformat() if month == current else "",
            no_plan=declared is None and lines_plan is None,
            lines_plan=lines_plan, outside=outside, lines=out_lines,
            earlier_than_version=month < day[:7]))
    latest = schedule(days[-1])
    return {"state": "ok", "versions": days, "rss": rss.stem,
            "sources": {day: {"where": place["where"], "file": place["file"]}
                        for day, place in found.items()}, "fact_at": fact_at.isoformat(),
            "unit": latest["unit"], "months": months,
            "paid_check": paid_check(latest, days[-1], payments, estimate),
            "total": latest["total"], "uncoded": latest["uncoded"]}


def paid_check(plan: dict[str, Any], day: str, payments: dict[str, Any],
               estimate: dict[str, Any]) -> dict[str, Any]:
    """Колонка «Выплачено» графика против реестра по тем же статьям на дату графика."""
    cut = datetime.date.fromisoformat(day)
    attribution = _attribution(plan, estimate)
    dated = [row for row in payments["rows"] if row.get("date") and row["date"] <= cut]
    subs = {row["code"] for row in plan["rows"] if row["parent"]}
    own = {attribution.get(row.get("estimate_code") or "", "") for row in dated} & subs
    lines, target = _lines(plan, own)
    paid: dict[str, float] = {}
    for row in dated:
        code = attribution.get(row.get("estimate_code") or "", "")
        if code:
            paid[target[code]] = paid.get(target[code], 0.0) + float(row["amount"] or 0.0)
    out = []
    for line in lines:
        graph = _sum([row["paid"] for row in line["rows"]])
        register = paid.get(line["code"], 0.0)
        gap = register - (graph or 0.0)
        notes = [note for row in line["rows"] for note in row["notes"] if "Выплачено" in note]
        # Текст вместо числа в «Выплачено» — тоже расхождение: числа нет.
        if abs(gap) >= PAID_TOLERANCE or (graph is None and (register or notes)):
            out.append({"code": line["code"], "work": line["work"], "graph": graph,
                        "register": register, "gap": gap, "notes": notes})
    register_total = sum(paid.values())
    graph_total = plan["total"]["paid"]
    return {"at": day, "graph": graph_total, "register": register_total,
            "gap": register_total - (graph_total or 0.0), "lines": out}
